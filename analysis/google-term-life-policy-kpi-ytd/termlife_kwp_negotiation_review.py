"""
termlife_kwp_negotiation_review.py  (v1)
=========================================
Pre/post review of the TWO keyword lists the CEO used in the payout
negotiation (from the coworker's RPL_Analysis_Life.xlsx, Analysis
sheet, rows 38+): Competitor Terms (A38:D100) and Generic Terms
(F38:I53). Replicates each list with its original Keyword Planner
columns (avg monthly searches, top-of-page CPC, estimated CPL) and
adds actual pre/post performance from paid_search.db: Leads, CPL,
RPL (= conv value / leads, the SQ payment per lead).

Grains (KWP terms are not match-type specific):
  KEYWORD grain  - all keyword instances whose text matches the KWP
                   phrase, ANY match type / ad group, campaigns rolled
                   up (complete aggregation; instances & match types
                   listed per row).
  QUERY grain    - search queries whose text equals the KWP phrase
                   (sqp). The only grain that can answer the GENERIC
                   list, which is mostly not bid as keywords; also
                   catches close-variant/broad serving on competitor
                   phrases. sqp hides low-volume terms -> lower bound.

Windows: PRE Mar 1 - Jun 1 | POST Jun 2 - Sep 30 (full September,
per 10/20 ruling) | Oct MTD separate. Windows differ in length ->
leads/day is the comparable volume column.

Inputs in ANALYSIS_DIR\\input\\: the coworker workbook
(*RPL_Analysis*.xlsx). Output: KWP_Negotiation_Keyword_Review.xlsx.

Caveats (written into Read_Me): KWP columns are Google Planner
estimates at negotiation time, not our data - kept for side-by-side
only; estimated CPL there = top-of-page CPC / 0.186 CVR assumption;
RPL here = SQ payment per routed lead (quality-blind); query grain is
an sqp-visible lower bound; conv value = SQ payment.
"""
import glob
import os
import sqlite3
import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.drawing.text import RichTextProperties
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

DB_PATH = r"C:\Projects\google-ads-ops\db\paid_search.db"
ANALYSIS_DIR = (r"C:\Projects\google-ads-ops\analysis"
                r"\google-term-life-policy-kpi-ytd")
OUT_XLSX = os.path.join(ANALYSIS_DIR, "output",
                        "KWP_Negotiation_Keyword_Review.xlsx")
SRC_GLOB = "*RPL_Analysis*.xlsx"

ACCOUNT_ID = 29
PRE0, PRE1 = "2026-03-01", "2026-06-02"        # Mar 1 - Jun 1
POST0, POST1 = "2026-06-02", "2026-10-01"      # Jun 2 - Sep 30 (full)
OCT0, OCT1 = "2026-10-01", "2026-11-01"

NAVY, LIGHT = "1F3864", "D9E2F2"
SERIES_HEX = ["2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4", "008300"]

pd.set_option("display.width", 240)
pd.set_option("display.max_columns", 44)
pd.set_option("display.max_rows", 250)
pd.set_option("display.float_format", lambda x: f"{x:,.1f}")


def hr(t):
    print("\n" + "=" * 110)
    print(t)
    print("=" * 110)


def meta(ws, title, source, window, caveat, takeaway=None):
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=13, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A2"] = f"Source: {source}   |   Window: {window}"
    ws["A3"] = f"Caveats: {caveat}"
    for r in (2, 3):
        ws[f"A{r}"].font = Font(italic=True, size=9)
    if takeaway:
        ws["A4"] = f"TAKEAWAY: {takeaway}"
        ws["A4"].font = Font(bold=True, size=10)
        ws["A4"].fill = PatternFill("solid", fgColor=LIGHT)
    ws.freeze_panes = "A6"


def write_block(ws, row, subtitle, df):
    ws.cell(row=row, column=1, value=subtitle).font = Font(bold=True,
                                                           size=11)
    ws.cell(row=row, column=1).fill = PatternFill("solid", fgColor=LIGHT)
    row += 1
    for j, c in enumerate(df.columns, start=1):
        cell = ws.cell(row=row, column=j, value=str(c))
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
    row += 1
    for _, r in df.iterrows():
        for j, v in enumerate(r, start=1):
            if isinstance(v, (np.floating, float)):
                v = None if pd.isna(v) else round(float(v), 2)
            elif isinstance(v, (np.integer,)):
                v = int(v)
            ws.cell(row=row, column=j, value=v)
        row += 1
    for j, c in enumerate(df.columns, start=1):
        width = max([len(str(c))] + [len(str(x)) for x in df[c].head(80)])
        ws.column_dimensions[get_column_letter(j)].width = min(width + 2, 44)
    return row + 2


def write_series_block(ws, start_col, df):
    for j, cname in enumerate(df.columns):
        ws.cell(row=1, column=start_col + j, value=str(cname))
        for i, v in enumerate(df[cname].tolist(), start=2):
            if isinstance(v, (np.floating, float)):
                v = None if pd.isna(v) else round(float(v), 2)
            elif isinstance(v, (np.integer,)):
                v = int(v)
            elif v is not None and not isinstance(v, (int, float)):
                v = str(v)
            ws.cell(row=i, column=start_col + j, value=v)
    return len(df)


def make_line_chart(ws_data, start_col, n_series, n_rows, title, y_title,
                    skip=1):
    ch = LineChart()
    ch.title = title
    data = Reference(ws_data, min_col=start_col + 1,
                     max_col=start_col + n_series,
                     min_row=1, max_row=n_rows + 1)
    cats = Reference(ws_data, min_col=start_col, min_row=2,
                     max_row=n_rows + 1)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    if ch.title is not None:
        ch.title.overlay = False
    ch.legend.position = "r"
    ch.legend.overlay = False
    ch.y_axis.title = y_title
    try:
        ch.y_axis.title.tx.rich.bodyPr = RichTextProperties(
            rot=-5400000, vert="horz")
    except AttributeError:
        pass
    xn = ws_data.cell(row=1, column=start_col).value
    if xn:
        ch.x_axis.title = str(xn)
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.x_axis.tickLblSkip = skip
    ch.x_axis.tickMarkSkip = skip
    ch.layout = Layout(manualLayout=ManualLayout(
        layoutTarget="inner", xMode="edge", yMode="edge",
        x=0.10, y=0.14, w=0.72, h=0.66))
    ch.height = 9.5
    ch.width = 30
    for i, s in enumerate(ch.series):
        s.smooth = False
        s.graphicalProperties.line.width = 20000
        s.graphicalProperties.line.solidFill = SERIES_HEX[i % len(SERIES_HEX)]
    return ch


def norm(s):
    return " ".join(str(s).strip().lower()
                    .removeprefix("[").removesuffix("]").split())


def load_lists():
    paths = glob.glob(os.path.join(ANALYSIS_DIR, "input", SRC_GLOB))
    if not paths:
        raise SystemExit(f"coworker workbook not found in input\\ "
                         f"({SRC_GLOB})")
    path = sorted(paths)[-1]
    ws = load_workbook(path, data_only=True)["Analysis"]

    def read_list(kw_col, first_data_col, r0=38, r1=120):
        rows = []
        for r in range(r0, r1):
            kw = ws.cell(row=r, column=kw_col).value
            if kw is None or str(kw).strip() == "":
                continue
            s = str(kw)
            if any(t in s for t in ("Blinded", "Show Ethos", "Remove CPC")):
                continue  # presentation notes, not keywords
            rows.append({
                "keyword": norm(s),
                "KWP_avg_monthly_searches":
                    ws.cell(row=r, column=first_data_col).value,
                "KWP_top_of_page_CPC":
                    ws.cell(row=r, column=first_data_col + 1).value,
                "KWP_estimated_CPL":
                    ws.cell(row=r, column=first_data_col + 2).value})
        return pd.DataFrame(rows).drop_duplicates(subset="keyword")

    comp = read_list(1, 2)     # A: keyword; B,C,D
    gen = read_list(6, 7)      # F: keyword; G,H,I
    return comp, gen, os.path.basename(path)


def win_cols(df, dcol="report_date"):
    rows = {}
    for d0, d1, lab in [(PRE0, PRE1, "pre"), (POST0, POST1, "post"),
                        (OCT0, OCT1, "oct")]:
        w = df[(df[dcol] >= d0) & (df[dcol] < d1)]
        days = (pd.Timestamp(d1) - pd.Timestamp(d0)).days
        leads, cost, val = (w["leads"].sum(), w["spend"].sum(),
                            w["value"].sum())
        rows[f"{lab}_leads"] = leads
        rows[f"{lab}_leads_per_day"] = leads / days
        rows[f"{lab}_spend"] = cost
        rows[f"{lab}_CPL"] = cost / leads if leads else np.nan
        rows[f"{lab}_RPL"] = val / leads if leads else np.nan
    for mcol in ["leads_per_day", "CPL", "RPL"]:
        a, b = rows.get(f"pre_{mcol}"), rows.get(f"post_{mcol}")
        rows[f"chg_{mcol}_%"] = (100 * (b - a) / a
                                 if a and not pd.isna(a) else np.nan)
    return rows


def main():
    con = sqlite3.connect(f"file:{DB_PATH}?mode=ro", uri=True)
    comp, gen, src_name = load_lists()
    hr(f"LISTS from {src_name}: competitor={len(comp)} keywords, "
       f"generic={len(gen)} keywords")

    ki = pd.read_sql_query(
        "SELECT id, keyword_text, match_type, campaign_id, ad_group_name "
        "FROM keyword_instances WHERE account_id=?", con,
        params=(ACCOUNT_ID,))
    ki["kw_norm"] = ki["keyword_text"].map(norm)

    all_terms = set(comp["keyword"]) | set(gen["keyword"])
    ki_m = ki[ki["kw_norm"].isin(all_terms)]
    hr(f"KEYWORD-GRAIN MATCH: {ki_m['kw_norm'].nunique()} of "
       f"{len(all_terms)} KWP terms exist as keywords "
       f"({len(ki_m)} instances, any match type)")

    # keyword-grain daily
    k = pd.DataFrame(columns=["kw_norm", "report_date", "impressions",
                              "spend", "leads", "value"])
    if len(ki_m):
        ids = ",".join(str(i) for i in ki_m["id"])
        k = pd.read_sql_query(f"""
            SELECT p.instance_id, p.report_date,
                   SUM(p.impressions) AS impressions,
                   SUM(p.cost_micros)/1e6 AS spend,
                   SUM(p.primary_conv_count) AS leads,
                   SUM(p.primary_conv_value) AS value
            FROM keyword_performance_daily p
            WHERE p.instance_id IN ({ids}) AND p.device='ALL'
            GROUP BY 1,2""", con)
        k = k.merge(ki_m[["id", "kw_norm"]], left_on="instance_id",
                    right_on="id")

    # query-grain daily (exact text equality against sqp)
    q = pd.read_sql_query(f"""
        SELECT LOWER(TRIM(sq.query_text)) AS qtext, p.report_date,
               SUM(p.impressions) AS impressions,
               SUM(p.cost_micros)/1e6 AS spend,
               SUM(p.primary_conv_count) AS leads,
               SUM(p.primary_conv_value) AS value
        FROM search_query_performance p
        JOIN keyword_instances ki ON ki.id=p.instance_id
                                  AND ki.account_id={ACCOUNT_ID}
        JOIN search_queries sq ON sq.id=p.query_id
        WHERE p.platform='google' AND p.device='ALL'
        GROUP BY 1,2""", con)
    q["qtext"] = q["qtext"].map(norm)
    q = q[q["qtext"].isin(all_terms)]
    hr(f"QUERY-GRAIN MATCH: {q['qtext'].nunique()} of {len(all_terms)} "
       f"KWP terms appear as visible search queries")

    def build(list_df, list_name):
        kw_rows, q_rows = [], []
        for _, r in list_df.iterrows():
            t = r["keyword"]
            base = {"keyword": t,
                    "KWP_avg_monthly_searches":
                        r["KWP_avg_monthly_searches"],
                    "KWP_top_of_page_CPC": r["KWP_top_of_page_CPC"],
                    "KWP_estimated_CPL": r["KWP_estimated_CPL"]}
            sub = ki_m[ki_m["kw_norm"] == t]
            rec = dict(base)
            rec["bid_as_keyword"] = "YES" if len(sub) else "NO"
            rec["n_instances"] = len(sub)
            rec["match_types"] = ",".join(
                sorted(sub["match_type"].str.upper().unique()))
            rec.update(win_cols(k[k["kw_norm"] == t]))
            kw_rows.append(rec)
            rec2 = dict(base)
            rec2.update(win_cols(q[q["qtext"] == t]))
            q_rows.append(rec2)
        kw_tbl = pd.DataFrame(kw_rows).sort_values(
            "post_spend", ascending=False)
        q_tbl = pd.DataFrame(q_rows).sort_values(
            "post_spend", ascending=False)
        # list totals (both grains)
        tot = []
        for lab, df_, keycol in [("keyword grain", k, "kw_norm"),
                                 ("query grain", q, "qtext")]:
            sub = df_[df_[keycol].isin(set(list_df["keyword"]))]
            rec = {"list": f"{list_name} TOTAL ({lab})"}
            rec.update(win_cols(sub))
            tot.append(rec)
        return kw_tbl, q_tbl, pd.DataFrame(tot)

    comp_kw, comp_q, comp_tot = build(comp, "Competitor")
    gen_kw, gen_q, gen_tot = build(gen, "Generic")

    show = ["keyword", "KWP_estimated_CPL", "bid_as_keyword",
            "pre_leads", "post_leads", "pre_CPL", "post_CPL",
            "pre_RPL", "post_RPL", "chg_leads_per_day_%", "chg_CPL_%",
            "chg_RPL_%"]
    hr("COMPETITOR LIST — keyword grain (top 25 by post spend)")
    print(comp_kw[show].head(25).to_string(index=False))
    hr("GENERIC LIST — keyword grain")
    print(gen_kw[show].to_string(index=False))
    hr("LIST TOTALS")
    tshow = ["list", "pre_leads", "post_leads", "pre_leads_per_day",
             "post_leads_per_day", "pre_CPL", "post_CPL", "pre_RPL",
             "post_RPL"]
    print(pd.concat([comp_tot, gen_tot])[tshow].to_string(index=False))

    # monthly trail per list (query grain = the honest common grain)
    q["month"] = q["report_date"].str[:7]
    trail = pd.DataFrame({"month": sorted(q["month"].unique())})
    for nm_, lst in [("Competitor", comp), ("Generic", gen)]:
        sub = (q[q["qtext"].isin(set(lst["keyword"]))]
               .groupby("month")
               .agg(leads=("leads", "sum"), spend=("spend", "sum"),
                    value=("value", "sum")).reset_index())
        sub[f"{nm_} leads"] = sub["leads"]
        sub[f"{nm_} RPL"] = sub["value"] / sub["leads"].replace(0, np.nan)
        trail = trail.merge(sub[["month", f"{nm_} leads", f"{nm_} RPL"]],
                            on="month", how="left")
    trail = trail[trail["month"] >= "2026-01"]
    hr("MONTHLY TRAIL (query grain)")
    print(trail.to_string(index=False))

    # ======================================================== workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Read_Me"
    meta(ws, "KWP NEGOTIATION LISTS — PRE/POST LEADS, CPL, RPL",
         f"lists + KWP columns from {src_name} (Analysis rows 38+); "
         "actuals from paid_search.db (read-only)",
         f"PRE {PRE0}..Jun 1 | POST Jun 2..Sep 30 (full) | Oct MTD",
         "KWP columns are Google Keyword Planner estimates at "
         "negotiation time (their estimated CPL = top-of-page CPC / "
         "0.186 CVR assumption) - side-by-side reference only; KWP "
         "terms are not match-type specific -> keyword grain = ALL "
         "match types aggregated, query grain = exact search-term "
         "match (sqp-visible lower bound); RPL = SQ payment per "
         "routed lead (quality-blind); windows differ in length -> "
         "compare leads/day",
         "The negotiation case was: these terms are priced out at the "
         "old payout. This workbook shows what the June 2 raise "
         "actually bought on them: leads, CPL, and the RPL that pays "
         "for it.")
    gl = pd.DataFrame([
        ["Competitor_Terms", "The 63-term competitor list (A38:A100), "
         "KWP columns replicated + pre/post actuals at both grains."],
        ["Generic_Terms", "The 16-term generic list (F38:F53). Mostly "
         "NOT bid as keywords - the query grain is the meaningful "
         "view (close variants/broad serving on those exact "
         "searches)."],
        ["Keyword grain", "All keyword instances matching the phrase, "
         "any match type/ad group, campaigns rolled up; "
         "bid_as_keyword + n_instances + match_types show what is "
         "behind each row."],
        ["Query grain", "Search queries equal to the phrase. Catches "
         "serving without a matching keyword; hidden low-volume "
         "terms excluded by Google -> lower bound."],
        ["RPL", "conv value / leads = SQ payment per routed lead. The "
         "metric the payout raise moved."],
        ["Totals & trail", "List-level pre/post totals at both grains; "
         "monthly leads + RPL trail (query grain) + charts."]],
        columns=["item", "meaning"])
    write_block(ws, 6, "How to read this workbook", gl)

    for nm_, kw_tbl, q_tbl, tot in [
            ("Competitor_Terms", comp_kw, comp_q, comp_tot),
            ("Generic_Terms", gen_kw, gen_q, gen_tot)]:
        wsx = wb.create_sheet(nm_)
        meta(wsx, f"{nm_.replace('_', ' ').upper()} — KWP columns + "
             "pre/post Leads, CPL, RPL",
             "KWP list + paid_search.db",
             f"pre {PRE0}..Jun 1 | post Jun 2..Sep 30 | oct MTD",
             "keyword grain = all match types aggregated; query grain "
             "= exact search-term match, sqp-visible lower bound",
             "Sorted by post spend - top rows are where the raise's "
             "money landed on this list.")
        r = write_block(wsx, 6, "List totals (both grains)",
                        tot)
        r = write_block(wsx, r, "Keyword grain (all match types "
                        "aggregated)", kw_tbl)
        write_block(wsx, r, "Query grain (exact search-term match)",
                    q_tbl)

    wsc = wb.create_sheet("Charts")
    wsd = wb.create_sheet("Chart_Data")
    col = 1
    cb = trail[["month", "Competitor leads", "Generic leads"]]
    n = write_series_block(wsd, col, cb)
    wsc.add_chart(make_line_chart(
        wsd, col, 2, n, "Monthly leads on the negotiation lists "
        "(query grain; raise = Jun 2)", "leads"), "A2")
    col += len(cb.columns) + 1
    cb2 = trail[["month", "Competitor RPL", "Generic RPL"]]
    n = write_series_block(wsd, col, cb2)
    wsc.add_chart(make_line_chart(
        wsd, col, 2, n, "Monthly RPL on the negotiation lists "
        "(SQ payment per lead)", "RPL ($)"), "A24")
    col += len(cb2.columns) + 1
    meta(wsc, "CHARTS — the negotiation lists, monthly",
         "query grain", "Jan -> Oct MTD",
         "query grain lower bound; Oct partial",
         "Look for the June step in leads and the RPL level shift - "
         "that pairing is the negotiation's scoreboard.")
    wsd.sheet_state = "hidden"

    os.makedirs(os.path.dirname(OUT_XLSX), exist_ok=True)
    wb.save(OUT_XLSX)
    hr(f"WORKBOOK SAVED -> {OUT_XLSX}")
    con.close()


if __name__ == "__main__":
    main()
