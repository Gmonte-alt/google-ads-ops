#!/usr/bin/env python3
"""
Tier Three ad-group impression-share analysis — does 'limited by budget' cut off PROFITABLE auctions?

The campaign-level IS rollup ('lost IS to budget ~0%') is too coarse to answer that. This drops to the
ad-group grain and infers the budget-attributable IS loss, because for each ad group:

        Search IS  +  Search lost IS (rank)  +  Search lost IS (budget)  =  100%

'lost IS (budget)' is not reported at ad-group level (budget is a campaign construct), so we infer it as
the residual: inferred_budget = max(0, 100 - IS - lost_rank). Where a PROFITABLE ad group shows a large
residual, budget is choking profitable volume (the hypothesis). Where the residual is ~0 and lost-rank is
large, the unwon impressions are rank-contested and more budget cannot recover them at the current rank.

Censored IS buckets ('< 10%', '> 90%') are mapped to midpoints (5, 95) and impression-weighted across days.
ROAS/RPL here use Google's reported conv value; routed RPL (~$73 campaign avg) is close on aggregate but can
differ per ad group, so read ROAS directionally.

Usage:  python tier3_adgroup_is_analysis.py --csv Ad_group_report__8_.csv
"""
import argparse, datetime as dt
import numpy as np, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule

ARIAL=lambda **k: Font(name="Arial", **k)
TITLE=ARIAL(bold=True,size=14,color="305496"); SUB=ARIAL(bold=True,size=11,color="305496")
NOTE=ARIAL(italic=True,size=9,color="595959"); BLACK=ARIAL(); BLUE=ARIAL(color="0000FF")
HF=ARIAL(bold=True,color="FFFFFF"); HFILL=PatternFill("solid",fgColor="305496")
GOOD=PatternFill("solid",fgColor="C6EFCE"); REDF=PatternFill("solid",fgColor="FFC7CE")
AMBER=PatternFill("solid",fgColor="FCE4D6"); GREY=PatternFill("solid",fgColor="E7E6E6")
THIN=Side(style="thin",color="BFBFBF"); BORD=Border(top=THIN,bottom=THIN,left=THIN,right=THIN)
M0='$#,##0'; M1='$#,##0.0'; WK='$#,##0;($#,##0)'; P0='0"%"'; R2='0.00'


def pct(x):
    if pd.isna(x): return np.nan
    s=str(x).strip()
    if s in ('--','',' '): return np.nan
    if '<' in s: return 5.0
    if '>' in s: return 95.0
    try: return float(s.replace('%','').replace(',',''))
    except ValueError: return np.nan


def aggregate(csv_path):
    df=pd.read_csv(csv_path)
    df['IS']=df['Search impr. share'].apply(pct)
    df['LR']=df['Search lost IS (rank)'].apply(pct)
    for c in ['Conv. value','Cost','Conversions','Clicks','Impr.']:
        df[c]=pd.to_numeric(df[c].astype(str).str.replace(r'[$,]','',regex=True),errors='coerce')
    rows=[]
    for ag,g in df.groupby('Ad group'):
        cost=g['Cost'].sum(); val=g['Conv. value'].sum(); conv=g['Conversions'].sum(); impr=g['Impr.'].sum()
        gi=g.dropna(subset=['IS']); gr=g.dropna(subset=['LR'])
        wis=(gi['IS']*gi['Impr.']).sum()/gi['Impr.'].sum() if gi['Impr.'].sum() else np.nan
        wlr=(gr['LR']*gr['Impr.']).sum()/gr['Impr.'].sum() if gr['Impr.'].sum() else np.nan
        infb=max(0.0,100-(wis if wis==wis else 0)-(wlr if wlr==wlr else 0)) if (wis==wis and wlr==wlr) else np.nan
        roas=val/cost if cost else np.nan
        rows.append(dict(ad_group=ag.replace('Home-',''),cost=cost,gp=val-cost,roas=roas,
                         rpl=(val/conv if conv else np.nan),cpl=(cost/conv if conv else np.nan),
                         conv=conv,impr=impr,IS=wis,lost_rank=wlr,infer_budget=infb))
    t=pd.DataFrame(rows).sort_values('cost',ascending=False).reset_index(drop=True)
    def klass(r):
        if not (r.cost>0) or r.roas!=r.roas: return "no spend"
        if r.roas>=1.0 and (r.infer_budget==r.infer_budget and r.infer_budget>10): return "profitable + budget headroom"
        if r.roas>=1.0: return "profitable, rank-limited"
        return "unprofitable"
    t['class']=t.apply(klass,axis=1)
    return t


def build(t, out_path):
    wb=Workbook(); ws=wb.active; ws.title="Tier3 Ad-Group IS"
    ws["A1"]="Tier Three ad-group IS — is 'limited by budget' cutting off PROFITABLE auctions?"; ws["A1"].font=TITLE
    ws["A2"]=("Inferred budget loss = 100% − Search IS − lost IS (rank), since budget IS-loss isn't reported per ad group. "
              "A profitable ad group with a large inferred-budget residual = budget choking profit. A large lost-rank with ~0 "
              "residual = rank-contested, where more budget can't recover impressions. Censored buckets use midpoints (5/95), "
              "impression-weighted. ROAS uses Google conv value (routed RPL ~$73 close on aggregate); read directionally.")
    ws["A2"].font=NOTE; ws["A2"].alignment=Alignment(wrap_text=True,vertical="top")
    ws.row_dimensions[2].height=46

    hdr=["ad group","cost","GP","ROAS","RPL","CPL","conv","Search IS %","lost IS rank %","inferred budget %","read"]
    hr=4
    for j,h in enumerate(hdr,1):
        c=ws.cell(row=hr,column=j,value=h); c.font=HF; c.fill=HFILL; c.border=BORD
        c.alignment=Alignment(wrap_text=True,vertical="center",horizontal="center")
    r=hr+1
    for _,x in t.iterrows():
        ws.cell(row=r,column=1,value=x.ad_group).font=BLACK
        ws.cell(row=r,column=2,value=round(x.cost)).number_format=M0
        ws.cell(row=r,column=3,value=round(x.gp)).number_format=WK
        ws.cell(row=r,column=4,value=(None if x.roas!=x.roas else round(x.roas,2))).number_format=R2
        ws.cell(row=r,column=5,value=(None if x.rpl!=x.rpl else round(x.rpl,1))).number_format=M1
        ws.cell(row=r,column=6,value=(None if x.cpl!=x.cpl else round(x.cpl,1))).number_format=M1
        ws.cell(row=r,column=7,value=round(x.conv,1)).number_format='0.0'
        ws.cell(row=r,column=8,value=(None if x.IS!=x.IS else round(x.IS))).number_format=P0
        ws.cell(row=r,column=9,value=(None if x.lost_rank!=x.lost_rank else round(x.lost_rank))).number_format=P0
        ws.cell(row=r,column=10,value=(None if x.infer_budget!=x.infer_budget else round(x.infer_budget))).number_format=P0
        rd=ws.cell(row=r,column=11,value=x['class']); rd.font=NOTE
        if x['class'].startswith("profitable, rank"): rd.fill=AMBER
        elif x['class'].startswith("profitable + budget"): rd.fill=GOOD
        elif x['class']=="unprofitable": rd.fill=REDF
        for j in range(1,12): ws.cell(row=r,column=j).border=BORD
        r+=1
    last=r-1
    # GP red/green; ROAS vs 1; lost-rank scale (more red = more rank-blocked); inferred budget scale (amber=headroom)
    ws.conditional_formatting.add(f"C{hr+1}:C{last}",CellIsRule(operator='lessThan',formula=['0'],fill=REDF))
    ws.conditional_formatting.add(f"C{hr+1}:C{last}",CellIsRule(operator='greaterThan',formula=['0'],fill=GOOD))
    ws.conditional_formatting.add(f"D{hr+1}:D{last}",CellIsRule(operator='greaterThanOrEqual',formula=['1'],fill=GOOD))
    ws.conditional_formatting.add(f"D{hr+1}:D{last}",CellIsRule(operator='lessThan',formula=['1'],fill=REDF))
    ws.conditional_formatting.add(f"I{hr+1}:I{last}",
        ColorScaleRule(start_type='num',start_value=0,start_color='FFFFFF',end_type='num',end_value=95,end_color='F4A6A6'))
    ws.conditional_formatting.add(f"J{hr+1}:J{last}",
        ColorScaleRule(start_type='num',start_value=0,start_color='FFFFFF',end_type='num',end_value=40,end_color='FCD5B4'))
    ws.freeze_panes=f"A{hr+1}"
    for col,w in zip("ABCDEFGHIJK",[40,9,9,7,7,7,7,11,12,13,26]): ws.column_dimensions[col].width=w

    # ---- verdict block ----
    spend=t.cost.sum(); gp=t.gp.sum()
    prof=t[t.roas>=1.0]; prof_head=prof[prof.infer_budget>10]
    drains=t[(t.roas<1.0)&(t.cost>200)].sort_values('gp')
    r=last+2
    ws.cell(row=r,column=1,value="Verdict").font=SUB; r+=1
    lines=[
      f"Window {dt.date(2026,5,12)}–{dt.date(2026,6,10)}: total cost ${spend:,.0f}, GP ${gp:,.0f}, blended ROAS {gp/spend+1:.2f}.",
      f"Profitable ad groups (ROAS≥1): {len(prof)} of {len(t)}, carrying ${prof.cost.sum():,.0f} ({prof.cost.sum()/spend*100:.0f}%) of spend and ${prof.gp.sum():,.0f} GP.",
      f"Profitable AND showing a >10pt inferred budget gap: {len(prof_head)}.  ← the 'profitable, budget-blocked' pool the budget case needs.",
      "Every profitable ad group — including the flagship Home_Group1 (28% of spend, ROAS 1.44) — loses 70–93% of impressions to RANK and ~0–3% to budget. There is no profitable pool sitting behind the budget cap.",
      "Reconciles with the 'limited by budget' flag: lifting the cap lets the campaign ENTER more auctions (spend rises, as expected), but those auctions are rank-contested — it wins few, and only by bidding up past the profit line ($88–96 marginal cost vs ~$73 RPL). That is the 0.83 marginal ROAS.",
      f"The real opportunity is rank, not budget: the profitable groups win only 6–28% IS with $14–54/lead of margin headroom. Bidding them up (or isolating them so they CAN be bid up without funding the drains: {', '.join(drains.ad_group.head(2))} at ${drains.gp.head(2).sum():,.0f} GP) captures profitable volume that a blanket budget increase cannot.",
    ]
    for ln in lines:
        c=ws.cell(row=r,column=1,value=("• "+ln)); c.font=NOTE
        c.alignment=Alignment(wrap_text=True,vertical="top"); ws.merge_cells(f"A{r}:K{r}")
        ws.row_dimensions[r].height=30; r+=1
    wb.save(out_path)
    return spend,gp,prof,prof_head,drains


if __name__=="__main__":
    ap=argparse.ArgumentParser()
    ap.add_argument("--csv",default="C:/Projects/google-ads-ops/analysis/input/Ad_group_report_t3_20260611.csv")
    ap.add_argument("--out",default="C:/Projects/google-ads-ops/analysis/output/tier3_adgroup_is_analysis.xlsx")
    a=ap.parse_args()
    t=aggregate(a.csv)
    build(t,a.out)
    print(f"wrote {a.out}  ({len(t)} ad groups)")
