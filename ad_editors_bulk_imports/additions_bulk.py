#!/usr/bin/env python3
"""Generate the Google Ads Editor bulk sheet for carrier keyword additions, MIRRORED across all active
campaigns. Carry-over keywords for split carriers are DERIVED from the actual grab-bag contents
(pause_split_review.csv from inventory_assessment.py) so the new ad group re-adds exactly what the pause
sheet removes -- no silent coverage loss -- minus geo variants (handled by the geo-campaign + base term),
plus the new bare-brand terms.

Inputs:
  --campaigns-csv       inventory campaign_summary.csv (exact campaign names)
  --pause-review-csv    inventory pause_split_review.csv (real grab-bag keywords for the split carriers)
Emits: Ad Groups, Keywords (mirrored x campaigns), Dropped_Geo (report of intentional geo drops).
Smart Bidding -> no per-keyword bids. Ad-group naming follows 'Home-NonGeneric-<Carrier> Insurance-Core'.
"""
import argparse
import csv
import os
import re

# --- split carriers: new dedicated ad group + the NEW bare-brand terms to add (carry-over is derived) ---
SPLIT_CARRIERS = {
    "kin": {
        "ad_group": "Home-NonGeneric-Kin Insurance-Core",
        "bare_exact": ["kin insurance", "kin homeowners insurance", "kin insurance quote",
                       "kin insurance reviews"],
        "bare_broad": ["kin insurance"],
    },
    "costco": {
        "ad_group": "Home-NonGeneric-Costco Insurance-Core",
        "bare_exact": ["costco insurance", "costco house insurance", "costco home and auto insurance"],
        "bare_broad": ["costco insurance"],
    },
    "american family": {
        "ad_group": "Home-NonGeneric-American Family Insurance-Core",
        "bare_exact": ["american family insurance", "american family",
                       "american family home insurance"],
        "bare_broad": ["american family insurance"],
    },
}
# existing ad groups that only need bare anchors added (not split; carry-over N/A)
EXISTING_ADGROUP_ADDS = {
    "Home-NonGeneric-Hartford Insurance-Core": {
        "exact": ["the hartford", "hartford insurance", "the hartford insurance company"],
        "broad": ["hartford"],
    },
    "Home-NonGeneric-Travelers Insurance-Core": {
        "exact": ["travelers", "travelers insurance", "travelers insurance company"],
        "broad": ["travelers insurance"],
    },
}

_STATE = re.compile(r"\b(alabama|alaska|arizona|arkansas|california|colorado|connecticut|delaware|"
    r"florida|georgia|hawaii|idaho|illinois|indiana|iowa|kansas|kentucky|louisiana|maine|maryland|"
    r"massachusetts|michigan|minnesota|mississippi|missouri|montana|nebraska|nevada|new hampshire|"
    r"new jersey|new mexico|new york|north carolina|north dakota|ohio|oklahoma|oregon|pennsylvania|"
    r"rhode island|south carolina|south dakota|tennessee|texas|utah|vermont|virginia|washington|"
    r"west virginia|wisconsin|wyoming)\b")


def is_geo(kw):
    return bool(_STATE.search(kw.lower()))


def carrier_key(brand):
    b = (brand or "").lower()
    for k in SPLIT_CARRIERS:
        if k in b:
            return k
    return None


def load_campaigns(path):
    with open(path, newline="", encoding="utf-8") as f:
        return [r["campaign"] for r in csv.DictReader(f) if r.get("campaign")]


def derive_carryover(pause_review_path):
    """From the real grab-bag keywords, distinct (keyword, match) per split carrier, splitting geo out."""
    carry = {k: set() for k in SPLIT_CARRIERS}
    dropped_geo = {k: set() for k in SPLIT_CARRIERS}
    if not pause_review_path or not os.path.exists(pause_review_path):
        print("  [WARN] --pause-review-csv missing; carry-over NOT derived -> new groups get bare terms "
              "ONLY, and pausing the grab-bag would drop the home terms. Provide it.")
        return carry, dropped_geo
    with open(pause_review_path, newline="", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            k = carrier_key(r.get("Brand"))
            if not k:
                continue
            item = (r["Keyword"].strip(), r["Match Type"].strip())
            (dropped_geo if is_geo(item[0]) else carry)[k].add(item)
    return carry, dropped_geo


def build_keyword_rows(campaigns, carry):
    rows = []
    for camp in campaigns:
        # split carriers: carry-over (derived) + new bare terms, deduped
        for k, spec in SPLIT_CARRIERS.items():
            kws = set(carry.get(k, set()))
            kws |= {(t, "Exact") for t in spec["bare_exact"]}
            kws |= {(t, "Broad") for t in spec["bare_broad"]}
            for kw, mt in sorted(kws):
                rows.append((camp, spec["ad_group"], kw, mt, "Enabled"))
        # existing-group bare adds
        for ag, spec in EXISTING_ADGROUP_ADDS.items():
            for kw in spec.get("exact", []):
                rows.append((camp, ag, kw, "Exact", "Enabled"))
            for kw in spec.get("broad", []):
                rows.append((camp, ag, kw, "Broad", "Enabled"))
    return rows


def build_adgroup_rows(campaigns):
    return [(camp, spec["ad_group"], "Enabled") for camp in campaigns for spec in SPLIT_CARRIERS.values()]


def dropped_geo_rows(dropped_geo):
    out = []
    for k, items in dropped_geo.items():
        for kw, mt in sorted(items):
            out.append((SPLIT_CARRIERS[k]["ad_group"], kw, mt,
                        "paused in grab-bag, NOT re-added (base term + geo-campaign covers it)"))
    return out


def write_csv(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)


def write_xlsx(path, adg, kws, dropped):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        print("  (openpyxl unavailable; CSVs only)")
        return
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    hf, fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="365F91")

    def sheet(name, header, rows, note=None):
        ws = wb.create_sheet(name)
        start = 1
        if note:
            ws.cell(row=1, column=1, value=note).font = Font(italic=True, color="808080")
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(header))
            ws.row_dimensions[1].height = 28
            ws.cell(row=1, column=1).alignment = Alignment(wrap_text=True, vertical="top")
            start = 2
        for j, h in enumerate(header, 1):
            c = ws.cell(row=start, column=j, value=h)
            c.font = hf
            c.fill = fill
        for r in rows:
            ws.append(list(r))
        ws.freeze_panes = "A%d" % (start + 1)
        for col, w in {1: 52, 2: 40, 3: 30, 4: 22, 5: 12}.items():
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = w

    sheet("Ad Groups", ["Campaign", "Ad Group", "Status"], adg,
          note="Create these NEW dedicated ad groups in every campaign (Hartford already exists).")
    sheet("Keywords", ["Campaign", "Ad Group", "Keyword", "Match Type", "Status"], kws,
          note="Split carriers = grab-bag carry-over (derived) + new bare terms; mirrored across all "
               "campaigns. Everything the pause sheet removes is re-added here except geo (see Dropped_Geo).")
    sheet("Dropped_Geo", ["Ad Group", "Keyword", "Match Type", "Reason"], dropped,
          note="Geo variants paused in the grab-bag and intentionally NOT re-added: the base term inside "
               "each geo-campaign already covers these queries. Review to confirm you're OK dropping them.")
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description="Carrier-additions Editor bulk sheet (carry-over derived)")
    ap.add_argument("--campaigns-csv", default="campaign_summary.csv")
    ap.add_argument("--pause-review-csv", default=None,
                    help="inventory pause_split_review.csv -> derive grab-bag carry-over (avoids coverage loss)")
    ap.add_argument("--out-dir", default="additions_output")
    a = ap.parse_args()
    campaigns = load_campaigns(a.campaigns_csv)
    if len(campaigns) != 7:
        print("  [note] expected 7 campaigns, found %d -- verify the list." % len(campaigns))
    carry, dropped_geo = derive_carryover(a.pause_review_csv)
    for k in SPLIT_CARRIERS:
        if not carry.get(k):
            print("  [WARN] no grab-bag carry-over derived for '%s' -> its new group gets BARE terms only. "
                  "Re-run inventory_assessment.py --pause-split-carriers to include it, or pausing its "
                  "grab-bag will drop its home terms." % k)
    os.makedirs(a.out_dir, exist_ok=True)
    adg = build_adgroup_rows(campaigns)
    kws = build_keyword_rows(campaigns, carry)
    dropped = dropped_geo_rows(dropped_geo)
    write_csv(os.path.join(a.out_dir, "bulk_ad_groups.csv"), ["Campaign", "Ad Group", "Status"], adg)
    write_csv(os.path.join(a.out_dir, "bulk_keywords.csv"),
              ["Campaign", "Ad Group", "Keyword", "Match Type", "Status"], kws)
    write_csv(os.path.join(a.out_dir, "dropped_geo_REPORT_do_not_import.csv"),
              ["Ad Group", "Keyword", "Match Type", "Reason"], dropped)
    write_xlsx(os.path.join(a.out_dir, "carrier_additions_bulk.xlsx"), adg, kws, dropped)
    for k in SPLIT_CARRIERS:
        print("  %-8s carry-over kept: %d | geo dropped: %d | bare added: %d"
              % (k, len(carry.get(k, set())), len(dropped_geo.get(k, set())),
                 len(SPLIT_CARRIERS[k]["bare_exact"]) + len(SPLIT_CARRIERS[k]["bare_broad"])))
    print("campaigns: %d | keyword rows: %d | new ad groups/campaign: %d | geo-dropped (distinct): %d"
          % (len(campaigns), len(kws), len(SPLIT_CARRIERS), len(dropped)))
    print("wrote -> %s" % os.path.abspath(a.out_dir))


if __name__ == "__main__":
    main()
