"""
Current Tiering Assessment — Microsoft Home geo-tiering, activity-based. (v2)

Answers: "Given the campaigns that are ACTUALLY live (activity in the last
30 days), how does the current committed tiering hold up over an arbitrary
analysis window — and do multiple windows agree?"

v2 additions (motivated by the first real runs):
  A. --compare-windows "84,42,28" — runs the assessment across several windows
     in one invocation and produces a WINDOW AGREEMENT sheet: per state, the
     fixed-centroid tier under each window, an agreement verdict, and a
     re-baseline recommendation (RETIER_CANDIDATE / COMMIT_CANDIDATE /
     ON_TARGET / WATCH / REVIEW_VOLUME). States that agree across windows are
     robust to the window choice; states that disagree are still noise.
  B. DROPPED states now show their sub-threshold volumes and (low-confidence)
     economics instead of blanks — a 190-click near-miss reads differently
     from a true zero, and eligibility-mode flips between windows (clicks-only
     vs clicks-OR-conversions) become explainable.
  C. --emit-retier-draft — writes a PROPOSED commitments YAML (draft only;
     the real ledger is never touched). Reviewed and pasted by a human, per
     the ledger's design principle.

Core behavior unchanged from v1:
  1. Campaign selection is activity-based (search-network clicks/spend in a
     trailing activity window, default 30d), not status='ENABLED'. Paused-
     with-activity campaigns are included and flagged; dormant shells excluded.
  2. Window fully parameterized; thresholds, eligibility mode, RPL blend, and
     stability sub-windows auto-scale (see Methodology sheet).
  3. Assessment scores fresh output against the committed ledger on the
     FIXED-CENTROID basis (frozen baseline centroids), with reclustered tiers
     alongside and boundary-shift flags. Classifications follow the
     commitment-ledger taxonomy, optionally enriched with Trajectory Watch.

Method (clustering, network definitions, RPL/1.5, label normalization) is
imported verbatim from microsoft_geo_model_comparison.py.

Usage:
    # Single window (same as v1)
    python current_tiering_assessment.py --end-date 2026-09-07 --window-days 42

    # Multi-window agreement + re-baseline draft
    python current_tiering_assessment.py --end-date 2026-09-07 \\
        --compare-windows "84,42,28" --emit-retier-draft
"""

from __future__ import annotations

import argparse
import sqlite3
import sys
from dataclasses import dataclass, field
from datetime import date, timedelta
from pathlib import Path
from typing import Optional

import numpy as np
import pandas as pd

import microsoft_geo_model_comparison as core

# ----------------------------------------------------------------------------
# Configuration
# ----------------------------------------------------------------------------

DEFAULT_ACTIVITY_DAYS = 30
DEFAULT_WINDOW_DAYS = 84
BLEND_MIN_WINDOW_DAYS = 56
BLEND_RECENT_DAYS = 28
DEFAULT_BLEND_WEIGHT = 0.6
SHORT_WINDOW_CUTOFF = 56
CORE_CLICKS_PER_84D = 200
MIN_CLICK_THRESHOLD = 50
MIN_STABILITY_DAYS = 14

_TIER_RANK = {t: i + 1 for i, t in enumerate(core.TIER_LABEL_BASE)}
_RANK_TIER = {v: k for k, v in _TIER_RANK.items()}


# ----------------------------------------------------------------------------
# Windows
# ----------------------------------------------------------------------------

@dataclass
class AssessmentWindows:
    end_date: date
    activity_start: date
    full_start: date
    mid_start: date
    short_start: date
    window_days: int
    mid_days: int
    short_days: int


def build_assessment_windows(
    end_date: date, window_days: int, activity_days: int,
) -> AssessmentWindows:
    mid_days = max(MIN_STABILITY_DAYS, round(window_days * 2 / 3))
    short_days = max(MIN_STABILITY_DAYS, round(window_days / 2))
    return AssessmentWindows(
        end_date=end_date,
        activity_start=end_date - timedelta(days=activity_days - 1),
        full_start=end_date - timedelta(days=window_days - 1),
        mid_start=end_date - timedelta(days=mid_days - 1),
        short_start=end_date - timedelta(days=short_days - 1),
        window_days=window_days,
        mid_days=mid_days,
        short_days=short_days,
    )


def scaled_click_threshold(window_days: int) -> int:
    return max(MIN_CLICK_THRESHOLD, round(CORE_CLICKS_PER_84D * window_days / 84))


def scaled_conv_threshold(window_days: int) -> int:
    return max(5, round(10 * window_days / 30))


# ----------------------------------------------------------------------------
# Activity-based campaign selection
# ----------------------------------------------------------------------------

def pull_active_campaigns(
    conn: sqlite3.Connection,
    activity_start: date,
    end_date: date,
    networks: set[str],
) -> pd.DataFrame:
    """Campaigns with search-network activity (clicks or spend) in the activity
    window, joined to the campaigns table. campaign_id normalized to str."""
    np_ = ",".join("?" for _ in networks)
    act = pd.read_sql_query(
        f"""
        SELECT campaign_id,
               SUM(clicks) AS activity_clicks,
               SUM(cost_micros) / 1e6 AS activity_spend,
               SUM(conversions) AS activity_conversions,
               MAX(date) AS last_active_date
        FROM geo_performance_daily
        WHERE account_id = ? AND location_type = ? AND device = ?
          AND date BETWEEN ? AND ?
          AND network IN ({np_})
        GROUP BY campaign_id
        HAVING SUM(clicks) > 0 OR SUM(cost_micros) > 0
        """,
        conn,
        params=[
            core.ACCOUNT_ID, core.LOCATION_TYPE, core.DEVICE,
            activity_start.isoformat(), end_date.isoformat(),
            *sorted(networks),
        ],
    )
    if act.empty:
        return act
    camps = pd.read_sql_query(
        """SELECT campaign_id, campaign_name, status, bidding_strategy_type,
                  target_cpa_micros
           FROM campaigns WHERE account_id = ?""",
        conn, params=(core.ACCOUNT_ID,),
    )
    camps["current_tcpa"] = camps["target_cpa_micros"].apply(
        lambda x: x / 1e6 if pd.notna(x) and x > 0 else None
    )
    act["campaign_id"] = act["campaign_id"].astype(str)
    camps["campaign_id"] = camps["campaign_id"].astype(str)
    out = act.merge(
        camps[["campaign_id", "campaign_name", "status",
               "bidding_strategy_type", "current_tcpa"]],
        on="campaign_id", how="left",
    )
    out["status"] = out["status"].fillna("(not in campaigns table)")
    return out.sort_values("activity_spend", ascending=False).reset_index(drop=True)


# ----------------------------------------------------------------------------
# Eligibility (window-adaptive)
# ----------------------------------------------------------------------------

def eligible_pool_adaptive(
    df: pd.DataFrame,
    window_days: int,
    click_threshold: int,
    conv_threshold: int,
) -> tuple[pd.DataFrame, str]:
    if window_days >= SHORT_WINDOW_CUTOFF:
        pool = core.eligible_pool(df, click_threshold=click_threshold)
        label = f">={click_threshold} search clicks ({window_days}-day window)"
    else:
        clicks_ok = df["clicks"] >= click_threshold
        conv_ok = df["conversions"] >= conv_threshold
        pool = df[
            (clicks_ok | conv_ok)
            & df["rpl"].notna() & df["cpc"].notna()
            & (df["rpl"] > 0) & (df["cpc"] > 0)
            & (df["state"].str.upper() != "UNKNOWN")
        ].copy().reset_index(drop=True)
        label = (f">={click_threshold} clicks OR >={conv_threshold} conversions "
                 f"({window_days}-day window; short-window permissive mode)")
    return pool, label


def run_models_for_subwindow(
    conn, start: date, end: date, campaign_ids: list, k: int,
    window_days: int, click_threshold: int, conv_threshold: int,
) -> pd.DataFrame | None:
    agg = core.pull_search_state_aggregates(conn, start, end, campaign_ids)
    if agg.empty:
        return None
    agg = core.compute_derived(agg)
    pool, _ = eligible_pool_adaptive(agg, window_days, click_threshold, conv_threshold)
    if len(pool) < k:
        return None
    result, _ = core.apply_all_models(pool, k)
    return result


# ----------------------------------------------------------------------------
# Ledger / centroid / trajectory loading
# ----------------------------------------------------------------------------

def load_ledger_commitments(ledger_path: Path) -> tuple[dict[str, str], dict]:
    try:
        import yaml
    except ImportError:
        print("WARNING: PyYAML not installed; skipping ledger comparison.", file=sys.stderr)
        return {}, {}
    if not ledger_path.exists():
        print(f"WARNING: ledger not found at {ledger_path}; skipping comparison.", file=sys.stderr)
        return {}, {}
    with open(ledger_path) as f:
        data = yaml.safe_load(f) or {}
    root = data.get("microsoft_home_geo_tiers", data)
    meta = {
        "baseline_date": root.get("baseline_date"),
        "baseline_model": root.get("baseline_model"),
        "baseline_blend": root.get("baseline_blend"),
        "baseline_k": root.get("baseline_k"),
    }
    return dict(root.get("commitments", {})), meta


def load_baseline_centroids(centroids_path: Path) -> dict[str, float]:
    try:
        import yaml
    except ImportError:
        return {}
    if not centroids_path.exists():
        print(f"WARNING: centroids not found at {centroids_path}; "
              f"fixed-centroid view skipped.", file=sys.stderr)
        return {}
    with open(centroids_path) as f:
        data = yaml.safe_load(f) or {}
    root = data.get("baseline_centroids", data)
    return {t: float(c) for t, c in root.get("centroids", {}).items()}


def load_trajectory_watch(backtest_workbook: Optional[Path]) -> dict[str, dict]:
    if not backtest_workbook:
        return {}
    if not backtest_workbook.exists():
        print(f"WARNING: backtest workbook not found: {backtest_workbook}", file=sys.stderr)
        return {}
    try:
        tw = pd.read_excel(backtest_workbook, sheet_name="Trajectory Watch", header=3)
    except Exception as e:
        print(f"WARNING: couldn't read Trajectory Watch: {e}", file=sys.stderr)
        return {}
    out = {}
    for _, r in tw.iterrows():
        out[r["state"]] = {
            "classification": r.get("classification"),
            "rpl_slope_pct_per_week": r.get("rpl_slope_pct_per_week"),
            "r_squared": r.get("r_squared"),
        }
    return out


# ----------------------------------------------------------------------------
# Assessment (single window)
# ----------------------------------------------------------------------------

def assign_to_fixed_centroids(
    result: pd.DataFrame, baseline_centroids: dict[str, float],
) -> dict[str, str]:
    if not baseline_centroids:
        return {}
    items = sorted(baseline_centroids.items(), key=lambda x: x[1], reverse=True)

    def _nearest(ideal: float):
        if pd.isna(ideal):
            return None
        return min(items, key=lambda tc: abs(tc[1] - ideal))[0]

    return {s: _nearest(i) for s, i in zip(result["state"], result["ideal_tcpa"])}


def classify_deviation(deviation: int | None, traj_class: str | None) -> tuple[str, str]:
    if deviation is None:
        return "N/A", ""
    if deviation == 0:
        return "ON_TARGET", "No action."
    if abs(deviation) == 1:
        if traj_class == "OSCILLATING":
            return ("OSCILLATING_BOUNDARY",
                    "Hold commitment. Boundary noise, not a real move.")
        if traj_class == "DRIFTING_CONFIRMED":
            return ("DRIFT_INTO_NEXT_TIER",
                    "Watch closely; sustained trajectory has crossed the boundary. "
                    "Consider explicit re-tier event if trend continues.")
        if traj_class == "DRIFTING_UNSTABLE":
            return ("DRIFT_UNCONFIRMED",
                    "Watch. Crossed adjacent tier but trajectory is noisy; "
                    "wait for more anchors.")
        return ("MINOR_DEVIATION",
                "Monitor. 1-tier deviation, no trajectory evidence available.")
    return ("MATERIAL_DEVIATION",
            f"Investigate. State is {abs(deviation)} tiers off commitment; "
            f"confirm data quality, consider re-tier event.")


def build_assessment(
    result_full: pd.DataFrame,
    all_states: pd.DataFrame,
    ledger_commitments: dict[str, str],
    fixed_map: dict[str, str],
    traj_lookup: dict[str, dict],
    click_threshold: int,
    headline_model: str = "C2_weighted_kmeans",
) -> pd.DataFrame:
    """One row per state in (fresh pool ∪ ledger). Deviation/classification
    use the FIXED-CENTROID basis. DROPPED states are enriched with their
    sub-threshold volumes and low-confidence economics from all_states."""
    recluster_map = dict(zip(result_full["state"], result_full[headline_model]))
    econ = result_full.set_index("state")
    raw = all_states.set_index("state") if all_states is not None else None
    all_names = sorted(set(recluster_map) | set(ledger_commitments))

    rows = []
    for state in all_names:
        committed = ledger_commitments.get(state)
        tier_fixed = fixed_map.get(state)
        tier_reclust = recluster_map.get(state)
        traj = traj_lookup.get(state, {})
        traj_class = traj.get("classification")

        in_pool = state in recluster_map
        committed_rank = _TIER_RANK.get(committed) if committed else None
        basis_tier = tier_fixed if tier_fixed else tier_reclust
        basis_rank = _TIER_RANK.get(basis_tier) if basis_tier else None

        sub_note = ""
        rpl = ideal = clicks = conv = spend = None
        if in_pool:
            rpl = float(econ.loc[state, "rpl"])
            ideal = float(econ.loc[state, "ideal_tcpa"])
            clicks = int(econ.loc[state, "clicks"])
            conv = float(econ.loc[state, "conversions"])
            spend = float(econ.loc[state, "spend"])
        elif raw is not None and state in raw.index:
            # Sub-threshold enrichment: show what volume/economics exist
            rr = raw.loc[state]
            clicks = int(rr.get("clicks", 0) or 0)
            conv = float(rr.get("conversions", 0) or 0)
            spend = float(rr.get("spend", 0) or 0)
            r_rpl = rr.get("rpl")
            rpl = float(r_rpl) if pd.notna(r_rpl) else None
            ideal = rpl / core.TARGET_ROAS if rpl else None
            gap = click_threshold - clicks
            sub_note = (f"below threshold by {gap} clicks"
                        if gap > 0 else "below threshold (other criteria)")
            if conv == 0:
                sub_note += "; 0 conversions (RPL unmeasurable)"
            else:
                sub_note += "; economics low-confidence"

        if committed is None and in_pool:
            classification, action, deviation = (
                "NEW_STATE",
                "Eligible in this window but not in the ledger. "
                "Candidate for a commitment if volume persists.", None)
        elif committed is not None and not in_pool:
            classification, deviation = "DROPPED", None
            action = ("Committed but below eligibility in this window. "
                      + (sub_note if sub_note else "No data at all in window."))
        else:
            deviation = (basis_rank - committed_rank
                         if basis_rank is not None and committed_rank is not None
                         else None)
            classification, action = classify_deviation(deviation, traj_class)

        boundary_note = ""
        if tier_fixed and tier_reclust and tier_fixed != tier_reclust:
            boundary_note = f"re-cluster says {tier_reclust} (boundary shift)"

        rows.append({
            "state": state,
            "committed_tier": committed or "(not committed)",
            "tier_fixed_centroid": tier_fixed or ("(not eligible)" if not in_pool else "(no centroids)"),
            "tier_reclustered": tier_reclust or "(not eligible)",
            "deviation_ranks": deviation,
            "classification": classification,
            "recommended_action": action,
            "boundary_shift_note": boundary_note,
            "rpl": rpl,
            "ideal_tcpa": ideal,
            "clicks": clicks,
            "conversions": conv,
            "spend": spend,
            "sub_threshold_note": sub_note,
            "trajectory_class": traj_class or "N/A",
            "slope_pct_per_week": traj.get("rpl_slope_pct_per_week"),
            "r_squared": traj.get("r_squared"),
        })

    df = pd.DataFrame(rows)
    sort_order = {
        "MATERIAL_DEVIATION": 0, "DRIFT_INTO_NEXT_TIER": 1,
        "DRIFT_UNCONFIRMED": 2, "MINOR_DEVIATION": 3,
        "OSCILLATING_BOUNDARY": 4, "DROPPED": 5, "NEW_STATE": 6,
        "ON_TARGET": 7, "N/A": 8,
    }
    df["_sort"] = df["classification"].map(sort_order).fillna(9)
    df["_spend"] = df["spend"].fillna(0)
    df = (df.sort_values(["_sort", "_spend"], ascending=[True, False])
            .drop(columns=["_sort", "_spend"]).reset_index(drop=True))
    return df


def build_campaign_view(
    active_campaigns: pd.DataFrame,
    result_full: pd.DataFrame,
    state_to_camp: dict,
    headline_model: str,
) -> pd.DataFrame:
    s = result_full.copy()
    s["campaign_id"] = s["state"].map({k: str(v) for k, v in state_to_camp.items()})
    rows = []
    for _, camp in active_campaigns.iterrows():
        cid = str(camp["campaign_id"])
        grp = s[s["campaign_id"] == cid]
        if grp.empty:
            rows.append({
                "campaign_id": cid, "campaign_name": camp.get("campaign_name"),
                "status": camp.get("status"), "current_tcpa": camp.get("current_tcpa"),
                "n_primary_states": 0, "states": "",
                "weighted_rpl": None, "implied_tcpa_150": None,
                "tcpa_gap_pct": None, "tier_mix": "",
                "note": "no eligible state has this campaign as primary",
            })
            continue
        w_rpl = ((grp["rpl"] * grp["spend"]).sum() / grp["spend"].sum()
                 if grp["spend"].sum() > 0 else grp["rpl"].mean())
        implied = w_rpl / core.TARGET_ROAS
        cur = camp.get("current_tcpa")
        gap = ((implied - cur) / cur if cur and pd.notna(cur) and cur > 0 else None)
        mix = grp[headline_model].value_counts()
        rows.append({
            "campaign_id": cid, "campaign_name": camp.get("campaign_name"),
            "status": camp.get("status"), "current_tcpa": cur,
            "n_primary_states": len(grp),
            "states": ", ".join(sorted(grp["state"])),
            "weighted_rpl": w_rpl, "implied_tcpa_150": implied,
            "tcpa_gap_pct": gap,
            "tier_mix": ", ".join(f"{t}×{n}" for t, n in mix.items()),
            "note": ("heterogeneous grouping" if len(mix) > 1 else ""),
        })
    return pd.DataFrame(rows)


# ----------------------------------------------------------------------------
# Per-window run bundle
# ----------------------------------------------------------------------------

@dataclass
class WindowRun:
    window_days: int
    windows: AssessmentWindows
    eligibility_label: str
    blend_label: str
    result_full: pd.DataFrame
    assessment_df: Optional[pd.DataFrame]
    stability_df: Optional[pd.DataFrame]
    tier_targets_df: pd.DataFrame
    campaign_view_df: pd.DataFrame
    all_states: pd.DataFrame
    click_threshold: int


def run_window_assessment(
    conn, end_date: date, window_days: int, activity_days: int,
    campaign_ids: list, active_campaigns: pd.DataFrame,
    args, ledger_commitments: dict, baseline_centroids: dict,
    traj_lookup: dict,
) -> WindowRun | None:
    windows = build_assessment_windows(end_date, window_days, activity_days)
    click_threshold = args.click_threshold or scaled_click_threshold(window_days)
    conv_threshold = args.conv_threshold or scaled_conv_threshold(window_days)
    blend_weight = (args.rpl_blend if args.rpl_blend is not None
                    else (DEFAULT_BLEND_WEIGHT
                          if window_days >= BLEND_MIN_WINDOW_DAYS else 0.0))
    if window_days < BLEND_MIN_WINDOW_DAYS:
        blend_weight = 0.0

    raw_full = core.pull_search_state_aggregates(
        conn, windows.full_start, end_date, campaign_ids)
    if raw_full.empty:
        print(f"  [{window_days}d] no rows in window — skipped.", file=sys.stderr)
        return None
    all_states = core.compute_derived(raw_full)

    blend_label = "none (straight window RPL)"
    if blend_weight > 0:
        recent_start = end_date - timedelta(days=BLEND_RECENT_DAYS - 1)
        raw_recent = core.pull_search_state_aggregates(
            conn, recent_start, end_date, campaign_ids)
        raw_recent = core.compute_derived(raw_recent) if not raw_recent.empty else None
        all_states = core.apply_rpl_blend(all_states, raw_recent, blend_weight)
        n_blended = int(all_states["blend_applied"].sum())
        blend_label = (f"{int((1-blend_weight)*100)}/{int(blend_weight*100)} "
                       f"window/{BLEND_RECENT_DAYS}d recent "
                       f"({n_blended}/{len(all_states)} states blended)")

    pool, eligibility_label = eligible_pool_adaptive(
        all_states, window_days, click_threshold, conv_threshold)
    if len(pool) < args.k:
        print(f"  [{window_days}d] eligible pool ({len(pool)}) < k ({args.k}) — skipped.",
              file=sys.stderr)
        return None
    result_full, _ = core.apply_all_models(pool, args.k)

    result_mid = run_models_for_subwindow(
        conn, windows.mid_start, end_date, campaign_ids, args.k,
        windows.mid_days, scaled_click_threshold(windows.mid_days),
        scaled_conv_threshold(windows.mid_days))
    result_short = run_models_for_subwindow(
        conn, windows.short_start, end_date, campaign_ids, args.k,
        windows.short_days, scaled_click_threshold(windows.short_days),
        scaled_conv_threshold(windows.short_days))
    stability_df = core.build_stability_table(result_full, result_mid, result_short)

    fixed_map = assign_to_fixed_centroids(result_full, baseline_centroids)
    assessment_df = None
    if ledger_commitments:
        assessment_df = build_assessment(
            result_full, all_states, ledger_commitments, fixed_map,
            traj_lookup, click_threshold, args.headline_model)

    state_to_camp = core.pull_state_to_campaign_map(
        conn, windows.full_start, end_date, campaign_ids)
    tier_targets_df = core.tier_aggregate_metrics(result_full, args.headline_model)
    campaign_view_df = build_campaign_view(
        active_campaigns, result_full, state_to_camp, args.headline_model)

    print(f"  [{window_days}d] eligible: {len(pool)}  |  {eligibility_label}  |  blend: {blend_label}")
    return WindowRun(
        window_days=window_days, windows=windows,
        eligibility_label=eligibility_label, blend_label=blend_label,
        result_full=result_full, assessment_df=assessment_df,
        stability_df=stability_df, tier_targets_df=tier_targets_df,
        campaign_view_df=campaign_view_df, all_states=all_states,
        click_threshold=click_threshold,
    )


# ----------------------------------------------------------------------------
# Window agreement (multi-window mode)
# ----------------------------------------------------------------------------

def build_window_agreement(
    runs: list[WindowRun],
    ledger_commitments: dict[str, str],
) -> pd.DataFrame:
    """Per state across all windows: fixed-centroid tier per window, agreement
    verdict, and a re-baseline recommendation.

    Verdicts (windows ordered longest -> shortest; shortest = most recent-weighted):
      AGREE            — same fixed-centroid tier in every window it's eligible in,
                         eligible in all windows
      AGREE_DIRECTION  — eligible in all windows; deviations vs commitment share
                         a sign but differ in magnitude
      MIXED            — eligible in all windows; deviations conflict in sign
      PARTIAL          — eligible in some windows only (often an eligibility-mode
                         artifact between long clicks-only and short OR windows)
      DROPPED_ALL      — committed but eligible in no window

    Recommendations:
      ON_TARGET         — consensus == committed
      RETIER_CANDIDATE  — AGREE with consensus != committed (proposed = consensus),
                          or AGREE_DIRECTION (proposed = shortest-window tier,
                          marked directional — verify)
      COMMIT_CANDIDATE  — NEW state, AGREE across all windows
      WATCH             — MIXED, PARTIAL, or new state without full agreement
      REVIEW_VOLUME     — DROPPED_ALL
    """
    runs_sorted = sorted(runs, key=lambda r: -r.window_days)
    win_labels = [f"{r.window_days}d" for r in runs_sorted]

    per_window_fixed: list[dict] = []
    per_window_raw: list[pd.DataFrame] = []
    for r in runs_sorted:
        if r.assessment_df is not None:
            fm = dict(zip(r.assessment_df["state"], r.assessment_df["tier_fixed_centroid"]))
        else:
            fm = {}
        per_window_fixed.append(fm)
        per_window_raw.append(r.all_states.set_index("state"))

    all_names = sorted(set(ledger_commitments) | set().union(*[set(m) for m in per_window_fixed]))
    shortest = runs_sorted[-1]
    shortest_econ = shortest.result_full.set_index("state")

    rows = []
    for state in all_names:
        committed = ledger_commitments.get(state)
        tiers = []
        for fm in per_window_fixed:
            t = fm.get(state)
            tiers.append(t if (t and not str(t).startswith("(")) else None)
        eligible_in = [t is not None for t in tiers]
        n_elig = sum(eligible_in)
        valid_tiers = [t for t in tiers if t is not None]
        ranks = [_TIER_RANK.get(t) for t in valid_tiers]
        committed_rank = _TIER_RANK.get(committed) if committed else None

        if n_elig == 0:
            verdict = "DROPPED_ALL" if committed else "NO_DATA"
        elif n_elig < len(runs_sorted):
            verdict = "PARTIAL"
        elif len(set(valid_tiers)) == 1:
            verdict = "AGREE"
        elif committed_rank is not None:
            devs = [rk - committed_rank for rk in ranks]
            signs = {np.sign(d) for d in devs if d != 0}
            verdict = "AGREE_DIRECTION" if len(signs) == 1 and 0 not in {d for d in devs} else "MIXED"
        else:
            verdict = "MIXED"

        consensus = valid_tiers[-1] if valid_tiers else None  # shortest window = most current
        if verdict == "AGREE":
            consensus = valid_tiers[0]

        if verdict in ("DROPPED_ALL",):
            rec, proposed = "REVIEW_VOLUME", None
        elif verdict in ("PARTIAL", "MIXED", "NO_DATA"):
            rec, proposed = "WATCH", None
        elif committed is None:
            rec, proposed = ("COMMIT_CANDIDATE", consensus) if verdict == "AGREE" else ("WATCH", None)
        elif verdict == "AGREE" and consensus == committed:
            rec, proposed = "ON_TARGET", committed
        elif verdict == "AGREE":
            rec, proposed = "RETIER_CANDIDATE", consensus
        else:  # AGREE_DIRECTION
            rec, proposed = "RETIER_CANDIDATE (directional — verify)", consensus

        spend_short = (float(shortest_econ.loc[state, "spend"])
                       if state in shortest_econ.index else None)
        rpl_short = (float(shortest_econ.loc[state, "rpl"])
                     if state in shortest_econ.index else None)

        row = {"state": state, "committed_tier": committed or "(not committed)"}
        for lbl, t, raw in zip(win_labels, tiers, per_window_raw):
            if t is not None:
                row[f"tier_{lbl}"] = t
            elif state in raw.index:
                row[f"tier_{lbl}"] = f"(below thresh: {int(raw.loc[state].get('clicks', 0) or 0)} clicks)"
            else:
                row[f"tier_{lbl}"] = "(no data)"
        row.update({
            "windows_eligible": f"{n_elig}/{len(runs_sorted)}",
            "verdict": verdict,
            "recommendation": rec,
            "proposed_tier": proposed or "",
            "rpl_shortest_window": rpl_short,
            "spend_shortest_window": spend_short,
        })
        rows.append(row)

    df = pd.DataFrame(rows)
    rec_order = {
        "RETIER_CANDIDATE": 0, "RETIER_CANDIDATE (directional — verify)": 1,
        "COMMIT_CANDIDATE": 2, "WATCH": 3, "REVIEW_VOLUME": 4, "ON_TARGET": 5,
    }
    df["_sort"] = df["recommendation"].map(rec_order).fillna(6)
    df["_spend"] = df["spend_shortest_window"].fillna(0)
    return (df.sort_values(["_sort", "_spend"], ascending=[True, False])
              .drop(columns=["_sort", "_spend"]).reset_index(drop=True))


def emit_retier_draft(
    agreement_df: pd.DataFrame,
    ledger_meta: dict,
    end_date: date,
    win_labels: list[str],
    output_path: Path,
) -> None:
    """Write a PROPOSED commitments YAML — a draft for human review.
    Never touches the real ledger."""
    import yaml
    proposed = {}
    changes, holds, watches, drops, commits = [], [], [], [], []
    for _, r in agreement_df.iterrows():
        rec = r["recommendation"]
        state = r["state"]
        if rec == "ON_TARGET":
            proposed[state] = r["committed_tier"]
            holds.append(state)
        elif rec.startswith("RETIER_CANDIDATE") and r["proposed_tier"]:
            proposed[state] = r["proposed_tier"]
            changes.append(f"{state}: {r['committed_tier']} -> {r['proposed_tier']}"
                           + (" (directional)" if "directional" in rec else ""))
        elif rec == "COMMIT_CANDIDATE" and r["proposed_tier"]:
            proposed[state] = r["proposed_tier"]
            commits.append(f"{state}: NEW -> {r['proposed_tier']}")
        elif rec == "WATCH":
            if r["committed_tier"] != "(not committed)":
                proposed[state] = r["committed_tier"]  # hold existing commitment
            watches.append(state)
        elif rec == "REVIEW_VOLUME":
            drops.append(state)  # deliberately omitted — decide explicitly

    draft = {
        "PROPOSED_microsoft_home_geo_tiers": {
            "NOTE": ("DRAFT ONLY — generated by current_tiering_assessment.py. "
                     "Review, edit, and paste into the real ledger deliberately. "
                     "This file is never read by any pipeline."),
            "proposed_date": end_date.isoformat(),
            "based_on_windows": win_labels,
            "prior_baseline_date": str(ledger_meta.get("baseline_date")),
            "commitments": dict(sorted(proposed.items(),
                                       key=lambda x: (_TIER_RANK.get(x[1], 99), x[0]))),
            "retier_history_draft_entry": {
                "date": end_date.isoformat(),
                "reason": "full_rebaseline_after_structural_changes",
                "changes": changes,
                "new_commitments": commits,
                "held_on_watch": watches,
                "omitted_pending_volume_review": drops,
            },
        }
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with open(output_path, "w") as f:
        yaml.dump(draft, f, default_flow_style=False, sort_keys=False, allow_unicode=True)
    print(f"Wrote re-tier DRAFT (review before adopting): {output_path}")


# ----------------------------------------------------------------------------
# Workbook writers
# ----------------------------------------------------------------------------

def _make_styles():
    from openpyxl.styles import Alignment, Font, PatternFill
    return {
        "HEADER_FILL": PatternFill("solid", start_color="1F4E79", end_color="1F4E79"),
        "HEADER_FONT": Font(bold=True, color="FFFFFF", name="Arial"),
        "TITLE_FONT": Font(bold=True, size=14, name="Arial"),
        "SECTION_FONT": Font(bold=True, size=11, name="Arial"),
        "BODY_FONT": Font(name="Arial"),
        "ALIGN_CENTER": Alignment(horizontal="center"),
        "CLASS_FILLS": {
            "MATERIAL_DEVIATION": PatternFill("solid", start_color="FFCCCC", end_color="FFCCCC"),
            "DRIFT_INTO_NEXT_TIER": PatternFill("solid", start_color="FFE4B5", end_color="FFE4B5"),
            "DRIFT_UNCONFIRMED": PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC"),
            "MINOR_DEVIATION": PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC"),
            "OSCILLATING_BOUNDARY": PatternFill("solid", start_color="E8F4F8", end_color="E8F4F8"),
            "DROPPED": PatternFill("solid", start_color="FFCCCC", end_color="FFCCCC"),
            "NEW_STATE": PatternFill("solid", start_color="E2EFDA", end_color="E2EFDA"),
            "ON_TARGET": PatternFill("solid", start_color="E8F5E9", end_color="E8F5E9"),
            # agreement-sheet recommendation tints
            "RETIER_CANDIDATE": PatternFill("solid", start_color="FFE4B5", end_color="FFE4B5"),
            "RETIER_CANDIDATE (directional — verify)": PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC"),
            "COMMIT_CANDIDATE": PatternFill("solid", start_color="E2EFDA", end_color="E2EFDA"),
            "WATCH": PatternFill("solid", start_color="E8F4F8", end_color="E8F4F8"),
            "REVIEW_VOLUME": PatternFill("solid", start_color="FFCCCC", end_color="FFCCCC"),
        },
    }


def _write_df(ws, df: pd.DataFrame, header_row: int, styles: dict,
              tint_col: str | None = None):
    from openpyxl.utils import get_column_letter
    for c, col in enumerate(df.columns, start=1):
        cell = ws.cell(row=header_row, column=c, value=str(col))
        cell.font = styles["HEADER_FONT"]
        cell.fill = styles["HEADER_FILL"]
        cell.alignment = styles["ALIGN_CENTER"]
    for r_off, rec in enumerate(df.to_dict("records"), start=1):
        tint_key = rec.get(tint_col, "") if tint_col else ""
        for c_idx, col in enumerate(df.columns, start=1):
            v = rec[col]
            if isinstance(v, float) and np.isnan(v):
                v = None
            elif isinstance(v, (np.integer, np.floating)):
                v = v.item()
            cell = ws.cell(row=header_row + r_off, column=c_idx, value=v)
            cell.font = styles["BODY_FONT"]
            col_str = str(col).lower()
            if any(x in col_str for x in ("rpl", "tcpa", "spend")) and "pct" not in col_str:
                cell.number_format = "$#,##0.00"
            elif "pct" in col_str:
                cell.number_format = "0.0%"
            elif "r_squared" in col_str or "slope" in col_str:
                cell.number_format = "0.000"
            elif "clicks" in col_str or "conversion" in col_str:
                cell.number_format = "#,##0"
            if tint_key in styles["CLASS_FILLS"]:
                cell.fill = styles["CLASS_FILLS"][tint_key]
    for c_idx, col in enumerate(df.columns, start=1):
        ws.column_dimensions[get_column_letter(c_idx)].width = max(
            14, min(48, len(str(col)) + 2))


def _write_notes(ws, notes: list[tuple[str, bool]], styles: dict, width: int = 95):
    for r_idx, (text, is_hdr) in enumerate(notes, start=1):
        ws.cell(row=r_idx, column=1, value=text).font = (
            styles["SECTION_FONT"] if is_hdr else styles["BODY_FONT"])
    ws.column_dimensions["A"].width = width


def _write_single_window_sheets(
    wb, run: WindowRun, active_campaigns: pd.DataFrame,
    ledger_meta: dict, n_centroids: int, headline_model: str,
    styles: dict, k: int, suffix: str = "",
):
    windows = run.windows
    ws = wb.create_sheet(f"Overview{suffix}")
    ws["A1"] = f"Current Tiering Assessment — Microsoft Home ({windows.window_days}d window)"
    ws["A1"].font = styles["TITLE_FONT"]
    config = [
        ("Analysis window", f"{windows.full_start} to {windows.end_date} ({windows.window_days} days)"),
        ("Activity window (campaign selection)", f"{windows.activity_start} to {windows.end_date}"),
        ("Active campaigns selected", str(len(active_campaigns))),
        ("  of which ENABLED", str(int((active_campaigns["status"] == "ENABLED").sum()))),
        ("Eligibility", run.eligibility_label),
        ("RPL blend", run.blend_label),
        ("Headline model", headline_model),
        ("k (tiers)", str(k)),
        ("Stability sub-windows", f"{windows.window_days}d / {windows.mid_days}d / {windows.short_days}d"),
        ("Eligible states", str(len(run.result_full))),
        ("Ledger baseline", f"{ledger_meta.get('baseline_date')} | {ledger_meta.get('baseline_model')} "
                            f"| k={ledger_meta.get('baseline_k')} | blend={ledger_meta.get('baseline_blend')}"
                            if ledger_meta else "(no ledger loaded)"),
        ("Frozen centroids loaded", str(n_centroids)),
    ]
    r = 3
    for label, val in config:
        ws.cell(r, 1, label).font = styles["BODY_FONT"]
        ws.cell(r, 2, val).font = styles["BODY_FONT"]
        r += 1
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 75

    if run.assessment_df is not None and not run.assessment_df.empty:
        ws = wb.create_sheet(f"Assessment vs Ledger{suffix}")
        ws["A1"] = f"Fresh Model Output vs Committed Tiers ({windows.window_days}d)"
        ws["A1"].font = styles["TITLE_FONT"]
        dist = run.assessment_df["classification"].value_counts().to_dict()
        ws["A2"] = "  ".join(f"{c}: {n}" for c, n in dist.items())
        ws["A2"].font = styles["BODY_FONT"]
        _write_df(ws, run.assessment_df, header_row=4, styles=styles,
                  tint_col="classification")
        ws.freeze_panes = "B5"

    ws = wb.create_sheet(f"Tier Targets{suffix}")
    ws["A1"] = f"Per-Tier Aggregates ({windows.window_days}d)"
    ws["A1"].font = styles["TITLE_FONT"]
    _write_df(ws, run.tier_targets_df, header_row=3, styles=styles)

    ws = wb.create_sheet(f"Campaign Fit{suffix}")
    ws["A1"] = "Active Campaigns — Selection Basis and Current-tCPA Fit"
    ws["A1"].font = styles["TITLE_FONT"]
    ws["A2"] = ("CAUTION: current_tcpa comes from the campaigns table snapshot; "
                "verify it reflects the latest raises before reading tcpa_gap_pct.")
    ws["A2"].font = styles["BODY_FONT"]
    merged = active_campaigns.merge(
        run.campaign_view_df.drop(columns=["campaign_name", "status", "current_tcpa"]),
        on="campaign_id", how="left")
    _write_df(ws, merged, header_row=4, styles=styles)

    cols = ["state", "clicks", "conversions", "spend", "rpl", "cpc", "ideal_tcpa",
            *core.MODEL_NAMES]
    cols = [c for c in cols if c in run.result_full.columns]
    ws = wb.create_sheet(f"Tier Assignments{suffix}")
    ws["A1"] = "Full Model Output (all four methods)"
    ws["A1"].font = styles["TITLE_FONT"]
    _write_df(ws, run.result_full[cols].sort_values("rpl", ascending=False),
              header_row=3, styles=styles)

    if run.stability_df is not None and not run.stability_df.empty:
        ws = wb.create_sheet(f"Stability{suffix}")
        ws["A1"] = (f"Tier Stability Across {windows.window_days}d / "
                    f"{windows.mid_days}d / {windows.short_days}d")
        ws["A1"].font = styles["TITLE_FONT"]
        ws["A2"] = "Columns tier_12w/tier_8w/tier_4w map to full/mid/short sub-windows."
        ws["A2"].font = styles["BODY_FONT"]
        _write_df(ws, run.stability_df.sort_values(["model", "state"]),
                  header_row=4, styles=styles)


METHODOLOGY_NOTES = [
    ("Current Tiering Assessment — Methodology (v2)", True),
    ("", False),
    ("Imports the core model (microsoft_geo_model_comparison.py) verbatim: clustering,", False),
    ("network definitions, RPL and ideal-tCPA formulas, tier-label normalization.", False),
    ("", False),
    ("Deliberate differences from the core model:", True),
    ("1. CAMPAIGN SELECTION — activity-based (search-network clicks/spend in the trailing", False),
    ("   activity window, default 30d), not status='ENABLED'. Paused-with-activity", False),
    ("   campaigns included and flagged; enabled-but-dormant shells excluded.", False),
    ("2. WINDOW — parameterized; thresholds auto-scale:", False),
    (f"   click threshold = max({MIN_CLICK_THRESHOLD}, 200 x window/84);", False),
    (f"   windows < {SHORT_WINDOW_CUTOFF}d use clicks-OR-conversions eligibility;", False),
    (f"   RPL blend ({DEFAULT_BLEND_WEIGHT} recent weight, {BLEND_RECENT_DAYS}d recent side) only when window >= {BLEND_MIN_WINDOW_DAYS}d;", False),
    (f"   stability sub-windows = full / 2/3 / 1/2 (min {MIN_STABILITY_DAYS}d).", False),
    ("", False),
    ("Assessment basis", True),
    ("Deviation and classification use the FIXED-CENTROID assignment (fresh ideal_tcpa", False),
    ("scored against frozen baseline centroids) — isolates state-economics movement", False),
    ("from cluster-boundary movement. Re-clustered tiers shown alongside; disagreements", False),
    ("flagged as boundary shifts. DROPPED states show sub-threshold volumes and", False),
    ("low-confidence economics so near-misses are distinguishable from true zeroes.", False),
    ("", False),
    ("Window Agreement (multi-window mode)", True),
    ("Each window is an independent full run over the same activity-selected campaign", False),
    ("set. The agreement sheet compares fixed-centroid tiers across windows:", False),
    ("  AGREE = same tier everywhere -> robust to window choice.", False),
    ("  AGREE_DIRECTION = deviations share a sign; proposed tier taken from the", False),
    ("    SHORTEST window (most current-regime-weighted) and marked for verification.", False),
    ("  MIXED = windows conflict -> noise; hold.", False),
    ("  PARTIAL = eligible in some windows only. NOTE: long windows use clicks-only", False),
    ("    eligibility while short windows use clicks-OR-conversions, so PARTIAL can be", False),
    ("    an eligibility-mode artifact rather than a volume change — check the", False),
    ("    below-thresh annotations before reading it as signal.", False),
    ("", False),
    ("The ledger remains the operational truth. The re-tier draft (--emit-retier-draft)", False),
    ("is a proposal file for human review; no pipeline reads it, and this script never", False),
    ("modifies the real ledger.", False),
]


def write_workbook_single(
    output_path: Path, run: WindowRun, active_campaigns: pd.DataFrame,
    ledger_meta: dict, n_centroids: int, headline_model: str, k: int,
) -> None:
    from openpyxl import Workbook
    styles = _make_styles()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook(); wb.remove(wb.active)
    _write_single_window_sheets(wb, run, active_campaigns, ledger_meta,
                                n_centroids, headline_model, styles, k)
    ws = wb.create_sheet("Methodology")
    _write_notes(ws, METHODOLOGY_NOTES, styles)
    wb.save(output_path)


def write_workbook_comparison(
    output_path: Path, runs: list[WindowRun], agreement_df: pd.DataFrame,
    active_campaigns: pd.DataFrame, ledger_meta: dict, n_centroids: int,
    headline_model: str, k: int, end_date: date,
) -> None:
    from openpyxl import Workbook
    styles = _make_styles()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook(); wb.remove(wb.active)
    runs_sorted = sorted(runs, key=lambda r: -r.window_days)
    win_labels = [f"{r.window_days}d" for r in runs_sorted]

    # ---- Headline: Window Agreement ----
    ws = wb.create_sheet("Window Agreement")
    ws["A1"] = f"Cross-Window Tier Agreement — windows: {', '.join(win_labels)} (end {end_date})"
    ws["A1"].font = styles["TITLE_FONT"]
    dist = agreement_df["recommendation"].value_counts().to_dict()
    ws["A2"] = "  ".join(f"{c}: {n}" for c, n in dist.items())
    ws["A2"].font = styles["BODY_FONT"]
    ws["A3"] = ("Tiers are FIXED-CENTROID per window. AGREE across windows = robust to "
                "window choice. proposed_tier for directional rows comes from the "
                "shortest window (most current-regime-weighted) — verify before adopting.")
    ws["A3"].font = styles["BODY_FONT"]
    _write_df(ws, agreement_df, header_row=5, styles=styles, tint_col="recommendation")
    ws.freeze_panes = "B6"

    # ---- Per-window sheets ----
    for run in runs_sorted:
        _write_single_window_sheets(
            wb, run, active_campaigns, ledger_meta, n_centroids,
            headline_model, styles, k, suffix=f" ({run.window_days}d)")

    ws = wb.create_sheet("Methodology")
    _write_notes(ws, METHODOLOGY_NOTES, styles)
    wb.save(output_path)


# ----------------------------------------------------------------------------
# Main
# ----------------------------------------------------------------------------

def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--end-date", type=str, default=None)
    parser.add_argument("--window-days", type=int, default=DEFAULT_WINDOW_DAYS)
    parser.add_argument("--compare-windows", type=str, default=None,
                        help='Comma-separated window lengths, e.g. "84,42,28". '
                             "Runs each and adds the Window Agreement sheet. "
                             "Overrides --window-days.")
    parser.add_argument("--activity-days", type=int, default=DEFAULT_ACTIVITY_DAYS)
    parser.add_argument("--k", type=int, default=6)
    parser.add_argument("--db-path", type=str, default=core.DEFAULT_DB_PATH)
    parser.add_argument("--output-dir", type=str,
                        default=str(Path(core.DEFAULT_OUTPUT_DIR) / "assessment"))
    parser.add_argument("--click-threshold", type=int, default=None)
    parser.add_argument("--conv-threshold", type=int, default=None)
    parser.add_argument("--rpl-blend", type=float, default=None)
    parser.add_argument("--ledger", type=str,
                        default="commitments/microsoft_home_baseline.yaml")
    parser.add_argument("--baseline-centroids", type=str,
                        default="commitments/baseline_centroids.yaml")
    parser.add_argument("--backtest-workbook", type=str, default=None)
    parser.add_argument("--headline-model", type=str, default="C2_weighted_kmeans")
    parser.add_argument("--emit-retier-draft", action="store_true",
                        help="Write a proposed-commitments YAML draft "
                             "(multi-window mode only; never touches the ledger).")
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args(argv)

    db_path = Path(args.db_path)
    if not db_path.exists():
        print(f"ERROR: DB not found: {db_path}", file=sys.stderr)
        return 1
    conn = sqlite3.connect(str(db_path))
    end_date = core.resolve_end_date(conn, args.end_date)

    window_list = ([int(x.strip()) for x in args.compare_windows.split(",") if x.strip()]
                   if args.compare_windows else [args.window_days])
    window_list = sorted(set(window_list), reverse=True)
    compare_mode = len(window_list) > 1

    activity_start = end_date - timedelta(days=args.activity_days - 1)
    print("CURRENT TIERING ASSESSMENT (v2)")
    print(f"  End date:         {end_date}")
    print(f"  Windows:          {', '.join(str(w) + 'd' for w in window_list)}")
    print(f"  Activity window:  {activity_start} to {end_date} ({args.activity_days}d)")
    print(f"  k={args.k}  headline={args.headline_model}\n")

    active = pull_active_campaigns(conn, activity_start, end_date, core.SEARCH_NETWORKS)
    if active.empty:
        print("ERROR: no campaigns with activity in the activity window.", file=sys.stderr)
        return 1
    campaign_ids = active["campaign_id"].tolist()
    n_enabled = int((active["status"] == "ENABLED").sum())
    print(f"Active campaigns: {len(active)} ({n_enabled} ENABLED, "
          f"{len(active) - n_enabled} paused/other with activity)")
    for _, c in active.iterrows():
        tag = "" if c["status"] == "ENABLED" else f"  [{c['status']}]"
        print(f"  {c['campaign_id']}: {c['campaign_name']}  "
              f"${c['activity_spend']:,.0f} / {int(c['activity_clicks'])} clicks{tag}")
    print()

    ledger_commitments, ledger_meta = load_ledger_commitments(Path(args.ledger))
    baseline_centroids = load_baseline_centroids(Path(args.baseline_centroids))
    traj_lookup = load_trajectory_watch(
        Path(args.backtest_workbook) if args.backtest_workbook else None)

    runs: list[WindowRun] = []
    for wd in window_list:
        run = run_window_assessment(
            conn, end_date, wd, args.activity_days, campaign_ids, active,
            args, ledger_commitments, baseline_centroids, traj_lookup)
        if run is not None:
            runs.append(run)
    if not runs:
        print("ERROR: no window produced a usable pool.", file=sys.stderr)
        return 1

    for run in runs:
        if run.assessment_df is not None:
            dist = run.assessment_df["classification"].value_counts()
            print(f"\n[{run.window_days}d] assessment vs ledger:")
            for cls, n in dist.items():
                print(f"  {cls}: {n}")

    agreement_df = None
    if compare_mode and len(runs) > 1 and ledger_commitments:
        agreement_df = build_window_agreement(runs, ledger_commitments)
        print("\nWindow agreement recommendations:")
        for rec, n in agreement_df["recommendation"].value_counts().items():
            print(f"  {rec}: {n}")

    if args.dry_run:
        print("\n--dry-run set, skipping workbook write.")
        return 0

    win_labels = [f"{r.window_days}d" for r in sorted(runs, key=lambda r: -r.window_days)]
    if compare_mode and agreement_df is not None:
        output_path = (Path(args.output_dir)
                       / f"current_tiering_assessment_{end_date}_compare_{'-'.join(win_labels)}.xlsx")
        write_workbook_comparison(
            output_path, runs, agreement_df, active, ledger_meta,
            len(baseline_centroids), args.headline_model, args.k, end_date)
        print(f"\nWrote: {output_path}")
        if args.emit_retier_draft:
            draft_path = (Path(args.output_dir)
                          / f"retier_draft_{end_date}.yaml")
            emit_retier_draft(agreement_df, ledger_meta, end_date, win_labels, draft_path)
    else:
        if args.emit_retier_draft:
            print("NOTE: --emit-retier-draft requires --compare-windows with 2+ windows "
                  "and a loaded ledger; skipped.", file=sys.stderr)
        run = runs[0]
        output_path = (Path(args.output_dir)
                       / f"current_tiering_assessment_{end_date}_{run.window_days}d.xlsx")
        write_workbook_single(
            output_path, run, active, ledger_meta,
            len(baseline_centroids), args.headline_model, args.k)
        print(f"\nWrote: {output_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
