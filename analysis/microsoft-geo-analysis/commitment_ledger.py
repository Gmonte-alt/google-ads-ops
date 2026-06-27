"""
Commitment Ledger — Manages locked tier assignments for Microsoft Home geo tiering.

Reads a YAML commitment file (the operational truth of which state belongs in
which tier) and compares it against the latest model output. Surfaces deviations
that warrant attention without forcing tier changes based purely on model drift.

Design principle: the model produces an analytical recommendation. The ledger
holds the operational commitment. They coexist; the ledger overrides only via
explicit re-tier events recorded in retier_history.

CLI usage (standalone):
    python commitment_ledger.py \\
        --ledger commitments\\microsoft_home_baseline.yaml \\
        --workbook output\\microsoft_geo_model_comparison_2026-05-24.xlsx \\
        --output output\\commitment_status_2026-05-24.xlsx

The ledger YAML structure:
    microsoft_home_geo_tiers:
      baseline_date: 2026-05-24
      baseline_model: C2_weighted_kmeans
      baseline_blend: 0.6
      baseline_k: 6
      commitments:
        Colorado: TIER ONE
        Illinois: TIER ONE
        # ...
      retier_history:
        - date: 2026-05-24
          reason: initial_baseline_commitment
          changes: initial
"""

from __future__ import annotations

import argparse
import sys
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Optional

import numpy as np
import pandas as pd

try:
    import yaml
    _YAML_AVAILABLE = True
except ImportError:
    _YAML_AVAILABLE = False


# ----------------------------------------------------------------------------
# YAML I/O
# ----------------------------------------------------------------------------

@dataclass
class CommitmentLedger:
    baseline_date: date
    baseline_model: str
    baseline_blend: float
    baseline_k: int
    commitments: dict[str, str]  # state -> tier
    retier_history: list[dict]


def load_ledger(path: Path) -> CommitmentLedger:
    if not _YAML_AVAILABLE:
        raise RuntimeError(
            "PyYAML not installed. Run: pip install pyyaml"
        )
    with open(path) as f:
        data = yaml.safe_load(f)
    root = data.get("microsoft_home_geo_tiers", data)  # tolerate flat or nested
    return CommitmentLedger(
        baseline_date=root["baseline_date"]
            if isinstance(root["baseline_date"], date)
            else date.fromisoformat(str(root["baseline_date"])),
        baseline_model=root.get("baseline_model", "C2_weighted_kmeans"),
        baseline_blend=float(root.get("baseline_blend", 0.6)),
        baseline_k=int(root.get("baseline_k", 6)),
        commitments=dict(root.get("commitments", {})),
        retier_history=list(root.get("retier_history", [])),
    )


def save_ledger(path: Path, ledger: CommitmentLedger) -> None:
    if not _YAML_AVAILABLE:
        raise RuntimeError("PyYAML not installed.")
    data = {
        "microsoft_home_geo_tiers": {
            "baseline_date": ledger.baseline_date.isoformat(),
            "baseline_model": ledger.baseline_model,
            "baseline_blend": ledger.baseline_blend,
            "baseline_k": ledger.baseline_k,
            "commitments": ledger.commitments,
            "retier_history": ledger.retier_history,
        }
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w") as f:
        yaml.dump(data, f, default_flow_style=False, sort_keys=False)


def create_ledger_from_workbook(
    workbook_path: Path, baseline_date: date,
    headline_model: str = "C2_weighted_kmeans",
    blend: float = 0.6, k: int = 6,
) -> CommitmentLedger:
    """Build an initial ledger by reading the current tier assignments from
    a workbook's Master Tier Review sheet.
    """
    mtr = pd.read_excel(workbook_path, sheet_name="Master Tier Review", header=13)
    above = mtr[mtr["below_threshold"] == False]
    commitments = dict(zip(above["state"], above["tier"]))
    return CommitmentLedger(
        baseline_date=baseline_date,
        baseline_model=headline_model,
        baseline_blend=blend,
        baseline_k=k,
        commitments=commitments,
        retier_history=[{
            "date": baseline_date.isoformat(),
            "reason": "initial_baseline_commitment",
            "changes": "initial",
            "states_committed": len(commitments),
        }],
    )


# ----------------------------------------------------------------------------
# Status computation
# ----------------------------------------------------------------------------

_TIER_RANK = {
    "TIER ONE": 1, "TIER TWO": 2, "TIER THREE": 3,
    "TIER FOUR": 4, "TIER FIVE": 5, "TIER SIX": 6,
    "TIER SEVEN": 7, "TIER EIGHT": 8, "TIER NINE": 9,
    "TIER TEN": 10, "TIER ELEVEN": 11, "TIER TWELVE": 12,
}


def compute_commitment_status(
    ledger: CommitmentLedger,
    current_workbook_path: Path,
    trajectory_df: Optional[pd.DataFrame] = None,
) -> pd.DataFrame:
    """For each committed state, compare commitment vs current model output.

    Returns one row per committed state with:
      state, committed_tier, current_model_tier, current_rpl,
      deviation_magnitude (in tier ranks), classification, recommended_action,
      trajectory_class (if trajectory provided)
    """
    mtr = pd.read_excel(current_workbook_path, sheet_name="Master Tier Review", header=13)

    traj_lookup = {}
    if trajectory_df is not None and not trajectory_df.empty:
        for _, r in trajectory_df.iterrows():
            traj_lookup[r["state"]] = {
                "classification": r["classification"],
                "rpl_slope_pct_per_week": r.get("rpl_slope_pct_per_week"),
                "r_squared": r.get("r_squared"),
                "projected_tier_4w": r.get("projected_tier_4w"),
                "weeks_to_next_tier_boundary": r.get("weeks_to_next_tier_boundary"),
            }

    rows = []
    for state, committed_tier in ledger.commitments.items():
        match = mtr[mtr["state"] == state]
        if match.empty:
            rows.append({
                "state": state,
                "committed_tier": committed_tier,
                "current_model_tier": "NOT IN WORKBOOK",
                "current_rpl": None,
                "deviation_ranks": None,
                "classification": "MISSING",
                "recommended_action": "Verify state data; state was committed but no longer in pool.",
                "trajectory_class": traj_lookup.get(state, {}).get("classification", "N/A"),
                "slope_pct_per_week": None,
                "r_squared": None,
                "projected_tier_4w": None,
                "weeks_to_boundary": None,
            })
            continue

        m = match.iloc[0]
        current_tier = m["tier"]
        current_rpl = m.get("rpl")
        committed_rank = _TIER_RANK.get(committed_tier, 99)
        current_rank = _TIER_RANK.get(current_tier, 99)
        deviation = current_rank - committed_rank  # positive = model says lower tier than committed

        # Pull trajectory info if available
        traj_info = traj_lookup.get(state, {})
        traj_class = traj_info.get("classification", "N/A")

        # Classification + recommended action
        if deviation == 0:
            classification = "ON_TARGET"
            action = "No action."
        elif abs(deviation) == 1 and traj_class == "OSCILLATING":
            classification = "OSCILLATING_BOUNDARY"
            action = "Hold commitment. State oscillates around tier boundary; not a real move."
        elif abs(deviation) == 1 and traj_class == "DRIFTING_CONFIRMED":
            classification = "DRIFT_INTO_NEXT_TIER"
            wks = traj_info.get("weeks_to_next_tier_boundary")
            slope_pct = traj_info.get("rpl_slope_pct_per_week", 0)
            direction = "up" if (slope_pct or 0) > 0 else "down"
            wks_str = f"{wks:.1f}w" if wks else "n/a"
            action = (
                f"Watch closely. Sustained trajectory ({direction}); state has crossed into "
                f"adjacent tier under current model. Consider explicit re-tier event "
                f"if trend continues. Boundary reached in ~{wks_str}."
            )
        elif abs(deviation) == 1 and traj_class == "DRIFTING_UNSTABLE":
            classification = "DRIFT_UNCONFIRMED"
            action = (
                "Watch. State has crossed adjacent tier but trajectory is noisy. "
                "Wait for more anchors before considering re-tier."
            )
        elif abs(deviation) >= 2:
            classification = "MATERIAL_DEVIATION"
            action = (
                f"Investigate. State is {abs(deviation)} tiers off commitment. "
                f"Confirm data quality; consider re-tier event."
            )
        else:
            classification = "MINOR_DEVIATION"
            action = "Monitor. Small deviation; hold commitment."

        rows.append({
            "state": state,
            "committed_tier": committed_tier,
            "current_model_tier": current_tier,
            "current_rpl": float(current_rpl) if pd.notna(current_rpl) else None,
            "deviation_ranks": int(deviation),
            "classification": classification,
            "recommended_action": action,
            "trajectory_class": traj_class,
            "slope_pct_per_week": traj_info.get("rpl_slope_pct_per_week"),
            "r_squared": traj_info.get("r_squared"),
            "projected_tier_4w": traj_info.get("projected_tier_4w"),
            "weeks_to_boundary": traj_info.get("weeks_to_next_tier_boundary"),
        })

    df = pd.DataFrame(rows)
    if df.empty:
        return df
    # Sort: MATERIAL_DEVIATION first, then DRIFT_INTO_NEXT_TIER, then others
    sort_order = {
        "MATERIAL_DEVIATION": 0, "DRIFT_INTO_NEXT_TIER": 1,
        "DRIFT_UNCONFIRMED": 2, "OSCILLATING_BOUNDARY": 3,
        "MINOR_DEVIATION": 4, "MISSING": 5, "ON_TARGET": 6,
    }
    df["_sort"] = df["classification"].map(sort_order)
    df = df.sort_values(["_sort", "state"]).drop(columns=["_sort"]).reset_index(drop=True)
    return df


# ----------------------------------------------------------------------------
# Workbook output (or appending to existing workbook)
# ----------------------------------------------------------------------------

def write_commitment_status_workbook(
    output_path: Path, ledger: CommitmentLedger,
    status_df: pd.DataFrame,
) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    HEADER_FILL = PatternFill("solid", start_color="7030A0", end_color="7030A0")
    HEADER_FONT = Font(bold=True, color="FFFFFF", name="Arial")
    TITLE_FONT = Font(bold=True, size=14, name="Arial")
    SECTION_FONT = Font(bold=True, size=11, name="Arial")
    BODY_FONT = Font(name="Arial")

    # Classification fill colors
    CLASS_FILLS = {
        "MATERIAL_DEVIATION": PatternFill("solid", start_color="FFCCCC", end_color="FFCCCC"),
        "DRIFT_INTO_NEXT_TIER": PatternFill("solid", start_color="FFE4B5", end_color="FFE4B5"),
        "DRIFT_UNCONFIRMED": PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC"),
        "OSCILLATING_BOUNDARY": PatternFill("solid", start_color="E8F4F8", end_color="E8F4F8"),
        "ON_TARGET": PatternFill("solid", start_color="E8F5E9", end_color="E8F5E9"),
        "MINOR_DEVIATION": PatternFill("solid", start_color="F5F5F5", end_color="F5F5F5"),
        "MISSING": PatternFill("solid", start_color="FFCCCC", end_color="FFCCCC"),
    }

    def _write_df(ws, df: pd.DataFrame, header_row: int):
        for c, col in enumerate(df.columns, start=1):
            cell = ws.cell(row=header_row, column=c, value=str(col))
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center")
        for r_off, rec in enumerate(df.to_dict("records"), start=1):
            classification = rec.get("classification", "")
            for c_idx, col in enumerate(df.columns, start=1):
                v = rec[col]
                if isinstance(v, float) and np.isnan(v):
                    v = None
                elif isinstance(v, (np.integer, np.floating)):
                    v = v.item()
                cell = ws.cell(row=header_row + r_off, column=c_idx, value=v)
                cell.font = BODY_FONT
                col_str = str(col).lower()
                if "rpl" in col_str and "slope" not in col_str:
                    cell.number_format = "$#,##0.00"
                elif "slope" in col_str or "r_squared" in col_str:
                    cell.number_format = "0.000"
                # Apply row tint based on classification
                if classification in CLASS_FILLS:
                    cell.fill = CLASS_FILLS[classification]
        for c_idx, col in enumerate(df.columns, start=1):
            width = max(14, min(50, len(str(col)) + 2))
            ws.column_dimensions[get_column_letter(c_idx)].width = width

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)

    # ---- Sheet 1: Commitment Status ----
    ws = wb.create_sheet("Commitment Status")
    ws["A1"] = "Tier Commitment vs Current Model Output"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = (
        f"Baseline: {ledger.baseline_date}  |  Model: {ledger.baseline_model}  |  "
        f"k={ledger.baseline_k}, blend={ledger.baseline_blend}  |  "
        f"States committed: {len(ledger.commitments)}"
    )
    ws["A2"].font = BODY_FONT

    # Distribution summary
    if not status_df.empty:
        dist = status_df["classification"].value_counts().to_dict()
        ws["A4"] = "Status distribution:"
        ws["A4"].font = SECTION_FONT
        priority = ["MATERIAL_DEVIATION", "DRIFT_INTO_NEXT_TIER", "DRIFT_UNCONFIRMED",
                    "OSCILLATING_BOUNDARY", "MINOR_DEVIATION", "ON_TARGET", "MISSING"]
        row = 5
        for cls in priority:
            if cls in dist:
                ws[f"A{row}"] = f"  {cls}: {dist[cls]} state(s)"
                ws[f"A{row}"].font = BODY_FONT
                if cls in CLASS_FILLS:
                    ws[f"A{row}"].fill = CLASS_FILLS[cls]
                row += 1

        _write_df(ws, status_df, header_row=row + 1)

    # ---- Sheet 2: Re-tier History ----
    ws = wb.create_sheet("Retier History")
    ws["A1"] = "Re-tier Event History"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = (
        "Each row records a deliberate operational decision to update commitments. "
        "Model output alone does not trigger commitment changes."
    )
    ws["A2"].font = BODY_FONT
    if ledger.retier_history:
        history_df = pd.DataFrame(ledger.retier_history)
        _write_df(ws, history_df, header_row=4)
    else:
        ws["A4"] = "(No re-tier events recorded yet.)"

    # ---- Sheet 3: Committed Tiers (the full ledger) ----
    ws = wb.create_sheet("Committed Tiers")
    ws["A1"] = "Full Tier Commitments (the operational truth)"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = (
        "This is the source of truth for which tier each state belongs to. "
        "Edit the ledger YAML file to change commitments — do not modify this sheet."
    )
    ws["A2"].font = BODY_FONT
    commitments_df = pd.DataFrame([
        {"state": s, "committed_tier": t}
        for s, t in sorted(ledger.commitments.items(),
                            key=lambda x: (_TIER_RANK.get(x[1], 99), x[0]))
    ])
    _write_df(ws, commitments_df, header_row=4)

    # ---- Sheet 4: Methodology ----
    ws = wb.create_sheet("Methodology")
    notes = [
        ("Commitment Ledger — Methodology", True),
        ("", False),
        ("Design principle: separate analytical recommendation from operational commitment.", False),
        ("", False),
        ("The model (C2 weighted k-means) produces a tier assignment per state on each", False),
        ("anchor date. The ledger holds the operational truth — the tier we have committed", False),
        ("to bid each state at. The two coexist; the ledger does not update automatically", False),
        ("based on weekly model output.", False),
        ("", False),
        ("Status classifications", True),
        ("ON_TARGET — model agrees with commitment. No action.", False),
        ("OSCILLATING_BOUNDARY — state deviates 1 tier but trajectory classifies as", False),
        ("  OSCILLATING. Boundary noise, not real movement. Hold commitment.", False),
        ("DRIFT_INTO_NEXT_TIER — state deviates 1 tier and trajectory confirms sustained", False),
        ("  movement (R² ≥ 0.5, slope ≥ 2%/week). Watch closely; consider explicit re-tier.", False),
        ("DRIFT_UNCONFIRMED — state deviates 1 tier; trajectory is directional but noisy.", False),
        ("  Wait for more anchors before acting.", False),
        ("MATERIAL_DEVIATION — state is 2+ tiers off commitment. Investigate immediately.", False),
        ("MINOR_DEVIATION — small deviation in tier rank; monitor.", False),
        ("MISSING — committed state not in current model output. Data quality issue.", False),
        ("", False),
        ("Re-tier events", True),
        ("Re-tier events are deliberate operational decisions, recorded in the ledger's", False),
        ("retier_history. Each event records: date, reason, which states moved, and from", False),
        ("which tier to which tier. This creates an audit trail.", False),
        ("", False),
        ("To execute a re-tier, edit the YAML ledger file directly: update the commitments", False),
        ("section and add a new retier_history entry.", False),
    ]
    for r_idx, (text, is_hdr) in enumerate(notes, start=1):
        c = ws.cell(row=r_idx, column=1, value=text)
        c.font = SECTION_FONT if is_hdr else BODY_FONT
    ws.column_dimensions["A"].width = 95

    wb.save(output_path)


# ----------------------------------------------------------------------------
# Main CLI
# ----------------------------------------------------------------------------

def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--ledger", type=str, required=True,
                        help="Path to YAML commitment ledger file")
    parser.add_argument("--workbook", type=str, required=True,
                        help="Path to current model output workbook")
    parser.add_argument("--output", type=str, required=True,
                        help="Path to commitment status workbook to produce")
    parser.add_argument("--backtest-workbook", type=str, default=None,
                        help="Optional path to backtest workbook for trajectory enrichment")
    parser.add_argument("--init-from-workbook", action="store_true",
                        help="Create a new ledger from the current workbook's tier assignments")
    parser.add_argument("--baseline-date", type=str, default=None,
                        help="If --init-from-workbook, the baseline date to record")
    args = parser.parse_args(argv)

    ledger_path = Path(args.ledger)
    wb_path = Path(args.workbook)

    if args.init_from_workbook:
        if not args.baseline_date:
            print("ERROR: --baseline-date required with --init-from-workbook", file=sys.stderr)
            return 1
        if ledger_path.exists():
            print(f"ERROR: ledger already exists at {ledger_path}. Refusing to overwrite.", file=sys.stderr)
            return 1
        baseline_date = date.fromisoformat(args.baseline_date)
        ledger = create_ledger_from_workbook(wb_path, baseline_date)
        save_ledger(ledger_path, ledger)
        print(f"Created initial ledger: {ledger_path}")
        print(f"  Baseline date: {baseline_date}")
        print(f"  States committed: {len(ledger.commitments)}")

    if not ledger_path.exists():
        print(f"ERROR: ledger not found at {ledger_path}", file=sys.stderr)
        print("       Use --init-from-workbook to create one.", file=sys.stderr)
        return 1

    ledger = load_ledger(ledger_path)
    print(f"Loaded ledger from {ledger_path}")
    print(f"  Baseline: {ledger.baseline_date}, {ledger.baseline_model}, "
          f"k={ledger.baseline_k}, blend={ledger.baseline_blend}")
    print(f"  States committed: {len(ledger.commitments)}")

    # Optional trajectory enrichment
    trajectory_df = None
    if args.backtest_workbook:
        bt_path = Path(args.backtest_workbook)
        if bt_path.exists():
            try:
                trajectory_df = pd.read_excel(bt_path, sheet_name="Trajectory Watch", header=3)
                print(f"  Loaded trajectory data ({len(trajectory_df)} states) from {bt_path}")
            except Exception as e:
                print(f"  WARNING: couldn't load trajectory data: {e}")

    status_df = compute_commitment_status(ledger, wb_path, trajectory_df)
    if status_df.empty:
        print("ERROR: no commitment status rows produced", file=sys.stderr)
        return 1

    # Summary to terminal
    print("\nStatus distribution:")
    for cls, n in status_df["classification"].value_counts().items():
        print(f"  {cls}: {n}")

    output_path = Path(args.output)
    write_commitment_status_workbook(output_path, ledger, status_df)
    print(f"\nWrote: {output_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
