"""
tROAS Target Gradient v3 — Google Home Insurance
═══════════════════════════════════════════════════════════════════════
Per-campaign target-change gradient with an ELASTICITY MODEL (v3): each
rung now carries modeled conversions, spend, and GP — not just expected
ROAS — so the volume/efficiency trade is visible: raising the target
delivers FEWER conversions at higher ROAS; lowering it buys volume at
lower ROAS.

MODEL (two moving parts, both directional):
  1. ROAS response:  expected ROAS(t) = delivery_ratio × t
     (post-freeze delivery ratio realized/current held constant).
  2. Volume response: volume_mult(t) = (t / current) ^ (-ε)
     applied to post-freeze spend/day and conv/day. ε > 0 means higher
     targets shed volume. Modeled revenue/day = spend/day × mult ×
     expected ROAS; modeled GP/day = revenue/day − spend/day (DB terms).

ELASTICITY SOURCE per campaign:
  Weekly ln-ln fit of conversions vs realized ROAS over the trailing
  --fit-weeks (default 26) Mon-Sun weeks: slope of ln(conv) on
  ln(ROAS); ε = −slope, clamped to [0.3, 6.0]. Falls back to
  --elasticity (default 2.0) when the fit has <8 usable weeks, r² <
  0.10, or the wrong sign. The source ("fit ε=2.31 r²=0.42 n=24" or
  "default ε=2.00") is stamped per campaign in the table and Notes.
  CAVEAT: the fit reads historical covariation (which mixes target
  changes with market movement) as if causal — directional only. The
  Google tROAS simulator (Layer 3, UI, after the first clean week) is
  the arbiter for rung-vs-rung curvature; reconcile fitted ε against
  elasticity_regression_google.py where available.

The GP/day column is the payoff: with ε > 1 the modeled GP curve has an
interior maximum — the rung marked ★ maximizes modeled GP in true
economics (revenue corrected by (1+correction) for the decision, DB
numbers never restated in the data columns).

TARGET RETRIEVAL (unchanged from v2): latest TARGET_ROAS event per
campaign_id in campaign_change_events (snapshot campaigns.target_roas
NULL for these ids; fallback only). NOTE: cutover-day detachments from
portfolio strategies may not emit TARGET_ROAS events — spot-check
resolved targets against the UI once.

DATA RULES (unchanged): campaign_performance_daily_search device='ALL',
revenue = native conversion_value, summed numerators;
campaign_performance_geo_daily_search device rows summed; no keyword or
base tables; DB opened mode=ro.

USAGE:
  python troas_target_gradient_v3.py
  python troas_target_gradient_v3.py --elasticity 2.5 --fit-weeks 26
  python troas_target_gradient_v3.py --no-fit          # force default ε
  python troas_target_gradient_v3.py --span-points 4 --step-points 1

OUTPUT: troas_target_gradient_<YYYY-MM-DD>.xlsx
  Charts / Gradient / State Decomp / Notes
"""

import argparse
import io
import math
import sqlite3
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

try:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl (with Pillow) is required: pip install openpyxl pillow")

DEFAULT_DB_PATH = Path(r"C:\Projects\google-ads-ops\db\paid_search.db")
ACCOUNT_ID = 1

ENABLED_CAMPAIGN_IDS = [
    "23385188622",   # CALIFORNIA-NB
    "23394808297",   # TEXAS-NB
    "23394808300",   # FLORIDA-NB
    "23428389081",   # TIER1 (RPL shell)
    "23496334752",   # TIER2
    "23530860942",   # TIER3
    "23586637678",   # TIER4
]

EPS_MIN, EPS_MAX = 0.3, 6.0
FIT_MIN_WEEKS, FIT_MIN_R2 = 8, 0.10

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
CAMP_FILL = PatternFill("solid", fgColor="D9E2F3")
CURR_FILL = PatternFill("solid", fgColor="E2EFDA")
BEN_FILL = PatternFill("solid", fgColor="FFF2CC")
NEG_FILL = PatternFill("solid", fgColor="F8CBAD")
STAR_FILL = PatternFill("solid", fgColor="C6E0B4")     # GP-max rung
SOFT_FONT = Font(italic=True, color="808080", size=9)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


# ─────────────────────────────────────────────────────────────────────
# Data access
# ─────────────────────────────────────────────────────────────────────

def connect_ro(db_path: Path) -> sqlite3.Connection:
    conn = sqlite3.connect(f"file:{db_path.as_posix()}?mode=ro", uri=True)
    conn.row_factory = sqlite3.Row
    return conn


def normalize_target(raw):
    if raw is None:
        return None
    s = str(raw).strip().rstrip(",").rstrip('"').strip()
    if not s:
        return None
    if s.endswith("%"):
        try:
            return float(s.rstrip("%")) / 100.0
        except ValueError:
            return None
    try:
        v = float(s)
    except ValueError:
        return None
    return v / 100.0 if v > 10 else v


def fetch_campaigns(conn):
    ph = ",".join("?" for _ in ENABLED_CAMPAIGN_IDS)
    sql = f"""
        SELECT campaign_id, campaign_name, target_roas
        FROM campaigns
        WHERE account_id = ? AND campaign_id IN ({ph})
          AND status = 'ENABLED'
        ORDER BY campaign_name
    """
    rows = conn.execute(sql, [ACCOUNT_ID, *ENABLED_CAMPAIGN_IDS]).fetchall()
    return [{
        "campaign_id": str(r["campaign_id"]),
        "name": r["campaign_name"],
        "snapshot_target": normalize_target(r["target_roas"]),
    } for r in rows]


def fetch_latest_targets(conn, campaigns):
    ids = [c["campaign_id"] for c in campaigns]
    ph = ",".join("?" for _ in ids)
    sql = f"""
        SELECT e.campaign_id, e.change_date, e.new_value
        FROM campaign_change_events e
        WHERE e.platform = 'google' AND e.account_id = ?
          AND e.change_type = 'TARGET_ROAS' AND e.campaign_id IN ({ph})
        ORDER BY e.change_date
    """
    latest = {}
    try:
        for r in conn.execute(sql, [ACCOUNT_ID, *ids]):
            v = normalize_target(r["new_value"])
            if v is not None:
                latest[str(r["campaign_id"])] = (v, str(r["change_date"])[:10])
    except sqlite3.Error as exc:
        print(f"  WARN: could not read campaign_change_events: {exc}")
    for c in campaigns:
        if c["campaign_id"] in latest:
            c["target"], evt_date = latest[c["campaign_id"]]
            c["target_source"] = f"event {evt_date}"
        elif c["snapshot_target"] is not None:
            c["target"] = c["snapshot_target"]
            c["target_source"] = "snapshot (no TARGET_ROAS events)"
        else:
            c["target"] = None
            c["target_source"] = "UNRESOLVED — supply manually"
    return campaigns


def max_data_date(conn):
    row = conn.execute(
        """SELECT MAX(p.date) FROM campaign_performance_daily_search p
           JOIN campaigns c ON c.campaign_id = p.campaign_id
           WHERE c.account_id = ?""", [ACCOUNT_ID]).fetchone()
    return datetime.strptime(row[0][:10], "%Y-%m-%d").date()


def campaign_perf(conn, ids, start, end):
    ph = ",".join("?" for _ in ids)
    sql = f"""
        SELECT campaign_id,
               SUM(cost_micros) / 1e6 AS cost,
               SUM(conversions)       AS conv,
               SUM(conversion_value)  AS revenue
        FROM campaign_performance_daily_search
        WHERE campaign_id IN ({ph}) AND device = 'ALL'
          AND date BETWEEN ? AND ?
        GROUP BY campaign_id
    """
    rows = conn.execute(sql, [*ids, start.isoformat(), end.isoformat()])
    return {str(r["campaign_id"]): dict(r) for r in rows}


def state_perf(conn, ids, start, end):
    ph = ",".join("?" for _ in ids)
    sql = f"""
        SELECT campaign_id, geo_state,
               SUM(cost_micros) / 1e6 AS cost,
               SUM(conversions)       AS conv,
               SUM(conversion_value)  AS revenue
        FROM campaign_performance_geo_daily_search
        WHERE campaign_id IN ({ph}) AND date BETWEEN ? AND ?
        GROUP BY campaign_id, geo_state
    """
    rows = conn.execute(sql, [*ids, start.isoformat(), end.isoformat()])
    out = {}
    for r in rows:
        out.setdefault(str(r["campaign_id"]), []).append(dict(r))
    return out


def weekly_history(conn, ids, end, weeks):
    """Weekly (Mon-Sun) cost/conv/revenue per campaign for the fit."""
    # snap end back to the last Sunday on/before end
    last_sun = end - timedelta(days=(end.weekday() + 1) % 7)
    start = last_sun - timedelta(days=7 * weeks - 1)
    ph = ",".join("?" for _ in ids)
    sql = f"""
        SELECT campaign_id, date,
               SUM(cost_micros) / 1e6 AS cost,
               SUM(conversions)       AS conv,
               SUM(conversion_value)  AS revenue
        FROM campaign_performance_daily_search
        WHERE campaign_id IN ({ph}) AND device = 'ALL'
          AND date BETWEEN ? AND ?
        GROUP BY campaign_id, date
    """
    rows = conn.execute(sql, [*ids, start.isoformat(), last_sun.isoformat()])
    weekly = {}
    for r in rows:
        d = datetime.strptime(str(r["date"])[:10], "%Y-%m-%d").date()
        wk = d - timedelta(days=d.weekday())        # Monday key
        agg = weekly.setdefault(str(r["campaign_id"]), {}).setdefault(
            wk, {"cost": 0.0, "conv": 0.0, "revenue": 0.0})
        agg["cost"] += r["cost"] or 0
        agg["conv"] += r["conv"] or 0
        agg["revenue"] += r["revenue"] or 0
    return weekly


# ─────────────────────────────────────────────────────────────────────
# Elasticity fit
# ─────────────────────────────────────────────────────────────────────

def fit_elasticity(weeks_dict, default_eps):
    """ln(conv) ~ a + b*ln(roas) per weekly point; eps = -b.
    Returns (eps, source_label)."""
    pts = []
    for agg in weeks_dict.values():
        if agg["cost"] > 0 and agg["conv"] > 0 and agg["revenue"] > 0:
            roas = agg["revenue"] / agg["cost"]
            if roas > 0:
                pts.append((math.log(roas), math.log(agg["conv"])))
    n = len(pts)
    if n < FIT_MIN_WEEKS:
        return default_eps, f"default ε={default_eps:.2f} (fit n={n}<{FIT_MIN_WEEKS})"
    mx = sum(p[0] for p in pts) / n
    my = sum(p[1] for p in pts) / n
    sxx = sum((p[0] - mx) ** 2 for p in pts)
    sxy = sum((p[0] - mx) * (p[1] - my) for p in pts)
    syy = sum((p[1] - my) ** 2 for p in pts)
    if sxx <= 0 or syy <= 0:
        return default_eps, f"default ε={default_eps:.2f} (degenerate fit)"
    b = sxy / sxx
    r2 = (sxy * sxy) / (sxx * syy)
    if b >= 0 or r2 < FIT_MIN_R2:
        return default_eps, (f"default ε={default_eps:.2f} "
                             f"(fit rejected: slope={b:+.2f}, r²={r2:.2f}, n={n})")
    eps = min(max(-b, EPS_MIN), EPS_MAX)
    return eps, f"fit ε={eps:.2f} (r²={r2:.2f}, n={n})"


# ─────────────────────────────────────────────────────────────────────
# Gradient math
# ─────────────────────────────────────────────────────────────────────

def safe_div(a, b):
    return a / b if b else None


def build_ladder(camp, perf, args, db_breakeven, n_days, eps, eps_label):
    cost = perf.get("cost", 0) or 0
    revenue = perf.get("revenue", 0) or 0
    conv = perf.get("conv", 0) or 0
    realized = safe_div(revenue, cost)
    target = camp["target"]
    ratio = safe_div(realized, target) if (realized and target) else None
    gp = revenue - cost
    spend_day, conv_day = cost / n_days, conv / n_days

    rungs = []
    if target is not None and ratio:
        n_steps = int(round(args.span_points / args.step_points))
        for k in range(-n_steps, n_steps + 1):
            pts = k * args.step_points
            t = target + pts / 100.0
            exp_roas = ratio * t
            mult = (t / target) ** (-eps)
            e_spend = spend_day * mult
            e_conv = conv_day * mult
            e_rev = e_spend * exp_roas
            e_gp_db = e_rev - e_spend
            e_gp_true = e_rev * (1 + args.correction_pct / 100.0) - e_spend
            rungs.append({
                "points": pts, "target": t, "expected_roas": exp_roas,
                "e_spend_day": e_spend, "e_conv_day": e_conv,
                "e_gp_day_db": e_gp_db, "e_gp_day_true": e_gp_true,
                "below_breakeven": exp_roas < db_breakeven,
                "is_current": k == 0,
                "is_ben": abs(pts + args.ben_points) < 1e-9,
                "is_gp_max": False,
            })
        best = max(rungs, key=lambda r: r["e_gp_day_true"])
        best["is_gp_max"] = True
        best["gp_max_edge"] = best is rungs[0] or best is rungs[-1]
    return {"cost": cost, "revenue": revenue, "conv": conv, "gp": gp,
            "realized": realized, "ratio": ratio, "rungs": rungs,
            "spend_day": spend_day, "conv_day": conv_day,
            "eps": eps, "eps_label": eps_label}


def short_name(name):
    n = name.replace("HOME INSURANCE", "").replace("_STATE-GROUP_", "")
    n = n.replace("-NB-SEARCH-ALL-USA-STE", "").replace(
        "_COMBINED_TROAS_SEARCH_ALLDEVICE_USA-STE", "")
    return n.strip("-_ ") or name


# ─────────────────────────────────────────────────────────────────────
# Chart — small multiples: ROAS line (left axis) + GP/day (right axis)
# ─────────────────────────────────────────────────────────────────────

def build_gradient_chart_png(campaigns, ladders, args, db_breakeven):
    plotable = [c for c in campaigns if ladders[c["campaign_id"]]["rungs"]]
    if not plotable:
        return None
    n = len(plotable)
    ncols = 2
    nrows = (n + ncols - 1) // ncols
    fig, axes = plt.subplots(nrows, ncols, figsize=(11.5, 3.3 * nrows))
    axes = axes.flatten() if n > 1 else [axes]

    for ax, camp in zip(axes, plotable):
        lad = ladders[camp["campaign_id"]]
        xs = [r["target"] for r in lad["rungs"]]
        roas = [r["expected_roas"] for r in lad["rungs"]]
        gp = [r["e_gp_day_db"] for r in lad["rungs"]]
        conv = [r["e_conv_day"] for r in lad["rungs"]]
        target = camp["target"]
        ben_t = target - args.ben_points / 100.0
        gp_max = next(r for r in lad["rungs"] if r["is_gp_max"])

        ax.axhspan(min(min(roas), db_breakeven) - 0.02, db_breakeven,
                   color="#F8CBAD", alpha=0.4, zorder=0)
        ax.axhline(db_breakeven, color="#C0504D", linestyle="--",
                   linewidth=1.1, label=f"break-even (DB {db_breakeven:.2f})")
        ax.plot(xs, roas, marker="o", markersize=4, color="#1F4E79",
                linewidth=1.7, label="expected ROAS", zorder=3)
        ax.axvline(target, color="#404040", linestyle="-.", linewidth=1.1,
                   label=f"current {target:.2f}")
        ax.axvline(ben_t, color="#BF8F00", linestyle=":", linewidth=1.5,
                   label=f"Ben -{args.ben_points:g}pt")
        ax.set_ylabel("expected ROAS (DB)", fontsize=7.5, color="#1F4E79")
        ax.tick_params(labelsize=7)
        ax.grid(True, alpha=0.22)

        ax2 = ax.twinx()
        ax2.plot(xs, gp, marker="s", markersize=3.5, color="#2E7D32",
                 linewidth=1.6, linestyle="--", label="modeled GP/day")
        ax2.plot(xs, conv, marker="^", markersize=3.5, color="#7F7F7F",
                 linewidth=1.2, linestyle=":", label="modeled conv/day")
        ax2.plot([gp_max["target"]], [gp_max["e_gp_day_db"]], marker="*",
                 markersize=13, color="#2E7D32", zorder=5,
                 label=(f"GP-max {gp_max['points']:+g}pt"
                        + (" (edge)" if gp_max.get("gp_max_edge") else "")))
        ax2.set_ylabel("GP $/day (DB) · conv/day", fontsize=7.5,
                       color="#2E7D32")
        ax2.tick_params(labelsize=7)

        h1, l1 = ax.get_legend_handles_labels()
        h2, l2 = ax2.get_legend_handles_labels()
        ax.legend(h1 + h2, l1 + l2, fontsize=5.6, loc="upper left",
                  framealpha=0.85)
        ax.set_title(f"{short_name(camp['name'])}   [{lad['eps_label']}]",
                     fontsize=8.5, fontweight="bold")
        ax.set_xlabel("candidate tROAS target", fontsize=7.5)

    for ax in axes[len(plotable):]:
        ax.set_visible(False)
    fig.suptitle(
        f"tROAS gradient ±{args.span_points:g}pts — volume model: "
        f"mult=(t/current)^(-ε); higher target ⇒ fewer conversions. "
        f"Directional, not a forecast.", fontsize=9.5)
    fig.tight_layout(rect=[0, 0, 1, 0.965])
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=150)
    plt.close(fig)
    buf.seek(0)
    return buf


# ─────────────────────────────────────────────────────────────────────
# Workbook
# ─────────────────────────────────────────────────────────────────────

def style_header(ws, row, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center")


def write_gradient_sheet(ws, campaigns, ladders, args, db_breakeven,
                         win_start, win_end):
    ws.title = "Gradient"
    ws.sheet_view.showGridLines = False
    headers = ["Campaign", "Rung (pts)", "Target",
               "Expected ROAS", "vs break-even",
               "Modeled conv/day", "Modeled spend/day", "Modeled GP/day (DB)",
               "GP-max", "Elasticity source", "Target source",
               "Actual: realized", "Actual: conv (win)", "Actual: spend (win)",
               "Actual: GP (win)", "Rollback trigger"]
    ws.append([f"tROAS Target Gradient — Google Home — window {win_start} → "
               f"{win_end} (trailing {args.lag_days}d lag-soft excluded) — "
               f"PROVISIONAL until a full clean post-freeze week"])
    ws.cell(row=1, column=1).font = Font(bold=True, size=11)
    ws.append([f"Volume model: mult = (t/current)^(-ε). Modeled columns are "
               f"directional (ε from weekly history fit or default "
               f"{args.elasticity:g}); simulator adjudicates. Break-even "
               f"in DB terms = 1/(1+{args.correction_pct:g}%) = "
               f"{db_breakeven:.3f}. ★ = modeled GP-max rung (true econ). "
               f"Actual columns are the post-freeze window at the CURRENT "
               f"target — identical on every rung by construction."])
    ws.cell(row=2, column=1).font = SOFT_FONT
    ws.append([])
    ws.append(headers)
    style_header(ws, 4, len(headers))

    r = 5
    for camp in campaigns:
        lad = ladders[camp["campaign_id"]]
        if not lad["rungs"]:
            ws.append([camp["name"], "—", None, None, "target unresolved",
                       None, None, None, "", "", camp["target_source"],
                       lad["realized"], lad["conv"], lad["cost"], lad["gp"], ""])
            ws.cell(row=r, column=1).fill = CAMP_FILL
            r += 2
            ws.append([])
            continue
        first = True
        for rung in lad["rungs"]:
            rollback = ""
            if first and lad["realized"]:
                floor = max(db_breakeven, 0.90 * lad["realized"])
                rollback = (f"Revert if rolling-7d realized ROAS < "
                            f"{floor:.2f} (DB) for 5+ days, or CPL rises "
                            f">25% vs post-freeze baseline")
            ws.append([
                camp["name"] if first else "",
                f"{rung['points']:+g}" + (" (current)" if rung["is_current"]
                                          else " (Ben)" if rung["is_ben"] else ""),
                rung["target"],
                rung["expected_roas"],
                "BELOW" if rung["below_breakeven"] else "above",
                rung["e_conv_day"],
                rung["e_spend_day"],
                rung["e_gp_day_db"],
                ("★ edge→" if rung.get("gp_max_edge") else "★")
                if rung["is_gp_max"] else "",
                lad["eps_label"] if first else "",
                camp["target_source"] if first else "",
                lad["realized"],
                lad["conv"],
                lad["cost"],
                lad["gp"],
                rollback,
            ])
            for c in range(1, len(headers) + 1):
                cell = ws.cell(row=r, column=c)
                cell.border = BORDER
                if c in (3, 4, 12):
                    cell.number_format = "0.00"
                elif c == 6:
                    cell.number_format = "0.0"
                elif c in (7, 8, 13, 14, 15):
                    cell.number_format = "#,##0"
            if first:
                ws.cell(row=r, column=1).fill = CAMP_FILL
                ws.cell(row=r, column=1).font = Font(bold=True, size=9)
            fill = (CURR_FILL if rung["is_current"]
                    else BEN_FILL if rung["is_ben"] else None)
            if fill:
                for c in range(2, 9):
                    ws.cell(row=r, column=c).fill = fill
            if rung["is_gp_max"]:
                ws.cell(row=r, column=9).fill = STAR_FILL
                ws.cell(row=r, column=8).fill = STAR_FILL
            if rung["below_breakeven"]:
                ws.cell(row=r, column=4).fill = NEG_FILL
                ws.cell(row=r, column=5).fill = NEG_FILL
            first = False
            r += 1
        r += 1
        ws.append([])

    widths = [46, 12, 8, 10, 11, 11, 12, 13, 7, 26, 20, 11, 11, 12, 11, 48]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_state_sheet(ws, campaigns, states):
    ws.title = "State Decomp"
    ws.sheet_view.showGridLines = False
    headers = ["Campaign", "State", "Spend", "Spend share", "Conv",
               "Revenue (DB $)", "Realized ROAS", "GP (DB $)"]
    ws.append(headers)
    style_header(ws, 1, len(headers))
    r = 2
    for camp in campaigns:
        rows = states.get(camp["campaign_id"], [])
        total_cost = sum(x["cost"] or 0 for x in rows)
        rows.sort(key=lambda x: (x["revenue"] or 0) - (x["cost"] or 0))
        first = True
        for x in rows:
            cost, rev = x["cost"] or 0, x["revenue"] or 0
            gp = rev - cost
            ws.append([camp["name"] if first else "", x["geo_state"], cost,
                       safe_div(cost, total_cost), x["conv"], rev,
                       safe_div(rev, cost), gp])
            for c in range(1, len(headers) + 1):
                cell = ws.cell(row=r, column=c)
                cell.border = BORDER
                if c in (3, 5, 6, 8):
                    cell.number_format = "#,##0"
                elif c == 4:
                    cell.number_format = "0.0%"
                elif c == 7:
                    cell.number_format = "0.00"
            if first:
                ws.cell(row=r, column=1).fill = CAMP_FILL
                ws.cell(row=r, column=1).font = Font(bold=True, size=9)
            if gp < 0:
                ws.cell(row=r, column=8).fill = NEG_FILL
            first = False
            r += 1
        r += 1
        ws.append([])
    for i, w in enumerate([46, 8, 10, 11, 8, 12, 12, 10], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_notes_sheet(ws, args, win_start, win_end, soft_read,
                      db_breakeven, campaigns, ladders):
    ws.title = "Notes"
    lines = [
        ("Run", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Realized window", f"{win_start} → {win_end} (trailing "
                            f"{args.lag_days}d excluded as lag-soft)"),
        ("Provisional flag", "numbers provisional until a full clean "
                             "Mon–Sun post-freeze week exists"),
        ("Target authority", "campaign_change_events latest TARGET_ROAS "
                             "per campaign_id; snapshot fallback. "
                             "Cutover-day portfolio detachments may not "
                             "emit events — spot-check vs UI once."),
        ("Revenue basis", "campaign_performance_daily_search."
                          "conversion_value, device='ALL' (ruled basis)"),
        ("State basis", "campaign_performance_geo_daily_search, device "
                        "rows summed (no ALL rollup)"),
        ("ROAS response", "expected ROAS = delivery ratio × candidate "
                          "target (post-freeze ratio held constant)"),
        ("Volume model", f"mult = (t/current)^(-ε); ε per campaign from "
                         f"weekly ln(conv)~ln(ROAS) fit over trailing "
                         f"{args.fit_weeks} wks (accepted if n≥"
                         f"{FIT_MIN_WEEKS}, slope<0, r²≥{FIT_MIN_R2}; "
                         f"clamped {EPS_MIN}–{EPS_MAX}); else default "
                         f"ε={args.elasticity:g}. Historical covariation "
                         f"treated as causal — DIRECTIONAL ONLY; "
                         f"reconcile vs elasticity_regression_google.py "
                         f"and the UI simulator (Layer 3)."),
        ("GP-max ★", f"rung maximizing modeled GP/day in true economics "
                     f"(revenue × (1+{args.correction_pct:g}%) − spend); "
                     f"displayed GP/day column stays in DB terms. "
                     f"'★ edge→' means the maximum sits at the ladder "
                     f"boundary — the modeled optimum lies beyond the "
                     f"±span, so widen --span-points or treat as 'keep "
                     f"moving in this direction in scheduled steps'"),
        ("Correction", f"{args.correction_pct:g}% applied to break-even "
                       f"threshold (1/(1+c) = {db_breakeven:.3f}) and the "
                       f"GP-max decision only; DB realized never restated"),
        ("Schedule (Layer 4)", "one move per campaign, ≥1 week apart, "
                               "batched same-day; TARGET_ROAS events "
                               "auto-land in change history"),
        ("Soft trailing window", soft_read),
    ]
    for c in campaigns:
        lad = ladders[c["campaign_id"]]
        t = f"{c['target']:.2f}" if c["target"] is not None else "—"
        lines.append((f"{short_name(c['name'])}",
                      f"target {t} [{c['target_source']}] · {lad['eps_label']}"))
    for k, v in lines:
        ws.append([k, v])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=9)
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 115
    for row in ws.iter_rows():
        row[1].alignment = Alignment(wrap_text=True, vertical="top")


# ─────────────────────────────────────────────────────────────────────

def main():
    p = argparse.ArgumentParser()
    p.add_argument("--db-path", default=str(DEFAULT_DB_PATH))
    p.add_argument("--window-start", default="2026-07-30")
    p.add_argument("--lag-days", type=int, default=3)
    p.add_argument("--span-points", type=float, default=4.0)
    p.add_argument("--step-points", type=float, default=1.0)
    p.add_argument("--ben-points", type=float, default=3.0)
    p.add_argument("--correction-pct", type=float, default=11.5)
    p.add_argument("--elasticity", type=float, default=2.0,
                   help="fallback ε when the per-campaign fit is rejected")
    p.add_argument("--fit-weeks", type=int, default=26)
    p.add_argument("--no-fit", action="store_true",
                   help="skip fitting; use --elasticity for all campaigns")
    p.add_argument("--output", default=None)
    args = p.parse_args()

    db_path = Path(args.db_path)
    if not db_path.exists():
        sys.exit(f"DB not found: {db_path}")
    conn = connect_ro(db_path)

    db_breakeven = 1.0 / (1.0 + args.correction_pct / 100.0)

    campaigns = fetch_campaigns(conn)
    if not campaigns:
        sys.exit("No campaigns resolved — check ENABLED_CAMPAIGN_IDS.")
    fetch_latest_targets(conn, campaigns)

    max_date = max_data_date(conn)
    win_start = datetime.strptime(args.window_start, "%Y-%m-%d").date()
    win_end = max_date - timedelta(days=args.lag_days)
    if win_end <= win_start:
        sys.exit(f"Window empty: {win_start} → {win_end}")
    n_days = (win_end - win_start).days + 1
    print(f"Realized window: {win_start} → {win_end} ({n_days}d; max data "
          f"date {max_date}, trailing {args.lag_days}d soft)")

    ids = [c["campaign_id"] for c in campaigns]
    perf = campaign_perf(conn, ids, win_start, win_end)
    states = state_perf(conn, ids, win_start, win_end)
    hist = {} if args.no_fit else weekly_history(conn, ids, win_end,
                                                 args.fit_weeks)
    soft = campaign_perf(conn, ids, win_end + timedelta(days=1), max_date)
    soft_cost = sum(v["cost"] or 0 for v in soft.values())
    soft_rev = sum(v["revenue"] or 0 for v in soft.values())
    soft_read = (f"{win_end + timedelta(days=1)} → {max_date}: spend "
                 f"{soft_cost:,.0f}, revenue {soft_rev:,.0f}, ROAS "
                 f"{safe_div(soft_rev, soft_cost) or 0:.2f} (SOFT)")

    ladders = {}
    print(f"\nGradient (break-even {db_breakeven:.3f}):")
    for camp in campaigns:
        if args.no_fit:
            eps, eps_label = args.elasticity, f"default ε={args.elasticity:.2f} (--no-fit)"
        else:
            eps, eps_label = fit_elasticity(
                hist.get(camp["campaign_id"], {}), args.elasticity)
        lad = build_ladder(camp, perf.get(camp["campaign_id"], {}),
                           args, db_breakeven, n_days, eps, eps_label)
        ladders[camp["campaign_id"]] = lad
        rz = f"{lad['realized']:.2f}" if lad["realized"] else "—"
        gp_max = next((r for r in lad["rungs"] if r["is_gp_max"]), None)
        gm = (f"GP-max {gp_max['points']:+g}pt"
              + (" (EDGE — optimum beyond ladder span)"
                 if gp_max.get("gp_max_edge") else "")) if gp_max else "no ladder"
        print(f"  {short_name(camp['name']):12s} target "
              f"{camp['target'] if camp['target'] else 0:.2f} "
              f"[{camp['target_source']}] realized {rz}  {lad['eps_label']}"
              f"  {gm}")

    wb = Workbook()
    charts_ws = wb.active
    charts_ws.title = "Charts"
    charts_ws.sheet_view.showGridLines = False
    png = build_gradient_chart_png(campaigns, ladders, args, db_breakeven)
    if png:
        charts_ws.add_image(XLImage(png), "A1")
    else:
        charts_ws["A1"] = "No plottable campaigns (targets unresolved)."

    write_gradient_sheet(wb.create_sheet(), campaigns, ladders, args,
                         db_breakeven, win_start, win_end)
    write_state_sheet(wb.create_sheet(), campaigns, states)
    write_notes_sheet(wb.create_sheet(), args, win_start, win_end,
                      soft_read, db_breakeven, campaigns, ladders)

    out = args.output or f"troas_target_gradient_{date.today().isoformat()}.xlsx"
    wb.save(out)
    conn.close()
    print(f"\nSaved: {out}")


if __name__ == "__main__":
    main()
