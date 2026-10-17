"""
termlife_day1_kpi_tracker.py
============================
"The only fair comparator for a one-day-old Policy KPI is other days at the
same age." This script builds that comparator from the batch history in
life_disposition_snapshot:

  A. DAY-1 SERIES   - every cohort day reconstructed AT AGE ONE (from the
                      batch dated day+1): ops, CRIs, prem, rate-card spend,
                      day-1 Policy KPI -- paid segments (CA/CB/CC) only.
  B. BENCHMARK      - yesterday vs the weekday day-1 history: median / min /
                      max / rank, per segment and total. Confirms or refutes
                      "the dip is across all brackets" at equal age.
  C. TRIANGLE       - prem by cohort day x age (how each day matured across
                      batches). Yields the typical age-1 -> age-4 growth
                      multiple and a PROJECTED mature KPI for yesterday.

Read-only. Console summary + small workbook in OUT_DIR.
Weekend cohorts are shown but excluded from benchmarks (ads off Sunday;
Saturday leads largely unworked -- see 9/25 co-worker notes).
"""

import sqlite3
import sys
from datetime import date, datetime
from pathlib import Path
from statistics import median

from openpyxl import Workbook
from openpyxl.styles import Font

DB_PATH = Path(r"C:\Projects\google-ads-ops\db\paid_search.db")
OUT_DIR = Path(r"C:\Projects\google-ads-ops\reports\quotecom-life-diposition\output")

PAID = ("CA", "CB", "CC")
SEG_NAME = {"CA": "Low", "CB": "Med", "CC": "High"}

# Same rate card as the main report -- keep the two in sync.
SEG_OP_RATES = {
    "CC": {"2026-08": 149.67, "2026-09": 144.02, "2026-10": 148.53,
           "default": 148.53},
    "CB": {"2026-08": 86.55, "2026-09": 84.34, "2026-10": 86.75,
           "default": 86.75},
    "CA": {"2026-08": 30.98, "2026-09": 31.51, "2026-10": 33.24,
           "default": 33.24},
}

TRIANGLE_DAYS = 14      # cohort days to show in the maturation triangle
TRIANGLE_MAX_AGE = 8
MATURE_AGE = 4          # age used for the "projected mature KPI" multiple

BOLD = Font(bold=True)


def rate(seg, d: date) -> float:
    r = SEG_OP_RATES[seg]
    return r.get(d.strftime("%Y-%m"), r["default"])


def main():
    if not DB_PATH.exists():
        sys.exit(f"DB not found: {DB_PATH}")
    con = sqlite3.connect(DB_PATH)

    # ---------------- A. day-1 series ------------------------------------
    rows = con.execute(f"""
        SELECT batch_date, day, substr(source_code,1,2) AS seg,
               SUM(CASE WHEN substr(source_code,3,1)='M' THEN 1 ELSE 0 END),
               SUM(CASE WHEN substr(source_code,3,1)='M' THEN op_count ELSE 0 END),
               SUM(sqo), SUM(cris), SUM(capped_cri_dollars)
        FROM life_disposition_snapshot
        WHERE day = date(batch_date, '-1 day')
          AND substr(source_code,1,2) IN {PAID}
        GROUP BY 1, 2, 3 ORDER BY 1, 3""").fetchall()
    if not rows:
        sys.exit("no day-1 observations (need batches where day = batch-1)")

    days = {}   # cohort day -> {seg: dict}
    for batch, day, seg, leads, ops_m, sqo, cris, prem in rows:
        d = date.fromisoformat(day)
        days.setdefault(d, {})[seg] = dict(
            leads=leads or 0, ops=ops_m or 0, sqo=sqo or 0,
            cris=cris or 0, prem=prem or 0.0,
            spend=(ops_m or 0) * rate(seg, d))

    def totals(segmap):
        t = dict(ops=0, sqo=0, cris=0, prem=0.0, spend=0.0)
        for s in segmap.values():
            for k in t:
                t[k] += s[k]
        t["kpi"] = t["prem"] / t["spend"] if t["spend"] else None
        return t

    all_days = sorted(days)
    yday = all_days[-1]
    print(f"== A. DAY-1 KPI SERIES (paid segments, each day AT AGE 1) ==")
    print(f"  {'cohort':>10} {'wd':>3} {'ops':>4} {'CRI':>4} {'prem$':>8} "
          f"{'spend$':>8} {'KPI1':>6}   Hi/Md/Lo KPI1")
    for d in all_days:
        t = totals(days[d])
        seg_k = "/".join(
            f"{(days[d].get(s, {}).get('prem', 0) / days[d][s]['spend']):.2f}"
            if days[d].get(s, {}).get("spend") else "-"
            for s in ("CC", "CB", "CA"))
        mark = "  <-- yesterday" if d == yday else ""
        print(f"  {d.isoformat():>10} {d.strftime('%a'):>3} {t['ops']:4d} "
              f"{t['cris']:4d} {t['prem']:8,.0f} {t['spend']:8,.0f} "
              f"{t['kpi'] if t['kpi'] is not None else 0:6.2f}   {seg_k}{mark}")

    # ---------------- B. benchmark ----------------------------------------
    bench_days = [d for d in all_days
                  if d.weekday() < 5 and d != yday and totals(days[d])["spend"]]
    print(f"\n== B. YESTERDAY ({yday} {yday.strftime('%a')}) vs WEEKDAY DAY-1 "
          f"HISTORY ({len(bench_days)} days) ==")
    if not bench_days:
        print("  not enough weekday history yet")
    else:
        yt = totals(days[yday])
        hist = sorted(totals(days[d])["kpi"] for d in bench_days)
        below = sum(1 for k in hist if k < (yt["kpi"] or 0))
        print(f"  TOTAL  KPI1: yesterday {yt['kpi']:.2f} | weekday median "
              f"{median(hist):.2f} | range {hist[0]:.2f}-{hist[-1]:.2f} | "
              f"rank {below + 1} of {len(hist) + 1} (1 = worst)")
        for s in ("CC", "CB", "CA"):
            hseg = sorted(
                days[d][s]["prem"] / days[d][s]["spend"]
                for d in bench_days if days[d].get(s, {}).get("spend"))
            y = days[yday].get(s, {})
            yk = y["prem"] / y["spend"] if y.get("spend") else None
            if hseg and yk is not None:
                print(f"  {SEG_NAME[s]:>5}  KPI1: yesterday {yk:.2f} | median "
                      f"{median(hseg):.2f} | range {hseg[0]:.2f}-{hseg[-1]:.2f}")
        med_ops = median(totals(days[d])["ops"] for d in bench_days)
        med_hi = median(
            days[d].get("CC", {}).get("ops", 0) / totals(days[d])["ops"]
            for d in bench_days if totals(days[d])["ops"])
        y_hi = (days[yday].get("CC", {}).get("ops", 0) / yt["ops"]
                if yt["ops"] else 0)
        med_cri = median(totals(days[d])["cris"] for d in bench_days)
        print(f"  drivers: ops {yt['ops']} vs median {med_ops:.0f} | High-ops "
              f"share {y_hi:.0%} vs {med_hi:.0%} | day-1 CRIs {yt['cris']} vs "
              f"median {med_cri:.0f}")

    # ---------------- C. maturation triangle ------------------------------
    tri = con.execute(f"""
        SELECT day, CAST(julianday(batch_date) - julianday(day) AS INT) AS age,
               SUM(capped_cri_dollars)
        FROM life_disposition_snapshot
        WHERE substr(source_code,1,2) IN {PAID}
          AND day >= date((SELECT MAX(batch_date)
                           FROM life_disposition_snapshot), '-{TRIANGLE_DAYS} day')
          AND julianday(batch_date) - julianday(day)
              BETWEEN 1 AND {TRIANGLE_MAX_AGE}
        GROUP BY 1, 2 ORDER BY 1, 2""").fetchall()
    cells = {}
    for day, age, prem in tri:
        cells[(date.fromisoformat(day), int(age))] = prem or 0.0
    tdays = sorted({d for d, _ in cells})
    print(f"\n== C. PREM MATURATION TRIANGLE (paid segs; cohort day x age) ==")
    print("  " + f"{'cohort':>10} {'wd':>3} " +
          " ".join(f"a{a:>7}" for a in range(1, TRIANGLE_MAX_AGE + 1)))
    for d in tdays:
        line = " ".join(
            f"{cells[(d, a)]:8,.0f}" if (d, a) in cells else f"{'.':>8}"
            for a in range(1, TRIANGLE_MAX_AGE + 1))
        print(f"  {d.isoformat():>10} {d.strftime('%a'):>3} {line}")

    mults = [cells[(d, MATURE_AGE)] / cells[(d, 1)]
             for d in tdays
             if cells.get((d, 1)) and cells.get((d, MATURE_AGE))
             and d.weekday() < 5]
    if mults and yday in days and totals(days[yday])["spend"]:
        m = median(mults)
        yk = totals(days[yday])["kpi"]
        print(f"\n  typical age-1 -> age-{MATURE_AGE} prem multiple "
              f"(weekday cohorts, n={len(mults)}): x{m:.2f}")
        print(f"  PROJECTED age-{MATURE_AGE} KPI for yesterday: "
              f"{yk:.2f} x {m:.2f} = {yk * m:.2f}   "
              f"(SQ break-even ~2.96, target 4)")

    # ---------------- workbook --------------------------------------------
    wb = Workbook()
    ws = wb.active
    ws.title = "Day1_Series"
    hdr = ["Cohort day", "Weekday", "Ops(M)", "SQOs", "CRIs", "Prem",
           "Spend(est)", "KPI1", "High KPI1", "Med KPI1", "Low KPI1"]
    for j, h in enumerate(hdr, 1):
        ws.cell(1, j, h).font = BOLD
    for i, d in enumerate(all_days, 2):
        t = totals(days[d])
        ws.cell(i, 1, d.isoformat())
        ws.cell(i, 2, d.strftime("%a"))
        ws.cell(i, 3, t["ops"])
        ws.cell(i, 4, t["sqo"])
        ws.cell(i, 5, t["cris"])
        ws.cell(i, 6, round(t["prem"], 2)).number_format = "$#,##0"
        ws.cell(i, 7, round(t["spend"], 2)).number_format = "$#,##0"
        ws.cell(i, 8, t["kpi"]).number_format = "0.00"
        for j, s in enumerate(("CC", "CB", "CA"), 9):
            seg = days[d].get(s, {})
            if seg.get("spend"):
                ws.cell(i, j, seg["prem"] / seg["spend"]).number_format = "0.00"
    ws2 = wb.create_sheet("Prem_Triangle")
    ws2.cell(1, 1, "Cohort day").font = BOLD
    for a in range(1, TRIANGLE_MAX_AGE + 1):
        ws2.cell(1, a + 1, f"age {a}").font = BOLD
    for i, d in enumerate(tdays, 2):
        ws2.cell(i, 1, d.isoformat())
        for a in range(1, TRIANGLE_MAX_AGE + 1):
            if (d, a) in cells:
                ws2.cell(i, a + 1, round(cells[(d, a)], 2)
                         ).number_format = "$#,##0"
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out = OUT_DIR / f"termlife_day1_kpi_{datetime.now():%Y%m%d}.xlsx"
    wb.save(out)
    print(f"\n  workbook -> {out}")
    con.close()


if __name__ == "__main__":
    main()
