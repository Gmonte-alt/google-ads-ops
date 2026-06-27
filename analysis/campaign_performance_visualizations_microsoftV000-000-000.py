"""
Campaign Performance Visualizations — Microsoft Home Insurance
═══════════════════════════════════════════════════════════════════════
Generates a per-campaign weekly performance workbook covering the last
12 weeks of operating equilibrium (default: 12 weeks ending 2026-04-24).

Designed to answer the leadership-side question: "Why isn't the model
finding more headroom?" — by showing each campaign's recent trajectory
on conversions, CPA, RPL, and profit. Each metric is plotted weekly so
the eye picks up trends without being distracted by daily noise.

USAGE:
  python campaign_performance_visualizations_microsoft.py
  python analysis/campaign_performance_visualizations_microsoft.py --end-date 2026-04-24
  python campaign_performance_visualizations_microsoft.py --weeks 16

OUTPUT: campaign_performance_microsoft_<YYYY-MM-DD>.xlsx
  - Overview (small multiples: 12 weekly profit charts on one sheet)
  - <CAMPAIGN_SHORT> sheets (one per campaign, 4-panel weekly chart):
      Panel 1: weekly conversions
      Panel 2: weekly CPA + current tCPA reference line
      Panel 3: weekly RPL
      Panel 4: weekly profit
"""

import argparse
import io
import sqlite3
from datetime import date, timedelta
from pathlib import Path
import sys

import pandas as pd
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.drawing.image import Image as XLImage

# ── CONFIG ─────────────────────────────────────────────────────────
ACCOUNT_ID = 2  # Microsoft Home Insurance
ACCOUNT_LABEL = "microsoft_home"
DEFAULT_DB_PATH = Path(r"C:\Projects\google-ads-ops\db\paid_search.db")
DEFAULT_CHANGE_CSV = Path(r"C:\Projects\google-ads-ops\analysis\input\Microsoft-bids-budget-changehistory-YTD.csv")
DEFAULT_OUTPUT_DIR = Path(r"C:\Projects\google-ads-ops\analysis\output")

DEFAULT_WEEKS = 12
DEFAULT_END_DATE = date(2026, 4, 24)


def banner(text: str) -> None:
    print("═" * 70)
    print(text)
    print("═" * 70)


# ═══════════════════════════════════════════════════════════════════
# DB QUERIES
# ═══════════════════════════════════════════════════════════════════

def fetch_enabled_campaigns(conn) -> pd.DataFrame:
    """Enabled campaigns + their current tCPA. Filters to ENABLED status."""
    sql = """
    SELECT
        campaign_id,
        campaign_name,
        target_cpa_micros / 1.0e6 AS current_tcpa
    FROM campaigns
    WHERE account_id = ?
      AND status = 'ENABLED'
    ORDER BY campaign_name
    """
    df = pd.read_sql(sql, conn, params=(ACCOUNT_ID,))
    return df


def fetch_campaign_daily(conn, campaign_name: str,
                         start_date: date, end_date: date) -> pd.DataFrame:
    """Daily campaign metrics. Aggregates across devices.

    Returns columns: report_date, spend_usd, conv, conv_value.
    """
    sql = """
    SELECT
        cpd.date AS report_date,
        SUM(cpd.cost_micros) / 1e6 AS spend_usd,
        SUM(cpd.conversions) AS conv,
        SUM(cpd.conversion_value) AS conv_value
    FROM campaign_performance_daily cpd
    JOIN campaigns c ON cpd.campaign_id = c.campaign_id
    WHERE c.account_id = ?
      AND c.campaign_name = ?
      AND cpd.date >= ?
      AND cpd.date <= ?
    GROUP BY cpd.date
    ORDER BY cpd.date
    """
    df = pd.read_sql(sql, conn, params=(ACCOUNT_ID, campaign_name,
                                         start_date.isoformat(),
                                         end_date.isoformat()))
    df["report_date"] = pd.to_datetime(df["report_date"]).dt.date
    return df


# ═══════════════════════════════════════════════════════════════════
# WEEKLY AGGREGATION
# ═══════════════════════════════════════════════════════════════════

def aggregate_weekly(daily_df: pd.DataFrame) -> pd.DataFrame:
    """Roll up daily data to ISO weeks (Mon-Sun).

    Returns columns: week_start (Mon date), week_end (Sun date),
    spend, conv, conv_value, cpa, rpl, profit.

    CPA and RPL are weighted ratios (sum spend / sum conv, etc.) — NOT
    averages of daily values. This matches how the strategic script
    computes weekly aggregates.
    """
    if daily_df.empty:
        return pd.DataFrame()

    df = daily_df.copy()
    df["report_date"] = pd.to_datetime(df["report_date"])
    # ISO week: Monday is start of week
    df["week_start"] = df["report_date"] - pd.to_timedelta(
        df["report_date"].dt.weekday, unit="D"
    )
    df["week_start"] = df["week_start"].dt.date

    agg = df.groupby("week_start").agg(
        spend=("spend_usd", "sum"),
        conv=("conv", "sum"),
        conv_value=("conv_value", "sum"),
    ).reset_index()

    agg["week_end"] = agg["week_start"].apply(lambda d: d + timedelta(days=6))
    agg["cpa"] = agg.apply(lambda r: r["spend"] / r["conv"] if r["conv"] > 0 else None, axis=1)
    agg["rpl"] = agg.apply(lambda r: r["conv_value"] / r["conv"] if r["conv"] > 0 else None, axis=1)
    agg["profit"] = agg["conv_value"] - agg["spend"]

    return agg


# ═══════════════════════════════════════════════════════════════════
# CHANGE-HISTORY LOADER (lightweight subset of elasticity script)
# ═══════════════════════════════════════════════════════════════════

def load_change_events(csv_path: Path,
                       start_date: date, end_date: date) -> pd.DataFrame:
    """Load tCPA + budget change events from the YTD change CSV, filtered
    to the visualization window. Used to overlay vertical change-event
    lines on the per-campaign charts.
    """
    if not csv_path.exists():
        return pd.DataFrame(columns=["change_date", "campaign_name",
                                      "attribute", "old_value", "new_value"])
    df = pd.read_csv(csv_path)
    df["change_date"] = pd.to_datetime(df["Date"], format="%m/%d/%Y").dt.date
    df = df.rename(columns={
        "Campaign": "campaign_name",
        "Attribute changed": "attribute",
        "Old value": "old_value",
        "New value": "new_value",
    })
    df = df[(df["change_date"] >= start_date) & (df["change_date"] <= end_date)].copy()
    return df[["change_date", "campaign_name", "attribute",
               "old_value", "new_value"]].reset_index(drop=True)


# ═══════════════════════════════════════════════════════════════════
# CHART BUILDERS
# ═══════════════════════════════════════════════════════════════════

def build_4panel_campaign_chart_png(
    weekly_df: pd.DataFrame,
    campaign_name: str,
    current_tcpa: float | None,
    change_events: list[dict],
    window_start: date,
    window_end: date,
) -> bytes:
    """4-panel weekly chart for one campaign:
       Panel 1 (top):    weekly conversions (bar chart)
       Panel 2:          weekly CPA + current tCPA reference line
       Panel 3:          weekly RPL
       Panel 4 (bottom): weekly profit (bar chart, color-coded green/red)

    Plus vertical lines at any change events within the window.
    """
    if weekly_df.empty:
        fig, ax = plt.subplots(figsize=(11, 8.5))
        ax.text(0.5, 0.5, f"No data available for\n{campaign_name}\n"
                          f"between {window_start} and {window_end}",
                ha="center", va="center", fontsize=12, transform=ax.transAxes)
        ax.set_axis_off()
        buf = io.BytesIO()
        fig.savefig(buf, format="png", dpi=100, bbox_inches="tight")
        plt.close(fig)
        return buf.getvalue()

    df = weekly_df.copy()
    df["week_dt"] = pd.to_datetime(df["week_start"])
    bar_width = 6  # days; leaves a 1-day gap between weekly bars

    fig, (ax1, ax2, ax3, ax4) = plt.subplots(
        4, 1, figsize=(11, 9.5), sharex=True,
        gridspec_kw={"hspace": 0.18},
    )

    # ── Panel 1: weekly conversions ────────────────────────────────
    ax1.bar(df["week_dt"], df["conv"], width=bar_width,
            color="#4F81BD", alpha=0.75, edgecolor="#1F4E79", linewidth=0.5)
    # Annotate each bar with its value
    for _, row in df.iterrows():
        if row["conv"] > 0:
            ax1.annotate(f"{int(round(row['conv']))}",
                         xy=(row["week_dt"], row["conv"]),
                         xytext=(0, 3), textcoords="offset points",
                         ha="center", fontsize=8, color="#1F4E79")
    ax1.set_ylabel("Weekly conversions", fontsize=10)
    ax1.set_title(f"{campaign_name}\n12-week performance: {window_start} to {window_end}",
                  fontsize=11, fontweight="bold", loc="left")
    ax1.grid(True, alpha=0.3, axis="y")

    # ── Panel 2: weekly CPA ────────────────────────────────────────
    ax2.plot(df["week_dt"], df["cpa"], marker="o", color="#C0504D",
             linewidth=2, markersize=7, label="Weekly CPA")
    if current_tcpa is not None:
        ax2.axhline(current_tcpa, color="#1F4E79", linestyle="-.", linewidth=1.5,
                    alpha=0.7, label=f"Current tCPA target (${current_tcpa:.0f})")
    # Annotate CPA values
    for _, row in df.iterrows():
        if pd.notna(row["cpa"]):
            ax2.annotate(f"${row['cpa']:.0f}",
                         xy=(row["week_dt"], row["cpa"]),
                         xytext=(0, 5), textcoords="offset points",
                         ha="center", fontsize=8, color="#C0504D")
    ax2.set_ylabel("Weekly CPA ($)", fontsize=10)
    ax2.legend(loc="best", fontsize=8, framealpha=0.9)
    ax2.grid(True, alpha=0.3)

    # ── Panel 3: weekly RPL ────────────────────────────────────────
    ax3.plot(df["week_dt"], df["rpl"], marker="s", color="#9BBB59",
             linewidth=2, markersize=7, label="Weekly RPL")
    for _, row in df.iterrows():
        if pd.notna(row["rpl"]):
            ax3.annotate(f"${row['rpl']:.0f}",
                         xy=(row["week_dt"], row["rpl"]),
                         xytext=(0, 5), textcoords="offset points",
                         ha="center", fontsize=8, color="#5F7F2F")
    ax3.set_ylabel("Weekly RPL ($)", fontsize=10)
    ax3.legend(loc="best", fontsize=8, framealpha=0.9)
    ax3.grid(True, alpha=0.3)

    # ── Panel 4: weekly profit ─────────────────────────────────────
    profit_colors = ["#9BBB59" if p >= 0 else "#C0504D" for p in df["profit"]]
    ax4.bar(df["week_dt"], df["profit"], width=bar_width,
            color=profit_colors, alpha=0.75, edgecolor="#404040", linewidth=0.5)
    ax4.axhline(0, color="black", linewidth=0.8, alpha=0.5)
    for _, row in df.iterrows():
        if pd.notna(row["profit"]):
            offset = 8 if row["profit"] >= 0 else -12
            ax4.annotate(f"${row['profit']:,.0f}",
                         xy=(row["week_dt"], row["profit"]),
                         xytext=(0, offset), textcoords="offset points",
                         ha="center", fontsize=8,
                         color="#5F7F2F" if row["profit"] >= 0 else "#8B0000")
    ax4.set_ylabel("Weekly profit ($)", fontsize=10)
    ax4.set_xlabel("Week starting (Mon)", fontsize=10)
    ax4.grid(True, alpha=0.3, axis="y")

    # Overlay change-event markers on all panels
    for ev in change_events:
        change_dt = pd.Timestamp(ev["change_date"])
        # Choose color by attribute type
        if "TargetCPA" in str(ev.get("attribute", "")):
            color, ls = "red", "--"
        elif "Budget" in str(ev.get("attribute", "")):
            color, ls = "#FF8C00", ":"  # orange dotted for budget changes
        else:
            color, ls = "gray", "--"
        for ax in (ax1, ax2, ax3, ax4):
            ax.axvline(change_dt, color=color, linestyle=ls, linewidth=1.2, alpha=0.6)

    # Format x-axis
    ax4.xaxis.set_major_locator(mdates.WeekdayLocator(byweekday=0))  # Mondays
    ax4.xaxis.set_major_formatter(mdates.DateFormatter("%b %d"))
    plt.setp(ax4.xaxis.get_majorticklabels(), rotation=30, ha="right")

    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=100, bbox_inches="tight")
    plt.close(fig)
    return buf.getvalue()


def build_overview_small_multiples_png(
    campaign_weekly: dict[str, pd.DataFrame],
    window_start: date,
    window_end: date,
) -> bytes:
    """One sheet showing all 12 campaigns' weekly profit as small multiples.

    Layout: 4 columns × 3 rows. Each panel is a single campaign's weekly
    profit bar chart. Y-axis is independent per panel (profit magnitudes
    differ by 50× across campaigns; shared y-axis would make small
    campaigns invisible). Each panel labels its total 12-week profit.
    """
    n_campaigns = len(campaign_weekly)
    if n_campaigns == 0:
        fig, ax = plt.subplots(figsize=(11, 8.5))
        ax.text(0.5, 0.5, "No campaign data available",
                ha="center", va="center", fontsize=12, transform=ax.transAxes)
        ax.set_axis_off()
        buf = io.BytesIO()
        fig.savefig(buf, format="png", dpi=100, bbox_inches="tight")
        plt.close(fig)
        return buf.getvalue()

    # Compute grid dimensions
    n_cols = 4
    n_rows = max(1, (n_campaigns + n_cols - 1) // n_cols)

    fig, axes = plt.subplots(
        n_rows, n_cols, figsize=(16, 3.2 * n_rows),
        sharex=True, gridspec_kw={"hspace": 0.55, "wspace": 0.25},
    )
    axes = axes.flatten() if n_rows * n_cols > 1 else [axes]

    # Sort campaigns by total 12-week profit descending so biggest are top-left
    sorted_items = sorted(
        campaign_weekly.items(),
        key=lambda kv: -(kv[1]["profit"].sum() if not kv[1].empty else 0),
    )

    for idx, (campaign_name, weekly_df) in enumerate(sorted_items):
        ax = axes[idx]
        if weekly_df.empty:
            ax.text(0.5, 0.5, "(no data)", ha="center", va="center",
                    transform=ax.transAxes, fontsize=9, color="gray")
            ax.set_axis_off()
            continue

        weeks = pd.to_datetime(weekly_df["week_start"])
        profits = weekly_df["profit"].values
        colors = ["#9BBB59" if p >= 0 else "#C0504D" for p in profits]
        ax.bar(weeks, profits, width=6, color=colors, alpha=0.8,
               edgecolor="#404040", linewidth=0.4)
        ax.axhline(0, color="black", linewidth=0.5, alpha=0.5)

        # Title: short campaign name + total profit
        short_name = campaign_name.replace("HOME INSURANCE-", "").split("-SEARCH")[0]
        total_profit = profits.sum()
        ax.set_title(f"{short_name}\n${total_profit:,.0f} total",
                     fontsize=9, fontweight="bold", loc="left")
        ax.tick_params(axis="x", labelsize=7)
        ax.tick_params(axis="y", labelsize=7)
        ax.xaxis.set_major_locator(mdates.MonthLocator())
        ax.xaxis.set_major_formatter(mdates.DateFormatter("%b"))
        ax.grid(True, alpha=0.25, axis="y")
        # Format y-axis as currency
        ax.yaxis.set_major_formatter(plt.FuncFormatter(lambda v, p: f"${v:,.0f}"))

    # Hide unused subplots
    for idx in range(len(sorted_items), len(axes)):
        axes[idx].set_axis_off()

    fig.suptitle(
        f"Weekly Profit by Campaign — {window_start} to {window_end} "
        f"(sorted by total 12-week profit)",
        fontsize=13, fontweight="bold", y=0.995,
    )

    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=100, bbox_inches="tight")
    plt.close(fig)
    return buf.getvalue()


# ═══════════════════════════════════════════════════════════════════
# MAIN
# ═══════════════════════════════════════════════════════════════════

def parse_args():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--db-path", default=str(DEFAULT_DB_PATH))
    p.add_argument("--change-history", default=str(DEFAULT_CHANGE_CSV))
    p.add_argument("--output-dir", default=str(DEFAULT_OUTPUT_DIR))
    p.add_argument("--end-date", default=DEFAULT_END_DATE.isoformat(),
                   help=f"Last day of analysis window. Default: "
                        f"{DEFAULT_END_DATE.isoformat()}")
    p.add_argument("--weeks", type=int, default=DEFAULT_WEEKS,
                   help=f"Number of weeks back from --end-date. Default: "
                        f"{DEFAULT_WEEKS}")
    return p.parse_args()


def main():
    args = parse_args()
    db_path = Path(args.db_path)
    change_csv = Path(args.change_history)
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    end_date = date.fromisoformat(args.end_date)
    n_weeks = args.weeks
    # Compute window start: Monday of the week N weeks before end_date's week
    end_week_monday = end_date - timedelta(days=end_date.weekday())
    window_start = end_week_monday - timedelta(weeks=n_weeks - 1)
    window_end = end_date

    banner("CAMPAIGN PERFORMANCE VISUALIZATIONS — Microsoft Home")
    print(f"  DB:                  {db_path}")
    print(f"  Change history:      {change_csv}")
    print(f"  Output:              {output_dir}")
    print(f"  Window:              {window_start} (Mon) to {window_end}")
    print(f"  Weeks:               {n_weeks}")
    print()

    if not db_path.exists():
        print(f"ERROR: DB not found at {db_path}")
        sys.exit(1)

    conn = sqlite3.connect(db_path)
    enabled = fetch_enabled_campaigns(conn)
    print(f"Found {len(enabled)} enabled campaigns")
    print()

    # Load change events filtered to the window
    change_events_df = load_change_events(change_csv, window_start, window_end)
    print(f"Loaded {len(change_events_df)} change events within window "
          f"({(change_events_df['attribute'] == 'TargetCPA').sum()} tCPA, "
          f"{change_events_df['attribute'].str.contains('Budget', na=False).sum()} budget)")
    print()

    # Process each campaign: fetch daily, aggregate weekly
    campaign_weekly: dict[str, pd.DataFrame] = {}
    for _, camp in enabled.iterrows():
        cname = camp["campaign_name"]
        daily = fetch_campaign_daily(conn, cname, window_start, window_end)
        weekly = aggregate_weekly(daily)
        campaign_weekly[cname] = weekly

    # Build workbook
    print("Building workbook...")
    wb = Workbook()
    wb.remove(wb.active)

    # ── Overview sheet (small multiples) ──
    ws = wb.create_sheet("Overview")
    ws.cell(row=1, column=1, value="WEEKLY PROFIT — ALL CAMPAIGNS").font = Font(bold=True, size=14)
    ws.cell(row=2, column=1, value=f"Window: {window_start} to {window_end} "
                                    f"({n_weeks} weeks). Sorted by total 12-week profit.")
    print("  Generating small-multiples overview...")
    overview_png = build_overview_small_multiples_png(
        campaign_weekly, window_start, window_end,
    )
    img = XLImage(io.BytesIO(overview_png))
    ws.add_image(img, "A4")
    print("  ✓ Overview")

    # ── Per-campaign 4-panel sheets ──
    sheet_count = 0
    print("  Generating per-campaign 4-panel charts...")
    for _, camp in enabled.iterrows():
        cname = camp["campaign_name"]
        current_tcpa = camp["current_tcpa"]
        if pd.isna(current_tcpa):
            current_tcpa = None
        weekly_df = campaign_weekly.get(cname, pd.DataFrame())
        # Filter change events to this campaign
        camp_events = change_events_df[
            change_events_df["campaign_name"] == cname
        ].to_dict("records")

        # Sheet name: short campaign name (Excel limit 31 chars)
        short_name = cname.replace("HOME INSURANCE-", "").split("-SEARCH")[0][:25]
        sheet_name = short_name[:31]
        # Avoid duplicate sheet names
        if sheet_name in [s.title for s in wb.worksheets]:
            sheet_name = f"{short_name[:28]}_2"[:31]

        try:
            png_bytes = build_4panel_campaign_chart_png(
                weekly_df=weekly_df,
                campaign_name=cname,
                current_tcpa=current_tcpa,
                change_events=camp_events,
                window_start=window_start,
                window_end=window_end,
            )
            ws_chart = wb.create_sheet(sheet_name)
            img = XLImage(io.BytesIO(png_bytes))
            ws_chart.add_image(img, "A1")
            sheet_count += 1
            print(f"  ✓ {sheet_name}")
        except Exception as exc:
            print(f"  ✗ {sheet_name}: chart failed ({exc})")

    conn.close()
    print(f"Generated {sheet_count} per-campaign chart sheets")

    out_path = output_dir / f"campaign_performance_microsoft_{date.today().isoformat()}.xlsx"
    wb.save(out_path)
    print()
    print(f"  ✓ Wrote {out_path}")


if __name__ == "__main__":
    main()
