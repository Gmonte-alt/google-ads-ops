#!/usr/bin/env python3
r"""
ben_meeting_workbook.py -- Ben meeting workbook (Google Home, account_id=1)

Builds ben_meeting_workbook_<date>.xlsx: one sheet per question, chart-first,
VALUES ONLY (no formulas), ordered by narrative sequence Q4 -> Q1 -> Q2 -> Q3
(cause -> constraint -> remaining gap -> path back).

Inputs (all read from --tracker-dir, graceful skip where marked optional):
  REQUIRED
    march_vs_now_output.txt                      (redirect of march_vs_now_diagnostic.py)
    decomp_campaign_rolling7d_2026-06-13.csv     (cliff week, campaign lens)
    decomp_match_type_rolling7d_2026-06-13.csv
    decomp_category_rolling7d_2026-06-13.csv
  OPTIONAL
    decomp_campaign_month_2026-04-30.csv         (Mar->Apr bridge)  \
    decomp_campaign_month_2026-05-31.csv         (Apr->May bridge)   } Q3 driver chart
    decomp_campaign_month_2026-06-30.csv         (May->Jun bridge)  /
    decomp_campaign_rolling19d_2026-07-19.csv    (Jun->Jul bridge)
    decomp_category_match_month_2026-03-31.csv   \  Q2 cell CvR cross-check
    decomp_category_match_rolling19d_2026-07-19.csv /
    --funnel-csv <path>  weekly funnel CSV with columns:
        week, click_to_zip, zip_to_lead   (rates as % or fractions)  -> Q2 funnel chart
    --weekly-csv <path>  Ben's weekly table CSV with columns incl.:
        week, roas  (or: week, revenue, cost)                        -> Q4 drift chart

Usage:
  python ben_meeting_workbook.py --tracker-dir . [--output ben_meeting_workbook_2026-07-20.xlsx]
"""

import argparse
import re
import sys
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference, Series
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.drawing.fill import PatternFillProperties  # noqa: F401 (openpyxl quirk import)
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------- constants
CLIFF_STAMP = "rolling7d_2026-06-13"
RSA_TOUCHED_PAT = re.compile(r"Tier2|CALIFORNIA", re.I)  # RSA swap: Tier2 6/5, CA 6/8

# ROAS drift fallback (used only when --weekly-csv absent).
# Source: Ben's weekly Jan-Jul table / KT 2026-07-20 sec.4 ("ROAS slides
# 1.32 (5/23) -> 1.17 -> 1.13 (6/6) with NO changes 5/15-6/4").
ROAS_DRIFT_FALLBACK = [("wk 5/23", 1.32), ("wk 5/30", 1.17), ("wk 6/6", 1.13)]

CHANGE_TIMELINE = [
    ("5/6",      "Targets TIGHTENED (Tier2 113->120 etc.)"),
    ("5/14",     "Partial easing (-3pts; net Tier2 +4)"),
    ("5/15-6/4", "NO bid/budget/target/status changes (change history)"),
    ("6/5",      "RSA challenger swap, Tier2 (conservative: champions kept, ~20% imp share)"),
    ("6/8",      "RSA challenger swap, CA (same conservative design)"),
    ("6/7-13",   "THE CLIFF: account clicks -41%, cost -48% WoW"),
]

F_TITLE = Font(name="Arial", size=14, bold=True, color="1F3864")
F_SUB   = Font(name="Arial", size=10, italic=True, color="595959")
F_HDR   = Font(name="Arial", size=10, bold=True, color="FFFFFF")
F_BODY  = Font(name="Arial", size=10)
F_BOLD  = Font(name="Arial", size=10, bold=True)
F_NOTE  = Font(name="Arial", size=9, italic=True, color="7F7F7F")
F_ANS   = Font(name="Arial", size=10, bold=True, color="1F3864")
FILL_HDR    = PatternFill("solid", fgColor="1F3864")
FILL_ANSWER = PatternFill("solid", fgColor="DEEAF6")
FILL_TOUCH  = PatternFill("solid", fgColor="C55A11")   # RSA-touched accent
FILL_UNTOUCH= PatternFill("solid", fgColor="8496B0")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

C_TOUCHED   = "C55A11"  # orange
C_UNTOUCHED = "8496B0"  # slate
C_START_END = "808080"  # waterfall endpoints
C_UP        = "70AD47"
C_DOWN      = "C00000"
C_RPL       = "2E75B6"
C_CPC       = "ED7D31"
C_ROAS      = "C00000"
C_GP        = "2E75B6"
DRIVER_COLORS = {"value": "2E75B6", "efficiency": "70AD47",
                 "volume": "ED7D31", "interaction": "A6A6A6"}


# ---------------------------------------------------------------- parsing
def parse_march_vs_now(path: Path) -> dict:
    """Parse the three sections of march_vs_now_diagnostic.py terminal output."""
    txt = path.read_text(encoding="utf-8", errors="replace")
    out = {}

    # --- section 1: monthly table
    rows = []
    for m in re.finditer(
        r"^(\d{4}-\d{2})\s+(\d+)\s+([\d.]+)\s+([\d.]+)\s+([\d.]+)\s+([\d.]+)"
        r"\s+([\d.]+)\s+([\d.]+)\s+([\d.]+)\s*$", txt, re.M):
        month, days, clicks_d, cpc, cvr, leads_d, rpl, cpl, gp_d = m.groups()
        rows.append(dict(month=month, days=int(days), clicks_d=float(clicks_d),
                         cpc=float(cpc), cvr=float(cvr) / 100.0,
                         leads_d=float(leads_d), rpl=float(rpl),
                         cpl=float(cpl), gp_d=float(gp_d)))
    monthly = pd.DataFrame(rows)
    if monthly.empty:
        raise ValueError(f"could not parse monthly series from {path}")
    imp = re.search(r"impressions/day:\s*(\{[^}]*\})", txt)
    imps = eval(imp.group(1)) if imp else {}  # trusted local diagnostic output
    monthly["imps_d"] = monthly["month"].map(imps)
    monthly["rpc"] = monthly["rpl"] * monthly["cvr"]          # revenue per click
    monthly["roas_auction"] = monthly["rpc"] / monthly["cpc"]  # per-auction ROAS
    monthly["roas_realized"] = (monthly["leads_d"] * monthly["rpl"]) / \
                               (monthly["clicks_d"] * monthly["cpc"])
    out["monthly"] = monthly

    # --- section 2: mix-fixed RPL
    def grab(label):
        m = re.search(label + r"\s*:\s*([\d.]+)", txt)
        return float(m.group(1)) if m else None
    out["rpl_baseline"] = grab(r"baseline measured RPL")
    out["rpl_measured"] = grab(r"current\s+measured RPL")
    out["rpl_mixfixed"] = grab(r"current\s+MIX-FIXED RPL")

    # --- section 3: shift-share
    m = re.search(r"CvR:\s*([\d.]+)%\s*->\s*([\d.]+)%", txt)
    out["cvr_from"], out["cvr_to"] = float(m.group(1)), float(m.group(2))
    out["cvr_mix"] = float(re.search(r"MIX\s+component[^:]*:\s*([+-][\d.]+)pts", txt).group(1))
    out["cvr_rate"] = float(re.search(r"RATE component[^:]*:\s*([+-][\d.]+)pts", txt).group(1))
    cells = []
    for m in re.finditer(
        r"^\s{4}(\S+)\s+x\s+(\S+)\s+CvR\s+([\d.]+)%\s*->\s*([\d.]+)%"
        r"\s+contribution\s+([+-][\d.]+)pts\s*$", txt, re.M):
        cat, match, c0, c1, contrib = m.groups()
        cells.append(dict(cell=f"{cat} x {match}", cvr_from=float(c0),
                          cvr_to=float(c1), contribution=float(contrib)))
    out["cells"] = pd.DataFrame(cells)
    return out


def load_decomp(path: Path) -> pd.DataFrame:
    d = pd.read_csv(path)
    d["clicks_prev"] = d["leads_prev"] / d["cvr_prev"]
    d["clicks_cur"] = d["leads_cur"] / d["cvr_cur"]
    d["clicks_pct"] = d["clicks_cur"] / d["clicks_prev"] - 1.0
    d["cost_pct"] = d["cost_cur"] / d["cost_prev"] - 1.0
    d["cpc_pct"] = d["cpc_cur"] / d["cpc_prev"] - 1.0
    return d


def short_campaign(name: str) -> str:
    s = name.replace("HOME INSURANCE", "HI")
    s = re.sub(r"_COMBINED_TROAS_SEARCH_ALLDEVICE_USA-STE", "", s)
    s = re.sub(r"-NB-SEARCH-ALL-USA-STE", "", s)
    s = s.replace("HI_STATE-GROUP_", "").replace("zz-RETIRED-HI-", "zz-")
    s = s.replace("Tier1RPL80Plus", "Tier1 (80+)")
    s = s.replace("Tier2RPL70Plus", "Tier2 (70+)")
    s = s.replace("Tier3RPL50Plus", "Tier3 (50+)")
    return s


# ---------------------------------------------------------------- sheet helpers
def sheet_title(ws, row, question, answer, source):
    ws.cell(row=row, column=1, value=question).font = F_TITLE
    row += 1
    c = ws.cell(row=row, column=1, value="ANSWER: " + answer)
    c.font = F_ANS
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[row].height = 46
    for col in range(1, 13):
        ws.cell(row=row, column=col).fill = FILL_ANSWER
    row += 1
    ws.cell(row=row, column=1, value="Source: " + source).font = F_NOTE
    return row + 2


def write_table(ws, row, col, headers, data_rows, fmts=None, widths=None):
    """Write a values-only table; returns (first_data_row, last_row)."""
    for j, h in enumerate(headers):
        c = ws.cell(row=row, column=col + j, value=h)
        c.font, c.fill, c.border = F_HDR, FILL_HDR, BORDER
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    first = row + 1
    for i, r in enumerate(data_rows):
        for j, v in enumerate(r):
            c = ws.cell(row=first + i, column=col + j, value=v)
            c.font, c.border = F_BODY, BORDER
            if fmts and fmts[j]:
                c.number_format = fmts[j]
    if widths:
        for j, w in enumerate(widths):
            L = get_column_letter(col + j)
            ws.column_dimensions[L].width = max(ws.column_dimensions[L].width or 0, w)
    return first, first + len(data_rows) - 1


def notes_block(ws, row, col, lines, title=None):
    if title:
        ws.cell(row=row, column=col, value=title).font = F_BOLD
        row += 1
    for ln in lines:
        c = ws.cell(row=row, column=col, value=ln)
        c.font = F_NOTE if ln.startswith(("Source", "Note", "Caveat")) else F_BODY
        c.alignment = Alignment(wrap_text=False)
        row += 1
    return row + 1


def style_chart(ch, title, x_title=None, y_title=None, height=8.5, width=17):
    ch.title = title
    ch.height, ch.width = height, width
    ch.style = 2
    if x_title:
        ch.x_axis.title = x_title
    if y_title:
        ch.y_axis.title = y_title
    ch.y_axis.majorGridlines = ChartLines()
    return ch


def color_series(ser, hexcolor, line=False):
    if line:
        ser.graphicalProperties.line.solidFill = hexcolor
        ser.graphicalProperties.line.width = 22000
        ser.smooth = False
    else:
        ser.graphicalProperties.solidFill = hexcolor
        ser.graphicalProperties.line.noFill = True


# ---------------------------------------------------------------- Q4 sheet
def build_q4(wb, tracker: Path, weekly_csv: Path | None, manifest: list):
    ws = wb.create_sheet("Q4 Cliff Attribution")
    camp = load_decomp(tracker / f"decomp_campaign_{CLIFF_STAMP}.csv")
    mt = load_decomp(tracker / f"decomp_match_type_{CLIFF_STAMP}.csv")
    cat = load_decomp(tracker / f"decomp_category_{CLIFF_STAMP}.csv")
    camp["touched"] = camp["campaign"].str.contains(RSA_TOUCHED_PAT)
    camp["short"] = camp["campaign"].map(short_campaign)
    camp = camp.sort_values("clicks_pct")  # deepest decline first

    t = camp[camp.touched]
    u = camp[~camp.touched]
    t_pct = t.clicks_cur.sum() / t.clicks_prev.sum() - 1
    u_pct = u.clicks_cur.sum() / u.clicks_prev.sum() - 1
    acct_pct = camp.clicks_cur.sum() / camp.clicks_prev.sum() - 1
    cost_pct = camp.cost_cur.sum() / camp.cost_prev.sum() - 1

    row = sheet_title(
        ws, 1,
        "Q4 - What did we do in the two weeks before the volume drop?",
        f"Nothing on bidding (no target/budget/status changes 5/15-6/4). The RSA test is exculpated: "
        f"untouched campaigns fell {u_pct:+.1%} vs RSA-touched {t_pct:+.1%} - the swapped cohort "
        f"dropped LEAST. The cliff (account clicks {acct_pct:+.0%}, cost {cost_pct:+.0%} WoW) was the "
        f"tROAS controller autonomously correcting sub-target drift, shedding value-ranked inventory.",
        f"Change_history_report__14_.csv; decomp_campaign/match_type/category_{CLIFF_STAMP}.csv; "
        f"Ben's weekly table (ROAS drift)")

    # -- campaign table
    hdr = ["Campaign", "RSA-touched", "Clicks (prior 7d)", "Clicks (cliff 7d)",
           "Clicks %ch", "Cost %ch", "CPC prior", "CPC cliff", "CPC %ch", "dGP ($)"]
    data = [[r.short, "TOUCHED" if r.touched else "untouched",
             round(r.clicks_prev), round(r.clicks_cur), r.clicks_pct, r.cost_pct,
             r.cpc_prev, r.cpc_cur, r.cpc_pct, r.delta_gp]
            for r in camp.itertuples()]
    data.append(["TOUCHED cohort (Tier2 + CA)", "", round(t.clicks_prev.sum()),
                 round(t.clicks_cur.sum()), t_pct, t.cost_cur.sum() / t.cost_prev.sum() - 1,
                 None, None, None, t.delta_gp.sum()])
    data.append(["UNTOUCHED cohort (TX/FL/T1/T3/legacy)", "", round(u.clicks_prev.sum()),
                 round(u.clicks_cur.sum()), u_pct, u.cost_cur.sum() / u.cost_prev.sum() - 1,
                 None, None, None, u.delta_gp.sum()])
    fmts = [None, None, "#,##0", "#,##0", "+0.0%;-0.0%", "+0.0%;-0.0%",
            "$0.00", "$0.00", "+0.0%;-0.0%", "$#,##0;($#,##0)"]
    first, last = write_table(ws, row, 1, hdr, data, fmts,
                              widths=[34, 12, 13, 13, 10, 10, 9, 9, 9, 11])
    for rr in (last - 1, last):
        for cc in range(1, 11):
            ws.cell(row=rr, column=cc).font = F_BOLD
    for i, r in enumerate(camp.itertuples()):
        ws.cell(row=first + i, column=2).fill = FILL_TOUCH if r.touched else FILL_UNTOUCH
        ws.cell(row=first + i, column=2).font = Font(name="Arial", size=9, bold=True, color="FFFFFF")

    # -- killer chart: horizontal bars, per-point color touched vs untouched
    ch = BarChart()
    ch.type = "bar"
    ch.gapWidth = 60
    n = len(camp)
    ref = Reference(ws, min_col=5, min_row=first, max_row=first + n - 1)
    cats = Reference(ws, min_col=1, min_row=first, max_row=first + n - 1)
    ser = Series(ref, title="Clicks % change (cliff wk vs prior wk)")
    ch.series.append(ser)
    ch.set_categories(cats)
    for i, r in enumerate(camp.itertuples()):
        dp = DataPoint(idx=i)
        dp.graphicalProperties.solidFill = C_TOUCHED if r.touched else C_UNTOUCHED
        ser.data_points.append(dp)
    ch.dLbls = DataLabelList()
    ch.dLbls.showVal = True
    ch.dLbls.numFmt = "+0%;-0%"
    ch.legend = None
    style_chart(ch, "Cliff-week click decline by campaign - RSA-touched (orange) fell LEAST",
                y_title="Clicks % change", height=10.5, width=19)
    ws.add_chart(ch, f"A{last + 3}")
    chart_bottom = last + 3 + 22

    # -- ROAS drift mini-chart (weekly csv if given, else KT fallback points)
    drift_rows, drift_src = None, None
    if weekly_csv and weekly_csv.exists():
        wk = pd.read_csv(weekly_csv)
        wk.columns = [c.strip().lower() for c in wk.columns]
        if "roas" not in wk.columns and {"revenue", "cost"} <= set(wk.columns):
            wk["roas"] = wk["revenue"] / wk["cost"]
        if {"week", "roas"} <= set(wk.columns):
            drift_rows = list(wk[["week", "roas"]].itertuples(index=False, name=None))
            drift_src = f"--weekly-csv {weekly_csv.name}"
            manifest.append(("Q4", weekly_csv.name, "included (ROAS drift)"))
    if drift_rows is None:
        drift_rows = ROAS_DRIFT_FALLBACK
        drift_src = "Ben's weekly Jan-Jul table via KT 2026-07-20 (3 anchor weeks; pass --weekly-csv for full series)"
        manifest.append(("Q4", "weekly ROAS series", "FALLBACK - 3 KT anchor points (no --weekly-csv)"))
    r2 = chart_bottom
    ws.cell(row=r2, column=1, value="ROAS drift into the cliff (untouched period)").font = F_BOLD
    dfirst, dlast = write_table(ws, r2 + 1, 1, ["Week", "ROAS"],
                                [[w, v] for w, v in drift_rows],
                                [None, "0.00"], widths=[12, 8])
    ws.cell(row=dlast + 1, column=1, value="Source: " + drift_src).font = F_NOTE
    lc = LineChart()
    d = Reference(ws, min_col=2, min_row=r2 + 1, max_row=dlast)
    lc.add_data(d, titles_from_data=True)
    lc.set_categories(Reference(ws, min_col=1, min_row=dfirst, max_row=dlast))
    color_series(lc.series[0], C_ROAS, line=True)
    lc.series[0].marker.symbol = "circle"
    style_chart(lc, "ROAS drift 5/23-6/6 (no changes made 5/15-6/4)",
                y_title="ROAS", height=6.5, width=9)
    lc.legend = None
    ws.add_chart(lc, f"D{r2}")

    # -- change timeline
    r3 = dlast + 3
    r3 = notes_block(ws, r3, 1,
                     [f"  {d}:  {desc}" for d, desc in CHANGE_TIMELINE],
                     title="The only changes made (change history):")

    # -- value-ranked retreat evidence: match type + category declines
    ws.cell(row=r3, column=1, value="Value-ranked retreat (controller signature): "
            "EXACT/affinity held, generic/phrase shed hardest; CPC flat-to-down = 'we withdrew, not outbid'").font = F_BOLD
    r3 += 1
    mt_rows = [[r.match_type, round(r.clicks_prev), round(r.clicks_cur), r.clicks_pct, r.cpc_pct]
               for r in mt.sort_values("clicks_prev", ascending=False).itertuples()]
    write_table(ws, r3, 1, ["Match type", "Clicks prior", "Clicks cliff", "Clicks %ch", "CPC %ch"],
                mt_rows, [None, "#,##0", "#,##0", "+0.0%;-0.0%", "+0.0%;-0.0%"],
                widths=[16, 11, 11, 10, 9])
    key = cat.columns[0]
    cat_rows = [[getattr(r, key), round(r.clicks_prev), round(r.clicks_cur), r.clicks_pct, r.cpc_pct]
                for r in cat.sort_values("clicks_prev", ascending=False).itertuples()]
    write_table(ws, r3, 7, ["Category", "Clicks prior", "Clicks cliff", "Clicks %ch", "CPC %ch"],
                cat_rows, [None, "#,##0", "#,##0", "+0.0%;-0.0%", "+0.0%;-0.0%"],
                widths=[20, 11, 11, 10, 9])
    manifest.append(("Q4", f"decomp_campaign/match_type/category_{CLIFF_STAMP}.csv", "included"))
    return ws


# ---------------------------------------------------------------- Q1 sheet
def build_q1(wb, mvn: dict, manifest: list):
    ws = wb.create_sheet("Q1 Volume vs RPL")
    m = mvn["monthly"]
    jan = m.iloc[0]
    cur = m.iloc[-1]
    row = sheet_title(
        ws, 1,
        "Q1 - Clicks/impressions lowest all year but RPL is 2x prior lows. Why can't we get more volume?",
        f"The bidder buys per-auction economics, not RPL. Jan: RPL ${jan.rpl:.0f}, RPC ${jan.rpc:.2f} / "
        f"CPC ${jan.cpc:.2f} = {jan.roas_auction:.2f} per-auction ROAS. Now: RPL ${cur.rpl:.0f}, RPC "
        f"${cur.rpc:.2f} / ${cur.cpc:.2f} = {cur.roas_auction:.2f}. RPL ~2x but CPC "
        f"{cur.cpc / jan.cpc - 1:+.0%} and CvR {100 * (cur.cvr - jan.cvr):+.1f}pts - the marginal auction is "
        f"WORSE than January despite 2x the RPL (and January was pre-Geico).",
        "march_vs_now_diagnostic.py output, sections 1-2; Ben's weekly Jan-Jul table")

    hdr = ["Month", "RPL", "CvR", "RPC (RPL x CvR)", "CPC",
           "Per-auction ROAS (RPC/CPC)", "Clicks/day", "Imps/day", "GP/day"]
    data = [[r.month, r.rpl, r.cvr, r.rpc, r.cpc, r.roas_auction,
             r.clicks_d, r.imps_d, r.gp_d] for r in m.itertuples()]
    fmts = ["@", "$0.00", "0.0%", "$0.00", "$0.00", "0.00", "#,##0", "#,##0", "$#,##0"]
    first, last = write_table(ws, row, 1, hdr, data, fmts,
                              widths=[10, 9, 8, 14, 9, 20, 10, 10, 9])

    # chart: RPL + CPC primary ($), per-auction ROAS secondary
    ch = LineChart()
    ch.add_data(Reference(ws, min_col=2, min_row=row, max_row=last), titles_from_data=True)  # RPL
    ch.add_data(Reference(ws, min_col=5, min_row=row, max_row=last), titles_from_data=True)  # CPC
    ch.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
    color_series(ch.series[0], C_RPL, line=True)
    color_series(ch.series[1], C_CPC, line=True)
    ch2 = LineChart()
    ch2.add_data(Reference(ws, min_col=6, min_row=row, max_row=last), titles_from_data=True)  # ROAS
    color_series(ch2.series[0], C_ROAS, line=True)
    ch2.series[0].marker.symbol = "circle"
    ch2.y_axis.axId = 200
    ch2.y_axis.title = "Per-auction ROAS"
    ch2.y_axis.crosses = "max"
    style_chart(ch, "RPL doubled - but per-auction ROAS flattened toward 1.0 (that IS the answer)",
                y_title="$ per lead / $ per click")
    ch += ch2
    ws.add_chart(ch, f"A{last + 3}")
    r2 = last + 3 + 19

    # mix-fixed RPL nuance
    ws.cell(row=r2, column=1, value="Nuance vs March: the RPL uplift is REAL payout, not survivorship").font = F_BOLD
    mf_rows = [
        ["Baseline measured RPL (2026-03)", mvn["rpl_baseline"], None],
        ["Current measured RPL", mvn["rpl_measured"],
         mvn["rpl_measured"] / mvn["rpl_baseline"] - 1],
        ["Current MIX-FIXED RPL (today's payouts @ March mix)", mvn["rpl_mixfixed"],
         mvn["rpl_mixfixed"] / mvn["rpl_baseline"] - 1],
    ]
    write_table(ws, r2 + 1, 1, ["Measure", "RPL", "vs baseline"],
                mf_rows, [None, "$0.00", "+0%;-0%"], widths=[46, 9, 11])
    notes_block(ws, r2 + 6, 1, [
        "Mix-fixed > measured: today's mix is slightly ADVERSE vs March - payouts are genuinely ~10% richer.",
        "Implication: deliberate re-expansion buys volume on BETTER unit economics than March - once the funnel (Q2) is addressed.",
        "Caveat: mix-fixed sections cover surviving placements only (removed-placement baseline traffic excluded),",
        "which makes the 'payout-real' reading conservative.",
    ])
    manifest.append(("Q1", "march_vs_now_output.txt sec 1-2", "included"))
    return ws


# ---------------------------------------------------------------- Q2 sheet
def build_q2(wb, mvn: dict, tracker: Path, funnel_csv: Path | None, manifest: list):
    ws = wb.create_sheet("Q2 CvR Mix vs Rate")
    c0, c1 = mvn["cvr_from"], mvn["cvr_to"]
    mix, rate = mvn["cvr_mix"], mvn["cvr_rate"]
    row = sheet_title(
        ws, 1,
        "Q2 - CR has fallen. Is it keyword mix?",
        f"No. Shift-share: CvR {c0:.1f}% -> {c1:.1f}%, of which RATE {rate:+.1f}pts vs MIX {mix:+.1f}pts "
        f"(mix slightly HELPED). The rate decline concentrates in carrier EXACT cells - fixed queries "
        f"converting 6-9pts worse - triangulating with the zip->lead funnel drop (46%->38%, "
        f"match-type-agnostic). Downstream monetization, not keyword selection. This is the evidence "
        f"base for the gclid/partner-data escalation.",
        "march_vs_now output sec 3; decomp_category_match month 03-31 vs rolling19d 07-19; Looker Traffic export")

    # -- waterfall via stacked bars: [invisible base, visible delta]
    labels = [f"March CvR", f"Mix {mix:+.1f}", f"Rate {rate:+.1f}", "July CvR"]
    vals = [c0, mix, rate, c1]
    bases, heights, colors = [], [], []
    running = 0.0
    for i, (lab, v) in enumerate(zip(labels, vals)):
        if i in (0, len(vals) - 1):
            bases.append(0.0); heights.append(v); colors.append(C_START_END)
        else:
            new = running + v
            bases.append(min(running, new)); heights.append(abs(v))
            colors.append(C_UP if v >= 0 else C_DOWN)
            running = new
            continue
        running = v if i == 0 else running
    wf_first, wf_last = write_table(ws, row, 1, ["Step", "base", "delta"],
                                    [[l, b, h] for l, b, h in zip(labels, bases, heights)],
                                    [None, "0.0", "0.0"], widths=[14, 8, 8])
    ws.cell(row=wf_last + 1, column=1,
            value="Note: base/delta columns are chart plumbing for the waterfall.").font = F_NOTE
    ch = BarChart()
    ch.type = "col"
    ch.grouping = "stacked"
    ch.overlap = 100
    ch.gapWidth = 50
    base_ref = Reference(ws, min_col=2, min_row=row, max_row=wf_last)
    delt_ref = Reference(ws, min_col=3, min_row=row, max_row=wf_last)
    ch.add_data(base_ref, titles_from_data=True)
    ch.add_data(delt_ref, titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=wf_first, max_row=wf_last))
    ch.series[0].graphicalProperties.noFill = True
    for i, col_hex in enumerate(colors):
        dp = DataPoint(idx=i)
        dp.graphicalProperties.solidFill = col_hex
        ch.series[1].data_points.append(dp)
    ch.legend = None
    style_chart(ch, f"CvR bridge: {c0:.1f} -> {c1:.1f} - it's RATE, not MIX",
                y_title="CvR (pts)", height=8, width=10.5)
    ws.add_chart(ch, f"E{row - 1}")

    # -- cell contributions
    r2 = max(wf_last + 4, row + 16)
    cells = mvn["cells"]
    ws.cell(row=r2, column=1, value="Largest within-cell CvR declines (click-weighted)").font = F_BOLD
    cfirst, clast = write_table(
        ws, r2 + 1, 1, ["Cell (category x match)", "CvR March", "CvR July", "d pts", "Contribution (pts)"],
        [[r.cell, r.cvr_from / 100, r.cvr_to / 100, (r.cvr_to - r.cvr_from) / 100, r.contribution]
         for r in cells.itertuples()],
        [None, "0.0%", "0.0%", "+0.0%;-0.0%", "+0.0;-0.0"], widths=[30, 10, 10, 9, 15])
    ch2 = BarChart()
    ch2.type = "bar"
    ref = Reference(ws, min_col=5, min_row=cfirst, max_row=clast)
    ser = Series(ref, title="Contribution to CvR change (pts)")
    ch2.series.append(ser)
    ch2.set_categories(Reference(ws, min_col=1, min_row=cfirst, max_row=clast))
    color_series(ser, C_DOWN)
    ch2.legend = None
    style_chart(ch2, "Where the rate decline lives: carrier EXACT cells lead",
                y_title="pts", height=8, width=12)
    ws.add_chart(ch2, f"G{r2}")
    r3 = clast + 3

    # -- optional cell CvR cross-check from decomp files
    m03 = tracker / "decomp_category_match_month_2026-03-31.csv"
    r19 = tracker / "decomp_category_match_rolling19d_2026-07-19.csv"
    if not r19.exists():  # pre-patch stamp fallback
        legacy = tracker / "decomp_category_match_rolling_2026-07-19.csv"
        r19 = legacy if legacy.exists() else r19
    if m03.exists() and r19.exists():
        a = pd.read_csv(m03)
        b = pd.read_csv(r19)
        kc = [c for c in a.columns if "categ" in c][0]
        km = "match_type"
        a["cell"] = a[kc] + " x " + a[km]
        b["cell"] = b[kc] + " x " + b[km]
        j = a.set_index("cell")[["cvr_cur", "leads_cur"]].join(
            b.set_index("cell")[["cvr_cur"]], lsuffix="_mar", rsuffix="_jul", how="inner")
        j = j.sort_values("leads_cur", ascending=False).head(10)
        ws.cell(row=r3, column=1, value="Cross-check: cell CvRs straight from decomp files "
                "(March month-end vs July 19d)").font = F_BOLD
        write_table(ws, r3 + 1, 1, ["Cell", "CvR March", "CvR July"],
                    [[i, r.cvr_cur_mar, r.cvr_cur_jul] for i, r in j.iterrows()],
                    [None, "0.0%", "0.0%"], widths=[30, 10, 10])
        r3 += len(j) + 4
        manifest.append(("Q2", f"{m03.name} + {r19.name}", "included (cross-check)"))
    else:
        ws.cell(row=r3, column=1, value="(Cell-CvR cross-check skipped - decomp_category_match "
                "month_2026-03-31 / rolling19d_2026-07-19 not found in tracker dir.)").font = F_NOTE
        r3 += 2
        manifest.append(("Q2", "category_match cross-check files", "SKIPPED - not found"))

    # -- optional funnel chart
    if funnel_csv and funnel_csv.exists():
        f = pd.read_csv(funnel_csv)
        f.columns = [c.strip().lower() for c in f.columns]
        need = {"week", "click_to_zip", "zip_to_lead"}
        if need <= set(f.columns):
            for c in ("click_to_zip", "zip_to_lead"):
                if f[c].max() > 1.5:
                    f[c] = f[c] / 100.0
            ws.cell(row=r3, column=1, value="Funnel: click->zip vs zip->lead (weekly)").font = F_BOLD
            ffirst, flast = write_table(ws, r3 + 1, 1, ["Week", "Click->Zip", "Zip->Lead"],
                                        [[r.week, r.click_to_zip, r.zip_to_lead] for r in f.itertuples()],
                                        [None, "0.0%", "0.0%"], widths=[12, 11, 11])
            lc = LineChart()
            lc.add_data(Reference(ws, min_col=2, min_row=r3 + 1, max_row=flast), titles_from_data=True)
            lc.add_data(Reference(ws, min_col=3, min_row=r3 + 1, max_row=flast), titles_from_data=True)
            lc.set_categories(Reference(ws, min_col=1, min_row=ffirst, max_row=flast))
            color_series(lc.series[0], C_RPL, line=True)
            color_series(lc.series[1], C_DOWN, line=True)
            style_chart(lc, "The funnel story: zip->lead falls while click->zip recovers",
                        y_title="rate", height=8, width=13)
            ws.add_chart(lc, f"E{r3}")
            manifest.append(("Q2", funnel_csv.name, "included (funnel chart)"))
        else:
            ws.cell(row=r3, column=1, value=f"(Funnel CSV {funnel_csv.name} missing columns "
                    f"{need} - chart skipped.)").font = F_NOTE
            manifest.append(("Q2", funnel_csv.name, "SKIPPED - wrong columns"))
    else:
        ws.cell(row=r3, column=1, value="(Funnel chart skipped - pass --funnel-csv with the Looker "
                "Traffic export: week, click_to_zip, zip_to_lead. Known anchors: zip->lead 46% -> 38%.)").font = F_NOTE
        manifest.append(("Q2", "Looker funnel export", "SKIPPED - no --funnel-csv"))
    manifest.append(("Q2", "march_vs_now_output.txt sec 3", "included"))
    return ws


# ---------------------------------------------------------------- Q3 sheet
def build_q3(wb, mvn: dict, tracker: Path, manifest: list):
    ws = wb.create_sheet("Q3 March vs Now")
    m = mvn["monthly"]
    mar = m[m.month == "2026-03"].iloc[0]
    apr = m[m.month == "2026-04"].iloc[0]
    cur = m.iloc[-1]
    row = sheet_title(
        ws, 1,
        "Q3 - What were we doing in March that worked? What's different now?",
        f"March was a POSITION, not a playbook - volume at ~{mar.roas_realized:.2f} ROAS "
        f"(${mar.gp_d:,.0f} GP/day); APRIL was the actual peak (${apr.gp_d:,.0f} GP/day, "
        f"{apr.roas_realized:.2f} ROAS). Since then: May drift + 5/6 tightening, the June 7-13 "
        f"controller cliff, the 6/18 bid-down, July's CvR break. Recoverable: target position and "
        f"the funnel. Not recoverable by bidding: demand (imps/day {mar.imps_d:,.0f} -> "
        f"{cur.imps_d:,.0f}) and pre-cut payout levels.",
        "chained decomp_campaign_month_*.csv + category twins; march_vs_now monthly series")

    # -- monthly GP/day bars + realized ROAS line
    hdr = ["Month", "GP/day", "Realized ROAS", "Imps/day", "Leads/day"]
    data = [[r.month, r.gp_d, r.roas_realized, r.imps_d, r.leads_d] for r in m.itertuples()]
    first, last = write_table(ws, row, 1, hdr, data,
                              ["@", "$#,##0", "0.00", "#,##0", "#,##0"],
                              widths=[10, 10, 12, 10, 10])
    bc = BarChart()
    bc.type = "col"
    bc.gapWidth = 60
    bc.add_data(Reference(ws, min_col=2, min_row=row, max_row=last), titles_from_data=True)
    bc.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
    color_series(bc.series[0], C_GP)
    lc = LineChart()
    lc.add_data(Reference(ws, min_col=3, min_row=row, max_row=last), titles_from_data=True)
    color_series(lc.series[0], C_ROAS, line=True)
    lc.series[0].marker.symbol = "circle"
    lc.y_axis.axId = 200
    lc.y_axis.title = "Realized ROAS"
    lc.y_axis.crosses = "max"
    style_chart(bc, "March = volume at ~1.0 ROAS; April = the real sweet spot",
                y_title="GP / day ($)")
    bc += lc
    ws.add_chart(bc, f"A{last + 3}")
    r2 = last + 3 + 19

    # -- MoM driver bridges from chained campaign files
    bridges = [
        ("Mar->Apr", [tracker / "decomp_campaign_month_2026-04-30.csv"]),
        ("Apr->May", [tracker / "decomp_campaign_month_2026-05-31.csv"]),
        ("May->Jun", [tracker / "decomp_campaign_month_2026-06-30.csv"]),
        # 19d bridge may exist under the pre-patch stamp (no window-days in name)
        ("Jun->Jul (19d)", [tracker / "decomp_campaign_rolling19d_2026-07-19.csv",
                            tracker / "decomp_campaign_rolling_2026-07-19.csv"]),
    ]
    rows, found, missing = [], [], []
    for label, candidates in bridges:
        path = next((p for p in candidates if p.exists()), candidates[0])
        if path.exists():
            d = pd.read_csv(path)
            rows.append([label, d["value"].sum(), d["efficiency"].sum(),
                         d["volume"].sum(), d["interaction"].sum(), d["delta_gp"].sum()])
            found.append(path.name)
        else:
            missing.append(path.name)
    if rows:
        ws.cell(row=r2, column=1, value="Which driver broke in which month (account totals, campaign lens)").font = F_BOLD
        bfirst, blast = write_table(
            ws, r2 + 1, 1, ["Bridge", "value", "efficiency", "volume", "interaction", "dGP"],
            rows, [None] + ["$#,##0;($#,##0)"] * 5, widths=[14, 10, 10, 10, 10, 10])
        ch = BarChart()
        ch.type = "col"
        ch.grouping = "clustered"
        ch.gapWidth = 80
        ch.add_data(Reference(ws, min_col=2, max_col=5, min_row=r2 + 1, max_row=blast),
                    titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=bfirst, max_row=blast))
        for ser, name in zip(ch.series, ["value", "efficiency", "volume", "interaction"]):
            color_series(ser, DRIVER_COLORS[name])
        style_chart(ch, "MoM GP driver decomposition - which driver broke when",
                    y_title="dGP ($)", height=9, width=17)
        ws.add_chart(ch, f"A{blast + 2}")
        r2 = blast + 2 + 20
        manifest.append(("Q3", ", ".join(found), "included (driver bridges)"))
    if missing:
        ws.cell(row=r2, column=1,
                value="(Driver-bridge files not found, panels skipped: " + ", ".join(missing) + ")").font = F_NOTE
        r2 += 2
        manifest.append(("Q3", ", ".join(missing), "SKIPPED - not found"))

    notes_block(ws, r2, 1, [
        "Recoverable: target position (the easing decision) and the funnel (Q2).",
        "Not recoverable by bidding: demand (impressions/day) and pre-6/18 payout levels.",
        "Sequencing: July 22 re-read first (NC/CO promotions, tier1/tier3 band fixes), tier2 easing after.",
    ], title="The path back")
    manifest.append(("Q3", "march_vs_now_output.txt sec 1", "included"))
    return ws


# ---------------------------------------------------------------- appendix: cat x match detail
def _cell_metrics(df: pd.DataFrame, keycols) -> pd.DataFrame:
    """Full metric set for the CURRENT period of a decomp file, per cell."""
    import numpy as np
    d = df.copy()
    d["clicks"] = np.where(d["cvr_cur"] > 0, d["leads_cur"] / d["cvr_cur"],
                           np.where(d["cpc_cur"] > 0, d["cost_cur"] / d["cpc_cur"], 0.0))
    d["revenue"] = d["leads_cur"] * d["rpl_cur"]
    d["gp"] = d["gp_cur"]
    d["roas"] = d["revenue"] / d["cost_cur"]
    out = d[keycols + ["clicks", "leads_cur", "revenue", "cost_cur", "gp",
                       "cpc_cur", "cvr_cur", "cpl_cur", "rpl_cur", "roas"]]
    return out.rename(columns={"leads_cur": "leads", "cost_cur": "cost", "cpc_cur": "cpc",
                               "cvr_cur": "cvr", "cpl_cur": "cpl", "rpl_cur": "rpl"})


def _rollup(m: pd.DataFrame, by) -> pd.DataFrame:
    g = m.groupby(by, as_index=False)[["clicks", "leads", "revenue", "cost", "gp"]].sum()
    g["cpc"] = g["cost"] / g["clicks"]
    g["cvr"] = g["leads"] / g["clicks"]
    g["cpl"] = g["cost"] / g["leads"]
    g["rpl"] = g["revenue"] / g["leads"]
    g["roas"] = g["revenue"] / g["cost"]
    return g


METRICS = [  # (column, header, number format)
    ("clicks", "Clicks", "#,##0"), ("leads", "Leads", "#,##0"),
    ("revenue", "Revenue", "$#,##0"), ("cost", "Cost", "$#,##0"),
    ("gp", "GP", "$#,##0;($#,##0)"), ("cpc", "CPC", "$0.00"),
    ("cvr", "CvR", "0.0%"), ("cpl", "CPL", "$0.00"),
    ("rpl", "RPL", "$0.00"), ("roas", "ROAS", "0.00"),
]


def _paired_table(ws, row, label_hdr, labels, mar: pd.DataFrame, jul: pd.DataFrame, keyfn, metrics=None):
    """Interleaved March/July metric pairs + delta block. Returns next free row."""
    metrics = metrics or METRICS
    mar_i = mar.set_index(mar.apply(keyfn, axis=1))
    jul_i = jul.set_index(jul.apply(keyfn, axis=1))
    hdr = [label_hdr]
    fmts = [None]
    for col, h, f in metrics:
        hdr += [f"{h} Mar", f"{h} Jul"]
        fmts += [f, f]
    hdr += ["dClicks %", "dLeads %", "dGP $", "dCvR pts", "dCPL %", "dRPL %"]
    fmts += ["+0%;-0%", "+0%;-0%", "$#,##0;($#,##0)", "+0.0;-0.0", "+0%;-0%", "+0%;-0%"]
    def fin(x):
        try:
            import math
            return None if x is None or not math.isfinite(float(x)) else float(x)
        except (TypeError, ValueError):
            return x
    data = []
    for lab in labels:
        a = mar_i.loc[lab] if lab in mar_i.index else None
        b = jul_i.loc[lab] if lab in jul_i.index else None
        r = [lab]
        for col, _, _ in metrics:
            r += [None if a is None else a[col], None if b is None else b[col]]
        if a is not None and b is not None:
            r += [b.clicks / a.clicks - 1 if a.clicks else None,
                  b.leads / a.leads - 1 if a.leads else None,
                  b.gp - a.gp, 100 * (b.cvr - a.cvr),
                  b.cpl / a.cpl - 1 if a.cpl else None,
                  b.rpl / a.rpl - 1 if a.rpl else None]
        else:
            r += [None] * 6
        data.append([fin(x) if i > 0 else x for i, x in enumerate(r)])
    widths = [26] + [9, 9] * len(metrics) + [9, 9, 10, 9, 8, 8]
    first, last = write_table(ws, row, 1, hdr, data, fmts, widths)
    return last + 2, first, last


def build_catmatch_appendix(wb, tracker: Path, manifest: list):
    m03 = tracker / "decomp_category_match_month_2026-03-31.csv"
    r19 = tracker / "decomp_category_match_rolling19d_2026-07-19.csv"
    if not r19.exists():
        legacy = tracker / "decomp_category_match_rolling_2026-07-19.csv"
        r19 = legacy if legacy.exists() else r19
    if not (m03.exists() and r19.exists()):
        manifest.append(("Appendix", "category_match March/July files", "SKIPPED - not found"))
        return
    ws = wb.create_sheet("Appendix Cat x Match")
    ws.freeze_panes = "B1"   # keep row labels visible across the wide grid
    keys = ["keyword_category", "match_type"]
    mar = _cell_metrics(pd.read_csv(m03), keys)
    jul = _cell_metrics(pd.read_csv(r19), keys)
    for d in (mar, jul):
        d["cell"] = d["keyword_category"] + " x " + d["match_type"]

    ws.cell(row=1, column=1, value="Appendix - Category x Match detail (March 2026 vs July 1-19)").font = F_TITLE
    ws.cell(row=2, column=1, value="All metrics per period: Clicks, Leads (conversions), Revenue, Cost, GP, "
            "CPC, CvR, CPL, RPL, ROAS. Periods = 'current' side of decomp_category_match month_2026-03-31 "
            "(March) and rolling19d_2026-07-19 (July 1-19). July is 19 days vs March's 31 - compare rates "
            "and unit economics directly; compare volumes/$ with the day-count in mind.").font = F_SUB
    ws.cell(row=3, column=1, value="Note: impressions are not carried in the decomp lens (would need a "
            "readout extension against keyword_performance_daily); revenue = leads x RPL "
            "(primary_conv_value basis); ROAS = revenue/cost.").font = F_NOTE
    row = 5

    # (A) match type rollup
    ws.cell(row=row, column=1, value="A. By MATCH TYPE (account rollup)").font = F_BOLD
    mt_mar, mt_jul = _rollup(mar, "match_type"), _rollup(jul, "match_type")
    order = [m for m in ["BROAD", "EXACT", "PHRASE"] if m in set(mt_mar.match_type) | set(mt_jul.match_type)]
    row, _, _ = _paired_table(ws, row + 1, "Match type", order, mt_mar, mt_jul,
                              lambda r: r["match_type"])

    # (B) category rollup
    ws.cell(row=row, column=1, value="B. By CATEGORY (account rollup)").font = F_BOLD
    c_mar, c_jul = _rollup(mar, "keyword_category"), _rollup(jul, "keyword_category")
    cats = list(c_jul.sort_values("cost", ascending=False).keyword_category)
    cats += [c for c in c_mar.keyword_category if c not in cats]
    row, _, _ = _paired_table(ws, row + 1, "Category", cats, c_mar, c_jul,
                              lambda r: r["keyword_category"])

    # (C) full cell grid
    ws.cell(row=row, column=1, value="C. CATEGORY x MATCH cells (sorted by July cost)").font = F_BOLD
    cells = list(jul.sort_values("cost", ascending=False).cell)
    cells += [c for c in mar.cell if c not in cells]
    row, cfirst, clast = _paired_table(ws, row + 1, "Cell", cells, mar, jul, lambda r: r["cell"])

    # chart: GP March vs July by cell (top cells by July cost)
    ch = BarChart()
    ch.type = "bar"
    ch.grouping = "clustered"
    ch.gapWidth = 60
    n = min(len(cells), len(jul))
    gp_mar_col = 1 + 1 + (4 * 2)      # label + pairs before GP: 4 metrics -> GP Mar at col 10
    ch.add_data(Reference(ws, min_col=gp_mar_col, min_row=cfirst - 1, max_row=cfirst - 1 + n), titles_from_data=True)
    ch.add_data(Reference(ws, min_col=gp_mar_col + 1, min_row=cfirst - 1, max_row=cfirst - 1 + n), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=cfirst, max_row=cfirst - 1 + n))
    color_series(ch.series[0], "8496B0")
    color_series(ch.series[1], C_RPL)
    style_chart(ch, "GP by cell - March vs July 1-19", y_title="GP ($)", height=11, width=15)
    ws.add_chart(ch, f"A{row + 1}")
    manifest.append(("Appendix", f"{m03.name} + {r19.name}", "included (cat x match detail)"))


# ---------------------------------------------------------------- appendix: brand key (DB-backed)
BRAND_PERIODS = [("Mar", "2026-03-01", "2026-03-31"),
                 ("Jul", "2026-07-01", "2026-07-19")]
TREND_WEEKS = 8                     # trailing 7-day buckets ending at Jul end
COSTCO_NOTE = ("Costco ad group switched off DKI headlines wk of 7/13 "
               "(per Gilbert) -- clicks/revenue collapse visible in the trend; "
               "distinct, self-inflicted, recoverable cause vs controller/funnel.")


def _brand_pull(conn, vcs, resolver, registry, start, end, has_clicks, has_imps):
    """Per-brand metric rollup for one period, fragmentation-correct join."""
    import pandas as pd
    p, ki, k = vcs.PERF_COLS, vcs.KI_COLS, vcs.KW_COLS
    extra = (f", SUM(perf.clicks) AS clicks" if has_clicks else "") + \
            (f", SUM(perf.impressions) AS imps" if has_imps else "")
    sql = f"""
        SELECT ki.{ki['placement_id']} AS placement_id,
               kw.{k['keyword_text']}  AS keyword_text,
               SUM(perf.{p['cost_micros']})/1e6 AS cost,
               SUM(perf.{p['revenue']})         AS revenue,
               SUM(perf.{p['leads']})           AS leads{extra}
        FROM {vcs.PERF_TABLE} perf
        JOIN {vcs.KI_TABLE} ki ON ki.{ki['id']} = perf.{p['instance_id']}
        JOIN {vcs.KW_VIEW} kw
          ON kw.{k['placement_id']} = ki.{ki['placement_id']}
         AND kw.{k['account_id']}   = ki.{ki['account_id']}
        WHERE ki.{ki['account_id']} = :acct
          AND perf.{p['device']}    = :device
          AND perf.{p['date']} BETWEEN :start AND :end
        GROUP BY ki.{ki['placement_id']}, kw.{k['keyword_text']}"""
    df = pd.read_sql_query(sql, conn, params={
        "acct": vcs.ACCOUNT_ID, "device": vcs.DEVICE_ROLLUP_VALUE,
        "start": start, "end": end})
    if df.empty:
        return df
    df["brand_key"] = df["keyword_text"].map(
        lambda t: resolver(t) if isinstance(t, str) else None)
    df = df.merge(registry, on="brand_key", how="left")
    df = df[df["brand_category"].notna()].copy()      # home carriers only
    df["size_group"] = df["brand_key"].map(vcs.CARRIER_SIZE_MAP).fillna("unmapped_size")
    agg = {"cost": "sum", "revenue": "sum", "leads": "sum"}
    if has_clicks:
        agg["clicks"] = "sum"
    if has_imps:
        agg["imps"] = "sum"
    b = df.groupby(["brand_key", "size_group"], as_index=False).agg(agg)
    b["gp"] = b["revenue"] - b["cost"]
    if has_clicks:
        b["cpc"] = b["cost"] / b["clicks"]
        b["cvr"] = b["leads"] / b["clicks"]
    else:
        b["clicks"] = None; b["cpc"] = None; b["cvr"] = None
    b["cpl"] = b["cost"] / b["leads"]
    b["rpl"] = b["revenue"] / b["leads"]
    b["roas"] = b["revenue"] / b["cost"]
    return b


def build_brand_appendix(wb, tracker: Path, args, manifest: list):
    if not args.db:
        manifest.append(("Brands", "paid_search.db", "SKIPPED - no --db"))
        return
    import sqlite3
    sys.path.insert(0, str(tracker))
    try:
        import validate_carrier_size_axis as vcs
    except ImportError as e:
        manifest.append(("Brands", "validate_carrier_size_axis.py",
                         f"SKIPPED - import failed ({e})"))
        return
    try:
        conn = sqlite3.connect(args.db)
        conn.execute("PRAGMA query_only = ON")
        # brand resolver + registry, reused from the validation script
        registry = pd.read_sql_query(
            f"""SELECT {vcs.REGISTRY_COLS['brand_key']} AS brand_key,
                       {vcs.REGISTRY_COLS['brand_category']} AS brand_category
                FROM {vcs.REGISTRY_TABLE}
                WHERE {vcs.REGISTRY_COLS['brand_category']} LIKE 'carrier_%'
                  AND {vcs.REGISTRY_COLS['verticals']} LIKE '%home%'""", conn)
        matcher = vcs.load_brand_matcher(args.brand_matcher_dir)
        if matcher is not None:
            resolver = vcs.build_resolver(matcher, args.db)
        elif args.allow_registry_fallback:
            resolver = lambda t: vcs.registry_fallback_match(t, registry, "brand_key")
        else:
            manifest.append(("Brands", "brand_matcher",
                             "SKIPPED - matcher unavailable (pass --brand-matcher-dir "
                             "or --allow-registry-fallback)"))
            return
        perf_cols = vcs.table_columns(conn, vcs.PERF_TABLE)
        has_clicks, has_imps = "clicks" in perf_cols, "impressions" in perf_cols

        periods = {}
        for lab, s, e in BRAND_PERIODS:
            periods[lab] = _brand_pull(conn, vcs, resolver, registry, s, e,
                                       has_clicks, has_imps)
        mar, jul = periods["Mar"], periods["Jul"]
        if mar.empty and jul.empty:
            manifest.append(("Brands", "DB pull", "SKIPPED - zero rows both periods"))
            return

        ws = wb.create_sheet("Appendix Brand Key")
        ws.freeze_panes = "B1"
        ws.cell(row=1, column=1, value="Appendix - Brand key detail (March 2026 vs July 1-19)").font = F_TITLE
        ws.cell(row=2, column=1, value="Brand resolution + carrier-size axis reused from "
                "validate_carrier_size_axis.py (framework v1.1, axis frozen 7/14). "
                "Fragmentation-correct join. July is 19 days vs March's 31.").font = F_SUB
        if not has_clicks:
            ws.cell(row=3, column=1, value="Note: clicks column not found in perf table -- "
                    "CPC/CvR omitted.").font = F_NOTE
        row = 4

        metrics = ([("imps", "Imps", "#,##0")] if has_imps else []) + METRICS

        # A. size-group rollup
        def rollup_groups(b):
            agg = {"cost": "sum", "revenue": "sum", "leads": "sum", "gp": "sum"}
            if has_clicks:
                agg["clicks"] = "sum"
            if has_imps:
                agg["imps"] = "sum"
            g = b.groupby("size_group", as_index=False).agg(agg)
            if has_clicks:
                g["cpc"] = g["cost"] / g["clicks"]
                g["cvr"] = g["leads"] / g["clicks"]
            else:
                g["clicks"] = None; g["cpc"] = None; g["cvr"] = None
            g["cpl"] = g["cost"] / g["leads"]
            g["rpl"] = g["revenue"] / g["leads"]
            g["roas"] = g["revenue"] / g["cost"]
            return g
        gm, gj = rollup_groups(mar), rollup_groups(jul)
        ws.cell(row=row, column=1, value="A. By CARRIER SIZE GROUP (frozen axis)").font = F_BOLD
        order = [g for g in ["membership_affinity", "carrier_large", "carrier_small",
                             "unmapped_size"] if g in set(gm.size_group) | set(gj.size_group)]
        row, _, _ = _paired_table(ws, row + 1, "Size group", order, gm, gj,
                                  lambda r: r["size_group"], metrics)

        # B. per-brand
        ws.cell(row=row, column=1, value="B. By BRAND (sorted by March revenue)").font = F_BOLD
        brands = list(mar.sort_values("revenue", ascending=False).brand_key)
        brands += [b for b in jul.brand_key if b not in brands]
        row, bfirst, blast = _paired_table(ws, row + 1, "Brand", brands, mar, jul,
                                           lambda r: r["brand_key"], metrics)

        # C. weekly trend (clicks + revenue) for top brands -- the Costco cliff
        import datetime as dt
        end = dt.date.fromisoformat(BRAND_PERIODS[-1][2])
        weeks = []
        for i in range(TREND_WEEKS - 1, -1, -1):
            we = end - dt.timedelta(days=7 * i)
            wsdt = we - dt.timedelta(days=6)
            weeks.append((f"{wsdt.month}/{wsdt.day}-{we.month}/{we.day}",
                          wsdt.isoformat(), we.isoformat()))
        top = list((jul if not jul.empty else mar)
                   .sort_values("revenue", ascending=False).brand_key.head(6))
        for must in ("Costco",):
            if must not in top and must in set(mar.brand_key) | set(jul.brand_key):
                top = [must] + top[:5]
        series = {lab: _brand_pull(conn, vcs, resolver, registry, s, e,
                                   has_clicks, has_imps).set_index("brand_key")
                  for lab, s, e in weeks}
        ws.cell(row=row, column=1, value="C. WEEKLY TREND - top brands (7-day buckets)").font = F_BOLD
        ws.cell(row=row + 1, column=1, value=COSTCO_NOTE).font = F_NOTE
        row += 2
        for metric, title, fmt in ((("clicks" if has_clicks else "leads"),
                                    "Clicks" if has_clicks else "Leads", "#,##0"),
                                   ("revenue", "Revenue", "$#,##0")):
            hdr = [f"{title} / week"] + [w[0] for w in weeks]
            data = []
            for b in top:
                vals = []
                for lab, _, _ in weeks:
                    d = series[lab]
                    vals.append(float(d.loc[b, metric]) if b in d.index else 0.0)
                data.append([b] + vals)
            first, last = write_table(ws, row, 1, hdr, data,
                                      [None] + [fmt] * len(weeks),
                                      [22] + [10] * len(weeks))
            ch = LineChart()
            for i in range(len(top)):
                ref = Reference(ws, min_col=2, max_col=1 + len(weeks),
                                min_row=first + i, max_row=first + i)
                s = Series(ref, title_from_data=False, title=top[i])
                ch.series.append(s)
            ch.set_categories(Reference(ws, min_col=2, max_col=1 + len(weeks),
                                        min_row=row, max_row=row))
            for s in ch.series:
                s.smooth = False
            if "Costco" in top:
                color_series(ch.series[top.index("Costco")], C_DOWN, line=True)
            style_chart(ch, f"Weekly {title.lower()} by brand - the Costco DKI cliff",
                        y_title=title, height=8.5, width=17)
            ws.add_chart(ch, f"A{last + 2}")
            row = last + 2 + 19
        manifest.append(("Brands", f"{Path(args.db).name} via validate_carrier_size_axis "
                         f"(clicks={'y' if has_clicks else 'n'}, imps={'y' if has_imps else 'n'})",
                         "included (brand appendix)"))
    except Exception as e:
        manifest.append(("Brands", "brand appendix", f"SKIPPED - error: {e}"))
    finally:
        try:
            conn.close()
        except Exception:
            pass


# ---------------------------------------------------------------- cover
def build_cover(wb, manifest, out_name):
    ws = wb.create_sheet("Cover", 0)
    ws.cell(row=1, column=1, value="Google Home - Ben meeting workbook").font = F_TITLE
    ws.cell(row=2, column=1, value=f"Account: Google Ads (account_id=1) - Home Insurance   |   "
            f"Built {date.today().isoformat()}   |   {out_name}").font = F_SUB
    ws.cell(row=4, column=1, value="Narrative order (sheets left to right):").font = F_BOLD
    order = [
        ("Q4 Cliff Attribution", "CAUSE - the June 7-13 volume cliff was the tROAS controller, not creative"),
        ("Q1 Volume vs RPL", "CONSTRAINT - the bidder buys per-auction economics; marginal auction worse than Jan despite 2x RPL"),
        ("Q2 CvR Mix vs Rate", "REMAINING GAP - the CR fall is rate-not-mix; downstream monetization/funnel"),
        ("Q3 March vs Now", "PATH BACK - March was volume at breakeven, April the peak; the decision is deliberate target position"),
    ]
    r = 5
    for name, desc in order:
        ws.cell(row=r, column=1, value=name).font = F_BOLD
        ws.cell(row=r, column=3, value=desc).font = F_BODY
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Input manifest").font = F_BOLD
    write_table(ws, r + 1, 1, ["Sheet", "Input", "Status"],
                [list(x) for x in manifest], widths=[8, 58, 44])
    r2 = r + 2 + len(manifest) + 1
    notes_block(ws, r2, 1, [
        "Values-only workbook (no formulas) by design - a static meeting artifact.",
        "Note: all numbers trace to paid_search.db via the tracker scripts; revenue truth = primary_conv_value (partner actuals).",
    ])
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["C"].width = 90


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description="Build the Ben meeting workbook (Q4->Q1->Q2->Q3).")
    ap.add_argument("--tracker-dir", default=".", help="folder holding the decomp CSVs + march_vs_now_output.txt")
    ap.add_argument("--output", default=f"ben_meeting_workbook_{date.today().isoformat()}.xlsx")
    ap.add_argument("--funnel-csv", default=None, help="weekly Looker funnel export (week, click_to_zip, zip_to_lead)")
    ap.add_argument("--weekly-csv", default=None, help="Ben's weekly table as CSV (week, roas | week, revenue, cost)")
    ap.add_argument("--db", default=None, help="paid_search.db path -- enables the Brand Key appendix")
    ap.add_argument("--brand-matcher-dir", default=None, help="folder containing brand_matcher.py")
    ap.add_argument("--allow-registry-fallback", action="store_true",
                    help="if brand_matcher unavailable, substring-match against brand_registry (DIRECTIONAL ONLY)")
    args = ap.parse_args()

    tracker = Path(args.tracker_dir)
    mvn_path = tracker / "march_vs_now_output.txt"
    required = [mvn_path] + [tracker / f"decomp_{l}_{CLIFF_STAMP}.csv"
                             for l in ("campaign", "match_type", "category")]
    missing = [p.name for p in required if not p.exists()]
    if missing:
        sys.exit("Missing required inputs in tracker dir: " + ", ".join(missing) +
                 "\n(march_vs_now_output.txt = redirect of march_vs_now_diagnostic.py; "
                 "cliff CSVs = patched readout, rolling7d, end-date 2026-06-13)")

    mvn = parse_march_vs_now(mvn_path)
    manifest = [("Q4/Q1/Q2/Q3", "paid_search.db via tracker scripts", "source of record")]
    wb = Workbook()
    wb.remove(wb.active)
    build_q4(wb, tracker, Path(args.weekly_csv) if args.weekly_csv else None, manifest)
    build_q1(wb, mvn, manifest)
    build_q2(wb, mvn, tracker, Path(args.funnel_csv) if args.funnel_csv else None, manifest)
    build_q3(wb, mvn, tracker, manifest)
    build_catmatch_appendix(wb, tracker, manifest)
    build_brand_appendix(wb, tracker, args, manifest)
    build_cover(wb, manifest, args.output)
    wb.save(args.output)
    print(f"wrote {args.output}")
    print("manifest:")
    for sheet, inp, status in manifest:
        print(f"  [{sheet}] {inp}: {status}")


if __name__ == "__main__":
    main()
