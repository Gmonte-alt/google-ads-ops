#!/usr/bin/env python
r"""
keyword_census.py — Every targeted keyword, judged in the open
═══════════════════════════════════════════════════════════════════════
Full census of the charter campaigns' keyword book — including
enabled-but-never-serving rows the performance-driven instruments
can't see — with impressions, clicks, CTR, spend, ROAS, and
impression-weighted QS across SIX lookback windows (7 / 14 / 30 / 60 /
90 / YTD), and a transparent PAUSE/KEEP verdict per row. Answers: how
many keywords should be turned off, which are BASE, which are
INCREMENTAL.

VERDICTS (rules printed on the workbook; every verdict carries its
reason string):
  KEEP-PROTECTED  Farmers home-intent — standing rule, never pause.
  PAUSE           sustained bad economics: 90d ROAS < 0.85 AND 30d
                  ROAS < 0.85 AND not improving (30d not >= 90d+0.15
                  and 30d < 0.95), with >= --min-verdict-spend in 90d.
                  Bad and not getting better.
  KEEP-BASE       90d ROAS >= 1.00 with >= --base-spend in 90d — the
                  book that carries the P&L.
  KEEP-INCREMENTAL healthy but small (90d ROAS >= 1.00 below the base
                  spend bar, or thin-spend earners) — room to feed.
  WATCH           borderline (0.85–1.00), improving-unhealthy, or too
                  thin to judge — rerun next cycle.
  ZOMBIE          enabled, zero impressions in 90d — serving-
                  eligibility dead. DELIBERATELY NOT a pause
                  recommendation: the Oct-2025 mass pause is a
                  documented ratchet tread (footprint contraction fed
                  the spiral); zombies cost nothing and pausing them
                  shrinks the footprint further. They are the
                  exposure-death exhibit, not cleanup.
  ALREADY-PAUSED  instance status is paused (when a status column
                  exists).

DATA (house): census base = keyword_instances in the charter
campaigns (LEFT JOIN performance, so zero-data rows appear);
kpd device='ALL'; grain keyword x MATCH TYPE x campaign (instances
merged, count shown); QS impression-weighted per window (never
summed, unrated excluded); revenue/conversions columns auto-detected;
Leads = Conversions column; break-even 1.00; mode=ro, never during a
sync. Segment marginal class joined from the latest map CSV.

HONESTY: windows overlap the expansion waves — a keyword's 30d can sit
mostly inside a wave (the health workbook's phase-conditioning covers
that); verdicts here are level economics. PAUSE list is a CANDIDATE
set for the cut-list workflow (dated changes, tripwire-watched), not
an auto-execution list.

OUTPUT: keyword_census_<date>.xlsx (Summary / Census / Notes),
keyword_census_<date>.csv (full grain, all six windows),
keyword_census_<date>.png (verdict mix per campaign).

USAGE:
  python keyword_census.py
  python keyword_census.py --min-verdict-spend 100 --base-spend 500
"""

import argparse
import csv
import sqlite3
import sys
from datetime import datetime, timedelta
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl required: pip install openpyxl")

DEFAULT_DB_PATH = Path(r"C:\Projects\google-ads-ops\db\paid_search.db")
DEFAULT_OUT_DIR = Path(r"C:\Projects\google-ads-ops\analysis"
                       r"\profit-attribution-tracker\output"
                       r"\keyword-health")
DEFAULT_MAP_DIR = Path(r"C:\Projects\google-ads-ops\analysis"
                       r"\profit-attribution-tracker\output"
                       r"\marginal-roas-map")
ACCOUNT_ID = 1
CAMPAIGN_IDS = {
    "CALIFORNIA": 23385188622,
    "TEXAS":      23394808297,
    "FLORIDA":    23394808300,
    "TIER4":      23428389081,
    "TIER1":      23496334752,
    "TIER2":      23530860942,
    "TIER3":      23586637678,
}
CARRIER_PATTERNS = [
    ("Geico", ["geico"]), ("Farmers", ["farmers"]),
    ("State Farm", ["state farm", "statefarm"]),
    ("Travelers", ["travelers"]),
    ("Kin", ["kin insurance", " kin ", "kin home"]),
    ("AAA", ["aaa"]), ("Amica", ["amica"]),
    ("Hartford", ["hartford"]), ("Costco", ["costco"]),
]
OTHER_LABEL = "Other/generic"
HOME_TOKENS = ("home", "homeowner", "homeowners", "house")
REV_CANDIDATES = ["conversion_value", "conversions_value", "conv_value",
                  "primary_conv_value", "total_conv_value",
                  "all_conv_value"]
CONV_CANDIDATES = ["conversions", "primary_conv_count", "conv_count",
                   "all_conversions"]
STATUS_CANDIDATES = ["status", "keyword_status", "kw_status"]
WINDOWS = [7, 14, 30, 60, 90]          # + YTD appended at runtime
SHEET_WINDOWS = ["14", "30", "90", "YTD"]   # census sheet subset

VERDICT_ORDER = ["PAUSE", "WATCH", "KEEP-INCREMENTAL", "KEEP-BASE",
                 "KEEP-PROTECTED", "ZOMBIE", "ALREADY-PAUSED"]
VERDICT_COLORS = {"PAUSE": "#C00000", "WATCH": "#BF8F00",
                  "KEEP-INCREMENTAL": "#7CBF7C",
                  "KEEP-BASE": "#2E7D32",
                  "KEEP-PROTECTED": "#44709D", "ZOMBIE": "#808080",
                  "ALREADY-PAUSED": "#BFBFBF"}
VERDICT_FILL = {"PAUSE": PatternFill("solid", fgColor="F8CBAD"),
                "WATCH": PatternFill("solid", fgColor="FFF2CC"),
                "KEEP-BASE": PatternFill("solid", fgColor="C6E0B4"),
                "KEEP-INCREMENTAL":
                    PatternFill("solid", fgColor="E2EFDA"),
                "KEEP-PROTECTED":
                    PatternFill("solid", fgColor="D9E2F3"),
                "ZOMBIE": PatternFill("solid", fgColor="EDEDED"),
                "ALREADY-PAUSED":
                    PatternFill("solid", fgColor="EDEDED")}
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=9)
SOFT_FONT = Font(italic=True, color="808080", size=9)


def connect_ro(p):
    conn = sqlite3.connect(f"file:{Path(p).as_posix()}?mode=ro",
                           uri=True)
    conn.row_factory = sqlite3.Row
    return conn


def pick_col(conn, table, candidates):
    try:
        cols = {r[1] for r in conn.execute(f"PRAGMA table_info({table})")}
    except sqlite3.Error:
        return None
    for c in candidates:
        if c in cols:
            return c
    return None


def classify(kw):
    k = f" {kw.lower().strip()} "
    for label, pats in CARRIER_PATTERNS:
        for p in pats:
            if p in k:
                return label
    return OTHER_LABEL


def is_protected(kw):
    return classify(kw) == "Farmers" and \
        any(t in kw for t in HOME_TOKENS)


def load_map_classes(map_csv, map_dir):
    p = Path(map_csv) if map_csv else None
    if not p:
        d = Path(map_dir)
        if d.exists():
            cands = sorted(d.glob("marginal_roas_map_*.csv"),
                           key=lambda x: x.stat().st_mtime,
                           reverse=True)
            p = cands[0] if cands else None
    if not p or not p.exists():
        return {}
    out = {}
    with open(p, newline="", encoding="utf-8-sig") as f:
        for r in csv.DictReader(f):
            out[(r.get("campaign", ""), r.get("segment", ""))] = \
                r.get("class", "")
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--db-path", default=str(DEFAULT_DB_PATH))
    ap.add_argument("--out-dir", default=str(DEFAULT_OUT_DIR))
    ap.add_argument("--map-csv", default=None)
    ap.add_argument("--map-dir", default=str(DEFAULT_MAP_DIR))
    ap.add_argument("--min-verdict-spend", type=float, default=100.0,
                    help="90d spend needed for a PAUSE verdict")
    ap.add_argument("--base-spend", type=float, default=500.0,
                    help="90d spend for KEEP-BASE vs KEEP-INCREMENTAL")
    ap.add_argument("--detail-rows", type=int, default=2000)
    ap.add_argument("--map", action="append",
                    help='override ids: --map "TIER2=123"')
    args = ap.parse_args()

    camp_ids = dict(CAMPAIGN_IDS)
    for m in args.map or []:
        k, v = m.split("=", 1)
        camp_ids[k.strip().upper()] = int(v.strip())
    ids = list(camp_ids.values())
    short_of = {cid: s for s, cid in camp_ids.items()}

    conn = connect_ro(args.db_path)
    rev_col = pick_col(conn, "keyword_performance_daily",
                       REV_CANDIDATES)
    conv_col = pick_col(conn, "keyword_performance_daily",
                        CONV_CANDIDATES)
    impr_col = pick_col(conn, "keyword_performance_daily",
                        ["impressions", "impr"])
    click_col = pick_col(conn, "keyword_performance_daily",
                         ["clicks", "click_count"])
    status_col = pick_col(conn, "keyword_instances",
                          STATUS_CANDIDATES)
    has_kqd = pick_col(conn, "keyword_quality_daily",
                       ["historical_quality_score"]) is not None
    has_cid = pick_col(conn, "keyword_instances",
                       ["campaign_id"]) is not None
    if not (impr_col and click_col and rev_col and has_cid):
        sys.exit("kpd needs impressions/clicks/revenue and "
                 "ki.campaign_id.")
    print(f"kpd: impr={impr_col} clicks={click_col} rev={rev_col} "
          f"conv={conv_col}; ki status: {status_col or 'NONE'}; "
          f"QS layer: {'yes' if has_kqd else 'NO'}")

    row = conn.execute(
        """SELECT MAX(p.report_date)
           FROM keyword_performance_daily p
           JOIN keyword_instances ki ON ki.id = p.instance_id
           WHERE ki.account_id = ?""", [ACCOUNT_ID]).fetchone()
    dend = datetime.strptime(str(row[0])[:10], "%Y-%m-%d").date()
    ytd_start = dend.replace(month=1, day=1)
    wins = [(str(w), (dend - timedelta(days=w - 1)).isoformat())
            for w in WINDOWS] + [("YTD", ytd_start.isoformat())]
    earliest = min(s for _, s in wins)
    print(f"Data through {dend}; windows: "
          + ", ".join(f"{n}={s}" for n, s in wins))

    # per-window conditional aggregates in one pass
    agg_terms = []
    for name, start in wins:
        c = f"p.report_date >= '{start}'"
        agg_terms += [
            f"SUM(CASE WHEN {c} THEN p.{impr_col} ELSE 0 END) "
            f"im_{name}",
            f"SUM(CASE WHEN {c} THEN p.{click_col} ELSE 0 END) "
            f"ck_{name}",
            f"SUM(CASE WHEN {c} THEN p.cost_micros ELSE 0 END)/1e6 "
            f"s_{name}",
            f"SUM(CASE WHEN {c} THEN p.{rev_col} ELSE 0 END) "
            f"r_{name}",
        ]
        if conv_col:
            agg_terms.append(
                f"SUM(CASE WHEN {c} THEN p.{conv_col} ELSE 0 END) "
                f"cv_{name}")
        if has_kqd:
            agg_terms += [
                f"SUM(CASE WHEN {c} AND q.historical_quality_score "
                f"IS NOT NULL THEN p.{impr_col} ELSE 0 END) "
                f"ri_{name}",
                f"SUM(CASE WHEN {c} AND q.historical_quality_score "
                f"IS NOT NULL THEN q.historical_quality_score * "
                f"p.{impr_col} ELSE 0 END) qi_{name}",
            ]
    q = (f"SELECT ki.id iid, ki.campaign_id cid, "
         f"LOWER(TRIM(ki.keyword_text)) kw, "
         f"UPPER(TRIM(ki.match_type)) mt, "
         + (f"ki.{status_col} st, " if status_col else "NULL st, ")
         + ", ".join(agg_terms)
         + f" FROM keyword_instances ki "
         f"LEFT JOIN keyword_performance_daily p "
         f"ON p.instance_id = ki.id AND p.device = 'ALL' "
         f"AND p.report_date >= '{earliest}' "
         + ("LEFT JOIN keyword_quality_daily q "
            "ON q.instance_id = p.instance_id "
            "AND q.report_date = p.report_date "
            if has_kqd else "")
         + f"WHERE ki.account_id = ? AND ki.campaign_id IN "
         f"({','.join('?' * len(ids))}) "
         f"GROUP BY 1, 2, 3, 4, 5")
    rows = conn.execute(q, [ACCOUNT_ID] + ids).fetchall()
    conn.close()
    print(f"  instance rows: {len(rows):,}")

    # merge to keyword x match x campaign
    merged = {}
    win_names = [n for n, _ in wins]
    for r in rows:
        short = short_of.get(int(r["cid"]))
        if not short:
            continue
        key = (short, str(r["mt"] or "?"), r["kw"] or "")
        m = merged.setdefault(key, {
            "iids": set(), "st": set(),
            **{f"{f}_{n}": 0.0 for n in win_names
               for f in ("im", "ck", "s", "r", "cv", "ri", "qi")}})
        m["iids"].add(int(r["iid"]))
        if r["st"] is not None:
            m["st"].add(str(r["st"]).upper())
        for n in win_names:
            for f in ("im", "ck", "s", "r", "cv", "ri", "qi"):
                col = f"{f}_{n}"
                try:
                    m[col] += r[col] or 0.0
                except (IndexError, KeyError):
                    pass

    map_classes = load_map_classes(args.map_csv, args.map_dir)

    def wmetrics(m, n):
        im, ck = m[f"im_{n}"], m[f"ck_{n}"]
        s, rv = m[f"s_{n}"], m[f"r_{n}"]
        ri, qi = m[f"ri_{n}"], m[f"qi_{n}"]
        return {"im": im, "ck": ck,
                "ctr": (ck / im) if im > 0 else None,
                "s": s, "r": rv,
                "roas": (rv / s) if s > 0 else None,
                "qs": (qi / ri) if ri > 0 else None,
                "cv": m[f"cv_{n}"]}

    out = []
    for (short, mt, kw), m in merged.items():
        W = {n: wmetrics(m, n) for n in win_names}
        prot = is_protected(kw)
        seg = classify(kw)
        enabled = True
        if m["st"]:
            enabled = any("ENAB" in s for s in m["st"])
        r90, r30 = W["90"]["roas"], W["30"]["roas"]
        s90 = W["90"]["s"]
        improving = (r30 is not None and
                     ((r90 is not None and r30 - r90 >= 0.15)
                      or r30 >= 0.95))
        if not enabled:
            v, why = "ALREADY-PAUSED", "instance status paused"
        elif prot:
            v = "KEEP-PROTECTED"
            why = "Farmers home-intent standing rule"
        elif W["90"]["im"] == 0:
            v = "ZOMBIE"
            why = ("no impressions 90d — serving-eligibility dead; "
                   "not a pause (Oct-25 footprint lesson)")
        elif s90 < args.min_verdict_spend:
            if r90 is not None and r90 >= 1.0:
                v, why = "KEEP-INCREMENTAL", \
                    f"thin but earning (r90 {r90:.2f}, ${s90:,.0f})"
            else:
                v, why = "WATCH", \
                    f"thin (${s90:,.0f} 90d) — below verdict floor"
        elif r90 is not None and r90 < 0.85 and \
                (r30 is None or r30 < 0.85) and not improving:
            v = "PAUSE"
            r30s = f"{r30:.2f}" if r30 is not None else "n/a"
            why = (f"r90 {r90:.2f} / r30 {r30s} — bad both "
                   f"windows, not improving (${s90:,.0f} 90d)")
        elif r90 is not None and r90 >= 1.0:
            if s90 >= args.base_spend:
                v, why = "KEEP-BASE", \
                    f"r90 {r90:.2f} at ${s90:,.0f} 90d"
            else:
                v, why = "KEEP-INCREMENTAL", \
                    f"r90 {r90:.2f}, room below base bar"
        elif r90 is not None and r90 >= 0.85:
            v, why = "WATCH", f"borderline r90 {r90:.2f}"
        else:
            v, why = "WATCH", \
                (f"unhealthy r90 {r90:.2f} but improving "
                 f"(r30 {r30:.2f})" if r30 is not None
                 else f"unhealthy r90 {r90:.2f}, r30 n/a")
        rec = {"campaign": short, "keyword": kw, "match_type": mt,
               "n_instances": len(m["iids"]),
               "segment": seg, "protected": prot,
               "seg_marginal_class":
                   map_classes.get((short, seg), ""),
               "status": "/".join(sorted(m["st"])) or "",
               "verdict": v, "reason": why}
        for n in win_names:
            w = W[n]
            rec[f"impr_{n}"] = w["im"]
            rec[f"clicks_{n}"] = w["ck"]
            rec[f"ctr_{n}"] = w["ctr"]
            rec[f"spend_{n}"] = w["s"]
            rec[f"rev_{n}"] = w["r"]
            rec[f"roas_{n}"] = w["roas"]
            rec[f"iwqs_{n}"] = w["qs"]
            rec[f"conv_{n}"] = w["cv"]
        out.append(rec)

    stamp = dend.isoformat()
    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    csv_path = out_dir / f"keyword_census_{stamp}.csv"
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(out[0].keys()))
        w.writeheader()
        w.writerows(sorted(out, key=lambda x: -x["spend_90"]))

    # summary per campaign + account
    shorts = sorted(camp_ids)
    L = [f"KEYWORD CENSUS — data through {stamp} — "
         f"{len(out):,} keyword x match x campaign rows",
         f"PAUSE needs >= ${args.min_verdict_spend:,.0f} 90d spend, "
         f"<0.85 in 90d AND 30d, not improving; KEEP-BASE >= "
         f"${args.base_spend:,.0f} 90d at >= 1.00; ZOMBIE = enabled, "
         "zero 90d impressions (NOT a pause — Oct-25 footprint "
         "lesson).", ""]
    sums = {}
    for scope in shorts + ["ACCOUNT"]:
        sel = out if scope == "ACCOUNT" else \
            [x for x in out if x["campaign"] == scope]
        line = {v: (sum(1 for x in sel if x["verdict"] == v),
                    sum(x["spend_30"] for x in sel
                        if x["verdict"] == v) / 30.0,
                    sum(x["rev_30"] - x["spend_30"] for x in sel
                        if x["verdict"] == v) / 30.0)
                for v in VERDICT_ORDER}
        sums[scope] = line
        parts = [f"{v} {n}" + (f" (${s:,.0f}/d)" if s >= 1 else "")
                 for v, (n, s, _) in line.items() if n]
        L.append(f"{scope:11s} " + " · ".join(parts))
    p_n = sum(1 for x in out if x["verdict"] == "PAUSE")
    p_s = sum(x["spend_30"] for x in out
              if x["verdict"] == "PAUSE") / 30.0
    p_gp = sum(x["rev_30"] - x["spend_30"] for x in out
               if x["verdict"] == "PAUSE") / 30.0
    b_n = sum(1 for x in out if x["verdict"] == "KEEP-BASE")
    i_n = sum(1 for x in out if x["verdict"] == "KEEP-INCREMENTAL")
    z_n = sum(1 for x in out if x["verdict"] == "ZOMBIE")
    L += ["",
          f"HEADLINE: turn off {p_n} keywords (${p_s:,.0f}/day "
          f"spend = 30d total/30, {p_gp:+,.0f}/day GP same basis) · "
          f"BASE {b_n} · "
          f"INCREMENTAL {i_n} · ZOMBIES {z_n} (exposure-death "
          "exhibit, keep).",
          "PAUSE list = candidates for the cut-list workflow (dated, "
          "tripwire-watched) — not auto-execution."]
    print("\n" + "\n".join(L))

    # chart
    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(12.5, 4.4), dpi=150)
    bot1 = {c: 0 for c in shorts}
    bot2 = {c: 0.0 for c in shorts}
    for v in VERDICT_ORDER:
        ns = [sums[c][v][0] for c in shorts]
        ss = [sums[c][v][1] for c in shorts]
        if any(ns):
            ax1.bar(shorts, ns, bottom=[bot1[c] for c in shorts],
                    color=VERDICT_COLORS[v], label=v, width=0.62)
            for c, n in zip(shorts, ns):
                bot1[c] += n
        if any(ss):
            ax2.bar(shorts, ss, bottom=[bot2[c] for c in shorts],
                    color=VERDICT_COLORS[v], width=0.62)
            for c, s in zip(shorts, ss):
                bot2[c] += s
    ax1.set_title("Keyword count by verdict", fontsize=9.5)
    ax2.set_title("Spend $/day (30d total / 30) by verdict", fontsize=9.5)
    for ax in (ax1, ax2):
        ax.tick_params(labelsize=7.5)
        ax.grid(axis="y", alpha=0.25)
    ax1.legend(fontsize=6.5)
    fig.suptitle(f"Keyword census — {stamp}", fontsize=10)
    fig.tight_layout(rect=(0, 0, 1, 0.93))
    png_path = out_dir / f"keyword_census_{stamp}.png"
    fig.savefig(png_path)
    plt.close(fig)

    # workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False
    ws.cell(row=1, column=1, value=f"Keyword census — data through "
            f"{stamp}").font = Font(bold=True, size=13)
    ws.cell(row=2, column=1, value=L[1]).font = SOFT_FONT
    hdr = ["Scope"] + VERDICT_ORDER
    for c, h in enumerate(hdr, 1):
        cell = ws.cell(row=4, column=c, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    r = 5
    for scope in shorts + ["ACCOUNT"]:
        ws.cell(row=r, column=1, value=scope).font = \
            Font(bold=(scope == "ACCOUNT"))
        for c, v in enumerate(VERDICT_ORDER, 2):
            n, s, _ = sums[scope][v]
            ws.cell(row=r, column=c,
                    value=f"{n}" + (f" (${s:,.0f}/d)" if s >= 1
                                    else ""))
        r += 1
    for line in L[len(shorts) + 4:]:
        r += 1
        ws.cell(row=r, column=1, value=line).font = SOFT_FONT
    for c in range(1, len(hdr) + 1):
        ws.column_dimensions[get_column_letter(c)].width = \
            18 if c > 1 else 12

    wc = wb.create_sheet("Census")
    wc.sheet_view.showGridLines = False
    base_hdr = ["Campaign", "Keyword", "Match", "Inst", "Segment",
                "Prot", "SegMarginal", "Status", "Verdict", "Reason"]
    win_hdr = []
    for n in SHEET_WINDOWS:
        win_hdr += [f"Impr {n}", f"CTR {n}", f"iwQS {n}",
                    f"ROAS {n}", f"Spend {n}"]
    for c, h in enumerate(base_hdr + win_hdr, 1):
        cell = wc.cell(row=1, column=c, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    detail = sorted(out, key=lambda x: -x["spend_90"])
    shown = detail[:args.detail_rows]
    for i, x in enumerate(shown, 2):
        wc.cell(row=i, column=1, value=x["campaign"])
        wc.cell(row=i, column=2, value=x["keyword"][:55])
        wc.cell(row=i, column=3, value=x["match_type"][:8])
        wc.cell(row=i, column=4, value=x["n_instances"])
        wc.cell(row=i, column=5, value=x["segment"])
        wc.cell(row=i, column=6, value="YES" if x["protected"]
                else "")
        wc.cell(row=i, column=7, value=x["seg_marginal_class"])
        wc.cell(row=i, column=8, value=x["status"][:14])
        vcell = wc.cell(row=i, column=9, value=x["verdict"])
        if x["verdict"] in VERDICT_FILL:
            vcell.fill = VERDICT_FILL[x["verdict"]]
        wc.cell(row=i, column=10, value=x["reason"][:60]).font = \
            SOFT_FONT
        ci = 11
        for n in SHEET_WINDOWS:
            wc.cell(row=i, column=ci,
                    value=x[f"impr_{n}"]).number_format = "#,##0"
            wc.cell(row=i, column=ci + 1,
                    value=x[f"ctr_{n}"]).number_format = "0.0%"
            wc.cell(row=i, column=ci + 2,
                    value=x[f"iwqs_{n}"]).number_format = "0.0"
            wc.cell(row=i, column=ci + 3,
                    value=x[f"roas_{n}"]).number_format = "0.00"
            wc.cell(row=i, column=ci + 4,
                    value=x[f"spend_{n}"]).number_format = "$#,##0"
            ci += 5
    if len(detail) > len(shown):
        wc.cell(row=len(shown) + 3, column=1,
                value=f"… {len(detail) - len(shown):,} more rows in "
                      f"{csv_path.name} (all six windows there)"
                ).font = SOFT_FONT
    widths = [11, 36, 7, 5, 12, 5, 11, 9, 15, 34] + [8, 6, 6, 6, 8] \
        * len(SHEET_WINDOWS)
    for c, w in enumerate(widths, 1):
        wc.column_dimensions[get_column_letter(c)].width = w
    wc.freeze_panes = "C2"

    wn = wb.create_sheet("Notes")
    notes = [
        "Verdict rules",
        "  PAUSE: 90d ROAS < 0.85 AND 30d ROAS < 0.85, not improving "
        "(30d < 90d+0.15 and 30d < 0.95), 90d spend >= "
        f"${args.min_verdict_spend:,.0f}. Candidates for the "
        "cut-list workflow — dated, tripwire-watched; not "
        "auto-execution.",
        f"  KEEP-BASE: 90d ROAS >= 1.00 at >= "
        f"${args.base_spend:,.0f} 90d spend. KEEP-INCREMENTAL: "
        "earning below the base bar (room to feed). WATCH: "
        "borderline, improving-unhealthy, or thin.",
        "  ZOMBIE: enabled, zero 90d impressions — deliberately NOT "
        "a pause: the Oct-2025 mass pause is a documented ratchet "
        "tread (footprint contraction fed the spiral); zombies cost "
        "nothing and are the exposure-death exhibit.",
        "  KEEP-PROTECTED: Farmers home-intent standing rule.",
        "Data rules",
        "  Census base = keyword_instances in the charter campaigns "
        "(zero-data rows included); grain keyword x MATCH TYPE x "
        "campaign, instances merged (Inst column); kpd device='ALL'; "
        "Leads = Conversions column; iwQS per window (never summed, "
        "unrated excluded); break-even 1.00.",
        "  Windows 7/14/30/60/90/YTD end at the data date; the sheet "
        "shows 14/30/90/YTD, the CSV carries all six.",
        "  UNITS: all window columns (spend_30, rev_90, ...) are RAW "
        "TOTALS over that window; any figure written '$X/d' divides "
        "a 30d total by 30. A 30d total will be ~3x smaller than the "
        "90d total on a steady keyword — that is windows, not a bug.",
        "  Windows overlap expansion waves — verdicts are level "
        "economics; phase-conditioning lives in the health workbook.",
    ]
    for i, line in enumerate(notes, 1):
        c = wn.cell(row=i, column=1, value=line)
        if not line.startswith("  "):
            c.font = Font(bold=True, size=11)
    wn.column_dimensions["A"].width = 110

    xlsx_path = out_dir / f"keyword_census_{stamp}.xlsx"
    wb.save(xlsx_path)
    print(f"\nSaved: {xlsx_path}")
    print(f"Saved: {csv_path}")
    print(f"Saved: {png_path}")


if __name__ == "__main__":
    main()
