#!/usr/bin/env python3
"""farmers_term_analysis.xlsx: the geico-workbook pattern applied to Farmers.
Inputs: /home/claude/farmers_inputs/. QS series = q7 Farmers segment (q6 is geico-hardcoded).
Periods: baseline Jan-Aug 25 / GP-negative Sep-Nov 25 / contraction+QS-peak Dec-Feb / re-expansion Mar-Aug 26."""
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import LineChart, Reference
from openpyxl.utils import get_column_letter

def load(p):
    delim = '\t' if p.endswith('.tsv') else ','
    return list(csv.DictReader(open('/home/claude/farmers_inputs/%s' % p), delimiter=delim))
exact, fam, var = load('term_trend_monthly.csv'), load('family_trend_monthly.csv'), load('family_variants.tsv')

QS = [(r['month'], int(r['impressions']), int(r['rated_impressions']), float(r['rated_share']),
       float(r['iw_qs']), float(r['creative_below_share']), float(r['pred_ctr_below_share']))
      for r in load('farmers_qs_monthly.csv')]

ALL_MONTHS = ['2025-%02d'%m for m in range(1,13)] + ['2026-%02d'%m for m in range(1,9)]
em = {r['month']: r for r in exact}; fm = {r['month']: r for r in fam}

wb = Workbook(); ARIAL='Arial'
def F(bold=False,size=10,color='000000'): return Font(name=ARIAL,bold=bold,size=size,color=color)
HDR=PatternFill('solid',fgColor='1F3864'); CLIFF=PatternFill('solid',fgColor='FCE4E4')
thin=Side(style='thin',color='BFBFBF'); BORDER=Border(left=thin,right=thin,top=thin,bottom=thin)
CLIFF_M={'2025-09','2025-10','2025-11'}

MCOLS=[('month','Month','@'),('impressions','Impressions','#,##0'),('clicks','Clicks','#,##0'),
 ('ctr','CTR','0.0%'),('cpc','CPC','$#,##0.00'),('cost','Cost','$#,##0'),('zips','Zips','#,##0'),
 ('leads','Leads','#,##0'),('click_to_zip','Click→Zip','0.0%'),('zip_to_lead','Zip→Lead','0.0%'),
 ('click_to_lead','Click→Lead','0.0%'),('value','Conv. Value','$#,##0'),('roas','ROAS','0.00x'),
 ('gp','GP','$#,##0;($#,##0)'),('gp_per_lead','GP / Lead','$#,##0.00;($#,##0.00)'),('cpl','CPL','$#,##0.00')]

def monthly_sheet(name,mmap,note):
    ws=wb.create_sheet(name); ws['A1']=note; ws['A1'].font=F(bold=True,size=9,color='7F7F7F')
    r0=3
    for j,(k,label,fmt) in enumerate(MCOLS,start=1):
        c=ws.cell(row=r0,column=j,value=label); c.font=F(bold=True,color='FFFFFF'); c.fill=HDR; c.border=BORDER
        c.alignment=Alignment(horizontal='center')
    for i,m in enumerate(ALL_MONTHS,start=r0+1):
        row=mmap.get(m)
        for j,(k,label,fmt) in enumerate(MCOLS,start=1):
            if k=='month': v=m
            elif row is None: v=0
            else:
                try: v=float(row[k])
                except: v=0
            c=ws.cell(row=i,column=j,value=v); c.font=F(); c.border=BORDER
            if fmt!='@': c.number_format=fmt
            if m in CLIFF_M: c.fill=CLIFF
    ws.freeze_panes='A4'
    for j,w in enumerate([9,12,9,7,9,10,8,8,10,10,11,12,7,11,11,9],start=1):
        ws.column_dimensions[get_column_letter(j)].width=w
    return ws,r0+1,r0+len(ALL_MONTHS)

ws_e,e_first,e_last=monthly_sheet('Monthly_Farmers_Exact',em,
 'Exact "farmers" — monthly. Search only (PMax excluded). Low volume — family sheet is the primary view. Red = GP-negative stretch.')
ws_f,f_first,f_last=monthly_sheet('Monthly_Farmers_Family',fm,
 'All farmers-containing terms — monthly. Search only (PMax excluded). Red = GP-negative stretch (Sep–Nov 2025).')

# QS sheet: full history + aligned helper block for charts
ws_q=wb.create_sheet('Monthly_Farmers_QS')
ws_q['A1']=('Farmers keywords — Google QS history (q7 segment; keyword grain, impression-weighted). '
 '2024-04 is a partial backfill month — treat as unreliable. BELOW-average shares of RATED impressions; remainder = AVERAGE/ABOVE.')
ws_q['A1'].font=F(bold=True,size=9,color='7F7F7F')
qh=['Month','Impressions','Rated Impr','Rated Share','IW QS','Creative Below %','Pred CTR Below %']
for j,h in enumerate(qh,start=1):
    c=ws_q.cell(row=3,column=j,value=h); c.font=F(bold=True,color='FFFFFF'); c.fill=HDR; c.border=BORDER
qfmt=['@','#,##0','#,##0','0.0%','0.00','0.0%','0.0%']
for i,row in enumerate(QS,start=4):
    for j,(v,fmt) in enumerate(zip(row,qfmt),start=1):
        c=ws_q.cell(row=i,column=j,value=v); c.font=F(); c.border=BORDER
        if fmt!='@': c.number_format=fmt
        if row[0] in CLIFF_M: c.fill=CLIFF
ws_q.freeze_panes='A4'
for j,w in enumerate([9,12,11,11,8,15,15],start=1): ws_q.column_dimensions[get_column_letter(j)].width=w
# aligned helper (chart months only) in column I: header row 3, values rows 4..23
ws_q.cell(row=3,column=9,value='IW QS').font=F(bold=True,size=8,color='7F7F7F')
qs_by_m={r[0]:r[4] for r in QS}
for i,m in enumerate(ALL_MONTHS,start=4):
    c=ws_q.cell(row=i,column=9,value=qs_by_m.get(m)); c.number_format='0.00'; c.font=F(size=8,color='7F7F7F')
ws_q.cell(row=2,column=9,value='chart-aligned (2025-01..2026-08)').font=F(size=7,color='BFBFBF')

# Combined sheet (unchanged from v1)
ws_c=wb.create_sheet('Monthly_Combined')
ws_c['A1']=('Exact vs Family side-by-side. "Family excl. exact" = Family − Exact. Rates recomputed from component sums.')
ws_c['A1'].font=F(bold=True,size=9,color='7F7F7F')
groups=[('FARMERS EXACT',None),('FARMERS FAMILY (ALL)',None),('FAMILY EXCL. EXACT',None)]
sub=['Impr','Clicks','CTR','CPC','Leads','GP','CPL','ROAS']; r0=3
ws_c.cell(row=r0+1,column=1,value='Month').font=F(bold=True,color='FFFFFF'); ws_c.cell(row=r0+1,column=1).fill=HDR
col=2
for gname,_ in groups:
    c=ws_c.cell(row=r0,column=col,value=gname); c.font=F(bold=True,color='FFFFFF'); c.fill=HDR
    ws_c.merge_cells(start_row=r0,start_column=col,end_row=r0,end_column=col+len(sub)-1)
    c.alignment=Alignment(horizontal='center')
    for j,s2 in enumerate(sub):
        cc=ws_c.cell(row=r0+1,column=col+j,value=s2); cc.font=F(bold=True,color='FFFFFF'); cc.fill=HDR
    col+=len(sub)
for i,m in enumerate(ALL_MONTHS):
    r=r0+2+i
    ws_c.cell(row=r,column=1,value=m).font=F()
    def put(ci,fo,fmt):
        c=ws_c.cell(row=r,column=ci,value=fo); c.font=F(); c.number_format=fmt
    for base,sheet,fr in [(2,'Monthly_Farmers_Exact',e_first+i),(10,'Monthly_Farmers_Family',f_first+i)]:
        put(base+0,"='%s'!B%d"%(sheet,fr),'#,##0'); put(base+1,"='%s'!C%d"%(sheet,fr),'#,##0')
        put(base+2,"=IF('%s'!B%d=0,0,'%s'!C%d/'%s'!B%d)"%(sheet,fr,sheet,fr,sheet,fr),'0.0%')
        put(base+3,"=IF('%s'!C%d=0,0,'%s'!F%d/'%s'!C%d)"%(sheet,fr,sheet,fr,sheet,fr),'$#,##0.00')
        put(base+4,"='%s'!H%d"%(sheet,fr),'#,##0'); put(base+5,"='%s'!N%d"%(sheet,fr),'$#,##0;($#,##0)')
        put(base+6,"=IF('%s'!H%d=0,0,'%s'!F%d/'%s'!H%d)"%(sheet,fr,sheet,fr,sheet,fr),'$#,##0.00')
        put(base+7,"=IF('%s'!F%d=0,0,'%s'!L%d/'%s'!F%d)"%(sheet,fr,sheet,fr,sheet,fr),'0.00x')
    b=18; ef,ff=e_first+i,f_first+i
    put(b+0,"=J%d-B%d"%(r,r),'#,##0'); put(b+1,"=K%d-C%d"%(r,r),'#,##0')
    put(b+2,"=IF(R%d=0,0,S%d/R%d)"%(r,r,r),'0.0%')
    put(b+3,"=IF(S%d=0,0,('Monthly_Farmers_Family'!F%d-'Monthly_Farmers_Exact'!F%d)/S%d)"%(r,ff,ef,r),'$#,##0.00')
    put(b+4,"=N%d-F%d"%(r,r),'#,##0'); put(b+5,"=O%d-G%d"%(r,r),'$#,##0;($#,##0)')
    put(b+6,"=IF(V%d=0,0,('Monthly_Farmers_Family'!F%d-'Monthly_Farmers_Exact'!F%d)/V%d)"%(r,ff,ef,r),'$#,##0.00')
    put(b+7,"=IF(('Monthly_Farmers_Family'!F%d-'Monthly_Farmers_Exact'!F%d)=0,0,('Monthly_Farmers_Family'!L%d-'Monthly_Farmers_Exact'!L%d)/('Monthly_Farmers_Family'!F%d-'Monthly_Farmers_Exact'!F%d))"%(ff,ef,ff,ef,ff,ef),'0.00x')
ws_c.freeze_panes='B5'
for j in range(1,26): ws_c.column_dimensions[get_column_letter(j)].width=10

# Period summary (unchanged)
ws_p=wb.create_sheet('Period_Summary',0)
ws_p['A1']='Farmers: Period Summary — Baseline, GP-Negative Dip, Contraction, Re-expansion'; ws_p['A1'].font=F(bold=True,size=13)
ws_p['A2']=('Rates recomputed from period totals. Search only (PMax excluded). See Monthly_Farmers_QS for the quality-score history.')
ws_p['A2'].font=F(size=9,color='7F7F7F')
periods=[('Jan–Aug 2025 (baseline)',0,7),('Sep–Nov 2025 (GP-negative)',8,10),('Dec 25 – Feb 26 (contraction, QS peak)',11,13),('Mar–Aug 2026 (re-expansion, ~$9 CPC)',14,19)]
PC=['Impr/mo','Clicks/mo','CTR','CPC','Zips/mo','Leads/mo','Click→Zip','Zip→Lead','CPL','ROAS','GP/mo','GP/Lead']
def block(title,sheet,firstrow,startrow):
    ws_p.cell(row=startrow,column=1,value=title).font=F(bold=True,size=11)
    hr=startrow+1
    ws_p.cell(row=hr,column=1,value='Period').font=F(bold=True,color='FFFFFF'); ws_p.cell(row=hr,column=1).fill=HDR; ws_p.cell(row=hr,column=1).border=BORDER
    for j,h in enumerate(PC,start=2):
        c=ws_p.cell(row=hr,column=j,value=h); c.font=F(bold=True,color='FFFFFF'); c.fill=HDR; c.border=BORDER
    for i,(pname,m0,m1) in enumerate(periods):
        r=hr+1+i; n=m1-m0+1; a,b=firstrow+m0,firstrow+m1
        ws_p.cell(row=r,column=1,value=pname).font=F(bold=(i==0))
        S=lambda col:"SUM('%s'!%s%d:%s%d)"%(sheet,col,a,col,b)
        vals=[('=%s/%d'%(S('B'),n),'#,##0'),('=%s/%d'%(S('C'),n),'#,##0'),
              ('=IF(%s=0,0,%s/%s)'%(S('B'),S('C'),S('B')),'0.0%'),
              ('=IF(%s=0,0,%s/%s)'%(S('C'),S('F'),S('C')),'$#,##0.00'),
              ('=%s/%d'%(S('G'),n),'#,##0'),('=%s/%d'%(S('H'),n),'#,##0'),
              ('=IF(%s=0,0,%s/%s)'%(S('C'),S('G'),S('C')),'0.0%'),
              ('=IF(%s=0,0,%s/%s)'%(S('G'),S('H'),S('G')),'0.0%'),
              ('=IF(%s=0,0,%s/%s)'%(S('H'),S('F'),S('H')),'$#,##0.00'),
              ('=IF(%s=0,0,%s/%s)'%(S('F'),S('L'),S('F')),'0.00x'),
              ('=%s/%d'%(S('N'),n),'$#,##0;($#,##0)'),
              ('=IF(%s=0,0,%s/%s)'%(S('H'),S('N'),S('H')),'$#,##0.00;($#,##0.00)')]
        for j,(fo,fmt) in enumerate(vals,start=2):
            c=ws_p.cell(row=r,column=j,value=fo); c.font=F(); c.number_format=fmt; c.border=BORDER
            if i==1: c.fill=CLIFF
    return hr+len(periods)+2
r=4
r=block('FARMERS — EXACT TERM (low volume; directional)','Monthly_Farmers_Exact',e_first,r)
r=block('FARMERS — FAMILY (all farmers-containing terms)','Monthly_Farmers_Family',f_first,r)
ws_p.column_dimensions['A'].width=30
for j in range(2,15): ws_p.column_dimensions[get_column_letter(j)].width=11

# Variants (unchanged)
ws_v=wb.create_sheet('Variants_2025_v_2026')
ws_v['A1']='Farmers family variants — leads & GP, matched Jan–Aug windows 2025 vs 2026 (>=10 clicks)'; ws_v['A1'].font=F(bold=True,size=10)
vh=['Query','Leads 2025','Leads 2026','GP 2025','GP 2026','GP Change']
for j,h in enumerate(vh,start=1):
    c=ws_v.cell(row=3,column=j,value=h); c.font=F(bold=True,color='FFFFFF'); c.fill=HDR; c.border=BORDER
def flt(x):
    try: return float(x)
    except: return 0.0
var.sort(key=lambda r: flt(r['gp_change']))
for i,rw in enumerate(var,start=4):
    vals=[rw['query'],flt(rw['leads_prior']),flt(rw['leads_current']),flt(rw['gp_prior']),flt(rw['gp_current']),flt(rw['gp_change'])]
    fmts=['@','#,##0','#,##0','$#,##0;($#,##0)','$#,##0;($#,##0)','$#,##0;($#,##0)']
    for j,(v,fmt) in enumerate(zip(vals,fmts),start=1):
        c=ws_v.cell(row=i,column=j,value=v); c.font=F(); c.border=BORDER
        if fmt!='@': c.number_format=fmt
ws_v.freeze_panes='A4'; ws_v.column_dimensions['A'].width=42
for j in range(2,7): ws_v.column_dimensions[get_column_letter(j)].width=13
ws_v.auto_filter.ref='A3:F%d'%(3+len(var))

# Charts: 4 combo charts, each with QS on a secondary axis
ws_ch=wb.create_sheet('Charts')
ws_ch['A34']='Dashed line: impression-weighted QS (keyword grain, q7 Farmers segment). QS sheet covers 2024-04+; charts start 2025-01.'
ws_ch['A34'].font=F(size=8,color='7F7F7F')
def mkchart(title,ycol,y_title,anchor):
    c1=LineChart(); c1.title=title; c1.height=8; c1.width=17
    c1.y_axis.title=y_title; c1.x_axis.title='Month'
    d1=Reference(ws_e,min_col=ycol,min_row=e_first-1,max_row=e_last)
    d2=Reference(ws_f,min_col=ycol,min_row=f_first-1,max_row=f_last)
    cats=Reference(ws_e,min_col=1,min_row=e_first,max_row=e_last)
    c1.add_data(d1,titles_from_data=True); c1.add_data(d2,titles_from_data=True)
    c1.set_categories(cats)
    # secondary axis: iw_qs from the aligned helper column (I) on the QS sheet
    c2=LineChart()
    qref=Reference(ws_q,min_col=9,min_row=3,max_row=3+len(ALL_MONTHS))
    c2.add_data(qref,titles_from_data=True)
    c2.y_axis.axId=200; c2.y_axis.title='IW QS (1-10)'
    c2.y_axis.majorGridlines=None
    s=c2.series[0]
    s.graphicalProperties.line.dashStyle='dash'
    s.graphicalProperties.line.solidFill='7030A0'
    c1.y_axis.crosses='max'
    c1+=c2
    ws_ch.add_chart(c1,anchor)
mkchart('Farmers leads/month — exact vs family (QS overlay)',8,'Leads','A1')
mkchart('Farmers CTR — exact vs family (QS overlay)',4,'CTR','A18')
mkchart('Farmers CPC — exact vs family (QS overlay)',5,'CPC ($)','K1')
mkchart('Farmers GP/month — exact vs family (QS overlay)',14,'GP ($)','K18')

del wb['Sheet']
wb.save('/home/claude/farmers_term_analysis.xlsx')
print('saved farmers workbook')
