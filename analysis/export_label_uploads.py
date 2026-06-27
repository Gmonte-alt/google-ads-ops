"""
export_label_uploads.py [v2 — joins live campaigns/ad_groups for names]
=======================================================================
Reads the kw_target_brand_lookup output Excel and produces one
import-ready CSV per account, formatted for the appropriate ad platform's
editor (Google Ads Editor or Microsoft Ads Editor).

CHANGES FROM v1:
  - Names are now sourced from the live campaigns and ad_groups tables
    (joined on campaign_id / ad_group_id), NOT from the denormalized
    name columns in keyword_instances. This eliminates "ghost ad group"
    issues caused by stale or mis-synced names.
  - Rows whose campaign_id or ad_group_id can't be found in campaigns
    or ad_groups are SKIPPED with a warning — these likely refer to
    deleted entities that no longer exist in the live account.
  - Both Google and Microsoft outputs include ID columns (campaign_id,
    ad_group_id) as informational fields. Google's editor uses names
    for matching by default; the IDs are there as audit context and
    in case future versions of the editor support ID-based matching.

OUTPUT
──────
  reports/label_uploads_YYYY-MM-DD/
    google_3863660890_auto.csv          (UTF-16, tab-delimited)
    google_1006729211_auto_quote.csv    (UTF-16, tab-delimited)
    microsoft_F149Y8EN_auto.csv         (UTF-8,  comma-delimited)
    microsoft_F149SN63_home_insurance.csv (UTF-8, comma-delimited)
    skipped_rows.csv                    (audit log of rows that couldn't
                                         be matched to live campaigns/ad_groups)

USAGE
─────
    conda activate paid-search
    python analysis/export_label_uploads.py
    python analysis/export_label_uploads.py --tier1-only
"""

import argparse
import csv
import re
import sqlite3
from datetime import date
from pathlib import Path

import openpyxl


# ── Configuration ─────────────────────────────────────────────────────────────

DB_PATH     = Path(r"C:\Projects\google-ads-ops\db\paid_search.db")
OUTPUT_BASE = Path(r"C:\Projects\google-ads-ops\reports")
DEFAULT_LABEL = "Brand:Confie"


def find_default_xlsx() -> Path:
    matches = sorted(OUTPUT_BASE.glob("kw_target_brand_lookup_*.xlsx"))
    if not matches:
        raise FileNotFoundError(
            f"No kw_target_brand_lookup_*.xlsx found in {OUTPUT_BASE}"
        )
    return matches[-1]


EXPORT_ACCOUNTS = {
    ("google",    "3863660890"),
    ("google",    "1006729211"),
    ("microsoft", "F149Y8EN"),
    ("microsoft", "F149SN63"),
}

MATCH_TYPE_MAP = {"EXACT": "Exact", "PHRASE": "Phrase", "BROAD": "Broad"}
STATUS_MAP     = {"ENABLED": "Enabled", "PAUSED": "Paused", "REMOVED": "Removed"}


# ── Helpers ───────────────────────────────────────────────────────────────────

def slugify(s: str) -> str:
    s = s.strip().lower()
    s = re.sub(r"[^a-z0-9]+", "_", s)
    return s.strip("_")


def determine_tier(kw_status: str, camp_status: str, ag_status: str) -> str:
    if (kw_status   or "").upper() == "ENABLED" \
       and (camp_status or "").upper() == "ENABLED" \
       and (ag_status   or "").upper() == "ENABLED":
        return "Tier 1 (actively serving)"
    return "Tier 2 (dormant)"


# ── Live-name lookups from campaigns / ad_groups ──────────────────────────────

def build_live_name_maps(account_filter: list[tuple[str, str]]) -> tuple[dict, dict, dict]:
    """
    Returns three dicts keyed by (account_db_id, *_id):
      - campaign_name_map[(account_id, campaign_id)] -> live campaign_name
      - ad_group_name_map[(account_id, ad_group_id)] -> live ad_group_name
      - account_db_id_map[(platform, platform_account_id)] -> account.id

    These are the source of truth — denormalized names in keyword_instances
    are bypassed entirely.
    """
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    cur = conn.cursor()

    placeholders = ",".join(["(?, ?)"] * len(account_filter))
    flat = [v for pair in account_filter for v in pair]

    # Resolve account DB ids
    account_db_ids = cur.execute(
        f"""
        SELECT id, platform, platform_account_id
        FROM accounts
        WHERE (platform, platform_account_id) IN ({placeholders})
        """,
        flat,
    ).fetchall()
    account_db_id_map = {
        (r["platform"], str(r["platform_account_id"])): r["id"]
        for r in account_db_ids
    }
    target_db_ids = list(account_db_id_map.values())

    if not target_db_ids:
        conn.close()
        return {}, {}, {}

    db_ph = ",".join("?" * len(target_db_ids))

    camp_rows = cur.execute(
        f"SELECT account_id, campaign_id, campaign_name "
        f"FROM campaigns WHERE account_id IN ({db_ph})",
        target_db_ids,
    ).fetchall()
    campaign_name_map = {
        (r["account_id"], str(r["campaign_id"])): r["campaign_name"]
        for r in camp_rows
    }

    ag_rows = cur.execute(
        f"SELECT account_id, ad_group_id, ad_group_name "
        f"FROM ad_groups WHERE account_id IN ({db_ph})",
        target_db_ids,
    ).fetchall()
    ad_group_name_map = {
        (r["account_id"], str(r["ad_group_id"])): r["ad_group_name"]
        for r in ag_rows
    }

    conn.close()
    print(f"  Live names loaded: {len(campaign_name_map):,} campaigns, "
          f"{len(ad_group_name_map):,} ad groups")
    return campaign_name_map, ad_group_name_map, account_db_id_map


# ── Read instances sheet ──────────────────────────────────────────────────────

def load_instances(xlsx_path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    target_sheet = next(
        (s for s in wb.sheetnames if s.startswith("Instances ")), None
    )
    if target_sheet is None:
        raise ValueError(f"No 'Instances ' sheet found in {xlsx_path.name}")

    ws = wb[target_sheet]
    rows = list(ws.iter_rows(values_only=True))
    headers = list(rows[0])
    data = [dict(zip(headers, r)) for r in rows[1:]]
    print(f"  Loaded {len(data):,} instance rows from sheet '{target_sheet}'")
    return data


def dedupe_to_unique_keywords(instances: list[dict]) -> list[dict]:
    seen = set()
    unique = []
    for row in instances:
        iid = row.get("instance_id")
        if iid is None or iid in seen:
            continue
        seen.add(iid)
        unique.append(row)
    print(f"  Deduplicated to {len(unique):,} unique keyword instances")
    return unique


# ── Live-name overlay ─────────────────────────────────────────────────────────

def overlay_live_names(
    instances: list[dict],
    campaign_name_map: dict,
    ad_group_name_map: dict,
    account_db_id_map: dict,
) -> tuple[list[dict], list[dict]]:
    """
    For each instance row, replace campaign_name / ad_group_name with the
    live values from campaigns / ad_groups. Rows whose IDs aren't found in
    those tables are SKIPPED and returned in `skipped` for audit logging.
    """
    keep = []
    skipped = []
    overlaid_camp = 0
    overlaid_ag   = 0

    for r in instances:
        platform = (r.get("platform") or "").strip().lower()
        pid      = str(r.get("account_id_str") or "").strip()
        account_db_id = account_db_id_map.get((platform, pid))
        if account_db_id is None:
            skipped.append({**r, "skip_reason": "account not in DB filter"})
            continue

        campaign_id = str(r.get("campaign_id") or "").strip()
        ad_group_id = str(r.get("ad_group_id") or "").strip()

        live_campaign_name = campaign_name_map.get((account_db_id, campaign_id))
        live_ad_group_name = ad_group_name_map.get((account_db_id, ad_group_id))

        if live_campaign_name is None:
            skipped.append({**r, "skip_reason":
                            f"campaign_id {campaign_id} not in live campaigns table"})
            continue
        if live_ad_group_name is None:
            skipped.append({**r, "skip_reason":
                            f"ad_group_id {ad_group_id} not in live ad_groups table"})
            continue

        # Note any name overlay so we can report it
        if r.get("campaign_name") and r["campaign_name"] != live_campaign_name:
            overlaid_camp += 1
        if r.get("ad_group_name") and r["ad_group_name"] != live_ad_group_name:
            overlaid_ag += 1

        # Replace stored names with live names
        r = {**r,
             "campaign_name": live_campaign_name,
             "ad_group_name": live_ad_group_name}
        keep.append(r)

    print(f"  Live-name overlay: kept {len(keep):,}, skipped {len(skipped):,}")
    if overlaid_camp:
        print(f"    Overlaid {overlaid_camp:,} stale campaign names")
    if overlaid_ag:
        print(f"    Overlaid {overlaid_ag:,} stale ad group names")
    return keep, skipped


# ── Per-platform writers ──────────────────────────────────────────────────────

def write_google_csv(path: Path, rows: list[dict], label: str) -> int:
    headers = [
        "Campaign", "Ad Group", "Keyword", "Criterion Type", "Labels",
        "Status", "Campaign Status", "Ad Group Status",
        "Campaign ID", "Ad Group ID", "_tier",
    ]
    with open(path, "w", encoding="utf-16", newline="") as f:
        w = csv.writer(f, delimiter="\t")
        w.writerow(headers)
        for r in rows:
            mt = MATCH_TYPE_MAP.get((r.get("match_type") or "").upper(), "")
            kws = STATUS_MAP.get((r.get("kw_status")    or "").upper(), "")
            cs  = STATUS_MAP.get((r.get("campaign_status") or "").upper(), "")
            ags = STATUS_MAP.get((r.get("ad_group_status") or "").upper(), "")
            tier = determine_tier(
                r.get("kw_status"), r.get("campaign_status"), r.get("ad_group_status")
            )
            w.writerow([
                r.get("campaign_name") or "",
                r.get("ad_group_name") or "",
                r.get("keyword_text")  or "",
                mt,
                label,
                kws, cs, ags,
                str(r.get("campaign_id") or ""),
                str(r.get("ad_group_id") or ""),
                tier,
            ])
    return len(rows)


def write_microsoft_csv(path: Path, rows: list[dict], label: str) -> int:
    headers = [
        "Type", "Campaign", "Ad Group", "Keyword", "Match Type", "Labels",
        "Campaign ID", "Ad Group ID", "_tier",
    ]
    with open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(headers)
        for r in rows:
            mt = MATCH_TYPE_MAP.get((r.get("match_type") or "").upper(), "")
            tier = determine_tier(
                r.get("kw_status"), r.get("campaign_status"), r.get("ad_group_status")
            )
            w.writerow([
                "Keyword",
                r.get("campaign_name") or "",
                r.get("ad_group_name") or "",
                r.get("keyword_text")  or "",
                mt,
                label,
                str(r.get("campaign_id") or ""),
                str(r.get("ad_group_id") or ""),
                tier,
            ])
    return len(rows)


def write_skipped_log(path: Path, skipped: list[dict]) -> None:
    if not skipped:
        return
    fields = [
        "platform", "account_id_str", "account_name",
        "campaign_id", "campaign_name",
        "ad_group_id", "ad_group_name",
        "keyword_text", "match_type", "kw_status",
        "skip_reason",
    ]
    with open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        for r in skipped:
            w.writerow({k: r.get(k, "") for k in fields})


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    p = argparse.ArgumentParser()
    p.add_argument("--xlsx",  default=None)
    p.add_argument("--label", default=DEFAULT_LABEL)
    p.add_argument("--tier1-only", action="store_true")
    args = p.parse_args()

    xlsx_path = Path(args.xlsx) if args.xlsx else find_default_xlsx()

    print("=" * 70)
    print("export_label_uploads.py  [v2 — live-name overlay]")
    print(f"  Source : {xlsx_path}")
    print(f"  Label  : {args.label}")
    print(f"  Filter : {'Tier 1 only' if args.tier1_only else 'All tiers'}")
    print("=" * 70)

    print("\n[1] Loading live campaign/ad-group names from DB...")
    camp_map, ag_map, acct_map = build_live_name_maps(list(EXPORT_ACCOUNTS))

    print("\n[2] Reading instances sheet...")
    instances = load_instances(xlsx_path)

    print("\n[3] Deduplicating to unique keyword instances...")
    instances = dedupe_to_unique_keywords(instances)

    print("\n[4] Overlaying live campaign/ad-group names...")
    instances, skipped = overlay_live_names(instances, camp_map, ag_map, acct_map)

    if args.tier1_only:
        before = len(instances)
        instances = [
            r for r in instances
            if determine_tier(
                r.get("kw_status"), r.get("campaign_status"),
                r.get("ad_group_status")
            ).startswith("Tier 1")
        ]
        print(f"\n[5] Tier 1 filter: {before:,} → {len(instances):,}")

    # Group by account
    print("\n[6] Grouping by account...")
    groups: dict[tuple[str,str,str], list[dict]] = {}
    for r in instances:
        platform = (r.get("platform") or "").strip().lower()
        pid      = str(r.get("account_id_str") or "").strip()
        if (platform, pid) not in EXPORT_ACCOUNTS:
            continue
        name = r.get("account_name") or "unknown"
        groups.setdefault((platform, pid, name), []).append(r)

    out_dir = OUTPUT_BASE / f"label_uploads_{date.today().isoformat()}"
    out_dir.mkdir(parents=True, exist_ok=True)
    print(f"\n[7] Writing per-account files to:")
    print(f"      {out_dir}")

    total_t1 = total_t2 = 0
    for (platform, pid, name), rows in sorted(groups.items()):
        slug = slugify(name)
        out_path = out_dir / f"{platform}_{pid}_{slug}.csv"

        if platform == "google":
            n = write_google_csv(out_path, rows, args.label)
            fmt = "UTF-16 / tab"
        elif platform == "microsoft":
            n = write_microsoft_csv(out_path, rows, args.label)
            fmt = "UTF-8 / comma"
        else:
            continue

        t1 = sum(1 for r in rows if determine_tier(
            r.get("kw_status"), r.get("campaign_status"),
            r.get("ad_group_status")
        ).startswith("Tier 1"))
        t2 = n - t1
        total_t1 += t1
        total_t2 += t2
        print(f"  {platform:<10} {pid:<12} {name:<25} "
              f"{n:>5,} rows  (T1={t1:,}  T2={t2:,})  [{fmt}]")

    # Write skipped rows audit log
    if skipped:
        skip_path = out_dir / "skipped_rows.csv"
        write_skipped_log(skip_path, skipped)
        print(f"\n[8] Skipped-row audit log: {skip_path}")
        print(f"    {len(skipped):,} rows could not be matched to live "
              f"campaigns/ad_groups and were excluded.")

    print(f"\n[9] Summary:")
    print(f"  Tier 1 (actively serving): {total_t1:,}")
    print(f"  Tier 2 (dormant)         : {total_t2:,}")
    print(f"  Total exported           : {total_t1 + total_t2:,}")
    print(f"  Skipped (live-name miss) : {len(skipped):,}")
    print("\nDone.")


if __name__ == "__main__":
    main()