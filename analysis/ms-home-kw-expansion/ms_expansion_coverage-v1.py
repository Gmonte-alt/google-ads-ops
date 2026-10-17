"""
ms_expansion_coverage.py -- STRUCTURE-SIDE coverage check for the Microsoft Home
(account_id=2) generic keyword expansion. Read-only against paid_search.db.

For each keyword in the MS Keyword Planner export, answers two questions:

  (1) WHAT IS IT -- brand/DNB classification via the brand registry (brand_tagger.py,
      same logic as the USIA pipeline), plus a home-qualification flag: a term with no
      home/property token and no brand ("cheap insurance quotes") is UNQUALIFIED --
      cross-vertical leak risk at these CPCs, flagged not dropped.

  (2) IS IT ALREADY IN THE STRUCTURE -- resolve the text against keyword_current
      (acct 2) and issue one structure verdict:
        ACTIVE_TARGET   is_serving=1 instance exists (reports match types + live homes)
        PAUSED_TARGET   keyword PAUSED under a live hierarchy -> REACTIVATION candidate
        DORMANT_TARGET  exists only under paused/dead hierarchy -> history exists,
                        effectively net-new (rebuild or reactivate structure)
        NET_NEW         no instance at any status
      NEGATIVE_BLOCKED is layered on when --negatives-csv is supplied (MS negatives
      are NOT in the DB -- export from the platform; expected cols: keyword_text,
      match_type, level; extra cols ignored).

MATCHING (two tiers, per the 09 identity tagger):
  canonical  -- 09's normal form (lowercase + whitespace-collapse ONLY; punctuation
                KEPT), matched against keyword_identities.normalized_text via
                keyword_current.identity_id where available, else normalize_09 of
                keyword_text. This is the account's own identity layer -- zero drift.
  close-variant -- punctuation-stripped fallback ("homeowner's" ~ "homeowners"), since
                MS close-variant matching treats these as one target. Hits carry the
                same verdict but match_basis='close-variant' so they read softer.

  Plus a THEME suggestion mapped to the account's own ad-group naming taxonomy
  (Core/Quote/Best/Compare/Cheapest/Cost/Companies/Agent), with a marker for whether
  a LIVE generic ad group already exists for that theme (recomputed at runtime).

NOT COVERED here (needs the MS Search Term Report load): whether live BROAD keywords
already absorb the query. That is the served-side check, added when sqp has MS rows.

    python ms_expansion_coverage.py --planner-csv "Keyword_Planner_09-17-2026_at_14-39-58.csv"
    python ms_expansion_coverage.py --planner-csv ...csv --negatives-csv ms_negatives.csv
"""
import argparse
import csv
import os
import re
import sqlite3
import sys
from collections import defaultdict
from pathlib import Path
from datetime import date

sys.stdout.reconfigure(encoding="utf-8")
ACCT = 2  # Microsoft Home


# ---------------------------------------------------------------- normalizers
def normalize_09(text):
    """EXACT copy of 09_tag_keyword_identities.normalize -- the canonical identity form.
    Lowercase + whitespace-collapse ONLY. Punctuation/apostrophes are KEPT."""
    return re.sub(r"\s+", " ", (text or "").lower().strip())


def normalize_loose(text):
    """Close-variant form: drop apostrophes, punctuation->space, collapse. Used as the
    second-tier match (MS treats these as close variants) and for negatives/themes."""
    t = (text or "").lower().replace("\u2019", "").replace("'", "")
    t = re.sub(r"[^a-z0-9]+", " ", t)
    return re.sub(r"\s+", " ", t).strip()


# ---------------------------------------------------------------- classification
HOME_RE = re.compile(r"\b(home|homes|homeowner|homeowners|house|houses|property|condo|"
                     r"condos|dwelling|hazard|insure)\b")   # 'insure' kept: planner's like-group
VOL_MID = {"10 - 100": 55, "100 - 1K": 550, "1K - 10K": 5500,
           "10K - 100K": 55000, "100K - 1M": 550000}

# Theme ladder keyed to the ACCOUNT'S OWN ad-group naming (first match wins).
THEME_PATTERNS = [
    ("Agent",     r"\b(agent|agents|agency|agencies|broker|brokers|near me|local|find)\b"),
    ("Quote",     r"\b(quote|quotes|quotation|estimates?)\b"),
    ("Cheapest",  r"\b(cheap|cheaper|cheapest|affordable|low cost|budget|discount)\b"),
    ("Best",      r"\b(best|top|which .* is best|top-rated)\b"),
    ("Compare",   r"\b(compare|comparison|compares?)\b"),
    ("Companies", r"\b(company|companies)\b"),
    ("Cost",      r"\b(cost|costs|average|how much|rates?|premium|pricing)\b"),
    ("Policy",    r"\b(policy|policies|plans?|coverage)\b"),
]
THEME_RE = [(n, re.compile(p)) for n, p in THEME_PATTERNS]


def theme_of(loose_kw):
    for name, rx in THEME_RE:
        if rx.search(loose_kw):
            return name
    return "Core"


def load_tagger(tagger_dir, brand_db):
    """brand_tagger.py lives in kw-planner\\home_insurance -- import it if findable."""
    cands = [tagger_dir] if tagger_dir else []
    here = Path(__file__).resolve().parent
    cands += [Path(r"C:\Projects\google-ads-ops\kw-planner\home_insurance"),
              here, here.parent, Path.cwd()]
    for d in cands:
        try:
            if d and Path(d, "brand_tagger.py").exists():
                sys.path.insert(0, str(d))
                from brand_tagger import BrandTagger
                return BrandTagger(brand_db), str(d)
        except Exception as e:
            print("  [warn] brand_tagger import failed from %s: %s" % (d, e))
    return None, None


# ---------------------------------------------------------------- DB reads
def cols_of(conn, table):
    try:
        return [r[1] for r in conn.execute("PRAGMA table_info(%s)" % table)]
    except sqlite3.OperationalError:
        return []


def load_structure(conn):
    """acct-2 keyword_current rows indexed two ways:
         by_canon -- keyword_identities.normalized_text via identity_id (canonical),
                     falling back to normalize_09(keyword_text) where unlinked
         by_loose -- normalize_loose(keyword_text)
       Returns (by_canon, by_loose, n_rows, basis_note)."""
    kc_cols = cols_of(conn, "keyword_current")
    has_identity = "identity_id" in kc_cols
    id2norm = {}
    basis = "normalize_09(keyword_text) only (no identity link resolved)"
    if has_identity:
        ki_cols = cols_of(conn, "keyword_identities")
        idcol = next((c for c in ("id", "identity_id") if c in ki_cols), None)
        if idcol and "normalized_text" in ki_cols:
            for iid, nt in conn.execute(
                    "SELECT %s, normalized_text FROM keyword_identities "
                    "WHERE %s IN (SELECT DISTINCT identity_id FROM keyword_current "
                    "             WHERE account_id=? AND identity_id IS NOT NULL)"
                    % (idcol, idcol), (ACCT,)):
                id2norm[iid] = nt
            basis = ("keyword_identities.normalized_text via identity_id "
                     "(%s ids loaded; 09-form fallback for unlinked rows)"
                     % f"{len(id2norm):,}")
    rows = conn.execute(
        "SELECT keyword_text, match_type, keyword_status, hierarchy_enabled, is_serving, "
        "       campaign_name, ad_group_name%s "
        "FROM keyword_current WHERE account_id=?"
        % (", identity_id" if has_identity else ""), (ACCT,)).fetchall()
    by_canon, by_loose = defaultdict(list), defaultdict(list)
    linked = 0
    for r in rows:
        canon = None
        if has_identity and r["identity_id"] is not None:
            canon = id2norm.get(r["identity_id"])
            if canon:
                linked += 1
        if not canon:
            canon = normalize_09(r["keyword_text"])
        by_canon[canon].append(r)
        by_loose[normalize_loose(r["keyword_text"])].append(r)
    if id2norm:
        basis += "; %s of %s rows matched an identity" % (f"{linked:,}", f"{len(rows):,}")
    return by_canon, by_loose, len(rows), basis


def live_generic_themes(conn):
    """Live generic ad groups -> theme tokens present, from the naming convention."""
    themes = set()
    for (nm,) in conn.execute(
            "SELECT DISTINCT ag.ad_group_name FROM ad_groups ag "
            "JOIN campaigns c ON c.campaign_id=ag.campaign_id AND c.account_id=ag.account_id "
            "WHERE ag.account_id=? AND ag.status='ENABLED' AND c.status='ENABLED' "
            "AND lower(ag.ad_group_name) NOT LIKE '%non-generic%' "
            "AND lower(ag.ad_group_name) NOT LIKE '%nongeneric%' "
            "AND lower(ag.ad_group_name) NOT LIKE '%non generic%'", (ACCT,)):
        parts = [p for p in (nm or "").split("-") if p]
        # Home-Generic-{Product}-{Intent}[-suffix] -> intent slot; fall back to product
        if len(parts) >= 4:
            themes.add(parts[3].strip())
        elif len(parts) == 3:
            themes.add(parts[2].strip())
    return themes


def verdict_for(matches):
    if not matches:
        return "NET_NEW", []
    serving = [m for m in matches if m["is_serving"] == 1]
    if serving:
        return "ACTIVE_TARGET", serving
    paused_live = [m for m in matches
                   if m["hierarchy_enabled"] == 1 and m["keyword_status"] == "PAUSED"]
    if paused_live:
        return "PAUSED_TARGET", paused_live
    return "DORMANT_TARGET", matches


def load_negatives(path):
    negs = []
    if not path:
        return negs
    with open(path, newline="", encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            kt = (row.get("keyword_text") or row.get("Keyword") or row.get("keyword") or "").strip()
            if kt:
                negs.append({"norm": normalize_loose(kt),
                             "match": (row.get("match_type") or row.get("Match type") or "").lower(),
                             "level": row.get("level") or row.get("Level") or ""})
    return negs


def negative_hit(loose_kw, negs):
    """Conservative negative screen: exact -> equality; phrase/broad -> containment."""
    toks = loose_kw.split()
    for n in negs:
        if not n["norm"]:
            continue
        if n["match"].startswith("exact"):
            if n["norm"] == loose_kw:
                return n
        else:
            nt = n["norm"].split()
            if any(toks[i:i + len(nt)] == nt for i in range(len(toks) - len(nt) + 1)):
                return n
    return None


# ---------------------------------------------------------------- outputs
def write_workbook(out, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    HEAD = PatternFill("solid", fgColor="1F3864")
    WHITE = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    BOLD = Font(name="Arial", bold=True, size=10)
    REG = Font(name="Arial", size=10)
    TITLE = Font(name="Arial", bold=True, size=13, color="1F3864")
    VCOLOR = {"NET_NEW": "1F7A1F", "PAUSED_TARGET": "C07A00", "DORMANT_TARGET": "C07A00",
              "ACTIVE_TARGET": "C00000", "NEGATIVE_BLOCKED": "7030A0"}

    wb = Workbook(); ws = wb.active; ws.title = "Coverage"
    ws["A1"] = "MS Home expansion -- structure-side coverage (acct 2)"; ws["A1"].font = TITLE
    ws.merge_cells("A1:O1")
    cols = ["keyword", "planner_ad_group", "vol_band", "vol_mid", "suggested_bid",
            "impr_share", "brand", "bid_status", "dnb", "home_qualified",
            "theme", "theme_live", "structure_verdict", "match_basis", "matched_instances"]
    for j, h in enumerate(cols, 1):
        c = ws.cell(3, j, h); c.fill = HEAD; c.font = WHITE
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    r = 4
    for row in rows:
        for j, k in enumerate(cols, 1):
            c = ws.cell(r, j, row.get(k))
            c.font = BOLD if k == "keyword" else REG
            if k == "structure_verdict":
                c.font = Font(name="Arial", size=10, bold=True,
                              color=VCOLOR.get(row.get(k), "000000"))
        r += 1
    widths = {"A": 36, "B": 18, "C": 11, "D": 9, "E": 11, "F": 9, "G": 16, "H": 12,
              "I": 6, "J": 12, "K": 11, "L": 10, "M": 17, "N": 13, "O": 60}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = "A3:%s%d" % (get_column_letter(len(cols)), r - 1)
    wb.save(out)


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--planner-csv", required=True)
    p.add_argument("--db", default=r"C:\Projects\google-ads-ops\db\paid_search.db")
    p.add_argument("--negatives-csv", default=None,
                   help="MS negatives export (not in DB); adds NEGATIVE_BLOCKED verdicts")
    p.add_argument("--tagger-dir", default=None, help="folder containing brand_tagger.py")
    p.add_argument("--out", default=r"C:\Projects\google-ads-ops\analysis\ms-home-kw-expansion\output"
                                r"\ms_expansion_coverage_%s.xlsx" % date.today().isoformat())    
    a = p.parse_args()

    conn = sqlite3.connect("file:%s?mode=ro" % a.db, uri=True)  # READ-ONLY
    conn.row_factory = sqlite3.Row

    tagger, tdir = load_tagger(a.tagger_dir, a.db)
    print("brand tagger: %s" % (("loaded from %s (%s)" % (tdir, " | ".join(tagger.sources)))
                                if tagger else "NOT FOUND -- brand columns will be empty; "
                                "pass --tagger-dir"))
    by_canon, by_loose, n_struct, basis = load_structure(conn)
    print("acct-2 keyword_current rows: %s (%s distinct canonical texts)"
          % (f"{n_struct:,}", f"{len(by_canon):,}"))
    print("canonical match basis: %s" % basis)
    live_themes = live_generic_themes(conn)
    print("live generic themes (from ad-group naming): %s" % ", ".join(sorted(live_themes)))
    negs = load_negatives(a.negatives_csv)
    print("negatives overlay: %s" % ("%d rows from %s" % (len(negs), a.negatives_csv)
                                     if negs else "none supplied (MS negatives are NOT in the DB)"))

    out_rows, counts = [], defaultdict(int)
    n_canon = n_variant = 0
    with open(a.planner_csv, newline="", encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            kw = (row.get("Keyword") or "").strip()
            if not kw:
                continue
            n9, nl = normalize_09(kw), normalize_loose(kw)
            brand = bid = cat = src = ""
            dnb = False
            if tagger:
                brand, bid, cat, dnb, src = tagger.tag(n9)
                if not brand:                          # apostrophe-bearing edge cases
                    brand, bid, cat, dnb, src = tagger.tag(nl)
            homeq = bool(HOME_RE.search(nl)) or bool(brand)
            th = theme_of(nl)
            th_live = "LIVE" if any(th.lower() == t.lower() for t in live_themes) else "no live group"

            matches, mbasis = by_canon.get(n9, []), "09-canonical"
            if not matches:
                lv = by_loose.get(nl, [])
                if lv:
                    matches, mbasis = lv, "close-variant"
            if matches:
                n_canon += mbasis == "09-canonical"
                n_variant += mbasis == "close-variant"
            else:
                mbasis = ""
            verdict, hits = verdict_for(matches)
            neg = negative_hit(nl, negs)
            if neg and verdict == "NET_NEW":
                verdict = "NEGATIVE_BLOCKED"
            inst = "; ".join(sorted({"%s/%s%s [%s > %s]"
                                     % (m["match_type"], m["keyword_status"],
                                        "*SERVING*" if m["is_serving"] == 1 else "",
                                        (m["campaign_name"] or "")[:24],
                                        (m["ad_group_name"] or "")[:30]) for m in hits}))
            if mbasis == "close-variant" and hits:
                inst = "variant of: " + "; ".join(sorted({m["keyword_text"] for m in hits})[:3]) \
                       + " | " + inst
            if neg:
                inst = ("NEG %s/%s; " % (neg["match"], neg["level"])) + inst
            counts[verdict] += 1
            band = (row.get("Average monthly searches") or "").strip()
            out_rows.append({
                "keyword": kw, "planner_ad_group": row.get("Ad group"),
                "vol_band": band, "vol_mid": VOL_MID.get(band, 0),
                "suggested_bid": row.get("Suggested Bid"),
                "impr_share": row.get("Ad impr. share"),
                "brand": brand, "bid_status": bid, "dnb": int(bool(dnb)),
                "home_qualified": "yes" if homeq else "UNQUALIFIED",
                "theme": th, "theme_live": th_live,
                "structure_verdict": verdict, "match_basis": mbasis,
                "matched_instances": inst})
    conn.close()

    order = {"NET_NEW": 0, "NEGATIVE_BLOCKED": 1, "PAUSED_TARGET": 2,
             "DORMANT_TARGET": 3, "ACTIVE_TARGET": 4}
    out_rows.sort(key=lambda x: (order.get(x["structure_verdict"], 9),
                                 1 if x["brand"] else 0, -x["vol_mid"]))

    print("\n" + "=" * 78)
    print("  STRUCTURE VERDICTS (%d planner keywords)" % len(out_rows))
    print("=" * 78)
    for v in sorted(counts, key=lambda x: order.get(x, 9)):
        print("  %-18s %d" % (v, counts[v]))
    n_brand = sum(1 for x in out_rows if x["brand"])
    n_dnb = sum(1 for x in out_rows if x["dnb"])
    n_unq = sum(1 for x in out_rows if x["home_qualified"] == "UNQUALIFIED")
    print("  brand-tagged: %d (of which DNB: %d) | UNQUALIFIED generic: %d"
          % (n_brand, n_dnb, n_unq))
    print("  match tiers: 09-canonical=%d, close-variant only=%d" % (n_canon, n_variant))
    print("  (close-variant = punctuation/format differences MS treats as one target;")
    print("   same verdict, softer basis -- see match_basis column)")

    print("\n  NET_NEW, home-qualified, non-brand -- the true expansion shortlist:")
    for x in out_rows:
        if (x["structure_verdict"] == "NET_NEW" and not x["brand"]
                and x["home_qualified"] == "yes"):
            print("    %-42s %-10s theme=%-10s (%s)"
                  % (x["keyword"], x["vol_band"], x["theme"], x["theme_live"]))

    write_workbook(a.out, out_rows)
    base = os.path.splitext(a.out)[0]
    with open(base + ".csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(out_rows[0].keys()))
        w.writeheader(); w.writerows(out_rows)
    print("\nwrote %s and %s.csv" % (a.out, base))


if __name__ == "__main__":
    main()
