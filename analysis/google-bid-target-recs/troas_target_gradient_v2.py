"""
tROAS Target Gradient v2 — Google Home Insurance
═══════════════════════════════════════════════════════════════════════
Per-campaign target-change gradient for Ben: retrieves each campaign's
ACTUAL latest tROAS setting, then lays out a ladder of candidate targets
N points below and above it (default ±4pts in 1pt steps), with the
expected delivered economics of every rung. Ben's proposed -3pt move is
one highlighted rung. Output is a workbook with the ladder table, a
small-multiples chart (one panel per campaign visualizing the gradient),
the state decomposition, and a Notes audit sheet.

TARGET RETRIEVAL (v2 change — the fix for the 8/12 run):
  campaigns.target_roas came back NULL for all 7 campaigns in the
  2026-08-12 run — the snapshot column does not carry the campaign-level
  targets. Authority is therefore the change-history data model, same
  pattern as campaign_performance_visualizations_google V2:
    latest TARGET_ROAS event per campaign_id in campaign_change_events
    (platform='google', account_id=1), new_value normalized to a decimal
    (116% -> 1.16, 113 -> 1.13). campaigns.target_roas is fallback only.
  Campaigns resolve by immutable campaign_id (the 7 NB-SEARCH ids),
  never by name.

DATA RULES (unchanged from v1 / charter):
  - Realized: campaign_performance_daily_search, device='ALL' rows.
    Revenue = native conversion_value (ruled basis, <0.1% vs routed).
    ROAS = SUM(conversion_value)/SUM(cost_micros/1e6); summed numerators.
  - States: campaign_performance_geo_daily_search — no device='ALL'
    rollup there; device rows summed.
  - No keyword tables, no base (non-_search) tables. DB opened mode=ro.

GRADIENT SEMANTICS:
  Rungs run current-span .. current+span points (1pt = 0.01 on the
  ratio). Negative rungs loosen the target (volume direction); positive
  rungs tighten it (efficiency direction). Expected delivered ROAS per
  rung = post-freeze delivery ratio (realized/current) x candidate
  target — a first-order directional read, NOT a forecast; simulator
  curvature (Layer 3, from the Ads UI after the first clean week)
  adjudicates between adjacent rungs. Every rung is checked against the
  true break-even line expressed in DB terms, 1/(1+correction): the
  10-13%% value understatement is applied to the THRESHOLD, never to
  the DB's realized numbers.

USAGE:
  python troas_target_gradient_v2.py
  python troas_target_gradient_v2.py --span-points 4 --step-points 1
  python troas_target_gradient_v2.py --ben-points 3 --correction-pct 13
  python troas_target_gradient_v2.py --window-start 2026-07-30 --lag-days 3

OUTPUT: troas_target_gradient_<YYYY-MM-DD>.xlsx
  Charts / Gradient / State Decomp / Notes
"""

import argparse
import io
import sqlite3
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

try:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl (with Pillow) is required: pip install openpyxl pillow")

DEFAULT_DB_PATH = Path(r"C:\Projects\google-ads-ops\db\paid_search.db")
ACCOUNT_ID = 1

# The 7 NB-SEARCH campaigns, by immutable id (names are rename-prone).
ENABLED_CAMPAIGN_IDS = [
    "23385188622",   # CALIFORNIA-NB
    "23394808297",   # TEXAS-NB
    "23394808300",   # FLORIDA-NB
    "23428389081",   # TIER1 (RPL shell)
    "23496334752",   # TIER2
    "23530860942",   # TIER3
    "23586637678",   # TIER4
]

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
CAMP_FILL = PatternFill("solid", fgColor="D9E2F3")
CURR_FILL = PatternFill("solid", fgColor="E2EFDA")     # current target rung
BEN_FILL = PatternFill("solid", fgColor="FFF2CC")      # Ben's -3pt rung
NEG_FILL = PatternFill("solid", fgColor="F8CBAD")      # below true break-even
SOFT_FONT = Font(italic=True, color="808080", size=9)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


# ─────────────────────────────────────────────────────────────────────
# Data access
# ─────────────────────────────────────────────────────────────────────

def connect_ro(db_path: Path) -> sqlite3.Connection:
    conn = sqlite3.connect(f"file:{db_path.as_posix()}?mode=ro", uri=True)
    conn.row_factory = sqlite3.Row
    return conn


def normalize_target(raw):
    """116% -> 1.16, 1.16 -> 1.16, 113 -> 1.13, '' -> None."""
    if raw is None:
        return None
    s = str(raw).strip().rstrip(",").rstrip('"').strip()
    if not s:
        return None
    if s.endswith("%"):
        try:
            return float(s.rstrip("%")) / 100.0
        except ValueError:
            return None
    try:
        v = float(s)
    except ValueError:
        return None
    return v / 100.0 if v > 10 else v


def fetch_campaigns(conn):
    ph = ",".join("?" for _ in ENABLED_CAMPAIGN_IDS)
    sql = f"""
        SELECT campaign_id, campaign_name, target_roas
        FROM campaigns
        WHERE account_id = ?
          AND campaign_id IN ({ph})
          AND status = 'ENABLED'
        ORDER BY campaign_name
    """
    rows = conn.execute(sql, [ACCOUNT_ID, *ENABLED_CAMPAIGN_IDS]).fetchall()
    return [{
        "campaign_id": str(r["campaign_id"]),
        "name": r["campaign_name"],
        "snapshot_target": normalize_target(r["target_roas"]),
    } for r in rows]


def fetch_latest_targets(conn, campaigns):
    """Latest TARGET_ROAS event per campaign from campaign_change_events
    (the target authority). Sets camp['target'] and camp['target_source']
    ('event <date>' | 'snapshot' | None)."""
    ph = ",".join("?" for _ in campaigns)
    ids = [c["campaign_id"] for c in campaigns]
    sql = f"""
        SELECT e.campaign_id, e.change_date, e.new_value
        FROM campaign_change_events e
        WHERE e.platform = 'google'
          AND e.account_id = ?
          AND e.change_type = 'TARGET_ROAS'
          AND e.campaign_id IN ({ph})
        ORDER BY e.change_date
    """
    latest = {}
    try:
        for r in conn.execute(sql, [ACCOUNT_ID, *ids]):
            v = normalize_target(r["new_value"])
            if v is not None:
                latest[str(r["campaign_id"])] = (v, str(r["change_date"])[:10])
    except sqlite3.Error as exc:
        print(f"  WARN: could not read campaign_change_events: {exc}")
    for c in campaigns:
        if c["campaign_id"] in latest:
            c["target"], evt_date = latest[c["campaign_id"]]
            c["target_source"] = f"event {evt_date}"
        elif c["snapshot_target"] is not None:
            c["target"] = c["snapshot_target"]
            c["target_source"] = "snapshot (no TARGET_ROAS events)"
        else:
            c["target"] = None
            c["target_source"] = "UNRESOLVED — supply manually"
    return campaigns


def max_data_date(conn):
    row = conn.execute(
        """SELECT MAX(p.date) FROM campaign_performance_daily_search p
           JOIN campaigns c ON c.campaign_id = p.campaign_id
           WHERE c.account_id = ?""", [ACCOUNT_ID]).fetchone()
    return datetime.strptime(row[0][:10], "%Y-%m-%d").date()


def campaign_perf(conn, ids, start, end):
    ph = ",".join("?" for _ in ids)
    sql = f"""
        SELECT campaign_id,
               SUM(cost_micros) / 1e6 AS cost,
               SUM(conversions)       AS conv,
               SUM(conversion_value)  AS revenue
        FROM campaign_performance_daily_search
        WHERE campaign_id IN ({ph})
          AND device = 'ALL'
          AND date BETWEEN ? AND ?
        GROUP BY campaign_id
    """
    rows = conn.execute(sql, [*ids, start.isoformat(), end.isoformat()])
    return {str(r["campaign_id"]): dict(r) for r in rows}


def state_perf(conn, ids, start, end):
    ph = ",".join("?" for _ in ids)
    sql = f"""
        SELECT campaign_id, geo_state,
               SUM(cost_micros) / 1e6 AS cost,
               SUM(conversions)       AS conv,
               SUM(conversion_value)  AS revenue
        FROM campaign_performance_geo_daily_search
        WHERE campaign_id IN ({ph})
          AND date BETWEEN ? AND ?
        GROUP BY campaign_id, geo_state
    """
    rows = conn.execute(sql, [*ids, start.isoformat(), end.isoformat()])
    out = {}
    for r in rows:
        out.setdefault(str(r["campaign_id"]), []).append(dict(r))
    return out


# ─────────────────────────────────────────────────────────────────────
# Gradient math
# ─────────────────────────────────────────────────────────────────────

def safe_div(a, b):
    return a / b if b else None


def build_ladder(camp, perf, args, db_breakeven):
    cost = perf.get("cost", 0) or 0
    revenue = perf.get("revenue", 0) or 0
    realized = safe_div(revenue, cost)
    target = camp["target"]
    ratio = safe_div(realized, target) if (realized and target) else None
    gp = revenue - cost

    rungs = []
    if target is not None:
        n_steps = int(round(args.span_points / args.step_points))
        for k in range(-n_steps, n_steps + 1):
            pts = k * args.step_points
            t = target + pts / 100.0
            exp = ratio * t if ratio else None
            rungs.append({
                "points": pts,
                "target": t,
                "expected_roas": exp,
                "below_breakeven": exp is not None and exp < db_breakeven,
                "is_current": k == 0,
                "is_ben": abs(pts + args.ben_points) < 1e-9,
            })
    return {"cost": cost, "revenue": revenue, "conv": perf.get("conv"),
            "gp": gp, "realized": realized, "ratio": ratio, "rungs": rungs}


def short_name(name):
    n = name.replace("HOME INSURANCE", "").replace("_STATE-GROUP_", "")
    n = n.replace("-NB-SEARCH-ALL-USA-STE", "").replace("_COMBINED_TROAS_SEARCH_ALLDEVICE_USA-STE", "")
    return n.strip("-_ ") or name


# ─────────────────────────────────────────────────────────────────────
# Chart — small multiples, one panel per campaign
# ─────────────────────────────────────────────────────────────────────

def build_gradient_chart_png(campaigns, ladders, args, db_breakeven):
    plotable = [c for c in campaigns
                if ladders[c["campaign_id"]]["rungs"]
                and ladders[c["campaign_id"]]["ratio"]]
    if not plotable:
        return None
    n = len(plotable)
    ncols = 2
    nrows = (n + ncols - 1) // ncols
    fig, axes = plt.subplots(nrows, ncols, figsize=(11, 3.1 * nrows))
    axes = axes.flatten() if n > 1 else [axes]

    for ax, camp in zip(axes, plotable):
        lad = ladders[camp["campaign_id"]]
        xs = [r["target"] for r in lad["rungs"]]
        ys = [r["expected_roas"] for r in lad["rungs"]]
        target, realized = camp["target"], lad["realized"]
        ben_t = target - args.ben_points / 100.0

        # shaded GP-negative zone (true economics, in DB terms)
        ax.axhspan(min(min(ys), db_breakeven) - 0.02, db_breakeven,
                   color="#F8CBAD", alpha=0.45, zorder=0)
        ax.axhline(db_breakeven, color="#C0504D", linestyle="--",
                   linewidth=1.2,
                   label=f"true break-even (DB {db_breakeven:.2f})")
        ax.plot(xs, ys, marker="o", markersize=4, color="#1F4E79",
                linewidth=1.8, label="expected delivered ROAS", zorder=3)
        ax.axvline(target, color="#404040", linestyle="-.", linewidth=1.2,
                   label=f"current target {target:.2f}")
        ax.plot([target], [realized], marker="D", markersize=8,
                color="#2E7D32", zorder=4,
                label=f"realized {realized:.2f}")
        ax.axvline(ben_t, color="#BF8F00", linestyle=":", linewidth=1.6,
                   label=f"Ben -{args.ben_points:g}pt ({ben_t:.2f})")

        ax.set_title(short_name(camp["name"]), fontsize=9, fontweight="bold")
        ax.set_xlabel("candidate tROAS target", fontsize=7.5)
        ax.set_ylabel("expected ROAS (DB)", fontsize=7.5)
        ax.tick_params(labelsize=7)
        ax.grid(True, alpha=0.25)
        ax.legend(fontsize=6, loc="upper left", framealpha=0.85)

    for ax in axes[len(plotable):]:
        ax.set_visible(False)
    fig.suptitle(
        f"tROAS gradient ±{args.span_points:g}pts around current target — "
        f"expected delivered ROAS at delivery ratio held constant "
        f"(directional, not a forecast)", fontsize=10)
    fig.tight_layout(rect=[0, 0, 1, 0.97])
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=150)
    plt.close(fig)
    buf.seek(0)
    return buf


# ─────────────────────────────────────────────────────────────────────
# Workbook
# ─────────────────────────────────────────────────────────────────────

def style_header(ws, row, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center")


def write_gradient_sheet(ws, campaigns, ladders, args, db_breakeven,
                         win_start, win_end):
    ws.title = "Gradient"
    ws.sheet_view.showGridLines = False
    headers = ["Campaign", "Target source", "Rung (pts)", "Target",
               "Expected delivered ROAS", "vs true break-even",
               "Post-freeze realized", "Delivery ratio", "Spend (win)",
               "Conv (win)", "GP (win, DB $)", "Simulator read",
               "Rollback trigger"]
    ws.append([f"tROAS Target Gradient — Google Home — window {win_start} "
               f"→ {win_end} (trailing {args.lag_days}d excluded as "
               f"lag-soft) — PROVISIONAL until a full clean post-freeze "
               f"week"])
    ws.cell(row=1, column=1).font = Font(bold=True, size=11)
    ws.append([f"Targets from campaign_change_events (latest TARGET_ROAS "
               f"per id); true break-even in DB terms = "
               f"1/(1+{args.correction_pct:g}%) = {db_breakeven:.3f}. DB "
               f"realized never inflated — the threshold carries the "
               f"correction."])
    ws.cell(row=2, column=1).font = SOFT_FONT
    ws.append([])
    ws.append(headers)
    style_header(ws, 4, len(headers))

    r = 5
    for camp in campaigns:
        lad = ladders[camp["campaign_id"]]
        if not lad["rungs"]:
            ws.append([camp["name"], camp["target_source"], "—", None, None,
                       "target unresolved", lad["realized"], None,
                       lad["cost"], lad["conv"], lad["gp"], "", ""])
            ws.cell(row=r, column=1).fill = CAMP_FILL
            r += 2
            ws.append([])
            continue
        first = True
        for rung in lad["rungs"]:
            rollback = ""
            if first and lad["realized"]:
                floor = max(db_breakeven, 0.90 * lad["realized"])
                rollback = (f"Revert if rolling-7d realized ROAS < "
                            f"{floor:.2f} (DB) for 5+ days, or CPL "
                            f"rises >25% vs post-freeze baseline")
            ws.append([
                camp["name"] if first else "",
                camp["target_source"] if first else "",
                f"{rung['points']:+g}" + (" (current)" if rung["is_current"]
                                          else " (Ben)" if rung["is_ben"] else ""),
                rung["target"],
                rung["expected_roas"],
                "BELOW" if rung["below_breakeven"] else "above",
                lad["realized"],
                lad["ratio"],
                lad["cost"],
                lad["conv"],
                lad["gp"],
                "",
                rollback,
            ])
            for c in range(1, len(headers) + 1):
                cell = ws.cell(row=r, column=c)
                cell.border = BORDER
                if c in (4, 5, 7, 8):
                    cell.number_format = "0.00"
                elif c in (9, 10, 11):
                    cell.number_format = "#,##0"
            if first:
                ws.cell(row=r, column=1).fill = CAMP_FILL
                ws.cell(row=r, column=1).font = Font(bold=True, size=9)
            fill = (CURR_FILL if rung["is_current"]
                    else BEN_FILL if rung["is_ben"] else None)
            if fill:
                for c in range(3, 7):
                    ws.cell(row=r, column=c).fill = fill
            if rung["below_breakeven"]:
                ws.cell(row=r, column=5).fill = NEG_FILL
                ws.cell(row=r, column=6).fill = NEG_FILL
            first = False
            r += 1
        r += 1
        ws.append([])

    widths = [46, 22, 12, 9, 13, 12, 12, 11, 11, 9, 11, 14, 50]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_state_sheet(ws, campaigns, states):
    ws.title = "State Decomp"
    ws.sheet_view.showGridLines = False
    headers = ["Campaign", "State", "Spend", "Spend share", "Conv",
               "Revenue (DB $)", "Realized ROAS", "GP (DB $)"]
    ws.append(headers)
    style_header(ws, 1, len(headers))
    r = 2
    for camp in campaigns:
        rows = states.get(camp["campaign_id"], [])
        total_cost = sum(x["cost"] or 0 for x in rows)
        rows.sort(key=lambda x: (x["revenue"] or 0) - (x["cost"] or 0))
        first = True
        for x in rows:
            cost, rev = x["cost"] or 0, x["revenue"] or 0
            gp = rev - cost
            ws.append([camp["name"] if first else "", x["geo_state"], cost,
                       safe_div(cost, total_cost), x["conv"], rev,
                       safe_div(rev, cost), gp])
            for c in range(1, len(headers) + 1):
                cell = ws.cell(row=r, column=c)
                cell.border = BORDER
                if c in (3, 5, 6, 8):
                    cell.number_format = "#,##0"
                elif c == 4:
                    cell.number_format = "0.0%"
                elif c == 7:
                    cell.number_format = "0.00"
            if first:
                ws.cell(row=r, column=1).fill = CAMP_FILL
                ws.cell(row=r, column=1).font = Font(bold=True, size=9)
            if gp < 0:
                ws.cell(row=r, column=8).fill = NEG_FILL
            first = False
            r += 1
        r += 1
        ws.append([])
    for i, w in enumerate([46, 8, 10, 11, 8, 12, 12, 10], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_notes_sheet(ws, args, win_start, win_end, soft_read,
                      db_breakeven, campaigns):
    ws.title = "Notes"
    lines = [
        ("Run", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Realized window", f"{win_start} → {win_end} (trailing "
                            f"{args.lag_days}d excluded as lag-soft)"),
        ("Provisional flag", "numbers provisional until a full clean "
                             "Mon–Sun post-freeze week exists"),
        ("Target authority", "campaign_change_events latest TARGET_ROAS "
                             "new_value per campaign_id (snapshot "
                             "campaigns.target_roas NULL for these ids "
                             "as of 8/12; used as fallback only)"),
        ("Revenue basis", "campaign_performance_daily_search."
                          "conversion_value, device='ALL' (ruled basis, "
                          "<0.1% vs routed truth)"),
        ("State basis", "campaign_performance_geo_daily_search, device "
                        "rows summed (no ALL rollup)"),
        ("Gradient", f"±{args.span_points:g}pts around current target in "
                     f"{args.step_points:g}pt steps; Ben rung = "
                     f"-{args.ben_points:g}pts (highlighted amber; "
                     f"current rung green)"),
        ("Correction", f"{args.correction_pct:g}% value understatement "
                       f"applied to break-even threshold only "
                       f"(1/(1+c) = {db_breakeven:.3f} DB ROAS)"),
        ("Expected ROAS", "delivery ratio × candidate target — "
                          "directional first-order read, NOT a forecast; "
                          "simulator curvature (UI, after first clean "
                          "week) adjudicates adjacent rungs"),
        ("Schedule (Layer 4)", "one move per campaign, ≥1 week apart, "
                               "batched same-day; TARGET_ROAS events "
                               "auto-land in change history for "
                               "post-change anchoring"),
        ("Soft trailing window", soft_read),
    ]
    for c in campaigns:
        t = f"{c['target']:.2f}" if c["target"] is not None else "—"
        lines.append((f"Target: {short_name(c['name'])}",
                      f"{t}  [{c['target_source']}]"))
    for k, v in lines:
        ws.append([k, v])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=9)
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 110
    for row in ws.iter_rows():
        row[1].alignment = Alignment(wrap_text=True, vertical="top")


# ─────────────────────────────────────────────────────────────────────

def main():
    p = argparse.ArgumentParser()
    p.add_argument("--db-path", default=str(DEFAULT_DB_PATH))
    p.add_argument("--window-start", default="2026-07-30")
    p.add_argument("--lag-days", type=int, default=3)
    p.add_argument("--span-points", type=float, default=4.0,
                   help="points below AND above current target")
    p.add_argument("--step-points", type=float, default=1.0,
                   help="rung increment in points")
    p.add_argument("--ben-points", type=float, default=3.0,
                   help="Ben's proposed reduction, in points")
    p.add_argument("--correction-pct", type=float, default=11.5)
    p.add_argument("--output", default=None)
    args = p.parse_args()

    db_path = Path(args.db_path)
    if not db_path.exists():
        sys.exit(f"DB not found: {db_path}")
    conn = connect_ro(db_path)

    db_breakeven = 1.0 / (1.0 + args.correction_pct / 100.0)

    campaigns = fetch_campaigns(conn)
    if not campaigns:
        sys.exit("No campaigns resolved — check ENABLED_CAMPAIGN_IDS.")
    fetch_latest_targets(conn, campaigns)
    print(f"Resolved {len(campaigns)} campaigns:")
    for c in campaigns:
        t = f"{c['target']:.2f}" if c["target"] is not None else "—"
        print(f"  {c['campaign_id']}  target={t:6s} [{c['target_source']}]"
              f"  {short_name(c['name'])}")

    max_date = max_data_date(conn)
    win_start = datetime.strptime(args.window_start, "%Y-%m-%d").date()
    win_end = max_date - timedelta(days=args.lag_days)
    if win_end <= win_start:
        sys.exit(f"Window empty: {win_start} → {win_end}")
    print(f"\nRealized window: {win_start} → {win_end} "
          f"(max data date {max_date}, trailing {args.lag_days}d soft)")

    ids = [c["campaign_id"] for c in campaigns]
    perf = campaign_perf(conn, ids, win_start, win_end)
    states = state_perf(conn, ids, win_start, win_end)
    soft = campaign_perf(conn, ids, win_end + timedelta(days=1), max_date)
    soft_cost = sum(v["cost"] or 0 for v in soft.values())
    soft_rev = sum(v["revenue"] or 0 for v in soft.values())
    soft_read = (f"{win_end + timedelta(days=1)} → {max_date}: spend "
                 f"{soft_cost:,.0f}, revenue {soft_rev:,.0f}, ROAS "
                 f"{safe_div(soft_rev, soft_cost) or 0:.2f} (SOFT — lag "
                 f"restates upward)")

    ladders = {}
    print(f"\nGradient (DB terms; break-even {db_breakeven:.3f}):")
    for camp in campaigns:
        lad = build_ladder(camp, perf.get(camp["campaign_id"], {}),
                           args, db_breakeven)
        ladders[camp["campaign_id"]] = lad
        rz = f"{lad['realized']:.2f}" if lad["realized"] else "—"
        rt = f"{lad['ratio']:.2f}" if lad["ratio"] else "—"
        print(f"  {short_name(camp['name']):12s} realized {rz}  "
              f"delivery-ratio {rt}  GP {lad['gp']:>9,.0f}")

    wb = Workbook()
    charts_ws = wb.active
    charts_ws.title = "Charts"
    charts_ws.sheet_view.showGridLines = False
    png = build_gradient_chart_png(campaigns, ladders, args, db_breakeven)
    if png:
        charts_ws.add_image(XLImage(png), "A1")
    else:
        charts_ws["A1"] = "No plottable campaigns (targets unresolved)."

    write_gradient_sheet(wb.create_sheet(), campaigns, ladders, args,
                         db_breakeven, win_start, win_end)
    write_state_sheet(wb.create_sheet(), campaigns, states)
    write_notes_sheet(wb.create_sheet(), args, win_start, win_end,
                      soft_read, db_breakeven, campaigns)

    out = args.output or f"troas_target_gradient_{date.today().isoformat()}.xlsx"
    wb.save(out)
    conn.close()
    print(f"\nSaved: {out}")


if __name__ == "__main__":
    main()
