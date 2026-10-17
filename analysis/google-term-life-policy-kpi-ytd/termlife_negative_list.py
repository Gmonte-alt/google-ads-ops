"""
termlife_negative_list.py  (v1)
================================
Negative-keyword proposal builder for Term Life (account_id=29).

Produces Negative_Keyword_Proposal.xlsx — a RATIFICATION-GATED proposal.
Nothing in this script changes the account; it recommends.

Structure of the proposal
-------------------------
Campaigns (CR12 / CR3 / CR45) are zip-income tier splits of the SAME
keyword portfolios, so intent does not vary by campaign. Anything
campaign-level therefore goes into SHARED NEGATIVE LISTS applied to all
three campaigns, and the exercise includes re-mapping existing negatives
to the level where they belong:

  LIST 1  "TL - Disqualifying Product Intent"   (shared, all 3 campaigns)
          senior / final-expense / burial / guaranteed-acceptance /
          age-band etc. — term-life product mismatch.
  LIST 2  "TL - Service & Navigational"         (shared, all 3 campaigns)
          login / cancel / claims / phone number etc.
  AD-GROUP routing negatives (NOT shared — these depend on the ad
          group's own portfolio): carrier brands on the Ethos ad group;
          'ethos' on carrier ad groups (gap-fill of the existing
          patchwork).
  SURFACED, NOT NEGATED: whole life / universal life — sized in its own
          sheet as potential demand for a dedicated campaign.
  PROTECTED (never negated): no exam / no medical exam (Ethos selling
          proposition), reviews/legit/cost/calculator curiosity layer,
          close variants of own exact terms.

Inputs
------
  paid_search.db (read-only)
  Ads Editor negatives export (UTF-16 TSV) in input\  — current live
  negatives at account / campaign / ad-group level.

Caveats (also written into the workbook)
----------------------------------------
  sqp hides low-volume terms -> all sizing is a LOWER BOUND; token-level
  negatives also block the hidden tail, per-query negatives cannot.
  Our matching simulation approximates Google semantics (negative broad =
  all words present, any order, NO close variants -> misspellings are not
  blocked; phrase = contiguous sequence; exact = full match).
  Export has no add dates -> "last_served" infers how recently each
  existing negative became effective.
  Sep-29+ contaminated (schedule-cut incident); Oct MTD confirmation only.
"""
import glob
import os
import re
import sqlite3
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.drawing.text import RichTextProperties
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

DB_PATH = r"C:\Projects\google-ads-ops\db\paid_search.db"
ANALYSIS_DIR = (r"C:\Projects\google-ads-ops\analysis"
                r"\google-term-life-policy-kpi-ytd")
OUT_XLSX = os.path.join(ANALYSIS_DIR, "output",
                        "Negative_Keyword_Proposal.xlsx")

ACCOUNT_ID = 29
Y0 = "2026-01-01"            # full-YTD lookback (effective-since)
W0, W1 = "2026-06-01", "2026-09-29"   # primary evidence window
O0, O1 = "2026-10-01", "2026-11-01"   # Oct MTD confirmation
SPEND_FLOOR = 25.0           # min window spend for a token to be recommended
LIST1 = "TL - Disqualifying Product Intent"
LIST2 = "TL - Service & Navigational"

NAVY, LIGHT = "1F3864", "D9E2F2"
SERIES_HEX = ["2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4", "008300"]

pd.set_option("display.width", 240)
pd.set_option("display.max_columns", 40)
pd.set_option("display.max_rows", 300)
pd.set_option("display.float_format", lambda x: f"{x:,.1f}")

# ---------------------------------------------------------------- lexicons
# token -> regex (word-boundary). A token is RECOMMENDED only when queries
# matching it carry >= SPEND_FLOOR in the window AND it passes the
# collision check against the account's own keywords.
DISQUALIFY = {
    "senior": r"\bseniors?\b",
    "elderly": r"\belderly\b",
    "final expense": r"\bfinal expenses?\b",
    "burial": r"\bburial\b",
    "funeral": r"\bfuneral\b",
    "guaranteed acceptance": r"\bguaranteed acceptance\b",
    "guaranteed issue": r"\bguaranteed issue\b",
    "no questions": r"\bno questions\b",
    "over 50": r"\bover\s*50\b", "over 55": r"\bover\s*55\b",
    "over 60": r"\bover\s*60\b", "over 65": r"\bover\s*65\b",
    "over 70": r"\bover\s*70\b", "over 75": r"\bover\s*75\b",
    "over 80": r"\bover\s*80\b", "over 85": r"\bover\s*85\b",
    "for parents": r"\bfor (my )?parents?\b",
    "for mom": r"\bfor (my )?(mom|mother)\b",
    "for dad": r"\bfor (my )?(dad|father)\b",
}
SERVICE = {
    "login": r"\blog\s*in\b|\blogin\b",
    "sign in": r"\bsign\s*in\b",
    "portal": r"\bportal\b",
    "my account": r"\bmy account\b",
    "cancel": r"\bcancel(l?ation|l?ing|l?ed)?\b",
    "claim": r"\bclaims?\b",
    "customer service": r"\bcustomer (service|care|support)\b",
    "phone number": r"\bphone\b|\btelephone\b|\b1?800\b",
    "contact": r"\bcontact\b",
    "payment": r"\bpayment\b|\bpay (my )?bill\b|\bbill pay\b",
    "careers": r"\bcareers?\b|\bjobs?\b|\bsalary\b|\bhiring\b",
    "agent login": r"\bagent (login|portal)\b",
    "policy lookup": r"\bpolicy (lookup|number|status)\b",
    "address": r"\bheadquarters\b|\bmailing address\b",
}
SURFACE_ONLY = {            # sized, never negated (possible new campaign)
    "whole life": r"\bwhole life\b",
    "universal life": r"\buniversal life\b|\biul\b",
    "variable life": r"\bvariable life\b",
}
PROTECTED = {               # never allowed to become a negative
    "no exam": r"\bno (medical )?exam\b|\bwithout (a )?(medical )?exam\b"
               r"|\bno physical\b|\bno medical\b",
    "reviews": r"\breviews?\b|\bratings?\b",
    "legit": r"\blegit(imate)?\b",
    "cost": r"\bcost\b|\bcalculator\b|\brates?\b|\bprices?\b",
    "good": r"\bis .* good\b|\bworth it\b",
    "term": r"\bterm\b",
}
CARRIERS = {                # carrier/competitor brand tokens
    "colonial penn": r"\bcolonial penn\b|\bmycolonialpenn\b",
    "mutual of omaha": r"\bmutual of omaha\b",
    "aarp": r"\baarp\b",
    "globe life": r"\bglobe life\b",
    "global life": r"\bglobal life\b",
    "metlife": r"\bmet\s*life\b",
    "new york life": r"\bnew york life\b",
    "northwestern mutual": r"\bnorthwestern mutual\b",
    "prudential": r"\bprudential\b",
    "lincoln financial": r"\blincoln (financial|heritage)\b",
    "big lou": r"\bbig lou\b",
    "aflac": r"\baflac\b",
    "state farm": r"\bstate farm\b",
    "fidelity": r"\bfidelity\b",
    "transamerica": r"\btransamerica\b",
    "banner life": r"\bbanner life\b",
    "protective": r"\bprotective\b",
    "john hancock": r"\bjohn hancock\b",
    "primerica": r"\bprimerica\b",
    "gerber": r"\bgerber\b",
    "foresters": r"\bforesters\b",
    "haven life": r"\bhaven life\b",
    "bestow": r"\bbestow\b",
    "ladder": r"\bladder\b",
    "policygenius": r"\bpolicy\s*genius\b",
    "lemonade": r"\blemonade\b",
    "geico": r"\bgeico\b",
    "progressive": r"\bprogressive\b",
    "usaa": r"\busaa\b",
    "allstate": r"\ballstate\b",
    "nationwide": r"\bnationwide\b",
    "massmutual": r"\bmass\s*mutual\b",
    "american income": r"\bamerican income\b",
    "pacific life": r"\bpacific life\b|\bpacific\b",
    "cigna": r"\bcigna\b",
    "voya": r"\bvoya\b",
    "principal": r"\bprincipal\b",
    "brighthouse": r"\bbrighthouse\b",
    "americo": r"\bamerico\b",
    "sagicor": r"\bsagicor\b",
}
ETHOS_RE = re.compile(r"\bethos\b")
LUMINARY_RE = re.compile(r"\bluminary\b")
CARRIER_RE = {k: re.compile(v) for k, v in CARRIERS.items()}
LEX_RE = {k: re.compile(v) for k, v in
          {**DISQUALIFY, **SERVICE, **SURFACE_ONLY}.items()}
PROT_RE = {k: re.compile(v) for k, v in PROTECTED.items()}


def hr(t):
    print("\n" + "=" * 110)
    print(t)
    print("=" * 110)


def meta(ws, title, source, window, caveat, takeaway=None):
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=13, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A2"] = f"Source: {source}   |   Window: {window}"
    ws["A3"] = f"Caveats: {caveat}"
    for r in (2, 3):
        ws[f"A{r}"].font = Font(italic=True, size=9)
    if takeaway:
        ws["A4"] = f"TAKEAWAY: {takeaway}"
        ws["A4"].font = Font(bold=True, size=10)
        ws["A4"].fill = PatternFill("solid", fgColor=LIGHT)
    ws.freeze_panes = "A6"


def write_block(ws, row, subtitle, df):
    ws.cell(row=row, column=1, value=subtitle).font = Font(bold=True,
                                                           size=11)
    ws.cell(row=row, column=1).fill = PatternFill("solid", fgColor=LIGHT)
    row += 1
    for j, c in enumerate(df.columns, start=1):
        cell = ws.cell(row=row, column=j, value=str(c))
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
    row += 1
    for _, r in df.iterrows():
        for j, v in enumerate(r, start=1):
            if isinstance(v, (np.floating, float)):
                v = None if pd.isna(v) else round(float(v), 1)
            elif isinstance(v, (np.integer,)):
                v = int(v)
            ws.cell(row=row, column=j, value=v)
        row += 1
    for j, c in enumerate(df.columns, start=1):
        width = max([len(str(c))] + [len(str(x)) for x in df[c].head(60)])
        ws.column_dimensions[get_column_letter(j)].width = min(width + 2, 52)
    return row + 2


def write_series_block(ws, start_col, df):
    for j, cname in enumerate(df.columns):
        ws.cell(row=1, column=start_col + j, value=str(cname))
        for i, v in enumerate(df[cname].tolist(), start=2):
            if isinstance(v, (np.floating, float)):
                v = None if pd.isna(v) else round(float(v), 2)
            elif isinstance(v, (np.integer,)):
                v = int(v)
            elif v is not None and not isinstance(v, (int, float)):
                v = str(v)
            ws.cell(row=i, column=start_col + j, value=v)
    return len(df)


def make_line_chart(ws_data, start_col, n_series, n_rows, title, y_title,
                    skip=1):
    ch = LineChart()
    ch.title = title
    data = Reference(ws_data, min_col=start_col + 1,
                     max_col=start_col + n_series,
                     min_row=1, max_row=n_rows + 1)
    cats = Reference(ws_data, min_col=start_col, min_row=2,
                     max_row=n_rows + 1)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    if ch.title is not None:
        ch.title.overlay = False
    ch.legend.position = "r"
    ch.legend.overlay = False
    ch.y_axis.title = y_title
    try:
        ch.y_axis.title.tx.rich.bodyPr = RichTextProperties(
            rot=-5400000, vert="horz")
    except AttributeError:
        pass
    xn = ws_data.cell(row=1, column=start_col).value
    if xn:
        ch.x_axis.title = str(xn)
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.x_axis.tickLblSkip = skip
    ch.x_axis.tickMarkSkip = skip
    ch.layout = Layout(manualLayout=ManualLayout(
        layoutTarget="inner", xMode="edge", yMode="edge",
        x=0.10, y=0.14, w=0.72, h=0.66))
    ch.height = 9.5
    ch.width = 30
    for i, s in enumerate(ch.series):
        s.smooth = False
        s.graphicalProperties.line.width = 20000
        s.graphicalProperties.line.solidFill = SERIES_HEX[i % len(SERIES_HEX)]
    return ch


# -------------------------------------------------- negative-match engine
def toks(s):
    return tuple(re.findall(r"[a-z0-9']+", str(s).lower()))


def blocks(neg_tokens, ctype, q_tokens):
    """Does a negative (tokens, criterion type) block this query?
    Approximates Google semantics; negative broad has NO close variants."""
    if not neg_tokens:
        return False
    if "exact" in ctype:
        return q_tokens == neg_tokens
    if "phrase" in ctype:
        n = len(neg_tokens)
        return any(q_tokens[i:i + n] == neg_tokens
                   for i in range(len(q_tokens) - n + 1))
    return set(neg_tokens) <= set(q_tokens)       # broad


def load_existing(path):
    df = pd.read_csv(path, sep="\t", encoding="utf-16", dtype=str).fillna("")
    df = df[df["Keyword"].str.strip() != ""].copy()
    df["level"] = ["account" if c == "<Account-level>"
                   else ("campaign" if a.strip() == "" else "adgroup")
                   for c, a in zip(df["Campaign"], df["Ad Group"])]
    df["cr"] = df["Campaign"].str.extract(r"(CR\d+)")
    df["ctype"] = df["Criterion Type"].str.lower()
    df["neg_tokens"] = df["Keyword"].map(toks)
    df["ag_short"] = df["Ad Group"].str.replace("Life-NonGeneric-", "",
                                                regex=False)
    return df


def ag_brand(name):
    """Classify an ad group's portfolio from its name."""
    n = str(name)
    low = n.lower()
    if "ethos" in low:
        return "ethos"
    if "generic" in low and "nongeneric" not in low.replace("-", ""):
        return "generic"
    parts = n.split("-")
    if len(parts) >= 3 and parts[1].lower() == "nongeneric":
        return parts[2].strip().lower()           # carrier brand segment
    return "other"


def main():
    con = sqlite3.connect(f"file:{DB_PATH}?mode=ro", uri=True)

    # ---------------------------------------------------------- inputs
    exp = glob.glob(os.path.join(ANALYSIS_DIR, "input",
                                 "*negative-keywords*.csv"))
    if not exp:
        raise SystemExit("Ads Editor negatives export not found in input\\ "
                         "(expected *negative-keywords*.csv)")
    neg = load_existing(sorted(exp)[-1])
    hr(f"EXISTING NEGATIVES loaded: {len(neg)} rows "
       f"({(neg['level'] == 'account').sum()} account / "
       f"{(neg['level'] == 'campaign').sum()} campaign / "
       f"{(neg['level'] == 'adgroup').sum()} ad-group) "
       f"from {os.path.basename(sorted(exp)[-1])}")

    camps = pd.read_sql_query(
        "SELECT campaign_id, campaign_name FROM campaigns "
        "WHERE account_id=?", con, params=(ACCOUNT_ID,))
    camps["cr"] = camps["campaign_name"].str.extract(r"(CR\d+)")
    cr_of = {str(r.campaign_id): r.cr for r in camps.itertuples()}

    kws = pd.read_sql_query(
        "SELECT id, keyword_text, match_type, campaign_id, ad_group_name, "
        "status FROM keyword_instances WHERE account_id=?",
        con, params=(ACCOUNT_ID,))
    kws["cr"] = kws["campaign_id"].astype(str).map(cr_of)
    kws["ag_short"] = kws["ad_group_name"].str.replace(
        "Life-NonGeneric-", "", regex=False)
    kws["portfolio"] = kws["ad_group_name"].map(ag_brand)
    kws["kw_tokens"] = kws["keyword_text"].map(toks)
    enabled_kws = kws[kws["status"].str.upper() == "ENABLED"]
    hr(f"KEYWORD ROSTER: {len(kws)} instances, "
       f"{kws['ad_group_name'].nunique()} ad groups, portfolios: "
       f"{kws.groupby('portfolio')['id'].count().to_dict()}")

    # ------------------------------------------------- sqp (full YTD)
    q = pd.read_sql_query(f"""
        SELECT COALESCE(sq.query_text, CAST(p.query_id AS TEXT)) AS query,
               p.report_date, SUBSTR(p.report_date,1,7) AS month,
               p.instance_id, p.match_type_served,
               ki.ad_group_name, ki.keyword_text, ki.match_type,
               CAST(p.campaign_id AS TEXT) AS campaign_id,
               p.impressions, p.cost_micros/1e6 AS spend,
               p.primary_conv_count AS leads
        FROM search_query_performance p
        JOIN keyword_instances ki ON ki.id=p.instance_id
                                  AND ki.account_id={ACCOUNT_ID}
        LEFT JOIN search_queries sq ON sq.id=p.query_id
        WHERE p.platform='google' AND p.device='ALL'
          AND p.report_date>='{Y0}' AND p.report_date<'{O1}'""", con)
    q["cr"] = q["campaign_id"].map(cr_of)
    q["ag_short"] = q["ad_group_name"].str.replace("Life-NonGeneric-", "",
                                                   regex=False)
    q["portfolio"] = q["ad_group_name"].map(ag_brand)
    q["query"] = q["query"].str.lower().str.strip()
    q["in_window"] = (q["report_date"] >= W0) & (q["report_date"] < W1)
    q["in_oct"] = q["report_date"] >= O0
    hr(f"SQP loaded: {len(q):,} rows, {q['query'].nunique():,} distinct "
       f"visible queries, {Y0} -> {q['report_date'].max()}")

    # per-query classification (unique queries only)
    uq = pd.DataFrame({"query": q["query"].unique()})
    uq["q_tokens"] = uq["query"].map(toks)
    uq["has_ethos"] = uq["query"].str.contains(ETHOS_RE)
    uq["has_luminary"] = uq["query"].str.contains(LUMINARY_RE)
    uq["carrier"] = [next((k for k, rx in CARRIER_RE.items()
                           if rx.search(s)), None) for s in uq["query"]]
    for name, rx in LEX_RE.items():
        uq[f"lex_{name}"] = uq["query"].str.contains(rx)
    uq["protected"] = [any(rx.search(s) for rx in PROT_RE.values())
                       for s in uq["query"]]
    q = q.merge(uq, on="query", how="left")

    # -------------------------------------- existing-negative simulation
    # which visible queries would TODAY's negatives block, and when did
    # each negative's matching traffic last serve ("effective-since" clue)
    scoped = neg[neg["level"].isin(["campaign", "adgroup"])].copy()
    uq_pairs = list(zip(uq["query"], uq["q_tokens"]))
    blocked_mask = np.zeros(len(q), dtype=bool)
    last_served = []
    for nrow in scoped.itertuples():
        if nrow.level == "campaign":
            scope = (q["cr"] == nrow.cr).values
        else:
            scope = ((q["cr"] == nrow.cr)
                     & (q["ag_short"] == nrow.ag_short)).values
        blocked_qs = {s for s, t in uq_pairs
                      if blocks(nrow.neg_tokens, nrow.ctype, t)}
        hits = pd.Series(scope & q["query"].isin(blocked_qs).values,
                         index=q.index)
        blocked_mask |= hits.values
        sub = q[hits]
        last_served.append({
            "level": nrow.level, "cr": nrow.cr, "ad_group": nrow.ag_short,
            "negative": nrow.Keyword, "type": nrow.ctype,
            "window_spend_matching": sub.loc[sub["in_window"],
                                             "spend"].sum(),
            "window_leads_matching": sub.loc[sub["in_window"],
                                             "leads"].sum(),
            "last_matching_serve": sub["report_date"].max()
            if len(sub) else None})
    q["blocked_now"] = blocked_mask
    eff = pd.DataFrame(last_served).sort_values(
        "window_spend_matching", ascending=False)
    hr("EXISTING NEGATIVES vs WINDOW TRAFFIC (top 25 by matching spend "
       "— spend here means the negative was NOT live during the window, "
       "or is not blocking as assumed)")
    print(eff.head(25).to_string(index=False))

    # account-level negatives: summarized only (legacy auto/Progressive)
    acct_n = neg[neg["level"] == "account"]

    # ----------------------------------------- shared-list candidates
    def token_evidence(lex, list_name):
        rows = []
        for name in lex:
            col = f"lex_{name}"
            hit = q[q[col].fillna(False)]
            w = hit[hit["in_window"]]
            o = hit[hit["in_oct"]]
            if w["spend"].sum() == 0 and o["spend"].sum() == 0:
                continue
            # collision vs the account's own enabled keywords
            nt = toks(name)
            coll = enabled_kws[enabled_kws["kw_tokens"].map(
                lambda k: set(nt) <= set(k))]
            ex = (w.groupby("query")["spend"].sum()
                  .sort_values(ascending=False).head(3).index.tolist())
            already = w.loc[w["blocked_now"], "spend"].sum()
            rows.append({
                "proposed_negative": name, "match_type": "phrase",
                "destination_list": list_name,
                "window_spend": w["spend"].sum(),
                "window_leads": w["leads"].sum(),
                "window_CPL": (w["spend"].sum() / w["leads"].sum())
                if w["leads"].sum() else None,
                "sep_spend": w.loc[w["month"] == "2026-09", "spend"].sum(),
                "oct_mtd_spend": o["spend"].sum(),
                "distinct_queries": w["query"].nunique(),
                "%_already_blocked_by_live_negs":
                    100 * already / w["spend"].sum()
                    if w["spend"].sum() else 0.0,
                "keyword_collision": "; ".join(
                    coll["keyword_text"].unique()[:3]) or "none",
                "action": ("DO NOT ADD - collides with own keywords; "
                           "handle at ad-group level")
                if len(coll) else
                ("RECOMMEND ADD" if w["spend"].sum() >= SPEND_FLOOR
                 else "WATCHLIST (below spend floor)"),
                "example_queries": " | ".join(ex),
            })
        return (pd.DataFrame(rows)
                .sort_values("window_spend", ascending=False)
                if rows else pd.DataFrame())

    cand1 = token_evidence(DISQUALIFY, LIST1)
    cand2 = token_evidence(SERVICE, LIST2)
    hr(f"LIST 1 CANDIDATES — {LIST1}")
    print(cand1.to_string(index=False) if len(cand1) else "  (none)")
    hr(f"LIST 2 CANDIDATES — {LIST2}")
    print(cand2.to_string(index=False) if len(cand2) else "  (none)")

    # ------------------------------------------------- remap plan
    camp_neg = neg[neg["level"] == "campaign"].copy()
    DEST = {"login": LIST2, "log in": LIST2, "policy lookup": LIST2,
            "complaint": LIST2, "complaints": LIST2, "legit": LIST2,
            "dental": LIST1, "geico": LIST1, "india": LIST1, "uk": LIST1,
            "stadium": LIST1, "mycolonialpenn": LIST1}
    remap = (camp_neg.groupby(["Keyword", "ctype"])["cr"]
             .apply(lambda s: ",".join(sorted(set(s)))).reset_index())
    remap.columns = ["existing_negative", "type", "live_in_campaigns"]
    remap["destination"] = remap["existing_negative"].map(DEST).fillna(
        "review")
    remap["action"] = ("MIGRATE: add to '" + remap["destination"]
                       + "', then remove the per-campaign copies")
    remap.loc[remap["existing_negative"] == "legit", "action"] = (
        "REVIEW BEFORE MIGRATING: 'legit' blocks the converting "
        "brand-curiosity layer ('is ethos legit') — confirm intent; "
        "consider narrowing to exact negatives on offending queries")
    hr("REMAP PLAN — existing campaign-level negatives -> shared lists")
    print(remap.to_string(index=False))

    # ----------------------------------- ad-group routing: direction 1
    # carrier/senior traffic on the ETHOS ad group(s)
    eth_q = q[(q["portfolio"] == "ethos") & q["in_window"]]
    d1 = []
    for ctok, rx in CARRIER_RE.items():
        hit = eth_q[eth_q["carrier"] == ctok]
        if hit["spend"].sum() <= 0:
            continue
        for cr in sorted(hit["cr"].dropna().unique()):
            sub = hit[hit["cr"] == cr]
            existing = neg[(neg["level"] == "adgroup") & (neg["cr"] == cr)
                           & (neg["ag_short"].str.contains("Ethos"))
                           & (neg["Keyword"].str.lower()
                              .str.replace(" ", "")
                              == ctok.replace(" ", ""))]
            d1.append({
                "campaign": cr, "ad_group": "Ethos-Life Insurance",
                "proposed_negative": ctok, "match_type": "phrase",
                "window_spend": sub["spend"].sum(),
                "window_leads": sub["leads"].sum(),
                "sep_spend": sub.loc[sub["month"] == "2026-09",
                                     "spend"].sum(),
                "existing_negative": "YES (broad)" if len(existing)
                else "none",
                "action": "already live - verify add date" if len(existing)
                else ("RECOMMEND ADD" if sub["spend"].sum() >= SPEND_FLOOR
                      else "WATCHLIST"),
                "example_queries": " | ".join(
                    sub.groupby("query")["spend"].sum()
                    .sort_values(ascending=False).head(2).index)})
    D1 = (pd.DataFrame(d1).sort_values("window_spend", ascending=False)
          if d1 else pd.DataFrame())
    hr("AD-GROUP ROUTING 1 — carrier brands serving on ETHOS ad groups")
    print(D1.to_string(index=False) if len(D1) else "  (none visible)")

    # ----------------------------------- ad-group routing: direction 2
    # 'ethos' queries on carrier ad groups: the gap matrix, in dollars
    car_q = q[(~q["portfolio"].isin(["ethos", "generic", "other"]))
              & q["has_ethos"].fillna(False) & q["in_window"]]
    d2 = (car_q.groupby(["cr", "ag_short"])
          .agg(window_spend=("spend", "sum"),
               window_leads=("leads", "sum"),
               sep_spend=("spend",
                          lambda s: s[car_q.loc[s.index, "month"]
                                      == "2026-09"].sum()),
               distinct_queries=("query", "nunique")).reset_index())
    have = set((r.cr, r.ag_short) for r in neg[
        (neg["level"] == "adgroup")
        & (neg["Keyword"].str.strip() == "ethos")].itertuples())
    d2["existing_ethos_negative"] = [
        "YES (broad)" if (r.cr, r.ag_short) in have else "GAP"
        for r in d2.itertuples()]
    d2["proposed_negative"] = "ethos"
    d2["match_type"] = "broad"
    d2["action"] = np.where(
        d2["existing_ethos_negative"] == "GAP",
        np.where(d2["window_spend"] >= SPEND_FLOOR,
                 "RECOMMEND ADD", "ADD (consistency - low spend)"),
        "already live - verify add date")
    d2 = d2.sort_values("window_spend", ascending=False)
    hr("AD-GROUP ROUTING 2 — 'ethos' queries on CARRIER ad groups "
       "(gap-fill; shadow-demand repatriation)")
    print(d2.to_string(index=False) if len(d2) else "  (none visible)")

    # --------------------------- the gap MATRIX (coverage x leaked $)
    # rows = every carrier ad group in the live keyword roster (incl.
    # groups with zero negatives, e.g. HighCPC); cells = live / GAP /
    # n/a (group absent from that campaign); + leaked ethos-query $.
    carrier_ags = (kws[~kws["portfolio"].isin(["ethos", "generic",
                                               "other"])]
                   [["ag_short", "cr"]].drop_duplicates())
    crs = ["CR12", "CR3", "CR45"]
    mat_rows = []
    leak_w = car_q.groupby(["ag_short", "cr"])["spend"].sum()
    leak_s = (car_q[car_q["month"] == "2026-09"]
              .groupby(["ag_short", "cr"])["spend"].sum())
    for ag in sorted(carrier_ags["ag_short"].unique()):
        in_cr = set(carrier_ags.loc[carrier_ags["ag_short"] == ag, "cr"])
        row = {"carrier_ad_group": ag}
        for cr in crs:
            if cr not in in_cr:
                row[cr] = "n/a"
            elif (cr, ag) in have:
                row[cr] = "live"
            else:
                leak = leak_w.get((ag, cr), 0.0)
                row[cr] = f"GAP (${leak:,.0f})" if leak else "GAP"
        row["ethos_query_window_$"] = float(
            leak_w.loc[ag].sum()) if ag in leak_w.index.get_level_values(
            0) else 0.0
        row["ethos_query_sep_$"] = float(
            leak_s.loc[ag].sum()) if ag in leak_s.index.get_level_values(
            0) else 0.0
        mat_rows.append(row)
    gap_matrix = pd.DataFrame(mat_rows).sort_values(
        "ethos_query_sep_$", ascending=False)
    total = {"carrier_ad_group": "TOTAL (the carrier->ethos leak)",
             "CR12": "", "CR3": "", "CR45": "",
             "ethos_query_window_$":
             gap_matrix["ethos_query_window_$"].sum(),
             "ethos_query_sep_$": gap_matrix["ethos_query_sep_$"].sum()}
    gap_matrix = pd.concat([gap_matrix, pd.DataFrame([total])],
                           ignore_index=True)
    hr("GAP MATRIX — 'ethos' negative coverage by carrier ad group x "
       "campaign, with leaked ethos-query spend")
    print(gap_matrix.to_string(index=False))

    # luminary isolation note
    lum = q[q["has_luminary"].fillna(False) & q["in_window"]]
    lum_tab = (lum.groupby(["cr", "ag_short"])
               .agg(spend=("spend", "sum"), leads=("leads", "sum"))
               .reset_index().sort_values("spend", ascending=False))

    # ------------------------------------------------- whole-life block
    wl = q[q[[f"lex_{k}" for k in SURFACE_ONLY]].fillna(False).any(axis=1)]
    wl_m = (wl[wl["report_date"] >= W0].groupby(["month"])
            .agg(spend=("spend", "sum"), leads=("leads", "sum"),
                 impressions=("impressions", "sum")).reset_index())
    wl_m["CPL"] = wl_m["spend"] / wl_m["leads"].replace(0, np.nan)
    wl_top = (wl[wl["in_window"]].groupby("query")
              .agg(spend=("spend", "sum"), leads=("leads", "sum"))
              .sort_values("spend", ascending=False).head(25)
              .reset_index())
    hr("WHOLE/UNIVERSAL LIFE — surfaced, NOT negated (possible dedicated "
       "campaign)")
    print(wl_m.to_string(index=False))

    # ------------------------------------------------- impact summary
    def rec_spend(df_):
        return df_.loc[df_.get("action", pd.Series(dtype=str))
                       .astype(str).str.startswith("RECOMMEND"),
                       "sep_spend"].sum() if len(df_) else 0.0
    impact = pd.DataFrame([
        {"component": f"Shared list 1 ({LIST1})",
         "sep_spend_affected": rec_spend(cand1)},
        {"component": f"Shared list 2 ({LIST2})",
         "sep_spend_affected": rec_spend(cand2)},
        {"component": "Ad-group: carriers on Ethos groups",
         "sep_spend_affected": rec_spend(D1)},
        {"component": "Ad-group: 'ethos' on carrier groups (gaps)",
         "sep_spend_affected":
         d2.loc[d2["existing_ethos_negative"] == "GAP",
                "sep_spend"].sum() if len(d2) else 0.0}])
    impact["note"] = ("VISIBLE spend only - lower bound; token negatives "
                      "also block the hidden ~15%")
    hr("IMPACT SUMMARY (September, visible $)")
    print(impact.to_string(index=False))

    # monthly class spend (chart data)
    def class_month(mask, label):
        s = (q[mask & (q["report_date"] >= W0)]
             .groupby("month")["spend"].sum())
        s.name = label
        return s
    dis_mask = q[[f"lex_{k}" for k in DISQUALIFY]].fillna(False).any(axis=1)
    ser_mask = q[[f"lex_{k}" for k in SERVICE]].fillna(False).any(axis=1)
    rt1 = (q["portfolio"] == "ethos") & q["carrier"].notna()
    rt2 = (~q["portfolio"].isin(["ethos", "generic", "other"])
           ) & q["has_ethos"].fillna(False)
    cm = pd.concat([class_month(dis_mask, "Disqualifying intent"),
                    class_month(ser_mask, "Service/navigational"),
                    class_month(rt1, "Carrier-on-Ethos (routing)"),
                    class_month(rt2, "Ethos-on-carrier (routing)")],
                   axis=1).fillna(0).reset_index()

    # ======================================================== workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Read_Me"
    meta(ws, "NEGATIVE KEYWORD PROPOSAL — Term Life (acct 29) — "
             "RATIFICATION GATED",
         "paid_search.db (read-only) + Ads Editor negatives export",
         f"evidence {W0}..{W1} (Sep-29+ excluded); Oct MTD confirmation",
         "sqp hides low-volume terms: all sizing is a LOWER BOUND; "
         "matching simulation approximates Google semantics (negative "
         "broad has NO close variants - misspellings are not blocked); "
         "export carries no add dates - 'last_matching_serve' infers "
         "effectiveness",
         "Nothing ships from this workbook. Review -> ratify -> build in "
         "Ads Editor.")
    glossary = pd.DataFrame([
        ["Proposal structure", "Campaigns are zip-income tiers of the "
         "same portfolios, so intent-level negatives go into SHARED "
         "LISTS applied to all 3; only routing negatives stay at ad-group "
         "level."],
        [LIST1, "NEW shared list: senior/final-expense/burial/age-band "
         "etc. - term-life product mismatch. Blocks nothing we want."],
        [LIST2, "NEW shared list: login/cancel/claims/phone etc. - "
         "existing-customer traffic."],
        ["Remap_Plan", "Existing campaign-level negatives migrate into "
         "the two shared lists (one place to manage); ad-group routing "
         "negatives stay put; 560 account-level legacy auto/Progressive "
         "entries left as-is."],
        ["AdGroup_Routing", "Direction 1: carrier brands on Ethos ad "
         "groups. Direction 2: 'ethos' broad negative gap-fill on "
         "carrier ad groups - repatriates shadow demand, gives up "
         "nothing."],
        ["Whole_Life", "Sized but NOT negated - candidate for a "
         "dedicated campaign."],
        ["PROTECTED (never negated)", "no exam / no medical exam (Ethos "
         "proposition); reviews/legit/cost/calculator curiosity layer; "
         "close variants of own exact terms ($22 CPL, 20% of leads)."],
        ["Success metric", "mis-routed % of visible broad spend "
         "(Broad_Match_Deep_Dive E2) falling from ~30% toward 0."],
        ["already_blocked / last_matching_serve", "If an existing "
         "negative shows window spend, it was added after that spend "
         "ran (export has no dates) - or is not matching as assumed."]],
        columns=["item", "meaning"])
    write_block(ws, 6, "How to read this workbook", glossary)

    ws1 = wb.create_sheet("List1_Disqualifying")
    meta(ws1, f"SHARED LIST 1 — {LIST1}",
         "sqp x lexicon", f"{W0}..{W1} + Oct MTD",
         "phrase negatives; collision-checked against own keywords; "
         "spend floor ${:,.0f}".format(SPEND_FLOOR),
         "Senior/final-expense demand is the concrete prem/op mechanism "
         "- term-life searchers these are not. Apply to all 3 campaigns "
         "via the shared list.")
    r = write_block(ws1, 6, "Candidates (evidence-ranked)",
                    cand1 if len(cand1) else
                    pd.DataFrame({"note": ["no visible spend matched"]}))

    ws2 = wb.create_sheet("List2_Service")
    meta(ws2, f"SHARED LIST 2 — {LIST2}",
         "sqp x lexicon", f"{W0}..{W1} + Oct MTD",
         "phrase negatives; collision-checked; existing campaign-level "
         "service negatives migrate here (see Remap_Plan)",
         "Existing-customer and navigational traffic - no new-policy "
         "value. Centralize in one shared list.")
    write_block(ws2, 6, "Candidates (evidence-ranked)",
                cand2 if len(cand2) else
                pd.DataFrame({"note": ["no visible spend matched"]}))

    ws3 = wb.create_sheet("Remap_Plan")
    meta(ws3, "REMAP PLAN — existing negatives re-mapped to the right "
              "level",
         "Ads Editor export", "current live negatives",
         "export has no add dates; 'legit' flagged for review before "
         "migration",
         "12 campaign-level tokens x 3 campaigns collapse into the two "
         "shared lists - one place to manage; ad-group routing "
         "negatives stay at ad-group level.")
    r = write_block(ws3, 6, "Campaign-level -> shared lists", remap)
    r = write_block(ws3, r, "Existing negatives vs window traffic "
                    "(matching spend means added late or not matching)",
                    eff.head(40))
    write_block(ws3, r, "Account-level legacy list (left as-is)",
                pd.DataFrame({"summary": [
                    f"{len(acct_n)} account-level excluded negatives - "
                    "inherited auto/Progressive typo list; no overlap "
                    "with life intent; recommend leaving untouched"]}))

    ws4 = wb.create_sheet("AdGroup_Routing")
    meta(ws4, "AD-GROUP ROUTING NEGATIVES — both leak directions",
         "sqp x ad-group portfolio x existing negatives",
         f"{W0}..{W1}",
         "ad-group level because the negative depends on the group's own "
         "portfolio; shared lists would collide with carrier keywords "
         "(e.g. Mutual of Omaha ad group bids 'mutual of omaha')",
         "Direction 2 gives up nothing: ethos demand on carrier groups "
         "is re-routed to the right keywords/ads/tracking, not lost.")
    r = write_block(ws4, 6, "GAP MATRIX - 'ethos' negative coverage by "
                    "carrier ad group x campaign (live / GAP with leaked "
                    "window $ / n/a = group absent); last two columns = "
                    "ethos-query spend the group carried in the window "
                    "and in September - the carrier->ethos leak",
                    gap_matrix)
    r = write_block(ws4, r, "Direction 1 - carrier/senior brands on "
                    "ETHOS ad groups",
                    D1 if len(D1) else
                    pd.DataFrame({"note": ["no visible spend"]}))
    r = write_block(ws4, r, "Direction 2 - 'ethos' negative gap-fill on "
                    "CARRIER ad groups (detail behind the matrix)",
                    d2 if len(d2) else
                    pd.DataFrame({"note": ["no visible spend"]}))
    write_block(ws4, r, "Luminary exposure (isolate - do not blanket-"
                "negate without strategy call)",
                lum_tab if len(lum_tab) else
                pd.DataFrame({"note": ["no luminary queries visible"]}))

    ws5 = wb.create_sheet("Whole_Life")
    meta(ws5, "WHOLE / UNIVERSAL LIFE — SURFACED, NOT NEGATED",
         "sqp x lexicon", f"{W0} onward",
         "visible lower bound",
         "Sized as potential demand for a dedicated campaign - excluded "
         "from every proposed list above.")
    r = write_block(ws5, 6, "Monthly", wl_m)
    write_block(ws5, r, "Top queries (window)", wl_top)

    ws6 = wb.create_sheet("Impact_Summary")
    meta(ws6, "IMPACT SUMMARY — September visible spend affected",
         "all blocks above", "Sep 1-28",
         "visible-only lower bound; token negatives also block the "
         "hidden tail",
         "Success metric: mis-routed % of broad spend (~30% in Sep) "
         "toward 0 after rollout.")
    write_block(ws6, 6, "By component", impact)

    wsc = wb.create_sheet("Charts")
    wsd = wb.create_sheet("Chart_Data")
    n = write_series_block(wsd, 1, cm)
    wsc.add_chart(make_line_chart(wsd, 1, len(cm.columns) - 1, n,
                                  "Candidate-class visible spend by month",
                                  "spend ($)"), "A2")
    meta(wsc, "CANDIDATE-CLASS SPEND — monthly",
         "sqp x lexicon/portfolio", f"{W0} onward",
         "visible lower bound; Oct MTD partial",
         "All four classes step up in September with the budget raises - "
         "the negatives close the doors the raise opened.")
    wsd.sheet_state = "hidden"

    os.makedirs(os.path.dirname(OUT_XLSX), exist_ok=True)
    wb.save(OUT_XLSX)
    hr(f"WORKBOOK SAVED -> {OUT_XLSX}")
    con.close()


if __name__ == "__main__":
    main()
