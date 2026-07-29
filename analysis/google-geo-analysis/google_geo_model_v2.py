"""
google_geo_model.py — Google Home geo-tiering. Ported from microsoft_geo_model_comparison.py;
the clustering method (A2 quantile, B1 2-D k-means, B4 Ward hierarchical, C2 weighted k-means)
is unchanged. What's configurable is the FEATURE the tiers are built on.

CLUSTERING FEATURE (--cluster-feature):
  gp_per_lead  (DEFAULT) = RPL - CPL  (margin per lead). Matches the current Google tier
                            philosophy, which is a CPC x lead-value PROFILE with margin as the
                            spine (Tier 1 cheap/efficient/positive-margin; Tier 2 premium/high-value/
                            thin-margin; Tier 3 volatile mid-pack; Tier 4 break-even/volume). GP/lead
                            folds cost and value into the margin you bid for. TIER ONE = highest GP/lead.
  roas         = revenue / spend (margin as a ratio).
  rpl          = revenue / conversion (pure value; the original Microsoft basis — ignores CPC).

GOOGLE DATA LAYER (vs the Microsoft core):
  - account_id = 1; source campaign_performance_geo_daily_search (sums device rows; no 'ALL' row;
    no network/location_type). Native conversion basis (same as the workbook RPL_Redistribution):
    RPL/GP ordering sound; absolute $ run ~10-13% conservative.
  - search-only (no audience network); single-state campaigns (--exclude-states) are carved out
    of the clustering pool since they stay standalone.
  - tROAS account: tier_target_tcpa_analytical (= tier RPL / 1.5) is reported for reference only,
    NOT an operational bid. The operational output is tier MEMBERSHIP.

Usage:
    python google_geo_model.py --rpl-blend 0.6 --k 7
    python google_geo_model.py --rpl-blend 0.6 --elbow            # sweep k, write elbow chart
    python google_geo_model.py --rpl-blend 0.6 --k 7 --cluster-feature rpl

SHORT-WINDOW VALIDATION MODE (--start-date):
    Explicit window (e.g. post partner bid-down). Stability check becomes a half-split
    fragility flag; eligibility should switch to conversions (--conv-threshold, RPL is
    conversion-count-limited on short windows); --seed-sweep adds per-state seed stability
    for the stochastic models; the committed --ledger (not name-based campaign inference)
    supplies the reference column; output gets a 'validation' filename prefix so gp_bridge
    auto-discovery never picks it up as the authoritative zone source.

    # Director's B1 (RPL, CPC) read, post-bid-down:
    python google_geo_model.py --cluster-feature rpl --headline-model B1_kmeans_2d --k 4 \\
        --start-date 2026-06-19 --end-date 2026-07-07 --conv-threshold 9 \\
        --seed-sweep 42,7,13,99,2026 --ledger commitments\\google_home_committed.yaml
    # C2 gp_per_lead companion, same window, separate workbook:
    python google_geo_model.py --cluster-feature gp_per_lead --k 4 \\
        --start-date 2026-06-19 --end-date 2026-07-07 --conv-threshold 9 \\
        --seed-sweep 42,7,13,99,2026 --ledger commitments\\google_home_committed.yaml
"""
from __future__ import annotations

import argparse
import sqlite3
import sys
from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path

import numpy as np
import pandas as pd

DEFAULT_DB_PATH = r"C:\Projects\google-ads-ops\db\paid_search.db"
DEFAULT_OUTPUT_DIR = r"C:\Projects\google-ads-ops\analysis\google-geo-analysis\output"
ACCOUNT_ID = 1
GEO_TABLE = "campaign_performance_geo_daily_search"
CLICK_THRESHOLD = 200
TARGET_ROAS = 1.50
DEFAULT_K = 7
DEFAULT_EXCLUDE = ["California", "Texas", "Florida"]   # standalone single-state campaigns

# Set in main() from --cluster-feature
FEATURE = "gp_per_lead"
FEATURE_LABEL = "GP/lead"

TIER_LABEL_BASE = [
    "TIER ONE", "TIER TWO", "TIER THREE", "TIER FOUR", "TIER FIVE",
    "TIER SIX", "TIER SEVEN", "TIER EIGHT", "TIER NINE", "TIER TEN",
    "TIER ELEVEN", "TIER TWELVE",
]
MODEL_NAMES = ["A2_quantile", "B1_kmeans_2d", "B4_hierarchical", "C2_weighted_kmeans"]
CURRENT_TIER_NAMES = ["Tier One", "Tier Two", "Tier Three", "Tier Four",
                      "California", "Texas", "Florida"]
# Re-tier completion dates (end-of-day), keyed by COMMITTED tier. Used to annotate
# restructure exposure on short-window runs. Update as the rollout progresses.
MOVE_DATES = {"TIER ONE": "2026-06-25", "TIER TWO": "2026-06-30", "TIER THREE": "2026-07-08"}


@dataclass
class Windows:
    end_date: date
    w12_start: date
    w8_start: date
    w4_start: date


def build_windows(end_date: date) -> Windows:
    return Windows(end_date=end_date,
                   w12_start=end_date - timedelta(days=83),
                   w8_start=end_date - timedelta(days=55),
                   w4_start=end_date - timedelta(days=27))


def current_label(name: str) -> str:
    u = (name or "").upper()
    for nm in CURRENT_TIER_NAMES:
        if nm.upper() in u:
            return nm
    return name or ""


# ----------------------------------------------------------------------------
# Data pull (Google)
# ----------------------------------------------------------------------------

def resolve_end_date(conn, override):
    if override:
        return date.fromisoformat(override)
    row = conn.execute(f"SELECT MAX(date) FROM {GEO_TABLE} WHERE account_id = ?", (ACCOUNT_ID,)).fetchone()
    if not row or not row[0]:
        raise RuntimeError(f"No rows in {GEO_TABLE} for account_id={ACCOUNT_ID}")
    max_date = date.fromisoformat(row[0])
    yesterday = date.today() - timedelta(days=1)
    return yesterday if max_date == yesterday else max_date


def pull_state_aggregates(conn, start, end):
    sql = f"""
        SELECT geo_state AS state, SUM(clicks) AS clicks,
               SUM(cost_micros) / 1e6 AS spend,
               SUM(conversions) AS conversions, SUM(conversion_value) AS revenue
        FROM {GEO_TABLE}
        WHERE account_id = ? AND device <> 'ALL' AND geo_state IS NOT NULL
          AND date BETWEEN ? AND ?
        GROUP BY geo_state
    """
    return pd.read_sql_query(sql, conn, params=(ACCOUNT_ID, start.isoformat(), end.isoformat()))


def pull_state_daily(conn, start, end):
    """Per (state, date) aggregates — the resampling grain for the day-level bootstrap."""
    sql = f"""
        SELECT geo_state AS state, date, SUM(clicks) AS clicks,
               SUM(cost_micros) / 1e6 AS spend,
               SUM(conversions) AS conversions, SUM(conversion_value) AS revenue
        FROM {GEO_TABLE}
        WHERE account_id = ? AND device <> 'ALL' AND geo_state IS NOT NULL
          AND date BETWEEN ? AND ?
        GROUP BY geo_state, date
    """
    return pd.read_sql_query(sql, conn, params=(ACCOUNT_ID, start.isoformat(), end.isoformat()))


def pull_state_to_current_tier(conn, start, end):
    try:
        df = pd.read_sql_query(
            f"""SELECT g.geo_state AS state, c.campaign_name AS campaign,
                       SUM(g.cost_micros)/1e6 AS spend
                FROM {GEO_TABLE} g
                JOIN campaigns c ON c.campaign_id = g.campaign_id AND c.account_id = g.account_id
                WHERE g.account_id = ? AND g.device <> 'ALL' AND g.geo_state IS NOT NULL
                  AND g.date BETWEEN ? AND ?
                GROUP BY g.geo_state, c.campaign_name""",
            conn, params=(ACCOUNT_ID, start.isoformat(), end.isoformat()))
        if df.empty:
            return {}
        top = df.sort_values("spend", ascending=False).groupby("state").head(1)
        return {s: current_label(c) for s, c in zip(top["state"], top["campaign"])}
    except Exception as e:
        print(f"  (current-tier map skipped: {e})", file=sys.stderr)
        return {}


# ============================================================================
# Method (clustering identical to the Microsoft core; value axis = FEATURE)
# ============================================================================

def compute_derived(df):
    df = df.copy()
    df["cpc"] = df["spend"] / df["clicks"].where(df["clicks"] > 0)
    df["cpa"] = df["spend"] / df["conversions"].where(df["conversions"] > 0)
    df["roas"] = df["revenue"] / df["spend"].where(df["spend"] > 0)
    df["rpl"] = df["revenue"] / df["conversions"].where(df["conversions"] > 0)
    df["cpl"] = df["spend"] / df["conversions"].where(df["conversions"] > 0)
    df["gp_per_lead"] = df["rpl"] - df["cpl"]
    df["ideal_tcpa"] = df["rpl"] / TARGET_ROAS
    return df


def apply_rpl_blend(df_12w, df_4w, blend_weight, min_blend_clicks=75):
    """Blend 4w into 12w RPL, then re-derive cpl/gp_per_lead/roas-dependent fields are
    based on the (blended) RPL and the 12w cost structure."""
    out = df_12w.copy()
    out["rpl_12w_original"] = out["rpl"]
    out["blend_applied"] = False
    if blend_weight <= 0.0 or df_4w is None or df_4w.empty:
        return out
    blend_weight = min(max(blend_weight, 0.0), 1.0)
    idx4 = df_4w.set_index("state")

    def _blend(row):
        st = row["state"]; r12 = row["rpl"]
        if st not in idx4.index:
            return r12, False
        r4 = idx4.loc[st]
        if pd.isna(r4["rpl"]) or r4["rpl"] <= 0 or r4["clicks"] < min_blend_clicks:
            return r12, False
        return (1.0 - blend_weight) * r12 + blend_weight * r4["rpl"], True

    res = out.apply(_blend, axis=1)
    out["rpl"] = [b[0] for b in res]
    out["blend_applied"] = [b[1] for b in res]
    # re-derive value-dependent fields from blended RPL (cpl/cpc unchanged — cost-side)
    out["gp_per_lead"] = out["rpl"] - out["cpl"]
    out["ideal_tcpa"] = out["rpl"] / TARGET_ROAS
    return out


def eligible_pool(df, click_threshold=CLICK_THRESHOLD, conv_threshold=None):
    """Eligible = enough volume, finite RPL/CPC, not UNKNOWN. (GP/lead may be negative —
    a real state we bid on; eligibility is about volume/validity, not profitability.)
    Volume gate: conversions when conv_threshold is set (short windows — RPL is
    conversion-count-limited, not click-limited), else clicks (the 12w default)."""
    vol = (df["conversions"] >= conv_threshold) if conv_threshold is not None \
        else (df["clicks"] >= click_threshold)
    return df[
        vol
        & df["rpl"].notna() & df["cpc"].notna() & df[FEATURE].notna()
        & (df["rpl"] > 0) & (df["cpc"] > 0)
        & (df["state"].str.upper() != "UNKNOWN")
    ].copy().reset_index(drop=True)


def relabel_by_feature(states, cluster_col, k):
    """Map cluster IDs to TIER labels ordered by mean FEATURE descending.
    TIER ONE = best (highest GP/lead / ROAS / RPL)."""
    means = states.groupby(cluster_col)[FEATURE].mean().sort_values(ascending=False)
    rank_map = {old: TIER_LABEL_BASE[i] for i, old in enumerate(means.index[:k])}
    return states[cluster_col].map(rank_map).astype(object)


def model_a2_quantile(states, k):
    ranks = states[FEATURE].rank(method="first", ascending=False)
    out = states.copy(); out["_c"] = pd.cut(ranks, bins=k, labels=False, include_lowest=True)
    return relabel_by_feature(out, "_c", k)


def model_b1_kmeans_2d(states, k, seed=42):
    from sklearn.cluster import KMeans
    from sklearn.preprocessing import StandardScaler
    X = StandardScaler().fit_transform(states[[FEATURE, "cpc"]].values)
    km = KMeans(n_clusters=k, random_state=seed, n_init=10).fit(X)
    out = states.copy(); out["_c"] = km.labels_
    return relabel_by_feature(out, "_c", k)


def model_b4_hierarchical(states, k):
    from scipy.cluster.hierarchy import linkage, fcluster
    from sklearn.preprocessing import StandardScaler
    feat = pd.DataFrame({FEATURE: states[FEATURE].values, "cpc": states["cpc"].values,
                         "log_clicks": np.log(states["clicks"].values)})
    X = StandardScaler().fit_transform(feat.values)
    Z = linkage(X, method="ward")
    out = states.copy(); out["_c"] = fcluster(Z, t=k, criterion="maxclust")
    return relabel_by_feature(out, "_c", k), Z


def model_c2_weighted_kmeans(states, k, seed=42):
    """1D weighted k-means on FEATURE, sample_weight=spend (headline). Minimizes
    Σ((feature_i - tier_feature)^2 * spend_i)."""
    from sklearn.cluster import KMeans
    X = states[[FEATURE]].values; w = states["spend"].values
    km = KMeans(n_clusters=k, random_state=seed, n_init=10).fit(X, sample_weight=w)
    out = states.copy(); out["_c"] = km.labels_
    return relabel_by_feature(out, "_c", k)


def run_single_model(elig, k, model_name, seed=42):
    """Dispatch one model by name (used by the bootstrap so it re-runs only the headline)."""
    if model_name == "A2_quantile":
        return model_a2_quantile(elig, k)
    if model_name == "B1_kmeans_2d":
        return model_b1_kmeans_2d(elig, k, seed=seed)
    if model_name == "B4_hierarchical":
        return model_b4_hierarchical(elig, k)[0]
    return model_c2_weighted_kmeans(elig, k, seed=seed)


def bootstrap_assignments(daily, k, headline, n_boot, click_threshold, conv_threshold,
                          excluded, seed=2026):
    """Day-level bootstrap: resample the window's DATES with replacement n_boot times,
    re-aggregate states, re-run the headline model, and tally each state's tier.
    P(modal tier) is the decision-legible upgrade over the 2-fold half-split: it answers
    'if this window's days had fallen slightly differently, how often does this state
    keep its tier?' Eligibility is re-applied per resample; thin states drop in and out,
    which boot_eligible_pct records."""
    from collections import Counter
    if excluded:
        daily = daily[~daily["state"].isin(excluded)]
    daily = daily[daily["state"].str.upper() != "UNKNOWN"]
    dates = sorted(daily["date"].unique())
    states = sorted(daily["state"].unique())
    mats = {}
    for c in ("clicks", "spend", "conversions", "revenue"):
        mats[c] = (daily.pivot_table(index="state", columns="date", values=c,
                                     aggfunc="sum", fill_value=0.0)
                   .reindex(index=states, columns=dates, fill_value=0.0).values)
    rng = np.random.default_rng(seed)
    nd = len(dates)
    tally = {s: Counter() for s in states}
    used = 0
    for _ in range(n_boot):
        w = np.bincount(rng.integers(0, nd, nd), minlength=nd).astype(float)
        agg = pd.DataFrame({"state": states})
        for c in ("clicks", "spend", "conversions", "revenue"):
            agg[c] = mats[c] @ w
        agg = compute_derived(agg)
        elig = eligible_pool(agg, click_threshold=click_threshold, conv_threshold=conv_threshold)
        if len(elig) < k:
            continue
        used += 1
        lab = run_single_model(elig.reset_index(drop=True), k, headline)
        for s, t in zip(elig["state"], lab):
            tally[s][t] += 1
    rows = []
    for s in states:
        c = tally[s]; n = sum(c.values())
        if n == 0:
            rows.append({"state": s, "boot_rounds": 0, "boot_eligible_pct": 0.0,
                         "boot_modal_tier": None, "boot_p_modal": np.nan, "boot_distribution": ""})
            continue
        modal, cnt = c.most_common(1)[0]
        dist = "  ".join(f"{t.replace('TIER ', 'T')}:{v / n:.0%}" for t, v in c.most_common(3))
        rows.append({"state": s, "boot_rounds": n,
                     "boot_eligible_pct": n / used if used else 0.0,
                     "boot_modal_tier": modal, "boot_p_modal": cnt / n,
                     "boot_distribution": dist})
    return pd.DataFrame(rows), used


def bid_down_incidence(conn, pre_range, post_start, post_end, excluded, min_conv=8):
    """Per-state RPL(post)/RPL(pre) around the partner bid-down. Tests whether the hit was
    UNIFORM across states: tight ratios -> pre-period data preserves state ordering (a
    recency blend is defensible); dispersed ratios -> the partner's demand mix hit states
    differentially, and blending pre-period data biases states relative to each other.
    Also identifies where that partner's demand concentrates (the most-hit states)."""
    a0, a1 = pre_range.split(":")
    d0, d1 = date.fromisoformat(a0), date.fromisoformat(a1)
    pre_days = (d1 - d0).days + 1
    post_days = (post_end - post_start).days + 1
    pre = compute_derived(pull_state_aggregates(conn, d0, d1))
    post = compute_derived(pull_state_aggregates(conn, post_start, post_end))
    m = pre.merge(post, on="state", suffixes=("_pre", "_post"))
    if excluded:
        m = m[~m["state"].isin(excluded)]
    m = m[m["state"].str.upper() != "UNKNOWN"].copy()
    m["rpl_ratio"] = m["rpl_post"] / m["rpl_pre"]
    for side, nd in (("pre", pre_days), ("post", post_days)):
        for metric in ("clicks", "spend", "revenue", "conversions"):
            m[f"{metric}_day_{side}"] = m[f"{metric}_{side}"] / nd
    ok = m[(m["conversions_pre"] >= min_conv) & (m["conversions_post"] >= min_conv)
           & m["rpl_ratio"].notna() & np.isfinite(m["rpl_ratio"])]
    if ok.empty:
        return pd.DataFrame(), {}
    med = float(ok["rpl_ratio"].median())
    q1, q3 = ok["rpl_ratio"].quantile([0.25, 0.75])
    m["readable"] = m["state"].isin(ok["state"])
    m["incidence_flag"] = np.where(
        m["readable"] & ((m["rpl_ratio"] / med - 1).abs() > 0.15),
        "DEVIANT (>15% vs median)",
        np.where(m["readable"], "", "thin (<{} conv one side)".format(min_conv)))
    cols = (["state", "conversions_pre", "rpl_pre", "conversions_post", "rpl_post",
             "rpl_ratio", "incidence_flag"]
            + [f"{met}_{side}" for side in ("pre", "post") for met in ("cpc", "clicks", "spend", "revenue")]
            + [f"{met}_day_{side}" for side in ("pre", "post")
               for met in ("clicks", "spend", "revenue", "conversions")])
    out = m[cols].sort_values("rpl_ratio").reset_index(drop=True)
    summary = {"median_ratio": med, "iqr": float(q3 - q1), "iqr_over_median": float((q3 - q1) / med),
               "n_readable": int(len(ok)),
               "n_deviant": int((m["incidence_flag"].str.startswith("DEVIANT")).sum()),
               "pre_window": f"{a0}..{a1}", "pre_days": pre_days, "post_days": post_days}
    return out, summary


def apply_shrinkage(alls, m_prior):
    """Empirical-Bayes shrinkage: each state's RPL/CPL pulled toward the pool mean with
    m_prior pseudo-conversions. Thin states become partial reads instead of raw noise;
    high-volume states barely move. Diagnostic columns only — the headline clustering
    stays on the raw feature."""
    pool_conv = alls["conversions"].sum()
    if pool_conv <= 0:
        return alls
    pool_rpl = alls["revenue"].sum() / pool_conv
    pool_cpl = alls["spend"].sum() / pool_conv
    alls = alls.copy()
    alls["shrunk_rpl"] = (alls["revenue"] + m_prior * pool_rpl) / (alls["conversions"] + m_prior)
    alls["shrunk_cpl"] = (alls["spend"] + m_prior * pool_cpl) / (alls["conversions"] + m_prior)
    alls["shrunk_gp_per_lead"] = alls["shrunk_rpl"] - alls["shrunk_cpl"]
    return alls


def apply_all_models(states, k):
    out = states.copy().reset_index(drop=True)
    out["A2_quantile"] = model_a2_quantile(out, k)
    out["B1_kmeans_2d"] = model_b1_kmeans_2d(out, k)
    b4, Z = model_b4_hierarchical(out, k)
    out["B4_hierarchical"] = b4
    out["C2_weighted_kmeans"] = model_c2_weighted_kmeans(out, k)
    return out, {"b4_linkage": Z, "b4_state_order": out["state"].tolist()}


def _wtd(grp, col):
    sp = grp["spend"].sum()
    return (grp[col] * grp["spend"]).sum() / sp if sp > 0 else np.nan


def tier_aggregate_metrics(states, tier_col):
    """Per-tier profile so each tier maps to a margin archetype."""
    rows = []
    for tier, grp in states.groupby(tier_col, observed=True):
        if pd.isna(tier) or tier == "":
            continue
        rows.append({
            "tier": tier, "n_states": len(grp),
            "states": ", ".join(sorted(grp["state"].tolist())),
            "tier_spend": grp["spend"].sum(),
            "wtd_gp_per_lead": _wtd(grp, "gp_per_lead"),
            "wtd_rpl": _wtd(grp, "rpl"),
            "wtd_cpc": _wtd(grp, "cpc"),
            "wtd_cpl": _wtd(grp, "cpl"),
            "wtd_roas": _wtd(grp, "roas"),
            "tier_target_tcpa_analytical": _wtd(grp, "rpl") / TARGET_ROAS,
            "feat_min": grp[FEATURE].min(), "feat_max": grp[FEATURE].max(),
            "feat_range": grp[FEATURE].max() - grp[FEATURE].min(),
            "feat_stdev": grp[FEATURE].std(ddof=0),
        })
    df = pd.DataFrame(rows)
    if df.empty:
        return df
    order = {t: i for i, t in enumerate(TIER_LABEL_BASE)}
    df["_o"] = df["tier"].map(order)
    return df.sort_values("_o").drop(columns=["_o"]).reset_index(drop=True)


def model_summary_metrics(states, tier_col):
    ta = tier_aggregate_metrics(states, tier_col)
    if ta.empty:
        return {}
    feat_map = dict(zip(ta["tier"], ta[f"wtd_{FEATURE}"] if f"wtd_{FEATURE}" in ta else ta["wtd_gp_per_lead"]))
    # tier-center of the FEATURE (spend-weighted mean within tier)
    center = states.groupby(tier_col)[FEATURE].apply(lambda g: np.average(g, weights=states.loc[g.index, "spend"]) if states.loc[g.index, "spend"].sum() > 0 else g.mean())
    s = states.copy()
    s["_tc"] = s[tier_col].map(center)
    bid_loss = ((s[FEATURE] - s["_tc"]) ** 2 * s["spend"]).sum()
    vstd = ta.dropna(subset=["feat_stdev"])
    avg_stdev = ((vstd["feat_stdev"] * vstd["n_states"]).sum() / vstd["n_states"].sum()) if not vstd.empty else np.nan
    return {"model": tier_col, "n_tiers_populated": len(ta),
            "avg_within_tier_feature_stdev": avg_stdev,
            "spend_weighted_bid_loss": bid_loss}


def run_models_for_window(conn, start, end, k, click_threshold=CLICK_THRESHOLD, exclude=None,
                          conv_threshold=None):
    agg = pull_state_aggregates(conn, start, end)
    if agg.empty:
        return None
    agg = compute_derived(agg)
    if exclude:
        agg = agg[~agg["state"].isin(exclude)]
    elig = eligible_pool(agg, click_threshold=click_threshold, conv_threshold=conv_threshold)
    if len(elig) < k:
        return None
    res, _ = apply_all_models(elig, k)
    return res


def build_stability_table(r12, r8, r4):
    rows = []
    states = sorted(set(r12["state"]))
    for r in (r8, r4):
        if r is not None:
            states = sorted(set(states) | set(r["state"]))
    i12 = r12.set_index("state") if r12 is not None else None
    i8 = r8.set_index("state") if r8 is not None else None
    i4 = r4.set_index("state") if r4 is not None else None

    def rank(t):
        if t is None or (isinstance(t, float) and np.isnan(t)):
            return None
        try:
            return TIER_LABEL_BASE.index(t)
        except ValueError:
            return None

    for st in states:
        for m in MODEL_NAMES:
            t12 = i12.loc[st, m] if (i12 is not None and st in i12.index) else None
            t8 = i8.loc[st, m] if (i8 is not None and st in i8.index) else None
            t4 = i4.loc[st, m] if (i4 is not None and st in i4.index) else None
            present = [t for t in (t12, t8, t4) if t is not None and not (isinstance(t, float) and pd.isna(t))]
            nu = len(set(present)) if present else 0
            rk = [rank(t) for t in (t12, t8, t4)]
            valid = [r for r in rk if r is not None]
            if not valid:
                d = "no_data"
            elif nu == 1:
                d = "stable"
            elif len(valid) < 3:
                d = "partial"
            elif rk[2] < rk[0]:
                d = "drifting_up"
            elif rk[2] > rk[0]:
                d = "drifting_down"
            else:
                d = "volatile"
            rows.append({"state": st, "model": m, "tier_12w": t12, "tier_8w": t8,
                         "tier_4w": t4, "stability_score": nu, "direction": d})
    return pd.DataFrame(rows)


def build_halfsplit_table(conn, start, end, k, headline, click_threshold, conv_threshold,
                          exclude, full_result):
    """Short-window fragility check (replaces the 12w/8w/4w stability check, which is not
    meaningful at ~20 days): re-cluster the first and second half of the window independently
    and compare the headline assignment. 'fragile' means the read is data-thin — a flag,
    not a stability claim. Half-window volume gates are halved (floor 3 conv / 25 clicks)."""
    days = (end - start).days
    mid = start + timedelta(days=days // 2)
    half_conv = max(3, int(np.ceil(conv_threshold / 2.0))) if conv_threshold else None
    half_clicks = max(25, click_threshold // 2)
    rH1 = run_models_for_window(conn, start, mid, k, half_clicks, exclude, half_conv)
    rH2 = run_models_for_window(conn, mid + timedelta(days=1), end, k, half_clicks, exclude, half_conv)
    full = dict(zip(full_result["state"], full_result[headline]))
    h1 = dict(zip(rH1["state"], rH1[headline])) if rH1 is not None else {}
    h2 = dict(zip(rH2["state"], rH2[headline])) if rH2 is not None else {}
    rows = []
    for st in sorted(full):
        t_full, t1, t2 = full[st], h1.get(st), h2.get(st)
        present = [t for t in (t1, t2) if t]
        if not present:
            flag = "thin (no half-window read)"
        elif len(present) < 2:
            flag = "partial (one half only)" if present[0] == t_full else "fragile (half disagrees)"
        elif t1 == t2 == t_full:
            flag = "consistent"
        elif t1 == t2:
            flag = "halves agree, differ from full"
        else:
            flag = "fragile (halves disagree)"
        rows.append({"state": st, "tier_full": t_full, "tier_H1": t1, "tier_H2": t2, "flag": flag})
    return pd.DataFrame(rows), mid


def seed_sweep_stability(elig, k, seeds):
    """Re-run the stochastic models (B1, C2) across seeds and count distinct assignments per
    state. 'stable' = same tier at every seed; anything else is seed-dependent — on a thin
    window that is itself the finding. A2/B4 are deterministic and excluded."""
    base = elig.reset_index(drop=True)
    out = pd.DataFrame({"state": base["state"].values})
    for name, fn in (("B1_kmeans_2d", model_b1_kmeans_2d),
                     ("C2_weighted_kmeans", model_c2_weighted_kmeans)):
        runs = np.vstack([fn(base, k, seed=s).values for s in seeds])
        n_distinct = [len(set(runs[:, i])) for i in range(runs.shape[1])]
        out[f"{name}_seed_stability"] = ["stable" if n == 1 else f"flips across {n} tiers"
                                         for n in n_distinct]
    return out


def assign_below_threshold(all_states, eligible_result, headline, click_threshold,
                           conv_threshold=None):
    if eligible_result is None or eligible_result.empty:
        return pd.DataFrame()
    under = (all_states["conversions"] < conv_threshold) if conv_threshold is not None \
        else (all_states["clicks"] < click_threshold)
    below = all_states[
        under
        & all_states["rpl"].notna() & (all_states["rpl"] > 0)
        & all_states[FEATURE].notna()
        & (all_states["state"].str.upper() != "UNKNOWN")
    ].copy()
    if below.empty:
        return pd.DataFrame()
    cent = eligible_result.groupby(headline)[FEATURE].mean()
    below["assigned_tier_nearest"] = below[FEATURE].apply(lambda t: (cent - t).abs().idxmin())
    sh = {"rpl": "shrunk_rpl", "gp_per_lead": "shrunk_gp_per_lead"}.get(FEATURE)
    if sh and sh in below.columns:
        below["assigned_tier_nearest_shrunk"] = below[sh].apply(lambda t: (cent - t).abs().idxmin())
    cols = ["state", "clicks", "conversions", "spend", "rpl", "cpc", "cpl", "gp_per_lead",
            "roas", "shrunk_rpl", "shrunk_gp_per_lead", "assigned_tier_nearest",
            "assigned_tier_nearest_shrunk"]
    cols = [c for c in cols if c in below.columns]
    return below[cols].sort_values(FEATURE, ascending=False).reset_index(drop=True)


def elbow_sweep(elig, k_values, out_png: Path | None, model="C2_weighted_kmeans"):
    """Sweep k on the headline model. Adds marginal_gain_pct = % reduction in
    spend-weighted bid-loss vs k-1 (the 'performance gain' of adding one more tier)."""
    rows = []
    for k in sorted(set(k_values)):
        if k < 2 or len(elig) < k:
            continue
        res, _ = apply_all_models(elig, k)
        m = model_summary_metrics(res, model)
        rows.append({"k": k, "spend_weighted_bid_loss": m["spend_weighted_bid_loss"],
                     "avg_within_tier_feature_stdev": m["avg_within_tier_feature_stdev"]})
    df = pd.DataFrame(rows)
    if df.empty:
        return df
    prev = df["spend_weighted_bid_loss"].shift(1)
    df["marginal_gain_pct"] = np.where(prev > 0, 1.0 - df["spend_weighted_bid_loss"] / prev, np.nan)
    if out_png is not None:
        try:
            import matplotlib
            matplotlib.use("Agg")
            import matplotlib.pyplot as plt
            fig, (ax1, ax2) = plt.subplots(2, 1, figsize=(7.5, 6.4), sharex=True,
                                           gridspec_kw={"height_ratios": [3, 2]})
            ax1.plot(df["k"], df["spend_weighted_bid_loss"], "o-", color="#305496", lw=2)
            ax1.set_ylabel(f"spend-wtd bid-loss on {FEATURE_LABEL}")
            ax1.set_title(f"Elbow — {model} on {FEATURE_LABEL}: dispersion vs k, and the gain per added tier")
            ax1.grid(alpha=0.3)
            for _, r in df.iterrows():
                ax1.annotate(f"{int(r['k'])}", (r["k"], r["spend_weighted_bid_loss"]),
                             textcoords="offset points", xytext=(0, 8), ha="center", fontsize=8)
            g = df.dropna(subset=["marginal_gain_pct"])
            ax2.bar(g["k"], g["marginal_gain_pct"] * 100.0, width=0.6, color="#9bbb59")
            for _, r in g.iterrows():
                ax2.annotate(f"{r['marginal_gain_pct']*100:.0f}%", (r["k"], r["marginal_gain_pct"] * 100.0),
                             textcoords="offset points", xytext=(0, 3), ha="center", fontsize=8)
            ax2.set_xlabel("k (number of tiers)")
            ax2.set_ylabel("% bid-loss reduction vs k\u22121")
            ax2.grid(alpha=0.3, axis="y")
            ax2.set_xticks(df["k"].astype(int).tolist())
            fig.tight_layout(); fig.savefig(out_png, dpi=150); plt.close(fig)
        except Exception as e:
            print(f"  (elbow chart skipped: {e})", file=sys.stderr)
    return df


# ----------------------------------------------------------------------------
# Workbook
# ----------------------------------------------------------------------------

def _write_df(ws, df, header_row=1):
    from openpyxl.styles import Font, PatternFill
    hf = Font(bold=True, color="FFFFFF"); hfill = PatternFill("solid", fgColor="305496")
    money = {"spend", "tier_spend", "revenue", "spend_pre", "spend_post",
             "revenue_pre", "revenue_post"}
    twod = {"rpl", "cpc", "cpl", "gp_per_lead", "roas", "ideal_tcpa", "wtd_gp_per_lead",
            "wtd_rpl", "wtd_cpc", "wtd_cpl", "wtd_roas", "tier_target_tcpa_analytical",
            "feat_min", "feat_max", "feat_range", "feat_stdev", "avg_within_tier_feature_stdev",
            "shrunk_rpl", "shrunk_cpl", "shrunk_gp_per_lead", "rpl_ratio",
            "rpl_pre", "rpl_post"}
    pct = {"boot_p_modal", "boot_eligible_pct"}
    cols = list(df.columns)
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=header_row, column=j, value=str(c)); cell.font = hf; cell.fill = hfill
    for i, (_, row) in enumerate(df.iterrows(), header_row + 1):
        for j, c in enumerate(cols, 1):
            v = row[c]
            if isinstance(v, np.floating):
                v = float(v)
            elif isinstance(v, np.integer):
                v = int(v)
            if isinstance(v, float) and pd.isna(v):
                v = None
            cell = ws.cell(row=i, column=j, value=v)
            if c in money:
                cell.number_format = '"$"#,##0'
            elif c == "spend_weighted_bid_loss":
                cell.number_format = "#,##0"
            elif c.endswith("_pct") or c in pct:
                cell.number_format = "0.0%"
            elif "_day_" in c:
                cell.number_format = "0.0"
            elif c in ("cpc_pre", "cpc_post"):
                cell.number_format = "0.00"
            elif c in twod:
                cell.number_format = "0.00"
    ws.freeze_panes = ws.cell(row=header_row + 1, column=2).coordinate
    return header_row + len(df)


def write_workbook(path, windows, k, rpl_blend, click_threshold, feature, result_12w, headline,
                   tier_metrics, model_compare, stability, below, current_map, ledger_cmp,
                   elbow_df, elbow_png, excluded, halfsplit=None, halfsplit_mid=None,
                   window_note="", conv_threshold=None, boot_df=None, boot_used=0,
                   incidence=None, incidence_summary=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    TITLE = Font(bold=True, size=14, color="305496"); NOTE = Font(italic=True, color="595959", size=9)
    wb = Workbook(); wb.remove(wb.active)

    gate_desc = (f">= {conv_threshold:g} conversions" if conv_threshold is not None
                 else f">= {click_threshold} clicks")
    ws = wb.create_sheet("Proposed Tiers")
    ws["A1"] = f"Google Home geo tiers — {headline} on {FEATURE_LABEL}, k={k}"; ws["A1"].font = TITLE
    ws["A2"] = (f"Window {windows.w12_start}..{windows.end_date}. RPL blend={rpl_blend:.2f}. "
                f"Eligibility {gate_desc}. TIER ONE = best {FEATURE_LABEL}. "
                f"Excluded (standalone campaigns): {', '.join(excluded) if excluded else 'none'}. "
                f"Per-tier profile (wtd_cpc / wtd_rpl / wtd_gp_per_lead / wtd_roas) maps each tier "
                f"to a margin archetype. tier_target_tcpa_analytical = wtd RPL / {TARGET_ROAS} (reference, not a bid).")
    ws["A2"].font = NOTE
    if window_note:
        ws["A3"] = window_note; ws["A3"].font = Font(bold=True, italic=True, color="9C0006", size=9)
        _write_df(ws, tier_metrics, header_row=5)
    else:
        _write_df(ws, tier_metrics, header_row=4)

    ws = wb.create_sheet("State Assignments")
    ws["A1"] = "Per-state assignments — all models + current campaign"; ws["A1"].font = TITLE
    ws["A2"] = f"Headline = {headline} on {FEATURE_LABEL}. gp_per_lead = rpl - cpl; roas = revenue/spend."
    ws["A2"].font = NOTE
    assign = result_12w.copy()
    base = ["state"]
    if current_map:
        assign["current_tier"] = assign["state"].map(current_map); base.append("current_tier")
    extra = [c for c in ("shrunk_rpl", "shrunk_gp_per_lead", "committed_tier",
                         "restructure_exposure", "boot_modal_tier", "boot_p_modal",
                         "B1_kmeans_2d_seed_stability", "C2_weighted_kmeans_seed_stability")
             if c in assign.columns]
    cols = base + ["clicks", "conversions", "spend", "rpl", "cpc", "cpl", "gp_per_lead", "roas"] \
        + MODEL_NAMES + extra
    cols = [c for c in cols if c in assign.columns]
    _write_df(ws, assign[cols].sort_values(FEATURE, ascending=False).reset_index(drop=True), header_row=4)

    ws = wb.create_sheet("Model Comparison")
    ws["A1"] = f"Model comparison (lower bid-loss / stdev on {FEATURE_LABEL} = tighter tiers)"; ws["A1"].font = TITLE
    _write_df(ws, model_compare, header_row=3)

    if elbow_df is not None and not elbow_df.empty:
        ws = wb.create_sheet("Elbow")
        ws["A1"] = f"Elbow — within-tier {FEATURE_LABEL} dispersion vs k (C2)"; ws["A1"].font = TITLE
        ws["A2"] = ("Swept on the headline model. Top: spend-weighted bid-loss (lower = tighter tiers). "
                    "Bottom: marginal_gain_pct = % bid-loss reduction vs k-1 — the performance gain of adding "
                    "one more tier. Pick k where that gain stops paying for the operational cost of another shell."); ws["A2"].font = NOTE
        _write_df(ws, elbow_df, header_row=4)
        if elbow_png and Path(elbow_png).exists():
            from openpyxl.drawing.image import Image as XLImage
            ws.add_image(XLImage(str(elbow_png)), "F4")

    if stability is not None and not stability.empty:
        ws = wb.create_sheet("Stability (12w_8w_4w)")
        ws["A1"] = "Tier stability across 12w / 8w / 4w windows"; ws["A1"].font = TITLE
        ws["A2"] = "stability_score = # distinct tiers across windows (1 = rock-stable)."; ws["A2"].font = NOTE
        st = stability[stability["model"] == headline].drop(columns=["model"])
        _write_df(ws, st.reset_index(drop=True), header_row=4)

    if halfsplit is not None and not halfsplit.empty:
        ws = wb.create_sheet("Half-Split Fragility")
        ws["A1"] = "Short-window fragility check — first vs second half"; ws["A1"].font = TITLE
        ws["A2"] = (f"H1 = {windows.w12_start}..{halfsplit_mid}, H2 = {halfsplit_mid + timedelta(days=1)}"
                    f"..{windows.end_date}, each re-clustered independently on the {headline} headline "
                    f"(halved volume gates). This replaces the 12w/8w/4w stability check, which is not "
                    f"meaningful at this window length. 'fragile' = the read is data-thin — a flag on the "
                    f"state's assignment, not a stability claim.")
        ws["A2"].font = NOTE
        _write_df(ws, halfsplit, header_row=4)

    if below is not None and not below.empty:
        ws = wb.create_sheet("Below Threshold")
        ws["A1"] = f"States below {click_threshold} clicks — nearest-tier on {FEATURE_LABEL}"; ws["A1"].font = TITLE
        _write_df(ws, below, header_row=3)

    if boot_df is not None and not boot_df.empty:
        ws = wb.create_sheet("Bootstrap Stability")
        ws["A1"] = f"Day-level bootstrap — P(tier) per state, {headline}"; ws["A1"].font = TITLE
        ws["A2"] = (f"{boot_used} resamples of the window's days (with replacement); eligibility "
                    f"re-applied per resample. boot_p_modal = share of eligible resamples landing on "
                    f"the modal tier: >=80% = firm read, <60% = the assignment is an artifact of which "
                    f"days happened to fall in the window. boot_eligible_pct < 100% means the state "
                    f"drifts below the volume gate under resampling — a thin read by construction.")
        ws["A2"].font = NOTE
        _write_df(ws, boot_df.sort_values("boot_p_modal", ascending=False, na_position="last")
                  .reset_index(drop=True), header_row=4)

    if incidence is not None and not incidence.empty:
        ws = wb.create_sheet("Bid-Down Incidence")
        ws["A1"] = "Partner bid-down incidence — per-state RPL(post)/RPL(pre)"; ws["A1"].font = TITLE
        s = incidence_summary or {}
        ws["A2"] = (f"Pre {s.get('pre_window','?')} ({s.get('pre_days','?')}d) vs the run window "
                    f"({s.get('post_days','?')}d). Median ratio "
                    f"{s.get('median_ratio', float('nan')):.3f}, IQR/median {s.get('iqr_over_median', float('nan')):.1%}, "
                    f"{s.get('n_deviant', '?')} deviant of {s.get('n_readable', '?')} readable. Read: tight ratios "
                    f"(low IQR/median, few DEVIANT) = the hit was ~uniform, so pre-period data preserves state "
                    f"ordering and a recency blend is defensible; dispersed ratios = the partner's demand mix hit "
                    f"states differentially, and blending pre-period data biases states relative to each other. "
                    f"DEVIANT states with low ratios are where that partner's demand concentrates. Per-day columns "
                    f"normalize volume for the unequal window lengths. Caution on volume reads: post-period "
                    f"clicks/day and spend/day also carry the tROAS tightenings and the tier-move learning resets, "
                    f"not just the bid-down — only the RPL ratio isolates the payout change.")
        ws["A2"].font = NOTE
        _write_df(ws, incidence, header_row=4)

    if ledger_cmp is not None and not ledger_cmp.empty:
        ws = wb.create_sheet("Ledger Comparison")
        ws["A1"] = "Proposed tier vs committed ledger"; ws["A1"].font = TITLE
        end_row = _write_df(ws, ledger_cmp, header_row=3)
        rd = ledger_cmp.dropna(subset=["proposed_tier", "committed_tier"])
        if not rd.empty:
            n_same = int((rd["agreement"] == "SAME").sum()); n = len(rd)
            line = f"Agreement among readable states: {n_same}/{n} ({n_same / n:.0%}) by state count"
            if "spend" in rd.columns and rd["spend"].notna().any():
                sp = rd["spend"].sum(); sp_same = rd.loc[rd["agreement"] == "SAME", "spend"].sum()
                if sp > 0:
                    line += (f"; ${sp_same:,.0f} of ${sp:,.0f} ({sp_same / sp:.0%}) by spend — "
                             f"the spend-weighted figure is the decision-relevant one.")
            ws.cell(row=end_row + 2, column=1, value=line).font = NOTE
            ws.cell(row=end_row + 3, column=1,
                    value=f"{len(ledger_cmp) - n} committed states had no proposal (below threshold "
                          f"or untargeted) and are excluded from both figures.").font = NOTE

    ws = wb.create_sheet("Methodology")
    notes = [
        ("Google Home geo-tiering — methodology", True), ("", False),
        (f"Clustering feature: {feature} ({FEATURE_LABEL}). TIER ONE = best {FEATURE_LABEL}.", False),
        ("Models: A2 quantile, B1 k-means(feature,cpc), B4 Ward(feature,cpc,log_clicks),", False),
        ("C2 weighted k-means on the feature (sample_weight=spend) = headline.", False),
        ("RPL blend mixes 4w into 12w per --rpl-blend, then gp_per_lead is re-derived.", False),
        ("", False),
        ("gp_per_lead = RPL - CPL (margin per lead): encodes the current CPC x lead-value tier", False),
        ("profile (Tier 1 cheap/efficient/positive-margin; Tier 2 premium/high-value/thin-margin;", False),
        ("Tier 3 volatile mid-pack; Tier 4 break-even/volume) in one number.", False),
        ("", False),
        ("Google data: account_id=1; campaign_performance_geo_daily_search (sum device rows; native conv,", False),
        ("~10-13% conservative on $); search-only; single-state campaigns excluded from the pool.", False),
        ("tROAS account: tier_target_tcpa_analytical is reference only; the output is tier membership.", False),
    ]
    if headline.startswith("B1"):
        notes += [
            ("", False),
            (f"B1 headline note: 2-D k-means on ({FEATURE_LABEL}, CPC), standardized. CPC shapes tier", False),
            (f"MEMBERSHIP, but tier LABELS are ordered by mean {FEATURE_LABEL} only — TIER ONE = highest", False),
            (f"{FEATURE_LABEL}, NOT best margin. Do not read these labels as the committed gp_per_lead tiers.", False),
        ]
    if window_note:
        notes += [
            ("", False),
            ("SHORT-WINDOW RUN — validation read against the committed ledger, NOT a re-tiering candidate.", False),
            ("CPC in this window carries restructure noise: tier moves completed in-window (see", False),
            ("restructure_exposure on State Assignments) each reset bidder learning (~1-2 wk). A state's", False),
            ("window-CPC can straddle two tROAS regimes plus a learning reset, so a (RPL, CPC) clustering", False),
            ("partially groups states by WHEN they moved, not what they are. Clean re-read: ~2 weeks after", False),
            ("the last in-window move settles. RPL is payout-driven and largely survives the restructure.", False),
        ]
    for r, (t, h) in enumerate(notes, 1):
        ws.cell(row=r, column=1, value=t).font = TITLE if h else NOTE
    ws.column_dimensions["A"].width = 112

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def load_ledger(path):
    try:
        import yaml
    except ImportError:
        return {}
    if not path.exists():
        return {}
    with open(path) as f:
        data = yaml.safe_load(f) or {}
    root = data.get("google_home_geo_tiers", data)
    return dict(root.get("commitments", {}))


def load_reference_workbook(path, tier_col="C2_weighted_kmeans"):
    """Read {state: tier} from a prior model workbook — the de facto commitment when no
    YAML ledger exists (same pattern as gp_bridge.read_zones_from_model, so the reference
    can never drift from what actually drove the rollout). Pulls the headline column from
    'State Assignments' plus 'Below Threshold' nearest-tier."""
    p = Path(path)
    if not p.exists():
        print(f"  WARNING: reference workbook not found: {p}", file=sys.stderr)
        return {}
    from openpyxl import load_workbook
    wb = load_workbook(p, data_only=True)
    out = {}

    def grab(sheet, col):
        if sheet not in wb.sheetnames:
            return
        rows = list(wb[sheet].iter_rows(values_only=True))
        hdr = next((r for r in rows if r and r[0] == "state"), None)
        if not hdr:
            return
        H = {h: i for i, h in enumerate(hdr)}
        if "state" not in H or col not in H:
            return
        for r in rows:
            if r and r[0] not in (None, "state") and r[H[col]]:
                out[r[H["state"]]] = r[H[col]]

    grab("State Assignments", tier_col)
    grab("Below Threshold", "assigned_tier_nearest")
    return out


def load_reference_locations(path):
    """Read an AS-BUILT {state: tier} map from a prepared campaign-location export (the
    account is the source of truth when the plan workbook drifted during execution).
    Expected columns (case-insensitive, one row per campaign x targeted state):
      state (or location)  -- 'Ohio' or 'Ohio, United States'
      tier                 -- TIER ONE..TIER FOUR / Standalone; if absent, derived from
                              campaign_name and the derivation is printed for confirmation
      campaign_id, campaign_name -- optional, audit trail
    Rows with negative/excluded locations should be removed before export."""
    p = Path(path)
    if not p.exists():
        print(f"  WARNING: reference locations file not found: {p}", file=sys.stderr)
        return {}
    df = pd.read_excel(p) if p.suffix.lower() in (".xlsx", ".xlsm") else pd.read_csv(p)
    df.columns = [str(c).strip().lower().replace(" ", "_") for c in df.columns]
    state_col = next((c for c in df.columns
                      if c.startswith("state") or "location" in c or c == "geo_state"), None)
    if state_col is None:
        print("  ERROR: reference locations needs a 'state' (or 'location') column.", file=sys.stderr)
        return {}
    df["_state"] = df[state_col].astype(str).str.split(",").str[0].str.strip()

    valid = set(TIER_LABEL_BASE) | {"STANDALONE"}
    if "tier" in df.columns:
        df["_tier"] = df["tier"].astype(str).str.strip().str.upper()
    elif "campaign_name" in df.columns:
        def derive(nm):
            u = str(nm).upper()
            for t in TIER_LABEL_BASE:
                if t in u:
                    return t
            for s in ("CALIFORNIA", "TEXAS", "FLORIDA"):
                if s in u:
                    return "STANDALONE"
            return ""
        df["_tier"] = df["campaign_name"].map(derive)
        print("  tier derived from campaign_name — confirm this mapping:")
        for nm, t in df.groupby("campaign_name")["_tier"].first().items():
            print(f"    {nm}  ->  {t or 'UNMAPPED'}")
    else:
        print("  ERROR: reference locations needs a 'tier' or 'campaign_name' column.", file=sys.stderr)
        return {}

    bad = df[~df["_tier"].isin(valid)]
    if not bad.empty:
        print(f"  WARNING: {len(bad)} rows with unrecognized tier dropped: "
              f"{sorted(bad['_tier'].unique())}", file=sys.stderr)
        df = df[df["_tier"].isin(valid)]
    df = df[df["_tier"] != "STANDALONE"]

    dup = df.groupby("_state")["_tier"].nunique()
    conflicts = sorted(dup[dup > 1].index)
    if conflicts:
        print(f"  WARNING: states targeted in multiple tiers (fix the export; keeping first "
              f"occurrence): {', '.join(conflicts)}", file=sys.stderr)
    out = dict(zip(df["_state"], df["_tier"]))  # later rows win; conflicts warned above
    for s in conflicts:
        first = df[df["_state"] == s]["_tier"].iloc[0]
        out[s] = first
    return out


def emit_ledger(path, ref_map):
    """Freeze the loaded reference map as a YAML ledger, whenever you're ready to formalize."""
    import yaml
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    with open(p, "w") as f:
        yaml.safe_dump({"google_home_geo_tiers": {"commitments": dict(sorted(ref_map.items()))}},
                       f, sort_keys=False, allow_unicode=True)
    print(f"  ledger frozen -> {p} ({len(ref_map)} states)")


# ----------------------------------------------------------------------------
# Main
# ----------------------------------------------------------------------------

def main(argv=None):
    global FEATURE, FEATURE_LABEL
    p = argparse.ArgumentParser()
    p.add_argument("--end-date", default=None)
    p.add_argument("--start-date", default=None,
                   help="Explicit window start (YYYY-MM-DD). Overrides the 12w lookback; the "
                        "12w/8w/4w stability check is replaced by a half-split fragility check, "
                        "the RPL blend is skipped, and the output is stamped as a validation run.")
    p.add_argument("--k", type=int, default=DEFAULT_K)
    p.add_argument("--rpl-blend", type=float, default=0.0)
    p.add_argument("--click-threshold", type=int, default=CLICK_THRESHOLD)
    p.add_argument("--conv-threshold", type=float, default=None,
                   help="Eligibility gate on conversions instead of clicks (recommended for "
                        "short windows, where RPL is conversion-count-limited; e.g. 9).")
    p.add_argument("--seed-sweep", default=None,
                   help="Comma-separated seeds (e.g. '42,7,13,99,2026'); adds per-state "
                        "seed-stability columns for the stochastic models (B1, C2).")
    p.add_argument("--bootstrap", type=int, default=0,
                   help="Day-level bootstrap iterations (e.g. 200): resample the window's "
                        "days, re-run the headline model, report P(modal tier) per state. "
                        "The probabilistic upgrade over the half-split fragility flag.")
    p.add_argument("--shrinkage", type=float, default=0.0,
                   help="Empirical-Bayes prior strength in pseudo-conversions (e.g. 10): adds "
                        "shrunk_rpl / shrunk_gp_per_lead diagnostic columns pulling thin states "
                        "toward the pool mean, plus a shrunk nearest-tier for below-threshold "
                        "states. Headline clustering stays on the raw feature.")
    p.add_argument("--incidence-pre", default=None,
                   help="Pre-bid-down window as START:END (e.g. 2026-05-01:2026-06-17): adds a "
                        "Bid-Down Incidence sheet testing per-state RPL(post)/RPL(pre) ratio "
                        "uniformity — the empirical answer to whether a recency blend is safe.")
    p.add_argument("--cluster-feature", choices=["gp_per_lead", "roas", "rpl"], default="gp_per_lead")
    p.add_argument("--exclude-states", default=",".join(DEFAULT_EXCLUDE),
                   help="Comma-separated states to drop from the pool (standalone campaigns). "
                        "Pass '' to include all.")
    p.add_argument("--elbow", action="store_true", help="Sweep k and write an elbow chart.")
    p.add_argument("--k-range", default="2,3,4,5,6,7,8,9,10,11,12")
    p.add_argument("--headline-model", default="C2_weighted_kmeans")
    p.add_argument("--db-path", default=DEFAULT_DB_PATH)
    p.add_argument("--output-dir", default=DEFAULT_OUTPUT_DIR)
    p.add_argument("--ledger", default=None)
    p.add_argument("--reference-xlsx", default=None,
                   help="Prior model workbook (e.g. google_geo_model_gp_per_lead_2026-06-09.xlsx) "
                        "to use as the reference/committed tier map when no YAML ledger exists. "
                        "Reads State Assignments (C2 column) + Below Threshold, like gp_bridge.")
    p.add_argument("--reference-locations", default=None,
                   help="AS-BUILT campaign-location export (csv/xlsx; one row per campaign x "
                        "targeted state, columns: state, tier[, campaign_id, campaign_name]). "
                        "Preferred reference when the plan workbook drifted during execution. "
                        "Precedence: --ledger > --reference-locations > --reference-xlsx.")
    p.add_argument("--emit-ledger", default=None,
                   help="Write the loaded reference map to this YAML path (freeze the ledger "
                        "whenever you're ready to formalize). Requires --reference-xlsx or --ledger.")
    args = p.parse_args(argv)

    FEATURE = args.cluster_feature
    FEATURE_LABEL = {"gp_per_lead": "GP/lead", "roas": "ROAS", "rpl": "RPL"}[FEATURE]
    excluded = [s.strip() for s in args.exclude_states.split(",") if s.strip()]

    db = Path(args.db_path)
    if not db.exists():
        print(f"ERROR: DB not found: {db}", file=sys.stderr); return 1
    conn = sqlite3.connect(str(db))
    end_date = resolve_end_date(conn, args.end_date)
    custom_start = date.fromisoformat(args.start_date) if args.start_date else None
    if custom_start:
        if custom_start >= end_date:
            print(f"ERROR: --start-date {custom_start} >= end date {end_date}.", file=sys.stderr)
            return 1
        windows = Windows(end_date=end_date, w12_start=custom_start,
                          w8_start=custom_start, w4_start=custom_start)
    else:
        windows = build_windows(end_date)
    n_days = (windows.end_date - windows.w12_start).days + 1
    k = args.k
    gate = (f"conv_threshold={args.conv_threshold:g}" if args.conv_threshold is not None
            else f"click_threshold={args.click_threshold}")
    print("GOOGLE HOME GEO-TIERING" + ("  [SHORT-WINDOW VALIDATION RUN]" if custom_start else ""))
    print(f"  feature={FEATURE} ({FEATURE_LABEL})  window={windows.w12_start}..{end_date} ({n_days}d)  "
          f"k={k}  rpl_blend={args.rpl_blend}  {gate}")
    print(f"  excluded (standalone): {excluded or 'none'}")

    raw12 = pull_state_aggregates(conn, windows.w12_start, windows.end_date)
    if raw12.empty:
        print(f"ERROR: no rows in {GEO_TABLE} for the window.", file=sys.stderr); return 1
    raw4 = None
    use_blend = args.rpl_blend > 0.0 and not custom_start
    if args.rpl_blend > 0.0 and custom_start:
        print("  NOTE: --rpl-blend skipped — with --start-date the 4w window equals the full "
              "window, so the blend is a no-op.")
    if use_blend:
        raw4 = pull_state_aggregates(conn, windows.w4_start, windows.end_date)
        raw4 = compute_derived(raw4) if not raw4.empty else None

    alls = compute_derived(raw12)
    if use_blend:
        alls = apply_rpl_blend(alls, raw4, args.rpl_blend)
    if excluded:
        alls = alls[~alls["state"].isin(excluded)].reset_index(drop=True)
    if args.shrinkage > 0:
        alls = apply_shrinkage(alls, args.shrinkage)

    elig = eligible_pool(alls, click_threshold=args.click_threshold,
                         conv_threshold=args.conv_threshold)
    print(f"  eligible states: {len(elig)} of {len(alls)} (after exclusions)")
    if len(elig) < k:
        print(f"ERROR: eligible pool ({len(elig)}) < k ({k}).", file=sys.stderr); return 1

    result12, _ = apply_all_models(elig, k)
    headline = args.headline_model

    halfsplit = None; halfsplit_mid = None
    if custom_start:
        stability = pd.DataFrame()
        halfsplit, halfsplit_mid = build_halfsplit_table(
            conn, windows.w12_start, windows.end_date, k, headline,
            args.click_threshold, args.conv_threshold, excluded, result12)
        n_frag = int(halfsplit["flag"].str.startswith(("fragile", "thin")).sum())
        print(f"  half-split fragility: {n_frag} of {len(halfsplit)} states flagged")
    else:
        r8 = run_models_for_window(conn, windows.w8_start, windows.end_date, k, 75, excluded)
        r4 = run_models_for_window(conn, windows.w4_start, windows.end_date, k, 75, excluded)
        stability = build_stability_table(result12, r8, r4)

    if args.seed_sweep:
        seeds = [int(x) for x in args.seed_sweep.split(",") if x.strip()]
        seed_df = seed_sweep_stability(elig, k, seeds)
        result12 = result12.merge(seed_df, on="state", how="left")
        for m in ("B1_kmeans_2d", "C2_weighted_kmeans"):
            col = f"{m}_seed_stability"
            n_flip = int((result12[col] != "stable").sum())
            print(f"  seed sweep ({len(seeds)} seeds) {m}: {n_flip} of {len(result12)} states flip")

    boot_df = None; boot_used = 0
    if args.bootstrap > 0:
        daily = pull_state_daily(conn, windows.w12_start, windows.end_date)
        boot_df, boot_used = bootstrap_assignments(
            daily, k, headline, args.bootstrap, args.click_threshold,
            args.conv_threshold, excluded)
        result12 = result12.merge(
            boot_df[["state", "boot_modal_tier", "boot_p_modal", "boot_eligible_pct"]],
            on="state", how="left")
        firm = int((result12["boot_p_modal"] >= 0.8).sum())
        soft = int((result12["boot_p_modal"] < 0.6).sum())
        print(f"  bootstrap ({boot_used} usable of {args.bootstrap}): "
              f"{firm} of {len(result12)} states P(modal)>=80%, {soft} below 60%")

    incidence = None; incidence_summary = {}
    if args.incidence_pre:
        incidence, incidence_summary = bid_down_incidence(
            conn, args.incidence_pre, windows.w12_start, windows.end_date, excluded)
        if incidence_summary:
            print(f"  bid-down incidence: median ratio {incidence_summary['median_ratio']:.3f}, "
                  f"IQR/median {incidence_summary['iqr_over_median']:.1%}, "
                  f"{incidence_summary['n_deviant']} deviant of {incidence_summary['n_readable']} readable")

    tier_metrics = tier_aggregate_metrics(result12, headline)
    model_compare = pd.DataFrame([model_summary_metrics(result12, m) for m in MODEL_NAMES])
    below = assign_below_threshold(alls, result12, headline, args.click_threshold,
                                   conv_threshold=args.conv_threshold)

    # Reference column: on short windows the spend-share/name-based current-tier map straddles
    # the T1/T2/T3 moves and post-rename labels, so it is skipped — the committed ledger is
    # the reference instead (state-level performance is campaign-agnostic either way).
    commit = load_ledger(Path(args.ledger)) if args.ledger else {}
    if not commit and args.reference_locations:
        commit = load_reference_locations(args.reference_locations)
        if commit:
            print(f"  reference tiers from AS-BUILT location export: {args.reference_locations} "
                  f"({len(commit)} states)")
    if not commit and args.reference_xlsx:
        commit = load_reference_workbook(args.reference_xlsx)
        if commit:
            print(f"  reference tiers from workbook: {args.reference_xlsx} ({len(commit)} states)")
    commit = {s: t for s, t in commit.items() if s not in excluded}
    if args.emit_ledger:
        if commit:
            emit_ledger(args.emit_ledger, commit)
        else:
            print("  NOTE: --emit-ledger skipped — no reference loaded.", file=sys.stderr)
    if custom_start:
        current_map = {}
        if not commit:
            print("  NOTE: name-based current-tier map skipped (window straddles the re-tier "
                  "moves); pass --reference-xlsx (prior model workbook) or --ledger for the "
                  "committed reference column.")
    else:
        current_map = pull_state_to_current_tier(conn, windows.w12_start, windows.end_date)
    if commit:
        result12["committed_tier"] = result12["state"].map(commit)
        def _exposure(t):
            d = MOVE_DATES.get(t)
            if d and windows.w12_start <= date.fromisoformat(d) <= windows.end_date:
                return f"moved {d} (in-window)"
            return ""
        result12["restructure_exposure"] = result12["committed_tier"].map(_exposure)

    elbow_df = None; elbow_png = None
    if args.elbow:
        kv = [int(x) for x in args.k_range.split(",")]
        elbow_png = Path(args.output_dir) / f"google_geo_elbow_{FEATURE}_{headline}_{end_date}.png"
        elbow_df = elbow_sweep(elig, kv, elbow_png, model=headline)

    ledger_cmp = None
    if commit:
        prop = dict(zip(result12["state"], result12[headline]))
        ledger_cmp = pd.DataFrame([
            {"state": s, "proposed_tier": prop.get(s), "committed_tier": commit.get(s),
             "agreement": "SAME" if prop.get(s) == commit.get(s) else "MOVED"}
            for s in sorted(set(prop) | set(commit))])
        ledger_cmp["spend"] = ledger_cmp["state"].map(dict(zip(alls["state"], alls["spend"])))

    conn.close()
    print(f"\n  Proposed tiers ({FEATURE_LABEL}):")
    for _, r in tier_metrics.iterrows():
        print(f"    {r['tier']:12} n={int(r['n_states']):2}  {FEATURE_LABEL} "
              f"{r['feat_min']:.0f}..{r['feat_max']:.0f}  wtd_cpc ${r['wtd_cpc']:.0f}  "
              f"wtd_rpl ${r['wtd_rpl']:.0f}  wtd_roas {r['wtd_roas']:.2f}")
        print(f"                 {r['states']}")
    if elbow_df is not None and not elbow_df.empty:
        print("\n  Elbow (k -> bid-loss):")
        for _, r in elbow_df.iterrows():
            print(f"    k={int(r['k'])}: {r['spend_weighted_bid_loss']:,.0f}")

    window_note = ""
    if custom_start:
        in_window = {t: d for t, d in MOVE_DATES.items()
                     if windows.w12_start <= date.fromisoformat(d) <= windows.end_date}
        moves_txt = ", ".join(f"{t} {d}" for t, d in in_window.items()) or "none"
        window_note = (f"SHORT-WINDOW VALIDATION RUN ({n_days}d, post partner bid-down): read this "
                       f"against the committed ledger — NOT a re-tiering candidate. CPC carries "
                       f"restructure noise from in-window tier moves ({moves_txt}); a (RPL, CPC) "
                       f"clustering partially groups states by when they moved. Clean re-read ~2 "
                       f"weeks after the last move settles.")

    if custom_start:
        # 'validation' prefix keeps gp_bridge's google_geo_model_{feature}_* auto-discovery
        # from ever treating a short-window run as the authoritative zone source.
        out = Path(args.output_dir) / (f"google_geo_model_validation_{FEATURE}_{headline}_"
                                       f"{windows.w12_start}_{end_date}.xlsx")
    else:
        out = Path(args.output_dir) / f"google_geo_model_{FEATURE}_{end_date}.xlsx"
    write_workbook(out, windows, k, args.rpl_blend, args.click_threshold, FEATURE, result12,
                   headline, tier_metrics, model_compare, stability, below, current_map,
                   ledger_cmp, elbow_df, elbow_png, excluded, halfsplit=halfsplit,
                   halfsplit_mid=halfsplit_mid, window_note=window_note,
                   conv_threshold=args.conv_threshold, boot_df=boot_df, boot_used=boot_used,
                   incidence=incidence, incidence_summary=incidence_summary)
    print(f"\n  Wrote: {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
