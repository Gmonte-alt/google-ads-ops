"""xlsx_out.py -- shared Excel writer for the change-history scripts (un-numbered module)."""

def write_xlsx(rows, path, cols, sheet_title="Changes"):
    """rows = list of dicts; cols = ordered list of dict keys to write."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("[!] openpyxl not installed. Run:  pip install openpyxl   (or use --csv)")
        return False
    import collections

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append([c.replace("_", " ").upper() for c in cols])
    for i in range(1, len(cols) + 1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
        c.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append([r.get(k, "") for k in cols])
    for i, k in enumerate(cols, 1):
        w = {"platform": 11, "date": 12, "campaign": 46, "ad_group": 24, "category": 20,
             "klass": 14, "header": 36, "matches": 9, "sample": 80, "detail": 60,
             "count": 8, "time": 10}.get(k, 18)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions

    s = wb.create_sheet("Summary")
    s.append(["METRIC", "VALUE"])
    s["A1"].font = s["B1"].font = Font(bold=True, color="FFFFFF")
    s["A1"].fill = s["B1"].fill = PatternFill("solid", fgColor="1F4E79")
    s.append(["total records", len(rows)])
    for label, key in (("by platform", "platform"), ("by class", "klass"), ("by category", "category")):
        s.append([])
        s.append([label.upper(), ""])
        s.cell(row=s.max_row, column=1).font = Font(bold=True)
        for k, n in collections.Counter(r.get(key, "") for r in rows).most_common():
            s.append([k, n])
    s.column_dimensions["A"].width = 34
    s.column_dimensions["B"].width = 12

    wb.save(path)
    print("[ok] wrote %d rows -> %s" % (len(rows), path))
    return True