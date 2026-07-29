"""
state_lifecycle_study.py — longitudinal state performance study for Google Home (account 1).

Answers the management questions directly:
  - "What were good states by a longer lookback?"  -> era-by-era metrics per state
  - "Was Delaware good once? Can it come back?"    -> per-state deep-dive sheets + charts
  - "Known good / known bad performers?"           -> lifecycle classification
  - "Focus on RPL / future revenue"               -> RPL stability (structural vs repriced
                                                     premium), the credibility input a
                                                     recency-blended RPL model needs.

Method:
  The window is cut into ERAS at the known regime boundaries (edit the ERAS constant as
  needed). Per state x era: conv/day, clicks/day, spend/day, RPL, CPC, CPL, GP/lead, ROAS.
  Each state is classified by baseline-era vs current-era economics (consistent_good,
  consistent_bad, decayed, recovered, emerged, thin) with a cause tag. For decayed states,
  the baseline->current GP/day change is decomposed into volume / value / efficiency /
  interaction (the gp_bridge identity, per-day normalized). Deep-dive states get a weekly
  series sheet plus a 3-panel chart with era boundaries marked.

Basis: native conversions (sum device rows, device<>'ALL'), ~10-13% conservative on $ with
~24h routing lag — end date defaults to (max date in DB) - 1.

Usage:
    python state_lifecycle_study.py
    python state_lifecycle_study.py --deep-dive "Delaware,Colorado,Oregon" --end-date 2026-07-21
"""
from __future__ import annotations
import argparse, sqlite3, sys
from datetime import date, timedelta
from pathlib import Path
import numpy as np
import pandas as pd

DEFAULT_DB = r"C:\Projects\google-ads-ops\db\paid_search.db"
DEFAULT_OUT = r"C:\Projects\google-ads-ops\analysis\google-geo-analysis\output"
ACCOUNT_ID = 1
GEO_TABLE = "campaign_performance_geo_daily_search"

# Regime boundaries (edit as history accrues). Baseline start is clamped to data start.
ERAS = [
    ("E1 baseline",     None,           "2026-04-30"),   # pre-tightening
    ("E2 tightening",   "2026-05-01",   "2026-06-08"),   # May tROAS tightening -> margin loss
    ("E3 cliff",        "2026-06-09",   "2026-06-18"),   # overcorrection cliff -> partner bid-down
    ("E4 new regime",   "2026-06-19",   None),           # post bid-down (current)
]
BASELINE_ERA, CURRENT_ERA = "E1 baseline", "E4 new regime"

GOOD_ROAS, BAD_ROAS = 1.05, 0.95
CONV_PER_DAY_MIN = 0.12          # era readability floor (~1 conv / 8 days)
REPRICED_BELOW, IMPROVED_ABOVE = 0.85, 1.10


def pull_daily(conn, start, end):
    sql = f"""
        SELECT geo_state AS state, date,
               SUM(clicks) AS clicks, SUM(cost_micros)/1e6 AS spend,
               SUM(conversions) AS conversions, SUM(conversion_value) AS revenue
        FROM {GEO_TABLE}
        WHERE account_id=? AND device<>'ALL' AND geo_state IS NOT NULL
          AND date BETWEEN ? AND ?
        GROUP BY geo_state, date
    """
    df = pd.read_sql_query(sql, conn, params=(ACCOUNT_ID, start.isoformat(), end.isoformat()))
    df = df[df["state"].str.upper() != "UNKNOWN"].copy()
    df["date"] = pd.to_datetime(df["date"]).dt.date
    return df


def derive(g):
    out = {}
    out["conv_day"] = g["conversions"] / g["days"]
    out["clicks_day"] = g["clicks"] / g["days"]
    out["spend_day"] = g["spend"] / g["days"]
    out["rpl"] = np.where(g["conversions"] > 0, g["revenue"] / g["conversions"], np.nan)
    out["cpc"] = np.where(g["clicks"] > 0, g["spend"] / g["clicks"], np.nan)
    out["cpl"] = np.where(g["conversions"] > 0, g["spend"] / g["conversions"], np.nan)
    out["gp_per_lead"] = out["rpl"] - out["cpl"]
    out["roas"] = np.where(g["spend"] > 0, g["revenue"] / g["spend"], np.nan)
    return pd.DataFrame(out, index=g.index)


def era_table(daily, eras):
    """states x eras metric table."""
    rows = []
    for name, lo, hi in eras:
        sel = daily[(daily["date"] >= lo) & (daily["date"] <= hi)]
        days = (hi - lo).days + 1
        g = sel.groupby("state")[["clicks", "spend", "conversions", "revenue"]].sum()
        g["days"] = days
        g = pd.concat([g, derive(g)], axis=1).reset_index()
        g.insert(1, "era", name)
        rows.append(g)
    return pd.concat(rows, ignore_index=True)


def classify(et):
    """Per-state lifecycle class from baseline vs current era."""
    base = et[et["era"] == BASELINE_ERA].set_index("state")
    cur = et[et["era"] == CURRENT_ERA].set_index("state")
    states = sorted(set(base.index) | set(cur.index))
    rows = []
    for s in states:
        b = base.loc[s] if s in base.index else None
        c = cur.loc[s] if s in cur.index else None
        b_ok = b is not None and b["conv_day"] >= CONV_PER_DAY_MIN
        c_ok = c is not None and c["conv_day"] >= CONV_PER_DAY_MIN
        rpl_ratio = (c["rpl"] / b["rpl"]) if (b_ok and c_ok and b["rpl"] > 0) else np.nan
        if not b_ok and not c_ok:
            klass, cause = "thin", "insufficient volume both eras"
        elif not b_ok and c_ok:
            klass = "emerged" if c["roas"] >= GOOD_ROAS else "thin_baseline"
            cause = "no readable baseline"
        elif b_ok and not c_ok:
            klass, cause = "collapsed", "volume gone in current era"
        else:
            bg, cg = b["roas"] >= GOOD_ROAS, c["roas"] >= GOOD_ROAS
            bb, cb = b["roas"] < BAD_ROAS, c["roas"] < BAD_ROAS
            if bg and cg: klass = "consistent_good"
            elif bb and cb: klass = "consistent_bad"
            elif bg and cb: klass = "decayed"
            elif bb and cg: klass = "recovered"
            else: klass = "borderline"
            if not np.isnan(rpl_ratio) and rpl_ratio < REPRICED_BELOW:
                cause = f"value repriced ({rpl_ratio:.2f}x)"
            elif not np.isnan(rpl_ratio) and rpl_ratio > IMPROVED_ABOVE:
                cause = f"value improved ({rpl_ratio:.2f}x)"
            else:
                cause = "value held" if not np.isnan(rpl_ratio) else ""
        premium = ("repriced" if (not np.isnan(rpl_ratio) and rpl_ratio < REPRICED_BELOW)
                   else "improved" if (not np.isnan(rpl_ratio) and rpl_ratio > IMPROVED_ABOVE)
                   else "structural" if not np.isnan(rpl_ratio) else "")
        rows.append({
            "state": s, "class": klass, "cause": cause, "rpl_premium": premium,
            "base_conv_day": b["conv_day"] if b is not None else np.nan,
            "base_rpl": b["rpl"] if b is not None else np.nan,
            "base_roas": b["roas"] if b is not None else np.nan,
            "cur_conv_day": c["conv_day"] if c is not None else np.nan,
            "cur_rpl": c["rpl"] if c is not None else np.nan,
            "cur_roas": c["roas"] if c is not None else np.nan,
            "rpl_ratio": rpl_ratio,
            "cur_spend_day": c["spend_day"] if c is not None else 0.0,
        })
    df = pd.DataFrame(rows)
    order = {"consistent_good": 0, "recovered": 1, "emerged": 2, "borderline": 3,
             "decayed": 4, "consistent_bad": 5, "collapsed": 6, "thin_baseline": 7, "thin": 8}
    df["_o"] = df["class"].map(order)
    return df.sort_values(["_o", "cur_spend_day"], ascending=[True, False]).drop(columns="_o")


def decompose(et, states):
    """Baseline->current GP/day bridge per state: volume + value + efficiency + interaction."""
    base = et[et["era"] == BASELINE_ERA].set_index("state")
    cur = et[et["era"] == CURRENT_ERA].set_index("state")
    rows = []
    for s in states:
        if s not in base.index or s not in cur.index:
            continue
        b, c = base.loc[s], cur.loc[s]
        if not (b["conversions"] > 0 and c["conversions"] > 0):
            continue
        gplB, gplC = b["gp_per_lead"], c["gp_per_lead"]
        dGP = c["conv_day"] * gplC - b["conv_day"] * gplB
        vol = (c["conv_day"] - b["conv_day"]) * gplB
        val = b["conv_day"] * (c["rpl"] - b["rpl"])
        eff = -b["conv_day"] * (c["cpl"] - b["cpl"])
        inter = dGP - vol - val - eff
        comps = {"volume": vol, "value": val, "efficiency": eff}
        driver = max(comps, key=lambda k: abs(comps[k]))
        rows.append({"state": s, "gp_day_base": b["conv_day"] * gplB, "gp_day_cur": c["conv_day"] * gplC,
                     "delta_gp_day": dGP, "volume_effect": vol, "value_effect": val,
                     "efficiency_effect": eff, "interaction": inter, "primary_driver": driver})
    return pd.DataFrame(rows).sort_values("delta_gp_day")


def weekly_series(daily, state):
    d = daily[daily["state"] == state].copy()
    if d.empty:
        return d
    d["dt"] = pd.to_datetime(d["date"])
    d["week"] = (d["dt"] - pd.to_timedelta(d["dt"].dt.weekday, unit="D")).dt.date
    g = d.groupby("week")[["clicks", "spend", "conversions", "revenue"]].sum()
    g["days"] = d.groupby("week")["date"].nunique()
    g = pd.concat([g, derive(g)], axis=1).reset_index()
    return g


def deep_dive_chart(wk, state, eras, out_png):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    fig, axes = plt.subplots(3, 1, figsize=(9.5, 7.5), sharex=True)
    x = pd.to_datetime(wk["week"])
    panels = [("conv_day", "conversions/day", "#c0504d"),
              ("rpl", "RPL ($)", "#4f81bd"),
              ("roas", "ROAS", "#2e7d32")]
    for ax, (col, lbl, color) in zip(axes, panels):
        ax.plot(x, wk[col], "-o", ms=3, lw=1.8, color=color)
        ax.set_ylabel(lbl, fontsize=9); ax.grid(alpha=0.25)
        for name, lo, hi in eras[1:]:
            ax.axvline(pd.to_datetime(lo), color="#999", lw=0.9, ls=":")
    axes[2].axhline(1.0, color="#999", lw=0.8, ls="--")
    for name, lo, hi in eras:
        mid = pd.to_datetime(lo) + (pd.to_datetime(hi) - pd.to_datetime(lo)) / 2
        axes[0].annotate(name.split(" ", 1)[1], (mid, axes[0].get_ylim()[1]),
                         ha="center", va="bottom", fontsize=7.5, color="#666")
    axes[0].set_title(f"{state} — lifecycle across regimes (native conv basis)", fontsize=10.5)
    fig.tight_layout(); fig.savefig(out_png, dpi=140); plt.close(fig)


def _write_df(ws, df, header_row):
    from openpyxl.styles import Font, PatternFill
    HF = Font(bold=True, color="FFFFFF"); HFill = PatternFill("solid", fgColor="305496")
    twod = {"rpl", "cpc", "cpl", "gp_per_lead", "roas", "base_rpl", "cur_rpl", "base_roas",
            "cur_roas", "rpl_ratio", "conv_day", "clicks_day", "base_conv_day", "cur_conv_day",
            "gp_day_base", "gp_day_cur", "delta_gp_day", "volume_effect", "value_effect",
            "efficiency_effect", "interaction"}
    money = {"spend", "revenue", "spend_day", "cur_spend_day"}
    for j, c in enumerate(df.columns, 1):
        x = ws.cell(row=header_row, column=j, value=str(c)); x.font = HF; x.fill = HFill
    for i, (_, r) in enumerate(df.iterrows(), header_row + 1):
        for j, c in enumerate(df.columns, 1):
            v = r[c]
            if isinstance(v, (np.floating, np.integer)):
                v = None if pd.isna(v) else float(v)
            cell = ws.cell(row=i, column=j, value=v)
            if c in money:
                cell.number_format = "$#,##0"
            elif c in twod:
                cell.number_format = "0.00"
    ws.freeze_panes = ws.cell(row=header_row + 1, column=2)
    return header_row + len(df)


def write_workbook(path, start, end, eras, cls, et, decomp, dives, charts):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.drawing.image import Image as XLImage
    TITLE = Font(bold=True, size=14, color="305496"); NOTE = Font(italic=True, color="595959", size=9)
    wb = Workbook(); wb.remove(wb.active)

    ws = wb.create_sheet("Lifecycle Summary")
    ws["A1"] = "State lifecycle classification — longer-lookback study"; ws["A1"].font = TITLE
    era_txt = "; ".join(f"{n}: {lo}..{hi}" for n, lo, hi in eras)
    ws["A2"] = (f"Window {start}..{end}, native conv basis (~10-13% conservative on $). Eras: {era_txt}. "
                f"Class compares {BASELINE_ERA} vs {CURRENT_ERA} (good >= {GOOD_ROAS} ROAS, bad < {BAD_ROAS}; "
                f"readable >= {CONV_PER_DAY_MIN} conv/day). rpl_premium: structural = value held across the "
                f"regime change; repriced = current RPL < {REPRICED_BELOW:.0%} of baseline (partner-fragile "
                f"premium — weight history cautiously in any RPL model); improved = > {IMPROVED_ABOVE:.0%}.")
    ws["A2"].font = NOTE
    _write_df(ws, cls, header_row=4)

    for metric, sheet in (("roas", "Trend ROAS"), ("rpl", "Trend RPL"), ("conv_day", "Trend Volume")):
        ws = wb.create_sheet(sheet)
        ws["A1"] = f"{sheet} — states x eras"; ws["A1"].font = TITLE
        piv = et.pivot_table(index="state", columns="era", values=metric, aggfunc="first")
        piv = piv[[n for n, _, _ in eras]].reset_index()
        _write_df(ws, piv, header_row=3)

    if not decomp.empty:
        ws = wb.create_sheet("Decay Decomposition")
        ws["A1"] = "Why GP/day changed — baseline vs current era bridge"; ws["A1"].font = TITLE
        ws["A2"] = ("delta_gp_day = volume + value + efficiency + interaction (exact identity). "
                    "primary_driver names the largest component: 'value' = payout repricing; "
                    "'efficiency' = cost-side (CPL) movement, incl. our own bid policy; "
                    "'volume' = scale change at old margins.")
        ws["A2"].font = NOTE
        _write_df(ws, decomp, header_row=4)

    for state, wk in dives.items():
        ws = wb.create_sheet(f"Dive {state[:24]}")
        ws["A1"] = f"{state} — weekly series"; ws["A1"].font = TITLE
        cols = ["week", "days", "clicks", "spend", "conversions", "revenue", "conv_day",
                "rpl", "cpc", "cpl", "gp_per_lead", "roas"]
        endrow = _write_df(ws, wk[cols], header_row=3)
        if state in charts:
            ws.add_image(XLImage(str(charts[state])), f"N3")

    ws = wb.create_sheet("Methodology")
    notes = [
        "Longer-lookback lifecycle study. Grain: state x era (regime-bounded), plus weekly deep dives.",
        "Native conversion basis; ~10-13% conservative on $; end date trailing 1 day for routing lag.",
        "Era boundaries encode the known regime changes: May tROAS tightening, June overcorrection",
        "cliff, 6/18 partner bid-down. Classification is baseline-vs-current; the trend sheets show",
        "the path between. 'repriced' premiums were partner-fragile — an RPL tiering model that",
        "blends history should discount those states' baseline RPL; 'structural' premiums survived",
        "the regime change and can be trusted at longer lookbacks.",
    ]
    for i, t in enumerate(notes, 1):
        ws.cell(row=i, column=1, value=t).font = NOTE
    wb.save(path)


def main(argv=None):
    p = argparse.ArgumentParser()
    p.add_argument("--db-path", default=DEFAULT_DB)
    p.add_argument("--output-dir", default=DEFAULT_OUT)
    p.add_argument("--start-date", default=None, help="Default: first date in the DB.")
    p.add_argument("--end-date", default=None, help="Default: max date in DB minus 1 (routing lag).")
    p.add_argument("--deep-dive", default="Delaware,Colorado",
                   help="Comma-separated states for weekly deep-dive sheets + charts.")
    args = p.parse_args(argv)

    db = Path(args.db_path)
    if not db.exists():
        print(f"ERROR: DB not found: {db}", file=sys.stderr); return 1
    conn = sqlite3.connect(str(db))
    lo, hi = pd.read_sql_query(
        f"SELECT MIN(date) lo, MAX(date) hi FROM {GEO_TABLE} WHERE account_id=?",
        conn, params=(ACCOUNT_ID,)).iloc[0]
    start = date.fromisoformat(args.start_date) if args.start_date else date.fromisoformat(lo)
    end = date.fromisoformat(args.end_date) if args.end_date else date.fromisoformat(hi) - timedelta(days=1)

    eras = []
    for name, elo, ehi in ERAS:
        elo_d = start if elo is None else max(start, date.fromisoformat(elo))
        ehi_d = end if ehi is None else min(end, date.fromisoformat(ehi))
        if elo_d <= ehi_d:
            eras.append((name, elo_d, ehi_d))
    print(f"STATE LIFECYCLE STUDY  window {start}..{end}")
    for n, elo_d, ehi_d in eras:
        print(f"  {n}: {elo_d}..{ehi_d} ({(ehi_d-elo_d).days+1}d)")

    daily = pull_daily(conn, start, end)
    conn.close()
    if daily.empty:
        print("ERROR: no rows.", file=sys.stderr); return 1

    et = era_table(daily, eras)
    cls = classify(et)
    counts = cls["class"].value_counts()
    print("\nLifecycle classes:")
    for k, v in counts.items():
        members = cls[cls["class"] == k]["state"].tolist()
        print(f"  {k:<16} {v:>2}  {', '.join(members[:8])}{' ...' if v > 8 else ''}")

    decay_states = cls[cls["class"].isin(["decayed", "consistent_bad", "collapsed"])]["state"].tolist()
    decomp = decompose(et, decay_states)
    if not decomp.empty:
        print("\nDecay drivers (largest GP/day losses):")
        for _, r in decomp.head(8).iterrows():
            print(f"  {r['state']:<16} dGP/day {r['delta_gp_day']:+7.1f}  driver: {r['primary_driver']}")

    outdir = Path(args.output_dir); outdir.mkdir(parents=True, exist_ok=True)
    dives, charts = {}, {}
    for s in [x.strip() for x in args.deep_dive.split(",") if x.strip()]:
        wk = weekly_series(daily, s)
        if wk.empty:
            print(f"  (deep-dive {s}: no data)"); continue
        dives[s] = wk
        png = outdir / f"_lifecycle_{s.replace(' ', '_')}.png"
        try:
            deep_dive_chart(wk, s, eras, png); charts[s] = png
        except Exception as ex:
            print(f"  (chart {s} skipped: {ex})")

    out = outdir / f"state_lifecycle_study_{end}.xlsx"
    write_workbook(out, start, end, eras, cls, et, decomp, dives, charts)
    print(f"\n  Wrote: {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
