"""
tROAS Target Gradient — Google Home Insurance
═══════════════════════════════════════════════════════════════════════
Produces the target-change recommendation workbook for Ben: per campaign,
a ladder ("gradient") of candidate tROAS targets with the expected
delivered economics of each rung, anchored on post-freeze realized
performance. Ben's proposed flat -3pt reduction is one rung on the
ladder; the framework steps (-15%/-20% relative) and the match-realized
rung bracket it so the table shows what each campaign can actually
support.

FOUR-LAYER FRAMEWORK MAPPING (floors framework superseded by Ben):
  Layer 1 (floors)     — 20/12/5/0 margin floors dropped per Ben.
                         Replaced by a true-break-even line expressed
                         in DB terms: DB-measured ROAS of
                         1/(1+correction) ≈ 0.885-0.909 is the real
                         GP=0 line, because DB conversion_value
                         understates true economics 10-13%. The
                         correction is applied to the THRESHOLD, never
                         to the DB's realized numbers.
  Layer 2 (realized)   — post-freeze realized ROAS per campaign +
                         state decomposition of each gap.
  Layer 3 (simulator)  — NOT in the DB. Simulator column left blank in
                         the output; fill directionally from the Ads UI
                         after the first clean post-freeze week.
  Layer 4 (schedule)   — steps of 15-20%, >=1 week apart, batched
                         same-day; rollback trigger auto-drafted per
                         campaign in the table.

DATA RULES (per charter):
  - Targets: campaigns.target_roas, account_id=1, campaign_type='SEARCH',
    resolved to campaign_id once and id-based thereafter.
  - Realized: campaign_performance_daily_search, device='ALL' rows only.
    Revenue = native conversion_value (ruled basis at campaign grain,
    verified <0.1% vs routed truth). ROAS = SUM(conversion_value) /
    SUM(cost_micros/1e6). Summed numerators, never averaged ratios.
  - States: campaign_performance_geo_daily_search — NO device='ALL'
    rollup exists there; device rows are summed.
  - No keyword tables, no base (non-_search) tables.
  - DB opened read-only (mode=ro). Do not run during a sync.

WINDOWS:
  - Realized window: --window-start (default 2026-07-30, first full
    post-cutover day; GBRAIDS routing consistent account-wide from
    7/30) through MAX(date) minus --lag-days (default 3) to keep
    conversion-lag-soft days out of the ROAS read.
  - The trailing soft days are reported separately, labeled SOFT.
  - Anything computed before a full clean post-freeze week is
    provisional; the workbook stamps the window so that's auditable.

USAGE:
  python troas_target_gradient.py
  python troas_target_gradient.py --window-start 2026-07-30 --lag-days 4
  python troas_target_gradient.py --points-reduction 3 --steps 15,20
  python troas_target_gradient.py --correction-pct 11.5

OUTPUT: troas_target_gradient_<YYYY-MM-DD>.xlsx
  - Gradient:    the one-page ladder table, one block per campaign
  - State Decomp: per-campaign state ROAS / spend share / GP
                  contribution, sorted by GP drag
  - Notes:       window, correction factor, conventions, soft-window read
"""

import argparse
import sqlite3
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl")

DEFAULT_DB_PATH = Path(r"C:\Projects\google-ads-ops\db\paid_search.db")
ACCOUNT_ID = 1

# Name patterns used ONCE to resolve the campaign set to ids.
NAME_PATTERNS = [
    "%TIER1%", "%TIER2%", "%TIER3%", "%TIER4%",
    "%CALIFORNIA%", "%FLORIDA%", "%TEXAS%",
]

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
CAMP_FILL = PatternFill("solid", fgColor="D9E2F3")
BEN_FILL = PatternFill("solid", fgColor="FFF2CC")     # Ben's -3pt rung
NEG_FILL = PatternFill("solid", fgColor="F8CBAD")     # below true break-even
SOFT_FONT = Font(italic=True, color="808080", size=9)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


# ─────────────────────────────────────────────────────────────────────
# Data access
# ─────────────────────────────────────────────────────────────────────

def connect_ro(db_path: Path) -> sqlite3.Connection:
    uri = f"file:{db_path.as_posix()}?mode=ro"
    conn = sqlite3.connect(uri, uri=True)
    conn.row_factory = sqlite3.Row
    return conn


def table_columns(conn, table):
    return {r[1] for r in conn.execute(f"PRAGMA table_info({table})")}


def resolve_campaigns(conn):
    """Resolve the tier shells + state campaigns to ids, once, by name
    pattern; everything downstream is id-based. Returns list of dicts."""
    cols = table_columns(conn, "campaigns")
    status_col = "status" if "status" in cols else None
    where_patterns = " OR ".join("campaign_name LIKE ?" for _ in NAME_PATTERNS)
    sql = f"""
        SELECT campaign_id, campaign_name, target_roas
               {', ' + status_col if status_col else ''}
        FROM campaigns
        WHERE account_id = ?
          AND campaign_type = 'SEARCH'
          AND ({where_patterns})
    """
    rows = conn.execute(sql, [ACCOUNT_ID, *NAME_PATTERNS]).fetchall()
    out = []
    for r in rows:
        status = r[status_col] if status_col else None
        if status and str(status).upper() not in ("ENABLED", "PAUSED", "2", "3"):
            continue
        tr = r["target_roas"]
        # normalize: stored as ratio (1.15) or percent (115)
        if tr is not None and tr > 10:
            tr = tr / 100.0
        out.append({
            "campaign_id": r["campaign_id"],
            "name": r["campaign_name"],
            "target": tr,
            "status": status,
        })
    out.sort(key=lambda c: c["name"])
    return out


def max_data_date(conn):
    row = conn.execute(
        """SELECT MAX(p.date) FROM campaign_performance_daily_search p
           JOIN campaigns c ON c.campaign_id = p.campaign_id
           WHERE c.account_id = ?""", [ACCOUNT_ID]).fetchone()
    return datetime.strptime(row[0][:10], "%Y-%m-%d").date()


def campaign_perf(conn, ids, start, end):
    """Realized perf per campaign, device='ALL', summed numerators."""
    ph = ",".join("?" for _ in ids)
    sql = f"""
        SELECT campaign_id,
               SUM(cost_micros) / 1e6            AS cost,
               SUM(conversions)                  AS conv,
               SUM(conversion_value)             AS revenue
        FROM campaign_performance_daily_search
        WHERE campaign_id IN ({ph})
          AND device = 'ALL'
          AND date BETWEEN ? AND ?
        GROUP BY campaign_id
    """
    rows = conn.execute(sql, [*ids, start.isoformat(), end.isoformat()])
    return {r["campaign_id"]: dict(r) for r in rows}


def state_perf(conn, ids, start, end):
    """State decomposition. Geo table has NO device='ALL' rollup —
    device rows are summed (verified pattern, distinct from campaign
    table)."""
    ph = ",".join("?" for _ in ids)
    sql = f"""
        SELECT campaign_id, geo_state,
               SUM(cost_micros) / 1e6            AS cost,
               SUM(conversions)                  AS conv,
               SUM(conversion_value)             AS revenue
        FROM campaign_performance_geo_daily_search
        WHERE campaign_id IN ({ph})
          AND date BETWEEN ? AND ?
        GROUP BY campaign_id, geo_state
    """
    rows = conn.execute(sql, [*ids, start.isoformat(), end.isoformat()])
    out = {}
    for r in rows:
        out.setdefault(r["campaign_id"], []).append(dict(r))
    return out


# ─────────────────────────────────────────────────────────────────────
# Gradient math
# ─────────────────────────────────────────────────────────────────────

def safe_div(a, b):
    return a / b if b else None


def build_ladder(camp, perf, args, db_breakeven):
    """Rungs: Hold / Ben -Npt / -15% / -20% / Match realized.
    Expected delivered ROAS assumes the post-freeze delivery ratio
    (realized / target) persists at the new target — a directional
    first-order read, NOT a forecast; the simulator column exists for
    the curvature judgment."""
    cost = perf.get("cost", 0) or 0
    revenue = perf.get("revenue", 0) or 0
    realized = safe_div(revenue, cost)
    target = camp["target"]
    ratio = safe_div(realized, target) if (realized and target) else None
    gp = revenue - cost

    rungs = []

    def add(label, new_target):
        if new_target is None:
            rungs.append({"label": label, "target": None})
            return
        exp = ratio * new_target if ratio else None
        rungs.append({
            "label": label,
            "target": new_target,
            "delta_pct": safe_div(new_target - target, target) if target else None,
            "expected_roas": exp,
            "below_breakeven": (exp is not None and exp < db_breakeven),
        })

    add("Hold (current)", target)
    if target is not None:
        add(f"Ben: -{args.points_reduction:g}pts",
            target - args.points_reduction / 100.0)
        for s in args.steps:
            add(f"Step -{s:g}% rel", target * (1 - s / 100.0))
    add("Match realized", realized)

    return {
        "cost": cost, "revenue": revenue, "conv": perf.get("conv"),
        "gp": gp, "realized": realized, "ratio": ratio, "rungs": rungs,
    }


def draft_rollback(realized, db_breakeven):
    if realized is None:
        return "n/a — no post-freeze spend"
    floor = max(db_breakeven, 0.90 * realized)
    return (f"Revert prior target if rolling-7d realized ROAS < "
            f"{floor:.2f} (DB terms) for 5+ consecutive days, or CPL "
            f"rises >25% vs post-freeze baseline")


# ─────────────────────────────────────────────────────────────────────
# Workbook
# ─────────────────────────────────────────────────────────────────────

def style_header(ws, row, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center")


def write_gradient_sheet(ws, campaigns, ladders, args, db_breakeven,
                         win_start, win_end):
    ws.title = "Gradient"
    ws.sheet_view.showGridLines = False
    headers = ["Campaign", "Rung", "Target", "Δ target", "Post-freeze "
               "realized", "Delivery ratio", "Expected delivered ROAS",
               "vs true break-even", "Spend (win)", "GP (win, DB $)",
               "Simulator read", "Evidence tier", "Rollback trigger"]
    ws.append([f"tROAS Target Gradient — Google Home — window "
               f"{win_start} → {win_end} (trailing {args.lag_days}d "
               f"excluded as lag-soft) — PROVISIONAL until a full clean "
               f"post-freeze week"])
    ws.cell(row=1, column=1).font = Font(bold=True, size=11)
    ws.append([f"True break-even in DB terms = 1/(1+{args.correction_pct:g}%) "
               f"= {db_breakeven:.3f} DB ROAS. DB realized numbers are "
               f"never inflated; the threshold carries the correction."])
    ws.cell(row=2, column=1).font = SOFT_FONT
    ws.append([])
    ws.append(headers)
    style_header(ws, 4, len(headers))

    r = 5
    for camp in campaigns:
        lad = ladders[camp["campaign_id"]]
        first = True
        for rung in lad["rungs"]:
            if rung["target"] is None:
                continue
            row = [
                camp["name"] if first else "",
                rung["label"],
                rung["target"],
                rung.get("delta_pct"),
                lad["realized"] if first else None,
                lad["ratio"] if first else None,
                rung.get("expected_roas"),
                ("BELOW" if rung.get("below_breakeven") else "above"),
                lad["cost"] if first else None,
                lad["gp"] if first else None,
                "",  # simulator — fill from UI after first clean week
                ("realized" if lad["realized"] else "policy"),
                draft_rollback(lad["realized"], db_breakeven) if first else "",
            ]
            ws.append(row)
            for c in range(1, len(headers) + 1):
                cell = ws.cell(row=r, column=c)
                cell.border = BORDER
                if c in (3, 5, 6, 7):
                    cell.number_format = "0.00"
                elif c == 4:
                    cell.number_format = "+0.0%;-0.0%"
                elif c in (9, 10):
                    cell.number_format = "#,##0"
            if first:
                ws.cell(row=r, column=1).fill = CAMP_FILL
                ws.cell(row=r, column=1).font = Font(bold=True, size=9)
            if rung["label"].startswith("Ben"):
                for c in range(2, 9):
                    ws.cell(row=r, column=c).fill = BEN_FILL
            if rung.get("below_breakeven"):
                ws.cell(row=r, column=7).fill = NEG_FILL
                ws.cell(row=r, column=8).fill = NEG_FILL
            first = False
            r += 1
        r += 1
        ws.append([])

    widths = [42, 16, 9, 9, 12, 11, 13, 12, 11, 11, 14, 10, 52]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_state_sheet(ws, campaigns, states, ladders):
    ws.title = "State Decomp"
    ws.sheet_view.showGridLines = False
    headers = ["Campaign", "State", "Spend", "Spend share", "Conv",
               "Revenue (DB $)", "Realized ROAS", "GP (DB $)",
               "GP contribution to gap"]
    ws.append(headers)
    style_header(ws, 1, len(headers))
    r = 2
    for camp in campaigns:
        rows = states.get(camp["campaign_id"], [])
        total_cost = sum(x["cost"] or 0 for x in rows)
        rows.sort(key=lambda x: (x["revenue"] or 0) - (x["cost"] or 0))
        first = True
        for x in rows:
            cost = x["cost"] or 0
            rev = x["revenue"] or 0
            gp = rev - cost
            ws.append([
                camp["name"] if first else "",
                x["geo_state"], cost,
                safe_div(cost, total_cost),
                x["conv"], rev, safe_div(rev, cost), gp, gp,
            ])
            for c in range(1, len(headers) + 1):
                cell = ws.cell(row=r, column=c)
                cell.border = BORDER
                if c in (3, 5, 6, 8, 9):
                    cell.number_format = "#,##0"
                elif c == 4:
                    cell.number_format = "0.0%"
                elif c == 7:
                    cell.number_format = "0.00"
            if first:
                ws.cell(row=r, column=1).fill = CAMP_FILL
                ws.cell(row=r, column=1).font = Font(bold=True, size=9)
            if gp < 0:
                ws.cell(row=r, column=8).fill = NEG_FILL
            first = False
            r += 1
        r += 1
        ws.append([])
    widths = [42, 8, 10, 11, 8, 12, 12, 10, 18]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_notes_sheet(ws, args, win_start, win_end, soft_read,
                      db_breakeven, camp_count):
    ws.title = "Notes"
    lines = [
        ("Run", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Realized window", f"{win_start} → {win_end}"),
        ("Lag exclusion", f"trailing {args.lag_days} days excluded "
                          "(conversion lag restates them upward)"),
        ("Provisional flag", "numbers provisional until a full clean "
                             "Mon–Sun post-freeze week exists"),
        ("Revenue basis", "campaign_performance_daily_search."
                          "conversion_value, device='ALL' (ruled basis, "
                          "verified <0.1% vs routed truth)"),
        ("State basis", "campaign_performance_geo_daily_search, device "
                        "rows summed (no ALL rollup in that table)"),
        ("Ratio convention", "summed numerators only; never averaged "
                             "daily ratios"),
        ("Correction", f"{args.correction_pct:g}% value understatement "
                       f"applied to the break-even THRESHOLD "
                       f"(1/(1+c) = {db_breakeven:.3f} DB ROAS); DB "
                       f"realized never inflated"),
        ("Floors", "20/12/5/0 margin-floor framework superseded per Ben "
                   "(8/12); gradient anchored on realized performance"),
        ("Expected ROAS", "directional first-order read: post-freeze "
                          "delivery ratio × candidate target; NOT a "
                          "forecast — simulator curvature (Layer 3) "
                          "adjudicates rung choice after first clean "
                          "week"),
        ("Schedule (Layer 4)", "one rung per move, ≥1 week apart, all "
                               "changes batched same-day; TARGET_ROAS "
                               "events land in change-history "
                               "automatically for post-change anchoring"),
        ("Campaigns resolved", str(camp_count)),
        ("Soft trailing window", soft_read),
    ]
    for k, v in lines:
        ws.append([k, v])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=9)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 110
    for row in ws.iter_rows():
        row[1].alignment = Alignment(wrap_text=True, vertical="top")


# ─────────────────────────────────────────────────────────────────────

def main():
    p = argparse.ArgumentParser()
    p.add_argument("--db-path", default=str(DEFAULT_DB_PATH))
    p.add_argument("--window-start", default="2026-07-30",
                   help="first day of the post-freeze realized window")
    p.add_argument("--lag-days", type=int, default=3,
                   help="trailing days excluded as conversion-lag-soft")
    p.add_argument("--points-reduction", type=float, default=3.0,
                   help="Ben's rung: percentage points off current target")
    p.add_argument("--steps", default="15,20",
                   help="relative-step rungs, comma-separated percents")
    p.add_argument("--correction-pct", type=float, default=11.5,
                   help="value-understatement correction (10-13 band)")
    p.add_argument("--output", default=None)
    args = p.parse_args()
    args.steps = [float(s) for s in str(args.steps).split(",") if s.strip()]

    db_path = Path(args.db_path)
    if not db_path.exists():
        sys.exit(f"DB not found: {db_path}")
    conn = connect_ro(db_path)

    db_breakeven = 1.0 / (1.0 + args.correction_pct / 100.0)

    campaigns = resolve_campaigns(conn)
    if not campaigns:
        sys.exit("No campaigns resolved — check NAME_PATTERNS / schema.")
    print(f"Resolved {len(campaigns)} campaigns:")
    for c in campaigns:
        t = f"{c['target']:.2f}" if c["target"] else "—"
        print(f"  {c['campaign_id']}  target={t}  {c['name']}")

    max_date = max_data_date(conn)
    win_start = datetime.strptime(args.window_start, "%Y-%m-%d").date()
    win_end = max_date - timedelta(days=args.lag_days)
    if win_end <= win_start:
        sys.exit(f"Window empty: {win_start} → {win_end} (max date "
                 f"{max_date}, lag {args.lag_days}d)")
    print(f"\nRealized window: {win_start} → {win_end} "
          f"(max data date {max_date}, trailing {args.lag_days}d soft)")

    ids = [c["campaign_id"] for c in campaigns]
    perf = campaign_perf(conn, ids, win_start, win_end)
    states = state_perf(conn, ids, win_start, win_end)

    # soft trailing read, reported but never mixed into the ladder
    soft = campaign_perf(conn, ids, win_end + timedelta(days=1), max_date)
    soft_cost = sum(v["cost"] or 0 for v in soft.values())
    soft_rev = sum(v["revenue"] or 0 for v in soft.values())
    soft_read = (f"{win_end + timedelta(days=1)} → {max_date}: "
                 f"spend {soft_cost:,.0f}, revenue {soft_rev:,.0f}, "
                 f"ROAS {safe_div(soft_rev, soft_cost) or 0:.2f} "
                 f"(SOFT — lag will restate upward)")

    ladders = {}
    print("\nGradient (DB terms; break-even line "
          f"{db_breakeven:.3f}):")
    for camp in campaigns:
        lad = build_ladder(camp, perf.get(camp["campaign_id"], {}),
                           args, db_breakeven)
        ladders[camp["campaign_id"]] = lad
        rz = f"{lad['realized']:.2f}" if lad["realized"] else "—"
        rt = f"{lad['ratio']:.2f}" if lad["ratio"] else "—"
        print(f"  {camp['name'][:44]:44s} realized {rz}  "
              f"delivery-ratio {rt}  GP {lad['gp']:,.0f}")

    wb = Workbook()
    write_gradient_sheet(wb.active, campaigns, ladders, args,
                         db_breakeven, win_start, win_end)
    write_state_sheet(wb.create_sheet(), campaigns, states, ladders)
    write_notes_sheet(wb.create_sheet(), args, win_start, win_end,
                      soft_read, db_breakeven, len(campaigns))

    out = args.output or f"troas_target_gradient_{date.today().isoformat()}.xlsx"
    wb.save(out)
    conn.close()
    print(f"\nSaved: {out}")


if __name__ == "__main__":
    main()
