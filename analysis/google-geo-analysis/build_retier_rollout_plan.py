#!/usr/bin/env python3
"""
Google Home re-tier rollout plan builder.

Reads the 'State Assignments' sheet of the geo-model workbook (the single source of truth Gilbert
maintains) and emits an exec-ready rollout-plan workbook. Re-run any time assignments or dates change.

What it answers (the exec asks):
  - what states are moving           -> Move Schedule
  - how much spend is disrupted       -> Rollout Summary (moving spend, % of total)
  - what's performing well            -> New Tier Profile (GP/lead, RPL, CPL, ROAS per new tier)
  - what tiers stack to current       -> Tier Transition Matrix (current x new)
  - timeline                          -> Rollout Summary + Move Schedule (by 'Change <date>' column)
  - execution clarity                 -> Shell Repurposing Map (which physical shell becomes which tier)

Expected columns on 'State Assignments' (header row auto-detected):
  state | current_tier | C2_weighted_kmeans (proposed tier) | clicks conversions spend rpl cpc cpl
  gp_per_lead roas | ... | Campaign Change? (0=No,1=Yes) | Destination campaign | Change <date> ...

Usage:
  python build_retier_rollout_plan.py --input google_geo_model_gp_per_lead_2026-06-09.xlsx
"""
import argparse, datetime as dt, re
import pandas as pd, numpy as np
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule

A=lambda **k: Font(name="Arial", **k)
TITLE=A(bold=True,size=14,color="1F4E79"); SUB=A(bold=True,size=11,color="1F4E79")
NOTE=A(italic=True,size=9,color="595959"); BLACK=A(); BOLD=A(bold=True); BLUE=A(color="0000FF")
HF=A(bold=True,color="FFFFFF"); HFILL=PatternFill("solid",fgColor="1F4E79")
SUBFILL=PatternFill("solid",fgColor="DDEBF7"); GREY=PatternFill("solid",fgColor="F2F2F2")
GOOD=PatternFill("solid",fgColor="C6EFCE"); REDF=PatternFill("solid",fgColor="FFC7CE"); AMBER=PatternFill("solid",fgColor="FCE4D6")
TH=Side(style="thin",color="BFBFBF"); BORD=Border(top=TH,bottom=TH,left=TH,right=TH)
WRAP=Alignment(wrap_text=True,vertical="top"); CTR=Alignment(horizontal="center",vertical="center",wrap_text=True)
M0='$#,##0'; M1='$#,##0.0'; GPF='$#,##0.0;($#,##0.0)'; R2='0.00'; P0='0%'
TIER_ORDER=['TIER ONE','TIER TWO','TIER THREE','TIER FOUR']
CUR_ORDER=['Tier One','Tier Two','Tier Three','Tier Four']


def load(path):
    wb=load_workbook(path, data_only=True); ws=wb['State Assignments']
    hr=next(r for r in range(1,8) if 'state' in [str(ws.cell(row=r,column=c).value).lower()
            for c in range(1,ws.max_column+1)])
    cols=[ws.cell(row=hr,column=c).value for c in range(1,ws.max_column+1)]
    recs=[]
    for r in range(hr+1,ws.max_row+1):
        v=[ws.cell(row=r,column=c).value for c in range(1,ws.max_column+1)]
        if not v[0] or str(v[0]).strip().lower().startswith("total"): continue
        recs.append(dict(zip(cols,v)))
    df=pd.DataFrame(recs)
    df.rename(columns={'C2_weighted_kmeans':'new_tier','Campaign Change? (0=No, 1=Yes)':'chg',
                       'Destination campaign':'dest'},inplace=True)
    for c in ['clicks','conversions','spend','rpl','cpc','cpl','gp_per_lead','roas']:
        df[c]=pd.to_numeric(df[c],errors='coerce')
    df['chg']=pd.to_numeric(df['chg'],errors='coerce').fillna(0).astype(int)
    date_cols=[c for c in df.columns if str(c).startswith('Change ')]
    def wk(row):
        for c in date_cols:
            if str(row[c]).strip().lower()=='yes': return c.replace('Change ','').strip()
        return 'unscheduled'
    df['week']=df.apply(wk,axis=1)
    # order weeks chronologically
    def keyf(w):
        try: return dt.datetime.strptime(w+'/2026','%m/%d/%Y')
        except Exception: return dt.datetime.max
    df['_wk']=df['week'].map(keyf)
    return df


def wavg(g,val,wt): 
    s=g[wt].sum(); return (g[val]*g[wt]).sum()/s if s else np.nan

def style_header(ws,row,headers,widths):
    for j,h in enumerate(headers,1):
        c=ws.cell(row=row,column=j,value=h); c.font=HF; c.fill=HFILL; c.border=BORD; c.alignment=CTR
    for col,w in zip("ABCDEFGHIJKLMN",widths): ws.column_dimensions[col].width=w

def posture(gpl):
    if gpl>=25: return "loosen tROAS / fund growth"
    if gpl>=10: return "hold, modest scale"
    if gpl>=0:  return "hold / watch"
    return "tighten / suppress"


def build(df, out):
    wb=Workbook(); wb.remove(wb.active)
    tot=df['spend'].sum(); mv=df[df['chg']==1]; st=df[df['chg']==0]
    weeks=df.sort_values('_wk')['week'].dropna().unique().tolist()

    # ---------- 1. Rollout Summary ----------
    ws=wb.create_sheet("Rollout Summary")
    ws["A1"]="Google Home — geo re-tier rollout plan"; ws["A1"].font=TITLE
    ws["A2"]=(f"Re-cluster states by GP/lead and consolidate into 4 margin-coherent tiers, phased best-tier-first. "
              f"California, Texas, Florida remain standalone and are untouched. Spend figures are the 12-week model window "
              f"(2026-03-18..06-09). Generated {dt.date.today():%b %d, %Y}.")
    ws["A2"].font=NOTE; ws["A2"].alignment=WRAP; ws.merge_cells("A2:H2"); ws.row_dimensions[2].height=42
    r=4
    ws.cell(row=r,column=1,value="At a glance").font=SUB; r+=1
    kpis=[("States re-tiering (ex-standalones)",f"{len(df)}"),
          ("States physically changing campaigns",f"{len(mv)}  ({len(mv)/len(df)*100:.0f}% of states)"),
          ("Spend in motion (disruption)",f"${mv['spend'].sum():,.0f}  ({mv['spend'].sum()/tot*100:.0f}% of ${tot:,.0f})"),
          ("States staying in place",f"{len(st)}  (${st['spend'].sum():,.0f}, {st['spend'].sum()/tot*100:.0f}%)"),
          ("Rollout window",f"{weeks[0]} → {weeks[-1]} (one tier per week)")]
    for k,v in kpis:
        ws.cell(row=r,column=1,value=k).font=BOLD
        c=ws.cell(row=r,column=3,value=v); c.font=BLACK; ws.merge_cells(start_row=r,start_column=3,end_row=r,end_column=6); r+=1
    r+=1
    ws.cell(row=r,column=1,value="Phased schedule (best-tier first)").font=SUB; r+=1
    style_header(ws,r,["week","new tier","destination shell","states","moving","moving spend","wtd GP/lead","wtd ROAS","focus"],
                 [10,12,16,8,8,13,12,10,30]); r+=1
    for w in weeks:
        g=df[df['week']==w]; gm=g[g['chg']==1]
        new_t=g['new_tier'].mode().iat[0] if len(g) else '?'
        dest=g['dest'].mode().iat[0] if g['dest'].notna().any() else '?'
        vals=[w,new_t,dest,len(g),len(gm),round(gm['spend'].sum()),round(wavg(g,'gp_per_lead','conversions'),1),
              round(wavg(g,'roas','spend'),2),posture(wavg(g,'gp_per_lead','conversions'))]
        for j,v in enumerate(vals,1):
            c=ws.cell(row=r,column=j,value=v); c.border=BORD; c.font=BLACK
            if j==6: c.number_format=M0
            if j==7: c.number_format=GPF
            if j==8: c.number_format=R2
            if j==9: c.font=NOTE
        r+=1
    gp_col=f"G{r-len(weeks)}:G{r-1}"; roas_col=f"H{r-len(weeks)}:H{r-1}"
    ws.conditional_formatting.add(gp_col,CellIsRule(operator='lessThan',formula=['0'],fill=REDF))
    ws.conditional_formatting.add(gp_col,CellIsRule(operator='greaterThanOrEqual',formula=['0'],fill=GOOD))
    ws.freeze_panes="A4"

    # ---------- 2. Move Schedule ----------
    ws=wb.create_sheet("Move Schedule")
    ws["A1"]="Move schedule — every state, by rollout week"; ws["A1"].font=TITLE
    ws["A2"]="from shell = current campaign; to shell = destination campaign (shell is repurposed, then relabeled). 'moving' = physically changes campaign."
    ws["A2"].font=NOTE; ws["A2"].alignment=WRAP; ws.merge_cells("A2:K2"); ws.row_dimensions[2].height=28
    hdr=["week","state","current tier","new tier","from shell","to shell","moving?","spend","conv","GP/lead","ROAS"]
    style_header(ws,4,hdr,[11,16,12,11,12,12,8,11,8,11,8]); r=5
    for w in weeks:
        g=df[df['week']==w].sort_values('spend',ascending=False)
        new_t=g['new_tier'].mode().iat[0]
        for _,x in g.iterrows():
            row=[w,x['state'],x['current_tier'],x['new_tier'],x['current_tier'],x['dest'],
                 "yes" if x['chg']==1 else "—",round(x['spend']),round(x['conversions']),
                 round(x['gp_per_lead'],1),round(x['roas'],2)]
            for j,v in enumerate(row,1):
                c=ws.cell(row=r,column=j,value=v); c.border=BORD; c.font=BLACK
                if j==8: c.number_format=M0
                if j==10: c.number_format=GPF
                if j==11: c.number_format=R2
                if j==7 and x['chg']==1: c.fill=AMBER
            r+=1
        # weekly subtotal
        gm=g[g['chg']==1]
        sub=["",f"{w} subtotal",new_t,"",f"{len(gm)} moving","","",round(g['spend'].sum()),round(g['conversions'].sum()),
             round(wavg(g,'gp_per_lead','conversions'),1),round(wavg(g,'roas','spend'),2)]
        for j,v in enumerate(sub,1):
            c=ws.cell(row=r,column=j,value=v); c.font=BOLD; c.fill=SUBFILL; c.border=BORD
            if j==8: c.number_format=M0
            if j==10: c.number_format=GPF
            if j==11: c.number_format=R2
        r+=1
    ws.conditional_formatting.add(f"J5:J{r-1}",CellIsRule(operator='lessThan',formula=['0'],fill=REDF))
    ws.freeze_panes="A5"

    # ---------- 3. Tier Transition Matrix ----------
    ws=wb.create_sheet("Tier Transition Matrix")
    ws["A1"]="Tier transition — current tier (rows) vs new tier (columns)"; ws["A1"].font=TITLE
    ws["A2"]="The case for the re-tier in one table: where each current tier's spend actually lands by margin. A diagonal would mean the old tiers were already margin-sorted; they are not."
    ws["A2"].font=NOTE; ws["A2"].alignment=WRAP; ws.merge_cells("A2:G2"); ws.row_dimensions[2].height=28
    def matrix(at_row,aggfunc,fmt,label):
        ws.cell(row=at_row,column=1,value=label).font=SUB
        piv=pd.pivot_table(df,values='spend' if aggfunc=='sum' else 'state',index='current_tier',
                           columns='new_tier',aggfunc=aggfunc,fill_value=0)
        piv=piv.reindex([c for c in CUR_ORDER if c in piv.index])
        piv=piv[[c for c in TIER_ORDER if c in piv.columns]]
        hrow=at_row+1
        ws.cell(row=hrow,column=1,value="current ↓ / new →").font=HF; ws.cell(row=hrow,column=1).fill=HFILL; ws.cell(row=hrow,column=1).border=BORD
        for j,t in enumerate(piv.columns,2):
            c=ws.cell(row=hrow,column=j,value=t); c.font=HF; c.fill=HFILL; c.border=BORD; c.alignment=CTR
        c=ws.cell(row=hrow,column=len(piv.columns)+2,value="total"); c.font=HF; c.fill=HFILL; c.border=BORD
        rr=hrow+1
        for idx,prow in piv.iterrows():
            ws.cell(row=rr,column=1,value=idx).font=BOLD; ws.cell(row=rr,column=1).fill=SUBFILL; ws.cell(row=rr,column=1).border=BORD
            for j,t in enumerate(piv.columns,2):
                c=ws.cell(row=rr,column=j,value=round(prow[t]) if aggfunc=='sum' else int(prow[t])); c.number_format=fmt; c.border=BORD
            c=ws.cell(row=rr,column=len(piv.columns)+2,value=round(prow.sum()) if aggfunc=='sum' else int(prow.sum())); c.number_format=fmt; c.border=BORD; c.font=BOLD
            rr+=1
        # column totals
        ws.cell(row=rr,column=1,value="total").font=BOLD; ws.cell(row=rr,column=1).fill=SUBFILL; ws.cell(row=rr,column=1).border=BORD
        for j,t in enumerate(piv.columns,2):
            c=ws.cell(row=rr,column=j,value=round(piv[t].sum()) if aggfunc=='sum' else int(piv[t].sum())); c.number_format=fmt; c.border=BORD; c.font=BOLD
        gt=ws.cell(row=rr,column=len(piv.columns)+2,value=round(piv.values.sum()) if aggfunc=='sum' else int(piv.values.sum())); gt.font=BOLD; gt.border=BORD; gt.number_format=fmt
        if aggfunc=='sum':
            ws.conditional_formatting.add(f"B{hrow+1}:E{rr-1}",
                ColorScaleRule(start_type='num',start_value=0,start_color='FFFFFF',end_type='max',end_color='9BC2E6'))
        return rr
    last=matrix(4,'sum',M0,"Spend ($)")
    last=matrix(last+3,'count','0',"State count")
    for col,w in zip("ABCDEF",[18,12,12,12,12,10]): ws.column_dimensions[col].width=w
    r=last+3
    ws.cell(row=r,column=1,value="Read-outs").font=SUB; r+=1
    one_to_one=df[(df['current_tier']=='Tier One')&(df['new_tier']=='TIER ONE')]['spend'].sum()
    cur1=df[df['current_tier']=='Tier One']; cur1_worst=cur1[cur1['new_tier']=='TIER FOUR']['spend'].sum()
    new1=df[df['new_tier']=='TIER ONE']; new1_fromlow=new1[new1['current_tier'].isin(['Tier Three','Tier Four'])]['spend'].sum()
    for t in [f"• Current Tier One puts ${one_to_one:,.0f} into the new best tier — it scatters into Tiers Two/Three/Four, with ${cur1_worst:,.0f} landing in the negative-margin tier.",
              f"• The new TIER ONE (best GP/lead) is ${new1_fromlow:,.0f}-of-${new1['spend'].sum():,.0f} drawn from current Tiers Three/Four — the most profitable states were under-bid in the lowest tiers.",
              "• Off-diagonal mass = how margin-incoherent the legacy tiering was. That mass is the re-tier opportunity."]:
        c=ws.cell(row=r,column=1,value=t); c.font=NOTE; c.alignment=WRAP; ws.merge_cells(f"A{r}:F{r}"); ws.row_dimensions[r].height=26; r+=1

    # ---------- 4. Shell Repurposing Map ----------
    ws=wb.create_sheet("Shell Repurposing Map")
    ws["A1"]="Shell repurposing — what each existing campaign becomes"; ws["A1"].font=TITLE
    ws["A2"]="Shells are reused to preserve history; names are mutable. Relabel each shell to its new tier immediately after its move. campaign_id is the stable key — name-parsing logic must update."
    ws["A2"].font=NOTE; ws["A2"].alignment=WRAP; ws.merge_cells("A2:H2"); ws.row_dimensions[2].height=28
    style_header(ws,4,["physical shell (current name)","becomes (new tier)","suggested new label","states staying","states leaving","states entering","resulting spend","resulting GP/lead"],
                 [22,14,20,11,11,11,12,13]); r=5
    shells=[s for s in CUR_ORDER if s in set(df['dest'].dropna())]
    for shell in shells:
        dest_states=df[df['dest']==shell]                       # states that will live here
        new_t=dest_states['new_tier'].mode().iat[0]
        staying=df[(df['current_tier']==shell)&(df['dest']==shell)]
        leaving=df[(df['current_tier']==shell)&(df['dest']!=shell)]
        entering=df[(df['current_tier']!=shell)&(df['dest']==shell)]
        row=[shell,new_t,f"{new_t} (GP/lead ${wavg(dest_states,'gp_per_lead','conversions'):.0f})",
             len(staying),len(leaving),len(entering),round(dest_states['spend'].sum()),
             round(wavg(dest_states,'gp_per_lead','conversions'),1)]
        for j,v in enumerate(row,1):
            c=ws.cell(row=r,column=j,value=v); c.border=BORD; c.font=BLACK
            if j==7: c.number_format=M0
            if j==8: c.number_format=GPF
        r+=1
    ws.conditional_formatting.add(f"H5:H{r-1}",CellIsRule(operator='lessThan',formula=['0'],fill=REDF))
    ws.conditional_formatting.add(f"H5:H{r-1}",CellIsRule(operator='greaterThanOrEqual',formula=['0'],fill=GOOD))

    # ---------- 5. New Tier Profile ----------
    ws=wb.create_sheet("New Tier Profile")
    ws["A1"]="New tier profile — what's performing well"; ws["A1"].font=TITLE
    ws["A2"]="Per consolidated tier on the 12-week window. Monotonic GP/lead gradient = margin-coherent tiers (the goal). Suggested posture is a starting bid stance, not a final target."
    ws["A2"].font=NOTE; ws["A2"].alignment=WRAP; ws.merge_cells("A2:I2"); ws.row_dimensions[2].height=28
    style_header(ws,4,["new tier","states","spend","conv","RPL","CPL","GP/lead","ROAS","suggested posture"],
                 [12,8,12,8,8,8,11,8,28]); r=5
    for t in TIER_ORDER:
        g=df[df['new_tier']==t]
        if not len(g): continue
        gpl=wavg(g,'gp_per_lead','conversions')
        row=[t,len(g),round(g['spend'].sum()),round(g['conversions'].sum()),round(wavg(g,'rpl','conversions'),1),
             round(wavg(g,'cpl','conversions'),1),round(gpl,1),round(wavg(g,'roas','spend'),2),posture(gpl)]
        for j,v in enumerate(row,1):
            c=ws.cell(row=r,column=j,value=v); c.border=BORD; c.font=BLACK
            if j==3: c.number_format=M0
            if j in (5,6,7): c.number_format=M1 if j<7 else GPF
            if j==8: c.number_format=R2
            if j==9: c.font=NOTE
        r+=1
    ws.conditional_formatting.add(f"G5:G{r-1}",CellIsRule(operator='lessThan',formula=['0'],fill=REDF))
    ws.conditional_formatting.add(f"G5:G{r-1}",CellIsRule(operator='greaterThanOrEqual',formula=['0'],fill=GOOD))

    # ---------- 6. GP Projection (scenario) ----------
    ws=wb.create_sheet("GP Projection")
    ws["A1"]="Forward GP lift — re-bidding each cohort to its margin (scenario)"; ws["A1"].font=TITLE
    ws["A2"]=("Planning estimate, not a forecast. Each tier is re-bid toward its margin: loosen the profitable tiers to "
              "capture volume (at some margin compression as CPL rises), suppress the negative-margin tier. The TIER FOUR "
              "piece is the most reliable (it stops measured negative-margin spend); the TIER ONE/TWO growth piece is "
              "rank-bounded — treat it as a ceiling and confirm with post-move marginal ROAS. 12wk window = 84 days. "
              "Blue cells are editable Base assumptions; projections recompute. Excludes CA/TX/FL standalones.")
    ws["A2"].font=NOTE; ws["A2"].alignment=WRAP; ws.merge_cells("A2:K2"); ws.row_dimensions[2].height=58
    # unified model: proj_GP = conv*(1+volΔ)*(GP/lead*retention).  Suppression = negative volΔ.
    SCN={  # (volΔ, retention) per scenario
      'TIER ONE':  {'cons':(0.15,0.90),'base':(0.30,0.85),'stretch':(0.50,0.80),'act':'loosen → capture profitable volume'},
      'TIER TWO':  {'cons':(0.05,0.95),'base':(0.15,0.90),'stretch':(0.25,0.88),'act':'modest scale'},
      'TIER THREE':{'cons':(0.00,1.00),'base':(0.00,1.00),'stretch':(0.05,0.95),'act':'hold / watch'},
      'TIER FOUR': {'cons':(-0.40,1.00),'base':(-0.60,1.00),'stretch':(-0.80,1.00),'act':'suppress → stop negative-margin spend'},
    }
    hdr=["new tier","action","curr conv","curr GP/lead","curr GP /12wk","Base volΔ%","Base GP/lead ret%",
         "Base proj GP /12wk","Base Δ/day","Cons Δ/day","Stretch Δ/day"]
    style_header(ws,4,hdr,[12,30,9,11,12,10,13,14,11,11,12]); r=5
    DAYS=84.0; first=r
    for t in TIER_ORDER:
        g=df[df['new_tier']==t]
        if not len(g): continue
        conv=g['conversions'].sum(); gpl=wavg(g,'gp_per_lead','conversions'); cur=conv*gpl
        bv,br=SCN[t]['base']; cv,cr=SCN[t]['cons']; sv,sr=SCN[t]['stretch']
        cons_d=(conv*(1+cv)*(gpl*cr)-cur)/DAYS; str_d=(conv*(1+sv)*(gpl*sr)-cur)/DAYS
        cells=[t,SCN[t]['act'],round(conv),round(gpl,1),f"=C{r}*D{r}",bv,br,
               f"=C{r}*(1+F{r})*(D{r}*G{r})",f"=(H{r}-E{r})/{DAYS}",round(cons_d,1),round(str_d,1)]
        for j,v in enumerate(cells,1):
            c=ws.cell(row=r,column=j,value=v); c.border=BORD; c.font=BLACK
            if j==2: c.font=NOTE
            if j in (5,8): c.number_format=M0
            if j in (9,10,11): c.number_format=GPF
            if j in (6,7): c.number_format=P0; c.font=BLUE; c.fill=PatternFill("solid",fgColor="FFF2CC")  # editable
        r+=1
    # totals
    ws.cell(row=r,column=2,value="TOTAL GP lift / day").font=BOLD
    for col,letter in [(9,'I'),(10,'J'),(11,'K')]:
        c=ws.cell(row=r,column=col,value=f"=SUM({letter}{first}:{letter}{r-1})"); c.number_format=GPF; c.font=BOLD; c.fill=SUBFILL
    tr=r; r+=1
    ws.cell(row=r,column=2,value="annualized (×365)").font=BOLD
    for col,letter in [(9,'I'),(10,'J'),(11,'K')]:
        c=ws.cell(row=r,column=col,value=f"={letter}{tr}*365"); c.number_format=M0; c.font=BOLD
    r+=2
    ws.cell(row=r,column=1,value="Read-outs").font=SUB; r+=1
    t4=df[df['new_tier']=='TIER FOUR']; t4gp=t4['conversions'].sum()*wavg(t4,'gp_per_lead','conversions')
    for tx in [f"• TIER FOUR is the grounded floor: it runs at ${t4gp/DAYS:,.0f}/day of negative GP today; suppression recovers most of that regardless of the growth assumptions — this is the same lever as the CA/Tier-Four tighten (~$943/wk with CA).",
               "• TIER ONE/TWO upside hinges on volume beating margin compression: at Base (+30% vol / 85% margin on TIER ONE) the growth tiers add a modest lift; if loosening can't capture volume because they're rank-limited, that piece trends to zero — so the conservative column is the safer plan number.",
               "• The durable win isn't this 12-week lift — it's margin-coherent tiers that let every future bid/budget move land on a clean cohort. This sheet sizes the near-term; the structure is the compounding part."]:
        c=ws.cell(row=r,column=1,value=tx); c.font=NOTE; c.alignment=WRAP; ws.merge_cells(f"A{r}:K{r}"); ws.row_dimensions[r].height=40; r+=1

    # ---------- 7. Risks & Sequencing ----------
    ws=wb.create_sheet("Risks & Sequencing")
    ws["A1"]="Risks, sequencing & open items"; ws["A1"].font=TITLE
    items=[
      ("Bidder learning reset","Every shell gains or loses states, so each resets tROAS/tCPA learning (~1–2 weeks). The one-tier-per-week phasing staggers the resets rather than resetting all four at once. Expect noisy tier metrics during the transition; don't judge a tier until ~2 weeks after its last state lands (full clean read ~late July)."),
      ("Relabel after move + id stability","Relabel each shell to its new tier the same day its move completes, so reporting isn't misread (the 'Tier Two' shell becomes the best tier). campaign_id is the stable key — id-based analysis survives the rename; any name-parsing logic must be updated."),
      ("Snapshot before 6/16","Capture the current state→campaign mapping before the first move (mirror the Microsoft pre_restructure snapshot) so post-move performance can be attributed cleanly."),
      ("Budget & target reset by archetype","Don't carry over each shell's old budget/target. Set the new TIER ONE (GP/lead ~$30, ROAS 1.38) looser / funded to capture profitable volume; tighten or suppress the new TIER FOUR (negative GP/lead). This is where the GP is recovered."),
      ("Standalones untouched","California, Texas, Florida stay as standalone campaigns and are excluded from all moves."),
      ("Below-threshold states — PENDING","Delaware (GP/lead $37, nearest TIER ONE), West Virginia (TIER THREE), New Hampshire / South Dakota / Wyoming (TIER FOUR) are under the 200-click eligibility floor and have no destination/date yet. Decide whether to fold Delaware into the 6/16 TIER ONE cohort and the rest into TIER FOUR on 7/7, or hold them out this round."),
      ("Don't stack levers","Move states OR change a tier's target in a given week, not both — otherwise the ROAS move can't be attributed and you stack two learning resets."),
    ]
    r=3
    for h,t in items:
        ws.cell(row=r,column=1,value=h).font=BOLD; ws.cell(row=r,column=1).fill=SUBFILL
        ws.merge_cells(f"A{r}:A{r}"); r+=1
        c=ws.cell(row=r,column=1,value=t); c.font=NOTE; c.alignment=WRAP; ws.merge_cells(f"A{r}:H{r}")
        ws.row_dimensions[r].height=44; r+=2
    ws.column_dimensions["A"].width=22
    for col in "BCDEFGH": ws.column_dimensions[col].width=12

    wb.save(out)


if __name__=="__main__":
    ap=argparse.ArgumentParser()
    ap.add_argument("--input",default="google_geo_model_gp_per_lead_2026-06-09.xlsx")
    ap.add_argument("--out",default=f"Google_Home_ReTier_Rollout_Plan_{dt.date.today():%Y-%m-%d}.xlsx")
    a=ap.parse_args()
    df=load(a.input)
    build(df,a.out)
    print(f"wrote {a.out}  ({len(df)} states, {df['chg'].sum()} moving)")
