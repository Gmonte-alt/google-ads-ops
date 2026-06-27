"""
google_geo_model.py — Google Home geo-tiering, ported from microsoft_geo_model_comparison.py.

METHOD IS IDENTICAL to the Microsoft core: the four models (A2 RPL-quantile,
B1 k-means on RPL/CPC, B4 Ward hierarchical, C2 weighted k-means on ideal_tcpa with
sample_weight=spend), the RPL 12w/4w blend (--rpl-blend), the eligibility gate, the
tier relabelling by RPL, and the 12w/8w/4w stability check are copied verbatim. The
headline model is C2_weighted_kmeans.

GOOGLE-SPECIFIC CHANGES (only the data layer + scope differ):
  - account_id = 1 (Google Home).
  - Source table: campaign_performance_geo_daily (NOT geo_performance_daily). It has no
    'ALL' device row, so state totals SUM the device rows (device <> 'ALL'), and it has
    no 'network' or 'location_type' columns. This mirrors analyze_geo_state_device.py,
    so the tiers stay consistent with the workbook's RPL_Redistribution.
  - Native conversion basis (conversions / conversion_value) — same caveat as the
    workbook: geo conversions are Google-native, not routed pixel-primary. RPL *shape*
    and ordering are sound; absolute $ run ~10-13% conservative. Tiering only needs the
    ordering/economics, so this is fine.
  - Search-only: Google Home has no Audience Network, so the Microsoft search-vs-audience
    network economics (and the whole Master Tier Review / k-sensitivity apparatus in
    model_compare_extras) are intentionally NOT ported. This script outputs the tier map,
    per-tier RPL bands, model comparison, and stability — the decision-grade pieces.
  - tROAS account: ideal_tcpa = RPL / TARGET_ROAS is kept as the clustering feature and an
    ANALYTICAL destination (ROAS-factored), per Gilbert. It is NOT an operational bid; the
    operational lever is tROAS per tier, set separately. The tier *membership* is the output.

Usage (Google analogue of the Microsoft run):
    python google_geo_model.py --rpl-blend 0.6 --k 6
    python google_geo_model.py --end-date 2026-06-07 --rpl-blend 0.6 --k 6
    python google_geo_model.py --rpl-blend 0.6 --k 6 --ledger commitments\\google_home_baseline.yaml
"""
from __future__ import annotations

import argparse
import sqlite3
import sys
from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path

import numpy as np
import pandas as pd

# ----------------------------------------------------------------------------
# Configuration (Google)
# ----------------------------------------------------------------------------
DEFAULT_DB_PATH = r"C:\Projects\google-ads-ops\db\paid_search.db"
DEFAULT_OUTPUT_DIR = r"C:\Projects\google-ads-ops\analysis\google-geo-analysis\output"
ACCOUNT_ID = 1
GEO_TABLE = "campaign_performance_geo_daily"
CLICK_THRESHOLD = 200          # 12w eligibility (tunable; Google's pool is ~26-28 states)
TARGET_ROAS = 1.50             # ideal_tcpa = RPL / TARGET_ROAS (analytical, ROAS-factored)
DEFAULT_K = 6

TIER_LABEL_BASE = [
    "TIER ONE", "TIER TWO", "TIER THREE", "TIER FOUR", "TIER FIVE",
    "TIER SIX", "TIER SEVEN", "TIER EIGHT", "TIER NINE", "TIER TEN",
    "TIER ELEVEN", "TIER TWELVE",
]
MODEL_NAMES = [
    "A2_rpl_quantile", "B1_kmeans_rpl_cpc", "B4_hierarchical", "C2_weighted_kmeans",
]
# Current Google campaign labels, for the current-vs-proposed comparison
CURRENT_TIER_NAMES = ["Tier One", "Tier Two", "Tier Three", "Tier Four",
                      "California", "Texas", "Florida"]


@dataclass
class Windows:
    end_date: date
    w12_start: date
    w8_start: date
    w4_start: date


def build_windows(end_date: date) -> Windows:
    return Windows(
        end_date=end_date,
        w12_start=end_date - timedelta(days=83),
        w8_start=end_date - timedelta(days=55),
        w4_start=end_date - timedelta(days=27),
    )


def current_label(name: str) -> str:
    u = (name or "").upper()
    for nm in CURRENT_TIER_NAMES:
        if nm.upper() in u:
            return nm
    return name or ""


# ----------------------------------------------------------------------------
# Data pull (Google: campaign_performance_geo_daily, sum device rows, search-only)
# ----------------------------------------------------------------------------

def resolve_end_date(conn: sqlite3.Connection, override: str | None) -> date:
    if override:
        return date.fromisoformat(override)
    row = conn.execute(
        f"SELECT MAX(date) FROM {GEO_TABLE} WHERE account_id = ?", (ACCOUNT_ID,)
    ).fetchone()
    if not row or not row[0]:
        raise RuntimeError(f"No rows in {GEO_TABLE} for account_id={ACCOUNT_ID}")
    max_date = date.fromisoformat(row[0])
    yesterday = date.today() - timedelta(days=1)
    return yesterday if max_date == yesterday else max_date


def pull_state_aggregates(conn: sqlite3.Connection, start: date, end: date) -> pd.DataFrame:
    """Per-state account-level aggregates (all Home search campaigns). Sums device
    rows (no 'ALL' row exists in this table); native conversion basis."""
    sql = f"""
        SELECT
            geo_state AS state,
            SUM(clicks) AS clicks,
            SUM(cost_micros) / 1e6 AS spend,
            SUM(conversions) AS conversions,
            SUM(conversion_value) AS revenue
        FROM {GEO_TABLE}
        WHERE account_id = ? AND device <> 'ALL' AND geo_state IS NOT NULL
          AND date BETWEEN ? AND ?
        GROUP BY geo_state
    """
    return pd.read_sql_query(sql, conn, params=(ACCOUNT_ID, start.isoformat(), end.isoformat()))


def pull_state_to_current_tier(conn: sqlite3.Connection, start: date, end: date) -> dict[str, str]:
    """Map each state to its current primary campaign (by spend) -> current tier label.
    Defensive: if campaign_id isn't in the geo table, returns {} and the comparison
    column is simply omitted."""
    try:
        df = pd.read_sql_query(
            f"""SELECT g.geo_state AS state, c.campaign_name AS campaign,
                       SUM(g.cost_micros)/1e6 AS spend
                FROM {GEO_TABLE} g
                JOIN campaigns c ON c.campaign_id = g.campaign_id AND c.account_id = g.account_id
                WHERE g.account_id = ? AND g.device <> 'ALL' AND g.geo_state IS NOT NULL
                  AND g.date BETWEEN ? AND ?
                GROUP BY g.geo_state, c.campaign_name""",
            conn, params=(ACCOUNT_ID, start.isoformat(), end.isoformat()),
        )
        if df.empty:
            return {}
        top = df.sort_values("spend", ascending=False).groupby("state").head(1)
        return {s: current_label(c) for s, c in zip(top["state"], top["campaign"])}
    except Exception as e:  # column/table shape differs — skip gracefully
        print(f"  (current-tier map skipped: {e})", file=sys.stderr)
        return {}


# ============================================================================
# METHOD — copied verbatim from microsoft_geo_model_comparison.py (unchanged)
# ============================================================================

def compute_derived(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df["cpc"] = df["spend"] / df["clicks"].where(df["clicks"] > 0)
    df["cpa"] = df["spend"] / df["conversions"].where(df["conversions"] > 0)
    df["roas"] = df["revenue"] / df["spend"].where(df["spend"] > 0)
    df["rpl"] = df["revenue"] / df["conversions"].where(df["conversions"] > 0)
    df["ideal_tcpa"] = df["rpl"] / TARGET_ROAS
    return df


def apply_rpl_blend(df_12w: pd.DataFrame, df_4w: pd.DataFrame | None,
                    blend_weight: float, min_blend_clicks: int = 75) -> pd.DataFrame:
    out = df_12w.copy()
    out["rpl_12w_original"] = out["rpl"]
    out["blend_applied"] = False
    if blend_weight <= 0.0 or df_4w is None or df_4w.empty:
        return out
    blend_weight = min(max(blend_weight, 0.0), 1.0)
    df_4w_indexed = df_4w.set_index("state")

    def _blend(row):
        state = row["state"]; rpl_12w = row["rpl"]
        if state not in df_4w_indexed.index:
            return rpl_12w, False
        r4 = df_4w_indexed.loc[state]
        c4 = r4["clicks"]; rpl_4w = r4["rpl"]
        if pd.isna(rpl_4w) or rpl_4w <= 0 or c4 < min_blend_clicks:
            return rpl_12w, False
        return (1.0 - blend_weight) * rpl_12w + blend_weight * rpl_4w, True

    blended = out.apply(_blend, axis=1)
    out["rpl"] = [b[0] for b in blended]
    out["blend_applied"] = [b[1] for b in blended]
    out["ideal_tcpa"] = out["rpl"] / TARGET_ROAS
    return out


def eligible_pool(df: pd.DataFrame, click_threshold: int = CLICK_THRESHOLD) -> pd.DataFrame:
    return df[
        (df["clicks"] >= click_threshold)
        & df["rpl"].notna() & df["cpc"].notna()
        & (df["rpl"] > 0) & (df["cpc"] > 0)
        & (df["state"].str.upper() != "UNKNOWN")
    ].copy().reset_index(drop=True)


def relabel_by_rpl(states: pd.DataFrame, cluster_col: str, k: int) -> pd.Series:
    means = states.groupby(cluster_col)["rpl"].mean().sort_values(ascending=False)
    rank_map = {old: TIER_LABEL_BASE[i] for i, old in enumerate(means.index[:k])}
    return states[cluster_col].map(rank_map).astype(object)


def model_a2_rpl_quantile(states: pd.DataFrame, k: int) -> pd.Series:
    ranks = states["rpl"].rank(method="first", ascending=False)
    bins = pd.cut(ranks, bins=k, labels=False, include_lowest=True)
    out = states.copy(); out["_cluster"] = bins
    return relabel_by_rpl(out, "_cluster", k)


def model_b1_kmeans_rpl_cpc(states: pd.DataFrame, k: int, seed: int = 42) -> pd.Series:
    from sklearn.cluster import KMeans
    from sklearn.preprocessing import StandardScaler
    X = StandardScaler().fit_transform(states[["rpl", "cpc"]].values)
    km = KMeans(n_clusters=k, random_state=seed, n_init=10).fit(X)
    out = states.copy(); out["_cluster"] = km.labels_
    return relabel_by_rpl(out, "_cluster", k)


def model_b4_hierarchical(states: pd.DataFrame, k: int):
    from scipy.cluster.hierarchy import linkage, fcluster
    from sklearn.preprocessing import StandardScaler
    feat = pd.DataFrame({
        "rpl": states["rpl"].values, "cpc": states["cpc"].values,
        "log_clicks": np.log(states["clicks"].values),
    })
    X = StandardScaler().fit_transform(feat.values)
    Z = linkage(X, method="ward")
    labels = fcluster(Z, t=k, criterion="maxclust")
    out = states.copy(); out["_cluster"] = labels
    return relabel_by_rpl(out, "_cluster", k), Z


def model_c2_weighted_kmeans(states: pd.DataFrame, k: int, seed: int = 42) -> pd.Series:
    """1D k-means on ideal_tcpa with sample_weight = spend (headline model)."""
    from sklearn.cluster import KMeans
    X = states[["ideal_tcpa"]].values
    w = states["spend"].values
    km = KMeans(n_clusters=k, random_state=seed, n_init=10).fit(X, sample_weight=w)
    out = states.copy(); out["_cluster"] = km.labels_
    return relabel_by_rpl(out, "_cluster", k)


def apply_all_models(states: pd.DataFrame, k: int) -> tuple[pd.DataFrame, dict]:
    out = states.copy().reset_index(drop=True)
    out["A2_rpl_quantile"] = model_a2_rpl_quantile(out, k)
    out["B1_kmeans_rpl_cpc"] = model_b1_kmeans_rpl_cpc(out, k)
    b4_labels, Z = model_b4_hierarchical(out, k)
    out["B4_hierarchical"] = b4_labels
    out["C2_weighted_kmeans"] = model_c2_weighted_kmeans(out, k)
    return out, {"b4_linkage": Z, "b4_state_order": out["state"].tolist()}


def tier_aggregate_metrics(states: pd.DataFrame, tier_col: str) -> pd.DataFrame:
    rows = []
    for tier, grp in states.groupby(tier_col, observed=True):
        if pd.isna(tier) or tier == "":
            continue
        sp = grp["spend"].sum()
        weighted_rpl = (grp["rpl"] * grp["spend"]).sum() / sp if sp > 0 else np.nan
        rpl_min, rpl_max = grp["rpl"].min(), grp["rpl"].max()
        rows.append({
            "tier": tier, "n_states": len(grp),
            "states": ", ".join(sorted(grp["state"].tolist())),
            "tier_spend": sp, "tier_weighted_rpl": weighted_rpl,
            "tier_target_tcpa_analytical": weighted_rpl / TARGET_ROAS if pd.notna(weighted_rpl) else np.nan,
            "rpl_min": rpl_min, "rpl_max": rpl_max,
            "rpl_spread_ratio": rpl_max / rpl_min if rpl_min > 0 else np.nan,
            "ideal_tcpa_stdev": grp["ideal_tcpa"].std(ddof=0),
        })
    df = pd.DataFrame(rows)
    if df.empty:
        return df
    order = {t: i for i, t in enumerate(TIER_LABEL_BASE)}
    df["_o"] = df["tier"].map(order)
    return df.sort_values("_o").drop(columns=["_o"]).reset_index(drop=True)


def model_summary_metrics(states: pd.DataFrame, tier_col: str) -> dict:
    tier_agg = tier_aggregate_metrics(states, tier_col)
    if tier_agg.empty:
        return {}
    tier_tcpa_map = dict(zip(tier_agg["tier"], tier_agg["tier_target_tcpa_analytical"]))
    s = states.copy()
    s["_tier_tcpa"] = s[tier_col].map(tier_tcpa_map)
    bid_loss = ((s["ideal_tcpa"] - s["_tier_tcpa"]) ** 2 * s["spend"]).sum()
    valid = tier_agg.dropna(subset=["rpl_spread_ratio"])
    avg_spread = ((valid["rpl_spread_ratio"] * valid["n_states"]).sum() / valid["n_states"].sum()
                  if not valid.empty else np.nan)
    vstd = tier_agg.dropna(subset=["ideal_tcpa_stdev"])
    avg_stdev = ((vstd["ideal_tcpa_stdev"] * vstd["n_states"]).sum() / vstd["n_states"].sum()
                 if not vstd.empty else np.nan)
    return {
        "model": tier_col, "n_tiers_populated": len(tier_agg),
        "avg_within_tier_rpl_spread": avg_spread,
        "avg_within_tier_tcpa_stdev": avg_stdev,
        "spend_weighted_bid_loss": bid_loss,
    }


def run_models_for_window(conn, start, end, k, click_threshold=CLICK_THRESHOLD):
    agg = pull_state_aggregates(conn, start, end)
    if agg.empty:
        return None
    agg = compute_derived(agg)
    elig = eligible_pool(agg, click_threshold=click_threshold)
    if len(elig) < k:
        return None
    result, _ = apply_all_models(elig, k)
    return result


def build_stability_table(result_12w, result_8w, result_4w) -> pd.DataFrame:
    rows = []
    states = sorted(set(result_12w["state"].tolist()))
    for r in (result_8w, result_4w):
        if r is not None:
            states = sorted(set(states) | set(r["state"].tolist()))
    idx_12w = result_12w.set_index("state") if result_12w is not None else None
    idx_8w = result_8w.set_index("state") if result_8w is not None else None
    idx_4w = result_4w.set_index("state") if result_4w is not None else None

    def trank(t):
        if t is None or (isinstance(t, float) and np.isnan(t)):
            return None
        try:
            return TIER_LABEL_BASE.index(t)
        except ValueError:
            return None

    for state in states:
        for model in MODEL_NAMES:
            t12 = idx_12w.loc[state, model] if (idx_12w is not None and state in idx_12w.index) else None
            t8 = idx_8w.loc[state, model] if (idx_8w is not None and state in idx_8w.index) else None
            t4 = idx_4w.loc[state, model] if (idx_4w is not None and state in idx_4w.index) else None
            present = [t for t in (t12, t8, t4) if t is not None and not (isinstance(t, float) and pd.isna(t))]
            n_unique = len(set(present)) if present else 0
            ranks = [trank(t) for t in (t12, t8, t4)]
            valid = [r for r in ranks if r is not None]
            if not valid:
                direction = "no_data"
            elif n_unique == 1:
                direction = "stable"
            elif len(valid) < 3:
                direction = "partial"
            elif ranks[2] < ranks[0]:
                direction = "drifting_up"
            elif ranks[2] > ranks[0]:
                direction = "drifting_down"
            else:
                direction = "volatile"
            rows.append({"state": state, "model": model, "tier_12w": t12,
                         "tier_8w": t8, "tier_4w": t4,
                         "stability_score": n_unique, "direction": direction})
    return pd.DataFrame(rows)


def assign_below_threshold(all_states: pd.DataFrame, eligible_result: pd.DataFrame,
                           headline: str, click_threshold: int) -> pd.DataFrame:
    """States below the click threshold, assigned to the nearest headline-model tier
    centroid (in ideal_tcpa). Lightweight analogue of the Microsoft below-threshold step."""
    if eligible_result is None or eligible_result.empty:
        return pd.DataFrame()
    below = all_states[
        (all_states["clicks"] < click_threshold)
        & all_states["rpl"].notna() & (all_states["rpl"] > 0)
        & (all_states["state"].str.upper() != "UNKNOWN")
    ].copy()
    if below.empty:
        return pd.DataFrame()
    cent = eligible_result.groupby(headline)["ideal_tcpa"].mean()
    below["assigned_tier_nearest"] = below["ideal_tcpa"].apply(lambda t: (cent - t).abs().idxmin())
    cols = ["state", "clicks", "conversions", "spend", "rpl", "cpc", "ideal_tcpa", "assigned_tier_nearest"]
    return below[cols].sort_values("rpl", ascending=False).reset_index(drop=True)


# ----------------------------------------------------------------------------
# Ledger comparison (optional)
# ----------------------------------------------------------------------------

def load_ledger(path: Path) -> dict[str, str]:
    try:
        import yaml
    except ImportError:
        print("  (PyYAML not installed; ledger comparison skipped)", file=sys.stderr)
        return {}
    if not path.exists():
        return {}
    with open(path) as f:
        data = yaml.safe_load(f) or {}
    root = data.get("google_home_geo_tiers", data)
    return dict(root.get("commitments", {}))


# ----------------------------------------------------------------------------
# Workbook
# ----------------------------------------------------------------------------

def _write_df(ws, df: pd.DataFrame, header_row: int = 1):
    from openpyxl.styles import Font, PatternFill
    hf = Font(bold=True, color="FFFFFF"); hfill = PatternFill("solid", fgColor="305496")
    money = {"spend", "tier_spend", "revenue"}
    twod = {"rpl", "cpc", "ideal_tcpa", "tier_weighted_rpl", "tier_target_tcpa_analytical",
            "rpl_min", "rpl_max", "rpl_spread_ratio", "ideal_tcpa_stdev",
            "avg_within_tier_rpl_spread", "avg_within_tier_tcpa_stdev"}
    cols = list(df.columns)
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=header_row, column=j, value=str(c)); cell.font = hf; cell.fill = hfill
    for i, (_, row) in enumerate(df.iterrows(), header_row + 1):
        for j, c in enumerate(cols, 1):
            v = row[c]
            if isinstance(v, (np.floating,)):
                v = float(v)
            elif isinstance(v, (np.integer,)):
                v = int(v)
            if isinstance(v, float) and pd.isna(v):
                v = None
            cell = ws.cell(row=i, column=j, value=v)
            if c in money:
                cell.number_format = '"$"#,##0'
            elif c == "spend_weighted_bid_loss":
                cell.number_format = '#,##0'
            elif c in twod:
                cell.number_format = "0.00"
    ws.freeze_panes = ws.cell(row=header_row + 1, column=2).coordinate
    return header_row + len(df)


def write_workbook(path: Path, windows: Windows, k: int, rpl_blend: float,
                   click_threshold: int, result_12w: pd.DataFrame, headline: str,
                   tier_metrics: pd.DataFrame, model_compare: pd.DataFrame,
                   stability: pd.DataFrame, below: pd.DataFrame,
                   current_map: dict, ledger_cmp: pd.DataFrame | None):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    TITLE = Font(bold=True, size=14, color="305496")
    NOTE = Font(italic=True, color="595959", size=9)
    wb = Workbook(); wb.remove(wb.active)

    # 1) Proposed tiers (headline C2) — the decision artifact
    ws = wb.create_sheet("Proposed Tiers (C2)")
    ws["A1"] = f"Google Home geo tiers — {headline}, k={k}"; ws["A1"].font = TITLE
    ws["A2"] = (f"Window {windows.w12_start}..{windows.end_date} (12w). "
                f"RPL blend weight={rpl_blend:.2f} (4w into 12w). Eligibility >= {click_threshold} clicks. "
                f"tier_target_tcpa_analytical = tier RPL / {TARGET_ROAS} (ROAS-factored; NOT an operational bid).")
    ws["A2"].font = NOTE
    _write_df(ws, tier_metrics, header_row=4)

    # 2) State assignments (all four models + current tier for comparison)
    ws = wb.create_sheet("State Assignments")
    ws["A1"] = "Per-state assignments — all four models + current campaign"; ws["A1"].font = TITLE
    ws["A2"] = "Headline = C2_weighted_kmeans. 'current_tier' is where the state sits today (by spend)."
    ws["A2"].font = NOTE
    base = ["state"]
    assign = result_12w.copy()
    if current_map:
        assign["current_tier"] = assign["state"].map(current_map)
        base.append("current_tier")
    cols = base + ["clicks", "conversions", "spend", "rpl", "cpc", "ideal_tcpa"] + MODEL_NAMES
    cols = [c for c in cols if c in assign.columns]
    _write_df(ws, assign[cols].sort_values("rpl", ascending=False).reset_index(drop=True), header_row=4)

    # 3) Model comparison
    ws = wb.create_sheet("Model Comparison")
    ws["A1"] = "Model comparison (lower bid-loss / spread / stdev = tighter tiers)"; ws["A1"].font = TITLE
    _write_df(ws, model_compare, header_row=3)

    # 4) Stability
    ws = wb.create_sheet("Stability (12w_8w_4w)")
    ws["A1"] = "Tier stability across 12w / 8w / 4w windows"; ws["A1"].font = TITLE
    ws["A2"] = "stability_score = # distinct tiers across windows (1 = rock-stable). direction flags drift."
    ws["A2"].font = NOTE
    st = stability[stability["model"] == headline].drop(columns=["model"]) if not stability.empty else stability
    _write_df(ws, st.reset_index(drop=True), header_row=4)

    # 5) Below threshold
    if below is not None and not below.empty:
        ws = wb.create_sheet("Below Threshold")
        ws["A1"] = f"States below {click_threshold} clicks — nearest-tier assignment"; ws["A1"].font = TITLE
        ws["A2"] = "Assigned to the nearest C2 tier centroid by ideal_tcpa. Lower confidence — thin volume."
        ws["A2"].font = NOTE
        _write_df(ws, below, header_row=4)

    # 6) Ledger comparison
    if ledger_cmp is not None and not ledger_cmp.empty:
        ws = wb.create_sheet("Ledger Comparison")
        ws["A1"] = "Proposed tier vs committed ledger"; ws["A1"].font = TITLE
        _write_df(ws, ledger_cmp, header_row=3)

    # 7) Methodology
    ws = wb.create_sheet("Methodology")
    notes = [
        ("Google Home geo-tiering — methodology", True),
        ("", False),
        ("Method is identical to microsoft_geo_model_comparison.py: A2 RPL-quantile, B1 k-means(RPL,CPC),", False),
        ("B4 Ward hierarchical, C2 weighted k-means on ideal_tcpa (sample_weight=spend). Headline = C2.", False),
        ("Tiers are relabelled TIER ONE..N by descending mean RPL. RPL blend mixes 4w into 12w per --rpl-blend.", False),
        ("", False),
        ("Google-specific:", True),
        ("- account_id=1; source campaign_performance_geo_daily (sums device rows; no 'ALL' row; no network/location_type).", False),
        ("- Native conversion basis (same as the workbook RPL_Redistribution): RPL ordering sound; $ ~10-13% conservative.", False),
        ("- Search-only: Microsoft search-vs-audience network economics and the Master Tier Review apparatus are NOT ported.", False),
        ("- tROAS account: ideal_tcpa = RPL/1.5 is the clustering feature and an analytical destination, NOT an operational bid.", False),
        ("  The output that matters operationally is the tier MEMBERSHIP (which states group together).", False),
    ]
    for r, (t, hdr) in enumerate(notes, 1):
        ws.cell(row=r, column=1, value=t).font = TITLE if hdr else NOTE
    ws.column_dimensions["A"].width = 110

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# ----------------------------------------------------------------------------
# Main
# ----------------------------------------------------------------------------

def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser()
    p.add_argument("--end-date", type=str, default=None)
    p.add_argument("--k", type=int, default=DEFAULT_K)
    p.add_argument("--rpl-blend", type=float, default=0.0,
                   help="Weight for blending 4w RPL into 12w (0=pure 12w; 0.6=40%% 12w + 60%% 4w).")
    p.add_argument("--click-threshold", type=int, default=CLICK_THRESHOLD)
    p.add_argument("--headline-model", type=str, default="C2_weighted_kmeans")
    p.add_argument("--db-path", type=str, default=DEFAULT_DB_PATH)
    p.add_argument("--output-dir", type=str, default=DEFAULT_OUTPUT_DIR)
    p.add_argument("--ledger", type=str, default=None,
                   help="Optional commitment ledger YAML for a proposed-vs-committed comparison.")
    args = p.parse_args(argv)

    db_path = Path(args.db_path)
    if not db_path.exists():
        print(f"ERROR: DB not found: {db_path}", file=sys.stderr)
        return 1
    conn = sqlite3.connect(str(db_path))

    end_date = resolve_end_date(conn, args.end_date)
    windows = build_windows(end_date)
    k = args.k
    print("GOOGLE HOME GEO-TIERING")
    print(f"  end_date={end_date}  12w={windows.w12_start}..{end_date}  k={k}  "
          f"rpl_blend={args.rpl_blend}  click_threshold={args.click_threshold}")

    raw_12w = pull_state_aggregates(conn, windows.w12_start, windows.end_date)
    if raw_12w.empty:
        print(f"ERROR: no rows in {GEO_TABLE} for the 12w window.", file=sys.stderr)
        return 1
    raw_4w = None
    if args.rpl_blend > 0.0:
        raw_4w = pull_state_aggregates(conn, windows.w4_start, windows.end_date)
        raw_4w = compute_derived(raw_4w) if not raw_4w.empty else None

    all_states_12w = compute_derived(raw_12w)
    if args.rpl_blend > 0.0:
        all_states_12w = apply_rpl_blend(all_states_12w, raw_4w, args.rpl_blend)
        print(f"  RPL blend active: {int(all_states_12w['blend_applied'].sum())}/{len(all_states_12w)} states blended")

    elig = eligible_pool(all_states_12w, click_threshold=args.click_threshold)
    print(f"  eligible states: {len(elig)} of {len(all_states_12w)} with data")
    if len(elig) < k:
        print(f"ERROR: eligible pool ({len(elig)}) < k ({k}). Lower --k or --click-threshold.", file=sys.stderr)
        return 1

    result_12w, _ = apply_all_models(elig, k)
    result_8w = run_models_for_window(conn, windows.w8_start, windows.end_date, k, click_threshold=75)
    result_4w = run_models_for_window(conn, windows.w4_start, windows.end_date, k, click_threshold=75)
    stability = build_stability_table(result_12w, result_8w, result_4w)

    headline = args.headline_model
    tier_metrics = tier_aggregate_metrics(result_12w, headline)
    model_compare = pd.DataFrame([model_summary_metrics(result_12w, m) for m in MODEL_NAMES])
    below = assign_below_threshold(all_states_12w, result_12w, headline, args.click_threshold)
    current_map = pull_state_to_current_tier(conn, windows.w12_start, windows.end_date)

    ledger_cmp = None
    commitments = load_ledger(Path(args.ledger)) if args.ledger else {}
    if commitments:
        prop = dict(zip(result_12w["state"], result_12w[headline]))
        rows = []
        for st in sorted(set(prop) | set(commitments)):
            pv, cv = prop.get(st), commitments.get(st)
            rows.append({"state": st, "proposed_tier": pv, "committed_tier": cv,
                         "agreement": "SAME" if pv == cv else "MOVED"})
        ledger_cmp = pd.DataFrame(rows)

    conn.close()
    print("\n  Proposed tiers (C2):")
    for _, r in tier_metrics.iterrows():
        print(f"    {r['tier']:12} n={int(r['n_states']):2}  RPL {r['rpl_min']:.0f}-{r['rpl_max']:.0f} "
              f"(wtd {r['tier_weighted_rpl']:.0f})  {r['states']}")

    out = Path(args.output_dir) / f"google_geo_model_{end_date}.xlsx"
    write_workbook(out, windows, k, args.rpl_blend, args.click_threshold,
                   result_12w, headline, tier_metrics, model_compare, stability,
                   below, current_map, ledger_cmp)
    print(f"\n  Wrote: {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
