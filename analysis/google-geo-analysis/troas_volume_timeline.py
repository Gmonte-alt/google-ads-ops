"""
troas_volume_timeline.py — visualize the downward pressure tROAS changes put on volume.

For each current campaign, plots weekly clicks/day and conversions/day across the period with
the reconstructed tROAS target level overlaid (step line) and vertical markers at each tROAS
change, so you can see volume bend after the bidder is tightened. Source: campaign_performance_geo_daily_search
(account 1, native conv, sum device rows) for volume; the Google Ads change-history CSV for tROAS events.

Usage:
    python troas_volume_timeline.py --bid-history "Change_history_report_google-home-bid-Mar1-Jun8.csv"
    python troas_volume_timeline.py --bid-history bid.csv --start 2026-03-01 --end 2026-06-08
"""
from __future__ import annotations
import argparse, re, sqlite3, sys
from pathlib import Path
import numpy as np, pandas as pd

DEFAULT_DB = r"C:\Projects\google-ads-ops\db\paid_search.db"
DEFAULT_OUT = r"C:\Projects\google-ads-ops\analysis\google-geo-analysis\output"
ACCOUNT_ID = 1
GEO = "campaign_performance_geo_daily_search"
CURRENT_CAMPAIGNS = ["Tier One", "Tier Two", "Tier Three", "Tier Four", "California", "Texas", "Florida"]
# Campaign-level logs as "...from 115% to 112%"; portfolio-level logs as "...from 1.19 to 1.2"
# (decimal ratio, strategy name in quotes). Match both; %? is optional.
TROAS_PAT = re.compile(r"Target ROAS changed from ([\d.]+)\s*%?\s*to\s*([\d.]+)\s*%?")
DATE_PAT = re.compile(r"([A-Z][a-z]{2,8}\s+\d{1,2},\s+\d{4})")


def _to_pct(v):
    """Normalize to percent: 115% stays 115; a 1.19 ratio becomes 119."""
    return v if v > 10 else v * 100.0


def _fmt_date(d):
    """Portable short date (e.g. 'May 6'); avoids the %-d / %#d strftime split across OSes."""
    return f"{d.strftime('%b')} {d.day}" if hasattr(d, "strftime") else str(d)


def current_label(name):
    u = (name or "").upper()
    for c in CURRENT_CAMPAIGNS:
        if c.upper() in u:
            return c
    return "Other"


def parse_events(csv_path):
    """Return DataFrame [date, campaign, old, new] of tROAS changes."""
    df = pd.read_csv(csv_path, skiprows=2)
    rows = []
    for _, r in df.iterrows():
        m = TROAS_PAT.search(str(r.get("Changes", "")))
        if not m:
            continue
        dm = DATE_PAT.search(str(r["Date & time"]))         # avoids NBSP/time-zone parsing issues
        d = pd.to_datetime(dm.group(1), errors="coerce") if dm else pd.NaT
        if pd.isna(d):
            continue
        rows.append({"date": d.normalize(), "campaign": current_label(str(r["Campaign"])),
                     "old": _to_pct(float(m.group(1))), "new": _to_pct(float(m.group(2)))})
    ev = pd.DataFrame(rows)
    return ev[ev["campaign"] != "Other"].sort_values("date").reset_index(drop=True) if not ev.empty else ev


def pull_weekly_volume(conn, start, end):
    sql = f"""SELECT g.date AS date, c.campaign_name AS campaign,
                     SUM(g.clicks) clicks, SUM(g.conversions) conv,
                     SUM(g.conversion_value) rev, SUM(g.cost_micros)/1e6 cost
              FROM {GEO} g JOIN campaigns c ON c.campaign_id=g.campaign_id AND c.account_id=g.account_id
              WHERE g.account_id=? AND g.device<>'ALL' AND g.date BETWEEN ? AND ?
              GROUP BY g.date, c.campaign_name"""
    df = pd.read_sql_query(sql, conn, params=(ACCOUNT_ID, start, end))
    if df.empty:
        return df
    df["campaign"] = df["campaign"].map(current_label)
    df = df[df["campaign"] != "Other"]
    df["date"] = pd.to_datetime(df["date"])
    df["week"] = df["date"] - pd.to_timedelta(df["date"].dt.weekday, unit="D")  # Monday
    wk = df.groupby(["campaign", "week"], as_index=False).agg(
        clicks=("clicks", "sum"), conv=("conv", "sum"), rev=("rev", "sum"), cost=("cost", "sum"),
        days=("date", "nunique"))
    wk["clicks_day"] = wk["clicks"] / 7.0
    wk["conv_day"] = wk["conv"] / 7.0
    wk["roas"] = np.where(wk["cost"] > 0, wk["rev"] / wk["cost"], np.nan)
    # complete week = fully inside [start, end] (drops the trailing partial week and any partial lead)
    s, e = pd.to_datetime(start), pd.to_datetime(end)
    wk["complete"] = (wk["week"] >= s) & (wk["week"] + pd.Timedelta(days=6) <= e)
    return wk


def troas_step(ev_camp, start, end):
    """Reconstruct (dates, levels) step series for one campaign's tROAS over [start, end]."""
    if ev_camp.empty:
        return None
    e = ev_camp.sort_values("date")
    xs = [pd.to_datetime(start)]
    ys = [e.iloc[0]["old"]]
    for _, r in e.iterrows():
        xs.append(r["date"]); ys.append(r["new"])
    xs.append(pd.to_datetime(end)); ys.append(ys[-1])
    return xs, ys


def make_campaign_png(camp, wk, ev_camp, start, end, out_png):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    fig, axes = plt.subplots(3, 1, figsize=(9, 7.2), sharex=True)
    step = troas_step(ev_camp, start, end)

    def vlines(ax):
        for _, r in ev_camp.iterrows():
            ax.axvline(r["date"], color="#999", lw=0.8, ls=":")

    # Panel 1 & 2 — volume, with the target tROAS step on a secondary axis
    for ax, (col, lbl, color) in zip(axes[:2], [("clicks_day", "clicks/day", "#4f81bd"),
                                                ("conv_day", "conversions/day", "#c0504d")]):
        ax.plot(wk["week"], wk[col], "-o", color=color, ms=3, lw=1.8)
        ax.set_ylabel(lbl, color=color, fontsize=9); ax.tick_params(axis="y", labelcolor=color)
        ax.grid(alpha=0.25); vlines(ax)
        if step is not None:
            ax2 = ax.twinx()
            ax2.step(step[0], step[1], where="post", color="#555", lw=1.4, ls="--", label="target tROAS %")
            ax2.set_ylabel("tROAS %", color="#555", fontsize=8)
            ax2.tick_params(axis="y", labelcolor="#555", labelsize=8)
            if ax is axes[0]:
                ax2.legend(fontsize=7.5, loc="lower right")

    # Panel 3 — achieved ROAS vs target tROAS (same %% axis, directly comparable)
    ax = axes[2]
    ax.plot(wk["week"], wk["roas"] * 100.0, "-o", color="#2e7d32", ms=3, lw=1.8, label="achieved ROAS %")
    if step is not None:
        ax.step(step[0], step[1], where="post", color="#555", lw=1.6, ls="--", label="target tROAS %")
    ax.set_ylabel("ROAS %", fontsize=9); ax.grid(alpha=0.25); vlines(ax)
    ax.legend(fontsize=8, loc="best")

    if ev_camp.empty:
        chs = "no tROAS changes logged"
    else:
        chs = "  ".join(f"{_fmt_date(r['date'])}: {r['old']:.0f}\u2192{r['new']:.0f}"
                        for _, r in ev_camp.iterrows())
    axes[0].set_title(f"{camp}   |   tROAS changes: {chs}", fontsize=9.5)
    fig.tight_layout()
    fig.savefig(out_png, dpi=140); plt.close(fig)


def compute_timeline(db_path, bid_history, start, end, output_dir):
    """Parse tROAS events, pull weekly volume/ROAS, render per-campaign PNGs. Returns (camps,pngs,wk,ev) or None."""
    ev = parse_events(bid_history)
    print(f"tROAS change events parsed: {len(ev)} across {ev['campaign'].nunique() if not ev.empty else 0} campaigns")
    db = Path(db_path)
    if not db.exists():
        print(f"ERROR: DB not found: {db}", file=sys.stderr); return None
    conn = sqlite3.connect(str(db))
    wk = pull_weekly_volume(conn, start, end)
    conn.close()
    if wk.empty:
        print("ERROR: no weekly volume rows.", file=sys.stderr); return None
    outdir = Path(output_dir); outdir.mkdir(parents=True, exist_ok=True)
    camps = [c for c in CURRENT_CAMPAIGNS if c in set(wk["campaign"])]
    pngs = {}
    for camp in camps:
        cw = wk[(wk["campaign"] == camp) & wk["complete"]].sort_values("week")  # drop partial weeks from charts
        ce = ev[ev["campaign"] == camp] if not ev.empty else ev
        png = outdir / f"_troas_vol_{camp.replace(' ', '_')}.png"
        make_campaign_png(camp, cw, ce, start, end, png)
        pngs[camp] = png
        print(f"  {camp}: {len(cw)} complete weeks, {len(ce)} tROAS changes")
    return camps, pngs, wk, ev


def append_to_workbook(wb, db_path, bid_history, start="2026-03-01", end="2026-06-08", output_dir="."):
    """Compute the timeline and add its sheets to an existing workbook (no standalone file)."""
    res = compute_timeline(db_path, bid_history, start, end, output_dir)
    if res is None:
        raise RuntimeError("troas_volume_timeline: no weekly volume rows or DB not found")
    camps, pngs, wk, ev = res
    write_wb(None, camps, pngs, wk, ev, start, end, wb=wb)
    return wb


def main(argv=None):
    p = argparse.ArgumentParser()
    p.add_argument("--bid-history", required=True, help="Google Ads change-history CSV (bid).")
    p.add_argument("--start", default="2026-03-01")
    p.add_argument("--end", default="2026-06-08")
    p.add_argument("--db-path", default=DEFAULT_DB)
    p.add_argument("--output-dir", default=DEFAULT_OUT)
    args = p.parse_args(argv)
    res = compute_timeline(args.db_path, args.bid_history, args.start, args.end, args.output_dir)
    if res is None:
        return 1
    camps, pngs, wk, ev = res
    out = Path(args.output_dir) / f"troas_volume_timeline_{args.end}.xlsx"
    write_wb(out, camps, pngs, wk, ev, args.start, args.end)
    print(f"\n  Wrote: {out}")
    return 0


def write_wb(path, camps, pngs, wk, ev, start, end, wb=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.drawing.image import Image as XLImage
    TITLE = Font(bold=True, size=14, color="305496"); NOTE = Font(italic=True, color="595959", size=9)
    HF = Font(bold=True, color="FFFFFF"); HFill = PatternFill("solid", fgColor="305496")
    standalone = wb is None
    if standalone:
        wb = Workbook(); wb.remove(wb.active)

    ws = wb.create_sheet("tROAS vs Volume")
    ws["A1"] = "tROAS changes vs weekly volume, by campaign"; ws["A1"].font = TITLE
    ws["A2"] = (f"{start} to {end}. Each panel: clicks/day & conversions/day (left) with the "
                f"reconstructed tROAS level (dashed step, right) and dotted markers at each change. "
                f"Bottom panel shows achieved ROAS vs target tROAS. Native conv; clicks don't lag. "
                f"Incomplete boundary weeks (incl. the trailing partial week) are excluded from the charts.")
    ws["A2"].font = NOTE
    row = 4
    for camp in camps:
        img = XLImage(str(pngs[camp]))
        ws.add_image(img, f"A{row}")
        row += 52   # ~7.2in 3-panel image needs ~52 rows of clearance

    ws = wb.create_sheet("Bid Changes")
    ws["A1"] = "Logged tROAS changes"; ws["A1"].font = TITLE
    for j, c in enumerate(["date", "campaign", "old %", "new %", "direction"], 1):
        x = ws.cell(row=3, column=j, value=c); x.font = HF; x.fill = HFill
    for i, (_, r) in enumerate(ev.iterrows(), 4):
        ws.cell(row=i, column=1, value=r["date"].strftime("%Y-%m-%d") if hasattr(r["date"], "strftime") else str(r["date"]))
        ws.cell(row=i, column=2, value=r["campaign"]); ws.cell(row=i, column=3, value=r["old"])
        ws.cell(row=i, column=4, value=r["new"])
        ws.cell(row=i, column=5, value="tighten" if r["new"] > r["old"] else "loosen")

    ws = wb.create_sheet("Weekly Volume & ROAS")
    for j, c in enumerate(["campaign", "week", "clicks/day", "conversions/day", "achieved ROAS %", "complete week"], 1):
        x = ws.cell(row=1, column=j, value=c); x.font = HF; x.fill = HFill
    for i, (_, r) in enumerate(wk.sort_values(["campaign", "week"]).iterrows(), 2):
        ws.cell(row=i, column=1, value=r["campaign"])
        ws.cell(row=i, column=2, value=r["week"].strftime("%Y-%m-%d"))
        ws.cell(row=i, column=3, value=round(r["clicks_day"], 1))
        ws.cell(row=i, column=4, value=round(r["conv_day"], 2))
        ws.cell(row=i, column=5, value=round(r["roas"] * 100.0, 1) if pd.notna(r["roas"]) else None)
        ws.cell(row=i, column=6, value="Y" if r["complete"] else "partial")
    ws.freeze_panes = "A2"

    if standalone:
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)
    return wb


if __name__ == "__main__":
    sys.exit(main())
