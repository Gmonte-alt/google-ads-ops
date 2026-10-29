"""
termlife_rpl_keyword_review.py  (v1)
=====================================
Pre/post review of the "RPL Increase 6.2.26" keyword cohort — the
high-CPC terms we told SQ we were priced out of, which the June 2
payout increase was negotiated to unlock.

Cohort definition: labels are NOT in paid_search.db, so the cohort is
pinned from the attached Google Ads UI keyword export (Label contains
"RPL Increase"). The export's metrics are a current-period snapshot;
every pre/post number here comes from the DB.

Questions answered
------------------
  1) Were these keywords in the account BEFORE the pay increase?
     (Presence sheet: first/last serve date per keyword x campaign)
  2) Activity before vs after, all platform metrics we store
     (impressions, clicks, CTR, CPC, cost, leads, CVR, CPL, conv
     value, value/lead, ROAS + per-day rates; QS where recorded)
  3) Benchmarks: account-wide and HighCPC-ad-group averages over the
     same windows, so keywords with thin/no pre data read against the
     account trend; relative-lift = cohort change / account change.

Windows (constants below): PRE Mar 1 - Jun 1 (last 3 full months
before the raise; Jan-Feb excluded to avoid the January category
spike flattering the pre period), POST Jun 2 - Sep 30 (full
September, incl. the Sep 29-30 schedule-cut days), Oct MTD separate. Windows differ in length -> every
table carries per-day rates.

Caveats written into the workbook: keyword-level impression share is
not stored in the DB (campaign grain only) - the UI export's IS
columns are a today-snapshot, not pre/post; conv value = SQ payment
(quality-blind); label membership is as-of the export date.
"""
import glob
import os
import sqlite3
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.drawing.text import RichTextProperties
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

DB_PATH = r"C:\Projects\google-ads-ops\db\paid_search.db"
ANALYSIS_DIR = (r"C:\Projects\google-ads-ops\analysis"
                r"\google-term-life-policy-kpi-ytd")
OUT_XLSX = os.path.join(ANALYSIS_DIR, "output",
                        "RPL_Increase_Keyword_Review.xlsx")
EXPORT_GLOB = "*Search_keyword_report*.csv"   # UI export in input\

ACCOUNT_ID = 29
RAISE_DATE = "2026-06-02"                      # from the label
PRE0, PRE1 = "2026-03-01", "2026-06-02"        # Mar 1 - Jun 1
POST0, POST1 = "2026-06-02", "2026-10-01"      # Jun 2 - Sep 30 (full)
OCT0, OCT1 = "2026-10-01", "2026-11-01"
TRAIL0 = "2026-01-01"

NAVY, LIGHT = "1F3864", "D9E2F2"
SERIES_HEX = ["2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4", "008300"]

pd.set_option("display.width", 240)
pd.set_option("display.max_columns", 44)
pd.set_option("display.max_rows", 250)
pd.set_option("display.float_format", lambda x: f"{x:,.1f}")


def hr(t):
    print("\n" + "=" * 110)
    print(t)
    print("=" * 110)


def meta(ws, title, source, window, caveat, takeaway=None):
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=13, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A2"] = f"Source: {source}   |   Window: {window}"
    ws["A3"] = f"Caveats: {caveat}"
    for r in (2, 3):
        ws[f"A{r}"].font = Font(italic=True, size=9)
    if takeaway:
        ws["A4"] = f"TAKEAWAY: {takeaway}"
        ws["A4"].font = Font(bold=True, size=10)
        ws["A4"].fill = PatternFill("solid", fgColor=LIGHT)
    ws.freeze_panes = "A6"


def write_block(ws, row, subtitle, df):
    ws.cell(row=row, column=1, value=subtitle).font = Font(bold=True,
                                                           size=11)
    ws.cell(row=row, column=1).fill = PatternFill("solid", fgColor=LIGHT)
    row += 1
    for j, c in enumerate(df.columns, start=1):
        cell = ws.cell(row=row, column=j, value=str(c))
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
    row += 1
    for _, r in df.iterrows():
        for j, v in enumerate(r, start=1):
            if isinstance(v, (np.floating, float)):
                v = None if pd.isna(v) else round(float(v), 2)
            elif isinstance(v, (np.integer,)):
                v = int(v)
            ws.cell(row=row, column=j, value=v)
        row += 1
    for j, c in enumerate(df.columns, start=1):
        width = max([len(str(c))] + [len(str(x)) for x in df[c].head(60)])
        ws.column_dimensions[get_column_letter(j)].width = min(width + 2, 46)
    return row + 2


def write_series_block(ws, start_col, df):
    for j, cname in enumerate(df.columns):
        ws.cell(row=1, column=start_col + j, value=str(cname))
        for i, v in enumerate(df[cname].tolist(), start=2):
            if isinstance(v, (np.floating, float)):
                v = None if pd.isna(v) else round(float(v), 2)
            elif isinstance(v, (np.integer,)):
                v = int(v)
            elif v is not None and not isinstance(v, (int, float)):
                v = str(v)
            ws.cell(row=i, column=start_col + j, value=v)
    return len(df)


def make_line_chart(ws_data, start_col, n_series, n_rows, title, y_title,
                    skip=1):
    ch = LineChart()
    ch.title = title
    data = Reference(ws_data, min_col=start_col + 1,
                     max_col=start_col + n_series,
                     min_row=1, max_row=n_rows + 1)
    cats = Reference(ws_data, min_col=start_col, min_row=2,
                     max_row=n_rows + 1)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    if ch.title is not None:
        ch.title.overlay = False
    ch.legend.position = "r"
    ch.legend.overlay = False
    ch.y_axis.title = y_title
    try:
        ch.y_axis.title.tx.rich.bodyPr = RichTextProperties(
            rot=-5400000, vert="horz")
    except AttributeError:
        pass
    xn = ws_data.cell(row=1, column=start_col).value
    if xn:
        ch.x_axis.title = str(xn)
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.x_axis.tickLblSkip = skip
    ch.x_axis.tickMarkSkip = skip
    ch.layout = Layout(manualLayout=ManualLayout(
        layoutTarget="inner", xMode="edge", yMode="edge",
        x=0.10, y=0.14, w=0.72, h=0.66))
    ch.height = 9.5
    ch.width = 30
    for i, s in enumerate(ch.series):
        s.smooth = False
        s.graphicalProperties.line.width = 20000
        s.graphicalProperties.line.solidFill = SERIES_HEX[i % len(SERIES_HEX)]
    return ch


def norm_kw(s):
    return (str(s).strip().lower()
            .removeprefix("[").removesuffix("]").strip())


def load_cohort():
    paths = glob.glob(os.path.join(ANALYSIS_DIR, "input", EXPORT_GLOB))
    if not paths:
        raise SystemExit(f"UI keyword export not found in input\\ "
                         f"({EXPORT_GLOB})")
    path = sorted(paths)[-1]
    for enc in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            raw = pd.read_csv(path, dtype=str, encoding=enc,
                              sep=None, engine="python")
            if "Keyword" in raw.columns:
                break
        except Exception:
            continue
    else:
        raise SystemExit("could not parse the UI export")
    raw = raw.fillna("")
    coh = raw[raw["Label"].str.contains("RPL Increase", case=False)].copy()
    coh["kw_norm"] = coh["Keyword"].map(norm_kw)
    coh["campaign_id"] = coh["Campaign ID"].str.strip()
    coh = coh[coh["kw_norm"] != ""]
    return coh, os.path.basename(path)


def window_agg(df, d0, d1, label):
    w = df[(df["report_date"] >= d0) & (df["report_date"] < d1)]
    days = (pd.Timestamp(d1) - pd.Timestamp(d0)).days
    g = w.agg(impr=("impressions", "sum"), clicks=("clicks", "sum"),
              cost=("spend", "sum"), leads=("leads", "sum"),
              value=("value", "sum"))
    return derive(g, days, label)


def derive(g, days, label):
    out = {"window": label,
           "impr": g["impr"], "clicks": g["clicks"], "cost": g["cost"],
           "leads": g["leads"], "conv_value": g["value"],
           "CTR_%": 100 * g["clicks"] / g["impr"] if g["impr"] else np.nan,
           "avg_CPC": g["cost"] / g["clicks"] if g["clicks"] else np.nan,
           "CVR_%": 100 * g["leads"] / g["clicks"] if g["clicks"]
           else np.nan,
           "CPL": g["cost"] / g["leads"] if g["leads"] else np.nan,
           "value_per_lead": g["value"] / g["leads"] if g["leads"]
           else np.nan,
           "ROAS": g["value"] / g["cost"] if g["cost"] else np.nan,
           "impr_per_day": g["impr"] / days,
           "clicks_per_day": g["clicks"] / days,
           "cost_per_day": g["cost"] / days,
           "leads_per_day": g["leads"] / days}
    return out


METRICS = ["impr", "clicks", "cost", "leads", "conv_value", "CTR_%",
           "avg_CPC", "CVR_%", "CPL", "value_per_lead", "ROAS",
           "impr_per_day", "clicks_per_day", "cost_per_day",
           "leads_per_day"]


def prepost_table(df, grain_cols, grain_label):
    """df: daily rows; returns one row per grain with pre/post/oct
    metric columns + % change on the per-day and efficiency metrics."""
    rows = []
    for key, g in df.groupby(grain_cols):
        if not isinstance(key, tuple):
            key = (key,)
        rec = dict(zip(grain_cols, key))
        for (d0, d1, lab) in [(PRE0, PRE1, "pre"), (POST0, POST1, "post"),
                              (OCT0, OCT1, "oct")]:
            w = g[(g["report_date"] >= d0) & (g["report_date"] < d1)]
            days = (pd.Timestamp(d1) - pd.Timestamp(d0)).days
            agg = {"impr": w["impressions"].sum(),
                   "clicks": w["clicks"].sum(), "cost": w["spend"].sum(),
                   "leads": w["leads"].sum(), "value": w["value"].sum()}
            dd = derive(agg, days, lab)
            for m in METRICS:
                rec[f"{lab}_{m}"] = dd[m]
        for m in ["impr_per_day", "clicks_per_day", "cost_per_day",
                  "leads_per_day", "avg_CPC", "CPL", "CTR_%", "CVR_%"]:
            a, b = rec.get(f"pre_{m}"), rec.get(f"post_{m}")
            rec[f"chg_{m}_%"] = (100 * (b - a) / a
                                 if a and not pd.isna(a) and a != 0
                                 else np.nan)
        rows.append(rec)
    out = pd.DataFrame(rows)
    out.insert(0, "grain", grain_label)
    return out


def main():
    con = sqlite3.connect(f"file:{DB_PATH}?mode=ro", uri=True)
    coh, exp_name = load_cohort()
    hr(f"COHORT from {exp_name}: {len(coh)} labeled rows, "
       f"{coh['kw_norm'].nunique()} unique keywords, campaigns: "
       f"{sorted(coh['campaign_id'].unique())}")
    print(sorted(coh["kw_norm"].unique()))

    camps = pd.read_sql_query(
        "SELECT campaign_id, campaign_name FROM campaigns "
        "WHERE account_id=?", con, params=(ACCOUNT_ID,))
    camps["cr"] = camps["campaign_name"].str.extract(r"(CR\d+)")
    cr_of = {str(r.campaign_id): r.cr for r in camps.itertuples()}

    # ------------- match cohort to keyword_instances (text+campaign)
    ki = pd.read_sql_query(
        "SELECT id, keyword_text, match_type, campaign_id, "
        "ad_group_name, status FROM keyword_instances "
        "WHERE account_id=?", con, params=(ACCOUNT_ID,))
    ki["kw_norm"] = ki["keyword_text"].map(norm_kw)
    ki["cid_s"] = ki["campaign_id"].astype(str)
    m = ki.merge(coh[["kw_norm", "campaign_id"]].drop_duplicates(),
                 left_on=["kw_norm", "cid_s"],
                 right_on=["kw_norm", "campaign_id"], how="inner",
                 suffixes=("", "_coh"))
    m = m[m["match_type"].str.upper() == "EXACT"]
    unmatched = (set(zip(coh["kw_norm"], coh["campaign_id"]))
                 - set(zip(m["kw_norm"], m["cid_s"])))
    hr(f"DB MATCH: {len(m)} keyword instances matched; "
       f"{len(unmatched)} cohort rows UNMATCHED")
    if unmatched:
        for u in sorted(unmatched):
            print("  unmatched:", u)
    m["cr"] = m["cid_s"].map(cr_of)
    m["is_ethos"] = m["kw_norm"].str.contains("ethos")
    inst_ids = ",".join(str(i) for i in m["id"])

    # ----------------------------------------- daily kpd for cohort
    k = pd.read_sql_query(f"""
        SELECT p.instance_id, p.report_date,
               SUM(p.impressions) AS impressions,
               SUM(p.clicks) AS clicks,
               SUM(p.cost_micros)/1e6 AS spend,
               SUM(p.primary_conv_count) AS leads,
               SUM(p.primary_conv_value) AS value
        FROM keyword_performance_daily p
        WHERE p.instance_id IN ({inst_ids}) AND p.device='ALL'
        GROUP BY 1,2""", con)
    k = k.merge(m[["id", "kw_norm", "cr", "is_ethos"]],
                left_on="instance_id", right_on="id")

    # -------------------------------------------------- 1) presence
    pres = (k.groupby(["kw_norm", "cr"])
            .agg(first_serve=("report_date", "min"),
                 last_serve=("report_date", "max"),
                 lifetime_impr=("impressions", "sum"),
                 lifetime_leads=("leads", "sum")).reset_index())
    pres["present_before_raise"] = np.where(
        pres["first_serve"] < RAISE_DATE, "YES", "NO (post-raise)")
    pre_act = k[(k["report_date"] >= PRE0) & (k["report_date"] < PRE1)] \
        .groupby(["kw_norm", "cr"])["impressions"].sum()
    pres["pre_window_impr"] = pres.set_index(["kw_norm", "cr"]).index.map(
        lambda t: pre_act.get(t, 0))
    pres = pres.sort_values(["kw_norm", "cr"])
    hr("PRESENCE — was each keyword in the account before Jun 2?")
    print(pres.to_string(index=False))
    n_new = (pres.groupby("kw_norm")["present_before_raise"]
             .apply(lambda s: all(v.startswith("NO") for v in s)))
    hr(f"Keywords with NO serving history before the raise: "
       f"{int(n_new.sum())} of {pres['kw_norm'].nunique()}")

    # ------------------------------------------- 2) pre/post tables
    kw_pp = prepost_table(k, ["kw_norm"], "keyword (all campaigns)")
    kw_pp = kw_pp.sort_values("post_cost", ascending=False)
    hr("PRE/POST BY KEYWORD (all campaigns; pre Mar1-Jun1, post "
       "Jun2-Sep30 full, oct MTD)")
    cshow = (["kw_norm"] + [f"pre_{x}" for x in
                            ["cost_per_day", "leads_per_day", "CPL",
                             "avg_CPC", "impr_per_day"]]
             + [f"post_{x}" for x in
                ["cost_per_day", "leads_per_day", "CPL", "avg_CPC",
                 "impr_per_day"]]
             + ["chg_leads_per_day_%", "chg_CPL_%", "chg_avg_CPC_%"])
    print(kw_pp[cshow].to_string(index=False))

    kwc_pp = prepost_table(k, ["kw_norm", "cr"], "keyword x campaign")
    camp_pp = prepost_table(k, ["cr"], "cohort by campaign")
    seg_pp = prepost_table(
        k.assign(seg=np.where(k["is_ethos"], "ethos-brand members",
                              "carrier/competitor members")),
        ["seg"], "cohort segment")
    tot_pp = prepost_table(k.assign(all="COHORT TOTAL"), ["all"],
                           "cohort total")
    hr("COHORT SEGMENTS — carrier terms are the clean read on the pay "
       "raise; the two ethos-quote members also ride the demand surge")
    pcols = (["pre_cost_per_day", "post_cost_per_day",
              "pre_leads_per_day", "post_leads_per_day", "pre_CPL",
              "post_CPL", "pre_avg_CPC", "post_avg_CPC",
              "pre_value_per_lead", "post_value_per_lead"])
    print(seg_pp[["seg"] + pcols].to_string(index=False))
    print(tot_pp[["all"] + pcols].to_string(index=False))

    # ------------------------------------------------ 3) benchmarks
    acct = pd.read_sql_query(f"""
        SELECT p.report_date, SUM(p.impressions) AS impressions,
               SUM(p.clicks) AS clicks, SUM(p.cost_micros)/1e6 AS spend,
               SUM(p.primary_conv_count) AS leads,
               SUM(p.primary_conv_value) AS value
        FROM keyword_performance_daily p
        JOIN keyword_instances ki ON ki.id=p.instance_id
        WHERE ki.account_id={ACCOUNT_ID} AND p.device='ALL'
        GROUP BY 1""", con)
    hc = pd.read_sql_query(f"""
        SELECT p.report_date, SUM(p.impressions) AS impressions,
               SUM(p.clicks) AS clicks, SUM(p.cost_micros)/1e6 AS spend,
               SUM(p.primary_conv_count) AS leads,
               SUM(p.primary_conv_value) AS value
        FROM keyword_performance_daily p
        JOIN keyword_instances ki ON ki.id=p.instance_id
        WHERE ki.account_id={ACCOUNT_ID} AND p.device='ALL'
          AND ki.ad_group_name LIKE '%HighCPC%'
        GROUP BY 1""", con)
    bench = pd.concat([
        prepost_table(acct.assign(b="ACCOUNT (all keywords)"), ["b"],
                      "benchmark"),
        prepost_table(hc.assign(b="HighCPC ad group (all keywords)"),
                      ["b"], "benchmark"),
        tot_pp.rename(columns={"all": "b"})])
    # relative lift: cohort post/pre ratio vs account post/pre ratio
    def ratio(row, mcol):
        a, b = row.get(f"pre_{mcol}"), row.get(f"post_{mcol}")
        return b / a if a and not pd.isna(a) else np.nan
    rel = {}
    for mcol in ["leads_per_day", "cost_per_day", "impr_per_day"]:
        r_acct = ratio(bench.iloc[0], mcol)
        r_coh = ratio(bench.iloc[2], mcol)
        rel[mcol] = (r_coh / r_acct if r_acct else np.nan)
    rel_tbl = pd.DataFrame([
        {"metric": mcol,
         "cohort_post/pre": ratio(bench.iloc[2], mcol),
         "account_post/pre": ratio(bench.iloc[0], mcol),
         "relative_lift_x": rel[mcol]} for mcol in rel])
    hr("BENCHMARKS + RELATIVE LIFT (cohort change vs account change - "
       ">1 means the cohort outgrew the account)")
    print(bench[["b"] + pcols].to_string(index=False))
    print()
    print(rel_tbl.to_string(index=False))

    # ----------------------------------------- QS trail (best-effort)
    qs_tbl = None
    try:
        qcols = pd.read_sql_query(
            "SELECT * FROM keyword_quality_daily LIMIT 1", con).columns
        qs_col = next(c for c in qcols if "quality" in c.lower())
        dt_col = next(c for c in qcols if "date" in c.lower())
        id_col = next(c for c in qcols if "instance" in c.lower())
        qs = pd.read_sql_query(
            f"SELECT {id_col} AS iid, {dt_col} AS d, {qs_col} AS qs "
            f"FROM keyword_quality_daily WHERE {id_col} IN ({inst_ids})",
            con)
        qs = qs.merge(m[["id", "kw_norm"]], left_on="iid", right_on="id")
        qs["win"] = np.where(qs["d"] < PRE1, "pre",
                             np.where(qs["d"] < POST1, "post", "oct"))
        qs_tbl = (qs.pivot_table(index="kw_norm", columns="win",
                                 values="qs", aggfunc="mean")
                  .reindex(columns=["pre", "post", "oct"]).reset_index())
        hr("QUALITY SCORE (mean recorded QS per window, where present)")
        print(qs_tbl.to_string(index=False))
    except Exception as e:
        print(f"\n(QS trail unavailable: {e} - escalate to architect "
              "if schema surprise)")

    # ------------------------------------------------- monthly trail
    k["month"] = k["report_date"].str[:7]
    acct["month"] = acct["report_date"].str[:7]
    trail = (k[k["report_date"] >= TRAIL0].groupby("month")
             .agg(cohort_spend=("spend", "sum"),
                  cohort_leads=("leads", "sum")).reset_index())
    trail["cohort_CPL"] = trail["cohort_spend"] / \
        trail["cohort_leads"].replace(0, np.nan)
    at = (acct[acct["report_date"] >= TRAIL0].groupby("month")
          .agg(acct_spend=("spend", "sum"),
               acct_leads=("leads", "sum")).reset_index())
    at["acct_CPL"] = at["acct_spend"] / at["acct_leads"].replace(0, np.nan)
    trail = trail.merge(at, on="month", how="left")
    hr("MONTHLY TRAIL (Jan -> Oct)")
    print(trail.to_string(index=False))

    # ======================================================== workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Read_Me"
    meta(ws, "RPL INCREASE 6.2.26 — KEYWORD COHORT PRE/POST REVIEW",
         f"cohort pinned from UI export {exp_name} (labels not in DB); "
         "all metrics from paid_search.db (read-only)",
         f"PRE {PRE0}..Jun 1 | POST {POST0}..Sep 30 (full) | Oct MTD separate",
         "Windows differ in length -> compare per-day rates, not "
         "totals; keyword-level impression share is NOT stored in the "
         "DB (the UI export's IS columns are a today-snapshot only); "
         "conv value = SQ payment (quality-blind); label membership "
         "as-of export date; Sep 29-30 (schedule-cut incident days) ARE included in the post window - two low-serving days, immaterial to per-day rates",
         "The question this answers: did the June 2 payout increase "
         "buy us the activity on these priced-out terms that we "
         "negotiated it for - and at what CPL?")
    gl = pd.DataFrame([
        ["Cohort", "Keywords labeled 'High CPC, RPL Increase 6.2.26' "
         "in the UI export - the terms the June payout raise was "
         "negotiated to unlock. Exact match, HighCPC-NA ad group, "
         "all 3 campaigns."],
        ["Presence", "First/last serve per keyword x campaign: answers "
         "whether the keyword existed/served before Jun 2."],
        ["Keyword_PrePost", "Every stored platform metric, pre vs post "
         "vs Oct, totals AND per-day, per keyword (campaigns rolled "
         "up) + keyword x campaign detail."],
        ["Benchmarks", "Account-wide and HighCPC-group averages over "
         "identical windows; relative lift = cohort change / account "
         "change (>1x = outgrew the account). Keywords with no pre "
         "data read against these."],
        ["Segments", "Carrier terms = the clean read on the raise; "
         "the two ethos-quote members also ride the Jul-Sep demand "
         "surge, so they are split out."],
        ["Monthly_Trail", "Jan->Oct monthly spend/leads/CPL, cohort vs "
         "account, + charts."]],
        columns=["sheet/term", "meaning"])
    write_block(ws, 6, "How to read this workbook", gl)

    ws1 = wb.create_sheet("Presence")
    meta(ws1, "PRESENCE — in the account before the raise?",
         "keyword_performance_daily full history", "all history",
         "first_serve = first day with any recorded serving at "
         "device=ALL; a keyword can exist but never serve",
         "Answers 'were these keywords already present before the pay "
         "increase' - YES rows mean the raise changed their economics, "
         "not their existence.")
    write_block(ws1, 6, "Per keyword x campaign", pres)

    ws2 = wb.create_sheet("Keyword_PrePost")
    meta(ws2, "KEYWORD PRE/POST — all stored platform metrics",
         "keyword_performance_daily",
         f"pre {PRE0}..Jun 1 | post Jun 2..Sep 30 | oct MTD",
         "per-day rates are the comparable columns (window lengths "
         "differ); CPL/CPC/CTR/CVR are window aggregates",
         "Sorted by post spend - the top rows are where the raise's "
         "money actually went.")
    r = write_block(ws2, 6, "By keyword (campaigns rolled up)",
                    kw_pp.drop(columns=["grain"]))
    write_block(ws2, r, "By keyword x campaign (detail)",
                kwc_pp.drop(columns=["grain"]))

    ws3 = wb.create_sheet("Benchmarks")
    meta(ws3, "BENCHMARKS + SEGMENTS + RELATIVE LIFT",
         "keyword_performance_daily", "same windows",
         "account benchmark includes the cohort itself (it is small "
         "enough not to dominate); relative lift >1 = cohort outgrew "
         "the account",
         "If relative lift on leads/day > 1, the raise bought real "
         "incremental activity on these terms beyond the account-wide "
         "surge tide.")
    r = write_block(ws3, 6, "Cohort vs account vs HighCPC group",
                    bench.drop(columns=["grain"], errors="ignore"))
    r = write_block(ws3, r, "Relative lift (cohort change / account "
                    "change)", rel_tbl)
    r = write_block(ws3, r, "Cohort segments (ethos members ride the "
                    "surge; carrier members = clean read)",
                    seg_pp.drop(columns=["grain"]))
    r = write_block(ws3, r, "Cohort by campaign",
                    camp_pp.drop(columns=["grain"]))
    if qs_tbl is not None and len(qs_tbl):
        write_block(ws3, r, "Quality Score, mean per window (where "
                    "recorded)", qs_tbl)

    ws4 = wb.create_sheet("Monthly_Trail")
    meta(ws4, "MONTHLY TRAIL — cohort vs account",
         "keyword_performance_daily", "Jan 2026 -> Oct MTD",
         "Oct partial; September is full in both the windows and here",
         "The shape check: did cohort activity step up AT the June 2 "
         "raise, or only at the September budget raises?")
    write_block(ws4, 6, "Monthly", trail)

    wsc = wb.create_sheet("Charts")
    wsd = wb.create_sheet("Chart_Data")
    jn = trail[trail["month"] == "2026-05"]
    base_c = jn["cohort_leads"].iloc[0] if len(jn) else np.nan
    base_a = jn["acct_leads"].iloc[0] if len(jn) else np.nan
    cb = pd.DataFrame({
        "month": trail["month"],
        "Cohort leads (May=100)": 100 * trail["cohort_leads"] / base_c
        if base_c else np.nan,
        "Account leads (May=100)": 100 * trail["acct_leads"] / base_a
        if base_a else np.nan})
    cb2 = trail[["month", "cohort_CPL", "acct_CPL"]].rename(columns={
        "cohort_CPL": "Cohort CPL ($)", "acct_CPL": "Account CPL ($)"})
    col = 1
    n = write_series_block(wsd, col, cb)
    wsc.add_chart(make_line_chart(
        wsd, col, 2, n, "Cohort vs account lead volume (May = 100; "
        "raise = Jun 2)", "Index (May=100)"), "A2")
    col += len(cb.columns) + 1
    n = write_series_block(wsd, col, cb2)
    wsc.add_chart(make_line_chart(
        wsd, col, 2, n, "Cohort vs account CPL by month", "CPL ($)"),
        "A24")
    col += len(cb2.columns) + 1
    meta(wsc, "CHARTS — the June 2 step test",
         "Monthly_Trail", "Jan -> Oct",
         "May = 100 base (last clean pre-raise month); Oct partial",
         "A June/July step in cohort leads that the account does not "
         "show = the raise worked; a step only in September = it was "
         "the budget raises, not the payout.")
    wsd.sheet_state = "hidden"

    os.makedirs(os.path.dirname(OUT_XLSX), exist_ok=True)
    wb.save(OUT_XLSX)
    hr(f"WORKBOOK SAVED -> {OUT_XLSX}")
    con.close()


if __name__ == "__main__":
    main()
