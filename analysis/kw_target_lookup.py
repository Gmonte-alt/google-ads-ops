"""
kw_target_lookup.py
====================
Looks up all keyword instances whose text CONTAINS any of the brand
terms in the lookup file, across a fixed set of accounts and a date
range. Useful for competitor brand research — finds every keyword
across the specified accounts that mentions a target brand.

Match logic: substring match on normalized keyword text (case-insensitive),
with a configurable false-positive exclusion list for ambiguous brands.

ACCOUNTS (multi-platform)
─────────────────────────
  google     3863660890  Auto MCC (autoinsurance.org)
  google     1006729211  Auto Quote (auto.quote.com)
  google     8886502147  Home Insurance
  microsoft  F149Y8EN    Auto Insurance
  microsoft  F149SN63    Home Insurance

INPUT FILE FORMAT
─────────────────
Plain CSV (no header) — one brand term per line.

OUTPUT
──────
Excel workbook with:
  Sheet 1 — All matched instances with performance data + campaign/ad group status
  Sheet 2 — Summary by brand × platform × account
  Sheet 3 — Review (short/generic terms flagged for false-positive risk)
  Sheet 4 — Unmatched brand terms

RUN
───
    conda activate paid-search
    python analysis/kw_target_lookup.py
    python analysis/kw_target_lookup.py --start 2026-03-29 --end 2026-04-27
    python analysis/kw_target_lookup.py --csv path/to/brands.csv
"""

import sqlite3
import argparse
import csv
from pathlib import Path
from datetime import date, timedelta

import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

DB_PATH    = Path(r"C:\Projects\google-ads-ops\db\paid_search.db")
OUTPUT_DIR = Path(r"C:\Projects\google-ads-ops\reports")
DEFAULT_CSV = Path(r"C:\Projects\google-ads-ops\analysis\input\confie-list-of-subsidiearies.csv")

# Hardcoded scope: 5 accounts across Google and Microsoft.
TARGET_ACCOUNTS = [
    ("google",    "3863660890"),  # Auto MCC (autoinsurance.org)
    ("google",    "1006729211"),  # Auto Quote (auto.quote.com)
    ("google",    "8886502147"),  # Home Insurance
    ("microsoft", "F149Y8EN"),    # Auto Insurance
    ("microsoft", "F149SN63"),    # Home Insurance
]

# Brand-level false-positive exclusions.
# Key   = brand term (must match an entry in the CSV exactly, lowercase)
# Value = list of substrings that, if found in the keyword's normalized
#         text, disqualify the match. Add to this dict over time as new
#         false positives are spotted.
BRAND_FALSE_POSITIVES = {
    "family auto": ["american family"],   # American Family Auto = unrelated insurer
}

# Default to last 30 complete days
END_DEFAULT   = (date.today() - timedelta(days=1)).isoformat()
START_DEFAULT = (date.today() - timedelta(days=30)).isoformat()

# Review-sheet heuristics
REVIEW_MAX_WORDS = 1
REVIEW_MAX_CHARS = 6

HEADER_FILL = "7030A0"
HEADER_FONT = "FFFFFF"


# ── DB connection ──────────────────────────────────────────────────────────────

def get_conn():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    conn.execute("PRAGMA journal_mode = WAL")
    return conn


# ── Load brand list from CSV ──────────────────────────────────────────────────

def load_brand_list(csv_path: Path) -> list[str]:
    if not csv_path.exists():
        raise FileNotFoundError(f"Brand list not found: {csv_path}")

    seen   = set()
    brands = []
    with open(csv_path, newline="", encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        for row in reader:
            if not row:
                continue
            term = row[0].strip().lower()
            if not term or term in seen:
                continue
            seen.add(term)
            brands.append(term)

    print(f"  Loaded {len(brands)} unique brand terms from {csv_path.name}")
    return brands


# ── Resolve instances via substring match ─────────────────────────────────────

def resolve_instances(conn, brands: list[str],
                      target_accounts: list[tuple[str, str]]
                      ) -> tuple[pd.DataFrame, list[str], int]:
    """
    For each (account, brand) pair, returns ALL keyword_instances whose
    normalized text CONTAINS the brand term — minus any false-positive
    exclusions defined in BRAND_FALSE_POSITIVES.

    JOINs campaigns and ad_groups for status fields. LEFT JOIN ensures
    keywords whose campaign/ad_group rows are missing still appear in
    output (with NULL status) rather than being silently dropped.
    """
    placeholders = ",".join(["(?, ?)"] * len(target_accounts))
    flat_params  = [v for pair in target_accounts for v in pair]

    acct_rows = conn.execute(
        f"""
        SELECT id, platform, platform_account_id, account_name
        FROM accounts
        WHERE (platform, platform_account_id) IN ({placeholders})
        """,
        flat_params,
    ).fetchall()

    if len(acct_rows) != len(target_accounts):
        found   = {(r["platform"], str(r["platform_account_id"])) for r in acct_rows}
        missing = [t for t in target_accounts if t not in found]
        print(f"  WARNING — accounts not found in DB: {missing}")

    acct_map = {
        (r["platform"], str(r["platform_account_id"])): {
            "account_id":   r["id"],
            "account_name": r["account_name"],
            "platform":     r["platform"],
        }
        for r in acct_rows
    }

    print(f"  Searching {len(acct_map)} accounts × {len(brands)} brands "
          f"= {len(acct_map) * len(brands)} sub-queries")

    results            = []
    matched_brands     = set()
    excluded_count     = 0
    excluded_examples  = {}  # brand -> list of dropped normalized_texts

    for (platform, pid), acct in acct_map.items():
        account_id   = acct["account_id"]
        account_name = acct["account_name"]

        for brand in brands:
            pattern    = f"%{brand}%"
            exclusions = BRAND_FALSE_POSITIVES.get(brand, [])

            ki_rows = conn.execute("""
                SELECT
                    ki.id            AS instance_id,
                    ki.criterion_id,
                    ki.ad_group_id,
                    ki.campaign_id,
                    ki.keyword_text,
                    ki.match_type,
                    ki.campaign_name,
                    ki.ad_group_name,
                    ki.status        AS kw_status,
                    ki_id.normalized_text,
                    c.status         AS campaign_status,
                    ag.status        AS ad_group_status
                FROM keyword_instances ki
                JOIN keyword_identities ki_id ON ki_id.id = ki.identity_id
                LEFT JOIN campaigns c
                       ON c.account_id = ki.account_id
                      AND c.campaign_id = ki.campaign_id
                LEFT JOIN ad_groups ag
                       ON ag.account_id = ki.account_id
                      AND ag.ad_group_id = ki.ad_group_id
                WHERE ki.account_id = ?
                  AND LOWER(ki_id.normalized_text) LIKE ?
            """, (account_id, pattern)).fetchall()

            for ki in ki_rows:
                ntext = (ki["normalized_text"] or "").lower()

                # Apply false-positive exclusions
                if exclusions and any(ex in ntext for ex in exclusions):
                    excluded_count += 1
                    excluded_examples.setdefault(brand, set()).add(ntext)
                    continue

                matched_brands.add(brand)

                results.append({
                    "platform":         platform,
                    "account_id_str":   pid,
                    "account_name":     account_name,
                    "account_id":       account_id,
                    "matched_brand":    brand,
                    "instance_id":      ki["instance_id"],
                    "criterion_id":     ki["criterion_id"],
                    "campaign_id":      ki["campaign_id"],
                    "campaign_name":    ki["campaign_name"],
                    "campaign_status":  ki["campaign_status"],
                    "ad_group_id":      ki["ad_group_id"],
                    "ad_group_name":    ki["ad_group_name"],
                    "ad_group_status":  ki["ad_group_status"],
                    "keyword_text":     ki["keyword_text"],
                    "match_type":       ki["match_type"],
                    "kw_status":        ki["kw_status"],
                    "normalized_text":  ki["normalized_text"],
                })

    unmatched_brands = [b for b in brands if b not in matched_brands]

    print(f"  Instance hits     : {len(results):,}")
    print(f"  Brands with hits  : {len(matched_brands)}")
    print(f"  Brands w/o hits   : {len(unmatched_brands)}")
    if excluded_count:
        print(f"  False-pos excluded: {excluded_count} rows")
        for brand, texts in excluded_examples.items():
            sample = ", ".join(sorted(texts)[:3])
            more   = f" (+{len(texts)-3} more)" if len(texts) > 3 else ""
            print(f"    '{brand}' excluded: {sample}{more}")

    return pd.DataFrame(results), unmatched_brands, excluded_count


# ── Fetch performance ─────────────────────────────────────────────────────────

def fetch_performance(conn, instance_ids: list, start: str, end: str) -> pd.DataFrame:
    """
    Pulls performance for the given instance_ids using the device='ALL'
    rollup, which is consistent across both Google and Microsoft (verified
    2026-04-28 — Microsoft's ALL rows are a true rollup of per-device rows).
    """
    if not instance_ids:
        return pd.DataFrame()

    all_rows = []
    chunk_size = 500
    for i in range(0, len(instance_ids), chunk_size):
        chunk = instance_ids[i:i + chunk_size]
        ph    = ",".join("?" * len(chunk))
        rows  = conn.execute(f"""
            SELECT
                kpd.instance_id,
                SUM(kpd.impressions)          AS impressions,
                SUM(kpd.clicks)               AS clicks,
                SUM(kpd.cost_micros) / 1e6    AS spend,
                SUM(kpd.primary_conv_count)   AS conversions,
                SUM(kpd.primary_conv_value)   AS revenue
            FROM keyword_performance_daily kpd
            WHERE kpd.instance_id IN ({ph})
              AND kpd.device      = 'ALL'
              AND kpd.geo_state   = 'ALL'
              AND kpd.report_date BETWEEN ? AND ?
            GROUP BY kpd.instance_id
        """, chunk + [start, end]).fetchall()
        all_rows.extend(rows)

    if not all_rows:
        return pd.DataFrame(columns=[
            "instance_id","impressions","clicks","spend","conversions","revenue"
        ])
    return pd.DataFrame([dict(r) for r in all_rows])


# ── Review-sheet heuristic ────────────────────────────────────────────────────

def build_review_df(results_df: pd.DataFrame, brands: list[str]) -> pd.DataFrame:
    rows = []
    for brand in brands:
        word_count = len(brand.split())
        char_count = len(brand)
        if word_count > REVIEW_MAX_WORDS and char_count > REVIEW_MAX_CHARS:
            continue

        sub = results_df[results_df["matched_brand"] == brand]
        if sub.empty:
            sample_texts = ""
            n_matches    = 0
            n_unique     = 0
        else:
            unique_texts = (sub["normalized_text"]
                            .drop_duplicates()
                            .head(8)
                            .tolist())
            sample_texts = " | ".join(unique_texts)
            n_matches    = len(sub)
            n_unique     = sub["normalized_text"].nunique()

        rows.append({
            "brand_term":         brand,
            "word_count":         word_count,
            "char_count":         char_count,
            "instance_matches":   n_matches,
            "unique_kw_texts":    n_unique,
            "exclusions_applied": ", ".join(BRAND_FALSE_POSITIVES.get(brand, [])),
            "sample_kw_texts":    sample_texts,
        })

    df = pd.DataFrame(rows)
    if not df.empty:
        df = df.sort_values(["instance_matches", "brand_term"],
                            ascending=[False, True]).reset_index(drop=True)
    return df


# ── Excel output ──────────────────────────────────────────────────────────────

def write_workbook(out_path: Path, results_df: pd.DataFrame,
                   summary_df: pd.DataFrame, review_df: pd.DataFrame,
                   unmatched_brands: list[str], start: str, end: str):

    wb = openpyxl.Workbook()
    hfill = PatternFill("solid", fgColor=HEADER_FILL)
    hfont = Font(color=HEADER_FONT, bold=True)

    def write_sheet(ws, df, title=None):
        if title:
            ws.title = title
        if df is None or df.empty:
            ws.append(["No data"])
            return
        for ci, col in enumerate(df.columns, 1):
            cell = ws.cell(row=1, column=ci, value=col)
            cell.fill = hfill
            cell.font = hfont
            cell.alignment = Alignment(horizontal="center")
        for ri, row in enumerate(df.itertuples(index=False), 2):
            for ci, val in enumerate(row, 1):
                ws.cell(row=ri, column=ci, value=val)
        for ci, col in enumerate(df.columns, 1):
            max_len = max(
                len(str(col)),
                df.iloc[:, ci-1].astype(str).str.len().max() if len(df) > 0 else 0
            )
            ws.column_dimensions[get_column_letter(ci)].width = min(max_len + 2, 60)
        ws.freeze_panes = ws.cell(row=2, column=1)

    ws1 = wb.active
    write_sheet(ws1, results_df, title=f"Instances {start} to {end}")

    ws2 = wb.create_sheet(title="Summary by Brand")
    write_sheet(ws2, summary_df)

    ws3 = wb.create_sheet(title="Review (FP risk)")
    if review_df is None or review_df.empty:
        ws3.append(["No short or single-word brand terms in lookup."])
    else:
        write_sheet(ws3, review_df)

    ws4 = wb.create_sheet(title="Unmatched Brands")
    if not unmatched_brands:
        ws4.append(["All brand terms matched at least one keyword."])
    else:
        ws4.append(["brand_term"])
        ws4["A1"].fill = hfill
        ws4["A1"].font = hfont
        for b in unmatched_brands:
            ws4.append([b])
        ws4.column_dimensions["A"].width = 40

    wb.save(out_path)
    print(f"\n  Saved: {out_path}")


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    p = argparse.ArgumentParser()
    p.add_argument("--csv",   default=str(DEFAULT_CSV))
    p.add_argument("--start", default=START_DEFAULT)
    p.add_argument("--end",   default=END_DEFAULT)
    args = p.parse_args()

    start = args.start
    end   = args.end

    print("=" * 70)
    print("kw_target_lookup.py  [substring match — multi-platform brand scan]")
    print(f"  Database  : {DB_PATH}")
    print(f"  Brand list: {Path(args.csv).name}")
    print(f"  Window    : {start} → {end}")
    print(f"  Accounts  : {len(TARGET_ACCOUNTS)} across "
          f"{len({p for p,_ in TARGET_ACCOUNTS})} platforms")
    for plat, pid in TARGET_ACCOUNTS:
        print(f"              {plat:<10} {pid}")
    if BRAND_FALSE_POSITIVES:
        print(f"  FP filters: {len(BRAND_FALSE_POSITIVES)} brand(s) with exclusions")
        for brand, exclusions in BRAND_FALSE_POSITIVES.items():
            print(f"              '{brand}' excludes: {', '.join(exclusions)}")
    print("=" * 70)

    conn = get_conn()

    print("\n[1] Loading brand list...")
    brands = load_brand_list(Path(args.csv))

    print("\n[2] Resolving instances (substring match)...")
    instances_df, unmatched_brands, excluded_count = resolve_instances(
        conn, brands, TARGET_ACCOUNTS
    )

    if instances_df.empty:
        print("  No instances matched any brand term.")
        conn.close()
        return

    print(f"\n[3] Fetching performance ({start} → {end})...")
    iids    = instances_df["instance_id"].drop_duplicates().tolist()
    perf_df = fetch_performance(conn, iids, start, end)
    print(f"  Performance rows: {len(perf_df):,}")

    results_df = instances_df.merge(perf_df, on="instance_id", how="left")
    for col in ["impressions","clicks","spend","conversions","revenue"]:
        results_df[col] = results_df[col].fillna(0)

    results_df["ctr"]  = (results_df["clicks"] / results_df["impressions"]
                          ).where(results_df["impressions"] > 0).round(4)
    results_df["cpc"]  = (results_df["spend"] / results_df["clicks"]
                          ).where(results_df["clicks"] > 0).round(2)
    results_df["cpa"]  = (results_df["spend"] / results_df["conversions"]
                          ).where(results_df["conversions"] > 0).round(2)
    results_df["roas"] = (results_df["revenue"] / results_df["spend"]
                          ).where(results_df["spend"] > 0).round(4)
    results_df["spend"]   = results_df["spend"].round(2)
    results_df["revenue"] = results_df["revenue"].round(2)

    out_cols = [
        "platform","account_id_str","account_name",
        "matched_brand",
        "normalized_text","keyword_text","match_type","kw_status",
        "campaign_name","campaign_status","campaign_id",
        "ad_group_name","ad_group_status","ad_group_id",
        "instance_id","criterion_id",
        "impressions","clicks","spend","conversions","revenue",
        "ctr","cpc","cpa","roas",
    ]
    results_df = results_df[[c for c in out_cols if c in results_df.columns]]

    summary_df = (
        results_df.groupby(
            ["platform","account_id_str","account_name","matched_brand"],
            as_index=False)
        .agg(
            instance_matches=("instance_id","count"),
            unique_kw_texts=("normalized_text","nunique"),
            enabled_instances=("kw_status", lambda x: (x=="ENABLED").sum()),
            impressions=("impressions","sum"),
            clicks=("clicks","sum"),
            spend=("spend","sum"),
            conversions=("conversions","sum"),
            revenue=("revenue","sum"),
        )
        .sort_values(["platform","account_name","spend"],
                     ascending=[True, True, False])
        .reset_index(drop=True)
    )
    summary_df["spend"]   = summary_df["spend"].round(2)
    summary_df["revenue"] = summary_df["revenue"].round(2)

    review_df = build_review_df(results_df, brands)

    unique_results = results_df.drop_duplicates(subset=["instance_id"])
    print(f"\n[4] Results:")
    print(f"  Unique instances matched   : {len(unique_results):,}")
    print(f"  Total instance × brand rows: {len(results_df):,}  "
          f"(higher because same kw can match multiple brands)")
    if excluded_count:
        print(f"  False-positive rows excluded: {excluded_count}")

    print(f"\n  By platform:")
    plat_totals = (unique_results.groupby("platform", as_index=False)
                   .agg(instances=("instance_id","count"),
                        spend=("spend","sum"),
                        revenue=("revenue","sum"))
                   .sort_values("platform"))
    for _, r in plat_totals.iterrows():
        print(f"    {r['platform']:<10} {r['instances']:>6,} instances  "
              f"spend ${r['spend']:>12,.2f}  rev ${r['revenue']:>12,.2f}")

    print(f"\n  Overall:")
    print(f"    ENABLED instances        : "
          f"{(unique_results['kw_status']=='ENABLED').sum():,}")
    print(f"    PAUSED instances         : "
          f"{(unique_results['kw_status']=='PAUSED').sum():,}")
    print(f"    Total spend (unique)     : ${unique_results['spend'].sum():,.2f}")
    print(f"    Total clicks (unique)    : {unique_results['clicks'].sum():,.0f}")
    print(f"    Total revenue (unique)   : ${unique_results['revenue'].sum():,.2f}")

    print(f"\n  Top 5 brands by spend (across platforms):")
    top = (summary_df.groupby("matched_brand", as_index=False)
           .agg(spend=("spend","sum"),
                instance_matches=("instance_matches","sum"))
           .sort_values("spend", ascending=False).head(5))
    for _, r in top.iterrows():
        print(f"    {r['matched_brand']:<30}  "
              f"${r['spend']:>12,.2f}  {r['instance_matches']} matches")

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    out_path = OUTPUT_DIR / f"kw_target_brand_lookup_{start}_to_{end}.xlsx"
    write_workbook(out_path, results_df, summary_df, review_df,
                   unmatched_brands, start, end)

    conn.close()
    print("Done.")


if __name__ == "__main__":
    main()