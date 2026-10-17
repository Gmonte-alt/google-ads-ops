"""
ms_expansion_absorption.py -- SERVED / BROAD-ABSORPTION check for the MS Home
expansion. Read-only. Resolves the pending_served_check flag on every proposed
target using the ms_search_terms staging table (UI Search Term Report backfill,
Jan-Aug 2026, 100% search network).

Per the analyst's guidance, this analysis is SERVED-SIDE ONLY:
  - impressions / clicks / whether-and-how a term served / which keyword absorbed it
  - the Conversions-column semantics are unconfirmed -> NO RPL, NO lead counts here
  - liveness of the absorbing keyword = TEXT join of the row's normalized keyword
    (09-canonical form) against the live book from ms_expansion_coverage's loader
    (keyword_identities -> keyword_current.is_serving) -- no criterion-ID join
  - "absent across all months" = VERY-LIKELY-not-absorbed (UI exports still bucket
    ultra-low-volume terms into 'other'), not a hard guarantee

Verdicts per proposed target:
  NOT_SERVED           no row matches the term or its close variants -> net-new reach
  VARIANT_SERVED_ONLY  only variant-signature relatives served -> effectively net-new;
                       the add mostly claims priority over loose matching
  ABSORBED_BY_LIVE     the query serves today via >=1 keyword that is LIVE right now
                       -> the add is a PRIORITY/CONTROL play, expect share shift more
                       than incremental reach
  ABSORBED_HISTORICAL  it served, but no absorbing keyword is in the live book now
                       -> currently unabsorbed; add is (re)claiming known demand

Plus the BUNDLE INTENT section: aggregate served volume for (auto|car|vehicle) x
(home|house|homeowner) queries and which keywords absorb them today -- the
first-order answer to whether live broad already catches bundle intent.

    python ms_expansion_absorption.py --keywords-csv output\\ms_expansion_selection_YYYY-MM-DD_keywords.csv
    python ms_expansion_absorption.py --keywords-csv ...csv --inspect-schema
"""
import argparse
import csv
import os
import re
import sqlite3
import sys
from collections import defaultdict
from datetime import date
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")
sys.path.insert(0, str(Path(__file__).resolve().parent))
from ms_expansion_coverage import normalize_09, normalize_loose, load_structure
from ms_expansion_select import variant_key

TABLE = "ms_search_terms"
# column-name candidates, resolved by PRAGMA (introspect, don't assume)
CANDS = {
    "term":     ["normalized_search_term", "search_term", "search_query", "query_text", "term"],
    "kw":       ["normalized_keyword", "keyword_text", "keyword"],
    "kw_match": ["match_type", "keyword_match_type", "delivered_match_type", "match"],
    "month":    ["report_month", "month", "report_date", "date"],
    "impr":     ["impressions", "impr"],
    "clicks":   ["clicks"],
    "campaign": ["campaign_name", "campaign"],
    "ad_group": ["ad_group_name", "ad_group"],
}
BUNDLE_A = r"(auto|car|cars|vehicle|vehicles|automobile)"
BUNDLE_H = r"(home|homes|house|houses|homeowner|homeowners)"


def cols_of(conn, table):
    try:
        return [r[1] for r in conn.execute("PRAGMA table_info(%s)" % table)]
    except sqlite3.OperationalError:
        return []


def resolve(conn):
    present = cols_of(conn, TABLE)
    if not present:
        sys.exit("ERROR: table '%s' not found. Has the backfill run (drop --dry-run)?" % TABLE)
    picked, missing = {}, []
    for key, cands in CANDS.items():
        picked[key] = next((c for c in cands if c in present), None)
        if picked[key] is None and key in ("term", "kw", "impr", "clicks", "month"):
            missing.append("%s (tried %s)" % (key, "/".join(cands)))
    if missing:
        sys.exit("SCHEMA ERROR in %s: cannot resolve %s.\nPresent columns: %s\n"
                 "Run --inspect-schema and adjust CANDS." % (TABLE, "; ".join(missing), present))
    return picked, present


def load_served(conn, c):
    """term-norm -> {impr, clicks, months set, absorbers: {(kw_norm, match): [impr, clicks]}}
    Also a variant-signature index for close-variant lookups. Conversions untouched."""
    sel = [c["term"], c["kw"], c["month"], c["impr"], c["clicks"]]
    sel.append(c["kw_match"] if c["kw_match"] else "''")
    q = "SELECT %s FROM %s" % (", ".join(sel), TABLE)
    served = {}
    vk_index = defaultdict(set)          # variant_key -> {term_norm}
    n = 0
    for term, kw, month, impr, clicks, mt in conn.execute(q):
        n += 1
        tn = normalize_09(term or "")
        if not tn:
            continue
        s = served.setdefault(tn, {"impr": 0, "clicks": 0, "months": set(),
                                   "absorbers": defaultdict(lambda: [0, 0])})
        s["impr"] += impr or 0
        s["clicks"] += clicks or 0
        s["months"].add(str(month)[:7])
        a = s["absorbers"][(normalize_09(kw or ""), (mt or "").lower())]
        a[0] += impr or 0
        a[1] += clicks or 0
        vk_index[variant_key(normalize_loose(term or ""))].add(tn)
    return served, vk_index, n


def live_text_set(by_canon):
    live = set()
    for canon, rows in by_canon.items():
        if any(r["is_serving"] == 1 for r in rows):
            live.add(canon)
    return live


def check_term(kw_loose, kw_n9, served, vk_index, live):
    rec = served.get(kw_n9) or served.get(kw_loose)
    variants = []
    if rec is None:
        for tn in vk_index.get(variant_key(kw_loose), set()):
            if tn not in (kw_n9, kw_loose):
                variants.append(tn)
        if not variants:
            return "NOT_SERVED", None, [], []
        vrec_terms = variants[:5]
        return "VARIANT_SERVED_ONLY", None, [], vrec_terms
    absorbers = sorted(rec["absorbers"].items(), key=lambda kv: -kv[1][1])
    live_abs = [(k, v) for k, v in absorbers if k[0] in live]
    verdict = "ABSORBED_BY_LIVE" if live_abs else "ABSORBED_HISTORICAL"
    return verdict, rec, absorbers, []


GUIDANCE = {
    "NOT_SERVED": "ADD = net-new reach (least-suppressed source; still not a hard guarantee)",
    "VARIANT_SERVED_ONLY": "ADD = effectively net-new; claims priority over loose matching",
    "ABSORBED_BY_LIVE": "ADD = priority/control play; expect share shift from the absorber, not pure increment",
    "ABSORBED_HISTORICAL": "served before, no live absorber now -> add (re)claims known demand",
}


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--keywords-csv", required=True,
                   help="the selection output *_keywords.csv (post keep/kill edits welcome)")
    p.add_argument("--db", default=r"C:\Projects\google-ads-ops\db\paid_search.db")
    p.add_argument("--inspect-schema", action="store_true")
    p.add_argument("--out", default=r"C:\Projects\google-ads-ops\analysis\ms-home-kw-expansion"
                                    r"\output\ms_expansion_absorption_%s.xlsx" % date.today().isoformat())
    a = p.parse_args()

    conn = sqlite3.connect("file:%s?mode=ro" % a.db, uri=True)
    conn.row_factory = sqlite3.Row
    if a.inspect_schema:
        print("%s columns: %s" % (TABLE, cols_of(conn, TABLE)))
        return
    c, present = resolve(conn)
    print("schema resolved: term=%s kw=%s match=%s month=%s (conversions columns present "
          "but DELIBERATELY unused -- semantics unconfirmed)"
          % (c["term"], c["kw"], c["kw_match"] or "(absent)", c["month"]))

    served, vk_index, n_rows = load_served(conn, c)
    months = sorted({m for s in served.values() for m in s["months"]})
    print("%s: %s rows -> %s distinct served terms | months %s..%s"
          % (TABLE, f"{n_rows:,}", f"{len(served):,}", months[0], months[-1]))
    by_canon, by_loose, n_struct, basis = load_structure(conn)
    live = live_text_set(by_canon)
    print("live book: %s serving canonical texts (%s)" % (f"{len(live):,}", basis.split(';')[0]))

    # ---- per-target check ----
    rows_by_term, n_killed_rows = {}, 0
    with open(a.keywords_csv, newline="", encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            kw = (row.get("keyword") or "").strip()
            if not kw:
                continue
            lo = normalize_loose(kw)
            killed = (row.get("decision") or "").strip().upper().startswith("KILL")
            n_killed_rows += killed
            t = rows_by_term.setdefault(lo, {"kw": kw, "loose": lo, "n9": normalize_09(kw),
                                             "ad_group": (row.get("ad_group_override") or "").strip()
                                                         or row.get("ad_group"),
                                             "module": row.get("module"), "kept": False})
            if not killed:
                t["kept"] = True      # a term is checked if ANY of its rows survives
    targets = [t for t in rows_by_term.values() if t["kept"]]
    print("targets to check: %d of %d distinct terms (%d killed rows via decision column)\n"
          % (len(targets), len(rows_by_term), n_killed_rows))

    out_rows, dist = [], defaultdict(int)
    for t in targets:
        verdict, rec, absorbers, variants = check_term(t["loose"], t["n9"], served, vk_index, live)
        dist[verdict] += 1
        top = "; ".join("%s (%s)%s impr=%s clk=%s"
                        % (k[0], k[1] or "?", " LIVE" if k[0] in live else "",
                           f"{v[0]:,}", f"{v[1]:,}")
                        for k, v in absorbers[:3])
        out_rows.append({"keyword": t["kw"], "ad_group": t["ad_group"], "module": t["module"],
                         "served_verdict": verdict,
                         "served_impr": rec["impr"] if rec else 0,
                         "served_clicks": rec["clicks"] if rec else 0,
                         "months_served": len(rec["months"]) if rec else 0,
                         "last_month": max(rec["months"]) if rec else "",
                         "top_absorbers": top,
                         "variant_terms_served": "; ".join(variants),
                         "guidance": GUIDANCE[verdict]})

    order = {"ABSORBED_BY_LIVE": 0, "ABSORBED_HISTORICAL": 1, "VARIANT_SERVED_ONLY": 2, "NOT_SERVED": 3}
    out_rows.sort(key=lambda x: (order[x["served_verdict"]], -x["served_clicks"]))

    print("=" * 78)
    print("  SERVED-CHECK VERDICTS (%d targets, %s..%s)" % (len(targets), months[0], months[-1]))
    print("=" * 78)
    for v in sorted(dist, key=lambda x: order[x]):
        print("  %-22s %d   %s" % (v, dist[v], GUIDANCE[v]))
    print("\n  absorbed targets (who's catching them today):")
    for r in out_rows:
        if r["served_verdict"].startswith("ABSORBED"):
            print("    %-40s impr=%-8s clk=%-6s %s"
                  % (r["keyword"], f"{r['served_impr']:,}", f"{r['served_clicks']:,}",
                     r["top_absorbers"][:70]))

    # ---- bundle intent aggregate ----
    A, H = re.compile(r"\b%s\b" % BUNDLE_A), re.compile(r"\b%s\b" % BUNDLE_H)
    b_impr = b_clicks = 0
    b_abs = defaultdict(lambda: [0, 0])
    b_terms = []
    for tn, s in served.items():
        if A.search(tn) and H.search(tn):
            b_impr += s["impr"]; b_clicks += s["clicks"]
            b_terms.append((tn, s["impr"], s["clicks"]))
            for k, v in s["absorbers"].items():
                b_abs[k][0] += v[0]; b_abs[k][1] += v[1]
    b_terms.sort(key=lambda x: -x[2])
    print("\n" + "=" * 78)
    print("  BUNDLE INTENT -- (auto|car|vehicle) x (home|house|homeowner) served queries")
    print("=" * 78)
    print("  served: %s impressions / %s clicks across %d distinct queries"
          % (f"{b_impr:,}", f"{b_clicks:,}", len(b_terms)))
    print("  top absorbing keywords:")
    for k, v in sorted(b_abs.items(), key=lambda kv: -kv[1][1])[:8]:
        print("    %-44s (%s)%s impr=%-9s clk=%s"
              % (k[0], k[1] or "?", " LIVE" if k[0] in live else "",
                 f"{v[0]:,}", f"{v[1]:,}"))

    # ---- workbook ----
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    HEADF = PatternFill("solid", fgColor="1F3864")
    WHITE = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    BOLD = Font(name="Arial", bold=True, size=10)
    REG = Font(name="Arial", size=10)
    NOTE = Font(name="Arial", size=9, italic=True, color="595959")
    TITLE = Font(name="Arial", bold=True, size=13, color="1F3864")

    def sheet_from(wb, title, heading, cols, dicts, widths, note=None):
        ws = wb.create_sheet(title)
        ws["A1"] = heading; ws["A1"].font = TITLE
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
        r0 = 3
        if note:
            ws["A2"] = note; ws["A2"].font = NOTE
            ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(cols))
            r0 = 4
        for j, h in enumerate(cols, 1):
            cell = ws.cell(r0, j, h); cell.fill = HEADF; cell.font = WHITE
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        r = r0 + 1
        for d in dicts:
            for j, k in enumerate(cols, 1):
                ws.cell(r, j, d.get(k)).font = BOLD if j == 1 else REG
            r += 1
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.freeze_panes = ws.cell(r0 + 1, 1).coordinate
        ws.auto_filter.ref = "A%d:%s%d" % (r0, get_column_letter(len(cols)), max(r - 1, r0))

    wb = Workbook(); wb.remove(wb.active)
    kcols = ["keyword", "ad_group", "module", "served_verdict", "served_impr", "served_clicks",
             "months_served", "last_month", "top_absorbers", "variant_terms_served", "guidance"]
    sheet_from(wb, "Served Check", "Per-target served / absorption verdicts (%s..%s)"
               % (months[0], months[-1]), kcols, out_rows,
               [34, 38, 8, 20, 10, 10, 8, 10, 70, 40, 56],
               note="Served-side only: conversions columns unused (semantics unconfirmed). "
                    "'NOT_SERVED' = very-likely-not-absorbed (UI export still buckets ultra-low "
                    "volume into 'other'), not a hard guarantee.")
    brows = [{"query": tn, "impr": im, "clicks": cl} for tn, im, cl in b_terms[:100]]
    sheet_from(wb, "Bundle Intent", "Bundle queries served Jan-Aug (top 100 by clicks): "
               "%s impr / %s clicks total" % (f"{b_impr:,}", f"{b_clicks:,}"),
               ["query", "impr", "clicks"], brows, [50, 12, 10])
    arows = [{"absorbing_keyword": k[0], "match": k[1] or "?",
              "live_now": "LIVE" if k[0] in live else "not live",
              "impr": v[0], "clicks": v[1]}
             for k, v in sorted(b_abs.items(), key=lambda kv: -kv[1][1])[:50]]
    sheet_from(wb, "Bundle Absorbers", "Keywords absorbing bundle intent (top 50 by clicks)",
               ["absorbing_keyword", "match", "live_now", "impr", "clicks"],
               arows, [46, 10, 10, 12, 10])
    os.makedirs(os.path.dirname(a.out) or ".", exist_ok=True)
    wb.save(a.out)

    base = os.path.splitext(a.keywords_csv)[0] + "_served.csv"
    with open(base, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=kcols)
        w.writeheader(); w.writerows(out_rows)
    print("\nwrote %s and %s" % (a.out, base))
    conn.close()


if __name__ == "__main__":
    main()
