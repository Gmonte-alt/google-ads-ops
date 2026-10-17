"""
geico_change_export.py -- Aggregated Geico change-history export for a window (default
Jul 1 - Oct 31 2025), described at the GENERAL level across campaigns (not campaign-by-campaign).

"Geico changes" = TWO things, both captured and tagged:
  * geico_adgroup  -- any change to a Geico ad group (ad_group_col contains 'geico')
  * geico_in_change-- 'geico' appears in the change TEXT but the ad group isn't a Geico one
                      (chiefly NEGATIVE KEYWORDS containing geico, added to generic ad groups
                      to fence geico traffic out) -- the user explicitly wants these included.

Reads google_change_csv_history_raw, splits each cell into blocks, categorizes, keeps only
Geico-relevant blocks, and rolls them up by (context, category, action) with counts, distinct
campaigns/ad-groups touched, item totals, date span, and a representative sample.

Output: an Excel workbook (Summary rollup + a prose paragraph + full detail sheet).

    python geico_change_export.py --db PATH                       # Jul-Oct 2025, brand Geico
    python geico_change_export.py --db PATH --start 2025-07-01 --end 2025-10-31 --out ./output
"""
import argparse
import collections
import os
import sqlite3
import sys

CHANGE_HISTORY_DIR = os.environ.get("CHANGE_HISTORY_DIR",
                                    r"C:\Projects\google-ads-ops\campaign_db_scripts\change_history")
ACCOUNT_ID = 1


def _import_taxonomy():
    if CHANGE_HISTORY_DIR not in sys.path:
        sys.path.insert(0, CHANGE_HISTORY_DIR)
    import change_taxonomy as T
    return T


def collect(conn, brand, start, end, account=ACCOUNT_ID):
    """Return (blocks, table) where blocks is the per-block detail list and table is the
    ledger table used. Each block dict: date, campaign, ad_group, context, category, action,
    klass, count, detail."""
    T = _import_taxonomy()
    b = brand.lower()
    table = "google_change_csv_history_raw"
    if not conn.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name=?", (table,)).fetchone():
        table = "google_change_csv_raw"  # fallback for a 2026 window
    q = ("SELECT change_datetime, campaign_col, ad_group_col, changes_raw FROM %s "
         "WHERE account_id=? AND substr(change_datetime,1,10) BETWEEN ? AND ?" % table)

    out = []
    for dt, camp, ag, raw in conn.execute(q, (account, start, end)):
        ag_is_geico = b in (ag or "").lower()
        blocks = T.split_blocks(raw)
        recs = T.categorize_cell(raw)
        for block, rec in zip(blocks, recs):
            text_has_geico = b in block.lower()
            if not (ag_is_geico or text_has_geico):
                continue
            context = "geico_adgroup" if ag_is_geico else "geico_in_change"
            out.append({
                "date": (dt or "")[:10],
                "campaign": camp or "",
                "ad_group": ag or "",
                "context": context,
                "category": rec["category"],
                "action": rec.get("action", ""),
                "klass": rec.get("klass", ""),
                "count": rec.get("count", 1),
                "detail": block.split("\n")[0].strip()[:120],
                "sample": " | ".join(l.strip() for l in block.split("\n")[1:] if l.strip())[:200],
            })
    return out, table


def rollup(blocks):
    """Aggregate by (context, category, action). Returns list of summary dicts."""
    agg = collections.defaultdict(lambda: {
        "n_blocks": 0, "n_items": 0, "campaigns": set(), "ad_groups": set(),
        "first": None, "last": None, "sample": ""})
    for bk in blocks:
        key = (bk["context"], bk["category"], bk["action"])
        a = agg[key]
        a["n_blocks"] += 1
        a["n_items"] += bk["count"] or 1
        if bk["campaign"]:
            a["campaigns"].add(bk["campaign"])
        if bk["ad_group"]:
            a["ad_groups"].add(bk["ad_group"])
        d = bk["date"]
        a["first"] = d if a["first"] is None or d < a["first"] else a["first"]
        a["last"] = d if a["last"] is None or d > a["last"] else a["last"]
        if not a["sample"] and (bk["sample"] or bk["detail"]):
            a["sample"] = (bk["sample"] or bk["detail"])[:180]
    rows = []
    for (context, category, action), a in agg.items():
        rows.append({
            "context": context, "category": category, "action": action,
            "n_blocks": a["n_blocks"], "n_items": a["n_items"],
            "n_campaigns": len(a["campaigns"]), "n_ad_groups": len(a["ad_groups"]),
            "first_date": a["first"], "last_date": a["last"], "sample": a["sample"]})
    rows.sort(key=lambda r: (r["context"], -r["n_blocks"]))
    return rows


def prose(blocks, summary, brand, start, end):
    """Build a short human paragraph describing the general change arc."""
    n = len(blocks)
    adg = [b for b in blocks if b["context"] == "geico_adgroup"]
    inch = [b for b in blocks if b["context"] == "geico_in_change"]
    neg = [b for b in inch if b["category"] == "negative_keyword"]
    cats = collections.Counter(b["category"] for b in adg)
    camps_adg = len({b["campaign"] for b in adg if b["campaign"]})
    top = ", ".join("%s (%d)" % (c, k) for c, k in cats.most_common(6))
    lines = []
    lines.append("%s change summary, %s to %s (account %d, Google Home)."
                 % (brand, start, end, ACCOUNT_ID))
    lines.append("")
    lines.append("%d Geico-relevant change blocks total: %d directly on Geico ad groups "
                 "across %d campaigns, and %d where 'geico' appeared in the change itself "
                 "(chiefly negative keywords fencing Geico traffic out of other ad groups)."
                 % (n, len(adg), camps_adg, len(inch)))
    lines.append("")
    if top:
        lines.append("On the Geico ad groups, the changes were: %s." % top)
    if neg:
        neg_items = sum(b["count"] or 1 for b in neg)
        neg_ags = len({b["ad_group"] for b in neg if b["ad_group"]})
        lines.append("Separately, %d negative-keyword blocks containing 'geico' (%d keywords) "
                     "were applied across %d ad groups." % (len(neg), neg_items, neg_ags))
    return "\n".join(lines)


def write_xlsx(path, prose_text, summary, blocks, brand, start, end):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("[!] openpyxl not installed. Run: pip install openpyxl"); return False
    wb = Workbook()

    s = wb.active
    s.title = "Summary"
    s["A1"] = "%s Change Summary — %s to %s" % (brand, start, end)
    s["A1"].font = Font(bold=True, size=13)
    for i, line in enumerate(prose_text.split("\n"), start=3):
        s.cell(row=i, column=1, value=line)
    start_row = 3 + prose_text.count("\n") + 3
    hdr = ["context", "category", "action", "# changes", "# items",
           "# campaigns", "# ad groups", "first", "last", "sample"]
    for j, h in enumerate(hdr, 1):
        ccell = s.cell(row=start_row, column=j, value=h.upper())
        ccell.font = Font(bold=True, color="FFFFFF")
        ccell.fill = PatternFill("solid", fgColor="1F4E79")
    for r, row in enumerate(summary, start=start_row + 1):
        for j, k in enumerate(["context", "category", "action", "n_blocks", "n_items",
                               "n_campaigns", "n_ad_groups", "first_date", "last_date", "sample"], 1):
            s.cell(row=r, column=j, value=row[k])
    widths = [16, 20, 10, 10, 9, 12, 12, 12, 12, 60]
    for j, w in enumerate(widths, 1):
        s.column_dimensions[get_column_letter(j)].width = w

    d = wb.create_sheet("All Geico changes")
    dh = ["date", "context", "campaign", "ad_group", "category", "action", "count", "detail", "sample"]
    for j, h in enumerate(dh, 1):
        c = d.cell(row=1, column=j, value=h.upper())
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
    for r, bk in enumerate(sorted(blocks, key=lambda b: (b["date"], b["context"])), start=2):
        for j, k in enumerate(dh, 1):
            d.cell(row=r, column=j, value=bk.get(k, ""))
    dwidths = [12, 16, 40, 34, 20, 10, 7, 44, 50]
    for j, w in enumerate(dwidths, 1):
        d.column_dimensions[get_column_letter(j)].width = w
    d.freeze_panes = "A2"
    d.auto_filter.ref = d.dimensions

    wb.save(path)
    return True


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--db", required=True)
    ap.add_argument("--brand", default="Geico")
    ap.add_argument("--start", default="2025-07-01")
    ap.add_argument("--end", default="2025-10-31")
    ap.add_argument("--out", default=".")
    args = ap.parse_args()

    conn = sqlite3.connect(args.db)
    blocks, table = collect(conn, args.brand, args.start, args.end)
    conn.close()
    print("read from %s | %d Geico-relevant blocks in %s..%s" % (table, len(blocks), args.start, args.end))
    if not blocks:
        print("(no Geico-relevant changes in window)"); return
    summary = rollup(blocks)
    prose_text = prose(blocks, summary, args.brand, args.start, args.end)
    print("\n" + prose_text + "\n")
    print("=== rollup (context / category / action : #changes | #items | #campaigns) ===")
    for r in summary:
        print("  %-15s %-20s %-8s : %4d | %5d | %d camps"
              % (r["context"], r["category"], r["action"], r["n_blocks"], r["n_items"], r["n_campaigns"]))

    os.makedirs(args.out, exist_ok=True)
    path = os.path.join(args.out, "%s_changes_%s_to_%s.xlsx"
                        % (args.brand.lower(), args.start, args.end))
    if write_xlsx(path, prose_text, summary, blocks, args.brand, args.start, args.end):
        print("\n[ok] wrote %s" % path)


if __name__ == "__main__":
    main()
