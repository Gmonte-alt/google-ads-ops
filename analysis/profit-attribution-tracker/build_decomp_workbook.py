#!/usr/bin/env python3
"""
build_decomp_workbook.py
========================
Packages the gp_decomposition_readout CSV outputs into one formatted Excel
workbook for review/sharing.

WHAT IT DOES
------------
Scans a directory for files named  decomp_<lens>_<grain>_<date>.csv
(any mix of lenses and grains: week / month / rolling), then builds:

  - "Summary" sheet: per grain -> account-level driver bridge table
    (VALUE / EFFICIENCY [cpc/cvr] / VOLUME / INTERACTION / NEW / EXITED)
    computed from that grain's category lens, plus a horizontal bar chart
    of the drivers. Reading notes included.
  - One sheet per CSV (e.g. roll_category, week_state_device): rows sorted
    bleeders -> gainers, red/white/green color scale on dGP and every
    driver column, currency/number formats, frozen header + autofilter,
    and LOW-VOLUME MUTING: rows where max(leads_prev, leads_cur) < 5 are
    grayed/italicized -- their driver splits (esp. interaction) are noise.

Values only (no formulas). Arial throughout. Charts built with explicit
fills (no invisible-marker repeats).

USAGE
-----
    python build_decomp_workbook.py                       # scans script dir
    python build_decomp_workbook.py --dir "C:\\...\\analysis\\profit-attribution-tracker"
    python build_decomp_workbook.py --dir ... --out my_readout.xlsx
"""

from __future__ import annotations

import argparse
import os
import re
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

FNAME_RE = re.compile(
    r"^decomp_(?P<lens>.+)_(?P<grain>week|month|rolling)_"
    r"(?P<date>\d{4}-\d{2}-\d{2})\.csv$")
PIVOT_RE = re.compile(
    r"^carrier_campaign_pivot_(?P<grain>week|month|rolling)_"
    r"(?P<date>\d{4}-\d{2}-\d{2})\.csv$")

DRIVERS = ["value", "efficiency", "volume", "interaction",
           "no_lead_spend", "new", "exited"]
SUB = ["eff_cpc", "eff_cvr", "eff_cross"]
LOW_LEAD_FLOOR = 5

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
BASE_FONT = Font(name="Arial", size=10)
MUTED_FONT = Font(name="Arial", size=10, italic=True, color="9E9E9E")
NOTE_FONT = Font(name="Arial", size=9, italic=True, color="595959")
TITLE_FONT = Font(name="Arial", size=14, bold=True, color="1F3864")
MONEY = '#,##0;(#,##0)'
NUM1 = '#,##0.0'
DRIVER_COLORS = {"value": "1F4E79", "efficiency": "C00000",
                 "eff_cpc": "E97132", "eff_cvr": "7030A0",
                 "volume": "2E7D32", "interaction": "7F7F7F",
                 "no_lead_spend": "BFBFBF", "new": "70AD47",
                 "exited": "8497B0"}

LENS_ORDER = ["category", "category_device", "match_type", "category_match",
              "carrier", "carrier_match", "carrier_match_device",
              "campaign", "state", "state_device"]
GRAIN_ORDER = {"rolling": 0, "week": 1, "month": 2}


def sheet_name(grain, lens):
    g = {"rolling": "roll", "week": "week", "month": "mnth"}[grain]
    return f"{g}_{lens}"[:31]


def col_order(df):
    keys = [c for c in ("carrier_brand", "keyword_category", "match_type",
                        "device", "campaign", "state") if c in df.columns]
    metrics = ["delta_gp", "value", "efficiency", "eff_cpc", "eff_cvr",
               "eff_cross", "volume", "interaction", "no_lead_spend",
               "new", "exited", "gp_prev", "gp_cur", "leads_prev",
               "leads_cur", "cost_prev", "cost_cur", "rpl_prev", "rpl_cur",
               "cpl_prev", "cpl_cur", "cpc_prev", "cpc_cur",
               "cvr_prev", "cvr_cur"]
    return keys + [c for c in metrics if c in df.columns], keys


def write_lens_sheet(wb, grain, lens, date, df):
    ws = wb.create_sheet(sheet_name(grain, lens))
    cols, keys = col_order(df)
    df = df[cols].sort_values("delta_gp")

    ws.cell(1, 1, f"{lens.replace('_', ' x ').upper()}  |  grain: {grain}"
                  f"  |  data through {date}").font = Font(
        name="Arial", size=11, bold=True, color="1F3864")
    ws.cell(2, 1, "Sorted bleeders -> gainers. Gray italic rows: fewer than "
                  f"{LOW_LEAD_FLOOR} leads in both periods -- driver splits "
                  "are noise; read dGP only.").font = NOTE_FONT

    hdr_row = 4
    for j, c in enumerate(cols, start=1):
        cell = ws.cell(hdr_row, j, c)
        cell.font, cell.fill = HDR_FONT, HDR_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for i, (_, r) in enumerate(df.iterrows(), start=hdr_row + 1):
        low = max(r.get("leads_prev", 0) or 0,
                  r.get("leads_cur", 0) or 0) < LOW_LEAD_FLOOR
        for j, c in enumerate(cols, start=1):
            v = r[c]
            cell = ws.cell(i, j, None if pd.isna(v) else
                           (round(float(v), 4) if isinstance(v, float) else v))
            cell.font = MUTED_FONT if low else BASE_FONT
            if c not in keys:
                cell.number_format = (NUM1 if "leads" in c or c.startswith(
                    ("rpl", "cpl", "cpc", "cvr")) else MONEY)
    last = hdr_row + len(df)

    # color scale on delta_gp + drivers + sub-split (numbers columns)
    for j, c in enumerate(cols, start=1):
        if c in DRIVERS + SUB + ["delta_gp"]:
            L = get_column_letter(j)
            ws.conditional_formatting.add(
                f"{L}{hdr_row+1}:{L}{last}",
                ColorScaleRule(start_type="min", start_color="F8696B",
                               mid_type="num", mid_value=0,
                               mid_color="FFFFFF",
                               end_type="max", end_color="63BE7B"))
    ws.auto_filter.ref = (f"A{hdr_row}:"
                          f"{get_column_letter(len(cols))}{last}")
    ws.freeze_panes = ws.cell(hdr_row + 1, len(keys) + 1)
    for j, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(j)].width = (
            max(12, min(46, int(df[c].astype(str).str.len().max() or 10) + 2))
            if c in keys else 11)
    return ws


def write_pivot_sheet(wb, grain, date, df):
    ws = wb.create_sheet(sheet_name(grain, "carrier_pivot"))
    ws.cell(1, 1, f"CARRIER PERFORMANCE x CAMPAIGN (GP by campaign; "
                  f"current {grain} period, through {date})").font = Font(
        name="Arial", size=11, bold=True, color="1F3864")
    ws.cell(2, 1, "Leading columns are the carrier's totals; remaining "
                  "columns are GP within each campaign.").font = NOTE_FONT
    hdr = 4
    cols = list(df.columns)
    for j, c in enumerate(cols, start=1):
        cell = ws.cell(hdr, j, c)
        cell.font, cell.fill = HDR_FONT, HDR_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for i, (_, r) in enumerate(df.iterrows(), start=hdr + 1):
        for j, c in enumerate(cols, start=1):
            v = r[c]
            cell = ws.cell(i, j, None if pd.isna(v) else
                           (round(float(v), 2) if isinstance(v, float) else v))
            cell.font = BASE_FONT
            if j > 1:
                cell.number_format = (NUM1 if "leads" in str(c) else MONEY)
    last = hdr + len(df)
    gp_cols = [j for j, c in enumerate(cols, start=1)
               if j > 1 and "cost" not in str(c) and "revenue" not in str(c)
               and "leads" not in str(c)]
    for j in gp_cols:
        L = get_column_letter(j)
        ws.conditional_formatting.add(
            f"{L}{hdr+1}:{L}{last}",
            ColorScaleRule(start_type="min", start_color="F8696B",
                           mid_type="num", mid_value=0, mid_color="FFFFFF",
                           end_type="max", end_color="63BE7B"))
    ws.freeze_panes = ws.cell(hdr + 1, 2)
    ws.auto_filter.ref = f"A{hdr}:{get_column_letter(len(cols))}{last}"
    ws.column_dimensions["A"].width = 24
    for j in range(2, len(cols) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 13
    return ws


def build_summary(wb, found):
    ws = wb.active
    ws.title = "Summary"
    ws.cell(1, 1, "Google Home -- GP Decomposition Readout").font = TITLE_FONT
    ws.cell(2, 1, "GP = leads x (RPL - CPL).  dGP = VALUE (dRPL) + "
                  "EFFICIENCY (-dCPL; split into cpc & cvr) + VOLUME "
                  "(dLeads) + INTERACTION.").font = NOTE_FONT
    ws.cell(3, 1, "Keyword lenses (category/match/campaign) and state "
                  "lenses are SIBLING views -- they reconcile in total but "
                  "Google provides no keyword x state cross.").font = NOTE_FONT

    row = 5
    grains = sorted({(g, d) for (_, g, d) in found},
                    key=lambda t: GRAIN_ORDER[t[0]])
    for grain, date in grains:
        cat = next((df for (l, g, d), df in found.items()
                    if l == "category" and g == grain and d == date), None)
        if cat is None:
            continue
        tot = cat[[c for c in DRIVERS + SUB + ["delta_gp", "gp_prev",
                                               "gp_cur"] if c in cat]].sum()
        ws.cell(row, 1, f"{grain.upper()} grain  (through {date})   "
                        f"GP {tot['gp_prev']:,.0f} -> {tot['gp_cur']:,.0f}"
                        f"   dGP {tot['delta_gp']:+,.0f}").font = Font(
            name="Arial", size=11, bold=True)
        row += 1
        hdr = row
        ws.cell(hdr, 1, "driver").font = HDR_FONT
        ws.cell(hdr, 1).fill = HDR_FILL
        ws.cell(hdr, 2, "dGP contribution").font = HDR_FONT
        ws.cell(hdr, 2).fill = HDR_FILL
        items = [(k, tot[k]) for k in DRIVERS if k in tot
                 and abs(tot[k]) > 0.5]
        # append the efficiency sub-split as indented rows
        subs = [(f"  {k.replace('eff_', 'eff: ')}", tot[k]) for k in SUB
                if k in tot and abs(tot[k]) > 0.5]
        for i, (k, v) in enumerate(items + subs, start=hdr + 1):
            ws.cell(i, 1, k).font = BASE_FONT
            c = ws.cell(i, 2, round(float(v), 2))
            c.font, c.number_format = BASE_FONT, MONEY
        n_main = len(items)
        chart = BarChart()
        chart.type = "bar"
        chart.title = f"{grain} -- dGP drivers"
        chart.legend = None
        chart.height, chart.width = 0.55 * (n_main + 2) + 2, 11
        data = Reference(ws, min_col=2, min_row=hdr + 1,
                         max_row=hdr + n_main)
        cats = Reference(ws, min_col=1, min_row=hdr + 1,
                         max_row=hdr + n_main)
        chart.add_data(data, titles_from_data=False)
        chart.set_categories(cats)
        s = chart.series[0]
        s.graphicalProperties.solidFill = "1F4E79"
        for idx, (k, _) in enumerate(items):
            dp = DataPoint(idx=idx)
            dp.graphicalProperties.solidFill = DRIVER_COLORS.get(k, "1F4E79")
            s.data_points.append(dp)
        ws.add_chart(chart, f"E{hdr - 1}")
        row = hdr + n_main + len(subs) + max(0, n_main + 3 - n_main
                                             - len(subs)) + 3
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 16
    return ws


def main():
    ap = argparse.ArgumentParser(
        description="Package decomp_*.csv files into a formatted workbook")
    ap.add_argument("--dir", default=None,
                    help="Directory containing decomp_*.csv (default: "
                         "this script's directory)")
    ap.add_argument("--out", default=None, help="Output .xlsx name")
    args = ap.parse_args()

    src = args.dir or os.path.dirname(os.path.abspath(__file__))
    found = {}
    pivots = {}
    for fn in sorted(os.listdir(src)):
        m = FNAME_RE.match(fn)
        if m:
            found[(m["lens"], m["grain"], m["date"])] = pd.read_csv(
                os.path.join(src, fn))
            continue
        pm = PIVOT_RE.match(fn)
        if pm:
            pivots[(pm["grain"], pm["date"])] = pd.read_csv(
                os.path.join(src, fn))
    if not found:
        print(f"[FAIL] no decomp_*.csv files found in {src}")
        sys.exit(1)
    # keep only the latest date per (lens, grain); older runs would collide
    # into duplicate sheets ("roll_category1", ...)
    latest_lg = {}
    for (l, g, d) in found:
        if (l, g) not in latest_lg or d > latest_lg[(l, g)]:
            latest_lg[(l, g)] = d
    dropped = [(l, g, d) for (l, g, d) in found if d != latest_lg[(l, g)]]
    for k in dropped:
        del found[k]
    if dropped:
        print(f"[NOTE] skipped {len(dropped)} older file(s) superseded by "
              f"a newer date for the same lens/grain.")
    latest_p = {}
    for (g, d) in pivots:
        if g not in latest_p or d > latest_p[g]:
            latest_p[g] = d
    pivots = {(g, d): v for (g, d), v in pivots.items()
              if d == latest_p[g]}
    print(f"found {len(found)} decomposition file(s):")
    for (l, g, d) in sorted(found, key=lambda t: (GRAIN_ORDER[t[1]],
                                                  LENS_ORDER.index(t[0])
                                                  if t[0] in LENS_ORDER
                                                  else 99)):
        print(f"  {g:<8} {l:<18} {d}")

    wb = Workbook()
    build_summary(wb, found)
    for (l, g, d) in sorted(found, key=lambda t: (GRAIN_ORDER[t[1]],
                                                  LENS_ORDER.index(t[0])
                                                  if t[0] in LENS_ORDER
                                                  else 99)):
        write_lens_sheet(wb, g, l, d, found[(l, g, d)])

    for (g, d), df in sorted(pivots.items(),
                             key=lambda t: GRAIN_ORDER[t[0][0]]):
        write_pivot_sheet(wb, g, d, df)
        print(f"  pivot    carrier_x_campaign {d} ({g})")

    latest = max(d for (_, _, d) in found)
    out = args.out or os.path.join(src, f"gp_decomposition_{latest}.xlsx")
    wb.save(out)
    print(f"\nWorkbook written: {out}")
    print("Sheets: Summary + one per lens/grain. No formulas (values "
          "readout); safe to share as-is.")


if __name__ == "__main__":
    main()
