#!/usr/bin/env python
"""
usia_keyword_select.py -- build the USIA Home ad-group keyword sets.

Reads home / on-vertical / not-DNB / national keywords from demand_research.db,
assigns each to one ad-group theme (first-match by priority), dedupes close
variants WITHIN each theme via TF-IDF + sklearn cosine (canonical = highest
avg_monthly_searches), caps each ad group at <=25 canonicals, and tags exact vs
phrase. Output feeds the launch shell's Ad Groups & Keywords tab.

Method ported from the bathroom-remodel KWP template: normalize+lemmatize ->
TF-IDF (1-2 grams) -> cosine cluster -> canonical per cluster. Adaptations:
theme-first assignment replaces the intent-category guardrail; cosine (not L2);
canonical by VOLUME (not row-count) to collapse KWP's shared-bucket synonyms.

    python usia_keyword_select.py
    python usia_keyword_select.py --sim 0.5 --cap 25 --exact-cutoff 10000 \
        --db C:\\Projects\\google-ads-ops\\db\\demand_research.db --out usia_home_keywords.xlsx
"""

import argparse
import csv
import re
import sqlite3
import string

from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity


def _make_lemmatizer():
    """Use WordNet if its corpus is already installed; otherwise a light
    rule-based plural folder. No runtime download (avoids proxy/SSL failures)."""
    try:
        from nltk.stem import WordNetLemmatizer
        wn = WordNetLemmatizer()
        wn.lemmatize("cars")          # probe -> LookupError if corpus absent
        return wn.lemmatize, "wordnet"
    except Exception:
        def _rule(t):
            if len(t) > 4 and t.endswith("ies"):
                return t[:-3] + "y"
            if len(t) > 4 and t.endswith(("ses", "xes", "zes", "ches", "shes")):
                return t[:-2]
            if len(t) > 3 and t.endswith("s") and not t.endswith("ss"):
                return t[:-1]
            return t
        return _rule, "rule-based"


LEMMA, LEMMA_KIND = _make_lemmatizer()

# ---- ad-group themes, priority order (first match wins); CORE is the catch-all ----
# Call channel wins first, then dwelling type (highest-leverage product modifier),
# then intent. Property-type terms ("cheap condo insurance") group by product, not intent.
THEME_PATTERNS = [
    ("AGENT_BROKER",        r"\b(agent|agents|agency|agencies|broker|brokers)\b"),
    ("NEAR_ME",             r"\b(near me|nearby|local|in my area)\b"),
    ("CONDO",               r"\b(condo|condos|condominium|condominiums|ho-?6)\b"),
    ("MOBILE_MANUFACTURED", r"\b(mobile home|mobile homes|manufactured home|manufactured homes|mobile homeowners?|manufactured homeowners?|trailer home|park home|park homes)\b"),
    ("PERIL",               r"\b(flood\w*|fire|wildfire|wind|windstorm|hail|hurricane|tornado|earthquake|seismic|sinkhole|landslide|mudslide|catastroph\w*|named peril|perils?|nfip|storm|disaster)\b"),
    ("HAZARD",              r"\bhazard\b"),
    ("PRICE",               r"\b(cheap|cheaper|cheapest|affordable|inexpensive|low cost|low-cost|lowest|lower|budget|discount|discounts)\b"),
    ("COST",                r"\b(cost|costs|price|prices|pricing|rate|rates|average|typical|how much|expensive|estimate|estimates)\b"),
    ("QUOTE",               r"\b(quote|quotes|quotation)\b"),
    ("COMPARE",             r"\b(best|good|great|top|top-rated|top rated|compar\w*|reviews?|ratings?|reputable|recommended)\b"),
    ("POLICY",              r"\b(ho-?[235789]|property and casualty|property casualty|personal liabilit\w*|personal property|personal belongings|bodily injury|property damage|casualty)\b"),
    ("COVERAGE",            r"\b(coverage|cover|policy|policies|dwelling|liability|deductible|replacement|renovation|remodel\w*|home improvement)\b"),
]
THEME_RE = [(name, re.compile(p)) for name, p in THEME_PATTERNS]
PRIORITY = [name for name, _ in THEME_PATTERNS] + ["BUNDLE", "HOMEOWNERS_EXT", "CORE"]
CALL_THEMES = {"AGENT_BROKER", "NEAR_ME"}   # these route to the Call campaign

# ---- exclusions: adjacent verticals (NOT home insurance) + geo-modified terms ----
# These are routed out of the home ad groups into a visible Excluded list (negatives source).
ADJACENT_PATTERNS = [
    ("home_warranty",   r"\b(warrant\w*|appliance\w*|protection plan\w*|home shield|american home shield)\b"),
    ("title",           r"\b(title insurance|home title)\b"),
    ("landlord_rental", r"\b(landlords?|land lord|let property|property let|buy to let|investment property|rental property|tenants?|vacant land|single family rental|sfr|lender.?placed|force.?placed|airbnb|short.?term rental|vrbo|vacation rental)\b"),
    ("business_event",  r"\b(business insurance|business property|business personal property|property insurance for business|insurance for business|business use|working from home|work from home|home business|home office insurance|office property|commercial|office building|commercial building|bounce house|event insurance|sba|eidl|occupational hazard|occupational accident|home inspector|property inspector|inspector insurance|boarding house|day ?care|home daycare|selling insurance|sell insurance)\b"),
    ("building",        r"\b(buildings? insurance|insurance for buildings?|build insurance|self build|new build insurance|builder'?s? risk|builders? risk|builders? insurance|home construction|construction insurance|dwelling under construction|under construction)\b"),
    ("hoa",             r"\b(hoa|homeowners association|condo association|master policy)\b"),
    ("farm_ranch",      r"\b(farm liability|farmowners|farm owners insurance|ranch insurance|agricultural insurance|crop insurance)\b"),
    ("renters",         r"\b(renters? insurance|tenants? insurance|cover my stuff)\b"),
    ("mortgage",        r"\b(mortgage insurance|mortgage protection|pmi)\b"),
    ("health_life",     r"\b(catastrophic (care|healthcare|health|coverage|plan|illness)|health insurance|health ?care|health plan|medicare|medicaid|medical insurance|nursing home|care home|care at home|home care insurance|key ?person|key ?man|life insurance|disability insurance|long.?term care|final expense|burial insurance)\b"),
]
ADJACENT_RE = [(name, re.compile(p)) for name, p in ADJACENT_PATTERNS]

# Auto handling is home-aware: a PURE auto term (no home signal) is adjacent and excluded, but
# a home+auto BUNDLE ("home and auto insurance quote") is a home-account prospect -- kept and
# routed to its own BUNDLE ad group. Brand bundles ("geico home and auto") fall through to the
# brand screen -> conquest; DNB carriers (USAA) drop via their tags.
AUTO_EXCLUDE_RE = re.compile(r"\b(auto insurance|car insurance|vehicle insurance|motorcycle insurance|rv insurance|boat insurance|atv insurance|auto property damage|auto liability|auto bodily injury)\b")
AUTO_SIGNAL_RE = re.compile(r"\b(auto|car|cars|vehicle|vehicles)\b")
HOME_TOKEN_RE = re.compile(r"\b(home|homes|homeowners?|house|houses|condo|condos)\b")


def is_bundle(kw):
    t = kw.lower()
    return bool(HOME_TOKEN_RE.search(t) and AUTO_SIGNAL_RE.search(t))

# Module negative for QBE (hard-avoid: exited the US market). The registry (v3.2) added
# QBE variants, but the demand pipeline's BrandTagger matches variants by EXACT string and
# embeds only canonical names of length >= 4, so "qbe" (3 chars) can't embed and suffixed
# forms like "qbe home insurance quote" miss the exact variants -> they'd still leak. One
# negative here catches every QBE form, which is cleaner than chasing each with a variant.
# (UHC/AIG share the short-name limitation but aren't leaking into the home selection.)
BRAND_INTERIM_RE = re.compile(r"\bqbe\b")

# Existing-customer / service intent -> route to negatives, not ad groups (e.g. "cal fair
# plan login"). Note: agent/phone/call/contact are NOT here -- those are call intent we want.
NAVIGATE_RE = re.compile(
    r"\b(log ?in|sign ?in|pay (my )?bill|bill pay|make a payment|cancel( my)? policy|"
    r"file a claim|claim status|claims? phone|customer service|policy number|"
    r"declarations? page|proof of insurance|account login)\b")

# Spanish-language terms -> not targeting Spanish at launch ("que es hazard insurance",
# "aseguranza blue fire"). Targeted so it doesn't catch English.
SPANISH_RE = re.compile(
    r"\b(que es|aseguranzas?|seguros? de|cuanto cuesta|para casa|de casa|de hogar|del hogar|hogar)\b")

# Too-generic / vague / informational terms with no clear home-purchase intent -> negatives.
# Includes how-to/informational ("home inventory") and year-stamped queries (stale; captured
# by the base term anyway).
GENERIC_NOISE_RE = re.compile(
    r"\b(additional insurance|extra insurance|more insurance|home inventory|inventory for insurance|"
    r"20[12]\d)\b")

# Insurance-company legal-name conventions ("X Fire and Casualty Insurance Company", "Empire Fire
# and Marine", "Ace Fire Underwriters"). These are carrier/brand intent that the PERIL theme would
# otherwise grab on the "fire" token. -> negatives / conquest, never generic peril heads.
CARRIER_NAME_RE = re.compile(
    r"\b(fire and (casualty|marine|liability|indemnity)|marine and fire|fire underwriters|"
    r"fire insurance compan\w*|casualty insurance compan\w*)\b")

# Non-home "X property" qualifiers: commercial / specialty / institutional property insurance that
# is NOT home insurance (the bare word "property" alone is ambiguous and stays; these qualifiers
# are not). Catches intellectual/church/farm/commercial/multifamily/building property, etc.
NONHOME_PROPERTY_RE = re.compile(
    r"\b(intellectual|church|farm|ranch|commercial|multifamily|multi family|apartment|business|"
    r"building|warehouse|industrial|municipal|government|federal|tribal) propert\w*\b")

# Non-insurance entities that match insurance jargon ("property and casualty") but are not
# insurers and not conquest targets -- educational (Kaplan), trade press (PropertyCasualty360).
NON_INSURANCE_RE = re.compile(
    r"\b(kaplan|propertycasualty ?360|property and casualty 360|property casualty 360)\b")

# Interim guard for brand terms leaking into generic: carriers/affinity not yet variant-tagged
# in the registry, plus DNB / out-of-scope (life) brands. Same role as UNTAGGED_CARRIER_RE;
# durable fix is registry variant-tagging (carriers->conquest, DNB->dnb). Specific phrases only,
# to avoid eating generic terms (e.g. "universal home", not bare "universal").
BRAND_LEAK_RE = re.compile(
    r"\b(njm|navy federal|true property insurance|liberty (home|house|homeowners|insurance)|"
    r"home first insurance|home insurance american|agi (condo|home|insurance)|"
    r"asi (home|insurance)|citizens? (home|homeowners|property)|first floridian|unitrin|"
    r"scotia|brighthouse|bright house|universal (home|casualty|property)|first american property|"
    r"state auto|nfp property|aig property|united property and casualty|ace property|"
    r"aon|alta (home|homeowner)|ameriprise|aegis|home ?site|wellington home|shelter (insurance|home)|"
    r"home protect|security first|garrison (property|home)|carolina property|primerica|"
    r"prudential home|horace mann|geovera|swyfft|old republic|homeowners insurance of america|"
    r"ocean harbor|western mutual)\b")

# Non-US property (overseas holiday homes) -- out of scope for a US home book. Country-adjective
# + property/holiday signal, scoped tight so it doesn't catch "french door" etc.
FOREIGN_RE = re.compile(
    r"\b(spanish|french|portuguese|italian|greek|german|overseas|abroad|spain|france|portugal)"
    r" (holiday|home|villa|propert\w*|house|second home)\b")

# HOLD: neither target nor add to negatives -- a parking lot for "decide later" (per review).
# Dropped silently from the build (not selected, not excluded->negative).
HOLD_RE = re.compile(r"\b(property damage exceeds|seeman holtz)\b")

# Homeowners-extension intent the taxonomy node doesn't yet enumerate (vacant/empty/holiday/
# second-home + vacant dwelling/property). Interim routing to HOMEOWNERS_EXT; durable fix is
# adding these as homeowners_extension seed terms. (vacant LAND is excluded upstream as
# landlord before this is reached.)
EXTENSION_RE = re.compile(
    r"\b(vacant|unoccupied|empty home|empty propert\w*|empty house|holiday home|holiday let|"
    r"second home|2nd home|seasonal home)\b")

COUNTY_RE = re.compile(r"\b\w+ county\b")   # county-level geo ("harris county home insurance")
# California-specific terms excluded at launch (CA not in the launch geo set; mine later).
CA_LAUNCH_RE = re.compile(r"\b(cal fire|calfire)\b")


def is_extension(kw):
    return bool(EXTENSION_RE.search(kw.lower()))

# Interim guard for carriers not yet in the registry, so they don't leak into the brand-free
# generic selection (same role as the QBE negative). DURABLE FIX is to add these to the brand
# registry + re-tag so they route correctly (carriers -> conquest/restricted, AAA -> affinity);
# see usia_peril_registry_additions.csv. Until then this keeps them out of the generic groups.
UNTAGGED_CARRIER_RE = re.compile(
    r"\b(wrights? (national )?flood|united (states )?fire|us fire insurance|twia|aaa|"
    r"national fire and marine|national liability and fire|selective flood|argenia|"
    r"farm state insurance|erty casualty|"
    r"twin city fire|charter oak fire|hanover|citizens? (property|home|insurance)|"
    r"honeycomb|fair plan|"
    r"blue fire|erie|kin (home|homeowners|insurance)|mpiua|foremost|triple a)\b")
    # subsidiaries/entities not yet in the registry: Twin City Fire (Hartford), Charter Oak
    # Fire (Travelers), Hanover, Citizens (state insurer), Honeycomb. FAIR plan lives in the
    # brand structure (California Fair Plan), so it is excluded from generic here.

US_STATES = {"alabama", "alaska", "arizona", "arkansas", "california", "colorado", "connecticut",
    "delaware", "florida", "georgia", "hawaii", "idaho", "illinois", "indiana", "iowa", "kansas",
    "kentucky", "louisiana", "maine", "maryland", "massachusetts", "michigan", "minnesota",
    "mississippi", "missouri", "montana", "nebraska", "nevada", "new hampshire", "new jersey",
    "new mexico", "new york", "north carolina", "north dakota", "ohio", "oklahoma", "oregon",
    "pennsylvania", "rhode island", "south carolina", "south dakota", "tennessee", "texas", "utah",
    "vermont", "virginia", "washington", "west virginia", "wisconsin", "wyoming"}
STATE_RE = re.compile(r"\b(" + "|".join(sorted(US_STATES, key=len, reverse=True)) + r")\b")
# safe 2-letter abbreviations; collision-prone ones (in/or/me/co/ma/pa/hi/ok/de/al/id) omitted
US_ABBR = {"az", "tx", "fl", "ny", "nj", "il", "oh", "ga", "nc", "sc", "va", "wa", "mn", "mo",
    "tn", "wi", "md", "ut", "nv", "nm", "ri", "wv", "mt", "sd", "nd", "vt", "wy", "ak"}

# Major US cities for city-level geo detection (the leaked terms were "houston home insurance",
# "home insurance dallas", etc.). DELIBERATELY EXCLUDES collision-prone names: "mobile" (would
# eat "mobile home insurance"), "jackson", "aurora", "mesa", "phoenix"->kept (low risk). Add as
# needed; this is launch-scoped (geo handled by campaign location targeting).
US_CITIES = {"new york", "los angeles", "chicago", "houston", "phoenix", "philadelphia",
    "san antonio", "san diego", "dallas", "austin", "fort worth", "san jose", "jacksonville",
    "columbus", "charlotte", "indianapolis", "seattle", "denver", "boston", "nashville",
    "baltimore", "oklahoma city", "louisville", "portland", "las vegas", "memphis", "detroit",
    "atlanta", "miami", "tampa", "orlando", "minneapolis", "cleveland", "kansas city",
    "st louis", "saint louis", "pittsburgh", "cincinnati", "sacramento", "san francisco",
    "fresno", "lubbock", "irving", "plano", "arlington", "corpus christi", "el paso", "tucson",
    "albuquerque", "omaha", "raleigh", "virginia beach", "long beach", "colorado springs",
    "milwaukee", "tulsa", "wichita", "new orleans", "honolulu", "buffalo", "rochester",
    "newark", "jersey city", "scottsdale", "henderson", "chesapeake", "norfolk", "richmond",
    "garland", "frisco", "galveston", "mckinney", "denton", "katy", "sugar land", "round rock", "ft worth"}
CITY_RE = re.compile(r"\b(" + "|".join(sorted(US_CITIES, key=len, reverse=True)) + r")\b")


# Collapse word-order + home-synonym duplicates that cosine clustering misses, e.g.
# "state farm home insurance quote" == "state farm quote home insurance", and
# "farmers home insurance" == "farmers insurance house". Key = sorted unique token set
# with home/homeowners/house folded to one token; representative = highest-volume surface.
HOME_SYNONYMS = {"homeowner": "home", "homeowners": "home", "homeowner's": "home",
                 "house": "home", "houses": "home", "homes": "home"}


def _stem(w):
    """Light inflectional stem so 'estimate'/'estimated' and plural forms collapse in the
    dedup key. Conservative -- folds home synonyms, plural -s/-ies, -ed/-ing, trailing -e."""
    w = HOME_SYNONYMS.get(w, w)
    if len(w) > 4 and w.endswith("ies"):
        w = w[:-3] + "y"
    elif len(w) > 3 and w.endswith("s") and not w.endswith("ss"):
        w = w[:-1]
    if len(w) > 4 and w.endswith("ed"):
        w = w[:-2]
    elif len(w) > 5 and w.endswith("ing"):
        w = w[:-3]
    if len(w) > 5 and w.endswith("e"):
        w = w[:-1]
    return w


def dedup_key(kw):
    toks = {_stem(w) for w in kw.lower().split() if w}
    return " ".join(sorted(toks))


def collapse_dupes(items):
    """items: list of (kw, vol) -> deduped list keeping the highest-volume surface form
    per token-set key (so permutations and home/house/homeowners variants merge to one)."""
    best = {}
    for kw, vol in items:
        k = dedup_key(kw)
        if k not in best or (vol or 0) > best[k][1]:
            best[k] = (kw, vol)
    return list(best.values())


# Single-word brand names that collide with generic insurance vocabulary -- the brand screen
# does NOT match generic candidates on these as bare tokens (their multi-word forms, e.g.
# "american family", "national general", still screen as phrases). Stops the screen from
# eating real generic terms like "standard home insurance" or "mutual home insurance".
COMMON_BRAND_COLLISIONS = {
    "universal", "national", "american", "general", "standard", "select", "selective",
    "pure", "mutual", "republic", "united", "western", "central", "family", "first",
    "liberty", "county", "alliance", "guard", "horace",
}


def load_brand_names(db):
    """Screen of known brand names (anything brand-tagged in the demand data) so the GENERIC
    selection drops brand terms even when a row is left untagged -- making it self-defending
    against tagging gaps instead of relying on a perfect retag. Multi-word names match as
    phrases; single-word names only if distinctive (len>=4, not a collision word). The retag
    is still what RECOVERS these into brand conquest; this only keeps them out of generic.
    Returns a compiled regex or None."""
    conn = sqlite3.connect(db)
    try:
        brands = [r[0] for r in conn.execute(
            "SELECT DISTINCT brand FROM kwp_ideas WHERE brand != ''").fetchall()]
    finally:
        conn.close()
    toks = set()
    for b in brands:
        bl = (b or "").lower().strip()
        if not bl:
            continue
        if " " in bl:                                   # multi-word: safe as a phrase
            toks.add(re.escape(bl))
        elif len(bl) >= 4 and bl not in COMMON_BRAND_COLLISIONS:
            toks.add(re.escape(bl))
    if not toks:
        return None
    return re.compile(r"\b(" + "|".join(sorted(toks, key=len, reverse=True)) + r")\b")


# USIA per-account product scope (the "mapping layer" the taxonomy deliberately leaves out --
# this is a USIA decision, not a property of the product). Flood IN (home-adjacent, converts in
# the sqp peril cut); dwelling_fire OUT (DP policies skew rental/vacant -> landlord territory).
USIA_IN_PRODUCTS = {"homeowners", "condo", "mobile_home", "flood", "homeowners_extension"}
USIA_OUT_PRODUCTS = {"landlord", "renters", "commercial_property", "bop", "dwelling_fire", "ranch_farm"}

# In-scope products that get their OWN ad group via the product tag (rather than routing by
# intent theme). homeowners_extension = vacant/unoccupied/second/seasonal home + land insurance.
PRODUCT_ADGROUP = {"homeowners_extension": "HOMEOWNERS_EXT"}


def load_product_map(db):
    """Read the product tags written by tag_products_demand_research.py. Returns
    {clean_keyword -> (product_code, vertical, ambiguous)}. Empty if not tagged yet, so the
    gate is a safe no-op until the tagging pass has been run (the regex fallback still works)."""
    conn = sqlite3.connect(db)
    try:
        cols = {r[1] for r in conn.execute("PRAGMA table_info(kwp_ideas)").fetchall()}
        if "product_code" not in cols:
            return {}
        rows = conn.execute(
            "SELECT keyword_text, product_code, product_vertical, product_ambiguous "
            "FROM kwp_ideas WHERE vertical='home' AND product_code IS NOT NULL").fetchall()
    finally:
        conn.close()
    m = {}
    for kw, code, vert, ambig in rows:
        m[clean_keyword(kw)] = (code, vert or "home", int(ambig or 0))
    return m


def classify_exclusion(kw, exclude_geo=True, brand_re=None, product=None):
    """Reason string if the keyword should leave the home ad groups, else None."""
    t = kw.lower()
    # 0) product taxonomy gate (data-driven, primary). Out-of-scope product or commercial
    # vertical -> exclude. Ambiguous-root matches ("building insurance") are NOT excluded here
    # -- they fall through to the regex fallback, which carries USIA's per-account scope calls.
    if product:
        code, vertical, ambiguous = product
        if not ambiguous and (code in USIA_OUT_PRODUCTS or vertical == "commercial"):
            return "product:" + code
    for name, rx in ADJACENT_RE:
        if rx.search(t):
            return "adjacent:" + name
    if NONHOME_PROPERTY_RE.search(t):
        return "adjacent:nonhome_property"   # commercial/specialty "X property" -- not home
    if (AUTO_EXCLUDE_RE.search(t) or AUTO_SIGNAL_RE.search(t)) and not HOME_TOKEN_RE.search(t):
        return "adjacent:auto"          # auto/vehicle signal, no home token (e.g. "hail damage car",
                                        # "car flood insurance"); home+auto bundles keep a home token
    if NAVIGATE_RE.search(t):
        return "navigate"
    if SPANISH_RE.search(t):
        return "spanish"
    if GENERIC_NOISE_RE.search(t):
        return "too_generic"
    if NON_INSURANCE_RE.search(t):
        return "non_insurance"
    if UNTAGGED_CARRIER_RE.search(t):
        return "untagged_carrier"
    if CARRIER_NAME_RE.search(t):
        return "carrier_name"               # "X fire and casualty/marine ... company" -- carrier names
    if BRAND_LEAK_RE.search(t):
        return "brand_leak"
    if BRAND_INTERIM_RE.search(t):
        return "brand_interim"
    if FOREIGN_RE.search(t):
        return "non_us"                     # overseas holiday/property -- out of scope
    if brand_re is not None and brand_re.search(t):
        return "branded"
    if exclude_geo and (STATE_RE.search(t) or CITY_RE.search(t) or COUNTY_RE.search(t)
                        or CA_LAUNCH_RE.search(t) or (set(t.split()) & US_ABBR)):
        return "geo_modified"
    return None

# words that survive normalization but shouldn't drive clustering
STOPWORDS = {"for", "and", "in", "of", "to", "a", "the", "my", "me", "is", "with"}

# home-synonym collapse so homeowners/house/property variants cluster together
NORM_REPLACEMENTS = {
    r"\bhome ?owner'?s?\b": " home ",
    r"\bhomeowner'?s?\b": " home ",
    r"\bhouse\b": " home ",
    r"\bhousehold\b": " home ",
    r"\bproperty\b": " home ",
    r"\bquotes\b": " quote ",
    r"\bquotation\b": " quote ",
}


def normalize_text(text):
    if not isinstance(text, str):
        return ""
    text = text.lower().translate(str.maketrans("", "", string.punctuation))
    for pat, repl in NORM_REPLACEMENTS.items():
        text = re.sub(pat, repl, text)
    toks = [LEMMA(t) for t in text.split()]
    toks = [t for t in toks if t not in STOPWORDS]
    return " ".join(toks)


def assign_theme(kw):
    """First match by priority on the original keyword text; default CORE."""
    t = kw.lower()
    for name, rx in THEME_RE:
        if rx.search(t):
            return name
    return "CORE"


def clean_keyword(kw):
    """Output hygiene on the biddable keyword text: '&' -> 'and', 'on line' -> 'online'."""
    if not isinstance(kw, str):
        return ""
    kw = kw.lower().strip()
    kw = kw.replace("&", " and ")
    kw = re.sub(r"\bon line\b", "online", kw)
    return re.sub(r"\s+", " ", kw).strip()


def cluster_theme(kws, vols, sim_threshold):
    """Greedy cluster by volume-desc seed. Returns clusters as lists of indices,
    each clusters[k][0] = canonical (highest-volume member)."""
    n = len(kws)
    if n == 0:
        return []
    if n == 1:
        return [[0]]
    norm = [normalize_text(k) for k in kws]
    X = TfidfVectorizer(ngram_range=(1, 2), min_df=1).fit_transform(norm)
    S = cosine_similarity(X)
    order = sorted(range(n), key=lambda i: -vols[i])
    assigned = [False] * n
    clusters = []
    for i in order:
        if assigned[i]:
            continue
        members = [i]
        assigned[i] = True
        for j in order:
            if not assigned[j] and S[i, j] >= sim_threshold:
                members.append(j)
                assigned[j] = True
        clusters.append(members)
    return clusters


def load_candidates(db, pull_dates=None):
    """National brand-free generic candidates. Unions across ALL generic-demand pulls
    by default (MAX volume per keyword), so a later seed batch (e.g. the peril pull)
    AUGMENTS the standard generic themes rather than replacing them -- scoping to a
    single pull_date silently drops the intent themes that live in the other pull.
    Pass an explicit list to pull_dates to restrict."""
    conn = sqlite3.connect(db)
    q = """
        SELECT keyword_text, MAX(avg_monthly_searches)
        FROM kwp_ideas
        WHERE vertical='home' AND geo_name='United States'
          AND vertical_match='on_vertical' AND dnb_flag=0 AND brand=''
          AND avg_monthly_searches > 0"""
    params = []
    if pull_dates:
        q += " AND pull_date IN (%s)" % ",".join("?" * len(pull_dates))
        params = list(pull_dates)
    q += " GROUP BY keyword_text"
    raw = conn.execute(q, params).fetchall()
    conn.close()
    echo = re.compile(r"^insurance\b.*\binsurance\b")   # KWP self-echo ("insurance home insurance")
    agg = {}
    for k, v in raw:
        ck = clean_keyword(k)
        if not ck or echo.search(ck):
            continue
        agg[ck] = max(agg.get(ck, 0), int(v or 0))
    return sorted(agg.items(), key=lambda x: -x[1])


def geoloop_date(db):
    conn = sqlite3.connect(db)
    d = conn.execute("""SELECT pull_date FROM kwp_ideas WHERE vertical='home'
                        GROUP BY pull_date ORDER BY COUNT(DISTINCT geo_constant_id) DESC LIMIT 1""").fetchone()
    conn.close()
    return d[0] if d else None


def build(db, out, sim_threshold, cap, exact_cutoff, floor, exclude_geo=True, pull_dates=None):
    gl = geoloop_date(db)
    cands = load_candidates(db, pull_dates)
    brand_re = load_brand_names(db)
    product_map = load_product_map(db)

    # 1) route adjacent/geo exclusions out, then assign theme (first-match by priority)
    excluded = {}     # reason -> list[(kw, vol)]
    by_theme = {t: [] for t in PRIORITY}
    for kw, vol in cands:
        if HOLD_RE.search(kw.lower()):
            continue                       # parking lot: neither targeted nor added to negatives
        pinfo = product_map.get(kw)
        reason = classify_exclusion(kw, exclude_geo, brand_re, pinfo)
        if reason:
            excluded.setdefault(reason, []).append((kw, vol))
            continue
        # home+auto bundles -> BUNDLE; in-scope specialty products / homeowners-extension intent
        # -> their own group; else theme
        if is_bundle(kw):
            ag = "BUNDLE"
        elif pinfo and PRODUCT_ADGROUP.get(pinfo[0]):
            ag = PRODUCT_ADGROUP[pinfo[0]]
        elif is_extension(kw):
            ag = "HOMEOWNERS_EXT"
        else:
            ag = assign_theme(kw)
        by_theme[ag].append((kw, vol))

    # BUNDLE earns its own ad group only if it clears the floor; otherwise fold each bundle
    # term back into its most relevant group (product node or intent theme) so the volume isn't
    # stranded in a sub-floor group.
    bundle_vol = sum(v for _, v in by_theme["BUNDLE"])
    if bundle_vol < floor:
        for kw, vol in by_theme["BUNDLE"]:
            pinfo = product_map.get(kw)
            ag = (PRODUCT_ADGROUP.get(pinfo[0]) if pinfo else None) or assign_theme(kw)
            by_theme[ag].append((kw, vol))
        by_theme["BUNDLE"] = []
        print("BUNDLE below floor (%s < %s) -> folded into themes" % (f"{bundle_vol:,}", f"{floor:,}"))
    selected = {}     # theme -> list of dict rows
    summary = []
    for theme in PRIORITY:
        pool = collapse_dupes(by_theme[theme])
        if not pool:
            summary.append((theme, 0, 0, 0, 0, 0, 0, "EMPTY"))
            continue
        kws = [k for k, _ in pool]
        vols = [v for _, v in pool]
        clusters = cluster_theme(kws, vols, sim_threshold)
        # canonical rows, ranked by canonical volume desc
        rows = []
        for members in clusters:
            ci = members[0]
            rows.append({
                "keyword": kws[ci],
                "volume": vols[ci],
                "cluster_size": len(members),
                "members": "; ".join(kws[m] for m in members),
            })
        rows.sort(key=lambda r: -r["volume"])
        kept = rows[:cap]
        # match type: >= cutoff -> exact; top canonical always exact; else phrase
        for idx, r in enumerate(kept):
            r["match_type"] = "exact" if (r["volume"] >= exact_cutoff or idx == 0) else "phrase"
            r["ad_group"] = theme
            r["campaign"] = "Call" if theme in CALL_THEMES else "Generic (tier)"
        selected[theme] = kept
        kept_vol = sum(r["volume"] for r in kept)
        n_exact = sum(1 for r in kept if r["match_type"] == "exact")
        summary.append((theme, len(pool), len(clusters), len(kept), kept_vol,
                        n_exact, len(kept) - n_exact, "OK" if kept_vol >= floor else "THIN"))

    write_outputs(out, selected, summary, gl, sim_threshold, cap, exact_cutoff, floor, excluded)
    return selected, summary, excluded


def write_outputs(out, selected, summary, gl, sim_threshold, cap, exact_cutoff, floor, excluded=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    HEAD = PatternFill("solid", fgColor="1F3864")
    SUB = PatternFill("solid", fgColor="D6DCE4")
    WHITE = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    BOLD = Font(name="Arial", bold=True, size=10)
    REG = Font(name="Arial", size=10)
    NOTE = Font(name="Arial", size=9, italic=True, color="595959")
    TITLE = Font(name="Arial", bold=True, size=14, color="1F3864")
    THIN = Side(style="thin", color="BFBFBF")
    BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    def hdr(ws, row, headers):
        for j, h in enumerate(headers, 1):
            c = ws.cell(row, j, h)
            c.fill = HEAD; c.font = WHITE
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = BORDER

    wb = Workbook(); wb.remove(wb.active)

    # Summary sheet
    ws = wb.create_sheet("Ad Group Summary")
    ws["A1"] = "USIA Home -- Keyword Selection Summary"; ws["A1"].font = TITLE
    ws.merge_cells("A1:H1")
    ws["A2"] = ("Source: demand_research.db (home, on-vertical, not-DNB, brand-free generic, national, pull %s). "
                "TF-IDF(1-2 gram) + sklearn cosine >= %.2f; canonical = highest volume; cap %d; "
                "exact if vol >= %s or top-of-group; floor %s. Carrier-brand terms excluded (Phase-2 conquest family)."
                % (gl, sim_threshold, cap, f"{exact_cutoff:,}", f"{floor:,}"))
    ws["A2"].font = NOTE; ws.merge_cells("A2:H2")
    hdr(ws, 4, ["Ad Group", "Campaign", "Raw kw", "Clusters", "Selected (<=cap)",
                "Selected Volume", "Exact / Phrase", "Floor"])
    r = 5
    for theme, raw_n, n_clu, n_sel, vol, n_ex, n_ph, flag in summary:
        ws.cell(r, 1, theme).font = BOLD
        ws.cell(r, 2, "Call" if theme in CALL_THEMES else "Generic (tier)").font = REG
        ws.cell(r, 3, raw_n).font = REG
        ws.cell(r, 4, n_clu).font = REG
        ws.cell(r, 5, n_sel).font = REG
        ws.cell(r, 6, vol).font = REG; ws.cell(r, 6).number_format = "#,##0"
        ws.cell(r, 7, "%d / %d" % (n_ex, n_ph)).font = REG
        c8 = ws.cell(r, 8, flag)
        c8.font = Font(name="Arial", size=10, color="C00000" if flag in ("THIN", "EMPTY") else "1F7A1F")
        r += 1
    for col, w in {"A": 14, "B": 16, "C": 9, "D": 10, "E": 16, "F": 16, "G": 14, "H": 9}.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"

    # Selected Keywords sheet
    ws = wb.create_sheet("Selected Keywords")
    ws["A1"] = "Selected Keywords -- by ad group"; ws["A1"].font = TITLE
    ws.merge_cells("A1:F1")
    ws["A2"] = ("One row per canonical keyword (the variant to bid). 'Cluster members' = the close "
                "variants it absorbs (KWP shared-bucket synonyms). Match type: exact heads + phrase mid-tail.")
    ws["A2"].font = NOTE; ws.merge_cells("A2:F2")
    hdr(ws, 4, ["Campaign", "Ad Group", "Keyword (canonical)", "Match Type", "Volume", "Cluster Members"])
    r = 5
    for theme in PRIORITY:
        for row in selected.get(theme, []):
            ws.cell(r, 1, row["campaign"]).font = REG
            ws.cell(r, 2, row["ad_group"]).font = BOLD
            ws.cell(r, 3, row["keyword"]).font = REG
            mt = ws.cell(r, 4, row["match_type"])
            mt.font = Font(name="Arial", size=10, color="1F3864" if row["match_type"] == "exact" else "000000")
            ws.cell(r, 5, row["volume"]).font = REG; ws.cell(r, 5).number_format = "#,##0"
            ws.cell(r, 6, row["members"]).font = NOTE
            ws.cell(r, 6).alignment = Alignment(wrap_text=True, vertical="top")
            r += 1
    for col, w in {"A": 15, "B": 14, "C": 40, "D": 11, "E": 11, "F": 70}.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"

    # Excluded sheet (adjacent verticals + geo-modified) -> negative-keyword source
    ws = wb.create_sheet("Excluded (negatives)")
    ws["A1"] = "Excluded from home ad groups"; ws["A1"].font = TITLE
    ws.merge_cells("A1:C1")
    ws["A2"] = ("Adjacent verticals (home warranty/appliance, title, landlord, business, mortgage) and "
                "geo-modified terms. Not home insurance / handled by location targeting. Use as negatives.")
    ws["A2"].font = NOTE; ws.merge_cells("A2:C2")
    hdr(ws, 4, ["Reason", "Keyword", "Volume"])
    r = 5
    for reason in sorted(excluded or {}):
        for kw, vol in sorted(excluded[reason], key=lambda x: -x[1]):
            ws.cell(r, 1, reason).font = REG
            ws.cell(r, 2, kw).font = REG
            ws.cell(r, 3, vol).font = REG; ws.cell(r, 3).number_format = "#,##0"
            r += 1
    for col, w in {"A": 22, "B": 44, "C": 11}.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    wb.save(out)

    # CSV (Editor-friendly) alongside
    csv_path = out.rsplit(".", 1)[0] + ".csv"
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["campaign", "ad_group", "keyword", "match_type", "volume", "cluster_size", "cluster_members"])
        for theme in PRIORITY:
            for row in selected.get(theme, []):
                w.writerow([row["campaign"], row["ad_group"], row["keyword"], row["match_type"],
                            row["volume"], row["cluster_size"], row["members"]])
    # excluded negatives CSV
    neg_path = out.rsplit(".", 1)[0] + "_excluded.csv"
    with open(neg_path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["reason", "keyword", "volume"])
        for reason in sorted(excluded or {}):
            for kw, vol in sorted(excluded[reason], key=lambda x: -x[1]):
                w.writerow([reason, kw, vol])
    print("wrote %s, %s, %s" % (out, csv_path, neg_path))


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--db", default=r"C:\Projects\google-ads-ops\db\demand_research.db")
    p.add_argument("--out", default="usia_home_keywords.xlsx")
    p.add_argument("--sim", type=float, default=0.5, help="cosine similarity to merge near-variants (default 0.5)")
    p.add_argument("--cap", type=int, default=25, help="max canonicals per ad group (default 25)")
    p.add_argument("--exact-cutoff", type=int, default=10000, help="vol >= this -> exact (default 10000)")
    p.add_argument("--floor", type=int, default=2000, help="ad-group volume floor for the THIN flag (default 2000)")
    p.add_argument("--keep-geo", action="store_true",
                   help="keep state/abbr geo-modified keywords in the ad groups (default: exclude, since location targeting handles geo)")
    p.add_argument("--pull-dates", nargs="*", default=None,
                   help="restrict to these pull_date(s); default unions ALL generic pulls so peril augments (not replaces) the standard intent themes.")
    a = p.parse_args()
    _, summary, excluded = build(a.db, a.out, a.sim, a.cap, a.exact_cutoff, a.floor,
                                 exclude_geo=not a.keep_geo, pull_dates=a.pull_dates)
    print("\nlemmatizer: %s" % LEMMA_KIND)
    print("ad groups:")
    for theme, raw_n, n_clu, n_sel, vol, n_ex, n_ph, flag in summary:
        print("  %-13s raw=%-4d clusters=%-4d kept=%-3d vol=%-9s exact/phrase=%d/%d  %s"
              % (theme, raw_n, n_clu, n_sel, f"{vol:,}", n_ex, n_ph, flag))
    if excluded:
        print("excluded (-> negatives):")
        for reason in sorted(excluded):
            tot = sum(v for _, v in excluded[reason])
            print("  %-22s %3d kw  vol=%s" % (reason, len(excluded[reason]), f"{tot:,}"))


if __name__ == "__main__":
    main()
