#!/usr/bin/env python
"""
usia_keyword_select_v2.py -- USIA Home ad-group keyword sets, v2 selection logic.

v2 changes vs usia_keyword_select.py (v1 is preserved untouched, for diffing):

  1) HEAD-FINDING is INTENT-SIGNATURE COVERAGE, not "highest-volume member".
     A keyword's intent signature = its normalized tokens MINUS the universal base
     {home, insurance}. So "cheap home insurance quotes" -> {cheap, quote};
     "home and auto phone number" -> {auto, phone, number}; "home insurance" -> {}.
     A head (a phrase keyword) consolidates every member whose signature it anchors.
     The bid keyword is always a REAL member surface (never a bare fragment like
     "phone number"), so a content-free signature {} is only ever the CORE catch-all
     "home insurance".

  2) FLOOR is applied to a head's SUMMED COVERED DEMAND, AFTER consolidation -- not to
     a single keyword's volume. Ten 10-volume members under one head sum to 100 and the
     head clears a floor of 100, even though no single member would. Sub-floor specific
     signatures fold UP by shedding their lowest-coverage (peripheral) token toward their
     anchor (e.g. {auto, phone, number} -> {auto} = "home and auto insurance"), so the
     phone-number-bundle intent rides inside BUNDLE when it can't earn its own group.

  3) Fold-up is PHRASE-CORRECT. A modifier intent ("cheap/quote/agent/near me") whose
     surface CONTAINS "home insurance" can ride under the CORE "home insurance" head when
     it's sub-floor (phrase "home insurance" matches "cheap home insurance"). A product
     intent ("condo insurance", "home and auto insurance", "flood insurance") has NO broader
     phrase that covers it, so a sub-floor product head KEEPS its own (THIN-flagged) head
     rather than mis-folding into "home insurance", which would not match it.

  4) MATCH TYPE by volume cutoff: a head clears EXACT when its summed covered demand is
     >= --exact-cutoff (high-demand, unambiguous core), else PHRASE. Each EXACT head gets
     1-2 phrase COMPANIONS -- a qualified-tail member whose intent signature strictly
     supersets the head's (e.g. {company} -> {company, phone, number}), since exact won't
     match the same query with added words. Default companion = highest-volume superset, plus
     a call-intent superset ("phone/number/contact/call") when one exists. Never a same-text
     twin (a bare phrase would cannibalize every other head), and empty-signature heads (the
     bare CORE "home insurance") get exact only. Phrase heads reach their own tail, so no
     companions. Promote proven heads/terms to additional exact later (~10 conv at target ROAS).

  5) CARRIER EXCLUSION (un-rollback). v2 uses v1's full exclusion verbatim: the carrier
     guards (UNTAGGED_CARRIER_RE / BRAND_LEAK_RE / BRAND_INTERIM_RE) stay in force, so real
     regional carriers (Erie, Citizens, Foremost, Hanover, Kin, Swyfft, FAIR Plan, Wright
     Flood, ...) are excluded from the generic book and routed to negatives -> conquest,
     instead of each minting its own brand-specific generic head. (An earlier rollback let
     them re-enter and form their own heads -- abandoned; it inflated CORE with conquest-by-
     accident.) Only addition: a USAA (DNB) guard for any untagged "usaa ..." variant.

  TF-IDF clustering is retained and reported as a secondary surface-cluster cross-check
  (the "Surface clusters" column), per the note that TF-IDF is right for grouping but
  wrong for head-finding -- so it informs, but does not drive, head choice.

    python usia_keyword_select_v2.py
    python usia_keyword_select_v2.py --floor 100 --exact-cutoff 10000 \
        --db C:\\Projects\\google-ads-ops\\db\\demand_research.db --out usia_home_keywords_v2.xlsx

Requires usia_keyword_select.py in the same folder (imports the stable load/exclusion/
normalize layers from it).
"""

import argparse
import csv
import re
import sqlite3
from collections import defaultdict

# Stable machinery imported from v1. v2's ONLY divergence from v1 is head selection
# (intent-signature coverage) + match typing -- the load, exclusion, normalize, theme,
# and cluster layers are reused unchanged. classify_exclusion brings v1's full carrier
# exclusion (UNTAGGED_CARRIER_RE / BRAND_LEAK_RE / BRAND_INTERIM_RE) back into force.
from usia_keyword_select import (
    LEMMA_KIND,
    PRIORITY, CALL_THEMES,
    is_bundle, HOLD_RE, is_extension,
    collapse_dupes, load_brand_names,
    PRODUCT_ADGROUP, load_product_map,
    normalize_text, assign_theme,
    cluster_theme, load_candidates, geoloop_date,
    classify_exclusion,
)


# ---------------------------------------------------------------------------
# CARRIER UN-ROLLBACK. v2 now uses v1's full exclusion verbatim, so the real regional
# carriers (Erie, Citizens, Foremost, Hanover, Kin, Swyfft, FAIR Plan, Wright Flood, ...)
# are excluded from the generic book and routed to negatives -> conquest, instead of each
# minting its own brand-specific generic head. The only addition is a USAA guard: USAA is
# DNB, normally dropped by the dnb_flag/brand filter in load_candidates, but this catches any
# untagged "usaa ..." variant that slips through with brand=''.
USAA_DNB_RE = re.compile(r"\busaa\b")


def classify_exclusion_v2(kw, exclude_geo=True, brand_re=None, product=None):
    """USAA (DNB) belt-and-suspenders, then v1's full exclusion (incl. all carrier guards).
    Strict-excluded terms never enter the corpus, so they are invisible to the n-gram math."""
    if USAA_DNB_RE.search(kw.lower()):
        return "brand_strict"
    return classify_exclusion(kw, exclude_geo, brand_re, product)


# ---------------------------------------------------------------------------
# Intent-signature head selection.
BASE_TOKENS = frozenset({"home", "insurance"})   # universal base (post-normalize fold)
# A sub-floor head can ride under the CORE "home insurance" phrase head ONLY if its surface
# actually contains that phrase (so phrase match would capture it). Product intents
# ("condo insurance", "home and auto insurance", "flood insurance") do NOT, and keep own head.
HOME_INS_PHRASE_RE = re.compile(r"\b(home|homeowner|homeowners) insurance\b")


def intent_sig(kw):
    """Normalized tokens minus the universal base -> the keyword's distinguishing intent.
    Empty signature == pure 'home/homeowners insurance' (the CORE catch-all)."""
    return frozenset(set(normalize_text(kw).split()) - BASE_TOKENS)


def _head_text(sig, members):
    """Bid surface for a head: a clean canonical surface carrying exactly this intent.
    Prefer surfaces containing 'insurance' (else any multiword surface), then pick the
    highest-individual-volume one (the most representative real query). The old 'shortest
    wins' rule let a bare possessive like "homeowner's" beat "home insurance" -- never again."""
    exact = [m for m in members if intent_sig(m["kw"]) == sig]
    pool = exact or members
    clean = [m for m in pool if "insurance" in m["kw"].split()]
    if not clean:
        clean = [m for m in pool if len(m["kw"].split()) >= 2]
    cand = clean or pool
    return max(cand, key=lambda m: (m["vol"], -len(m["kw"].split())))["kw"]


def select_heads(pool, floor, theme):
    """pool: list of (kw, vol) for ONE theme. Returns (heads, migrate):
      heads   -- list of dicts {keyword, sig, volume, n_members, members, thin}
      migrate -- list of (kw, vol) sub-floor modifier terms to ride under CORE 'home insurance'

    Method: each member starts at its full intent signature. We iteratively fold any
    below-floor signature UP by shedding its lowest-coverage (peripheral) token toward the
    high-coverage anchor (the consolidation step). After folding, members still stranded in
    a sub-floor signature either migrate to CORE (if their surface contains 'home insurance',
    so the CORE phrase head covers them) or keep their own THIN-flagged head (product intents
    with no broader phrase). Floor is on SUMMED covered demand, applied AFTER consolidation."""
    members = [{"kw": kw, "vol": vol, "sig": intent_sig(kw),
                "core": bool(HOME_INS_PHRASE_RE.search(kw))} for kw, vol in pool]

    # iterative token-shedding fold (consolidate scattered specifics into their anchor)
    while True:
        tokcov = defaultdict(int)
        for m in members:
            for t in m["sig"]:
                tokcov[t] += m["vol"]
        dem = defaultdict(int)
        for m in members:
            dem[m["sig"]] += m["vol"]
        moved = False
        for m in members:
            s = m["sig"]
            if len(s) >= 2 and dem[s] < floor:
                drop = min(s, key=lambda t: (tokcov[t], t))   # shed the most peripheral token
                m["sig"] = frozenset(s - {drop})
                moved = True
        if not moved:
            break

    # final partition demand; route sub-floor members
    dem = defaultdict(int)
    for m in members:
        dem[m["sig"]] += m["vol"]
    keep, migrate = [], []
    for m in members:
        sub_floor = m["sig"] and dem[m["sig"]] < floor and m["core"]
        if sub_floor and theme != "CORE":
            migrate.append((m["kw"], m["vol"]))   # modifier term -> rides under CORE 'home insurance'
        elif sub_floor:                            # already in CORE: absorb into the base head
            m["sig"] = frozenset()                 # ride under 'home insurance' (phrase covers it)
            keep.append(m)
        else:
            keep.append(m)

    groups = defaultdict(list)
    for m in keep:
        groups[m["sig"]].append(m)
    heads = []
    for sig, ms in groups.items():
        ms.sort(key=lambda m: -m["vol"])
        vol = sum(m["vol"] for m in ms)
        heads.append({
            "keyword": _head_text(sig, ms),
            "sig": sig,
            "volume": vol,
            "n_members": len(ms),
            "members": "; ".join(m["kw"] for m in ms),
            "member_list": [(m["kw"], m["vol"]) for m in ms],
            "thin": vol < floor,
        })
    heads.sort(key=lambda h: -h["volume"])
    return heads, migrate


# ---------------------------------------------------------------------------
# Match-type + phrase-companion pass.
#
# Exact above a volume cutoff (high-demand, unambiguous-intent core); phrase below it.
# An exact head locks the bare query (Google prioritizes an identical exact over phrase),
# but exact will NOT match the same query with added words -- so for each exact head we
# attach phrase COMPANIONS that recover that qualified tail. A companion must be a REAL
# member surface whose intent signature is a STRICT SUPERSET of the head's (it ADDS intent,
# e.g. {company} -> {company, phone, number}); a same-signature member ("household insurance
# companies") is a synonym the exact close-variant already covers, so it is NOT a companion,
# and we never emit a same-text phrase twin (that would be the catch-all we're avoiding).
#
# Heads with an EMPTY signature (the bare CORE "home insurance") get NO companion -- their
# members are bare synonyms, and a "home insurance" phrase would cannibalize every other head.
# Companion selection per exact head: the highest-volume strict-superset member, PLUS the
# highest-volume call-intent strict-superset member when one exists (1-2 companions).
CALL_TOKENS = frozenset({"phone", "number", "contact", "call"})


def pick_companions(head, brand_re=None):
    """Up to two phrase companions for an exact head: highest-volume strict-superset member,
    plus a call-intent strict-superset member if present. Empty-signature heads get none."""
    hs = head["sig"]
    if not hs:
        return []
    cands = [(kw, vol) for kw, vol in head["member_list"]
             if intent_sig(kw) > hs and (brand_re is None or not brand_re.search(kw.lower()))]
    if not cands:
        return []
    comps = []
    top = max(cands, key=lambda x: x[1])
    comps.append({"keyword": top[0], "volume": top[1], "call": bool(intent_sig(top[0]) & CALL_TOKENS)})
    calls = [(kw, vol) for kw, vol in cands if intent_sig(kw) & CALL_TOKENS]
    if calls:
        tc = max(calls, key=lambda x: x[1])
        if tc[0] != top[0]:
            comps.append({"keyword": tc[0], "volume": tc[1], "call": True})
    return comps


def apply_match_types(heads, exact_cutoff, brand_re=None):
    """Set each head exact/phrase by volume cutoff; attach phrase companions to exact heads."""
    for h in heads:
        h["match_type"] = "exact" if h["volume"] >= exact_cutoff else "phrase"
        h["companions"] = pick_companions(h, brand_re) if h["match_type"] == "exact" else []


# ---------------------------------------------------------------------------
def build(db, out, sim_threshold, floor, exact_cutoff, exclude_geo=True, pull_dates=None):
    gl = geoloop_date(db)
    cands = load_candidates(db, pull_dates)
    brand_re = load_brand_names(db)
    product_map = load_product_map(db)

    # 1) exclusions out (strict -> invisible to n-gram), then theme/product/bundle routing
    excluded = {}
    by_theme = {t: [] for t in PRIORITY}
    for kw, vol in cands:
        if HOLD_RE.search(kw.lower()):
            continue
        pinfo = product_map.get(kw)
        reason = classify_exclusion_v2(kw, exclude_geo, brand_re, pinfo)
        if reason:
            excluded.setdefault(reason, []).append((kw, vol))
            continue
        if is_bundle(kw):
            ag = "BUNDLE"
        elif pinfo and PRODUCT_ADGROUP.get(pinfo[0]):
            ag = PRODUCT_ADGROUP[pinfo[0]]
        elif is_extension(kw):
            ag = "HOMEOWNERS_EXT"
        else:
            ag = assign_theme(kw)
        by_theme[ag].append((kw, vol))

    # 2) head selection per theme. Non-CORE themes first; their sub-floor modifier terms
    # migrate into CORE so they ride under the 'home insurance' phrase head.
    selected = {}
    summary = []
    core_extra = []
    surface_clusters = {}
    for theme in PRIORITY:
        if theme == "CORE":
            continue
        pool = collapse_dupes(by_theme[theme])
        surface_clusters[theme] = len(cluster_theme([k for k, _ in pool],
                                                    [v for _, v in pool], sim_threshold)) if pool else 0
        heads, migrate = select_heads(pool, floor, theme)
        core_extra.extend(migrate)
        for h in heads:
            h["ad_group"] = theme
            h["campaign"] = "Call" if theme in CALL_THEMES else "Generic (tier)"
        selected[theme] = heads

    # CORE last, with the migrants folded in
    core_pool = collapse_dupes(by_theme["CORE"] + core_extra)
    surface_clusters["CORE"] = len(cluster_theme([k for k, _ in core_pool],
                                                 [v for _, v in core_pool], sim_threshold)) if core_pool else 0
    core_heads, _ = select_heads(core_pool, floor, "CORE")
    for h in core_heads:
        h["ad_group"] = "CORE"
        h["campaign"] = "Generic (tier)"
    selected["CORE"] = core_heads

    # 3) match types + phrase companions
    for theme in PRIORITY:
        apply_match_types(selected.get(theme, []), exact_cutoff, brand_re)

    for theme in PRIORITY:
        heads = selected.get(theme, [])
        vol = sum(h["volume"] for h in heads)
        n_thin = sum(1 for h in heads if h["thin"])
        n_exact = sum(1 for h in heads if h["match_type"] == "exact")
        n_comp = sum(len(h["companions"]) for h in heads)
        flag = "EMPTY" if not heads else ("THIN" if vol < floor else "OK")
        summary.append((theme, len(by_theme[theme]) + (len(core_extra) if theme == "CORE" else 0),
                        surface_clusters.get(theme, 0), len(heads), n_exact, n_comp, vol, n_thin, flag))

    write_outputs(out, selected, summary, gl, sim_threshold, floor, exact_cutoff, excluded, len(core_extra))
    return selected, summary, excluded


def write_outputs(out, selected, summary, gl, sim_threshold, floor, exact_cutoff, excluded, n_migrated):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    HEAD = PatternFill("solid", fgColor="1F3864")
    WHITE = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    BOLD = Font(name="Arial", bold=True, size=10)
    REG = Font(name="Arial", size=10)
    NOTE = Font(name="Arial", size=9, italic=True, color="595959")
    TITLE = Font(name="Arial", bold=True, size=14, color="1F3864")
    THIN = Side(style="thin", color="BFBFBF")
    BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    def hdr(ws, row, headers):
        for j, h in enumerate(headers, 1):
            c = ws.cell(row, j, h)
            c.fill = HEAD; c.font = WHITE
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = BORDER

    wb = Workbook(); wb.remove(wb.active)

    ws = wb.create_sheet("Ad Group Summary")
    ws["A1"] = "USIA Home -- Keyword Selection Summary (v2)"; ws["A1"].font = TITLE
    ws.merge_cells("A1:I1")
    ws["A2"] = ("Source: demand_research.db (home, on-vertical, not-DNB, brand-free generic, national, pull %s). "
                "v2: intent-signature coverage heads; floor %s on SUMMED covered demand after consolidation; "
                "exact >= %s else phrase; 1-2 phrase companions per exact head (highest-volume superset + a "
                "call-intent superset); carrier terms excluded to conquest (v1 guards in force). %d sub-floor "
                "modifier terms migrated to CORE. TF-IDF surface clusters shown as a cross-check only."
                % (gl, f"{floor:,}", f"{exact_cutoff:,}", n_migrated))
    ws["A2"].font = NOTE; ws.merge_cells("A2:I2")
    hdr(ws, 4, ["Ad Group", "Campaign", "Raw kw", "Surface clusters",
                "Heads", "Exact", "Companions", "Head Volume", "Floor"])
    r = 5
    for theme, raw_n, n_clu, n_sel, n_exact, n_comp, vol, n_thin, flag in summary:
        ws.cell(r, 1, theme).font = BOLD
        ws.cell(r, 2, "Call" if theme in CALL_THEMES else "Generic (tier)").font = REG
        ws.cell(r, 3, raw_n).font = REG
        ws.cell(r, 4, n_clu).font = REG
        ws.cell(r, 5, n_sel).font = REG
        ws.cell(r, 6, "%d / %d ph" % (n_exact, n_sel - n_exact)).font = REG
        ws.cell(r, 7, n_comp).font = REG
        ws.cell(r, 8, vol).font = REG; ws.cell(r, 8).number_format = "#,##0"
        c9 = ws.cell(r, 9, "%s (%d thin)" % (flag, n_thin))
        c9.font = Font(name="Arial", size=10, color="C00000" if flag in ("THIN", "EMPTY") else "1F7A1F")
        r += 1
    for col, w in {"A": 16, "B": 16, "C": 9, "D": 14, "E": 8, "F": 12, "G": 12, "H": 13, "I": 13}.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"

    ws = wb.create_sheet("Selected Keywords")
    ws["A1"] = "Selected Keywords -- heads (exact/phrase) + phrase companions"; ws["A1"].font = TITLE
    ws.merge_cells("A1:G1")
    ws["A2"] = ("One row per HEAD, plus indented phrase COMPANION rows under each exact head (a qualified "
                "tail the exact won't match -- never a same-text twin). 'Volume' on a head = summed covered "
                "demand (the floor test). THIN = below floor but kept (product intent, no broader phrase).")
    ws["A2"].font = NOTE; ws.merge_cells("A2:G2")
    hdr(ws, 4, ["Campaign", "Ad Group", "Keyword", "Match Type", "Volume", "Role", "Covered Members"])
    r = 5
    for theme in PRIORITY:
        for row in selected.get(theme, []):
            ws.cell(r, 1, row["campaign"]).font = REG
            ws.cell(r, 2, row["ad_group"]).font = BOLD
            kwc = ws.cell(r, 3, row["keyword"])
            kwc.font = Font(name="Arial", size=10, color="C00000" if row["thin"] else "000000")
            mt = ws.cell(r, 4, row["match_type"])
            mt.font = Font(name="Arial", size=10, bold=row["match_type"] == "exact",
                           color="1F3864" if row["match_type"] == "exact" else "000000")
            ws.cell(r, 5, row["volume"]).font = REG; ws.cell(r, 5).number_format = "#,##0"
            ws.cell(r, 6, "head").font = REG
            ws.cell(r, 7, row["members"]).font = NOTE
            ws.cell(r, 7).alignment = Alignment(wrap_text=True, vertical="top")
            r += 1
            for comp in row["companions"]:
                ws.cell(r, 1, row["campaign"]).font = NOTE
                ws.cell(r, 2, row["ad_group"]).font = NOTE
                ws.cell(r, 3, "   \u21b3 " + comp["keyword"]).font = REG
                ws.cell(r, 4, "phrase").font = REG
                ws.cell(r, 5, comp["volume"]).font = REG; ws.cell(r, 5).number_format = "#,##0"
                ws.cell(r, 6, "companion (call)" if comp["call"] else "companion").font = NOTE
                ws.cell(r, 7, "qualifies: " + row["keyword"]).font = NOTE
                r += 1
    for col, w in {"A": 15, "B": 16, "C": 44, "D": 11, "E": 12, "F": 16, "G": 70}.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"

    ws = wb.create_sheet("Excluded (negatives)")
    ws["A1"] = "Excluded from home ad groups (strict -> invisible to n-gram)"; ws["A1"].font = TITLE
    ws.merge_cells("A1:C1")
    ws["A2"] = ("Strict exclusions: adjacent verticals, geo, service-navigate, Spanish, non-US, plus all "
                "carrier/brand terms (Erie, Citizens, Foremost, Hanover, Kin, Swyfft, FAIR Plan, USAA, "
                "Liberty, AAA, ...) -- excluded from generic and routed here as the conquest / negative-keyword "
                "source. Use as the generic-campaign negative list.")
    ws["A2"].font = NOTE; ws.merge_cells("A2:C2")
    hdr(ws, 4, ["Reason", "Keyword", "Volume"])
    r = 5
    for reason in sorted(excluded or {}):
        for kw, vol in sorted(excluded[reason], key=lambda x: -x[1]):
            ws.cell(r, 1, reason).font = REG
            ws.cell(r, 2, kw).font = REG
            ws.cell(r, 3, vol).font = REG; ws.cell(r, 3).number_format = "#,##0"
            r += 1
    for col, w in {"A": 20, "B": 46, "C": 11}.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    wb.save(out)

    csv_path = out.rsplit(".", 1)[0] + ".csv"
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["campaign", "ad_group", "keyword", "match_type", "volume",
                    "role", "qualifies_head", "n_covered", "covered_members", "thin"])
        for theme in PRIORITY:
            for row in selected.get(theme, []):
                w.writerow([row["campaign"], row["ad_group"], row["keyword"], row["match_type"],
                            row["volume"], "head", "", row["n_members"], row["members"], int(row["thin"])])
                for comp in row["companions"]:
                    w.writerow([row["campaign"], row["ad_group"], comp["keyword"], "phrase",
                                comp["volume"], "companion_call" if comp["call"] else "companion",
                                row["keyword"], "", "", 0])
    neg_path = out.rsplit(".", 1)[0] + "_excluded.csv"
    with open(neg_path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["reason", "keyword", "volume"])
        for reason in sorted(excluded or {}):
            for kw, vol in sorted(excluded[reason], key=lambda x: -x[1]):
                w.writerow([reason, kw, vol])
    print("wrote %s, %s, %s" % (out, csv_path, neg_path))


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--db", default=r"C:\Projects\google-ads-ops\db\demand_research.db")
    p.add_argument("--out", default="usia_home_keywords_v2.xlsx")
    p.add_argument("--sim", type=float, default=0.5, help="cosine sim for the surface-cluster cross-check (default 0.5)")
    p.add_argument("--floor", type=int, default=100,
                   help="minimum SUMMED covered demand for a head to earn its own keyword (default 100)")
    p.add_argument("--exact-cutoff", type=int, default=10000,
                   help="head volume >= this -> exact (with phrase companions); else phrase (default 10000)")
    p.add_argument("--keep-geo", action="store_true",
                   help="keep geo-modified keywords (default: exclude; location targeting handles geo)")
    p.add_argument("--pull-dates", nargs="*", default=None,
                   help="restrict to these pull_date(s); default unions ALL generic pulls.")
    a = p.parse_args()
    _, summary, excluded = build(a.db, a.out, a.sim, a.floor, a.exact_cutoff,
                                 exclude_geo=not a.keep_geo, pull_dates=a.pull_dates)
    print("\nlemmatizer: %s" % LEMMA_KIND)
    print("ad groups (v2 intent-signature heads; exact >= %s, phrase below; + companions):" % f"{a.exact_cutoff:,}")
    for theme, raw_n, n_clu, n_sel, n_exact, n_comp, vol, n_thin, flag in summary:
        print("  %-14s raw=%-4d surf_clu=%-4d heads=%-3d exact=%-2d comp=%-2d vol=%-9s thin=%-2d  %s"
              % (theme, raw_n, n_clu, n_sel, n_exact, n_comp, f"{vol:,}", n_thin, flag))
    if excluded:
        print("excluded (strict -> negatives, invisible to n-gram):")
        for reason in sorted(excluded):
            tot = sum(v for _, v in excluded[reason])
            print("  %-22s %3d kw  vol=%s" % (reason, len(excluded[reason]), f"{tot:,}"))


if __name__ == "__main__":
    main()
