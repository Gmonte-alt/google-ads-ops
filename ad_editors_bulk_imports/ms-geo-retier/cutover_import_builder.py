"""
cutover_import_builder.py — Microsoft Home Re-Tier cutover: Editor bulk import build, gate & verify
Charter: one bulk session 2026-09-28 — 26 location deletes + 26 adds + 4 campaign tCPA raises.

This thread BUILDS AND POSTS; it does not re-decide the plan. Inputs are authoritative:
  - cutover_planB_2026-09-22_editor_import.csv   (52 validated location rows — the intent)
  - cutover_planB_2026-09-22_changeset.xlsx      (Changes / Full Map / Pre-flight Audit)
  - Home Insurance_PositiveLocations_current_view_2026-09-22.csv  (Gate 0 baseline export)

MODES (mutually exclusive):
  1. Validation only (default — run any time before Monday):
       python cutover_import_builder.py
     Re-runs every intake check + assembles the 56-row import IN MEMORY. Writes nothing.

  2. Build (Monday morning, AFTER pulling the fresh location export):
       python cutover_import_builder.py --fresh-export "input\\cutover_2026-09-28\\<fresh export>.csv"
     Gate 0: diffs fresh export vs the 9/22 baseline. ANY difference -> BLOCKED, nothing
     written, hand the fresh export back to the analysis thread. Identical -> writes the
     final 56-row import + run log to output\\cutover_2026-09-28\\.

  3. Post-verify (after posting, with a re-export pulled from the platform):
       python cutover_import_builder.py --post-export "<post-post export>.csv"
     Diffs actual platform state vs the Full Map expected end-state. Expect ZERO diffs.
     Writes a post-verification report for the weekly change log entry.

Explicitly excluded per charter: Tier2/Tier5 tCPAs (deferred cuts), all keyword/ad/
ad-group/budget/negative/audience changes, the expansion ad groups, the creative test.
"""

import argparse
import csv
import hashlib
import sys
from datetime import datetime
from pathlib import Path

# ============================================================
# CONFIGURATION — everything lives under the ms-geo-retier base
# ============================================================

BASE_DIR = Path(r"C:\Projects\google-ads-ops\ad_editors_bulk_imports\ms-geo-retier")
RUN_NAME = "cutover_2026-09-28"

INPUT_DIR = BASE_DIR / "input" / RUN_NAME
OUTPUT_DIR = BASE_DIR / "output" / RUN_NAME

HANDED_IMPORT = INPUT_DIR / "cutover_planB_2026-09-22_editor_import.csv"
CHANGESET_XLSX = INPUT_DIR / "cutover_planB_2026-09-22_changeset.xlsx"
BASELINE_EXPORT = INPUT_DIR / "Home Insurance_PositiveLocations_current_view_2026-09-22.csv"

# Campaign registry per charter (verify IDs against the account before posting).
# The name->ID map DERIVED from the baseline export must agree with this exactly.
REGISTRY = {
    "531494235": "HOME INSURANCE_STATE-GROUP_Tier1RPL80Plus_COMBINED_TCPA_SEARCH_ALLDEVICE_USA-STE",
    "531494261": "HOME INSURANCE_STATE-GROUP_Tier2RPL65Plus_COMBINED_TCPA_SEARCH_ALLDEVICE_USA-STE",
    "531494273": "HOME INSURANCE_STATE-GROUP_Tier3RPL55Plus_COMBINED_TCPA_SEARCH_ALLDEVICE_USA-STE",
    "531494275": "HOME INSURANCE_STATE-GROUP_Tier5RPL35Plus_COMBINED_TCPA_SEARCH_ALLDEVICE_USA-STE",
    "531494232": "HOME INSURANCE-CALIFORNIA-SEARCH-MOBILE-USA-STE",
    "531494260": "HOME INSURANCE-FLORIDA-SEARCH-MOBILE-USA-STE",
}

# Net-new build: the 4 tCPA raises posting WITH the cutover.
# (label, campaign_id, tcpa_today_for_manual_verify, tcpa_set_to)
TCPA_CHANGES = [
    ("Tier1", "531494235", "47.25", "54.34"),
    ("Tier3", "531494273", "36.58", "37.44"),
    ("California", "531494232", "16.10", "18.52"),
    ("Florida", "531494260", "12.00", "13.47"),
]
# Tier2 (531494261) and Tier5 (531494275) tCPAs are UNCHANGED — deferred by charter.
TCPA_FROZEN_IDS = {"531494261", "531494275"}

# Charter sentinels: (state, must_be_removed_from_id, must_be_added_to_id)
SENTINEL_MOVES = [
    ("Massachusetts", "531494261", "531494235"),  # Tier2 -> Tier1
    ("Hawaii", "531494260", "531494275"),         # Florida -> Tier5
    ("Louisiana", "531494260", "531494275"),      # Florida -> Tier5
    ("Oklahoma", "531494273", "531494261"),       # Tier3 -> Tier2
]
FORBIDDEN_STATES = ["Texas"]  # any Texas row anywhere is a build error

EXPECTED_STATE_COUNT = 49


# ============================================================
# HELPERS
# ============================================================

def unbr(v: str) -> str:
    return (v or "").strip().strip("[]").strip()


def state_of(name: str) -> str:
    return (name or "").split(",")[0].strip()


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def read_editor_export(path: Path):
    """Microsoft Ads Editor location export: UTF-16 (BOM-aware), tab-delimited."""
    with open(path, "rb") as f:
        magic = f.read(2)
    if magic == b"PK":
        raise SystemExit(
            f"FATAL: {path.name} is an Excel/zip file, not an Editor location export.\n"
            "  Pass the UTF-16 tab-delimited CSV exported from Ads Editor\n"
            "  (e.g. \"Home Insurance_PositiveLocations_current_view_*.csv\")."
        )
    with open(path, "r", encoding="utf-16", newline="") as f:
        rows = list(csv.DictReader(f, delimiter="\t"))
    if not rows:
        raise SystemExit(f"FATAL: export at {path} parsed to zero rows")
    return rows


def read_handed_import(path: Path):
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        return list(reader), list(reader.fieldnames)


class Report:
    def __init__(self):
        self.lines = []
        self.failed = False

    def check(self, ok: bool, label: str, detail: str = ""):
        tag = "PASS" if ok else "FAIL"
        if not ok:
            self.failed = True
        self.lines.append(f"  [{tag}] {label}" + (f" — {detail}" if detail else ""))

    def note(self, text: str):
        self.lines.append(text)

    def dump(self):
        print("\n".join(self.lines))


# ============================================================
# CORE VALIDATION + ASSEMBLY
# ============================================================

def derive_campaign_map(export_rows, rep: Report):
    """Campaign name -> ID map from the export's own structural columns."""
    name_to_ids = {}
    for r in export_rows:
        name, cid = r.get("Campaign", "").strip(), unbr(r.get("Parent Id", ""))
        if name and cid:
            name_to_ids.setdefault(name, set()).add(cid)
    ambiguous = {n: ids for n, ids in name_to_ids.items() if len(ids) > 1}
    rep.check(not ambiguous, "Export campaign names resolve to exactly one ID each",
              str(ambiguous) if ambiguous else "")
    name_to_id = {n: next(iter(ids)) for n, ids in name_to_ids.items()}
    derived = {cid: n for n, cid in name_to_id.items()}
    rep.check(derived == REGISTRY, "Derived name<->ID map matches charter registry",
              "" if derived == REGISTRY else f"derived={derived}")
    return name_to_id


def validate_and_assemble(rep: Report):
    """Runs all charter validations; returns (final_rows, final_headers, end_state)."""
    for p in (HANDED_IMPORT, CHANGESET_XLSX, BASELINE_EXPORT):
        if not p.exists():
            raise SystemExit(f"FATAL: required input missing: {p}")

    rep.note("INPUT FILES (sha256):")
    for p in (HANDED_IMPORT, CHANGESET_XLSX, BASELINE_EXPORT):
        rep.note(f"  {p.name}  {sha256(p)}")

    handed, headers = read_handed_import(HANDED_IMPORT)
    baseline = read_editor_export(BASELINE_EXPORT)

    rep.note("\n1) HANDED IMPORT CSV")
    dels = [r for r in handed if r["Status"] == "Deleted"]
    adds = [r for r in handed if r["Status"] == "Active"]
    rep.check(len(handed) == 52, "52 location rows", f"got {len(handed)}")
    rep.check(len(dels) == 26 and len(adds) == 26, "26 delete / 26 add",
              f"got {len(dels)}/{len(adds)}")
    rep.check(all(r["Type"] == "Campaign Location Criterion" for r in handed),
              "All rows are Campaign Location Criterion")
    rep.check(all(r["Location Type"] == "State/Province" for r in handed),
              "All rows are State/Province level")
    rep.check(all(r["Bid Adjustment"] in ("0", "") for r in handed),
              "All Bid Adjustments are 0")
    rep.check(all(unbr(r["ID"]) for r in dels), "Every delete carries a criterion ID")
    rep.check(all(r["Target"].strip() for r in handed), "Every row carries a Target ID")

    rep.note("\n2) BASELINE EXPORT (9/22) STRUCTURAL CROSS-CHECKS")
    rep.check(len(baseline) == EXPECTED_STATE_COUNT,
              f"Baseline export has {EXPECTED_STATE_COUNT} location rows",
              f"got {len(baseline)}")
    name_to_id = derive_campaign_map(baseline, rep)

    exp_pairs = {(unbr(r["ID"]), unbr(r["Parent Id"])) for r in baseline}
    missing = [r["Name"] for r in dels if (unbr(r["ID"]), unbr(r["Parent Id"])) not in exp_pairs]
    rep.check(not missing, "All 26 delete (criterion ID, campaign ID) pairs exist in export",
              str(missing))

    exp_tc = {}
    for r in baseline:
        exp_tc.setdefault(state_of(r["Name"]), set()).add(r["Target"].strip())
    tc_bad = [r["Name"] for r in adds
              if r["Target"].strip() not in exp_tc.get(state_of(r["Name"]), {r["Target"].strip()})]
    rep.check(not tc_bad, "Add Target IDs agree with export codes per state", str(tc_bad))

    rep.note("\n3) PARENT ID BACKFILL ON ADD ROWS (from export structural map)")
    unresolved = []
    for r in adds:
        cid = name_to_id.get(r["Campaign"])
        if cid:
            r["Parent Id"] = f"[{cid}]"
        else:
            unresolved.append(r["Campaign"])
    rep.check(not unresolved, "All 26 add destinations resolved name->ID via export",
              str(sorted(set(unresolved))))

    rep.note("\n4) CHARTER SENTINELS")
    for st in FORBIDDEN_STATES:
        hit = [r["Status"] for r in handed if state_of(r["Name"]) == st]
        rep.check(not hit, f"{st} appears NOWHERE in the import", f"found {hit}")
    for st, from_id, to_id in SENTINEL_MOVES:
        d_ok = any(state_of(r["Name"]) == st and unbr(r["Parent Id"]) == from_id for r in dels)
        a_ok = any(state_of(r["Name"]) == st and unbr(r["Parent Id"]) == to_id for r in adds)
        rep.check(d_ok and a_ok,
                  f"{st}: removed from {from_id}, added to {to_id}",
                  f"delete_ok={d_ok} add_ok={a_ok}")

    rep.note("\n5) CHANGESET WORKBOOK AGREEMENT")
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise SystemExit("FATAL: openpyxl required (pip install openpyxl)")
    wb = load_workbook(CHANGESET_XLSX, read_only=True)
    ch = list(wb["Changes (ADD-REMOVE)"].iter_rows(values_only=True))[1:]
    ch_rem = {(r[0], r[2]) for r in ch if r[3] == "REMOVE"}
    ch_add = {(r[0], r[4]) for r in ch if r[3] == "ADD"}
    csv_rem = {(state_of(r["Name"]), r["Campaign"]) for r in dels}
    csv_add = {(state_of(r["Name"]), r["Campaign"]) for r in adds}
    rep.check(ch_rem == csv_rem, "Changes-sheet REMOVE set == CSV delete set",
              str(ch_rem ^ csv_rem) if ch_rem != csv_rem else "")
    rep.check(ch_add == csv_add, "Changes-sheet ADD set == CSV add set",
              str(ch_add ^ csv_add) if ch_add != csv_add else "")
    audit = list(wb["Pre-flight Audit"].iter_rows(values_only=True))[1:]
    rep.check(len(audit) == 0, "Pre-flight Audit sheet has 0 findings", f"{len(audit)} rows")

    fm = list(wb["Full Map"].iter_rows(values_only=True))[1:]
    expected_end = {r[0]: r[4] for r in fm if r[3] in ("KEEP", "ADD")}  # state -> campaign name
    rep.check(len(expected_end) == EXPECTED_STATE_COUNT,
              f"Full Map end-state resolves to {EXPECTED_STATE_COUNT} states once each",
              f"got {len(expected_end)}")

    rep.note("\n6) SIMULATED END-STATE (baseline − deletes + adds)")
    del_keys = {(unbr(r["ID"]), unbr(r["Parent Id"])) for r in dels}
    sim = {state_of(r["Name"]): unbr(r["Parent Id"]) for r in baseline
           if (unbr(r["ID"]), unbr(r["Parent Id"])) not in del_keys}
    dup_hits = []
    for r in adds:
        st = state_of(r["Name"])
        if st in sim:
            dup_hits.append(st)
        sim[st] = unbr(r["Parent Id"])
    rep.check(not dup_hits, "No state ends dual-targeted in simulation", str(dup_hits))
    rep.check(len(sim) == EXPECTED_STATE_COUNT,
              f"Simulated end-state covers {EXPECTED_STATE_COUNT} states", f"got {len(sim)}")
    sim_names = {st: REGISTRY.get(cid, cid) for st, cid in sim.items()}
    mismatch = {st: (sim_names[st], expected_end.get(st))
                for st in sim_names if sim_names[st] != expected_end.get(st)}
    rep.check(not mismatch, "Simulated end-state == Full Map expected end-state",
              str(mismatch) if mismatch else "")
    fl_states = sorted(st for st, cid in sim.items() if cid == "531494260")
    ca_states = sorted(st for st, cid in sim.items() if cid == "531494232")
    rep.check(fl_states == ["Florida"], "Florida campaign ends targeting exactly {Florida}",
              str(fl_states))
    rep.check(ca_states == ["California"], "California campaign ends targeting exactly {California}",
              str(ca_states))
    empty = [cid for cid in REGISTRY if cid not in set(sim.values())]
    rep.check(not empty, "No campaign ends empty", str(empty))

    rep.note("\n7) ASSEMBLY — 52 location rows + 4 campaign tCPA rows = 56")
    final_headers = list(headers)
    for col in ("Bid Strategy Type", "Bid Strategy TargetCpa"):
        if col not in final_headers:
            final_headers.append(col)
    final_rows = []
    for r in handed:
        row = {h: r.get(h, "") for h in final_headers}
        final_rows.append(row)
    frozen_violation = [cid for _, cid, _, _ in TCPA_CHANGES if cid in TCPA_FROZEN_IDS]
    rep.check(not frozen_violation, "No tCPA change targets Tier2/Tier5 (frozen by charter)",
              str(frozen_violation))
    for label, cid, today, new in TCPA_CHANGES:
        row = {h: "" for h in final_headers}
        row["Type"] = "Campaign"
        row["Status"] = "Active"
        row["ID"] = f"[{cid}]"
        row["Campaign"] = REGISTRY[cid]
        row["Bid Strategy Type"] = "TargetCpa"
        row["Bid Strategy TargetCpa"] = new
        final_rows.append(row)
        rep.note(f"    + {label} ({cid}): tCPA -> ${new}  (verify today = ${today} in Editor first)")
    rep.check(len(final_rows) == 56, "Final import is exactly 56 rows", f"got {len(final_rows)}")
    touched = {unbr(r["Parent Id"]) or unbr(r["ID"]) for r in final_rows}
    rep.check(touched == set(REGISTRY), "Exactly 6 campaigns touched, matching registry",
              str(touched ^ set(REGISTRY)))

    return final_rows, final_headers, sim_names


# ============================================================
# GATE 0 — fresh-export freshness check (blocking)
# ============================================================

def gate0(fresh_path: Path, rep: Report) -> bool:
    rep.note("\nGATE 0 — FRESHNESS CHECK (blocking)")
    rep.note(f"  Fresh export: {fresh_path.name}  {sha256(fresh_path)}")
    fresh = read_editor_export(fresh_path)
    base = read_editor_export(BASELINE_EXPORT)

    base_hdr, fresh_hdr = set(base[0].keys()), set(fresh[0].keys())
    if base_hdr != fresh_hdr:
        rep.check(False, "Export headers identical",
                  f"only-in-baseline={sorted(base_hdr - fresh_hdr)} only-in-fresh={sorted(fresh_hdr - base_hdr)}")
        return False
    common = [h for h in base[0].keys()]

    def keyed(rows):
        return {(unbr(r["ID"]), unbr(r["Parent Id"])): tuple(r[h].strip() for h in common)
                for r in rows}

    b, f = keyed(base), keyed(fresh)
    only_base = sorted(set(b) - set(f))
    only_fresh = sorted(set(f) - set(b))
    changed = sorted(k for k in set(b) & set(f) if b[k] != f[k])

    identical = not (only_base or only_fresh or changed)
    rep.check(identical, "Fresh export identical to 9/22 baseline",
              f"removed={len(only_base)} new={len(only_fresh)} changed={len(changed)}")
    if not identical:
        for k in only_base[:10]:
            rep.note(f"    GONE from account since 9/22: criterion {k}")
        for k in only_fresh[:10]:
            rep.note(f"    NEW in account since 9/22:  criterion {k}")
        for k in changed[:10]:
            diffs = [f"{h}: '{bv}' -> '{fv}'" for h, bv, fv in
                     zip(common, b[k], f[k]) if bv != fv]
            rep.note(f"    CHANGED {k}: " + "; ".join(diffs))
        rep.note("\n  *** BLOCKED. Do NOT post. Do NOT hand-edit. ***")
        rep.note("  Hand the fresh export back to the analysis thread:")
        rep.note("    build_cutover_changeset.py --grouping B  (against the fresh export)")
    return identical


# ============================================================
# POST-VERIFY — platform re-export vs Full Map end-state
# ============================================================

def post_verify(post_path: Path, expected_end: dict, rep: Report) -> bool:
    rep.note("\nPOST-VERIFICATION — platform export vs Full Map end-state")
    rep.note(f"  Post export: {post_path.name}  {sha256(post_path)}")
    post = read_editor_export(post_path)
    actual = {}
    dup = []
    for r in post:
        st = state_of(r["Name"])
        if st in actual:
            dup.append(st)
        actual[st] = r["Campaign"].strip()
    rep.check(not dup, "No state dual-targeted in platform export", str(dup))
    rep.check(len(actual) == EXPECTED_STATE_COUNT,
              f"Platform targets {EXPECTED_STATE_COUNT} states", f"got {len(actual)}")
    mism = {st: (actual.get(st), expected_end.get(st))
            for st in set(actual) | set(expected_end)
            if actual.get(st) != expected_end.get(st)}
    rep.check(not mism, "Platform end-state matches Full Map EXACTLY (expect zero diffs)",
              str(mism) if mism else "")
    if not mism and not dup:
        rep.note("  ZERO diffs. Report back: post timestamp + this result + the 4 confirmed tCPAs.")
        rep.note("  Log in weekly change log — category: geo re-tier + target change, 6 campaigns.")
    return not (mism or dup)


# ============================================================
# MAIN
# ============================================================

def main():
    ap = argparse.ArgumentParser(description="MS Home re-tier cutover: build, gate, verify")
    mode = ap.add_mutually_exclusive_group()
    mode.add_argument("--fresh-export", type=Path,
                      help="Monday-morning location export -> Gate 0 + write final import")
    mode.add_argument("--post-export", type=Path,
                      help="Post-posting platform export -> verify against Full Map")
    args = ap.parse_args()

    stamp = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    rep = Report()
    rep.note(f"MS HOME RE-TIER CUTOVER BUILDER — run {stamp}")
    rep.note(f"Base: {BASE_DIR}\n")

    final_rows, final_headers, sim_names = validate_and_assemble(rep)

    if rep.failed:
        rep.dump()
        print("\nVALIDATION FAILED — nothing written. Fix inputs or escalate to analysis thread.")
        sys.exit(1)

    if args.post_export:
        ok = post_verify(args.post_export, sim_names, rep)
        OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        report_path = OUTPUT_DIR / f"post_verification_{stamp}.txt"
        report_path.write_text("\n".join(rep.lines), encoding="utf-8")
        rep.dump()
        print(f"\nReport written: {report_path}")
        sys.exit(0 if ok else 1)

    if not args.fresh_export:
        rep.dump()
        print("\nVALIDATION-ONLY RUN — all checks passed; nothing written.")
        print("Monday: pull fresh location export into the input folder, then run:")
        print(f'  python cutover_import_builder.py --fresh-export "input\\{RUN_NAME}\\<fresh export>.csv"')
        sys.exit(0)

    # Build mode: Gate 0 then write
    if not gate0(args.fresh_export, rep):
        rep.dump()
        sys.exit(2)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    import_path = OUTPUT_DIR / f"cutover_planB_FINAL_56row_import_{stamp}.csv"
    with open(import_path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=final_headers)
        w.writeheader()
        for r in final_rows:
            w.writerow(r)

    rep.note("\nPRE-POST CHECKLIST (verify in Editor's pending-changes pane before posting):")
    rep.note("  [ ] Exactly 56 pending changes: 52 location (26 del / 26 add) + 4 campaign edits")
    rep.note("  [ ] Exactly 6 campaigns touched")
    rep.note("  [ ] Texas appears NOWHERE in pending changes")
    rep.note("  [ ] Current tCPAs read Tier1 $47.25 / Tier3 $36.58 / CA $16.10 / FL $12.00 before edit")
    rep.note("  [ ] Tier2 and Tier5 tCPAs untouched")
    rep.note("  [ ] No keyword / ad / ad-group / budget / negative / audience rows pending")
    rep.note("  [ ] One session, morning; after posting, re-export and run --post-export")

    log_path = OUTPUT_DIR / f"run_log_{stamp}.txt"
    log_path.write_text("\n".join(rep.lines), encoding="utf-8")
    rep.dump()
    print(f"\nGATE 0 PASSED — FILES WRITTEN")
    print(f"  Import:  {import_path}")
    print(f"  Run log: {log_path}")


if __name__ == "__main__":
    main()
