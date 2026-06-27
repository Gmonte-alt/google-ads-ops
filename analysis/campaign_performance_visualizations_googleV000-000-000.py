"""
Campaign Performance Visualizations — Google Home Insurance
═══════════════════════════════════════════════════════════════════════
Generates a per-campaign weekly performance workbook covering the last
12 weeks of operating equilibrium. Designed to answer the leadership-
side question: "Why isn't the model finding more headroom?" — by
showing each campaign's recent trajectory on conversions, ROAS, CPA,
RPL, and profit. Each metric is plotted weekly so the eye picks up
trends without being distracted by daily noise.

PANEL LAYOUT (per campaign sheet, top to bottom):
  Panel 1: weekly conversions (bar chart)
  Panel 2: weekly ROAS + current tROAS reference line
  Panel 3: weekly CPA
  Panel 4: weekly RPL
  Panel 5: weekly profit (color-coded green/red bars)

Plus vertical lines at any change events within the window:
  - Red dashed: Target ROAS changes
  - Orange dotted: Budget changes

USAGE:
  python campaign_performance_visualizations_google.py
  python campaign_performance_visualizations_google.py --end-date 2026-04-24
  python campaign_performance_visualizations_google.py --weeks 16

OUTPUT: campaign_performance_google_<YYYY-MM-DD>.xlsx
  - Overview (small multiples: weekly profit charts on one sheet)
  - <CAMPAIGN_SHORT> sheets (one per campaign, 5-panel weekly chart)
"""

import argparse
import io
import sqlite3
import csv
import re
from datetime import date, datetime, timedelta
from pathlib import Path
import sys

import pandas as pd
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.drawing.image import Image as XLImage

# ── CONFIG ─────────────────────────────────────────────────────────
ACCOUNT_ID = 1  # Google Home Insurance
ACCOUNT_LABEL = "google_home"
DEFAULT_DB_PATH = Path(r"C:\Projects\google-ads-ops\db\paid_search.db")
DEFAULT_TARGETS_CSV = Path(r"C:\Projects\google-ads-ops\analysis\input\Change_history_report.csv")
DEFAULT_BUDGETS_CSV = Path(r"C:\Projects\google-ads-ops\analysis\input\Change_history_report__1_.csv")
DEFAULT_OUTPUT_DIR = Path(r"C:\Projects\google-ads-ops\analysis\output")

ENABLED_CAMPAIGN_IDS = [
    "23385188622",   # HOME INSURANCE-CALIFORNIA-NB-SEARCH-ALL-USA-STE
    "23394808297",   # HOME INSURANCE-TEXAS-NB-SEARCH-ALL-USA-STE
    "23394808300",   # HOME INSURANCE-FLORIDA-NB-SEARCH-ALL-USA-STE
    "23428389081",   # HOME INSURANCE-TIER ONE-NB-SEARCH-ALL-USA-STE
    "23496334752",   # HOME INSURANCE-TIER TWO-NB-SEARCH-ALL-USA-STE
    "23530860942",   # HOME INSURANCE-TIER THREE-NB-SEARCH-ALL-USA-STE
    "23586637678",   # HOME INSURANCE-TIER FOUR-NB-SEARCH-ALL-USA-STE
]

DEFAULT_WEEKS = 12

# Backfill buffer: number of days back from "today" to use as the default
# analysis end date. Google Ads offline-routed primary conversions can lag
# the click date by 1-2 days; running the chart against truly-yesterday
# data sometimes shows a phantom dip in the most recent week purely from
# unsettled attribution. The default of 1 day matches "yesterday"; bump
# to 2 if you regularly see suspicious recent-week dips. Override entirely
# with --end-date for reproducible historical runs.
DEFAULT_BACKFILL_BUFFER_DAYS = 1


def default_end_date() -> date:
    """Compute the default analysis end date dynamically. Equals
    today − DEFAULT_BACKFILL_BUFFER_DAYS, so a routine run on Monday
    May 4 with buffer=1 yields end_date = May 3 (Sunday) and the
    chart's most recent week is the just-completed Mon–Sun week.
    """
    return date.today() - timedelta(days=DEFAULT_BACKFILL_BUFFER_DAYS)


YEAR_START = date(2026, 1, 1)


def banner(text: str) -> None:
    print("═" * 70)
    print(text)
    print("═" * 70)


# ═══════════════════════════════════════════════════════════════════
# DB QUERIES
# ═══════════════════════════════════════════════════════════════════

def fetch_enabled_campaigns(conn) -> pd.DataFrame:
    """Enabled campaigns + their current tROAS. Filters to the 7
    NB-SEARCH-ALL-USA-STE campaigns from efficiency_frontier scope."""
    placeholders = ",".join("?" * len(ENABLED_CAMPAIGN_IDS))
    sql = f"""
    SELECT
        campaign_id,
        campaign_name,
        target_roas AS current_target
    FROM campaigns
    WHERE account_id = ?
      AND campaign_id IN ({placeholders})
      AND status = 'ENABLED'
    ORDER BY campaign_name
    """
    df = pd.read_sql(sql, conn, params=[ACCOUNT_ID] + ENABLED_CAMPAIGN_IDS)
    return df


def fetch_campaign_daily(conn, campaign_name: str,
                         start_date: date, end_date: date) -> pd.DataFrame:
    """Daily campaign metrics. Same dual-source pattern as the elasticity
    script and efficiency_frontier:
      • Spend from campaign_performance_daily (device='ALL' to skip the
        per-device breakdown rows, which would otherwise double-count).
      • Conversions from keyword_performance_daily aggregated to the
        campaign × date grain (device='ALL', geo_state='ALL'), summing
        primary_conv_count and primary_conv_value — the offline-routed
        source of truth for Google Home.

    Returns columns: report_date, spend_usd, conv, conv_value.
    """
    sql_cpd = """
    SELECT
        cpd.date AS report_date,
        SUM(cpd.cost_micros) / 1e6 AS spend_usd
    FROM campaign_performance_daily cpd
    JOIN campaigns c ON cpd.campaign_id = c.campaign_id
    WHERE c.account_id = ?
      AND c.campaign_name = ?
      AND cpd.date >= ?
      AND cpd.date <= ?
      AND cpd.device = 'ALL'
    GROUP BY cpd.date
    ORDER BY cpd.date
    """
    df_cpd = pd.read_sql(sql_cpd, conn, params=(ACCOUNT_ID, campaign_name,
                                                 start_date.isoformat(),
                                                 end_date.isoformat()))

    sql_kpd = """
    SELECT
        kpd.report_date AS report_date,
        SUM(kpd.primary_conv_count) AS conv,
        SUM(kpd.primary_conv_value) AS conv_value
    FROM keyword_performance_daily kpd
    JOIN keyword_instances ki ON ki.id = kpd.instance_id
    JOIN campaigns c ON ki.campaign_id = c.campaign_id
    WHERE c.account_id    = ?
      AND c.campaign_name = ?
      AND kpd.report_date >= ?
      AND kpd.report_date <= ?
      AND kpd.device      = 'ALL'
      AND kpd.geo_state   = 'ALL'
    GROUP BY kpd.report_date
    """
    df_kpd = pd.read_sql(sql_kpd, conn, params=(ACCOUNT_ID, campaign_name,
                                                 start_date.isoformat(),
                                                 end_date.isoformat()))

    df = df_cpd.merge(df_kpd, on="report_date", how="left")
    df["conv"] = df["conv"].fillna(0.0)
    df["conv_value"] = df["conv_value"].fillna(0.0)
    if not df.empty:
        df["report_date"] = pd.to_datetime(df["report_date"]).dt.date
    return df


# ═══════════════════════════════════════════════════════════════════
# WEEKLY AGGREGATION
# ═══════════════════════════════════════════════════════════════════

def aggregate_weekly(daily_df: pd.DataFrame) -> pd.DataFrame:
    """Roll up daily data to ISO weeks (Mon-Sun).

    Returns columns: week_start (Mon date), week_end (Sun date),
    spend, conv, conv_value, cpa, rpl, roas, profit.

    CPA, RPL, and ROAS are weighted ratios (sum-over-sum, NOT averages
    of daily values) — same convention as the Microsoft script and
    efficiency_frontier.
    """
    if daily_df.empty:
        return pd.DataFrame()

    df = daily_df.copy()
    df["report_date"] = pd.to_datetime(df["report_date"])
    df["week_start"] = df["report_date"] - pd.to_timedelta(
        df["report_date"].dt.weekday, unit="D"
    )
    df["week_start"] = df["week_start"].dt.date

    agg = df.groupby("week_start").agg(
        spend=("spend_usd", "sum"),
        conv=("conv", "sum"),
        conv_value=("conv_value", "sum"),
    ).reset_index()

    agg["week_end"] = agg["week_start"].apply(lambda d: d + timedelta(days=6))
    agg["cpa"] = agg.apply(lambda r: r["spend"] / r["conv"] if r["conv"] > 0 else None, axis=1)
    agg["rpl"] = agg.apply(lambda r: r["conv_value"] / r["conv"] if r["conv"] > 0 else None, axis=1)
    agg["roas"] = agg.apply(lambda r: r["conv_value"] / r["spend"] if r["spend"] > 0 else None, axis=1)
    agg["profit"] = agg["conv_value"] - agg["spend"]

    return agg


# ═══════════════════════════════════════════════════════════════════
# CHANGE-HISTORY LOADER (Google CSV format)
# ═══════════════════════════════════════════════════════════════════

def _parse_gads_datetime(s: str) -> datetime:
    """Google Ads CSV exports: 'Mar 30, 2026, 9:31:59 AM'."""
    return datetime.strptime(s.strip(), "%b %d, %Y, %I:%M:%S %p")


def _normalize_target_value(raw) -> float | None:
    """116% → 1.16, 1.16 → 1.16, 113 → 1.13, '' → None."""
    if raw is None:
        return None
    s = str(raw).strip().rstrip(",").rstrip('"').strip()
    if not s:
        return None
    if s.endswith("%"):
        try:
            return float(s.rstrip("%")) / 100.0
        except ValueError:
            return None
    try:
        v = float(s)
    except ValueError:
        return None
    if v > 10:
        return v / 100.0
    return v


def load_change_events(targets_csv: Path, budgets_csv: Path,
                       enabled_campaign_names: set[str],
                       start_date: date, end_date: date) -> pd.DataFrame:
    """Load tROAS + budget change events from the Google Ads change
    history CSVs, filtered to the visualization window AND tracked
    campaigns. Used to overlay vertical change-event lines on the
    per-campaign charts.

    Returns DataFrame with columns:
      change_date, campaign_name, attribute (TargetROAS|Budget),
      old_value, new_value
    """
    rows = []

    # ── Target ROAS changes (CAMPAIGN-level only) ──
    if targets_csv.exists():
        with open(targets_csv, encoding="utf-8-sig") as f:
            lines = f.read().splitlines()
        reader = csv.DictReader(lines[2:] if len(lines) > 2 else lines)
        for r in reader:
            ch = r.get("Changes", "") or ""
            if "Target ROAS" not in ch:
                continue
            try:
                ts = _parse_gads_datetime(r["Date & time"])
            except (ValueError, KeyError):
                continue
            cd = ts.date()
            if cd < start_date or cd > end_date:
                continue
            scope = "AG" if (r.get("Ad group") or "").strip() else "CAMPAIGN"
            if scope != "CAMPAIGN":
                continue
            cname = r.get("Campaign", "") or ""
            if cname not in enabled_campaign_names:
                continue
            m = re.search(r"changed from\s+(\S+?)\s+to\s+(\S+)", ch)
            if m:
                old = _normalize_target_value(m.group(1))
                new = _normalize_target_value(m.group(2))
            else:
                m2 = re.search(r"Target ROAS (?:is|was)\s+(\S+)", ch)
                if m2:
                    old, new = None, _normalize_target_value(m2.group(1))
                else:
                    old, new = None, None
            rows.append({
                "change_date": cd,
                "campaign_name": cname,
                "attribute": "TargetROAS",
                "old_value": old,
                "new_value": new,
            })
    else:
        print(f"  ℹ  Targets CSV not found at {targets_csv} — no target overlays")

    # ── Budget changes ──
    if budgets_csv.exists():
        with open(budgets_csv, encoding="utf-8-sig") as f:
            lines = f.read().splitlines()
        reader = csv.DictReader(lines[2:] if len(lines) > 2 else lines)
        for r in reader:
            ch = r.get("Changes", "") or ""
            if "Budget amount" not in ch:
                continue
            try:
                ts = _parse_gads_datetime(r["Date & time"])
            except (ValueError, KeyError):
                continue
            cd = ts.date()
            if cd < start_date or cd > end_date:
                continue
            cname = r.get("Campaign", "") or ""
            if cname not in enabled_campaign_names:
                continue
            m = re.search(r"from \$([\d,]+\.\d{2}) to \$([\d,]+\.\d{2})", ch)
            if m:
                old = float(m.group(1).replace(",", ""))
                new = float(m.group(2).replace(",", ""))
            else:
                old, new = None, None
            rows.append({
                "change_date": cd,
                "campaign_name": cname,
                "attribute": "Budget",
                "old_value": old,
                "new_value": new,
            })
    else:
        print(f"  ℹ  Budgets CSV not found at {budgets_csv} — no budget overlays")

    return pd.DataFrame(rows)


# ═══════════════════════════════════════════════════════════════════
# CHART BUILDERS
# ═══════════════════════════════════════════════════════════════════

def build_5panel_campaign_chart_png(
    weekly_df: pd.DataFrame,
    campaign_name: str,
    current_target: float | None,
    change_events: list[dict],
    window_start: date,
    window_end: date,
) -> bytes:
    """5-panel weekly chart for one campaign:
       Panel 1 (top):    weekly conversions (bar chart)
       Panel 2:          weekly ROAS + current tROAS reference line
       Panel 3:          weekly CPA
       Panel 4:          weekly RPL
       Panel 5 (bottom): weekly profit (bar chart, color-coded green/red)

    Plus vertical lines at any change events within the window
    (red dashed for Target ROAS, orange dotted for Budget).
    """
    if weekly_df.empty:
        fig, ax = plt.subplots(figsize=(11, 11))
        ax.text(0.5, 0.5, f"No data available for\n{campaign_name}\n"
                          f"between {window_start} and {window_end}",
                ha="center", va="center", fontsize=12, transform=ax.transAxes)
        ax.set_axis_off()
        buf = io.BytesIO()
        fig.savefig(buf, format="png", dpi=100, bbox_inches="tight")
        plt.close(fig)
        return buf.getvalue()

    df = weekly_df.copy()
    df["week_dt"] = pd.to_datetime(df["week_start"])
    bar_width = 6  # days; leaves a 1-day gap between weekly bars

    fig, (ax1, ax2, ax3, ax4, ax5) = plt.subplots(
        5, 1, figsize=(11, 12), sharex=True,
        gridspec_kw={"hspace": 0.18},
    )

    # ── Panel 1: weekly conversions ────────────────────────────────
    ax1.bar(df["week_dt"], df["conv"], width=bar_width,
            color="#4F81BD", alpha=0.75, edgecolor="#1F4E79", linewidth=0.5)
    for _, row in df.iterrows():
        if row["conv"] > 0:
            ax1.annotate(f"{int(round(row['conv']))}",
                         xy=(row["week_dt"], row["conv"]),
                         xytext=(0, 3), textcoords="offset points",
                         ha="center", fontsize=8, color="#1F4E79")
    ax1.set_ylabel("Weekly conversions", fontsize=10)
    ax1.set_title(f"{campaign_name}\n"
                  f"12-week performance: {window_start} to {window_end}",
                  fontsize=11, fontweight="bold", loc="left")
    ax1.grid(True, alpha=0.3, axis="y")

    # ── Panel 2: weekly ROAS + tROAS reference ─────────────────────
    ax2.plot(df["week_dt"], df["roas"], marker="o", color="#7030A0",
             linewidth=2, markersize=7, label="Weekly ROAS")
    if current_target is not None:
        ax2.axhline(current_target, color="#1F4E79", linestyle="-.", linewidth=1.5,
                    alpha=0.7, label=f"Current tROAS target ({current_target:.2f})")
    for _, row in df.iterrows():
        if pd.notna(row["roas"]):
            ax2.annotate(f"{row['roas']:.2f}",
                         xy=(row["week_dt"], row["roas"]),
                         xytext=(0, 5), textcoords="offset points",
                         ha="center", fontsize=8, color="#7030A0")
    ax2.set_ylabel("Weekly ROAS", fontsize=10)
    ax2.legend(loc="best", fontsize=8, framealpha=0.9)
    ax2.grid(True, alpha=0.3)

    # ── Panel 3: weekly CPA ────────────────────────────────────────
    ax3.plot(df["week_dt"], df["cpa"], marker="o", color="#C0504D",
             linewidth=2, markersize=7, label="Weekly CPA")
    for _, row in df.iterrows():
        if pd.notna(row["cpa"]):
            ax3.annotate(f"${row['cpa']:.0f}",
                         xy=(row["week_dt"], row["cpa"]),
                         xytext=(0, 5), textcoords="offset points",
                         ha="center", fontsize=8, color="#C0504D")
    ax3.set_ylabel("Weekly CPA ($)", fontsize=10)
    ax3.legend(loc="best", fontsize=8, framealpha=0.9)
    ax3.grid(True, alpha=0.3)

    # ── Panel 4: weekly RPL ────────────────────────────────────────
    ax4.plot(df["week_dt"], df["rpl"], marker="s", color="#9BBB59",
             linewidth=2, markersize=7, label="Weekly RPL")
    for _, row in df.iterrows():
        if pd.notna(row["rpl"]):
            ax4.annotate(f"${row['rpl']:.0f}",
                         xy=(row["week_dt"], row["rpl"]),
                         xytext=(0, 5), textcoords="offset points",
                         ha="center", fontsize=8, color="#5F7F2F")
    ax4.set_ylabel("Weekly RPL ($)", fontsize=10)
    ax4.legend(loc="best", fontsize=8, framealpha=0.9)
    ax4.grid(True, alpha=0.3)

    # ── Panel 5: weekly profit ─────────────────────────────────────
    profit_colors = ["#9BBB59" if p >= 0 else "#C0504D" for p in df["profit"]]
    ax5.bar(df["week_dt"], df["profit"], width=bar_width,
            color=profit_colors, alpha=0.75, edgecolor="#404040", linewidth=0.5)
    ax5.axhline(0, color="black", linewidth=0.8, alpha=0.5)
    for _, row in df.iterrows():
        if pd.notna(row["profit"]):
            offset = 8 if row["profit"] >= 0 else -12
            ax5.annotate(f"${row['profit']:,.0f}",
                         xy=(row["week_dt"], row["profit"]),
                         xytext=(0, offset), textcoords="offset points",
                         ha="center", fontsize=8,
                         color="#5F7F2F" if row["profit"] >= 0 else "#8B0000")
    ax5.set_ylabel("Weekly profit ($)", fontsize=10)
    ax5.set_xlabel("Week starting (Mon)", fontsize=10)
    ax5.grid(True, alpha=0.3, axis="y")

    # Overlay change-event markers on all panels
    for ev in change_events:
        change_dt = pd.Timestamp(ev["change_date"])
        attr = str(ev.get("attribute", ""))
        if "TargetROAS" in attr:
            color, ls = "red", "--"
        elif "Budget" in attr:
            color, ls = "#FF8C00", ":"  # orange dotted for budget
        else:
            color, ls = "gray", "--"
        for ax in (ax1, ax2, ax3, ax4, ax5):
            ax.axvline(change_dt, color=color, linestyle=ls, linewidth=1.2, alpha=0.6)

    ax5.xaxis.set_major_locator(mdates.WeekdayLocator(byweekday=0))  # Mondays
    ax5.xaxis.set_major_formatter(mdates.DateFormatter("%b %d"))
    plt.setp(ax5.xaxis.get_majorticklabels(), rotation=30, ha="right")

    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=100, bbox_inches="tight")
    plt.close(fig)
    return buf.getvalue()


def build_overview_small_multiples_png(
    campaign_weekly: dict[str, pd.DataFrame],
    window_start: date,
    window_end: date,
) -> bytes:
    """One sheet showing all enabled campaigns' weekly profit as small
    multiples. Layout: up to 4 columns, rows as needed. Each panel is a
    single campaign's weekly profit bar chart with independent y-axis
    (profit magnitudes differ enough that shared scale would make small
    campaigns invisible). Each panel labels its total window profit.
    """
    n_campaigns = len(campaign_weekly)
    if n_campaigns == 0:
        fig, ax = plt.subplots(figsize=(11, 8.5))
        ax.text(0.5, 0.5, "No campaign data available",
                ha="center", va="center", fontsize=12, transform=ax.transAxes)
        ax.set_axis_off()
        buf = io.BytesIO()
        fig.savefig(buf, format="png", dpi=100, bbox_inches="tight")
        plt.close(fig)
        return buf.getvalue()

    n_cols = min(4, n_campaigns)
    n_rows = max(1, (n_campaigns + n_cols - 1) // n_cols)

    fig, axes = plt.subplots(
        n_rows, n_cols, figsize=(16, 3.2 * n_rows),
        sharex=True, gridspec_kw={"hspace": 0.55, "wspace": 0.25},
    )
    if n_rows * n_cols == 1:
        axes = [axes]
    else:
        axes = axes.flatten() if hasattr(axes, "flatten") else list(axes)

    # Sort campaigns by total window profit descending so biggest are top-left
    sorted_items = sorted(
        campaign_weekly.items(),
        key=lambda kv: -(kv[1]["profit"].sum() if not kv[1].empty else 0),
    )

    for idx, (campaign_name, weekly_df) in enumerate(sorted_items):
        ax = axes[idx]
        if weekly_df.empty:
            ax.text(0.5, 0.5, "(no data)", ha="center", va="center",
                    transform=ax.transAxes, fontsize=9, color="gray")
            ax.set_axis_off()
            continue

        weeks = pd.to_datetime(weekly_df["week_start"])
        profits = weekly_df["profit"].values
        colors = ["#9BBB59" if p >= 0 else "#C0504D" for p in profits]
        ax.bar(weeks, profits, width=6, color=colors, alpha=0.8,
               edgecolor="#404040", linewidth=0.4)
        ax.axhline(0, color="black", linewidth=0.5, alpha=0.5)

        # Title: short campaign name + total profit
        short_name = (campaign_name
                      .replace("HOME INSURANCE-", "")
                      .split("-NB-SEARCH")[0])
        total_profit = profits.sum()
        ax.set_title(f"{short_name}\n${total_profit:,.0f} total",
                     fontsize=9, fontweight="bold", loc="left")
        ax.tick_params(axis="x", labelsize=7)
        ax.tick_params(axis="y", labelsize=7)
        ax.xaxis.set_major_locator(mdates.MonthLocator())
        ax.xaxis.set_major_formatter(mdates.DateFormatter("%b"))
        ax.grid(True, alpha=0.25, axis="y")
        ax.yaxis.set_major_formatter(plt.FuncFormatter(lambda v, p: f"${v:,.0f}"))

    # Hide unused subplots
    for idx in range(len(sorted_items), len(axes)):
        axes[idx].set_axis_off()

    fig.suptitle(
        f"Weekly Profit by Campaign — {window_start} to {window_end} "
        f"(sorted by total window profit)",
        fontsize=13, fontweight="bold", y=0.995,
    )

    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=100, bbox_inches="tight")
    plt.close(fig)
    return buf.getvalue()


# ═══════════════════════════════════════════════════════════════════
# MAIN
# ═══════════════════════════════════════════════════════════════════

def parse_args():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--db-path", default=str(DEFAULT_DB_PATH))
    p.add_argument("--targets-csv", default=str(DEFAULT_TARGETS_CSV),
                   help="Google Ads target-change CSV (for tROAS overlays)")
    p.add_argument("--budgets-csv", default=str(DEFAULT_BUDGETS_CSV),
                   help="Google Ads budget-change CSV (for budget overlays)")
    p.add_argument("--output-dir", default=str(DEFAULT_OUTPUT_DIR))
    p.add_argument("--end-date", default=None,
                   help=f"Last day of analysis window (ISO YYYY-MM-DD). "
                        f"Default: today − {DEFAULT_BACKFILL_BUFFER_DAYS}d "
                        f"(yesterday). Override for reproducible historical "
                        f"runs.")
    p.add_argument("--weeks", type=int, default=DEFAULT_WEEKS,
                   help=f"Number of weeks back from --end-date. Default: "
                        f"{DEFAULT_WEEKS}")
    return p.parse_args()


def main():
    args = parse_args()
    db_path = Path(args.db_path)
    targets_csv = Path(args.targets_csv)
    budgets_csv = Path(args.budgets_csv)
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    end_date = (date.fromisoformat(args.end_date)
                if args.end_date else default_end_date())
    n_weeks = args.weeks
    end_week_monday = end_date - timedelta(days=end_date.weekday())
    window_start = end_week_monday - timedelta(weeks=n_weeks - 1)
    window_end = end_date

    banner("CAMPAIGN PERFORMANCE VISUALIZATIONS — Google Home")
    print(f"  DB:                  {db_path}")
    print(f"  Targets CSV:         {targets_csv}")
    print(f"  Budgets CSV:         {budgets_csv}")
    print(f"  Output:              {output_dir}")
    print(f"  Window:              {window_start} (Mon) to {window_end}")
    print(f"  Weeks:               {n_weeks}")
    print()

    if not db_path.exists():
        print(f"ERROR: DB not found at {db_path}")
        sys.exit(1)

    conn = sqlite3.connect(db_path)
    enabled = fetch_enabled_campaigns(conn)
    if len(enabled) != len(ENABLED_CAMPAIGN_IDS):
        missing = set(ENABLED_CAMPAIGN_IDS) - set(enabled["campaign_id"].astype(str))
        print(f"⚠  Expected {len(ENABLED_CAMPAIGN_IDS)} enabled campaigns, DB returned {len(enabled)}. "
              f"Missing: {missing}")
    enabled_names = set(enabled["campaign_name"].tolist())
    print(f"Found {len(enabled)} enabled campaigns")
    print()

    # Load change events filtered to window + tracked campaigns
    change_events_df = load_change_events(targets_csv, budgets_csv,
                                          enabled_names, window_start, window_end)
    if change_events_df.empty:
        n_target = n_budget = 0
    else:
        n_target = (change_events_df["attribute"] == "TargetROAS").sum()
        n_budget = (change_events_df["attribute"] == "Budget").sum()
    print(f"Loaded {len(change_events_df)} change events within window "
          f"({n_target} tROAS, {n_budget} budget)")
    print()

    # Process each campaign: fetch daily, aggregate weekly
    campaign_weekly: dict[str, pd.DataFrame] = {}
    for _, camp in enabled.iterrows():
        cname = camp["campaign_name"]
        daily = fetch_campaign_daily(conn, cname, window_start, window_end)
        weekly = aggregate_weekly(daily)
        campaign_weekly[cname] = weekly

    # Build workbook
    print("Building workbook...")
    wb = Workbook()
    wb.remove(wb.active)

    # ── Overview sheet (small multiples) ──
    ws = wb.create_sheet("Overview")
    ws.cell(row=1, column=1, value="WEEKLY PROFIT — ALL CAMPAIGNS").font = Font(bold=True, size=14)
    ws.cell(row=2, column=1, value=f"Window: {window_start} to {window_end} "
                                    f"({n_weeks} weeks). Sorted by total window profit.")
    print("  Generating small-multiples overview...")
    overview_png = build_overview_small_multiples_png(
        campaign_weekly, window_start, window_end,
    )
    img = XLImage(io.BytesIO(overview_png))
    ws.add_image(img, "A4")
    print("  ✓ Overview")

    # ── Per-campaign 5-panel sheets ──
    sheet_count = 0
    print("  Generating per-campaign 5-panel charts...")
    used_sheet_names = {"Overview"}
    for _, camp in enabled.iterrows():
        cname = camp["campaign_name"]
        current_target = camp["current_target"]
        if pd.isna(current_target):
            current_target = None
        weekly_df = campaign_weekly.get(cname, pd.DataFrame())
        # Filter change events to this campaign
        if change_events_df.empty:
            camp_events = []
        else:
            camp_events = change_events_df[
                change_events_df["campaign_name"] == cname
            ].to_dict("records")

        # Sheet name: short campaign name (Excel limit 31 chars)
        short_name = (cname.replace("HOME INSURANCE-", "")
                      .split("-NB-SEARCH")[0][:25])
        sheet_name = short_name[:31]
        # Avoid duplicate sheet names defensively
        if sheet_name in used_sheet_names:
            sheet_name = f"{short_name[:28]}_2"[:31]
        used_sheet_names.add(sheet_name)

        try:
            png_bytes = build_5panel_campaign_chart_png(
                weekly_df=weekly_df,
                campaign_name=cname,
                current_target=current_target,
                change_events=camp_events,
                window_start=window_start,
                window_end=window_end,
            )
            ws_chart = wb.create_sheet(sheet_name)
            img = XLImage(io.BytesIO(png_bytes))
            ws_chart.add_image(img, "A1")
            sheet_count += 1
            print(f"  ✓ {sheet_name}")
        except Exception as exc:
            print(f"  ✗ {sheet_name}: chart failed ({exc})")

    conn.close()
    print(f"Generated {sheet_count} per-campaign chart sheets")

    out_path = output_dir / f"campaign_performance_google_{date.today().isoformat()}.xlsx"
    wb.save(out_path)
    print()
    print(f"  ✓ Wrote {out_path}")


if __name__ == "__main__":
    main()
