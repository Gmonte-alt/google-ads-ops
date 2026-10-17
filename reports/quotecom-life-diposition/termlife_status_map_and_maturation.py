"""
termlife_status_map_and_maturation.py
=====================================
Builds #1 and #2 of the SQLife disposition reporting layer
(see term_life_disposition_feed_notes.md).

  #1 STATUS MAP  - life_status_map: maps the raw Five9 status vocabulary
                   (77 strings and growing) onto a small PROVISIONAL stage
                   taxonomy via ordered pattern rules. Manual overrides survive
                   refreshes (set map_source='manual' on a row and the rules
                   will never touch it again). Unmapped statuses are inserted
                   as stage='UNMAPPED' and reported loudly.
                   Also (re)creates v_life_disposition_current_staged.

  #2 MATURATION  - cross-sectional maturation curve from the CURRENT batch:
                   SQO / CRI / revenue rates by lead AGE bucket, overall and
                   per segment (M-channel web leads). Answers: "how old must a
                   cohort be before its numbers are trustworthy?"
                   Output: Excel workbook + console summary incl. an estimated
                   maturity horizon.

PROVISIONAL taxonomy (until the partner handoff decodes the acronyms):
  NEW       lead created, not yet worked        (New Real Time, Send to SQT)
  QUEUE     in dialer queue                     (SQT Five9 Assignment, ...)
  CALLBACK  contact attempt cycle               (Callback - *)
  TRANSFER  warm-transfer path                  (LHA / SQ Other transfers)
  PIPELINE  post-app / post-sale processing     (CRI-*, COV, OFR, DEL, DEC,
                                                 Enroller, Upsell, Stuck in CRI)
  SOLD      explicit sale/enrollment statuses   (Sale Made, Enrollment Complete...)
  DEAD      terminal negative (+ sub_reason)    (1 - No Sale - *, DNC, No Sale)
  UNMAPPED  never-seen string -> review

HOUSE RULES honored throughout:
  - reads go through v_life_disposition_current (batch_date pinned by the view)
  - CRI / revenue truth = the cris / capped_cri_dollars FLAGS, never the status
    text (evidence: CRI dollars appear on 'No Sale'/'DNC'/'Upsell' rows)
"""

import re
import sqlite3
import sys
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------- CONFIG ----
DB_PATH  = Path(r"C:\Projects\google-ads-ops\db\paid_search.db")
OUT_DIR  = Path(r"C:\Projects\google-ads-ops\reports\quotecom-life-diposition\output")

AGE_BUCKETS = [("0-6", 0, 6), ("7-13", 7, 13), ("14-20", 14, 20),
               ("21-27", 21, 27), ("28-34", 28, 34), ("35-41", 35, 41),
               ("42-55", 42, 55), ("56-69", 56, 69), ("70-89", 70, 89)]

# reporting views (layouts confirmed 9/24): lead-creation basis, all channels
SEG_LABELS = {"CA": "Low Income", "CB": "Medium Income",
              "CC": "High Income", "QA": "Organic"}
SEG_ORDER = ["CC", "CB", "CA", "QA"]          # High -> Medium -> Low -> Organic
# Economics (corrected 10/6/26 -- Gilbert: Looker "Spend" = what SQ PAYS Quote):
#   SQ pays Quote PER OP at segment-tiered prices (~$150 High / $87 Med / $31
#   Low per op, Aug'26); avg_cost in the feed = account-BLENDED average price,
#   NOT the row's own price -- never use it for segment-level revenue.
#   SQ Actual Rev = 33.00% x CRI Prem; SQ margin line in Looker uses 33.75%.
#   Policy KPI = CRI Prem / payment-to-Quote = SQ's return multiple
#   (SQ break-even ~2.96, target 4). Quote's own margin = payment - Google
#   spend. Both live in the campaign-grain script (termlife_policy_kpi.py).
SQ_REV_SHARE = 0.33

# Per-op payment RATE CARD (what SQ pays Quote per M-channel Op), derived from
# the Looker by-advert export: WebM Spend / WebM Ops, per segment-month.
#   Aug'26: High 100,576/672  Med 44,398/513  Low 23,919/772
#   Sep'26: High 112,768/783  Med 51,866/615  Low 28,198/895
#   Oct'26: High  14,704/99   Med  7,547/87   Low  3,823/115
# Months not listed use "default" (latest known). Organic has no rate (SQ's
# Search view excludes it; organic lead sales ending per 9/25 notes). Referral
# payments exist on another basis (Looker shows R Spend with 0 Ops) and
# Transfers show $0 Spend -- both excluded from the estimate.
# REPLACE with the official rate card from the AM when available; until then
# refresh monthly from the Looker by-advert export.
SEG_OP_RATES = {
    "CC": {"2026-08": 149.67, "2026-09": 144.02, "2026-10": 148.53,
           "default": 148.53},
    "CB": {"2026-08": 86.55, "2026-09": 84.34, "2026-10": 86.75,
           "default": 86.75},
    "CA": {"2026-08": 30.98, "2026-09": 31.51, "2026-10": 33.24,
           "default": 33.24},
}

METRIC_HEADERS = ["Leads", "Ops", "SQOs", "CRIs", "CRI Prem",
                  "Spend (est)", "Policy KPI (est)", "SQ Rev (33%)",
                  "Op-SQO %", "SQO-CRI %", "Op-CRI %", "Avg Prem"]
NCOLS = 2 + len(METRIC_HEADERS)            # label cols + metric cols

# workbook color scheme (hex, no #). One family per segment, reused on all
# three reporting sheets so a segment is always the same color.
SEG_FILL_HEX = {"CC": "DDEBF7",            # High Income  - light blue
                "CB": "E2EFDA",            # Medium Income - light green
                "CA": "FCE4D6",            # Low Income   - light orange
                "QA": "E4DFEC"}            # Organic      - light purple
DEFAULT_SEG_FILL = "F2F2F2"                # any future segment - light gray
HEADER_HEX = "44546A"                      # title bar + column headers (navy)
TOTAL_HEX = "D9D9D9"                       # Daily 'Total' rows (gray)

RULE_VERSION = "rule:v1"

# ------------------------------------------------------- #1  STATUS RULES ---
DEAD_RE = re.compile(r"^1 - No Sale - NEW/PKG - (?P<reason>.+)$", re.I)

PIPELINE_PREFIXES = [  # (prefix, sub_reason label)
    ("1 - CRI", "CRI"), ("CRI -", "CRI"), ("Stuck in CRI", "CRI"),
    ("COV", "COV"), ("OFR/DEL", "OFR/DEL"), ("OFR", "OFR"),
    ("No OFA", "DEL"), ("DEC", "DEC"), ("Enroller", "Enroller"),
    ("Upsell", "Upsell"), ("Upsold", "Upsell"), ("UL Follow-Up", "UL"),
]


def classify(status: str) -> tuple[str, str | None]:
    """Return (stage, sub_reason) for a raw status string. Order matters."""
    s = (status or "").strip()
    if not s:
        return "UNMAPPED", None
    # 'SI Pod - X' behaves like X (SI = likely Simplified Issue pod)
    if s.lower().startswith("si pod - "):
        stage, sub = classify(s[9:])
        return stage, sub or "SI Pod"
    if re.fullmatch(r"new real time|send to sqt", s, re.I):
        return "NEW", None
    if re.fullmatch(r"sqt five9 assignment", s, re.I):
        return "QUEUE", None
    if re.fullmatch(r"failed to send to sqt", s, re.I):
        return "QUEUE", "routing issue"
    if re.fullmatch(r"dnc", s, re.I):
        return "DEAD", "Do Not Call"
    m = DEAD_RE.match(s)
    if m:
        return "DEAD", m["reason"].strip()
    # explicit sold BEFORE pipeline ('Del Sale Made') and BEFORE generic checks
    if re.fullmatch(r"sale made|enrollment complete|del sale made", s, re.I):
        return "SOLD", None
    if re.fullmatch(r"upsold", s, re.I):
        return "SOLD", "Upsell"
    # transfers (strip the XX. prefix pattern implicitly via search)
    if "lha transfer" in s.lower() or s.lower().startswith("no callback - transferred"):
        return "TRANSFER", s.split("-")[-1].strip() if "-" in s else None
    # post-app pipeline BEFORE the generic Callback rule (CRI - Callback - ...)
    for prefix, sub in PIPELINE_PREFIXES:
        if s.lower().startswith(prefix.lower()):
            return "PIPELINE", sub
    if s.lower().startswith("callback"):
        return "CALLBACK", None
    if re.fullmatch(r"no sale", s, re.I):
        return "DEAD", "Unspecified"
    return "UNMAPPED", None   # NTO and any new strings land here for review


def refresh_status_map(con: sqlite3.Connection) -> dict:
    con.execute("""
        CREATE TABLE IF NOT EXISTS life_status_map (
            status_raw  TEXT PRIMARY KEY,
            stage       TEXT NOT NULL,
            sub_reason  TEXT,
            map_source  TEXT NOT NULL,   -- 'rule:vN' or 'manual'
            first_seen_batch TEXT,
            updated_at  TEXT
        )""")
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    stats = {"new": 0, "updated": 0, "manual_kept": 0, "unmapped": []}
    rows = con.execute("""
        SELECT status, MIN(batch_date) FROM life_disposition_snapshot
        GROUP BY status""").fetchall()
    existing = {r[0]: (r[1], r[2], r[3]) for r in con.execute(
        "SELECT status_raw, stage, sub_reason, map_source FROM life_status_map")}
    for status, first_batch in rows:
        stage, sub = classify(status)
        if status not in existing:
            con.execute("""INSERT INTO life_status_map VALUES (?,?,?,?,?,?)""",
                        (status, stage, sub, RULE_VERSION, first_batch, now))
            stats["new"] += 1
        else:
            cur_stage, cur_sub, src = existing[status]
            if src == "manual":
                stats["manual_kept"] += 1
            elif (cur_stage, cur_sub) != (stage, sub):
                con.execute("""UPDATE life_status_map
                               SET stage=?, sub_reason=?, map_source=?, updated_at=?
                               WHERE status_raw=?""",
                            (stage, sub, RULE_VERSION, now, status))
                stats["updated"] += 1
        if stage == "UNMAPPED":
            stats["unmapped"].append(status)
    con.execute("DROP VIEW IF EXISTS v_life_disposition_current_staged")
    con.execute("""
        CREATE VIEW v_life_disposition_current_staged AS
        SELECT s.*, COALESCE(m.stage,'UNMAPPED') AS stage, m.sub_reason
        FROM v_life_disposition_current s
        LEFT JOIN life_status_map m ON m.status_raw = s.status""")
    con.commit()
    return stats


# ------------------------------------------------------ #2  MATURATION ------
def f(v) -> float:
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def bucket_of(age: int):
    for label, lo, hi in AGE_BUCKETS:
        if lo <= age <= hi:
            return label
    return None


def pull_current(con):
    batch = con.execute(
        "SELECT MAX(batch_date) FROM life_disposition_snapshot").fetchone()[0]
    if not batch:
        sys.exit("snapshot table is empty — run the ingest first")
    bd = date.fromisoformat(batch)
    rows = con.execute("""
        SELECT source_code, day, sqo, cris, capped_cri_dollars, stage, bound_date
        FROM v_life_disposition_current_staged""").fetchall()
    out = []
    for code, day, sqo, cris, rev, stage, bound in rows:
        try:
            age = (bd - date.fromisoformat((day or "").strip())).days
        except ValueError:
            continue
        b = bucket_of(age)
        if b is None:
            continue
        cri = 1 if f(cris) > 0 else 0
        lag = None
        if cri:
            try:
                lag = (date.fromisoformat((bound or "").strip())
                       - date.fromisoformat(day.strip())).days
            except ValueError:
                lag = None      # CRI without a parseable bound_date
        out.append(dict(seg=code[:2], ch=code[2], bucket=b,
                        sqo=1 if f(sqo) > 0 else 0,
                        cri=cri, lag=lag,
                        rev=f(rev), stage=stage))
    return batch, out


def aggregate(rows, seg=None):
    """bucket -> dict(leads, sqo, cri, rev). M-channel only (web leads)."""
    agg = {label: dict(leads=0, sqo=0, cri=0, rev=0.0) for label, *_ in AGE_BUCKETS}
    for r in rows:
        if r["ch"] != "M":
            continue
        if seg and r["seg"] != seg:
            continue
        a = agg[r["bucket"]]
        a["leads"] += 1
        a["sqo"] += r["sqo"]
        a["cri"] += r["cri"]
        a["rev"] += r["rev"]
    return agg


def maturity_horizon(agg):
    """First bucket whose CRI% reaches 80% of the mature plateau (last 2 buckets)."""
    labels = [b[0] for b in AGE_BUCKETS]
    rates = {l: (agg[l]["cri"] / agg[l]["leads"] if agg[l]["leads"] else 0.0)
             for l in labels}
    plat_buckets = [l for l in labels[-2:] if agg[l]["leads"] > 0]
    if not plat_buckets:
        return None, rates, 0.0
    plateau = sum(rates[l] for l in plat_buckets) / len(plat_buckets)
    if plateau == 0:
        return None, rates, plateau
    for l in labels:
        if agg[l]["leads"] and rates[l] >= 0.8 * plateau:
            return l, rates, plateau
    return None, rates, plateau


def lag_distribution(rows):
    """Distribution of CRI lag (bound_date − lead day) for M-channel CRIs.
    Unlike the maturation curve, a CRI's lag never changes across batches:
    it answers 'when does the money arrive', not 'how complete is a cohort'."""
    lags = [r["lag"] for r in rows if r["ch"] == "M" and r["cri"]]
    known = sorted(l for l in lags if l is not None and 0 <= l <= 89)
    neg = sum(1 for l in lags if l is not None and l < 0)
    over = sum(1 for l in lags if l is not None and l > 89)
    unknown = sum(1 for l in lags if l is None)
    counts = {label: 0 for label, *_ in AGE_BUCKETS}
    for l in known:
        counts[bucket_of(l)] += 1

    def q(pct):
        return known[min(len(known) - 1, int(pct * len(known)))] if known else None

    return dict(counts=counts, n=len(known), neg=neg, over=over,
                unknown=unknown, median=q(0.5), p80=q(0.8))


# ------------------------------------------------- period reporting views ---
# Definitions (confirmed): Leads = delivered rows; Ops = SUM(op_count);
# SQOs/CRIs = flag sums; CRI Prem = SUM(capped_cri_dollars); all three
# channels (M/R/T) rolled up per segment; time axis = lead-creation `day`,
# read from the CURRENT batch (cohort restatement, not booked-by-date).

def pull_period_rows(con):
    rows = con.execute("""
        SELECT source_code, day, op_count, sqo, cris, capped_cri_dollars
        FROM v_life_disposition_current""").fetchall()
    out, skipped = [], 0
    for code, day, op, sqo, cris, rev in rows:
        try:
            d = date.fromisoformat((day or "").strip())
        except ValueError:
            skipped += 1
            continue
        out.append(dict(seg=code[:2], ch=code[2], d=d, ops=int(f(op)),
                        sqo=1 if f(sqo) > 0 else 0,
                        cri=1 if f(cris) > 0 else 0, rev=f(rev)))
    return out, skipped


def seg_label(s):
    return SEG_LABELS.get(s, s)


def _seg_list(rows):
    present = {r["seg"] for r in rows}
    return [s for s in SEG_ORDER if s in present] + sorted(present - set(SEG_ORDER))


def _zero_acc():
    return dict(leads=0, ops=0, ops_m=0, sqo=0, cri=0, rev=0.0)


def _acc(a, r):
    a["leads"] += 1
    a["ops"] += r["ops"]
    if r.get("ch") == "M":
        a["ops_m"] += r["ops"]
    a["sqo"] += r["sqo"]
    a["cri"] += r["cri"]
    a["rev"] += r["rev"]


def _est_spend(seg, ym, a):
    """Estimated payment to Quote: M-channel Ops x per-op rate for that month.
    None when no rate exists (Organic / unknown segments)."""
    rates = SEG_OP_RATES.get(seg)
    if not rates or not a["ops_m"]:
        return None
    return round(a["ops_m"] * rates.get(ym, rates["default"]), 2)


def _merge(t, a):
    for k in t:
        t[k] += a[k]


def _write_metrics(ws, r, c, a, spend=None, bold=False):
    def put(col, val, fmt=None):
        cell = ws.cell(r, col, val)
        if fmt:
            cell.number_format = fmt
        if bold:
            cell.font = BOLD
    put(c, a["leads"])
    put(c + 1, a["ops"])
    put(c + 2, a["sqo"])
    put(c + 3, a["cri"])
    put(c + 4, round(a["rev"], 2), "$#,##0")
    put(c + 5, spend, "$#,##0")
    put(c + 6, a["rev"] / spend if spend else None, "0.00")
    put(c + 7, round(a["rev"] * SQ_REV_SHARE, 2), "$#,##0")
    put(c + 8, a["sqo"] / a["ops"] if a["ops"] else None, "0.0%")
    put(c + 9, a["cri"] / a["sqo"] if a["sqo"] else None, "0.0%")
    put(c + 10, a["cri"] / a["ops"] if a["ops"] else None, "0.0%")
    put(c + 11, a["rev"] / a["cri"] if a["cri"] else None, "$#,##0")


THIN = Side(style="thin", color="BFBFBF")
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _fill(hex_):
    return PatternFill("solid", fgColor=hex_)


def _seg_fill(seg):
    return SEG_FILL_HEX.get(seg, DEFAULT_SEG_FILL)


def _title_block(ws, title, info_lines):
    """Title bar (row 1) + italic info lines. Returns the header row index."""
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=NCOLS)
    c = ws.cell(1, 1, title)
    c.font = Font(bold=True, size=13, color="FFFFFF")
    c.fill = _fill(HEADER_HEX)
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 22
    r = 2
    for line in info_lines:
        ic = ws.cell(r, 1, line)
        ic.font = Font(italic=True, size=9, color="595959")
        r += 1
    return r + 1                            # one blank row, then the header


def _table_header(ws, header_row, first_cols):
    for j, h in enumerate(list(first_cols) + METRIC_HEADERS, 1):
        c = ws.cell(header_row, j, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = _fill(HEADER_HEX)
        c.border = CELL_BORDER
        c.alignment = Alignment(horizontal="center")
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 18
    for i in range(3, NCOLS + 1):
        ws.column_dimensions[get_column_letter(i)].width = 11
    ws.freeze_panes = f"A{header_row + 1}"  # title/info/header stay visible


def _style_row(ws, r, hex_, bold=False):
    for j in range(1, NCOLS + 1):
        c = ws.cell(r, j)
        c.fill = _fill(hex_)
        c.border = CELL_BORDER
        if bold:
            c.font = BOLD


def _feed_info(rows, bd):
    min_day = min(x["d"] for x in rows)
    max_day = max(x["d"] for x in rows)
    return min_day, max_day, [
        f"Source: SQLife Term disposition feed (SQL_GenericDisposition_Term_Clicks) — "
        f"daily snapshot batch {bd.isoformat()}, rolling 90-day window "
        f"{min_day.isoformat()} \u2192 {max_day.isoformat()}",
        "Basis: lead-creation date \u00b7 segments roll up Web + Referral + Transfer "
        "channels \u00b7 exact-duplicate feed rows removed on load",
        "Leads = delivered rows \u00b7 Ops = \u03a3 OP_Count \u00b7 SQOs/CRIs = flag "
        "sums \u00b7 CRI Prem = \u03a3 Capped_Cri$ \u00b7 '*' = partial period at "
        "window edge \u00b7 (MTD)/(WTD) = period in progress",
        "Spend (est) = M-channel Ops \u00d7 per-op RATE CARD (from Looker by-advert "
        "Aug\u2013Oct'26: ~$149/$86/$31 per op High/Med/Low; Organic & Referral/Transfer "
        "payments excluded \u2014 no rate). ESTIMATE: expect \u00b11\u20132% vs Looker.",
        "Policy KPI (est) = CRI Prem \u00f7 Spend (est) \u00b7 SQ break-even \u22482.96, "
        "target 4 \u00b7 Daily Totals include Organic prem but no Organic spend (Looker "
        "excludes Organic entirely) \u00b7 Quote-side margin \u2192 termlife_policy_kpi.py",
        "Recent periods are immature \u2014 CRIs still accruing; see the Maturation "
        "sheets for the horizon before comparing young cohorts to mature ones.",
    ]


def _monthly_sheet(wb, rows, bd):
    ws = wb.create_sheet("Monthly by Income")
    min_day, _, info = _feed_info(rows, bd)
    hr = _title_block(ws, "Term Life \u2014 Monthly Performance by Income Segment", info)
    _table_header(ws, hr, ["Segment", "Month"])
    groups = {}
    for x in rows:
        k = (x["seg"], x["d"].strftime("%Y-%m"))
        groups.setdefault(k, _zero_acc())
        _acc(groups[k], x)
    months = sorted({m for _, m in groups})
    cur = bd.strftime("%Y-%m")

    def mlabel(m):
        lab = datetime.strptime(m, "%Y-%m").strftime("%b %Y")
        if m == months[0] and min_day.day != 1:
            lab += " *"
        if m == cur:
            lab += " (MTD)"
        return lab

    segs = _seg_list(rows)
    r = hr + 1
    for seg in segs:
        for m in months:
            a = groups.get((seg, m)) or _zero_acc()
            ws.cell(r, 1, seg_label(seg))
            ws.cell(r, 2, mlabel(m))
            _write_metrics(ws, r, 3, a, spend=_est_spend(seg, m, a))
            _style_row(ws, r, _seg_fill(seg))
            r += 1
        r += 1                              # blank spacer between segments

    # Total section: all segments summed; Spend/KPI cover rated (paid) segs only
    for m in months:
        tot, tot_spend = _zero_acc(), 0.0
        for seg in segs:
            a = groups.get((seg, m)) or _zero_acc()
            _merge(tot, a)
            tot_spend += _est_spend(seg, m, a) or 0.0
        ws.cell(r, 1, "Total")
        ws.cell(r, 2, mlabel(m))
        _write_metrics(ws, r, 3, tot, spend=tot_spend or None)
        _style_row(ws, r, TOTAL_HEX, bold=True)
        r += 1


def _weekly_sheet(wb, rows, bd):
    ws = wb.create_sheet("Weekly by Income")
    min_day, _, info = _feed_info(rows, bd)
    hr = _title_block(ws, "Term Life \u2014 Weekly Performance by Income Segment "
                          "(weeks start Monday)", info)
    _table_header(ws, hr, ["Week of", "Segment"])

    def wk(d):
        return d - timedelta(days=d.weekday())

    groups = {}
    for x in rows:
        k = (wk(x["d"]).isoformat(), x["seg"])
        groups.setdefault(k, _zero_acc())
        _acc(groups[k], x)
    weeks = sorted({w for w, _ in groups})
    cur_week = wk(bd - timedelta(days=1)).isoformat()
    segs = _seg_list(rows)
    r = hr + 1
    for w in weeks:
        label = w
        if w == weeks[0] and min_day.isoformat() != w:
            label += " *"
        if w == cur_week:
            label += " (WTD)"
        tot, tot_spend = _zero_acc(), 0.0
        for seg in segs:
            a = groups.get((w, seg)) or _zero_acc()
            sp = _est_spend(seg, w[:7], a)
            ws.cell(r, 1, label)
            ws.cell(r, 2, seg_label(seg))
            _write_metrics(ws, r, 3, a, spend=sp)
            _merge(tot, a)
            tot_spend += sp or 0.0
            _style_row(ws, r, _seg_fill(seg))
            r += 1
        ws.cell(r, 1, label)
        ws.cell(r, 2, "Total")
        _write_metrics(ws, r, 3, tot, spend=tot_spend or None)
        _style_row(ws, r, TOTAL_HEX, bold=True)
        r += 1
        r += 1                              # blank spacer between weeks


def _daily_sheet(wb, rows, bd):
    ws = wb.create_sheet("Daily stats")
    end = bd - timedelta(days=1)            # feed covers through batch-1
    start = end.replace(day=1)
    _, _, info = _feed_info(rows, bd)
    hr = _title_block(ws, f"Term Life \u2014 Daily Stats by Income Segment, "
                          f"{start.strftime('%B %Y')} (current month)", info)
    _table_header(ws, hr, ["Date", "Segment"])
    groups = {}
    for x in rows:
        k = (x["d"].isoformat(), x["seg"])
        groups.setdefault(k, _zero_acc())
        _acc(groups[k], x)
    segs = _seg_list(rows)
    r, d = hr + 1, start
    while d <= end:
        tot = _zero_acc()
        tot_spend = 0.0
        for seg in segs:
            a = groups.get((d.isoformat(), seg)) or _zero_acc()
            sp = _est_spend(seg, d.strftime("%Y-%m"), a)
            ws.cell(r, 1, d.isoformat())
            ws.cell(r, 2, seg_label(seg))
            _write_metrics(ws, r, 3, a, spend=sp)
            _merge(tot, a)
            tot_spend += sp or 0.0
            _style_row(ws, r, _seg_fill(seg))
            r += 1
        ws.cell(r, 1, d.isoformat())
        ws.cell(r, 2, "Total")
        _write_metrics(ws, r, 3, tot, spend=tot_spend or None)
        _style_row(ws, r, TOTAL_HEX, bold=True)
        r += 1
        r += 1                              # blank spacer between dates
        d += timedelta(days=1)


# ----------------------------------------------------------- workbook -------
BOLD = Font(bold=True)


def write_block(ws, top, title, agg):
    ws.cell(top, 1, title).font = BOLD
    hdr = ["Age bucket", "Leads", "SQO", "SQO %", "CRI", "CRI %",
           "Revenue", "Rev / lead"]
    for j, h in enumerate(hdr, 1):
        ws.cell(top + 1, j, h).font = BOLD
    r = top + 2
    for label, *_ in AGE_BUCKETS:
        a = agg[label]
        ws.cell(r, 1, label)
        ws.cell(r, 2, a["leads"])
        ws.cell(r, 3, a["sqo"])
        ws.cell(r, 4, a["sqo"] / a["leads"] if a["leads"] else None).number_format = "0.0%"
        ws.cell(r, 5, a["cri"])
        ws.cell(r, 6, a["cri"] / a["leads"] if a["leads"] else None).number_format = "0.0%"
        ws.cell(r, 7, round(a["rev"], 2)).number_format = "$#,##0"
        ws.cell(r, 8, round(a["rev"] / a["leads"], 2) if a["leads"] else None
                ).number_format = "$#,##0"
        r += 1
    return r + 1


def build_workbook(batch, rows, per_seg, agg_all, lag, hz, stats, period_rows, path: Path):
    wb = Workbook()

    ws = wb.active
    ws.title = "ReadMe"
    for i, line in enumerate([
        f"SQLife Term disposition — maturation curve (cross-sectional)",
        f"Batch: {batch}   Generated: {datetime.now():%Y-%m-%d %H:%M}",
        "",
        "METHOD: single-batch cross-section. 'CRI % at age 30-day' means: leads that",
        "are CURRENTLY 30 days old convert at that rate — a proxy for how any cohort",
        "matures. CAVEAT: assumes stable mix/quality over the 90-day window; if traffic",
        "changed materially, young buckets reflect the new mix, old buckets the old one.",
        "True longitudinal curves replace this once enough daily batches accumulate.",
        "",
        "M-channel (web leads) only in curve sheets; R/T volumes too small for rates.",
        "CRI_Lag sheet = when the money arrives (bound_date − lead day). A CRI's lag",
        "never migrates across batches; the maturation buckets do. Complementary views.",
        "Monthly/Weekly/Daily views: lead-creation (day) basis from the current batch.",
        "Leads = delivered rows; Ops = SUM(op_count); all 3 channels rolled up per segment.",
        "Economics (10/6/26): SQ pays Quote PER OP, segment-tiered (~$150/$87/$31 per op",
        "High/Med/Low, Aug'26); avg_cost = blended average price, never segment revenue.",
        "Spend (est) applies the SEG_OP_RATES rate card (seeded from Looker by-advert,",
        "Aug-Oct'26) to M-channel Ops. It is an ESTIMATE (+/-1-2% vs Looker): Referral",
        "payments (other basis) and Transfers ($0 in Looker) are excluded; Organic has",
        "no rate. Update the rate card monthly or replace with the AM's official card.",
        "Policy KPI (est) = CRI Prem / Spend (est) = SQ's return multiple on buying our",
        "leads: break-even ~2.96, target 4. SQ Rev = 33% x CRI Prem. CRI Prem is premium",
        "volume, not anyone's revenue. Quote-side margin (payment - Google spend) is in",
        "termlife_policy_kpi.py.",
        "CRI / revenue = cris / capped_cri_dollars flags, never the status text.",
        "Status map is PROVISIONAL (rule:v1) pending partner acronym decode.",
    ], 1):
        ws.cell(i, 1, line)

    # ---- confirmed reporting views (front of workbook) ----
    bd = date.fromisoformat(batch)
    _monthly_sheet(wb, period_rows, bd)
    _weekly_sheet(wb, period_rows, bd)
    _daily_sheet(wb, period_rows, bd)

    ws = wb.create_sheet("Maturation_ALL")
    write_block(ws, 1, "ALL M-channel (CA+CB+CC+QA web leads)", agg_all)

    ws = wb.create_sheet("Maturation_bySeg")
    top = 1
    for seg, agg in per_seg.items():
        top = write_block(ws, top, f"Segment {seg} (M-channel)", agg)

    # CRI% curve data + chart
    ws = wb.create_sheet("CRI_Curve")
    labels = [b[0] for b in AGE_BUCKETS]
    ws.cell(1, 1, "Age bucket").font = BOLD
    series_names = ["ALL"] + list(per_seg)
    for j, name in enumerate(series_names, 2):
        ws.cell(1, j, name).font = BOLD
    for i, label in enumerate(labels, 2):
        ws.cell(i, 1, label)
        for j, name in enumerate(series_names, 2):
            agg = agg_all if name == "ALL" else per_seg[name]
            a = agg[label]
            c = ws.cell(i, j, a["cri"] / a["leads"] if a["leads"] else None)
            c.number_format = "0.0%"
    chart = LineChart()
    chart.title = f"CRI rate by lead age — batch {batch}"
    chart.y_axis.numFmt = "0%"
    chart.y_axis.title = "CRI %"
    chart.x_axis.title = "Lead age (days)"
    data = Reference(ws, min_col=2, max_col=1 + len(series_names),
                     min_row=1, max_row=1 + len(labels))
    cats = Reference(ws, min_col=1, min_row=2, max_row=1 + len(labels))
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height, chart.width = 10, 22
    ws.add_chart(chart, "H2")

    # stage mix by age (ALL M)
    ws = wb.create_sheet("StageMix")
    stages = sorted({r["stage"] for r in rows})
    mix = {l: {s: 0 for s in stages} for l in labels}
    tot = {l: 0 for l in labels}
    for r in rows:
        if r["ch"] != "M":
            continue
        mix[r["bucket"]][r["stage"]] += 1
        tot[r["bucket"]] += 1
    ws.cell(1, 1, "Age bucket").font = BOLD
    for j, s in enumerate(stages, 2):
        ws.cell(1, j, s).font = BOLD
    for i, l in enumerate(labels, 2):
        ws.cell(i, 1, l)
        for j, s in enumerate(stages, 2):
            c = ws.cell(i, j, mix[l][s] / tot[l] if tot[l] else None)
            c.number_format = "0.0%"

    # CRI lag distribution (fixed per lead — never migrates across batches)
    ws = wb.create_sheet("CRI_Lag")
    ws.cell(1, 1, "CRI lag = bound_date − lead creation day "
                  "(M-channel CRIs with parseable bound_date)").font = BOLD
    for j, h in enumerate(["Lag bucket (days)", "CRIs", "% of known",
                           "Cumulative %"], 1):
        ws.cell(3, j, h).font = BOLD
    cum, r = 0, 4
    for label, *_ in AGE_BUCKETS:
        c = lag["counts"][label]
        cum += c
        ws.cell(r, 1, label)
        ws.cell(r, 2, c)
        if lag["n"]:
            ws.cell(r, 3, c / lag["n"]).number_format = "0.0%"
            ws.cell(r, 4, cum / lag["n"]).number_format = "0.0%"
        r += 1
    if hz and lag["p80"] is not None:
        hz_hi = int(hz.split("-")[1])
        verdict = ("CONSISTENT — curve horizon and lag agree"
                   if lag["p80"] <= hz_hi
                   else "TENSION — stationarity of the cross-section is suspect")
        check = (f"cross-check vs maturity horizon '{hz}': p80 lag "
                 f"{lag['p80']}d vs horizon upper bound {hz_hi}d -> {verdict}")
    else:
        check = "cross-check vs maturity horizon: n/a"
    r += 1
    for line in [f"known CRIs: {lag['n']}",
                 f"CRIs missing/unparseable bound_date: {lag['unknown']}",
                 f"negative lags (bound before lead day): {lag['neg']}",
                 f"lags >89d: {lag['over']}",
                 f"median lag: {lag['median']}d    p80 lag: {lag['p80']}d",
                 check]:
        ws.cell(r, 1, line)
        r += 1

    # status map dump for review
    ws = wb.create_sheet("StatusMap")
    con = sqlite3.connect(DB_PATH)
    counts = dict(con.execute(
        "SELECT status, COUNT(*) FROM v_life_disposition_current GROUP BY status"))
    for j, h in enumerate(["status_raw", "stage", "sub_reason", "map_source",
                           "rows_in_current_batch"], 1):
        ws.cell(1, j, h).font = BOLD
    for i, (raw, stage, sub, src) in enumerate(con.execute(
            """SELECT status_raw, stage, sub_reason, map_source
               FROM life_status_map ORDER BY stage, status_raw"""), 2):
        ws.cell(i, 1, raw)
        ws.cell(i, 2, stage)
        ws.cell(i, 3, sub)
        ws.cell(i, 4, src)
        ws.cell(i, 5, counts.get(raw, 0))
        if stage == "UNMAPPED":
            for j in range(1, 6):
                ws.cell(i, j).font = Font(bold=True, color="CC0000")
    con.close()

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# ------------------------------------------------------------------- main ---
def main():
    if not DB_PATH.exists():
        sys.exit(f"DB not found: {DB_PATH}")
    con = sqlite3.connect(DB_PATH)

    print("== #1 STATUS MAP ==")
    stats = refresh_status_map(con)
    print(f"  statuses: {stats['new']} new, {stats['updated']} re-ruled, "
          f"{stats['manual_kept']} manual overrides preserved")
    n_stage = con.execute(
        "SELECT stage, COUNT(*) FROM life_status_map GROUP BY stage ORDER BY 2 DESC")
    for stage, n in n_stage:
        print(f"      {stage:9s} {n:3d} statuses")
    if stats["unmapped"]:
        print("  *** UNMAPPED statuses (review / add rules or manual rows): ***")
        for s in stats["unmapped"]:
            print(f"      {s!r}")
    print("  view v_life_disposition_current_staged (re)created")

    print("\n== #2 MATURATION (cross-sectional, M-channel) ==")
    batch, rows = pull_current(con)
    segs = sorted({r["seg"] for r in rows if r["ch"] == "M"})
    agg_all = aggregate(rows)
    per_seg = {s: aggregate(rows, s) for s in segs}
    lag = lag_distribution(rows)
    period_rows, period_skipped = pull_period_rows(con)

    print(f"  batch {batch}")
    print(f"  {'bucket':7s} {'leads':>6s} {'SQO%':>7s} {'CRI%':>7s} {'rev/lead':>9s}")
    for label, *_ in AGE_BUCKETS:
        a = agg_all[label]
        if a["leads"]:
            print(f"  {label:7s} {a['leads']:6d} {a['sqo']/a['leads']:7.1%} "
                  f"{a['cri']/a['leads']:7.1%} ${a['rev']/a['leads']:8,.0f}")
        else:
            print(f"  {label:7s} {0:6d}")

    hz, rates, plateau = maturity_horizon(agg_all)
    if hz:
        print(f"\n  mature-plateau CRI% (56-89d): {plateau:.1%}")
        print(f"  ESTIMATED MATURITY HORIZON: '{hz}' days — youngest bucket reaching "
              f"80% of plateau.\n  Cohorts younger than this are NOT comparable to "
              f"mature ones; flag them in every report.")
    else:
        print("\n  maturity horizon not estimable (sparse old buckets)")

    print("\n== CRI LAG (bound_date − lead day, M-channel CRIs) ==")
    tot_cri = lag["n"] + lag["unknown"] + lag["neg"] + lag["over"]
    if lag["n"]:
        print(f"  parseable lags: {lag['n']} of {tot_cri} CRIs "
              f"({lag['unknown']} missing bound_date, {lag['neg']} negative, "
              f"{lag['over']} >89d)")
        print(f"  median lag: {lag['median']}d | p80 lag: {lag['p80']}d")
        if hz and lag["p80"] is not None:
            hz_hi = int(hz.split("-")[1])
            verdict = ("CONSISTENT" if lag["p80"] <= hz_hi
                       else "TENSION — cross-sectional stationarity suspect")
            print(f"  cross-check vs maturity horizon '{hz}': p80 lag "
                  f"{lag['p80']}d vs bound {hz_hi}d -> {verdict}")
    else:
        print("  no CRIs with parseable bound_date — investigate the field")

    print("\n== PERIOD VIEWS ==")
    print(f"  Monthly/Weekly/Daily built from {len(period_rows)} rows "
          f"(all channels rolled up)"
          + (f" | {period_skipped} rows skipped (unparseable day)"
             if period_skipped else ""))

    out = OUT_DIR / f"termlife_disposition_report_{batch.replace('-', '')}.xlsx"
    build_workbook(batch, rows, per_seg, agg_all, lag, hz, stats, period_rows, out)
    print(f"\n  workbook -> {out}")
    con.close()


if __name__ == "__main__":
    main()
