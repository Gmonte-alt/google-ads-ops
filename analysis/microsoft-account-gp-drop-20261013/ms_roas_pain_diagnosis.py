"""
ms_roas_pain_diagnosis.py — ROAS pain diagnosis, Microsoft Home (account_id=2)
═══════════════════════════════════════════════════════════════════════════════
Trailing 14 / 7 / 3-day windows, each paired with its immediately preceding
equal-length baseline, drilled three ways:

  Level 1  Campaign
  Level 2  Campaign × geo state
  Level 3  Campaign × ad group  (+ generic vs non-generic rollup)

Goal: locate WHERE ROAS is deteriorating and how much gross profit each
pocket is costing. Pockets are ranked by ΔGP vs baseline (and ΔROAS), never
by raw ROAS, and every cell carries its spend share of its parent.

SOURCES (per the ratified prompt + pinned architect map):
  L1  campaign_performance_daily        — SEARCH ONLY BY DESIGN (no network
      column; audience excluded from all search campaigns 8/14), device='ALL',
      conversions / conversion_value. all_conversion_value is read alongside
      and adjudicated: expected identical on search rows; >0.5% divergence
      hard-warns (take to the architect).
  L2  geo_performance_daily             — device='ALL' (rollup present — never
      sum splits), network != 'Audience' (other values are all search
      networks), geo_state full names + 'UNKNOWN', location_type summed
      (INTEREST + PHYSICAL is not double-counting), carries campaign_id.
      SYNDICATED-PARTNER rows are excluded from the L2 state cells so they
      share L1's scope: the 10/13 smoke proved campaign_performance_daily
      EXCLUDES partner traffic (geo−campaign gap == the 9/18–21 partner
      slice to the dollar), a scope the script re-adjudicates every run.
      The partner slice is its own exhibit on 05_Change_Overlay.
  L3  verified at runtime, in order:
        (a) ad_group_performance_daily rows for acct 2
        (b) keyword_performance_daily via keyword_instances (ki.account_id=2)
        (c) keyword_performance_network_daily via ki, network != 'Audience'
      Leads = primary_conv_count; revenue = all_conv_value (MS offline revenue
      posts there; primary_conv_value is pixel-only and understates) when the
      source is keyword-grain; agp uses its native columns with the same
      resolution logic. Ad groups join by ad_group_id, NEVER by name (names
      are non-unique across ~50 legacy state campaigns + tier campaigns).

CONVENTIONS ENFORCED:
  - weighted metrics only (sum-then-divide): ROAS=Σrev/Σspend, GP=Σrev−Σspend,
    CPA=Σspend/Σleads, RPL=Σrev/Σleads, CvR=Σleads/Σclicks
  - campaign_id coerced to int before any Python-side keying/joins
    (TEXT-vs-int silently matches nothing in parts of the DB)
  - column names introspected per-run (PRAGMA), never assumed
  - read-only connection; parsed_* columns NOT used (Google-only, NULL on
    acct 2) — generic/branded = the Non-Generic name token

WINDOWS: END = latest date with positive account-level revenue (platform
revenue lags ~1 day; expected END ≈ today−2, verified not assumed). Then
3d = END−2…END, 7d = END−6…END, 14d = END−13…END, each with the prior
equal-length baseline.

OVERLAY (confounds, not findings — stamped onto Findings rows):
  - 9/18 generic keyword expansion launch (new Quote/Compare/Cost-Core/Bundle
    groups) — new groups isolated via first-served date >= 2026-09-18
  - 9/18–9/21 syndicated-search-partner incident (opt-out 9/21 PM; clean from
    9/22) — sits inside the 14d BASELINE window, but L1/L2 exclude partner
    traffic by scope; the slice is quantified per day on the overlay sheet
    and L3 (kpd, no network column) may still carry it — see 04_Reconcile
  - 9/23 re-tier cutover (26 state movers) + same-day tCPA raises — inside the
    14d baseline; moved states detected from the data (dominant campaign
    changed baseline→current)
  - NO Microsoft changes since the 9/23 cutover (confirmed 10/13): the 9/29
    T2 budget raise was Google-side; the 10/12 step-2 raises never posted
  - campaign_change_events pulled tolerantly (LIKE %CPA%/%BUDGET%/%ROAS%);
    MS rows may be sparse — the KNOWN_EVENTS operator log fills the gap

OUTPUT: one workbook —
  00_Params | 01_Campaign | 02_Campaign_State | 03_Campaign_AdGroup |
  03b_Generic_Rollup | 04_Reconcile | 05_Change_Overlay | 06_Findings
plus a console summary. --smoke runs END detection, source verdicts,
revenue adjudication and reconciliation gates only (no workbook).

USAGE (paid-search conda env):
  python ms_roas_pain_diagnosis.py --smoke
  python ms_roas_pain_diagnosis.py
  python ms_roas_pain_diagnosis.py --end-date 2026-10-11
"""

import argparse
import re
import sqlite3
import sys
from collections import defaultdict
from datetime import date, timedelta
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ── CONFIG ──────────────────────────────────────────────────────────────────
ACCOUNT_ID = 2
DEFAULT_DB = Path(r"C:\Projects\google-ads-ops\db\paid_search.db")
DEFAULT_OUT_DIR = Path(r"C:\Projects\google-ads-ops\analysis\ms-roas-pain\output")

WINDOWS = [14, 7, 3]                 # trailing window lengths, days
EXPANSION_LAUNCH = date(2026, 9, 18)  # new-group isolation threshold
REV_EQ_TOL = 0.005                   # conversion_value vs all_conversion_value
L2_RECON_WARN = 0.001                # geo should tie exactly to campaign
PARTNER_FLAG_SHARE = 0.10            # baseline partner spend share flag floor
THIN_SPEND_SHARE = 0.01              # <1% of parent spend = THIN tag
MIN_RANK_SPEND = 100.0               # cur+base spend floor for Findings rank

# Known operator actions (from the surge/expansion/re-tier logs) — overlay
# context where campaign_change_events is sparse on the Microsoft side.
KNOWN_EVENTS = [
    ("2026-09-18", "Generic keyword expansion LAUNCH (Quote/Compare/"
                   "Cost-Core/Bundle groups across live campaigns)"),
    ("2026-09-21", "Syndicated search partners opted OUT (PM; partner "
                   "traffic clean from 9/22 — 9/18–9/21 reads are "
                   "partner-dominated)"),
    ("2026-09-23", "Re-tier cutover posted (26 state movers) + same-day "
                   "tCPA raises: Tier1 $54.34, Tier3 $37.44, CA $18.52, "
                   "FL $13.47 (Tier2/Tier5 cuts deferred)"),
    ("2026-10-12", "Step-2 tCPA raises were SCHEDULED here but NOT POSTED "
                   "— confirmed 10/13: no Microsoft Home changes of any "
                   "kind since the 9/23 cutover (the 9/29 Tier 2 budget "
                   "raise was Google-side only)"),
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10, name="Arial")
SUBHDR_FILL = PatternFill("solid", fgColor="2E74B5")
TITLE_FONT = Font(bold=True, size=13, color="1F4E79", name="Arial")
NOTE_FONT = Font(italic=True, size=9, color="595959", name="Arial")
BODY_FONT = Font(size=10, name="Arial")
BOLD_FONT = Font(size=10, bold=True, name="Arial")
NEG_FONT = Font(size=10, color="C00000", name="Arial")
TOTAL_FILL = PatternFill("solid", fgColor="DDEBF7")
BASE_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)

MONEY = '"$"#,##0'
MONEY2 = '"$"#,##0.00'
NUM1 = "#,##0.0"
ROASF = "0.00"
PCT1 = "0.0%"

CONSOLE = []


def log(msg=""):
    print(msg)
    CONSOLE.append(str(msg))


def die(msg):
    log("")
    log(f"*** HARD FAIL: {msg}")
    sys.exit(1)


# ── introspection helpers ───────────────────────────────────────────────────

def table_columns(con, table):
    rows = con.execute(f"PRAGMA table_info({table})").fetchall()
    return [r[1] for r in rows]


def resolve_col(cols, candidates, table, what, required=True):
    for c in candidates:
        if c in cols:
            return c
    if required:
        die(f"{table}: no {what} column among {candidates}; columns={cols}")
    return None


def to_int_id(v):
    """Coerce campaign/ad-group ids to int — TEXT ids silently 0-match."""
    try:
        return int(str(v).strip())
    except (TypeError, ValueError):
        return None


BRANDED_RX = re.compile(r"non[\s_\-]?generic", re.IGNORECASE)


def is_branded(ad_group_name: str) -> bool:
    return bool(BRANDED_RX.search(ad_group_name or ""))


TIER_RX = re.compile(r"TIER\s*(\d+)", re.IGNORECASE)
STATE_TOKENS = {
    "CALIFORNIA": "CA", "FLORIDA": "FL", "TEXAS": "TX", "INDIANA": "IN",
    "KENTUCKY": "KY", "NEW JERSEY": "NJ", "MASSACHUSETTS": "MA",
}


def short_code(name: str) -> str:
    """Tier-aware campaign short code (tier branch FIRST — Tier 5 is the
    renamed Indiana shell, so tier names can embed a state token)."""
    up = (name or "").upper()
    m = TIER_RX.search(up)
    if m:
        return f"T{m.group(1)}"
    for tok, ab in STATE_TOKENS.items():
        if tok in up:
            return ab
    cleaned = "".join(c for c in up if c.isalnum())[:8]
    return cleaned or "CAMP"


# ── window machinery ────────────────────────────────────────────────────────

def build_windows(end: date):
    """{label: dict(cur=(a,b), base=(a,b))} for 14/7/3-day windows with the
    immediately preceding equal-length baseline."""
    out = {}
    for n in WINDOWS:
        cur_a = end - timedelta(days=n - 1)
        base_b = cur_a - timedelta(days=1)
        base_a = base_b - timedelta(days=n - 1)
        out[f"{n}d"] = {"cur": (cur_a, end), "base": (base_a, base_b)}
    return out


def detect_end(con, cpd_cols, cli_end):
    """Latest date with positive account-level revenue (search campaign
    table). Verified, never assumed — platform revenue lags ~1 day."""
    rows = con.execute(
        """SELECT t.date, SUM(t.conversion_value) rev, SUM(t.cost_micros)/1e6 sp
           FROM campaign_performance_daily t
           JOIN campaigns c ON c.campaign_id = t.campaign_id
           WHERE c.account_id = ? AND t.device = 'ALL'
           GROUP BY t.date ORDER BY t.date DESC LIMIT 8""",
        (ACCOUNT_ID,)).fetchall()
    if not rows:
        die("no campaign_performance_daily rows for account 2 (device='ALL')")
    log("[END] last 8 account days (date, revenue, spend):")
    for d, rv, sp in rows:
        log(f"       {d}  rev ${rv or 0:>10,.2f}   spend ${sp or 0:>9,.2f}")
    if cli_end:
        end = date.fromisoformat(cli_end)
        log(f"[END] CLI override: {end}")
        return end
    for d, rv, _ in rows:
        if (rv or 0) > 0:
            end = date.fromisoformat(str(d)[:10])
            log(f"[END] detected last complete-revenue day = {end}")
            tail = [r for r in rows if str(r[0])[:10] > end.isoformat()]
            if tail:
                log(f"      ({len(tail)} newer zero-revenue day(s) excluded "
                    f"— revenue still processing)")
            return end
    die("no positive-revenue day found in the last 8 days")


# ── metric helpers ──────────────────────────────────────────────────────────

def derive(spend, leads, rev):
    roas = rev / spend if spend > 0 else None
    gp = rev - spend
    return roas, gp


def win_mask(df, a, b):
    return (df["date"] >= a) & (df["date"] <= b)


def agg_window(df, keys, a, b, cols=("spend", "leads", "rev", "clicks")):
    cols = [c for c in cols if c in df.columns]
    sub = df[win_mask(df, a, b)]
    if sub.empty:
        return pd.DataFrame(columns=list(keys) + list(cols))
    # dropna=False: a null key must surface as a visible (none) group, never
    # silently delete its rows — pandas' default dropna=True is exactly how
    # acct-2 kpd rows with NULL campaign_id vanished into $0 sums (10/13)
    return sub.groupby(list(keys), as_index=False,
                       dropna=False)[list(cols)].sum()


def norm_key(v):
    """iterrows() upcasts pure-numeric key rows to float64, turning id 201
    into 201.0 — which str()s to '201.0' and silently misses TEXT ids in SQL
    (the str-vs-int 0-match bug in a new costume). Normalize every key
    component back to a plain Python int where it is integral."""
    if hasattr(v, "item"):          # numpy scalar -> python scalar
        v = v.item()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def block_lookup(df, keys):
    """-> {key_tuple: (spend, leads, rev)}"""
    out = {}
    for _, r in df.iterrows():
        k = tuple(norm_key(r[k]) for k in keys)
        out[k] = (float(r.get("spend", 0.0)), float(r.get("leads", 0.0)),
                  float(r.get("rev", 0.0)))
    return out


# ═══════════════════════════════════════════════════════════════════════════
# DATA PULLS
# ═══════════════════════════════════════════════════════════════════════════

def pull_l1(con, span_a, span_b):
    """Campaign×date daily from campaign_performance_daily (search-only by
    design). Both revenue candidates pulled for adjudication."""
    sql = """
    SELECT t.campaign_id, t.date,
           SUM(t.cost_micros)/1.0e6  AS spend,
           SUM(t.clicks)             AS clicks,
           SUM(t.conversions)        AS leads,
           SUM(t.conversion_value)   AS rev,
           SUM(t.all_conversion_value) AS rev_all
    FROM campaign_performance_daily t
    JOIN campaigns c ON c.campaign_id = t.campaign_id
    WHERE c.account_id = ? AND t.device = 'ALL'
      AND t.date BETWEEN ? AND ?
    GROUP BY t.campaign_id, t.date
    """
    df = pd.read_sql(sql, con, params=(ACCOUNT_ID, span_a.isoformat(),
                                       span_b.isoformat()))
    df["campaign_id"] = df["campaign_id"].map(to_int_id)
    df["date"] = pd.to_datetime(df["date"]).dt.date
    for c in ("spend", "clicks", "leads", "rev", "rev_all"):
        df[c] = df[c].fillna(0.0)
    return df


def adjudicate_l1_revenue(l1):
    tot_cv = l1["rev"].sum()
    tot_acv = l1["rev_all"].sum()
    gap = abs(tot_acv - tot_cv) / tot_cv if tot_cv else 0.0
    status = ("identical as expected; conversion_value used per prompt"
              if gap <= REV_EQ_TOL else "DIVERGENT — exceeds tolerance")
    verdict = (f"L1 revenue adjudication: conversion_value ${tot_cv:,.2f} vs "
               f"all_conversion_value ${tot_acv:,.2f} — gap {gap:.3%} "
               f"({status})")
    log(f"[ADJ] {verdict}")
    if gap > REV_EQ_TOL:
        log("      *** WARN: campaign-table revenue columns diverge on this "
            "span. conversion_value kept per the ratified prompt, but take "
            "this to the DB architect before trusting absolute GP. ***")
    return verdict, gap


def pull_geo(con, geo_cols, span_a, span_b):
    """Campaign×state×network×date from geo_performance_daily (search rows
    kept per-network so the syndicated-partner slice is separable)."""
    d_col = resolve_col(geo_cols, ["date", "report_date"],
                        "geo_performance_daily", "date")
    devs = [r[0] for r in con.execute(
        "SELECT DISTINCT device FROM geo_performance_daily "
        "WHERE account_id=?", (ACCOUNT_ID,))]
    all_lit = next((d for d in devs if str(d).strip().upper() == "ALL"), None)
    if all_lit is None:
        die(f"geo_performance_daily: no ALL device rollup (devices={devs}) — "
            f"pinned ruling says it exists")
    nets = [r[0] for r in con.execute(
        "SELECT DISTINCT network FROM geo_performance_daily "
        "WHERE account_id=?", (ACCOUNT_ID,))]
    log(f"[GEO] device literals: {devs}  (using '{all_lit}')")
    log(f"[GEO] network values:  {nets}")
    if "Audience" not in nets:
        log("      WARN: literal 'Audience' absent — verify the search-only "
            "filter still excludes the right rows")
    partner_nets = [n for n in nets if "syndicated" in str(n).lower()]
    log(f"[GEO] syndicated-partner network literal(s): "
        f"{partner_nets or 'NONE FOUND'}")

    sql = f"""
    SELECT t.campaign_id, t.{d_col} AS date, t.geo_state, t.network,
           SUM(t.cost_micros)/1.0e6 AS spend,
           SUM(t.clicks)            AS clicks,
           SUM(t.conversions)       AS leads,
           SUM(t.conversion_value)  AS rev
    FROM geo_performance_daily t
    WHERE t.account_id = ? AND t.device = ? AND t.network != 'Audience'
      AND t.{d_col} BETWEEN ? AND ?
    GROUP BY t.campaign_id, t.{d_col}, t.geo_state, t.network
    """
    df = pd.read_sql(sql, con, params=(ACCOUNT_ID, all_lit,
                                       span_a.isoformat(), span_b.isoformat()))
    df["campaign_id"] = df["campaign_id"].map(to_int_id)
    df["date"] = pd.to_datetime(df["date"]).dt.date
    df["geo_state"] = df["geo_state"].fillna("UNKNOWN")
    for c in ("spend", "clicks", "leads", "rev"):
        df[c] = df[c].fillna(0.0)
    df["is_partner"] = df["network"].isin(partner_nets)
    return df, partner_nets


def resolve_l3_source(con, span_a, span_b):
    """Verify, in order: agp acct-2 rows → kpd via ki → kpnd via ki.
    Returns (source_name, note)."""
    tables = [r[0] for r in con.execute(
        "SELECT name FROM sqlite_master WHERE type IN ('table','view')")]

    # (a) ad_group_performance_daily
    if "ad_group_performance_daily" in tables:
        agp_cols = table_columns(con, "ad_group_performance_daily")
        d_col = resolve_col(agp_cols, ["date", "report_date"],
                            "ad_group_performance_daily", "date")
        if "campaign_id" in agp_cols:
            n = con.execute(
                f"""SELECT COUNT(*) FROM ad_group_performance_daily t
                    JOIN campaigns c ON c.campaign_id = t.campaign_id
                    WHERE c.account_id = ? AND t.{d_col} BETWEEN ? AND ?""",
                (ACCOUNT_ID, span_a.isoformat(), span_b.isoformat())
            ).fetchone()[0]
            log(f"[L3] ad_group_performance_daily acct-2 rows in span: {n:,}")
            if n > 0:
                return "agp", (f"ad_group_performance_daily ({n:,} acct-2 "
                               f"rows in span; date col = {d_col})")
        else:
            log("[L3] agp lacks campaign_id — skipping candidate (a)")

    # (b) keyword_performance_daily via keyword_instances
    if "keyword_performance_daily" in tables:
        kpd_cols = table_columns(con, "keyword_performance_daily")
        d_col = resolve_col(kpd_cols, ["report_date", "date"],
                            "keyword_performance_daily", "date")
        n = con.execute(
            f"""SELECT COUNT(*) FROM keyword_performance_daily p
                JOIN keyword_instances k ON k.id = p.instance_id
                WHERE k.account_id = ? AND p.{d_col} BETWEEN ? AND ?""",
            (ACCOUNT_ID, span_a.isoformat(), span_b.isoformat())
        ).fetchone()[0]
        log(f"[L3] keyword_performance_daily acct-2 rows in span: {n:,}")
        if n > 0:
            return "kpd", (f"keyword_performance_daily via keyword_instances "
                           f"({n:,} acct-2 rows in span)")

    # (c) keyword_performance_network_daily via keyword_instances
    if "keyword_performance_network_daily" in tables:
        kn_cols = table_columns(con, "keyword_performance_network_daily")
        d_col = resolve_col(kn_cols, ["report_date", "date"],
                            "keyword_performance_network_daily", "date")
        n = con.execute(
            f"""SELECT COUNT(*) FROM keyword_performance_network_daily p
                JOIN keyword_instances k ON k.id = p.instance_id
                WHERE k.account_id = ? AND p.{d_col} BETWEEN ? AND ?""",
            (ACCOUNT_ID, span_a.isoformat(), span_b.isoformat())
        ).fetchone()[0]
        log(f"[L3] keyword_performance_network_daily acct-2 rows in span: {n:,}")
        if n > 0:
            return "kpnd", (f"keyword_performance_network_daily via "
                            f"keyword_instances ({n:,} acct-2 rows in span; "
                            f"value columns known not to reconcile to the "
                            f"campaign pair historically — recon gap "
                            f"REPORTED, not assumed)")

    die("no ad-group-grain source holds acct-2 rows in the span — this is "
        "the flagged open architect question; stop here and escalate")


def pull_l3(con, source, span_a, span_b, hist_a):
    """Ad-group×date daily per the verified source, plus first-served dates
    (from hist_a, well before the expansion launch) for new-group isolation.
    Leads/revenue semantics per source; everything joined by ad_group_id."""
    if source == "agp":
        cols = table_columns(con, "ad_group_performance_daily")
        d_col = resolve_col(cols, ["date", "report_date"],
                            "ad_group_performance_daily", "date")
        dev_clause = "AND t.device = 'ALL'" if "device" in cols else ""
        lead_col = resolve_col(cols, ["primary_conv_count", "conversions"],
                               "agp", "leads")
        rev_col = resolve_col(cols, ["all_conv_value", "all_conversion_value",
                                     "conversion_value", "conversions_value"],
                              "agp", "revenue")
        detail = (f"agp: leads={lead_col}, revenue={rev_col}, device "
                  f"filter={'ALL' if dev_clause else 'none present'}")
        log(f"[L3] {detail}")
        sql = f"""
        SELECT t.campaign_id, t.ad_group_id, t.{d_col} AS date,
               SUM(t.cost_micros)/1.0e6 AS spend,
               SUM(t.clicks)            AS clicks,
               SUM(t.{lead_col})        AS leads,
               SUM(t.{rev_col})         AS rev
        FROM ad_group_performance_daily t
        JOIN campaigns c ON c.campaign_id = t.campaign_id
        WHERE c.account_id = ? {dev_clause} AND t.{d_col} BETWEEN ? AND ?
        GROUP BY t.campaign_id, t.ad_group_id, t.{d_col}
        """
        df = pd.read_sql(sql, con, params=(ACCOUNT_ID, span_a.isoformat(),
                                           span_b.isoformat()))
        # first-served scans history, where the device vocab may drift —
        # impressions>0 on ANY device row is the right test, no filter
        first_sql = f"""
        SELECT t.ad_group_id, MIN(t.{d_col}) AS first_served
        FROM ad_group_performance_daily t
        JOIN campaigns c ON c.campaign_id = t.campaign_id
        WHERE c.account_id = ? AND t.impressions > 0
          AND t.{d_col} BETWEEN ? AND ?
        GROUP BY t.ad_group_id
        """
        fs = pd.read_sql(first_sql, con, params=(ACCOUNT_ID,
                                                 hist_a.isoformat(),
                                                 span_b.isoformat()))
    else:
        table = ("keyword_performance_daily" if source == "kpd"
                 else "keyword_performance_network_daily")
        cols = table_columns(con, table)
        d_col = resolve_col(cols, ["report_date", "date"], table, "date")
        lead_col = resolve_col(cols, ["primary_conv_count"], table, "leads")
        rev_col = resolve_col(cols, ["all_conv_value"], table, "revenue")
        net_clause = ("AND p.network != 'Audience'"
                      if source == "kpnd" and "network" in cols else "")

        def scoped_sums(where_extra, col_exprs):
            sql = (f"SELECT {', '.join(col_exprs)} FROM {table} p "
                   f"JOIN keyword_instances k ON k.id = p.instance_id "
                   f"WHERE k.account_id=? {net_clause} {where_extra} "
                   f"AND p.{d_col} BETWEEN ? AND ?")
            return con.execute(sql, (ACCOUNT_ID, span_a.isoformat(),
                                     span_b.isoformat())).fetchone()

        # DEVICE STRATEGY — adjudicated from the DATA, not the vocab. The
        # 10/13 smoke showed acct-2 kpd carries device='ALL' rows in-span
        # yet the 'ALL'-filtered pull summed to $0: rollup rows can exist
        # with NULL/zero metrics while the metrics live on the split rows
        # (or vice versa). Pick the scope by where the cost actually is,
        # and print the per-device table as evidence.
        dev_clause = ""
        if "device" in cols:
            probe = con.execute(
                f"SELECT p.device, COUNT(*), SUM(p.cost_micros), "
                f"SUM(p.{lead_col}), SUM(p.{rev_col}) "
                f"FROM {table} p "
                f"JOIN keyword_instances k ON k.id = p.instance_id "
                f"WHERE k.account_id=? {net_clause} "
                f"AND p.{d_col} BETWEEN ? AND ? "
                f"GROUP BY p.device ORDER BY 3 DESC",
                (ACCOUNT_ID, span_a.isoformat(), span_b.isoformat())
            ).fetchall()
            log(f"[L3] {table} in-span by device "
                f"(rows / cost / {lead_col} / {rev_col}):")
            for dv, n, co, ld, rv in probe:
                log(f"       {str(dv):14} {n:>8,}   ${(co or 0)/1e6:>11,.2f}"
                    f"   {ld or 0:>10,.1f}   ${rv or 0:>11,.2f}")
            all_rows = [r for r in probe
                        if str(r[0]).strip().upper() == "ALL"]
            all_cost = sum((r[2] or 0) for r in all_rows) / 1e6
            split_cost = sum((r[2] or 0) for r in probe
                             if str(r[0]).strip().upper() != "ALL") / 1e6
            if all_rows and all_cost > 0:
                esc = str(all_rows[0][0]).replace("'", "''")
                dev_clause = f"AND p.device = '{esc}'"
                dev_note = f"device = '{all_rows[0][0]}' rollup"
                if split_cost > 0:
                    g = abs(split_cost - all_cost) / all_cost
                    dev_note += f" (ALL vs Σsplits cost gap {g:.2%})"
                log(f"[L3] device strategy: {dev_note}")
            elif split_cost > 0:
                dev_note = "SUM of split devices"
                if all_rows:
                    esc = ",".join(
                        "'" + str(r[0]).replace("'", "''") + "'"
                        for r in all_rows)
                    dev_clause = f"AND p.device NOT IN ({esc})"
                    dev_note += (" (rollup rows exist but carry NO cost — "
                                 "excluded; relay to architect)")
                log(f"[L3] device strategy: {dev_note}")
            else:
                die(f"{table}: acct-2 rows in span carry ZERO cost on every "
                    f"device — column semantics need the architect before "
                    f"L3 can be built")
        else:
            dev_note = "no device column"

        # COLUMN ADEQUACY — pinned columns first; if a pinned column is
        # empty on this account's rows while the account plainly has the
        # metric, fall back with an explicit escalate note rather than
        # silently shipping zeros. GP ranking never uses the leads column,
        # so a leads fallback is display-only.
        rev_fb = lead_fb = False
        _, chk_ld, chk_rv = scoped_sums(
            dev_clause, ["COUNT(*)", f"SUM(p.{lead_col})",
                         f"SUM(p.{rev_col})"])
        if not (chk_rv or 0) > 0:
            for cand in ("primary_conv_value", "conversion_value",
                         "conversions_value"):
                if cand in cols and cand != rev_col:
                    alt = scoped_sums(dev_clause, [f"SUM(p.{cand})"])[0]
                    if (alt or 0) > 0:
                        log(f"[L3] *** {rev_col} is EMPTY on acct-2 rows — "
                            f"revenue falls back to {cand} "
                            f"(${alt:,.2f} in span). RELAY TO ARCHITECT: "
                            f"the pinned MS keyword-grain revenue column "
                            f"doesn't hold for this table/account. ***")
                        rev_col = cand
                        rev_fb = True
                        break
            else:
                die(f"{table}: no revenue column holds data for acct 2 "
                    f"in span — stop and escalate")
        if not (chk_ld or 0) > 0:
            if "all_conv_count" in cols:
                alt = scoped_sums(dev_clause,
                                  ["SUM(p.all_conv_count)"])[0]
                if (alt or 0) > 0:
                    log(f"[L3] *** {lead_col} is EMPTY on acct-2 rows — "
                        f"leads column falls back to all_conv_count "
                        f"({alt:,.1f} in span; includes Zip Submit, NOT a "
                        f"clean lead count — display-only, GP ranking "
                        f"unaffected). RELAY TO ARCHITECT. ***")
                    lead_col = "all_conv_count"
                    lead_fb = True
        col_note = (f"leads={lead_col}"
                    + (" (FALLBACK — display-only)" if lead_fb else "")
                    + f", revenue={rev_col}"
                    + (" (FALLBACK — escalate)" if rev_fb else ""))

        # CAMPAIGN LINKAGE — adjudicated from the data. On acct-2 rows
        # kpd.campaign_id can be NULL (the 7/30 campaign_id repair was
        # Google-side), and a NULL group key silently deletes its rows in
        # pandas: that is how a $47K span summed to $0 (10/13). COALESCE
        # through keyword_instances and print the null rates as evidence.
        ki_cols = table_columns(con, "keyword_instances")
        has_p_cid = "campaign_id" in cols
        has_k_cid = "campaign_id" in ki_cols
        if has_p_cid or has_k_cid:
            parts = (["p.campaign_id"] if has_p_cid else []) \
                + (["k.campaign_id"] if has_k_cid else [])
            camp_expr = (f"COALESCE({', '.join(parts)})"
                         if len(parts) > 1 else parts[0])
            probe_cols = []
            if has_p_cid:
                probe_cols.append("SUM(p.campaign_id IS NULL)")
            if has_k_cid:
                probe_cols.append("SUM(k.campaign_id IS NULL)")
            probe_cols.append(f"SUM(({camp_expr}) IS NULL)")
            probe_cols.append("SUM(k.ad_group_id IS NULL)")
            probe_cols.append("COUNT(*)")
            vals = scoped_sums(dev_clause, probe_cols)
            n_tot = vals[-1] or 0
            labels = (["p.campaign_id"] if has_p_cid else []) \
                + (["ki.campaign_id"] if has_k_cid else []) \
                + ["COALESCE", "ki.ad_group_id"]
            log(f"[L3] linkage null rates over {n_tot:,} pulled rows: "
                + ", ".join(f"{lab} {v or 0:,} null"
                            for lab, v in zip(labels, vals[:-1])))
            if n_tot and (vals[-3] or 0) == n_tot:
                die(f"{table}: campaign linkage is NULL on every acct-2 "
                    f"row even after COALESCE — escalate to the architect")
            if n_tot and (vals[-2] or 0) == n_tot:
                die(f"{table}: keyword_instances.ad_group_id is NULL on "
                    f"every acct-2 row — escalate to the architect")
        else:
            die(f"{table}/keyword_instances carry no campaign_id at all — "
                f"ad_groups-join fallback needs the architect")
        detail = (f"{table}: {col_note}; device: {dev_note}; campaign via "
                  f"{camp_expr}; "
                  f"net filter={'search-only' if net_clause else 'n/a'}")
        log(f"[L3] {table}: date={d_col}, {col_note}, "
            f"campaign via {camp_expr}, device: {dev_note}, "
            f"net filter={'search-only' if net_clause else 'n/a'}")
        sql = f"""
        SELECT {camp_expr} AS campaign_id, k.ad_group_id,
               p.{d_col} AS date,
               SUM(p.cost_micros)/1.0e6 AS spend,
               SUM(p.clicks)            AS clicks,
               SUM(p.{lead_col})        AS leads,
               SUM(p.{rev_col})         AS rev
        FROM {table} p
        JOIN keyword_instances k ON k.id = p.instance_id
        WHERE k.account_id = ? {dev_clause} {net_clause}
          AND p.{d_col} BETWEEN ? AND ?
        GROUP BY 1, 2, 3
        """
        df = pd.read_sql(sql, con, params=(ACCOUNT_ID, span_a.isoformat(),
                                           span_b.isoformat()))
        first_sql = f"""
        SELECT k.ad_group_id, MIN(p.{d_col}) AS first_served
        FROM {table} p
        JOIN keyword_instances k ON k.id = p.instance_id
        WHERE k.account_id = ? {net_clause}
          AND p.impressions > 0 AND p.{d_col} BETWEEN ? AND ?
        GROUP BY k.ad_group_id
        """
        fs = pd.read_sql(first_sql, con, params=(ACCOUNT_ID,
                                                 hist_a.isoformat(),
                                                 span_b.isoformat()))

    df["campaign_id"] = df["campaign_id"].map(to_int_id)
    df["ad_group_id"] = df["ad_group_id"].map(to_int_id)
    df["date"] = pd.to_datetime(df["date"]).dt.date
    for c in ("spend", "clicks", "leads", "rev"):
        df[c] = df[c].fillna(0.0)
    n_null_c = int(df["campaign_id"].isna().sum())
    n_null_a = int(df["ad_group_id"].isna().sum())
    if n_null_c or n_null_a:
        log(f"[L3] *** {n_null_c}/{len(df)} pulled rows lack campaign "
            f"linkage and {n_null_a}/{len(df)} lack ad-group linkage even "
            f"after COALESCE — kept under a visible (none) key so sums "
            f"stay honest; escalate if material ***")
    first_served = {to_int_id(r["ad_group_id"]):
                    date.fromisoformat(str(r["first_served"])[:10])
                    for _, r in fs.iterrows()
                    if to_int_id(r["ad_group_id"]) is not None}
    return df, first_served, detail


def pull_ad_group_names(con):
    """ad_group_id -> name. Key column resolved at runtime (ad_group_id,
    not id — v1 of another script guessed wrong and crashed)."""
    cols = table_columns(con, "ad_groups")
    key = resolve_col(cols, ["ad_group_id", "id"], "ad_groups", "key")
    name_col = resolve_col(cols, ["ad_group_name", "name"], "ad_groups",
                           "name")
    scope = "WHERE account_id = ?" if "account_id" in cols else ""
    params = (ACCOUNT_ID,) if scope else ()
    rows = con.execute(
        f"SELECT {key}, {name_col} FROM ad_groups {scope}", params).fetchall()
    return {to_int_id(k): (n or "") for k, n in rows}


def pull_change_events(con, span_a, span_b, roster_ids):
    """Tolerant pull of tCPA/budget/tROAS events for acct-2 campaigns.
    MS rows may be sparse — KNOWN_EVENTS fills the operator-log gap."""
    tables = [r[0] for r in con.execute(
        "SELECT name FROM sqlite_master WHERE type IN ('table','view')")]
    if "campaign_change_events" not in tables:
        log("[EVT] campaign_change_events not found — overlay uses the "
            "known-events log only")
        return []
    cols = table_columns(con, "campaign_change_events")
    d_col = resolve_col(cols, ["change_date", "event_date", "date"],
                        "campaign_change_events", "date", required=False)
    t_col = resolve_col(cols, ["change_type", "attribute", "type"],
                        "campaign_change_events", "type", required=False)
    if not (d_col and t_col):
        log(f"[EVT] unexpected schema ({cols}) — overlay uses known events")
        return []
    o_col = resolve_col(cols, ["old_value", "old"], "cce", "old",
                        required=False)
    n_col = resolve_col(cols, ["new_value", "new"], "cce", "new",
                        required=False)
    ph = ",".join("?" * len(roster_ids))
    sql = (f"SELECT campaign_id, {d_col}, {t_col}"
           + (f", {o_col}" if o_col else ", NULL")
           + (f", {n_col}" if n_col else ", NULL")
           + f" FROM campaign_change_events "
           f"WHERE campaign_id IN ({ph}) "
           f"AND (UPPER({t_col}) LIKE '%CPA%' OR UPPER({t_col}) LIKE "
           f"'%BUDGET%' OR UPPER({t_col}) LIKE '%ROAS%') "
           f"AND {d_col} BETWEEN ? AND ? ORDER BY {d_col}")
    rows = con.execute(sql, (*[str(i) for i in roster_ids],
                             span_a.isoformat(), span_b.isoformat())
                       ).fetchall()
    if not rows:
        # ids may be stored as ints; retry
        rows = con.execute(sql, (*roster_ids, span_a.isoformat(),
                                 span_b.isoformat())).fetchall()
    log(f"[EVT] campaign_change_events matched rows in span: {len(rows)}")
    return [(to_int_id(r[0]), str(r[1])[:10], str(r[2]),
             r[3], r[4]) for r in rows]


# ═══════════════════════════════════════════════════════════════════════════
# CELL / BLOCK ASSEMBLY
# ═══════════════════════════════════════════════════════════════════════════

def build_blocks(df, keys, windows):
    """{wlabel: {'cur': {key: (sp, ld, rv)}, 'base': {...}}}"""
    out = {}
    for w, rng in windows.items():
        out[w] = {}
        for side in ("cur", "base"):
            a, b = rng[side]
            out[w][side] = block_lookup(agg_window(df, keys, a, b), keys)
    return out


def all_keys(blocks):
    ks = set()
    for w in blocks.values():
        for side in w.values():
            ks.update(side.keys())
    return ks


def parent_totals(blocks, key_slice=None):
    """Per window/side totals; key_slice maps cell key -> parent key
    (None = single account-level parent)."""
    out = {}
    for w, sides in blocks.items():
        out[w] = {}
        for side, cells in sides.items():
            tot = defaultdict(lambda: [0.0, 0.0, 0.0])
            for k, (sp, ld, rv) in cells.items():
                pk = key_slice(k) if key_slice else "ACCOUNT"
                tot[pk][0] += sp
                tot[pk][1] += ld
                tot[pk][2] += rv
            out[w][side] = dict(tot)
    return out


def cell_metrics(cells, k):
    sp, ld, rv = cells.get(k, (0.0, 0.0, 0.0))
    roas, gp = derive(sp, ld, rv)
    return sp, ld, rv, roas, gp


# ═══════════════════════════════════════════════════════════════════════════
# WORKBOOK
# ═══════════════════════════════════════════════════════════════════════════

WIN_COLS = ["spend", "spd_shr", "leads", "rev", "ROAS", "GP",
            "b_spend", "b_leads", "b_rev", "b_ROAS", "b_GP",
            "dROAS", "dGP"]
WIN_FMTS = [MONEY, PCT1, NUM1, MONEY, ROASF, MONEY,
            MONEY, NUM1, MONEY, ROASF, MONEY, ROASF, MONEY]


def write_level_sheet(ws, title, id_headers, rows, windows, note):
    """rows: list of dicts with 'ids': [...], per-window metric tuples under
    row[w] = (sp, shr, ld, rv, roas, gp, bsp, bld, brv, broas, bgp), and
    optional 'flags'. Sorted by caller."""
    ws.cell(row=1, column=1, value=title).font = TITLE_FONT
    ws.cell(row=2, column=1, value=note).font = NOTE_FONT

    n_id = len(id_headers)
    # window group header
    for wi, w in enumerate(windows):
        c0 = n_id + 1 + wi * len(WIN_COLS)
        a, b = windows[w]["cur"]
        ba, bb = windows[w]["base"]
        cell = ws.cell(row=4, column=c0,
                       value=f"{w}: {a.strftime('%m/%d')}–{b.strftime('%m/%d')}"
                             f"  (base {ba.strftime('%m/%d')}–"
                             f"{bb.strftime('%m/%d')})")
        cell.fill = SUBHDR_FILL
        cell.font = HEADER_FONT
        ws.merge_cells(start_row=4, start_column=c0, end_row=4,
                       end_column=c0 + len(WIN_COLS) - 1)
    hdr = id_headers + [h for _ in windows for h in WIN_COLS] + ["flags"]
    for j, h in enumerate(hdr, 1):
        cell = ws.cell(row=5, column=j, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    r = 6
    for row in rows:
        for j, v in enumerate(row["ids"], 1):
            c = ws.cell(row=r, column=j, value=v)
            c.font = BOLD_FONT if row.get("total") else BODY_FONT
            if row.get("total"):
                c.fill = TOTAL_FILL
        for wi, w in enumerate(windows):
            (sp, shr, ld, rv, roas, gp,
             bsp, bld, brv, broas, bgp) = row[w]
            droas = (roas - broas) if (roas is not None and
                                       broas is not None) else None
            dgp = gp - bgp
            vals = [sp, shr, ld, rv, roas, gp,
                    bsp, bld, brv, broas, bgp, droas, dgp]
            for k, (v, fmt) in enumerate(zip(vals, WIN_FMTS)):
                col = len(row["ids"]) + 1 + wi * len(WIN_COLS) + k
                cell = ws.cell(row=r, column=col,
                               value=(round(v, 4) if isinstance(v, float)
                                      else v))
                cell.number_format = fmt
                cell.border = BORDER
                name = WIN_COLS[k]
                if name.startswith("b_"):
                    cell.fill = BASE_FILL
                if row.get("total"):
                    cell.fill = TOTAL_FILL
                    cell.font = BOLD_FONT
                elif name in ("dGP", "dROAS") and isinstance(v, (int, float)) \
                        and v < 0:
                    cell.font = NEG_FONT
                else:
                    cell.font = BODY_FONT
        fcol = len(row["ids"]) + len(windows) * len(WIN_COLS) + 1
        ws.cell(row=r, column=fcol,
                value=", ".join(row.get("flags", []))).font = NOTE_FONT
        r += 1

    for j in range(1, n_id + 1):
        ws.column_dimensions[get_column_letter(j)].width = \
            38 if j == n_id else 10
    for j in range(n_id + 1, n_id + 1 + len(windows) * len(WIN_COLS) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 10
    ws.freeze_panes = ws.cell(row=6, column=n_id + 1)
    return r - 1


def make_row(ids, blocks, k, parents, pkey, windows, flags=None, total=False):
    row = {"ids": ids, "flags": flags or [], "total": total}
    for w in windows:
        sp, ld, rv, roas, gp = cell_metrics(blocks[w]["cur"], k)
        psp = parents[w]["cur"].get(pkey, [0, 0, 0])[0] if parents else sp
        shr = sp / psp if psp > 0 else None
        bsp, bld, brv, broas, bgp = cell_metrics(blocks[w]["base"], k)
        row[w] = (sp, shr, ld, rv, roas, gp, bsp, bld, brv, broas, bgp)
    return row


def dgp_of(row, w):
    sp, shr, ld, rv, roas, gp, bsp, bld, brv, broas, bgp = row[w]
    return gp - bgp


# ═══════════════════════════════════════════════════════════════════════════
# MAIN
# ═══════════════════════════════════════════════════════════════════════════

def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--db", default=str(DEFAULT_DB))
    ap.add_argument("--end-date", default=None,
                    help="YYYY-MM-DD override; default = detected last "
                         "complete-revenue day")
    ap.add_argument("--out-dir", default=str(DEFAULT_OUT_DIR))
    ap.add_argument("--smoke", action="store_true",
                    help="verification pass only — END detection, source "
                         "verdicts, adjudication, recon gates; no workbook")
    args = ap.parse_args()

    db = Path(args.db)
    if not db.exists():
        die(f"DB not found at {db}")
    con = sqlite3.connect(f"file:{db}?mode=ro", uri=True)

    log("═" * 78)
    log("MS ROAS PAIN DIAGNOSIS — Microsoft Home (account_id=2)"
        + ("   [SMOKE]" if args.smoke else ""))
    log("═" * 78)

    cpd_cols = table_columns(con, "campaign_performance_daily")
    for need in ("date", "device", "conversions", "conversion_value",
                 "all_conversion_value", "cost_micros", "clicks"):
        if need not in cpd_cols:
            die(f"campaign_performance_daily lacks '{need}' "
                f"(columns={cpd_cols})")
    geo_cols = table_columns(con, "geo_performance_daily")

    # ── END + windows ───────────────────────────────────────────────────
    end = detect_end(con, cpd_cols, args.end_date)
    windows = build_windows(end)
    span_a = windows["14d"]["base"][0]
    hist_a = date(2026, 1, 1)         # first-served lookback for new groups
    log("[WIN] windows (current | baseline):")
    for w, rng in windows.items():
        log(f"       {w:>3}: {rng['cur'][0]} → {rng['cur'][1]}  |  "
            f"{rng['base'][0]} → {rng['base'][1]}")

    # ── roster ──────────────────────────────────────────────────────────
    camp_rows = con.execute(
        "SELECT campaign_id, campaign_name, status FROM campaigns "
        "WHERE account_id = ?", (ACCOUNT_ID,)).fetchall()
    camp_name = {to_int_id(cid): nm for cid, nm, _ in camp_rows}
    camp_status = {to_int_id(cid): st for cid, _, st in camp_rows}
    codes, seen = {}, set()
    for cid, nm in camp_name.items():
        c = short_code(nm)
        base, i = c, 2
        while c in seen:
            c = f"{base}{i}"
            i += 1
        seen.add(c)
        codes[cid] = c

    # ── Level 1 ─────────────────────────────────────────────────────────
    l1 = pull_l1(con, span_a, end)
    adj_verdict, _ = adjudicate_l1_revenue(l1)
    l1_blocks = build_blocks(l1, ["campaign_id"], windows)
    l1_parents = parent_totals(l1_blocks)       # ACCOUNT totals
    served = sorted(k[0] for k in all_keys(l1_blocks) if k[0] is not None)
    log(f"[L1] campaigns with activity in span: {len(served)}")

    # ── Level 2 (+ partner slice) ───────────────────────────────────────
    geo, partner_nets = pull_geo(con, geo_cols, span_a, end)
    geo_oo = geo[~geo["is_partner"]]
    # L2 state cells EXCLUDE syndicated-partner rows so they share L1's
    # scope (scope adjudication below proves which scope L1 carries); the
    # partner slice is its own exhibit, never mixed into state cells.
    l2_blocks = build_blocks(geo_oo, ["campaign_id", "geo_state"], windows)
    l2_parents = parent_totals(l2_blocks, key_slice=lambda k: k[0])
    partner_cmp = build_blocks(geo[geo["is_partner"]], ["campaign_id"],
                               windows)
    ptr_daily = (geo[geo["is_partner"]]
                 .groupby("date", as_index=False)
                 [["spend", "clicks", "leads", "rev"]].sum()
                 .sort_values("date"))

    # partner slice in the 14d BASELINE + SCOPE ADJUDICATION: does the
    # campaign table (L1 truth) include partner traffic, or exclude it?
    # Discovered 10/13: a 14d-base geo-vs-L1 gap that exactly equals the
    # partner slice means campaign_performance_daily EXCLUDES partners —
    # the 8/06 pair-sum identity was only ever verified with partners ~0.
    p_base = parent_totals(partner_cmp)["14d"]["base"].get(
        "ACCOUNT", [0, 0, 0])
    t_base = l1_parents["14d"]["base"].get("ACCOUNT", [0, 0, 0])
    gross_base = [0.0, 0.0, 0.0]
    for _k, (sp, ld, rv) in build_blocks(
            geo, ["campaign_id"], windows)["14d"]["base"].items():
        gross_base[0] += sp
        gross_base[1] += ld
        gross_base[2] += rv
    p_shr = p_base[0] / t_base[0] if t_base[0] else 0.0
    log(f"[PTR] syndicated-partner slice in 14d baseline (geo): spend "
        f"${p_base[0]:,.0f} (= {p_shr:.1%} of the L1 baseline), leads "
        f"{p_base[1]:,.1f}")
    scope_gap = gross_base[0] - t_base[0]
    if p_base[0] > 0 and abs(scope_gap - p_base[0]) <= 0.02 * max(t_base[0], 1):
        l1_scope = ("campaign_performance_daily EXCLUDES syndicated-partner "
                    f"rows: geo gross − L1 (${scope_gap:,.0f}) == the "
                    f"partner slice (${p_base[0]:,.0f}). L1/L2 are O&O-"
                    "scope and the 9/18–21 incident never inflated them; "
                    "the slice is quantified on 05_Change_Overlay. RELAY "
                    "TO ARCHITECT: the geo==campaign pair-sum identity "
                    "holds only while partner traffic is ~0.")
        log(f"[SCOPE] {l1_scope}")
    elif abs(scope_gap) <= 0.005 * max(t_base[0], 1):
        l1_scope = ("campaign_performance_daily INCLUDES partner traffic "
                    "(geo gross ties to L1); partner rows still kept out "
                    "of L2 state cells and shown separately")
        log(f"[SCOPE] {l1_scope}")
    else:
        l1_scope = (f"UNRESOLVED: geo gross − L1 = ${scope_gap:,.0f} but "
                    f"partner slice = ${p_base[0]:,.0f} — neither identity "
                    f"holds; take this to the architect before trusting L2")
        log(f"[SCOPE] *** {l1_scope} ***")

    # ── Level 3 ─────────────────────────────────────────────────────────
    l3_source, l3_note = resolve_l3_source(con, windows["14d"]["cur"][0], end)
    l3, first_served, l3_detail = pull_l3(con, l3_source, span_a, end,
                                          hist_a)
    ag_names = pull_ad_group_names(con)
    l3_blocks = build_blocks(l3, ["campaign_id", "ad_group_id"], windows)
    l3_parents = parent_totals(l3_blocks, key_slice=lambda k: k[0])

    def ag_flags_classify(agid):
        nm = ag_names.get(agid, "")
        branded = is_branded(nm)
        fs = first_served.get(agid)
        new_grp = fs is not None and fs >= EXPANSION_LAUNCH
        return nm, branded, new_grp

    # ── reconciliation ──────────────────────────────────────────────────
    log("")
    log("─" * 78)
    log("RECONCILIATION (window sums vs Level-1 campaign truth)")
    recon_rows = []
    for w in windows:
        for side in ("cur", "base"):
            t1 = l1_parents[w][side].get("ACCOUNT", [0, 0, 0])
            for lvl, blocks, thresh in (
                    ("L2 geo", l2_blocks, L2_RECON_WARN),
                    (f"L3 {l3_source}", l3_blocks, 0.03)):
                s = [0.0, 0.0, 0.0]
                for k, (sp, ld, rv) in blocks[w][side].items():
                    s[0] += sp
                    s[1] += ld
                    s[2] += rv
                gaps = []
                for i, mname in enumerate(("spend", "leads", "rev")):
                    g = (s[i] - t1[i]) / t1[i] if t1[i] else 0.0
                    gaps.append((mname, g))
                worst = max(abs(g) for _, g in gaps)
                tag = "OK" if worst <= thresh else "WARN"
                log(f"  [{tag}] {lvl:9} {w:>3} {side:4}: "
                    + "  ".join(f"{m} {g:+.2%}" for m, g in gaps)
                    + f"   [L1 ${t1[0]:,.0f} vs lvl ${s[0]:,.0f}]")
                recon_rows.append((lvl, w, side, t1[0], s[0], t1[1], s[1],
                                   t1[2], s[2], worst, tag))
    log("─" * 78)

    # ── L3 daily coverage vs L1 — acct-2 kpd is batch-loaded from UI
    # exports, so whole dates can be missing or partial; a short BASELINE
    # makes L3 ΔGP a load artifact, not a performance read ──────────────
    l1_daily = l1.groupby("date")["spend"].sum()
    l3_daily = l3.groupby("date")["spend"].sum()
    cov_rows = []
    for d in sorted(l1_daily.index):
        a = float(l1_daily.get(d, 0.0))
        b = float(l3_daily.get(d, 0.0))
        cov_rows.append((d, a, b, (b / a) if a > 0 else None))
    cov_short = [r for r in cov_rows
                 if r[3] is not None and not (0.99 <= r[3] <= 1.01)]
    if cov_short:
        log(f"[L3] DAILY kpd-vs-L1 spend coverage — {len(cov_short)} of "
            f"{len(cov_rows)} span days outside ±1%:")
        for d, a, b, c in cov_short:
            log(f"       {d}  L1 ${a:>8,.0f}   kpd ${b:>8,.0f}   "
                f"coverage {c:>4.0%}")
        log("[L3] dates under 100% = the MS keyword-performance loads are "
            "missing/partial there (refresh the export loads for those "
            "dates and rerun for clean L3 deltas); dates over 100% = kpd "
            "carries traffic L1 excludes (the partner slice)")
    l3_win_cov = {}
    for (lvl, w, side, t1_sp, s_sp, *_rest) in recon_rows:
        if lvl.startswith("L3"):
            l3_win_cov[(w, side)] = (s_sp / t1_sp) if t1_sp else None

    def _pc(v):
        return f"{v:.0%}" if v is not None else "n/a"
    l3_cov_note = "L3/L1 spend coverage (cur/base): " + "; ".join(
        f"{w} {_pc(l3_win_cov.get((w, 'cur')))}/"
        f"{_pc(l3_win_cov.get((w, 'base')))}" for w in windows)
    log(f"[L3] {l3_cov_note}")
    weak_l3 = {ws for ws, c in l3_win_cov.items()
               if c is not None and c < 0.95}

    # ── change events ───────────────────────────────────────────────────
    events = pull_change_events(con, span_a, end, served)

    # ── moved-state detection (dominant campaign changed base→cur, 14d) ─
    def dominant(cells):
        best = {}
        for (cid, st), (sp, _, _) in cells.items():
            if st == "UNKNOWN":
                continue
            if sp > best.get(st, (None, -1))[1]:
                best[st] = (cid, sp)
        return {st: cid for st, (cid, _) in best.items()}

    dom_base = dominant(l2_blocks["14d"]["base"])
    dom_cur = dominant(l2_blocks["14d"]["cur"])
    moved_states = {st for st in dom_base
                    if st in dom_cur and dom_base[st] != dom_cur[st]}
    log(f"[GEO] states whose dominant campaign changed 14d base→cur "
        f"(cutover movers, data-detected): {len(moved_states)}"
        + (f" — {sorted(moved_states)[:12]}"
           + (" ..." if len(moved_states) > 12 else "")
           if moved_states else ""))

    # campaigns with heavy partner baselines (per-campaign flag)
    partner_flag_cids = set()
    for cid in served:
        pb = partner_cmp["14d"]["base"].get((cid,), (0, 0, 0))[0]
        tb = l1_blocks["14d"]["base"].get((cid,), (0, 0, 0))[0]
        if tb > 0 and pb / tb >= PARTNER_FLAG_SHARE:
            partner_flag_cids.add(cid)
    if partner_flag_cids:
        log(f"[PTR] campaigns with ≥{PARTNER_FLAG_SHARE:.0%} partner share "
            f"in 14d baseline: "
            f"{[codes.get(c, c) for c in sorted(partner_flag_cids)]}")

    event_cids = {cid for cid, *_ in events}

    if args.smoke:
        log("")
        log("SMOKE COMPLETE — review END, adjudication verdict, L3 source "
            "verdict, partner share, recon gates above; then run full.")
        return

    # ═══════════════════ build rows per level ═══════════════════════════
    def camp_flags(cid):
        fl = []
        if cid in partner_flag_cids:
            fl.append("PARTNER_BASELINE")
        if cid in event_cids:
            fl.append("EVENT_IN_SPAN")
        if camp_status.get(cid) not in (None, "ENABLED"):
            fl.append(str(camp_status.get(cid)))
        return fl

    # L1 rows
    l1_rows = []
    for cid in served:
        l1_rows.append(make_row(
            [codes.get(cid, str(cid)), camp_name.get(cid, f"id {cid}")],
            l1_blocks, (cid,), l1_parents, "ACCOUNT", windows,
            flags=camp_flags(cid)))
    l1_rows.sort(key=lambda r: dgp_of(r, "7d"))
    acct_row = {"ids": ["", "ACCOUNT TOTAL"], "flags": [], "total": True}
    for w in windows:
        t = l1_parents[w]["cur"].get("ACCOUNT", [0, 0, 0])
        b = l1_parents[w]["base"].get("ACCOUNT", [0, 0, 0])
        roas, gp = derive(t[0], t[1], t[2])
        broas, bgp = derive(b[0], b[1], b[2])
        acct_row[w] = (t[0], 1.0, t[1], t[2], roas, gp,
                       b[0], b[1], b[2], broas, bgp)
    l1_rows.insert(0, acct_row)

    # platform-total context row: geo gross (search O&O + partner slice) —
    # what the platform actually spent/earned including the 9/18–21
    # incident, for comparison against the L1 scope
    gross_l1 = build_blocks(
        geo.groupby(["campaign_id", "date"], as_index=False)
           [["spend", "clicks", "leads", "rev"]].sum(),
        ["campaign_id"], windows)
    gross_parents = parent_totals(gross_l1)
    gross_row = {"ids": ["", "ACCOUNT + syndicated partners (geo gross)"],
                 "flags": ["context row — platform total incl the "
                           "9/18–21 partner incident"], "total": True}
    for w in windows:
        t = gross_parents[w]["cur"].get("ACCOUNT", [0, 0, 0])
        b = gross_parents[w]["base"].get("ACCOUNT", [0, 0, 0])
        roas, gp = derive(t[0], t[1], t[2])
        broas, bgp = derive(b[0], b[1], b[2])
        gross_row[w] = (t[0], None, t[1], t[2], roas, gp,
                        b[0], b[1], b[2], broas, bgp)
    l1_rows.insert(1, gross_row)

    # L2 rows
    l2_rows = []
    for (cid, st) in sorted(all_keys(l2_blocks),
                            key=lambda k: (codes.get(k[0], ""), k[1] or "")):
        fl = []
        if st in moved_states:
            fl.append("MOVED_STATE")
        if st == "UNKNOWN":
            fl.append("UNKNOWN_GEO")
        fl += camp_flags(cid)
        l2_rows.append(make_row(
            [codes.get(cid, str(cid)), st], l2_blocks, (cid, st),
            l2_parents, cid, windows, flags=fl))
    l2_rows.sort(key=lambda r: dgp_of(r, "7d"))

    # L3 rows
    l3_rows = []
    for (cid, agid) in all_keys(l3_blocks):
        nm, branded, new_grp = ag_flags_classify(agid)
        fl = ["BRANDED" if branded else "GENERIC"]
        if new_grp:
            fl.append("NEW_EXPANSION_GROUP")
        fl += camp_flags(cid)
        l3_rows.append(make_row(
            [codes.get(cid, str(cid)), nm or f"ad_group {agid}", agid],
            l3_blocks, (cid, agid), l3_parents, cid, windows, flags=fl))
    l3_rows.sort(key=lambda r: dgp_of(r, "7d"))

    # 03b generic rollup: account-wide + per campaign, split G/NG, and the
    # generic side further split into pre-existing vs new-expansion groups
    def roll_l3(filt):
        agg = {w: {s: defaultdict(lambda: [0.0, 0.0, 0.0])
                   for s in ("cur", "base")} for w in windows}
        for w in windows:
            for side in ("cur", "base"):
                for (cid, agid), (sp, ld, rv) in \
                        l3_blocks[w][side].items():
                    if filt(cid, agid):
                        agg[w][side]["ALL"][0] += sp
                        agg[w][side]["ALL"][1] += ld
                        agg[w][side]["ALL"][2] += rv
        row = {}
        for w in windows:
            t = agg[w]["cur"]["ALL"]
            b = agg[w]["base"]["ALL"]
            roas, gp = derive(t[0], t[1], t[2])
            broas, bgp = derive(b[0], b[1], b[2])
            row[w] = (t[0], None, t[1], t[2], roas, gp,
                      b[0], b[1], b[2], broas, bgp)
        return row

    def seg_row(label, filt, flags=None, total=False):
        r = roll_l3(filt)
        r["ids"] = ["ACCOUNT", label]
        r["flags"] = flags or []
        r["total"] = total
        return r

    def g_of(agid):
        return not is_branded(ag_names.get(agid, ""))

    def new_of(agid):
        fs = first_served.get(agid)
        return fs is not None and fs >= EXPANSION_LAUNCH

    rollup_rows = [
        seg_row("GENERIC (all)", lambda c, a: g_of(a), total=True),
        seg_row("GENERIC — pre-existing book",
                lambda c, a: g_of(a) and not new_of(a)),
        seg_row("GENERIC — 9/18 expansion groups",
                lambda c, a: g_of(a) and new_of(a),
                flags=["NEW_EXPANSION_GROUP"]),
        seg_row("NON-GENERIC (branded)", lambda c, a: not g_of(a),
                total=True),
    ]
    for cid in served:
        for lab, filt in (("GENERIC", lambda c, a, _cid=cid:
                           c == _cid and g_of(a)),
                          ("NON-GENERIC", lambda c, a, _cid=cid:
                           c == _cid and not g_of(a))):
            r = roll_l3(filt)
            r["ids"] = [codes.get(cid, str(cid)), lab]
            r["flags"] = camp_flags(cid)
            r["total"] = False
            rollup_rows.append(r)

    # ═══════════════════ findings ═══════════════════════════════════════
    def findings_for(rows, level_name, id_fmt):
        out = []
        for r in rows:
            if r.get("total"):
                continue
            sp = r["7d"][0]
            bsp = r["7d"][6]
            if sp + bsp < MIN_RANK_SPEND:
                continue
            d7 = dgp_of(r, "7d")
            d14 = dgp_of(r, "14d")
            shr = r["7d"][1]
            fl = list(r.get("flags", []))
            if shr is not None and shr < THIN_SPEND_SHARE:
                fl.append("THIN")
            out.append((d7, d14, level_name, id_fmt(r), shr, fl))
        out.sort(key=lambda x: x[0])
        return out[:12]

    findings = (
        findings_for(l1_rows, "L1 campaign",
                     lambda r: f"{r['ids'][0]} — {r['ids'][1]}")
        + findings_for(l2_rows, "L2 campaign×state",
                       lambda r: f"{r['ids'][0]} / {r['ids'][1]}")
        + findings_for(l3_rows, "L3 campaign×ad group",
                       lambda r: f"{r['ids'][0]} / {r['ids'][1]}")
    )
    if ("7d", "base") in weak_l3 or ("7d", "cur") in weak_l3:
        for ent in findings:
            if ent[2].startswith("L3"):
                ent[5].append("L3_COVERAGE_GAP")

    # ═══════════════════ workbook ═══════════════════════════════════════
    wb = Workbook()
    wb.remove(wb.active)

    # 00_Params
    ws = wb.create_sheet("00_Params")
    ws.cell(row=1, column=1,
            value="MS ROAS Pain Diagnosis — Microsoft Home "
                  "(account_id=2)").font = TITLE_FONT
    params = [
        ("Run date", date.today().isoformat()),
        ("END (last complete-revenue day, verified)", end.isoformat()),
        ("Windows", "; ".join(
            f"{w}: {r['cur'][0]}→{r['cur'][1]} (base {r['base'][0]}→"
            f"{r['base'][1]})" for w, r in windows.items())),
        ("L1 source", "campaign_performance_daily (search-only by design), "
                      "device='ALL', conversions / conversion_value"),
        ("L1 revenue adjudication", adj_verdict),
        ("L2 source", "geo_performance_daily, device='ALL', network != "
                      "'Audience'; location_type summed; full state names; "
                      "UNKNOWN kept as own row"),
        ("Partner slice", f"network literal(s): {partner_nets}; 14d-baseline "
                          f"partner spend ${p_base[0]:,.0f} = {p_shr:.1%} of "
                          f"the L1 baseline; excluded from L2 state cells"),
        ("L1 scope adjudication", l1_scope),
        ("L3 source (verified)", l3_note),
        ("L3 resolved columns / device strategy", l3_detail),
        ("Generic classifier", "ad-group NAME token: any 'Non-Generic'/"
                               "'NonGeneric' variant = branded; else "
                               "generic (parsed_* are Google-only)"),
        ("New-group isolation", f"first served >= {EXPANSION_LAUNCH} "
                                f"(lookback from {hist_a})"),
        ("Moved-state detection", "dominant campaign by spend changed "
                                  "14d base→cur (9/23 cutover)"),
        ("Metric rules", "weighted sum-then-divide only; ROAS=Σrev/Σspend; "
                         "GP=Σrev−Σspend; ranked by ΔGP, never raw ROAS"),
        ("Impression share", "NOT used for diagnosis (0.0999/0.9001 are "
                             "censored bounds)"),
        ("Looker caveat", "no Looker lead counts across 10/2 (partner "
                          "zip-submit definition change); platform "
                          "conversions only"),
    ]
    for i, (k, v) in enumerate(params, 3):
        ws.cell(row=i, column=1, value=k).font = BOLD_FONT
        ws.cell(row=i, column=2, value=v).font = BODY_FONT
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 120

    # 01 / 02 / 03
    write_level_sheet(
        wb.create_sheet("01_Campaign"),
        "Level 1 — Campaign (sorted by 7d ΔGP, worst drag first)",
        ["code", "campaign"], l1_rows, windows,
        "L1 = campaign_performance_daily, which per this run's scope "
        "adjudication (00_Params) carries O&O search only; the 'geo gross' "
        "row adds the syndicated-partner slice back for the platform-total "
        "view. spd_shr = share of account spend in that window.")
    write_level_sheet(
        wb.create_sheet("02_Campaign_State"),
        "Level 2 — Campaign × State (sorted by 7d ΔGP)",
        ["code", "state"], l2_rows, windows,
        "geo_performance_daily, search networks only, SYNDICATED-PARTNER "
        "ROWS EXCLUDED so state cells share L1's scope (the partner slice "
        "is on 05_Change_Overlay). spd_shr = share of the CAMPAIGN's spend "
        "in that window. MOVED_STATE = dominant campaign changed 14d "
        "base→cur (9/23 cutover, data-detected) — state movement, not "
        "auction deterioration. A state at zero may be an exclusion/move, "
        "not lost demand.")
    write_level_sheet(
        wb.create_sheet("03_Campaign_AdGroup"),
        f"Level 3 — Campaign × Ad Group (source: {l3_source}; sorted by "
        f"7d ΔGP)",
        ["code", "ad_group", "ad_group_id"], l3_rows, windows,
        "Joined by ad_group_id (names are non-unique on acct 2). "
        "NEW_EXPANSION_GROUP = first served on/after 9/18 — ramp, not "
        "deterioration; its baseline includes the partner-dominated launch "
        "window. spd_shr = share of the campaign's spend. "
        + l3_cov_note + ". A window side under ~95% coverage means its "
        "cells are short on load coverage (04_Reconcile daily table) — "
        "those ΔGPs are load artifacts as much as performance.")
    write_level_sheet(
        wb.create_sheet("03b_Generic_Rollup"),
        "Generic vs Non-Generic rollup (account + per campaign)",
        ["scope", "segment"], rollup_rows, windows,
        "Generic split further into pre-existing book vs 9/18 expansion "
        "groups so expansion ramp isn't read as account-wide generic "
        "deterioration.")

    # 04_Reconcile
    ws = wb.create_sheet("04_Reconcile")
    ws.cell(row=1, column=1, value="Reconciliation — level sums vs L1 "
                                   "campaign truth").font = TITLE_FONT
    ws.cell(row=2, column=1, value=(
        "Expectation: L2 geo (partner rows excluded, matching L1 scope) "
        "ties exactly. L3 threshold depends on source: agp ≤3%; kpnd gaps "
        "are reported as found. kpd has no network column, so if the "
        "9/18–21 partner traffic lives in kpd, the base-window L3 gaps "
        "will sit near the partner share — that pattern means kpd "
        "INCLUDES partners while L1 excludes them, and L3 baseline cells "
        "in the launch window carry the incident.")).font = NOTE_FONT
    hdr = ["level", "window", "side", "L1 spend", "lvl spend", "L1 leads",
           "lvl leads", "L1 rev", "lvl rev", "worst gap", "status"]
    for j, h in enumerate(hdr, 1):
        c = ws.cell(row=4, column=j, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
    for i, row in enumerate(recon_rows, 5):
        for j, v in enumerate(row, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = BODY_FONT
            if j in (4, 5, 8, 9):
                c.number_format = MONEY
            elif j in (6, 7):
                c.number_format = NUM1
            elif j == 10:
                c.number_format = "0.00%"
            if row[-1] == "WARN":
                c.font = NEG_FONT
    rr = 5 + len(recon_rows) + 2
    ws.cell(row=rr, column=1,
            value="L3 daily spend coverage vs L1 — under 100% = keyword "
                  "load missing/partial that date; over 100% = kpd carries "
                  "the partner slice L1 excludes").font = BOLD_FONT
    rr += 1
    for j, h in enumerate(["date", "L1 spend", "kpd spend", "coverage"], 1):
        c = ws.cell(row=rr, column=j, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
    rr += 1
    for d, a, b, c0 in cov_rows:
        ws.cell(row=rr, column=1, value=d.isoformat()).font = BODY_FONT
        x = ws.cell(row=rr, column=2, value=round(a, 2))
        x.number_format = MONEY
        x.font = BODY_FONT
        x = ws.cell(row=rr, column=3, value=round(b, 2))
        x.number_format = MONEY
        x.font = BODY_FONT
        x = ws.cell(row=rr, column=4, value=c0)
        x.number_format = "0%"
        x.font = (NEG_FONT if (c0 is not None and not 0.95 <= c0 <= 1.05)
                  else BODY_FONT)
        rr += 1
    for j in range(1, 12):
        ws.column_dimensions[get_column_letter(j)].width = 13

    # 05_Change_Overlay
    ws = wb.create_sheet("05_Change_Overlay")
    ws.cell(row=1, column=1, value="Change overlay — confounds, not "
                                   "findings").font = TITLE_FONT
    ws.cell(row=2, column=1, value=(
        "Known operator actions + campaign_change_events (tolerant "
        "tCPA/budget/tROAS match). Window membership shows which cells "
        "each event contaminates.")).font = NOTE_FONT
    hdr = ["date", "campaign", "event", "old", "new", "in windows"]
    for j, h in enumerate(hdr, 1):
        c = ws.cell(row=4, column=j, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT

    def window_membership(dstr):
        d = date.fromisoformat(dstr)
        tags = []
        for w, rng in windows.items():
            if rng["cur"][0] <= d <= rng["cur"][1]:
                tags.append(f"{w}-cur")
            elif rng["base"][0] <= d <= rng["base"][1]:
                tags.append(f"{w}-base")
        return ", ".join(tags) or "outside span"

    r = 5
    for dstr, desc in KNOWN_EVENTS:
        ws.cell(row=r, column=1, value=dstr).font = BOLD_FONT
        ws.cell(row=r, column=2, value="(operator log)").font = NOTE_FONT
        ws.cell(row=r, column=3, value=desc).font = BODY_FONT
        ws.cell(row=r, column=6, value=window_membership(dstr)).font = BODY_FONT
        r += 1
    for cid, dstr, typ, old, new in events:
        ws.cell(row=r, column=1, value=dstr).font = BODY_FONT
        ws.cell(row=r, column=2,
                value=f"{codes.get(cid, cid)} — "
                      f"{camp_name.get(cid, '')}"[:60]).font = BODY_FONT
        ws.cell(row=r, column=3, value=typ).font = BODY_FONT
        ws.cell(row=r, column=4, value=old).font = BODY_FONT
        ws.cell(row=r, column=5, value=new).font = BODY_FONT
        ws.cell(row=r, column=6, value=window_membership(dstr)).font = BODY_FONT
        r += 1

    # syndicated-partner slice, day by day (the quantified 9/18–21 incident)
    r += 1
    ws.cell(row=r, column=1,
            value="Syndicated-partner slice by day (geo; excluded from "
                  "L1/L2 cells per the scope adjudication)").font = BOLD_FONT
    r += 1
    for j, h in enumerate(["date", "spend", "clicks", "leads", "revenue",
                           "in windows"], 1):
        c = ws.cell(row=r, column=j, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
    r += 1
    for _, pr in ptr_daily.iterrows():
        if pr["spend"] <= 0:
            continue
        dstr = pr["date"].isoformat()
        ws.cell(row=r, column=1, value=dstr).font = BODY_FONT
        for j, (col, fmt) in enumerate(
                (("spend", MONEY2), ("clicks", "#,##0"),
                 ("leads", NUM1), ("rev", MONEY2)), 2):
            c = ws.cell(row=r, column=j, value=round(float(pr[col]), 2))
            c.number_format = fmt
            c.font = BODY_FONT
        ws.cell(row=r, column=6, value=window_membership(dstr)).font = BODY_FONT
        r += 1

    widths = [12, 44, 90, 12, 12, 24]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w

    # 06_Findings
    ws = wb.create_sheet("06_Findings")
    ws.cell(row=1, column=1, value="Top GP-drag pockets (ranked by 7d ΔGP "
                                   "within each level)").font = TITLE_FONT
    ws.cell(row=2, column=1, value=(
        f"Cells under ${MIN_RANK_SPEND:,.0f} combined cur+base 7d spend "
        f"excluded. THIN = <{THIN_SPEND_SHARE:.0%} of parent spend — bad "
        f"ROAS on trivial spend isn't where we're hurting. Read every "
        f"pocket against its flags: MOVED_STATE / NEW_EXPANSION_GROUP / "
        f"PARTNER_BASELINE / EVENT_IN_SPAN are overlay confounds, not "
        f"auction findings. PARTNER_BASELINE = the campaign ran ≥"
        f"{PARTNER_FLAG_SHARE:.0%} partner spend (geo) during the 14d "
        f"baseline — L1/L2 cells exclude it, but L3 (kpd) may include it. "
        f"L3_COVERAGE_GAP = the 7d window's kpd load coverage is <95% "
        f"({l3_cov_note}) — prefer L1/L2 for deltas there; see "
        f"04_Reconcile's daily table.")).font = NOTE_FONT
    hdr = ["level", "pocket", "7d ΔGP", "14d ΔGP", "7d spend share",
           "flags / caveats"]
    for j, h in enumerate(hdr, 1):
        c = ws.cell(row=4, column=j, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
    for i, (d7, d14, lvl, pocket, shr, fl) in enumerate(findings, 5):
        ws.cell(row=i, column=1, value=lvl).font = BODY_FONT
        ws.cell(row=i, column=2, value=pocket).font = BODY_FONT
        c = ws.cell(row=i, column=3, value=round(d7, 2))
        c.number_format = MONEY
        c.font = NEG_FONT if d7 < 0 else BODY_FONT
        c = ws.cell(row=i, column=4, value=round(d14, 2))
        c.number_format = MONEY
        c.font = NEG_FONT if d14 < 0 else BODY_FONT
        c = ws.cell(row=i, column=5, value=shr)
        c.number_format = PCT1
        ws.cell(row=i, column=6, value=", ".join(fl)).font = NOTE_FONT
    widths = [20, 56, 12, 12, 14, 48]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w

    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / f"ms_roas_pain_diagnosis_{end.isoformat()}.xlsx"
    wb.save(out)
    log("")
    log(f"[DONE] workbook written: {out}")

    # console topline
    t = l1_parents["7d"]["cur"]["ACCOUNT"]
    b = l1_parents["7d"]["base"]["ACCOUNT"]
    troas, tgp = derive(t[0], t[1], t[2])
    broas, bgp = derive(b[0], b[1], b[2])
    log(f"[7d]  spend ${t[0]:,.0f} (base ${b[0]:,.0f}) | leads {t[1]:,.1f} "
        f"(base {b[1]:,.1f}) | ROAS {troas or 0:.2f} (base {broas or 0:.2f}) "
        f"| GP ${tgp:,.0f} (base ${bgp:,.0f}) | ΔGP ${tgp - bgp:,.0f}")


if __name__ == "__main__":
    main()
