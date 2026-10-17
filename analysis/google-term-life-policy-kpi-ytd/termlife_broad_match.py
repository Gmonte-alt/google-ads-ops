#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
termlife_broad_match.py  (v1)
=============================
DB replication + enhancement of the teammate's Life_Insurance_Broad_Match
workbook. Term Life (account_id=29), May-Oct 2026. READ-ONLY vs
paid_search.db; writes output\\Broad_Match_Deep_Dive.xlsx + console.

Replicates (from the DB instead of UI pulls, so it reconciles and re-runs):
  R1  Monthly totals: account vs broad-keyword value/conversions
      (+ ENHANCEMENT: spend, CPL, broad vs non-broad efficiency)
  R2  By broad keyword x campaign, monthly conversion value
  R3  Converting search terms by broad keyword (with the keyword-vs-term
      reconciliation the original lacked: Google hides terms)

Enhancements:
  E1  THE QUALITY TIMELINE: broad % of conversions vs SQ premium/op by
      month — the test of "is broad the quality driver" (+ native chart)
  E2  CROSS-SERVING MATRIX: keyword side (ethos vs non-ethos broad) x
      query side (ethos-text vs off-brand) — the two leak directions

Scope/conventions: device='ALL'; broad = ki.match_type='BROAD' (targeted
match type); 'value' = Google conversion value = SQ payment proxy, BLIND
to premium quality; prem/op hardcoded from the Looker disposition series;
Sep-29+ contaminated; Oct MTD; term sheets undercount (Google hides
low-volume terms).
"""

import os
import sqlite3
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.drawing.text import RichTextProperties
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

DB_PATH = r"C:\Projects\google-ads-ops\db\paid_search.db"
OUT_XLSX = (r"C:\Projects\google-ads-ops\analysis"
            r"\google-term-life-policy-kpi-ytd\output"
            r"\Broad_Match_Deep_Dive.xlsx")

ACCOUNT_ID = 29
M0, M1 = "2026-05-01", "2026-11-01"
MONTHS = ["2026-05", "2026-06", "2026-07", "2026-08", "2026-09", "2026-10"]
# SQ premium per op, Looker disposition series (pulled ~10/6-7; Aug-Sep
# still maturing +10-15%; Oct MTD near-worthless; Sep-29+ contaminated)
PREM_PER_OP = {"2026-05": 443.47, "2026-06": 434.09, "2026-07": 383.55,
               "2026-08": 349.10, "2026-09": 285.50, "2026-10": 162.19}
NAVY, LIGHT = "1F3864", "D9E2F2"
SERIES_HEX = ["2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4",
              "008300"]
REF_HEX = "898781"

pd.set_option("display.width", 240)
pd.set_option("display.max_columns", 40)
pd.set_option("display.max_rows", 200)
pd.set_option("display.float_format", lambda x: f"{x:,.1f}")


def hr(t):
    print("\n" + "=" * 110)
    print(t)
    print("=" * 110)


def meta(ws, title, source, window, caveat, takeaway=None):
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=13, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A2"] = f"Source: {source}   |   Window: {window}"
    ws["A3"] = f"Caveats: {caveat}"
    for r in (2, 3):
        ws[f"A{r}"].font = Font(italic=True, size=9)
    if takeaway:
        ws["A4"] = f"TAKEAWAY: {takeaway}"
        ws["A4"].font = Font(bold=True, size=10)
        ws["A4"].fill = PatternFill("solid", fgColor=LIGHT)
    ws.freeze_panes = "A6"


def write_block(ws, row, subtitle, df):
    ws.cell(row=row, column=1, value=subtitle).font = Font(bold=True,
                                                           size=11)
    ws.cell(row=row, column=1).fill = PatternFill("solid", fgColor=LIGHT)
    row += 1
    for j, c in enumerate(df.columns, start=1):
        cell = ws.cell(row=row, column=j, value=str(c))
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
    row += 1
    for _, r in df.iterrows():
        for j, v in enumerate(r, start=1):
            if isinstance(v, (np.floating, float)):
                v = None if pd.isna(v) else round(float(v), 1)
            elif isinstance(v, (np.integer,)):
                v = int(v)
            ws.cell(row=row, column=j, value=v)
        row += 1
    for j, c in enumerate(df.columns, start=1):
        width = max([len(str(c))] + [len(str(x)) for x in df[c].head(60)])
        ws.column_dimensions[get_column_letter(j)].width = min(width + 2, 46)
    return row + 2


def write_series_block(ws, start_col, df):
    for j, cname in enumerate(df.columns):
        ws.cell(row=1, column=start_col + j, value=str(cname))
        for i, v in enumerate(df[cname].tolist(), start=2):
            if isinstance(v, (np.floating, float)):
                v = None if pd.isna(v) else round(float(v), 2)
            elif isinstance(v, (np.integer,)):
                v = int(v)
            elif v is not None and not isinstance(v, (int, float)):
                v = str(v)
            ws.cell(row=i, column=start_col + j, value=v)
    return len(df)


def make_line_chart(ws_data, start_col, n_series, n_rows, title, y_title,
                    skip=1):
    ch = LineChart()
    ch.title = title
    data = Reference(ws_data, min_col=start_col + 1,
                     max_col=start_col + n_series,
                     min_row=1, max_row=n_rows + 1)
    cats = Reference(ws_data, min_col=start_col, min_row=2,
                     max_row=n_rows + 1)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    if ch.title is not None:
        ch.title.overlay = False
    ch.legend.position = "r"
    ch.legend.overlay = False
    ch.y_axis.title = y_title
    try:
        ch.y_axis.title.tx.rich.bodyPr = RichTextProperties(
            rot=-5400000, vert="horz")
    except AttributeError:
        pass
    xn = ws_data.cell(row=1, column=start_col).value
    if xn:
        ch.x_axis.title = str(xn)
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.x_axis.tickLblSkip = skip
    ch.x_axis.tickMarkSkip = skip
    ch.layout = Layout(manualLayout=ManualLayout(
        layoutTarget="inner", xMode="edge", yMode="edge",
        x=0.10, y=0.14, w=0.72, h=0.66))
    ch.height = 9.5
    ch.width = 30
    for i, s in enumerate(ch.series):
        s.smooth = False
        s.graphicalProperties.line.width = 20000
        s.graphicalProperties.line.solidFill = SERIES_HEX[i % len(SERIES_HEX)]
    return ch


def main():
    con = sqlite3.connect(f"file:{DB_PATH}?mode=ro", uri=True)

    # ------------------------------------------------------------- R1
    hr("R1  MONTHLY TOTALS — account vs broad keywords "
       "(+ spend/CPL enhancement)")
    acct = pd.read_sql_query("""
        SELECT SUBSTR(date,1,7) AS month,
               SUM(conversion_value) AS acct_value,
               SUM(conversions) AS acct_conv,
               SUM(cost_micros)/1e6 AS acct_spend
        FROM campaign_performance_daily
        WHERE account_id=? AND device='ALL' AND date>=? AND date<?
        GROUP BY 1""", con, params=(ACCOUNT_ID, M0, M1))
    kw = pd.read_sql_query(f"""
        SELECT SUBSTR(p.report_date,1,7) AS month,
               CASE WHEN ki.match_type='BROAD' THEN 'broad'
                    ELSE 'nonbroad' END AS grp,
               CASE WHEN LOWER(ki.keyword_text) LIKE '%ethos%'
                    THEN 'ethos' ELSE 'rest' END AS brand,
               SUM(p.clicks) AS clicks,
               SUM(p.cost_micros)/1e6 AS spend,
               SUM(p.primary_conv_count) AS conv,
               SUM(p.primary_conv_value) AS value
        FROM keyword_performance_daily p
        JOIN keyword_instances ki ON ki.id=p.instance_id
        WHERE ki.account_id={ACCOUNT_ID} AND p.device='ALL'
          AND p.report_date>='{M0}' AND p.report_date<'{M1}'
        GROUP BY 1,2,3""", con)
    br = kw[kw["grp"] == "broad"].groupby("month").sum(numeric_only=True)
    nb = kw[kw["grp"] == "nonbroad"].groupby("month").sum(numeric_only=True)
    eb = kw[(kw["grp"] == "broad") & (kw["brand"] == "ethos")] \
        .groupby("month").sum(numeric_only=True)
    r1 = acct.set_index("month")
    for col, src_ in [("broad", br), ("ethos_broad", eb)]:
        for m_ in ["value", "conv", "clicks", "spend"]:
            r1[f"{col}_{m_}"] = src_[m_]
    r1["broad_%_value"] = 100 * r1["broad_value"] / r1["acct_value"]
    r1["broad_%_conv"] = 100 * r1["broad_conv"] / r1["acct_conv"]
    r1["broad_%_spend"] = 100 * r1["broad_spend"] / r1["acct_spend"]
    r1["broad_CPL"] = r1["broad_spend"] / r1["broad_conv"]
    r1["nonbroad_CPL"] = nb["spend"] / nb["conv"]
    r1["ethos_KW_%_of_broad_value"] = (100 * r1["ethos_broad_value"]
                                    / r1["broad_value"])
    R1 = r1.reset_index()[[
        "month", "acct_value", "acct_conv", "broad_value", "broad_conv",
        "broad_clicks", "broad_spend", "broad_%_value", "broad_%_conv",
        "broad_%_spend", "broad_CPL", "nonbroad_CPL",
        "ethos_KW_%_of_broad_value"]]
    print(R1.round(1).to_string(index=False))
    print("\n  'value' = Google conversion value = SQ payment proxy — "
          "BLIND to premium quality (proven: flat value/lead through the "
          "-34% prem/op collapse). It sizes the broad surface, nothing "
          "more.")

    # ------------------------------------------------------------- E1
    hr("E1  THE QUALITY TIMELINE — broad share vs SQ premium/op")
    e1 = R1[["month", "broad_%_conv"]].copy()
    e1["SQ_prem_per_op"] = e1["month"].map(PREM_PER_OP)
    base_b = e1.iloc[0]["broad_%_conv"]
    base_p = e1.iloc[0]["SQ_prem_per_op"]
    e1["broad_share_idx(May=100)"] = 100 * e1["broad_%_conv"] / base_b
    e1["prem_op_idx(May=100)"] = 100 * e1["SQ_prem_per_op"] / base_p
    print(e1.round(1).to_string(index=False))
    print("\n  THE TEST: if broad were the quality driver, prem/op should "
          "track broad share. It doesn't — July is broad's smallest month "
          "and the quality slide's first full month. Broad's footprint "
          "doubles only in Sep (post-raise): the amplifier, not the "
          "driver. Jun-Aug quality decline ran through exact-served core "
          "queries.")

    # ------------------------------------------------------------- R2
    hr("R2  BY BROAD KEYWORD x CAMPAIGN — monthly conversion value")
    r2 = pd.read_sql_query(f"""
        SELECT ki.keyword_text AS keyword, ki.campaign_name,
               SUBSTR(p.report_date,1,7) AS month,
               SUM(p.primary_conv_value) AS value
        FROM keyword_performance_daily p
        JOIN keyword_instances ki ON ki.id=p.instance_id
        WHERE ki.account_id={ACCOUNT_ID} AND p.device='ALL'
          AND ki.match_type='BROAD'
          AND p.report_date>='{M0}' AND p.report_date<'{M1}'
        GROUP BY 1,2,3""", con)
    r2["camp"] = r2["campaign_name"].str.extract(r"-(CR\d+|TIER\w*)-") \
        .fillna("OTHER")
    pv2 = r2.pivot_table(index=["keyword", "camp"], columns="month",
                         values="value", aggfunc="sum").fillna(0)
    pv2["total"] = pv2.sum(axis=1)
    pv2 = pv2.sort_values("total", ascending=False)
    R2 = pv2.reset_index()
    print(R2.round(0).to_string(index=False))

    # ------------------------------------------------------------- R3
    hr("R3  CONVERTING SEARCH TERMS on broad keywords "
       "(+ reconciliation vs keyword grain)")
    has_sq = con.execute("SELECT 1 FROM sqlite_master "
                         "WHERE name='search_queries'").fetchone()
    q_sel = ("COALESCE(q.query_text, CAST(p.query_id AS TEXT))"
             if has_sq else "CAST(p.query_id AS TEXT)")
    q_join = ("LEFT JOIN search_queries q ON q.id=p.query_id"
              if has_sq else "")
    r3 = pd.read_sql_query(f"""
        SELECT ki.keyword_text AS keyword, {q_sel} AS term,
               SUBSTR(p.report_date,1,7) AS month,
               SUM(p.primary_conv_value) AS value,
               SUM(p.primary_conv_count) AS conv
        FROM search_query_performance p
        JOIN keyword_instances ki ON ki.id=p.instance_id
                                  AND ki.account_id={ACCOUNT_ID}
        {q_join}
        WHERE p.platform='google' AND p.device='ALL'
          AND ki.match_type='BROAD'
          AND p.report_date>='{M0}' AND p.report_date<'{M1}'
        GROUP BY 1,2,3
        HAVING SUM(p.primary_conv_count) > 0""", con)
    pv3 = r3.pivot_table(index=["keyword", "term"], columns="month",
                         values="value", aggfunc="sum").fillna(0)
    pv3["total"] = pv3.sum(axis=1)
    pv3 = pv3.sort_values("total", ascending=False)
    n_terms = len(pv3)
    R3 = pv3.head(50).reset_index()
    print(f"  top 50 of {n_terms:,} converting broad terms:")
    print(R3.round(0).to_string(index=False))

    recon = (r3.groupby("month")["value"].sum()
             .to_frame("term_sheet_value"))
    recon["keyword_grain_value"] = br["value"]
    recon["visible_%"] = (100 * recon["term_sheet_value"]
                          / recon["keyword_grain_value"])
    recon = recon.reset_index()
    print("\n  RECONCILIATION (the original workbook's unlabeled gap): "
          "term-sheet value vs keyword-grain value —")
    print(recon.round(1).to_string(index=False))
    print("  the shortfall = Google's hidden search terms; the term "
          "sheets are a sample, not the total.")

    # ------------------------------------------------------------- R4/R5
    hr("R4  BY BROAD KEYWORD — conversion value by month (keyword grain, "
       "complete)")
    r4 = pd.read_sql_query(f"""
        SELECT ki.keyword_text AS keyword,
               SUBSTR(p.report_date,1,7) AS month,
               SUM(p.primary_conv_value) AS value
        FROM keyword_performance_daily p
        JOIN keyword_instances ki ON ki.id=p.instance_id
        WHERE ki.account_id={ACCOUNT_ID} AND p.device='ALL'
          AND ki.match_type='BROAD'
          AND p.report_date>='{M0}' AND p.report_date<'{M1}'
        GROUP BY 1,2""", con)
    pv4 = r4.pivot_table(index="keyword", columns="month", values="value",
                         aggfunc="sum").fillna(0)
    pv4["total_value"] = pv4.sum(axis=1)
    # total distinct terms per keyword (her parenthetical, from sqp)
    tc_tot = pd.read_sql_query(f"""
        SELECT ki.keyword_text AS keyword,
               COUNT(DISTINCT p.query_id) AS distinct_terms_total
        FROM search_query_performance p
        JOIN keyword_instances ki ON ki.id=p.instance_id
                                  AND ki.account_id={ACCOUNT_ID}
        WHERE p.platform='google' AND p.device='ALL'
          AND ki.match_type='BROAD'
          AND p.report_date>='{M0}' AND p.report_date<'{M1}'
        GROUP BY 1""", con).set_index("keyword")
    pv4 = pv4.join(tc_tot).sort_values("total_value", ascending=False)
    R4 = pv4.reset_index()
    print(R4.round(0).to_string(index=False))

    hr("R5  BY BROAD KEYWORD — DISTINCT SEARCH TERMS by month (matching "
       "width; sqp-visible lower bounds)")
    r5 = pd.read_sql_query(f"""
        SELECT ki.keyword_text AS keyword,
               SUBSTR(p.report_date,1,7) AS month,
               COUNT(DISTINCT p.query_id) AS distinct_terms
        FROM search_query_performance p
        JOIN keyword_instances ki ON ki.id=p.instance_id
                                  AND ki.account_id={ACCOUNT_ID}
        WHERE p.platform='google' AND p.device='ALL'
          AND ki.match_type='BROAD'
          AND p.report_date>='{M0}' AND p.report_date<'{M1}'
        GROUP BY 1,2""", con)
    pv5 = r5.pivot_table(index="keyword", columns="month",
                         values="distinct_terms", aggfunc="sum").fillna(0)
    pv5 = pv5.reindex(pv4.index.tolist()).fillna(0)
    R5 = pv5.reset_index()
    print(R5.round(0).to_string(index=False))
    print("\n  Reading: R5 is each broad keyword's MATCHING WIDTH — how "
          "many distinct searches it served that month. A keyword whose "
          "value (R4) and width (R5) jump together is expanding into new "
          "queries, not deepening on its core. Counts are sqp-visible "
          "lower bounds.")

    # ------------------------------------------------------------- E2
    hr("E2  CROSS-SERVING MATRIX — both leak directions, monthly")
    e2 = pd.read_sql_query(f"""
        SELECT SUBSTR(p.report_date,1,7) AS month,
               CASE WHEN LOWER(ki.keyword_text) LIKE '%ethos%'
                    THEN 'ethos keyword' ELSE 'carrier/other keyword'
               END AS kw_side,
               CASE WHEN LOWER({q_sel}) LIKE '%ethos%'
                    THEN 'ethos query' ELSE 'off-brand query' END AS q_side,
               SUM(p.cost_micros)/1e6 AS spend,
               SUM(p.primary_conv_count) AS conv,
               SUM(p.primary_conv_value) AS value
        FROM search_query_performance p
        JOIN keyword_instances ki ON ki.id=p.instance_id
                                  AND ki.account_id={ACCOUNT_ID}
        {q_join}
        WHERE p.platform='google' AND p.device='ALL'
          AND ki.match_type='BROAD'
          AND p.report_date>='{M0}' AND p.report_date<'{M1}'
        GROUP BY 1,2,3""", con)
    e2["cell"] = e2["kw_side"] + " -> " + e2["q_side"]
    ORDER = ["ethos keyword -> ethos query",
             "ethos keyword -> off-brand query",
             "carrier/other keyword -> ethos query",
             "carrier/other keyword -> off-brand query"]

    def e2_matrix(metric):
        """Cell matrix + the four ratio rows (Gilbert's formula rows,
        baked in so they survive re-runs)."""
        pv = e2.pivot_table(index="cell", columns="month", values=metric,
                            aggfunc="sum").fillna(0)
        pv = pv.reindex([o for o in ORDER if o in pv.index]).fillna(0)
        if not len(pv):
            return pd.DataFrame({"cell": ["(no broad sqp rows)"]})
        zero = pv.iloc[0] * 0

        def row(name):
            return pv.loc[name] if name in pv.index else zero
        ee, eo = row(ORDER[0]), row(ORDER[1])
        ce, co = row(ORDER[2]), row(ORDER[3])
        tot = ee + eo + ce + co
        ratios = pd.DataFrame({
            "% ethos-query of ethos kws": 100 * ee / (ee + eo),
            "% ethos-query of carrier kws (SHADOW DEMAND)":
                100 * ce / (ce + co),
            "% ethos-query of ALL broad": 100 * (ee + ce) / tot,
            "% mis-routed (both off-diagonal cells)":
                100 * (eo + ce) / tot,
        }).T
        return pd.concat([pv, ratios]).reset_index()                  .rename(columns={"index": "cell"})

    e2_blocks = {}
    for metric, label in [("value", "conversion VALUE"),
                          ("spend", "SPEND"),
                          ("conv", "CONVERSIONS")]:
        e2_blocks[metric] = e2_matrix(metric)
        print(f"\n  Serving cells + ratios — {label}:")
        print(e2_blocks[metric].round(1).to_string(index=False))

    # Leak sizing summary (the negatives-proposal number)
    leak = e2[e2["cell"].isin([ORDER[1], ORDER[2]])]
    ls = leak.groupby("month")[["spend", "conv", "value"]].sum()              .rename(columns={"spend": "leak_spend", "conv": "leak_conv",
                              "value": "leak_value"})
    ls["all_broad_spend_visible"] = e2.groupby("month")["spend"].sum()
    ls["leak_%_of_broad_spend"] = (100 * ls["leak_spend"]
                                   / ls["all_broad_spend_visible"])
    ls = ls.reset_index().round(1)
    print("\n  LEAK SIZING (both off-diagonal cells; sqp-visible lower "
          "bounds):")
    print(ls.to_string(index=False))
    print("\n  Reading: 'SHADOW DEMAND' row = ethos queries served by "
          "carrier/other broad keywords — the ethos surge inflating "
          "keyword groups that don't carry its name. 'Mis-routed' = both "
          "off-diagonal cells: ethos keywords buying off-brand traffic + "
          "carrier keywords buying ethos traffic. Negatives fix both "
          "directions; the carrier->ethos cell is routing correction "
          "(the demand is real, the ads/tracking are wrong).")

    # ------------------------------------------------------------- xlsx
    os.makedirs(os.path.dirname(OUT_XLSX), exist_ok=True)
    wb = Workbook()
    ws_rm = wb.active
    ws_rm.title = "Read_Me"
    meta(ws_rm, "How to read this workbook",
         "definitions for every block and key column", "n/a",
         "Applies to every sheet: 'value' = Google conversion value = the "
         "SQ payment proxy — it is BLIND to premium quality (held flat "
         "through the -34% prem/op collapse); 'broad' = targeted match "
         "type BROAD; term-grain tables undercount (Google hides "
         "low-volume terms); Sep-29+ contaminated by the schedule-cut "
         "incident; Oct is month-to-date.")
    guide = pd.DataFrame([
        ("R1", "How big is broad?",
         "Account totals vs broad-keyword totals. broad_%_value/conv/"
         "spend = broad keywords' share of the account. broad_CPL vs "
         "nonbroad_CPL = cost per Google-routed lead. "
         "ethos_KW_%_of_broad_value = share of broad value on keywords "
         "whose TEXT contains 'ethos' (keyword side). For the share of "
         "broad value from searches where the USER TYPED ethos, see "
         "E2's '% ethos-query of ALL broad' row — they differ by the "
         "two cross-serving cells."),
        ("E1", "Is broad the quality driver? (No.)",
         "Broad % of conversions beside SQ premium/op, both indexed "
         "May=100. If broad drove quality the lines would mirror; "
         "instead broad SHRANK into July while quality slid, and only "
         "doubled in Sep after the budget raises. See chart 1."),
        ("R2", "Which broad keywords, in which campaigns?",
         "Conversion value by broad keyword x campaign, monthly. "
         "Keyword grain (complete — includes Google-hidden terms)."),
        ("R3 / R3b", "What searches converted on broad?",
         "Top converting search terms (term grain = a SAMPLE: Google "
         "hides low-volume terms). R3b shows the visible % vs keyword "
         "grain each month — the gap is hidden terms, not an error."),
        ("R4", "Each broad keyword's value over time",
         "Conversion value by keyword x month (keyword grain, "
         "complete). distinct_terms_total = how many different "
         "searches it served over the window (visible only)."),
        ("R5", "Each broad keyword's MATCHING WIDTH over time",
         "Distinct search terms served per keyword per month. Width "
         "jumping with value (R4) = the keyword expanding into new "
         "queries, not deepening on its core. Lower bounds "
         "(suppression grew over the window)."),
        ("E2", "Where did broad traffic actually go? (the leaks)",
         "Four cells: keyword side (ethos vs carrier/other) x query "
         "side (user typed ethos or not). Off-diagonal = mis-routed: "
         "'ethos kw -> off-brand query' = brand keywords buying other "
         "traffic (senior/carrier terms); 'carrier kw -> ethos query' "
         "= SHADOW DEMAND, ethos searches served through carrier "
         "keywords' ads/tracking. Ratio rows: the two 'of ... kws' "
         "rows denominate within their keyword pair; the 'of ALL "
         "broad' and 'mis-routed' rows denominate across all four "
         "cells. Same matrix in value, spend, and conversions. LEAK "
         "SIZING totals the off-diagonal cells; leak_% is of visible "
         "broad spend. All sqp-visible = lower bounds."),
        ("Broad_Chart", "The four pictures",
         "1: broad share vs premium quality (the E1 test). 2: shadow "
         "demand % + mis-routed % of broad spend. 3: top-5 broad "
         "keywords' value by month. 4: the same keywords' matching "
         "width by month."),
    ], columns=["block", "question it answers", "how to read it"])
    write_block(ws_rm, 5, "Block-by-block guide", guide)
    ws_rm.column_dimensions["C"].width = 95
    for row_ in ws_rm.iter_rows(min_row=7, max_col=3):
        for c_ in row_:
            from openpyxl.styles import Alignment
            c_.alignment = Alignment(wrap_text=True, vertical="top")

    ws = wb.create_sheet("Broad_Match")
    meta(ws, "Broad match deep dive — DB replication + quality timeline",
         "kpd x keyword_instances (keyword grain) + sqp (term grain) + "
         "Looker prem/op series",
         "2026-05..2026-10 (Oct MTD)",
         "'value' = Google conversion value = SQ payment proxy, BLIND to "
         "premium quality; broad = targeted match type BROAD; Sep-29+ "
         "contaminated; term sheets undercount vs keyword grain (Google "
         "hides low-volume terms — see reconciliation); prem/op from "
         "Looker, Aug-Sep still maturing +10-15%",
         takeaway="Broad's footprint doubled in Sep (27% of conversions, "
         "post-raise) but quality does NOT track broad share — July is "
         "broad's smallest month and the slide's first full month. Broad "
         "is the September amplifier, not the driver. E2 adds the SHADOW "
         "DEMAND finding: by Sep ~40% of carrier-keyword broad value was "
         "ethos queries — the ethos surge pushed spend up across the "
         "whole account, part of it hidden behind keyword groups that "
         "don't carry its name. Fix = negatives both directions, not "
         "killing broad.")
    r = write_block(ws, 5, "R1 — Monthly totals: account vs broad "
                    "keywords (value/conv/spend/CPL)", R1.round(1))
    r = write_block(ws, r, "E1 — THE QUALITY TIMELINE: broad % of "
                    "conversions vs SQ premium/op (the 'is broad the "
                    "driver?' test — it isn't; see chart)", e1.round(1))
    r = write_block(ws, r, "R2 — Conversion value by broad keyword x "
                    "campaign, monthly", R2.round(0))
    r = write_block(ws, r, f"R3 — Top 50 of {n_terms:,} converting broad "
                    "search terms (value, monthly; sqp-visible only)",
                    R3.round(0))
    r = write_block(ws, r, "R3b — Reconciliation: term-grain vs "
                    "keyword-grain value (gap = Google's hidden terms)",
                    recon.round(1))
    r = write_block(ws, r, "R4 — Conversion value by broad keyword x "
                    "month (keyword grain; distinct_terms_total = "
                    "sqp-visible terms served over the window)",
                    R4.round(0))
    r = write_block(ws, r, "R5 — Distinct search terms by broad keyword "
                    "x month (matching width; sqp-visible lower bounds)",
                    R5.round(0))
    r = write_block(ws, r, "E2 — Cross-serving matrix + ratio rows, "
                    "conversion VALUE (bottom 4 rows are %; 'SHADOW "
                    "DEMAND' = ethos queries behind carrier keywords)",
                    e2_blocks["value"].round(1))
    r = write_block(ws, r, "E2 — Cross-serving matrix + ratio rows, SPEND",
                    e2_blocks["spend"].round(1))
    r = write_block(ws, r, "E2 — Cross-serving matrix + ratio rows, "
                    "CONVERSIONS", e2_blocks["conv"].round(1))
    write_block(ws, r, "E2 — LEAK SIZING: off-diagonal spend/conv/value "
                "and % of visible broad spend mis-routed", ls)

    ws_cd = wb.create_sheet("Chart_Data_B")
    n = write_series_block(
        ws_cd, 1,
        e1[["month", "broad_share_idx(May=100)",
            "prem_op_idx(May=100)"]])
    ws_tc = wb.create_sheet("Broad_Chart")
    meta(ws_tc, "Broad share vs premium quality",
         "R1 + Looker prem/op", "2026-05..2026-10",
         "Indexed May=100 (shared base, never dual-axis); Oct MTD + "
         "contaminated; prem/op maturing",
         takeaway="If broad drove quality, these lines would move "
         "together inverted. They don't: broad share FELL into July "
         "while quality slid, and only exploded in Sep (the raises). "
         "Broad = September amplifier.")
    hc = ws_tc.cell(row=5, column=1,
                    value="Broad % of conversions vs SQ premium/op, "
                          "indexed May 2026 = 100")
    hc.font = Font(bold=True, size=11)
    hc.fill = PatternFill("solid", fgColor=LIGHT)
    ws_tc.add_chart(make_line_chart(ws_cd, 1, 2, n,
                                    "Broad share vs premium quality "
                                    "(May=100)", "Index (May 2026 = 100)"),
                    "A6")
    vm = e2_blocks["value"].set_index("cell")
    sm_ = e2_blocks["spend"].set_index("cell")
    months_cols = [c for c in vm.columns]
    cd2 = pd.DataFrame({
        "month": months_cols,
        "Shadow demand: ethos-query % of carrier kws (value)":
            vm.loc["% ethos-query of carrier kws (SHADOW DEMAND)"].values,
        "Mis-routed % of broad spend":
            sm_.loc["% mis-routed (both off-diagonal cells)"].values,
    })
    n2 = write_series_block(ws_cd, 6, cd2)
    hc2 = ws_tc.cell(row=28, column=1,
                     value="Shadow demand & mis-routing, monthly — ethos "
                           "queries behind carrier keywords, and total "
                           "off-diagonal share of broad spend")
    hc2.font = Font(bold=True, size=11)
    hc2.fill = PatternFill("solid", fgColor=LIGHT)
    ws_tc.add_chart(make_line_chart(ws_cd, 6, 2, n2,
                                    "Shadow demand & mis-routed broad "
                                    "spend", "%"), "A29")
    top5 = pv4.head(5).index.tolist()
    mcols = [c for c in pv4.columns
             if c not in ("total_value", "distinct_terms_total")]

    def kw_chart_block(pv, label):
        d_ = pd.DataFrame({"month": mcols})
        for k in top5:
            if k in pv.index:
                d_[str(k)[:28]] = pv.loc[k, mcols].values
        return d_

    cdv = kw_chart_block(pv4, "value")
    n3 = write_series_block(ws_cd, 11, cdv)
    hc3 = ws_tc.cell(row=51, column=1,
                     value="Top 5 broad keywords — conversion value by "
                           "month")
    hc3.font = Font(bold=True, size=11)
    hc3.fill = PatternFill("solid", fgColor=LIGHT)
    ws_tc.add_chart(make_line_chart(ws_cd, 11, len(cdv.columns) - 1, n3,
                                    "Broad keywords: conversion value by "
                                    "month", "Conversion value ($)"),
                    "A52")

    cdw = kw_chart_block(pv5, "width")
    n4 = write_series_block(ws_cd, 19, cdw)
    hc4 = ws_tc.cell(row=74, column=1,
                     value="Top 5 broad keywords — matching width "
                           "(distinct search terms) by month")
    hc4.font = Font(bold=True, size=11)
    hc4.fill = PatternFill("solid", fgColor=LIGHT)
    ws_tc.add_chart(make_line_chart(ws_cd, 19, len(cdw.columns) - 1, n4,
                                    "Broad keywords: distinct search "
                                    "terms by month",
                                    "Distinct terms (visible)"),
                    "A75")
    ws_cd.sheet_state = "hidden"
    wb.save(OUT_XLSX)
    print(f"\nWorkbook written: {OUT_XLSX}")
    print("Done. Paste console output back to the analysis thread.")
    con.close()


if __name__ == "__main__":
    main()
