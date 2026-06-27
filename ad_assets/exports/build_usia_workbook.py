#!/usr/bin/env python3
"""Generate rsa-load-ready-copy_USIA_v2.xlsx (full 34-group copy library, v11 schema)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
import re
HL,DESC,SL=30,90,25
def vis(t):
    t=re.sub(r"\{KeyWord:([^}]*)\}",lambda m:m.group(1),t,flags=re.I)
    t=re.sub(r"\{Location\([^)]+\):([^}]*)\}",lambda m:m.group(1),t,flags=re.I)
    return len(t)

UB_H=["Get Matched to an Agent","A Licensed Agent Helps You","Find a Better Home Rate",
      "Compare Home Insurers","Trusted by Homeowners","Home Insurance Made Simple",
      "{LOCATION(State):Your State} Home Agents","No Obligation, Free Quote",
      "Talk to a Licensed Agent","{KeyWord:Home Insurance}"]
UB_H_ANGLE=["Agent (differentiator)","Agent/Trust","Comparison","Comparison","Social proof",
            "Anchor","Location+Agent","Trust","Agent","DKI"]
UB_D=["Answer a few questions and get matched with a licensed agent who finds your better rate.",
      "Compare home insurers with a licensed agent's help. Free, no obligation, no pressure.",
      "A licensed agent compares your options and finds coverage that fits your home."]

SWAPS=[("Rates for Seniors","Senior","Mature homeowner"),
       ("Home Rates for Seniors","Senior","Mature homeowner"),
       ("Senior Homeowner? Get a Quote","Senior","Mature homeowner"),
       ("An Agent Finds Your Discounts","Cost","Price-sensitive"),
       ("Right Coverage, Right Price","Cost","Price-sensitive"),
       ("Licensed in Your State","Trust","Trust/proof"),
       ("A-Rated Carriers Only","Trust","Trust/proof")]

LADDERS=[("ANCHOR: brand+product",1,"[Carrier] Home Insurance"),
 ("ANCHOR: brand+product",2,"[Carrier] Home Ins."),
 ("Compare rates",1,"Compare [Carrier] Home Rates"),("Compare rates",2,"[Carrier] Home Rates"),
 ("Compare rates",3,"Compare [Carrier] Rates"),
 ("Existing customer",1,"Have [Carrier]? Compare Home"),("Existing customer",2,"Got [Carrier]? Compare Home"),
 ("Existing customer",3,"[Carrier] Customer? Compare"),
 ("Renewal pain",1,"[Carrier] Home Renewal Up?"),("Renewal pain",2,"[Carrier] Home Rate Hike?"),
 ("Fit question",1,"Is [Carrier] Right for Home?"),("Fit question",2,"Right Home Pick: [Carrier]?"),
 ("Coverage (Does-Carrier-Have set)",1,"[Carrier] Home Coverage"),("Coverage (Does-Carrier-Have set)",2,"What [Carrier] Home Covers"),
 ("Side by side (bare-only)",1,"Compare [Carrier] Side by Side"),("How-compare (bare-only)",1,"How Does [Carrier] Compare?"),
 ("Membership compare ([member])",1,"Compare [member] Home Rates"),("Membership compare ([member])",2,"[member] Home Rates")]

INTENT=[]
def add(iset,atype,text,note=""): INTENT.append((iset,"Ad-group",atype,text,"Unpinned",note))
def iset_block(name,hls,descs,note=""):
    for h in hls: add(name,"Headline",h,note)
    for d in descs: add(name,"Description",d,note)

# Conquest template (ladder-resolved); descriptions tokenized
for t in ["Beat Your [Carrier] Home Quote","Have [Carrier]? Compare Home","[Carrier] Home Insurance",
          "Compare [Carrier] Home Rates","Is [Carrier] Right for Home?"]:
    add("Carrier Conquesting (template)","Headline",t,"ladder-resolved P1->P2->bare->drop")
iset_block("Carrier Conquesting (template)",[],
  ["Compare home insurance with a licensed agent. See how [Carrier] stacks up.",
   "Thinking about [Carrier]? A licensed agent compares your home options for you."])

iset_block("Generic Core",
  ["Home Insurance, Simplified","Compare & Save on Home","See More Than One Quote","Get the Right Coverage","Find Better Home Coverage"],
  ["Compare home insurance the easy way \u2014 a licensed agent finds your better option.",
   "Complete one short form and a licensed agent will compare home insurance options for you."])

iset_block("Compare",
  ["Compare Top Home Insurers","See More Than One Quote","An Agent Finds Your Match","Don't Settle for One Quote","Compare & Save on Home"],
  ["Compare top home insurers with a licensed agent who finds your better option.",
   "Why pick one? A licensed agent lines up your home options and points to the best fit."])

iset_block("Price",
  ["Affordable Home Coverage","An Agent Finds Discounts","Save, Keep Your Coverage","Lower Your Home Rate","Right Coverage, Right Price"],
  ["Lower your home rate without losing coverage \u2014 an agent finds discounts you qualify for.",
   "Cheapest isn't always right. An agent finds real savings on the coverage your home needs."])

iset_block("Cost",
  ["What Home Insurance Costs","See Your Home Estimate","An Agent Prices It for You","Home Insurance, Explained","Know Before You Buy"],
  ["See what home insurance really costs \u2014 a licensed agent prices it exactly for you.",
   "Researching home insurance? A licensed agent gives you a real estimate, no obligation."])

iset_block("Coverage",
  ["Is Your Home Covered?","Get the Right Home Cover","Home Coverage Reviewed","Home Coverage Done Right","Make Sure You're Covered"],
  ["Make sure your home is actually covered \u2014 a licensed agent reviews your policy with you.",
   "Not sure what you need? A licensed agent walks you through your home coverage."])

iset_block("Policy",
  ["Property & Casualty Help","Home Property Coverage","Property Damage Cover","Your Property, Covered","Property Insurance Help","An Agent Explains It"],
  ["Property and casualty coverage made clear \u2014 a licensed agent walks you through it.",
   "Confused by policy terms? A licensed agent explains your home coverage in plain English."],
  "band term = property")

iset_block("Bundle",
  ["Bundle Home & Auto","An Agent Bundles It","Home + Auto, One Agent","Save by Bundling Home","Bundle & Save on Home"],
  ["Bundle home and auto \u2014 a licensed agent sets up your multi-policy savings.",
   "Save more together. A licensed agent bundles your home and auto and finds the discount."])

iset_block("Quote",
  ["Get Your Home Quote","Free Home Insurance Quote","An Agent Quotes It Fast","Start Your Home Quote","Home Quote in Minutes"],
  ["Get your home insurance quote \u2014 or have a licensed agent do it for you in one step.",
   "Free, no-obligation home quote. A licensed agent finds your best rate fast."])

iset_block("Agent Broker",
  ["Talk to a Home Agent","Get Matched to an Agent","A Licensed Home Agent","Work With a Real Agent","Your Home Insurance Agent"],
  ["Get matched with a licensed home insurance agent who finds your better rate.",
   "Work with a real, licensed agent \u2014 they compare your home options and handle it."])

iset_block("Condo",
  ["Condo & HO6 Coverage","Condo Insurance Made Easy","An Agent Covers Your Condo","Compare Condo Insurance","Condo Insurance Quotes"],
  ["Condo or HO6 coverage made simple \u2014 a licensed agent finds the right condo policy.",
   "Compare condo insurance with a licensed agent who knows HO6 policies. Free quote."],
  "band term = condo")

iset_block("Hazard",
  ["Hazard Insurance for Lenders","Need Hazard for Your Loan?","Hazard Cover by Closing","Mortgage Hazard Insurance","Get Hazard Insurance Fast"],
  ["Need hazard insurance for your mortgage? A licensed agent sorts it before closing.",
   "Lender requires hazard coverage? A licensed agent gets your policy fast \u2014 no delays."],
  "band term = hazard; lender/closing intent")

iset_block("Homeowners Ext",
  ["Second Home Coverage","Insure Your Vacation Home","Add a Second-Home Policy","Seasonal Home Insurance","An Agent Covers 2nd Homes"],
  ["Have a second or seasonal home? A licensed agent adds the coverage you need.",
   "Vacation or vacant home? A licensed agent finds the right second-home policy for you."],
  "2nd/seasonal home; endorsement framing")

iset_block("Mobile Manufactured",
  ["Mobile Home Insurance","Manufactured Home Cover","Insure Your Mobile Home","Compare Mobile Home Rates","An Agent Covers Mobile Homes"],
  ["Mobile or manufactured home? A licensed agent helps you compare your coverage options.",
   "Compare mobile home insurance with a licensed agent. Free, no-obligation quote."])

# Peril (band term = peril family; DKI launch, customizer fast-follow)
iset_block("Peril",
  ["{KeyWord:Specialty Coverage} Help","Home Peril Insurance","Specialty Home Coverage","Fire, Flood, Quake & More","Agent Help for High-Risk Homes","High-Risk Home Coverage"],
  ["Need coverage for a specific risk? A licensed agent helps find the right policy.",
   "Fire, flood, quake or storm \u2014 a licensed agent helps you get the cover you need."],
  "DKI launch; customizer fast-follow. band=peril/specialty family")

# Near Me (location)
iset_block("Near Me",
  ["Local Home Insurance Agent","{LOCATION(State):Your State} Agents","Find an Agent Near You","An Agent in Your Area","Compare Local Home Rates"],
  ["Get matched with a licensed home insurance agent near you. Free, no-obligation quote.",
   "A local licensed agent compares your home options and finds a better rate."])

# Brand Generic (single brand; [Carrier] fallback + static Home suffix)
iset_block("Brand Generic (DKI suffix)",
  ["{KeyWord:[Carrier]} Home Insurance","{KeyWord:[Carrier]} Home Rates","{KeyWord:[Carrier]} Home Quotes","Shopping a Carrier? Compare","Compare Home Rates Easily"],
  ["Shopping for insurance? A licensed agent compares home rates and finds your fit.",
   "Tell us about your home and a licensed agent will help compare coverage and rates."],
  "[Carrier] fills DKI fallback; static 'Home X' carries vertical+band; ordinary HL (NOT is_dki)")

# Carrier Other (multi-brand generic conquest, no [Carrier])
iset_block("Carrier Other",
  ["Compare Home Carriers","Compare Your Home Rate","Switch & Save on Home","An Agent Compares Carriers","Better Home Rate, Found"],
  ["Compare home insurance carriers with a licensed agent who finds you a better rate.",
   "Thinking of switching? A licensed agent compares carriers and finds your home fit."])

# Membership / Affinity (AAA/AARP/Costco; AARP senior)
iset_block("Membership / Affinity",
  ["Beyond Your Member Rate","Compare Your Member Deal","Compare Home Member Rates","More Than One Home Option","An Agent Compares Rates"],
  ["Your membership points to one deal. A licensed agent compares it to the whole market.",
   "More than the member rate \u2014 a licensed agent compares your home options for you."])

# Aggregator
iset_block("Aggregator",
  ["Compare More Home Carriers","See All Your Home Options","An Agent, Not Just a List","Real Home Quotes, Compared","More Carriers, One Agent"],
  ["Compare more home carriers with a licensed agent \u2014 real quotes, not just a list.",
   "An agent compares your home options for you. See all your choices in one place, free."])

# ---- validate ----
errs=[]
for i,t in enumerate(UB_H,1):
    if vis(t)>HL: errs.append(("UB_H",t,vis(t)))
for t in UB_D:
    if vis(t)>DESC: errs.append(("UB_D",t,vis(t)))
for t,_,_ in SWAPS:
    if vis(t)>HL: errs.append(("SWAP",t,vis(t)))
for iset,h,atype,text,pin,note in INTENT:
    lim=HL if atype=="Headline" else DESC
    if vis(text)>lim: errs.append((iset,atype,text,vis(text),lim))
if errs:
    print("LENGTH ERRORS:");
    for e in errs: print("  ",e)
    raise SystemExit(1)
print("LENGTH ERRORS: NONE")

# ---- write ----
wb=Workbook(); hf=PatternFill("solid",start_color="1F3864"); hfont=Font(bold=True,color="FFFFFF")
def hdr(ws,row,n):
    for c in range(1,n+1):
        cell=ws.cell(row=row,column=c); cell.fill=hf; cell.font=hfont
ws=wb.active; ws.title="How to Use"
ws["B2"]="USIA Home Insurance RSA \u2014 Load-Ready Copy v2 (agent-anchored, 34 groups)"; ws["B2"].font=Font(bold=True,size=13)
ws["B4"]="Structure"; ws["C4"]="Two RSAs/ad group: RSA-A (agent-broad) + RSA-B (segment). Wiring in usia_cold_start_build.py; copy here."
ws["B5"]="Final URL"; ws["C5"]="https://www.usinsuranceagents.com/"
ws["B6"]="Limits"; ws["C6"]="HL<=30, Desc<=90, Sitelink<=25; DKI/Location measured on fallback."

ws=wb.create_sheet("RSA Ad-Group Load Asset")
cols=["Intent Set","Hierarchy","Asset Type","Asset Text","Chars","Limit","Status","Suggested Pin","Notes"]
ws.append(cols); hdr(ws,1,len(cols))
for iset,hier,atype,text,pin,note in INTENT:
    lim=HL if atype=="Headline" else DESC
    ws.append([iset,hier,atype,text,vis(text),lim,"OK",pin,note])

ws=wb.create_sheet("Carrier Variant Ladders")
ws["A1"]="Carrier Variant Ladders \u2014 length-tiered (USIA). Agent angle carried by backbone."; ws["A1"].font=Font(bold=True)
ws.append([None]*6); ws.append(["Concept","Priority","Headline template","Literal chars","Max name fits","Note"]); hdr(ws,3,6)
for c,p,tmpl in LADDERS:
    tok="[member]" if "[member]" in tmpl else "[Carrier]"
    ws.append([c,p,tmpl,len(tmpl),30-(len(tmpl)-len(tok)),"bare" if "home" not in tmpl.lower() else ""])

ws=wb.create_sheet("Build Script Rules")
for i,(k,v) in enumerate([("Resolution","Home P1->P2->bare->drop; one variant/concept; anchor mandatory."),
  ("Band","Per-group term: home (default), or condo/hazard/property/peril-family for sub-vertical groups. Floor=ceil(30% pool), >=3 families."),
  ("Bare cap","<=4 bare brand-token lines/ad."),
  ("DKI brand_generic","{KeyWord:[Carrier]} + static 'Home X' suffix; ordinary HL so static Home counts to band."),
  ("DKI brand_carrier","Standalone {KeyWord:Home Insurance}."),
  ("State Farm","DKI ALLOWED for USIA."),
  ("Peril","DKI {KeyWord:Specialty Coverage} launch; customizer fast-follow.")],3):
    ws.cell(row=i,column=1,value=k).font=Font(bold=True); ws.cell(row=i,column=2,value=v)

ws=wb.create_sheet("Universal Backbone")
ws["A1"]="Universal Backbone \u2014 agent-anchored; pairs with every intent set"; ws["A1"].font=Font(bold=True)
ws.append([None]*6); ws.append(["Asset Type","Asset Text","Angle (messaging axis)","Chars","Limit","Status"]); hdr(ws,3,6)
for t,a in zip(UB_H,UB_H_ANGLE): ws.append(["Headline",t,a,vis(t),HL,"OK"])
for t in UB_D: ws.append(["Description",t,"Agent/Comparison",vis(t),DESC,"OK"])

ws=wb.create_sheet("Headline Swaps")
ws["A1"]="Headline Swaps \u2014 add strong angles (senior, cost, trust)"; ws["A1"].font=Font(bold=True)
ws.append([None]*6); ws.append(["ADD \u2014 new headlines (distinct angles)",None,None,None,None,None])
ws.append(["Headline","Angle","Intent served","Chars","Limit","Status"]); hdr(ws,4,6)
for t,a,i in SWAPS: ws.append([t,a,i,vis(t),HL,"OK"])

ws=wb.create_sheet("Sitelink Build Plan")
ws["A1"]="Sitelink Build Plan \u2014 account-level core (<=25)"; ws["A1"].font=Font(bold=True)
ws.append([None]*5); ws.append(["Sitelink link text","Chars","Limit","Status","Role / note"]); hdr(ws,3,5)
for t,role in [("Get Matched to an Agent","Agent-as-outcome anchor"),("Rates For Seniors","#1 converter both platforms"),
   ("Compare Home Insurers","Comparison-via-agent"),("Condo & Mobile Homes","Property-type breadth"),
   ("Bundle Home & Auto","Sitelink tier only"),("Get Your Free Quote","Trust / no-obligation")]:
    ws.append([t,len(t),SL,"OK" if len(t)<=SL else "OVER",role])

for s in wb.sheetnames:
    wb[s].column_dimensions["A"].width=34
    if s not in ("How to Use","Build Script Rules"):
        wb[s].column_dimensions["B"].width=34; wb[s].column_dimensions["C"].width=26
wb.save("/home/claude/rsa-load-ready-copy_USIA_v2.xlsx")
print("Workbook v2 written. Intent sets:", len(set(i[0] for i in INTENT)))
