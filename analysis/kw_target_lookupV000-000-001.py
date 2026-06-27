"""
kw_target_lookup.py
====================
Looks up all keyword instances whose text CONTAINS any of the brand
terms in the lookup file, across a fixed set of accounts and a date
range. Useful for competitor brand research — finds every keyword
across the specified accounts that mentions a target brand.

Match logic: substring match on normalized keyword text (case-insensitive).
  Brand term 'freeway insurance'  matches  'cheap freeway insurance quote'
  Brand term 'wise'               matches  'wise insurance', 'wise auto'
                                  also     'penny wise insurance'  (review)

ACCOUNTS
────────
Hardcoded to:
  3863660890  Auto MCC (autoinsurance.org)
  1006729211  Auto Quote (auto.quote.com)
  8886502147  Home Insurance

INPUT FILE FORMAT
─────────────────
Plain CSV (no header) — one brand term per line. Whitespace and case
are normalised. Examples:
  freeway insurance
  bluefire
  cost u less

OUTPUT
──────
Excel workbook with:
  Sheet 1 — All matched instances with performance data
  Sheet 2 — Summary by brand × account
  Sheet 3 — Review (short/generic terms flagged for false-positive risk)
  Sheet 4 — Unmatched brand terms

RUN
───
    conda activate paid-search
    python analysis/kw_target_lookup.py
    python analysis/kw_target_lookup.py --start 2026-03-29 --end 2026-04-27
    python analysis/kw_target_lookup.py --csv path/to/brands.csv
"""

import sqlite3
import argparse
import csv
from pathlib import Path
from datetime import date, timedelta

import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

DB_PATH    = Path(r"C:\Projects\google-ads-ops\db\paid_search.db")
OUTPUT_DIR = Path(r"C:\Projects\google-ads-ops\reports")
DEFAULT_CSV = Path(r"C:\Projects\google-ads-ops\analysis\input\confie-list-of-subsidiearies.csv")

# Hardcoded scope: two auto Google accounts + Home
TARGET_ACCOUNT_IDS = ["3863660890", "1006729211", "8886502147"]

# Default to last 30 complete days
END_DEFAULT   = (date.today() - timedelta(days=1)).isoformat()
START_DEFAULT = (date.today() - timedelta(days=30)).isoformat()

# Review-sheet heuristics
REVIEW_MAX_WORDS = 1   # single-word brand terms get flagged for review
REVIEW_MAX_CHARS = 6   # any term ≤6 chars also flagged

HEADER_FILL = "7030A0"
HEADER_FONT = "FFFFFF"


# ── DB connection ──────────────────────────────────────────────────────────────

def get_conn():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    conn.execute("PRAGMA journal_mode = WAL")
    return conn


# ── Load brand list from CSV ──────────────────────────────────────────────────

def load_brand_list(csv_path: Path) -> list[str]:
    """
    Loads brand terms from a single-column CSV (no header).
    Strips whitespace, lowercases, drops empties, deduplicates while
    preserving order of first occurrence.
    """
    if not csv_path.exists():
        raise FileNotFoundError(f"Brand list not found: {csv_path}")

    seen     = set()
    brands   = []
    with open(csv_path, newline="", encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        for row in reader:
            if not row:
                continue
            term = row[0].strip().lower()
            if not term or term in seen:
                continue
            seen.add(term)
            brands.append(term)

    print(f"  Loaded {len(brands)} unique brand terms from {csv_path.name}")
    return brands


# ── Resolve instances via substring match ─────────────────────────────────────

def resolve_instances(conn, brands: list[str],
                      account_ids: list[str]) -> tuple[pd.DataFrame, list[str]]:
    """
    For each (account, brand) pair, returns ALL keyword_instances whose
    normalized text CONTAINS the brand term.
    """
    # Resolve target accounts
    acct_rows = conn.execute(
        f"""
        SELECT id, platform_account_id, account_name
        FROM accounts
        WHERE platform = 'google'
          AND platform_account_id IN ({','.join('?' * len(account_ids))})
        """,
        account_ids,
    ).fetchall()

    if len(acct_rows) != len(account_ids):
        found = {str(r["platform_account_id"]) for r in acct_rows}
        missing = [a for a in account_ids if a not in found]
        print(f"  WARNING — accounts not found in DB: {missing}")

    acct_map = {
        str(r["platform_account_id"]): {
            "account_id": r["id"], "account_name": r["account_name"],
        }
        for r in acct_rows
    }

    print(f"  Searching {len(acct_map)} accounts × {len(brands)} brands "
          f"= {len(acct_map) * len(brands)} sub-queries")

    results       = []
    matched_brands = set()

    for pid, acct in acct_map.items():
        account_id   = acct["account_id"]
        account_name = acct["account_name"]

        for brand in brands:
            pattern = f"%{brand}%"
            ki_rows = conn.execute("""
                SELECT
                    ki.id            AS instance_id,
                    ki.criterion_id,
                    ki.ad_group_id,
                    ki.campaign_id,
                    ki.keyword_text,
                    ki.match_type,
                    ki.campaign_name,
                    ki.ad_group_name,
                    ki.status        AS kw_status,
                    ki_id.normalized_text
                FROM keyword_instances ki
                JOIN keyword_identities ki_id ON ki_id.id = ki.identity_id
                WHERE ki.account_id = ?
                  AND LOWER(ki_id.normalized_text) LIKE ?
            """, (account_id, pattern)).fetchall()

            if ki_rows:
                matched_brands.add(brand)

            for ki in ki_rows:
                results.append({
                    "account_id_str":  pid,
                    "account_name":    account_name,
                    "account_id":      account_id,
                    "matched_brand":   brand,
                    "instance_id":     ki["instance_id"],
                    "criterion_id":    ki["criterion_id"],
                    "campaign_id":     ki["campaign_id"],
                    "campaign_name":   ki["campaign_name"],
                    "ad_group_id":     ki["ad_group_id"],
                    "ad_group_name":   ki["ad_group_name"],
                    "keyword_text":    ki["keyword_text"],
                    "match_type":      ki["match_type"],
                    "kw_status":       ki["kw_status"],
                    "normalized_text": ki["normalized_text"],
                })

    unmatched_brands = [b for b in brands if b not in matched_brands]

    print(f"  Instance hits     : {len(results):,}")
    print(f"  Brands with hits  : {len(matched_brands)}")
    print(f"  Brands w/o hits   : {len(unmatched_brands)}")

    return pd.DataFrame(results), unmatched_brands


# ── Fetch performance ─────────────────────────────────────────────────────────

def fetch_performance(conn, instance_ids: list, start: str, end: str) -> pd.DataFrame:
    if not instance_ids:
        return pd.DataFrame()

    all_rows = []
    chunk_size = 500
    for i in range(0, len(instance_ids), chunk_size):
        chunk = instance_ids[i:i + chunk_size]
        ph    = ",".join("?" * len(chunk))
        rows  = conn.execute(f"""
            SELECT
                kpd.instance_id,
                SUM(kpd.impressions)          AS impressions,
                SUM(kpd.clicks)               AS clicks,
                SUM(kpd.cost_micros) / 1e6    AS spend,
                SUM(kpd.primary_conv_count)   AS conversions,
                SUM(kpd.primary_conv_value)   AS revenue
            FROM keyword_performance_daily kpd
            WHERE kpd.instance_id IN ({ph})
              AND kpd.device      = 'ALL'
              AND kpd.geo_state   = 'ALL'
              AND kpd.report_date BETWEEN ? AND ?
            GROUP BY kpd.instance_id
        """, chunk + [start, end]).fetchall()
        all_rows.extend(rows)

    if not all_rows:
        return pd.DataFrame(columns=[
            "instance_id","impressions","clicks","spend","conversions","revenue"
        ])
    return pd.DataFrame([dict(r) for r in all_rows])


# ── Review-sheet heuristic ────────────────────────────────────────────────────

def build_review_df(results_df: pd.DataFrame, brands: list[str]) -> pd.DataFrame:
    """
    Flags brand terms that are short or single-word — false-positive risk
    candidates. Lists them with their match count and a sample of matched
    keyword texts so user can spot obvious false positives at a glance.
    """
    rows = []
    for brand in brands:
        word_count = len(brand.split())
        char_count = len(brand)
        if word_count > REVIEW_MAX_WORDS and char_count > REVIEW_MAX_CHARS:
            continue  # not a review candidate

        sub = results_df[results_df["matched_brand"] == brand]
        if sub.empty:
            sample_texts = ""
            n_matches    = 0
            n_unique     = 0
        else:
            unique_texts = (sub["normalized_text"]
                            .drop_duplicates()
                            .head(8)
                            .tolist())
            sample_texts = " | ".join(unique_texts)
            n_matches    = len(sub)
            n_unique     = sub["normalized_text"].nunique()

        rows.append({
            "brand_term":         brand,
            "word_count":         word_count,
            "char_count":         char_count,
            "instance_matches":   n_matches,
            "unique_kw_texts":    n_unique,
            "sample_kw_texts":    sample_texts,
        })

    df = pd.DataFrame(rows)
    if not df.empty:
        df = df.sort_values(["instance_matches", "brand_term"],
                            ascending=[False, True]).reset_index(drop=True)
    return df


# ── Excel output ──────────────────────────────────────────────────────────────

def write_workbook(out_path: Path, results_df: pd.DataFrame,
                   summary_df: pd.DataFrame, review_df: pd.DataFrame,
                   unmatched_brands: list[str], start: str, end: str):

    wb = openpyxl.Workbook()
    hfill = PatternFill("solid", fgColor=HEADER_FILL)
    hfont = Font(color=HEADER_FONT, bold=True)

    def write_sheet(ws, df, title=None):
        if title:
            ws.title = title
        if df is None or df.empty:
            ws.append(["No data"])
            return
        for ci, col in enumerate(df.columns, 1):
            cell = ws.cell(row=1, column=ci, value=col)
            cell.fill = hfill
            cell.font = hfont
            cell.alignment = Alignment(horizontal="center")
        for ri, row in enumerate(df.itertuples(index=False), 2):
            for ci, val in enumerate(row, 1):
                ws.cell(row=ri, column=ci, value=val)
        for ci, col in enumerate(df.columns, 1):
            max_len = max(
                len(str(col)),
                df.iloc[:, ci-1].astype(str).str.len().max() if len(df) > 0 else 0
            )
            ws.column_dimensions[get_column_letter(ci)].width = min(max_len + 2, 60)
        ws.freeze_panes = ws.cell(row=2, column=1)

    # Sheet 1 — all instances + performance
    ws1 = wb.active
    write_sheet(ws1, results_df, title=f"Instances {start} to {end}")

    # Sheet 2 — summary by brand × account
    ws2 = wb.create_sheet(title="Summary by Brand")
    write_sheet(ws2, summary_df)

    # Sheet 3 — review (short/generic terms)
    ws3 = wb.create_sheet(title="Review (FP risk)")
    if review_df is None or review_df.empty:
        ws3.append(["No short or single-word brand terms in lookup."])
    else:
        write_sheet(ws3, review_df)

    # Sheet 4 — unmatched
    ws4 = wb.create_sheet(title="Unmatched Brands")
    if not unmatched_brands:
        ws4.append(["All brand terms matched at least one keyword."])
    else:
        ws4.append(["brand_term"])
        ws4["A1"].fill = hfill
        ws4["A1"].font = hfont
        for b in unmatched_brands:
            ws4.append([b])
        ws4.column_dimensions["A"].width = 40

    wb.save(out_path)
    print(f"\n  Saved: {out_path}")


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    p = argparse.ArgumentParser()
    p.add_argument("--csv",   default=str(DEFAULT_CSV))
    p.add_argument("--start", default=START_DEFAULT)
    p.add_argument("--end",   default=END_DEFAULT)
    args = p.parse_args()

    start = args.start
    end   = args.end

    print("=" * 70)
    print("kw_target_lookup.py  [substring match — multi-account brand scan]")
    print(f"  Database : {DB_PATH}")
    print(f"  Brand list : {Path(args.csv).name}")
    print(f"  Accounts : {', '.join(TARGET_ACCOUNT_IDS)}")
    print(f"  Window   : {start} → {end}")
    print("=" * 70)

    conn = get_conn()

    print("\n[1] Loading brand list...")
    brands = load_brand_list(Path(args.csv))

    print("\n[2] Resolving instances (substring match)...")
    instances_df, unmatched_brands = resolve_instances(
        conn, brands, TARGET_ACCOUNT_IDS
    )

    if instances_df.empty:
        print("  No instances matched any brand term.")
        conn.close()
        return

    print(f"\n[3] Fetching performance ({start} → {end})...")
    iids    = instances_df["instance_id"].drop_duplicates().tolist()
    perf_df = fetch_performance(conn, iids, start, end)
    print(f"  Performance rows: {len(perf_df):,}")

    # Merge performance onto instances. Note: same instance can match
    # multiple brand terms (e.g. 'family auto insurance' matches both
    # 'family auto' and 'auto'), so a row per (instance × brand) is kept.
    results_df = instances_df.merge(perf_df, on="instance_id", how="left")
    for col in ["impressions","clicks","spend","conversions","revenue"]:
        results_df[col] = results_df[col].fillna(0)

    # Derived metrics
    results_df["ctr"]  = (results_df["clicks"] / results_df["impressions"]
                          ).where(results_df["impressions"] > 0).round(4)
    results_df["cpc"]  = (results_df["spend"] / results_df["clicks"]
                          ).where(results_df["clicks"] > 0).round(2)
    results_df["cpa"]  = (results_df["spend"] / results_df["conversions"]
                          ).where(results_df["conversions"] > 0).round(2)
    results_df["roas"] = (results_df["revenue"] / results_df["spend"]
                          ).where(results_df["spend"] > 0).round(4)
    results_df["spend"]   = results_df["spend"].round(2)
    results_df["revenue"] = results_df["revenue"].round(2)

    out_cols = [
        "account_id_str","account_name",
        "matched_brand",
        "normalized_text","keyword_text","match_type","kw_status",
        "campaign_name","campaign_id",
        "ad_group_name","ad_group_id",
        "instance_id","criterion_id",
        "impressions","clicks","spend","conversions","revenue",
        "ctr","cpc","cpa","roas",
    ]
    results_df = results_df[[c for c in out_cols if c in results_df.columns]]

    # Summary by brand × account
    # NOTE: spend/revenue here are summed across instance rows. If the
    # same instance matched two brand terms it appears in both brand
    # rollups — that's the right behaviour for "how much spend touched
    # brand X" but means brand totals can sum to MORE than the unique
    # instance total. The terminal totals below de-dupe to unique instances.
    summary_df = (
        results_df.groupby(
            ["account_id_str","account_name","matched_brand"], as_index=False)
        .agg(
            instance_matches=("instance_id","count"),
            unique_kw_texts=("normalized_text","nunique"),
            enabled_instances=("kw_status", lambda x: (x=="ENABLED").sum()),
            impressions=("impressions","sum"),
            clicks=("clicks","sum"),
            spend=("spend","sum"),
            conversions=("conversions","sum"),
            revenue=("revenue","sum"),
        )
        .sort_values(["account_name","spend"], ascending=[True, False])
        .reset_index(drop=True)
    )
    summary_df["spend"]   = summary_df["spend"].round(2)
    summary_df["revenue"] = summary_df["revenue"].round(2)

    # Review sheet for false-positive risk
    review_df = build_review_df(results_df, brands)

    # Terminal totals — de-duped to unique instances
    unique_results = results_df.drop_duplicates(subset=["instance_id"])
    print(f"\n[4] Results:")
    print(f"  Unique instances matched : {len(unique_results):,}")
    print(f"  Total instance × brand rows: {len(results_df):,}  "
          f"(higher because same kw can match multiple brands)")
    print(f"  ENABLED instances        : "
          f"{(unique_results['kw_status']=='ENABLED').sum():,}")
    print(f"  PAUSED instances         : "
          f"{(unique_results['kw_status']=='PAUSED').sum():,}")
    print(f"  Total spend (unique)     : ${unique_results['spend'].sum():,.2f}")
    print(f"  Total clicks (unique)    : {unique_results['clicks'].sum():,.0f}")
    print(f"  Total revenue (unique)   : ${unique_results['revenue'].sum():,.2f}")

    print(f"\n  Top 5 brands by spend:")
    top = (summary_df.groupby("matched_brand", as_index=False)
           .agg(spend=("spend","sum"),
                instance_matches=("instance_matches","sum"))
           .sort_values("spend", ascending=False).head(5))
    for _, r in top.iterrows():
        print(f"    {r['matched_brand']:<30}  "
              f"${r['spend']:>12,.2f}  {r['instance_matches']} matches")

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    out_path = OUTPUT_DIR / f"kw_target_brand_lookup_{start}_to_{end}.xlsx"
    write_workbook(out_path, results_df, summary_df, review_df,
                   unmatched_brands, start, end)

    conn.close()
    print("Done.")


if __name__ == "__main__":
    main()