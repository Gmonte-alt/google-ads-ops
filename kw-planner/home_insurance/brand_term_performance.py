#!/usr/bin/env python
"""
brand_term_performance.py -- existing-account performance on CARRIER BRAND search terms,
split BARE (no home token, e.g. "geico", "travelers insurance") vs HOME-QUALIFIED
("geico home insurance"), to size the bare-brand generic campaign before we commit it.

Source (real paid_search.db schema):
  search_query_performance  -- the search-TERM performance facts (cost_micros, primary_conv_*)
    JOIN search_queries     -- query_text + tags (modifier_brand, brand_category, home modifiers)
    account-scoped via the campaign set in keyword_instances (account_id + campaign_id)

Conventions: account scoping, device='ALL' only (auto-detected; if the table is device-
segmented with no ALL rollup, it sums across devices instead), cost = cost_micros/1e6, and
primary_conv_count / primary_conv_value as the truth metrics (pixel-as-primary). Read-only.

    python brand_term_performance.py --db ...\\paid_search.db --days 90 --account 1
    python brand_term_performance.py --db ...\\paid_search.db --scan        # dump schema

Brands come from the query tags (modifier_brand / brand_category), so aggregators and affinity
brands are dropped by category, not by a hardcoded list.
"""
import argparse
import csv
import re
import sqlite3
from collections import defaultdict

CONFIG = {
    "perf_table":     "search_query_performance",
    "perf_query_id":  "query_id",
    "perf_campaign":  "campaign_id",
    "perf_date":      "report_date",
    "perf_device":    "device",
    "cost_micros":    "cost_micros",
    "impr_col":       "impressions",
    "clicks_col":     "clicks",
    "conv_count_col": "primary_conv_count",
    "conv_value_col": "primary_conv_value",
    "cat_table":      "search_queries",
    "cat_id":         "id",
    "cat_text":       "query_text",
    "cat_brand":      "modifier_brand",
    "cat_category":   "brand_category",
    "inst_table":     "keyword_instances",
    "inst_account":   "account_id",
    "inst_campaign":  "campaign_id",
}

HOME_RE = re.compile(r"\b(home|homes|homeowners?|house|houses|condo|condos|dwelling|property)\b")
DROP_CATS = ("aggregator", "affinity")        # we want carriers, not lead-buyers / affinity


def cols(conn, table):
    try:
        return [r[1] for r in conn.execute("PRAGMA table_info(%s)" % table)]
    except sqlite3.OperationalError:
        return []


def scan_schema(conn):
    METRIC = ("cost_micros", "cost", "clicks", "impressions", "conversions",
              "conversion_value", "primary_conv_count", "primary_conv_value", "all_conv")
    KEY = ("query_text", "search_term", "keyword", "keyword_text", "criterion",
           "normalized_text", "criterion_id", "keyword_id", "query_id", "instance_id")
    tables = [r[0] for r in conn.execute(
        "SELECT name FROM sqlite_master WHERE type='table' ORDER BY name")]
    print("=== %d tables ===\n" % len(tables))
    for t in tables:
        cs = cols(conn, t)
        try:
            n = conn.execute("SELECT COUNT(*) FROM %s" % t).fetchone()[0]
        except sqlite3.OperationalError:
            n = "?"
        tags = []
        if any(any(m == c or m in c for m in METRIC) for c in cs):
            tags.append("PERF")
        if any(any(k == c for k in KEY) for c in cs):
            tags.append("KEY")
        print("  %-40s rows=%-9s %s" % (t, n, " ".join(tags)))
        if tags:
            print("      cols: %s" % ", ".join(cs))


def introspect(conn, C):
    tables = [r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")]
    problems = []
    checks = [("perf_table", ("perf_query_id", "perf_campaign", "perf_date", "perf_device",
                              "cost_micros", "impr_col", "clicks_col", "conv_count_col", "conv_value_col")),
              ("cat_table", ("cat_id", "cat_text", "cat_brand", "cat_category")),
              ("inst_table", ("inst_account", "inst_campaign"))]
    for tkey, colkeys in checks:
        if C[tkey] not in tables:
            problems.append("table '%s' not found. Run --scan." % C[tkey])
            continue
        tc = cols(conn, C[tkey])
        for ck in colkeys:
            if C[ck] not in tc:
                problems.append("%s '%s' not in %s. Columns: %s" % (ck, C[ck], C[tkey], tc))
    return (not problems), problems


def device_clause(conn, C):
    """device='ALL' if an ALL rollup exists; else no device filter (sum across devices)."""
    try:
        devs = {(r[0] or "").upper() for r in conn.execute(
            "SELECT DISTINCT %s FROM %s LIMIT 50" % (C["perf_device"], C["perf_table"]))}
    except sqlite3.OperationalError:
        devs = set()
    if "ALL" in devs:
        return " AND sqp.%s='ALL'" % C["perf_device"], "device='ALL' rollup"
    return "", "summed across devices (no ALL rollup found: %s)" % sorted(devs)


def build(db, days, account, out):
    conn = sqlite3.connect(db)
    C = CONFIG
    ok, problems = introspect(conn, C)
    if problems:
        print("SCHEMA NOTES:")
        for p in problems:
            print("  - " + p)
    if not ok:
        print("\nStopping: fix CONFIG and re-run (or --scan). Nothing queried.")
        conn.close()
        return

    dev_sql, dev_note = device_clause(conn, C)
    print("device handling: %s" % dev_note)

    q = """
        SELECT sq.{text}, LOWER(COALESCE(sq.{brand},'')) AS brand, LOWER(COALESCE(sq.{cat},'')) AS cat,
               SUM(sqp.{cost})/1000000.0 AS cost, SUM(sqp.{impr}) AS impr, SUM(sqp.{clicks}) AS clicks,
               SUM(sqp.{cc}) AS conv, SUM(sqp.{cv}) AS value
        FROM {perf} sqp
        JOIN {cat_t} sq ON sqp.{qid} = sq.{cid}
        WHERE sqp.{camp} IN (SELECT DISTINCT {ic} FROM {inst} WHERE {ia}=?)
          AND sqp.{date} >= date('now','-{days} days'){dev}
        GROUP BY sq.{cid}
        HAVING SUM(sqp.{clicks}) > 0
    """.format(text=C["cat_text"], brand=C["cat_brand"], cat=C["cat_category"],
               cost=C["cost_micros"], impr=C["impr_col"], clicks=C["clicks_col"],
               cc=C["conv_count_col"], cv=C["conv_value_col"], perf=C["perf_table"],
               cat_t=C["cat_table"], qid=C["perf_query_id"], cid=C["cat_id"],
               camp=C["perf_campaign"], ic=C["inst_campaign"], inst=C["inst_table"],
               ia=C["inst_account"], date=C["perf_date"], days=int(days), dev=dev_sql)
    rows = conn.execute(q, (account,)).fetchall()
    conn.close()

    # tag: carrier (from query tags, drop aggregator/affinity) + bare vs home-qualified
    agg = defaultdict(lambda: [0.0, 0, 0, 0.0, 0.0])
    per_term = []
    n_total, n_carrier = 0, 0
    for term, brand, cat, cost, impr, clicks, conv, value in rows:
        n_total += 1
        if not brand or cat in DROP_CATS:
            continue
        n_carrier += 1
        is_bare = not HOME_RE.search((term or "").lower())
        a = agg[(brand, is_bare)]
        a[0] += cost or 0; a[1] += impr or 0; a[2] += clicks or 0
        a[3] += conv or 0; a[4] += value or 0
        if is_bare:
            per_term.append((brand, term, cost or 0, clicks or 0, conv or 0, value or 0))

    print("queries with clicks: %d | carrier-brand queries (after dropping aggregator/affinity): %d"
          % (n_total, n_carrier))
    if n_carrier == 0:
        print("\nNo carrier-brand search terms returned. Either the account doesn't buy them, or the "
              "tags/sync are sparse -> commit the generic campaign conservatively.")
        return
    write_outputs(out, agg, per_term, days, account)


def _rates(cost, clicks, conv, value):
    return (cost / clicks if clicks else 0, conv / clicks if clicks else 0,
            cost / conv if conv else 0, value / cost if cost else 0)


def write_outputs(out, agg, per_term, days, account):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    HEAD = PatternFill("solid", fgColor="1F3864")
    WHITE = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    BOLD = Font(name="Arial", bold=True, size=10)
    REG = Font(name="Arial", size=10)
    TITLE = Font(name="Arial", bold=True, size=13, color="1F3864")

    def hdr(ws, row, hs):
        for j, h in enumerate(hs, 1):
            c = ws.cell(row, j, h); c.fill = HEAD; c.font = WHITE
            c.alignment = Alignment(horizontal="center", wrap_text=True)

    wb = Workbook(); wb.remove(wb.active)

    ws = wb.create_sheet("Bare vs Home-Qualified")
    ws["A1"] = "Carrier brand search terms -- BARE vs HOME-QUALIFIED (acct %s, last %dd)" % (account, days)
    ws["A1"].font = TITLE; ws.merge_cells("A1:I1")
    hdr(ws, 3, ["Segment", "Cost", "Impr", "Clicks", "Conv", "Value", "CvR", "CPA", "ROAS"])
    tot = defaultdict(lambda: [0.0, 0, 0, 0.0, 0.0])
    for (brand, is_bare), a in agg.items():
        t = tot["BARE" if is_bare else "HOME-QUALIFIED"]
        for i in range(5):
            t[i] += a[i]
    r = 4
    for seg in ("BARE", "HOME-QUALIFIED"):
        cost, impr, clicks, conv, value = tot[seg]
        cpc, cvr, cpa, roas = _rates(cost, clicks, conv, value)
        ws.cell(r, 1, seg).font = BOLD
        for j, v in enumerate([round(cost, 2), impr, clicks, round(conv, 1), round(value, 2)], 2):
            ws.cell(r, j, v).font = REG
        ws.cell(r, 7, round(cvr, 4)).font = REG
        ws.cell(r, 8, round(cpa, 2)).font = REG
        ws.cell(r, 9, round(roas, 2)).font = REG
        r += 1
    for col, w in {"A": 18, "B": 12, "C": 11, "D": 10, "E": 9, "F": 12, "G": 9, "H": 10, "I": 9}.items():
        ws.column_dimensions[col].width = w

    ws = wb.create_sheet("By Carrier (bare)")
    ws["A1"] = "Bare carrier terms by carrier -- the generic-campaign shortlist"; ws["A1"].font = TITLE
    ws.merge_cells("A1:H1")
    hdr(ws, 3, ["Carrier", "Cost", "Clicks", "Conv", "Value", "CvR", "CPA", "ROAS"])
    bycar = defaultdict(lambda: [0.0, 0, 0, 0.0, 0.0])
    for (brand, is_bare), a in agg.items():
        if is_bare:
            b = bycar[brand]
            for i in range(5):
                b[i] += a[i]
    r = 4
    for brand, b in sorted(bycar.items(), key=lambda x: -x[1][0]):
        cost, impr, clicks, conv, value = b
        cpc, cvr, cpa, roas = _rates(cost, clicks, conv, value)
        ws.cell(r, 1, brand).font = BOLD
        for j, v in enumerate([round(cost, 2), clicks, round(conv, 1), round(value, 2)], 2):
            ws.cell(r, j, v).font = REG
        ws.cell(r, 6, round(cvr, 4)).font = REG
        ws.cell(r, 7, round(cpa, 2)).font = REG
        ws.cell(r, 8, round(roas, 2)).font = REG
        r += 1
    for col, w in {"A": 20, "B": 12, "C": 10, "D": 9, "E": 12, "F": 9, "G": 10, "H": 9}.items():
        ws.column_dimensions[col].width = w
    wb.save(out)

    csv_path = out.rsplit(".", 1)[0] + "_bare_terms.csv"
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["carrier", "search_term", "cost", "clicks", "conv", "value"])
        for row in sorted(per_term, key=lambda x: -x[2]):
            w.writerow([row[0], row[1], round(row[2], 2), row[3], round(row[4], 1), round(row[5], 2)])

    print("\n=== BARE vs HOME-QUALIFIED (acct %s, last %dd) ===" % (account, days))
    for seg in ("BARE", "HOME-QUALIFIED"):
        cost, impr, clicks, conv, value = tot[seg]
        cpc, cvr, cpa, roas = _rates(cost, clicks, conv, value)
        print("  %-15s cost=$%-10s clicks=%-7d conv=%-7s CvR=%.2f%% CPA=$%-7s ROAS=%.2f"
              % (seg, f"{cost:,.0f}", int(clicks), f"{conv:,.0f}", 100 * cvr,
                 f"{cpa:,.0f}" if conv else "n/a", roas))
    print("\ntop bare-term carriers by spend:")
    for brand, b in sorted(bycar.items(), key=lambda x: -x[1][0])[:15]:
        cost, impr, clicks, conv, value = b
        _, cvr, cpa, roas = _rates(cost, clicks, conv, value)
        print("  %-18s cost=$%-9s conv=%-6s CvR=%.2f%% ROAS=%.2f"
              % (brand, f"{cost:,.0f}", f"{conv:,.0f}", 100 * cvr, roas))
    print("\nwrote %s and %s" % (out, csv_path))


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--db", default=r"C:\Projects\google-ads-ops\db\paid_search.db")
    p.add_argument("--days", type=int, default=90)
    p.add_argument("--account", type=int, default=1, help="account_id (Google Home = 1)")
    p.add_argument("--out", default="brand_term_performance.xlsx")
    p.add_argument("--scan", action="store_true", help="dump the full DB schema and exit")
    a = p.parse_args()
    if a.scan:
        conn = sqlite3.connect(a.db); scan_schema(conn); conn.close(); return
    build(a.db, a.days, a.account, a.out)


if __name__ == "__main__":
    main()
