"""
seam_negatives_final.py — B3 List 1 exact seam negatives + flagship-book prune (16a/16b)

Reads B3 List 1 (104 exact terms) from rollout_plan.xlsx, simulates the
POST-Stage-1 account (fresh export + the generated import files in --output-dir),
and runs the self-block gate: no campaign may receive an exact negative matching
its own still-enabled EXACT keyword.

Emits:
  16a_bh_flagship_book_prune_{date}.csv — pause-in-place rows for enabled EXACT
      keys in BRAND-HOME whose text is on B3 List 1 (the approved book prune;
      expected 12 on current state)
  16b_seam_negatives_exact_{date}.csv   — 104 Campaign Negative Exact rows x
      GENERIC + BRAND-HOME (208 rows), emitted ONLY if the gate passes with
      16a applied

Run AFTER builder_v2.py and rulings_addendum.py (it reads their output files).

Usage:
  python seam_negatives_final.py --export <fresh_export.csv> --plan rollout_plan.xlsx
                                 [--output-dir ./output_v2]
"""

import csv
import glob
import os
import argparse
from collections import defaultdict
from datetime import datetime

GEN = "HOME INSURANCE-GENERIC-HOME-INTENT-NB-SEARCH-ALL-USA-STE"
BH = "HOME INSURANCE_BRAND-HOME-INTENT_COMBINED_TROAS_SEARCH_ALLDEVICE_USA-STE"
L_PAUSE = "restr-2026-09-pause"



# ---- POSTED-NAME TRANSLATION (Editor import matches POSTED names) ----
# Local (pending-rename) name -> posted (live-account) name. Confirmed: TEXAS.
# TIER2RPL / CALIFORNIA presumed from section_e originals — CONFIRM IN WEB UI.
# NB Flagship intentionally unmapped: if it exists only locally (unposted
# creation), its local name IS its match key; if it exists posted under a
# different name, STOP and add it here.
POSTED_NAME_MAP = {
    "HOME INSURANCE-GENERIC-HOME-INTENT-NB-SEARCH-ALL-USA-STE":
        "HOME INSURANCE-TEXAS-NB-SEARCH-ALL-USA-STE",
    "HOME INSURANCE_BRAND-HOME-INTENT_COMBINED_TROAS_SEARCH_ALLDEVICE_USA-STE":
        "HOME INSURANCE_STATE-GROUP_TIER2RPL_COMBINED_TROAS_SEARCH_ALLDEVICE_USA-STE",
    "HOME INSURANCE_FLAGSHIP_CAFL_TROAS_SEARCH_ALLDEVICE_USA-STE":
        "HOME INSURANCE-CALIFORNIA-NB-SEARCH-ALL-USA-STE",
}

REV_NAME_MAP = {v: k for k, v in POSTED_NAME_MAP.items()}

def to_local(name):
    return REV_NAME_MAP.get(name, name)

def to_posted(rows):
    """Translate Campaign / Campaign#Original values to posted names."""
    for r in rows:
        for col in ("Campaign", "Campaign#Original"):
            if col in r and r[col] in POSTED_NAME_MAP:
                r[col] = POSTED_NAME_MAP[r[col]]
    return rows

def read16(path):
    return list(csv.DictReader(open(path, encoding="utf-16-le"), delimiter="\t"))


def latest(outdir, prefix):
    hits = sorted(glob.glob(os.path.join(outdir, prefix + "*.csv")))
    assert hits, f"GATE FAIL: no {prefix}* file in {outdir} — run the generator scripts first"
    return hits[-1]


def load_list1(plan_path):
    from openpyxl import load_workbook
    wb = load_workbook(plan_path, data_only=True, read_only=True)
    sheet = next((s for s in wb.sheetnames if "B3" in s), None)
    assert sheet, f"GATE FAIL: no B3 sheet in {plan_path}"
    terms, in1 = [], False
    for row in wb[sheet].iter_rows(values_only=True):
        v = row[0]
        if v is None:
            continue
        s = str(v).strip()
        if s.startswith("LIST 1"):
            in1 = True
            continue
        if s.startswith("LIST 2"):
            in1 = False
        if in1 and s.startswith("[") and s.endswith("]"):
            terms.append(s[1:-1])
    assert len(terms) == 104, f"GATE FAIL: B3 List 1 has {len(terms)} terms, expected 104"
    return terms


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--export", required=True)
    p.add_argument("--plan", required=True)
    p.add_argument("--output-dir", default="./output_v2")
    a = p.parse_args()
    o = a.output_dir
    ds = datetime.now().strftime("%Y-%m-%d")

    export = list(csv.DictReader(open(a.export, encoding="utf-16"), delimiter="\t"))
    list1 = load_list1(a.plan)
    print(f"B3 List 1 loaded: {len(list1)} exact terms")

    # simulate post-Stage-1: apply generated pauses/renames from the output dir
    kw_pauses, ag_pauses = set(), set()
    for pref in ["01_generic_keywords", "05_bh_keywords", "09_cafl_keywords",
                 "18_generic_keywords_addendum"]:
        for r in read16(latest(o, pref)):
            if r["Keyword status"] == "Paused":
                kw_pauses.add((to_local(r["Campaign"]), r["Ad group"],
                               r["Keyword"].lower(), r["Keyword match type"]))
    for pref in ["03_generic_adgroups", "07_bh_adgroup_status",
                 "10_cafl_adgroup_status", "19_generic_flood_pause"]:
        for r in read16(latest(o, pref)):
            ag_pauses.add((to_local(r["Campaign"]), r["Ad group"]))
    for pref in ["08_bh_renames", "11_cafl_renames"]:
        for r in read16(latest(o, pref)):
            if r.get("Ad group status") == "Paused":
                ag_pauses.add((to_local(r["Campaign"]), r["Ad group#Original"]))

    ag_status = {(r["Campaign"], r["Ad Group"]): r["Ad Group Status"]
                 for r in export if r.get("Ad Group") and r.get("Ad Group Status")}

    def serving_exact_rows(camp):
        out = []
        for r in export:
            k = r.get("Keyword", "").strip()
            if (r["Campaign"] != camp or not k
                    or "Negative" in r.get("Criterion Type", "")
                    or r["Criterion Type"] != "Exact"
                    or r.get("Status") != "Enabled"):
                continue
            c, g = r["Campaign"], r["Ad Group"]
            if ag_status.get((c, g)) != "Enabled" or (c, g) in ag_pauses:
                continue
            if (c, g, k.lower(), "Exact") in kw_pauses:
                continue
            out.append((g, k))
        return out

    l1 = {t.lower() for t in list1}

    # 16a: prune enabled EXACT List-1 keys in BRAND-HOME
    prune = [{"Campaign": BH, "Ad group": g, "Keyword": k,
              "Keyword match type": "Exact", "Keyword status": "Paused",
              "Labels": L_PAUSE}
             for g, k in serving_exact_rows(BH) if k.lower() in l1]
    print(f"\n16a flagship-book prune rows: {len(prune)}")
    for r in prune:
        print(f"  PAUSE [{r['Ad group']}] {r['Keyword']}")
    with open(f"{o}/16a_bh_flagship_book_prune_{ds}.csv", "w",
              encoding="utf-16-le", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["Campaign", "Ad group", "Keyword",
                                          "Keyword match type", "Keyword status",
                                          "Labels"], delimiter="\t")
        w.writeheader()
        w.writerows(to_posted(prune))

    # self-block gate with 16a applied
    for r in prune:
        kw_pauses.add((BH, r["Ad group"], r["Keyword"].lower(), "Exact"))
    viol = []
    for camp, nm in [(GEN, "GENERIC"), (BH, "BRAND-HOME")]:
        ex = {k.lower() for _, k in serving_exact_rows(camp)}
        viol += [(nm, t) for t in list1 if t.lower() in ex]
    print(f"\nSELF-BLOCK GATE (post-Stage-1 + 16a): {len(viol)} violation(s)")
    for v in viol:
        print(f"  VIOLATION: {v}")
    if viol:
        raise SystemExit("GATE FAIL — 16b NOT emitted. Account drifted; report to strategy.")

    with open(f"{o}/16b_seam_negatives_exact_{ds}.csv", "w",
              encoding="utf-16-le", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["Campaign", "Keyword", "Criterion Type",
                                          "Status"], delimiter="\t")
        w.writeheader()
        for camp in [GEN, BH]:
            for t in list1:
                w.writerow(to_posted([{"Campaign": camp, "Keyword": t,
                            "Criterion Type": "Campaign Negative Exact",
                            "Status": "Enabled"}])[0])
    print(f"\n16b written: 208 rows (104 x GENERIC + 104 x BRAND-HOME) — GATE PASS")
    print("Attach 16b CUTOVER MORNING only, coupled with the NB enable.")


if __name__ == "__main__":
    main()
