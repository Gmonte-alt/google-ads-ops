#!/usr/bin/env python
"""
usia_tier_matrix.py -- tier x ad-group demand matrix for the USIA Home launch.

For every ad group (generic AND brand), aggregates the per-state geo-loop search volume
of its SELECTED canonical keywords across each geo tier's state set, and gates each cell
at the impression floor (default 2,000/mo). Shows where an ad group has enough demand to
run in a given tier vs. where it's too thin and should consolidate or sit out.

Tiers are the launch tiers from usia_launch_shell.py; brand uses the same tier mix as
generic (per the launch decision). Demand basis = per-state KWP avg_monthly_searches
summed over the tier's states (search-volume proxy for impressions, same caveat as the
selection floor). National column is the US-level selection volume, shown for reference
only -- it is a different basis (US aggregate vs. sum of per-state) and won't equal the
tier columns.

    python usia_tier_matrix.py --generic-csv usia_home_keywords.csv \
        --brand-csv usia_home_brand.csv --out usia_tier_demand_matrix.xlsx
"""
import argparse
import csv
import os
import sqlite3
from collections import defaultdict, OrderedDict

from usia_keyword_select import clean_keyword

TIERS = OrderedDict([
    ("T1", ["Texas", "Illinois", "Pennsylvania", "Ohio", "New Jersey"]),
    ("T2", ["California", "Georgia", "North Carolina", "Tennessee", "Missouri", "Minnesota",
            "Alabama", "Wisconsin", "Oklahoma", "Connecticut", "Louisiana", "Oregon", "Kentucky",
            "Kansas", "Nevada", "Arkansas", "Utah", "Idaho", "West Virginia", "Delaware"]),
    ("T3", ["Florida", "New York", "Arizona", "Colorado", "Massachusetts", "Maryland",
            "South Carolina", "New Mexico", "Rhode Island"]),
])
ALL_STATES = [s for states in TIERS.values() for s in states]


def read_selected(path, family):
    rows = []
    if path and os.path.exists(path):
        for r in csv.DictReader(open(path, encoding="utf-8")):
            rows.append((family, r["campaign"], r["ad_group"],
                         clean_keyword(r["keyword"]), int(float(r.get("volume") or 0))))
    return rows


def build_lookup(db):
    conn = sqlite3.connect(db)
    ph = ",".join("?" * len(ALL_STATES))
    lookup = {}
    for kw, geo, av in conn.execute(
            "SELECT keyword_text, geo_name, avg_monthly_searches FROM kwp_ideas "
            "WHERE vertical='home' AND geo_name IN (%s) AND avg_monthly_searches > 0" % ph,
            ALL_STATES):
        k = (clean_keyword(kw), geo)
        if av and av > lookup.get(k, 0):     # MAX per (cleaned kw, state): shared buckets aren't additive
            lookup[k] = av
    conn.close()
    return lookup


def build(generic_csv, brand_csv, db, out, floor):
    sel = read_selected(generic_csv, "Generic") + read_selected(brand_csv, "Brand")
    lookup = build_lookup(db)

    agg = OrderedDict()      # (family, campaign, ad_group) -> {'kws': [], 'natl': int}
    for family, camp, ag, kw, vol in sel:
        a = agg.setdefault((family, camp, ag), {"kws": [], "natl": 0})
        a["kws"].append(kw)
        a["natl"] += vol

    matrix = []
    for (family, camp, ag), a in agg.items():
        tv = {}
        for t, states in TIERS.items():
            tv[t] = sum(lookup.get((kw, st), 0) for kw in a["kws"] for st in states)
        matrix.append({"family": family, "campaign": camp, "ad_group": ag,
                       "nkw": len(a["kws"]), "natl": a["natl"], **tv,
                       "all": tv["T1"] + tv["T2"] + tv["T3"]})

    fam_order = {"Generic": 0, "Brand": 1}
    matrix.sort(key=lambda r: (fam_order.get(r["family"], 9), r["campaign"], -r["T1"], -r["all"]))
    write_matrix(out, matrix, floor)
    return matrix


def write_matrix(out, matrix, floor):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    HEAD = PatternFill("solid", fgColor="1F3864")
    OKF = PatternFill("solid", fgColor="E2EFDA")
    THINF = PatternFill("solid", fgColor="FCE4D6")
    FAMF = PatternFill("solid", fgColor="D6DCE4")
    WHITE = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    BOLD = Font(name="Arial", bold=True, size=10)
    REG = Font(name="Arial", size=10)
    NOTE = Font(name="Arial", size=9, italic=True, color="595959")
    TITLE = Font(name="Arial", bold=True, size=14, color="1F3864")
    THIN = Side(style="thin", color="BFBFBF")
    BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    wb = Workbook(); ws = wb.active; ws.title = "Tier x Ad Group Demand"
    ws["A1"] = "USIA Home -- Tier x Ad Group Demand Matrix"; ws["A1"].font = TITLE
    ws.merge_cells("A1:I1")
    ws["A2"] = ("Per-state KWP search volume of each ad group's selected keywords, summed across the tier's "
                "states. Cells gated at %s/mo: green = clears, orange = THIN. T1=5, T2=20, T3=9 states. "
                "National = US-level selection volume (reference only; different basis). Volume is an "
                "impression proxy." % f"{floor:,}")
    ws["A2"].font = NOTE; ws.merge_cells("A2:I2")

    heads = ["Family", "Campaign", "Ad Group", "#KW", "National (ref)",
             "T1 (5)", "T2 (20)", "T3 (9)", "All Tiers"]
    hr = 4
    for j, h in enumerate(heads, 1):
        c = ws.cell(hr, j, h); c.fill = HEAD; c.font = WHITE
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = BORDER

    r = hr + 1
    cur_fam = None
    tier_tot = defaultdict(int)
    for row in matrix:
        if row["family"] != cur_fam:
            cur_fam = row["family"]
            cc = ws.cell(r, 1, "%s campaigns" % cur_fam); cc.font = BOLD; cc.fill = FAMF
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
            r += 1
        ws.cell(r, 1, row["family"]).font = REG
        ws.cell(r, 2, row["campaign"]).font = REG
        ws.cell(r, 3, row["ad_group"]).font = BOLD
        ws.cell(r, 4, row["nkw"]).font = REG
        ws.cell(r, 5, row["natl"]).font = REG; ws.cell(r, 5).number_format = "#,##0"
        for j, t in zip((6, 7, 8), ("T1", "T2", "T3")):
            c = ws.cell(r, j, row[t]); c.font = REG; c.number_format = "#,##0"
            c.fill = OKF if row[t] >= floor else THINF
            tier_tot[t] += row[t]
        ws.cell(r, 9, row["all"]).font = BOLD; ws.cell(r, 9).number_format = "#,##0"
        r += 1

    ws.cell(r, 3, "TIER TOTAL (all ad groups)").font = BOLD
    for j, t in zip((6, 7, 8), ("T1", "T2", "T3")):
        c = ws.cell(r, j, tier_tot[t]); c.font = BOLD; c.number_format = "#,##0"
    ws.cell(r, 9, sum(tier_tot.values())).font = BOLD; ws.cell(r, 9).number_format = "#,##0"

    for col, w in {"A": 10, "B": 26, "C": 20, "D": 7, "E": 14, "F": 11, "G": 11, "H": 11, "I": 12}.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    wb.save(out)
    print("wrote %s" % out)


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--db", default=r"C:\Projects\google-ads-ops\db\demand_research.db")
    p.add_argument("--generic-csv", default="usia_home_keywords.csv")
    p.add_argument("--brand-csv", default="usia_home_brand.csv")
    p.add_argument("--out", default="usia_tier_demand_matrix.xlsx")
    p.add_argument("--floor", type=int, default=2000)
    a = p.parse_args()
    matrix = build(a.generic_csv, a.brand_csv, a.db, a.out, a.floor)
    print("\ntier x ad group (THIN cells < %s):" % f"{a.floor:,}")
    cur = None
    for row in matrix:
        if row["family"] != cur:
            print("  [%s]" % row["family"]); cur = row["family"]
        flags = " ".join("%s%s" % (t, "" if row[t] >= a.floor else "*") for t in ("T1", "T2", "T3"))
        print("    %-18s T1=%-7s T2=%-7s T3=%-7s  (%s)"
              % (row["ad_group"], f"{row['T1']:,}", f"{row['T2']:,}", f"{row['T3']:,}",
                 " ".join(t for t in ("T1", "T2", "T3") if row[t] < a.floor) or "all clear"))


if __name__ == "__main__":
    main()
