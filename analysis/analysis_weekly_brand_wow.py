"""
analysis_weekly_brand_wow.py  [v2 — multi-account]
-----------------------------
Weekly WoW brand-level performance analysis.

Pass --account to select which account to analyze.
Pass --week and --prior to override the auto-detected weeks.

Metrics computed per brand (modifier_brand from search_queries):
    Revenue ($)         -- offline revenue (primary_conv_value)
    Cost ($)            -- ad spend
    GP ($)              -- Revenue - Cost
    GP %                -- GP / Revenue
    Rev/Click ($)       -- Revenue / Clicks
    ROAS                -- Revenue / Cost
    Impressions
    Clicks
    CTR %               -- Clicks / Impressions
    CPC ($)             -- Cost / Clicks
    CPL ($)             -- Cost / Leads (offline conversions)
    RPL ($)             -- Revenue / Leads
    Leads               -- offline conversions (primary_conv_count)
    Lead CVR %          -- Leads / Clicks

Two sheets:
    WoW by Brand (SQ)  -- search query grain, grouped by brand
    WoW by Keyword     -- keyword_performance_daily grain, grouped by keyword

Usage:
    conda activate paid-search
    python analysis/analysis_weekly_brand_wow.py --account home
    python analysis/analysis_weekly_brand_wow.py --account medicare
    python analysis/analysis_weekly_brand_wow.py --account medicare --week 2026-W14
    python analysis/analysis_weekly_brand_wow.py --account medicare --week 2026-W14 --prior 2026-W13

Output:
    analysis/output/weekly_brand_wow_{account}.xlsx
"""

import sqlite3
import argparse
import os
import numpy as np
import pandas as pd
from pathlib import Path
from datetime import date
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Account map
# ---------------------------------------------------------------------------

ACCOUNT_MAP = {
    "home": {
        "display_name": "Home Insurance",
        "google_id":    "8886502147",
    },
    "medicare": {
        "display_name": "Medicare",
        "google_id":    "5310820901",
    },
    "auto": {
        "display_name": "Auto Insurance (autoinsurance.org)",
        "google_id":    "3863660890",
    },
    "term_life": {
        "display_name": "Term Life",
        "google_id":    "1671155176",
    },
    "auto_quote": {
        "display_name": "Auto Insurance (auto.quote.com)",
        "google_id":    "1006729211",
    },
}

# ---------------------------------------------------------------------------
# Config
# ---------------------------------------------------------------------------

DB_PATH    = Path(r"C:\Projects\google-ads-ops\db\paid_search.db")
OUTPUT_DIR = Path(r"C:\Projects\google-ads-ops\analysis\output")
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

LOWER_IS_BETTER = {'Cost ($)', 'CPC ($)', 'CPL ($)'}
HIGHER_IS_BETTER = {
    'Revenue ($)', 'GP ($)', 'GP %', 'Rev/Click ($)', 'ROAS',
    'Impressions', 'Clicks', 'CTR %', 'RPL ($)', 'Leads', 'Lead CVR %',
}

METRIC_COLS = [
    ('revenue',    'Revenue ($)'),
    ('cost',       'Cost ($)'),
    ('gp',         'GP ($)'),
    ('gp_pct',     'GP %'),
    ('rev_click',  'Rev/Click ($)'),
    ('roas',       'ROAS'),
    ('impressions','Impressions'),
    ('clicks',     'Clicks'),
    ('ctr',        'CTR %'),
    ('cpc',        'CPC ($)'),
    ('cpl',        'CPL ($)'),
    ('rpl',        'RPL ($)'),
    ('leads',      'Leads'),
    ('lead_cvr',   'Lead CVR %'),
]


# ---------------------------------------------------------------------------
# Argument parsing
# ---------------------------------------------------------------------------

def parse_args():
    parser = argparse.ArgumentParser(
        description='Weekly WoW brand performance analysis.'
    )
    parser.add_argument(
        '--account', required=True, choices=list(ACCOUNT_MAP.keys()),
        help='Account to analyze.',
    )
    parser.add_argument('--week',  help='Current week label e.g. 2026-W14')
    parser.add_argument('--prior', help='Prior week label e.g. 2026-W13')
    return parser.parse_args()


# ---------------------------------------------------------------------------
# Database
# ---------------------------------------------------------------------------

def get_conn():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def get_most_recent_complete_weeks(
    conn, google_id: str, n: int = 2
) -> list[str]:
    """
    Returns the N most recent complete ISO weeks that have data for this
    account. Scoped to the specific account so Medicare's data window
    doesn't bleed into Home's week detection or vice versa.
    """
    rows = conn.execute("""
        SELECT DISTINCT sqp.report_date
        FROM search_query_performance sqp
        JOIN keyword_instances ki ON sqp.instance_id = ki.id
        JOIN accounts a           ON ki.account_id   = a.id
        WHERE a.platform_account_id = ?
          AND a.platform = 'google'
          AND sqp.device = 'ALL'
        ORDER BY sqp.report_date DESC
        LIMIT 200
    """, (google_id,)).fetchall()

    today        = date.today()
    current_iso  = today.isocalendar()
    current_year = current_iso.year
    current_week = current_iso.week

    seen  = set()
    weeks = []
    for r in rows:
        d     = date.fromisoformat(r['report_date'])
        iso   = d.isocalendar()
        label = f"{iso.year}-W{str(iso.week).zfill(2)}"
        if label in seen:
            continue
        seen.add(label)
        # Skip the current partial week
        if iso.year == current_year and iso.week == current_week:
            continue
        weeks.append(label)
        if len(weeks) == n:
            break

    return weeks


def fetch_brand_week_performance(
    conn, google_id: str, week_label: str
) -> pd.DataFrame:
    """
    Aggregates search query performance for one ISO week, grouped by
    modifier_brand + brand_category + vertical for the given account.
    """
    year, wnum = week_label.split('-W')
    week_start = date.fromisocalendar(int(year), int(wnum), 1)
    week_end   = date.fromisocalendar(int(year), int(wnum), 7)

    return pd.read_sql_query("""
        SELECT
            COALESCE(sq.modifier_brand, '(non-brand)') AS brand,
            sq.brand_category,
            sq.vertical,
            SUM(sqp.impressions)                        AS impressions,
            SUM(sqp.clicks)                             AS clicks,
            SUM(sqp.cost_micros) / 1e6                  AS cost,
            SUM(sqp.primary_conv_count)                 AS leads,
            SUM(sqp.primary_conv_value)                 AS revenue
        FROM search_query_performance sqp
        JOIN search_queries sq    ON sqp.query_id   = sq.id
        JOIN keyword_instances ki ON sqp.instance_id = ki.id
        JOIN accounts a           ON ki.account_id   = a.id
        WHERE a.platform_account_id = ?
          AND a.platform = 'google'
          AND sqp.device = 'ALL'
          AND sqp.report_date BETWEEN ? AND ?
        GROUP BY sq.modifier_brand, sq.brand_category, sq.vertical
        ORDER BY revenue DESC
    """, conn, params=(google_id, str(week_start), str(week_end)))


def fetch_keyword_week_performance(
    conn, google_id: str, week_label: str
) -> pd.DataFrame:
    """
    Aggregates keyword_performance_daily for one ISO week, grouped by
    keyword identity + brand. Uses keyword-level performance data rather
    than search query data, giving a cleaner CPL/CPC view that reflects
    actual bid-level spend.
    """
    year, wnum = week_label.split('-W')
    week_start = date.fromisocalendar(int(year), int(wnum), 1)
    week_end   = date.fromisocalendar(int(year), int(wnum), 7)

    return pd.read_sql_query("""
        SELECT
            ki_id.normalized_text                         AS keyword,
            ki_id.match_type,
            COALESCE(ki_id.modifier_brand, '(non-brand)') AS brand,
            ki_id.brand_category,
            ki_id.vertical,
            SUM(kpd.impressions)                          AS impressions,
            SUM(kpd.clicks)                               AS clicks,
            SUM(kpd.cost_micros) / 1e6                    AS cost,
            SUM(kpd.primary_conv_count)                   AS leads,
            SUM(kpd.primary_conv_value)                   AS revenue
        FROM keyword_performance_daily kpd
        JOIN keyword_instances ki     ON kpd.instance_id  = ki.id
        JOIN keyword_identities ki_id ON ki.identity_id   = ki_id.id
        JOIN accounts a               ON ki.account_id    = a.id
        WHERE a.platform_account_id = ?
          AND a.platform = 'google'
          AND kpd.device = 'ALL'
          AND kpd.report_date BETWEEN ? AND ?
        GROUP BY ki_id.normalized_text, ki_id.match_type,
                 ki_id.modifier_brand, ki_id.brand_category, ki_id.vertical
        ORDER BY revenue DESC
    """, conn, params=(google_id, str(week_start), str(week_end)))


# ---------------------------------------------------------------------------
# Metric computation
# ---------------------------------------------------------------------------

def compute_metrics(df: pd.DataFrame) -> pd.DataFrame:
    """Adds all derived metric columns to a raw performance DataFrame."""
    d = df.copy()
    d['gp']       = d['revenue'] - d['cost']
    d['gp_pct']   = np.where(d['revenue'] > 0,
                              d['gp'] / d['revenue'] * 100, np.nan)
    d['rev_click'] = np.where(d['clicks'] > 0,
                               d['revenue'] / d['clicks'], np.nan)
    d['roas']     = np.where(d['cost'] > 0,
                              d['revenue'] / d['cost'], np.nan)
    d['ctr']      = np.where(d['impressions'] > 0,
                              d['clicks'] / d['impressions'] * 100, np.nan)
    d['cpc']      = np.where(d['clicks'] > 0,
                              d['cost'] / d['clicks'], np.nan)
    d['cpl']      = np.where(d['leads'] > 0,
                              d['cost'] / d['leads'], np.nan)
    d['rpl']      = np.where(d['leads'] > 0,
                              d['revenue'] / d['leads'], np.nan)
    d['lead_cvr'] = np.where(d['clicks'] > 0,
                              d['leads'] / d['clicks'] * 100, np.nan)
    return d


# ---------------------------------------------------------------------------
# WoW table builders
# ---------------------------------------------------------------------------

def _wow_rows(merged: pd.DataFrame, key_cols: list,
              curr_week: str, prior_week: str) -> pd.DataFrame:
    """
    Shared WoW row builder. Takes a merged current/prior DataFrame and
    produces the wide WoW output with current, prior, delta, and delta%
    columns for every metric.
    """
    rows = []
    for _, r in merged.iterrows():
        row = {col: r[col] for col in key_cols if col in r}
        for raw_col, display_name in METRIC_COLS:
            cw  = r.get(f'{raw_col}_cw', 0) or 0
            pw  = r.get(f'{raw_col}_pw', 0) or 0
            delta     = cw - pw
            delta_pct = (delta / abs(pw) * 100) if pw != 0 else np.nan
            row[f'{display_name} ({curr_week})']  = round(cw, 2)
            row[f'{display_name} ({prior_week})'] = round(pw, 2)
            row[f'{display_name} Δ']              = round(delta, 2)
            row[f'{display_name} Δ%'] = (
                round(delta_pct, 1) if not np.isnan(delta_pct) else None
            )
        rows.append(row)
    return pd.DataFrame(rows)


def _summary_row(curr_df: pd.DataFrame, prior_df: pd.DataFrame,
                 key_cols: list, id_values: dict,
                 curr_week: str, prior_week: str) -> pd.DataFrame:
    """Builds a single TOTAL row, re-deriving ratio metrics from sums."""
    def totals(df):
        t = compute_metrics(df.sum(numeric_only=True).to_frame().T)
        # Re-derive ratio metrics from column totals
        t['gp_pct']    = (t['gp'] / t['revenue'] * 100).where(t['revenue'] > 0)
        t['rev_click']  = (t['revenue'] / t['clicks']).where(t['clicks'] > 0)
        t['roas']       = (t['revenue'] / t['cost']).where(t['cost'] > 0)
        t['ctr']        = (t['clicks'] / t['impressions'] * 100).where(t['impressions'] > 0)
        t['cpc']        = (t['cost'] / t['clicks']).where(t['clicks'] > 0)
        t['cpl']        = (t['cost'] / t['leads']).where(t['leads'] > 0)
        t['rpl']        = (t['revenue'] / t['leads']).where(t['leads'] > 0)
        t['lead_cvr']   = (t['leads'] / t['clicks'] * 100).where(t['clicks'] > 0)
        return t

    curr_tot  = totals(curr_df)
    prior_tot = totals(prior_df)

    row = {col: id_values.get(col, '') for col in key_cols}
    for raw_col, display_name in METRIC_COLS:
        cw  = float(curr_tot[raw_col].iloc[0]) if raw_col in curr_tot.columns else 0
        pw  = float(prior_tot[raw_col].iloc[0]) if raw_col in prior_tot.columns else 0
        cw  = 0 if np.isnan(cw)  else cw
        pw  = 0 if np.isnan(pw)  else pw
        delta     = cw - pw
        delta_pct = (delta / abs(pw) * 100) if pw != 0 else np.nan
        row[f'{display_name} ({curr_week})']  = round(cw, 2)
        row[f'{display_name} ({prior_week})'] = round(pw, 2)
        row[f'{display_name} Δ']              = round(delta, 2)
        row[f'{display_name} Δ%'] = (
            round(delta_pct, 1) if not np.isnan(delta_pct) else None
        )
    return pd.DataFrame([row])


def build_brand_wow(curr_df: pd.DataFrame, prior_df: pd.DataFrame,
                    curr_week: str, prior_week: str) -> pd.DataFrame:
    """
    Brand-level WoW table from search query performance.
    Joined on brand + brand_category + vertical.
    TOTAL row prepended.
    """
    curr  = compute_metrics(curr_df)
    prior = compute_metrics(prior_df)
    key   = ['brand', 'brand_category', 'vertical']

    merged = curr.merge(prior, on=key, how='outer',
                        suffixes=('_cw', '_pw')).fillna(0)

    detail = _wow_rows(merged, key, curr_week, prior_week)
    cw_rev = f'Revenue ($) ({curr_week})'
    if cw_rev in detail.columns:
        detail = detail.sort_values(cw_rev, ascending=False)

    summary = _summary_row(
        curr_df, prior_df, key,
        {'brand': 'TOTAL', 'brand_category': '', 'vertical': ''},
        curr_week, prior_week,
    )

    result = pd.concat([summary, detail], ignore_index=True)
    return result.rename(columns={
        'brand': 'Brand', 'brand_category': 'Brand Category', 'vertical': 'Vertical',
    })


def build_keyword_wow(curr_df: pd.DataFrame, prior_df: pd.DataFrame,
                      curr_week: str, prior_week: str) -> pd.DataFrame:
    """
    Keyword-level WoW table from keyword_performance_daily.
    Joined on keyword + match_type + brand + brand_category + vertical.
    TOTAL row prepended.
    """
    curr  = compute_metrics(curr_df)
    prior = compute_metrics(prior_df)
    key   = ['keyword', 'match_type', 'brand', 'brand_category', 'vertical']

    merged = curr.merge(prior, on=key, how='outer',
                        suffixes=('_cw', '_pw')).fillna(0)

    detail = _wow_rows(merged, key, curr_week, prior_week)
    cw_rev = f'Revenue ($) ({curr_week})'
    if cw_rev in detail.columns:
        detail = detail.sort_values(cw_rev, ascending=False)

    summary = _summary_row(
        curr_df, prior_df, key,
        {'keyword': 'TOTAL', 'match_type': '', 'brand': '',
         'brand_category': '', 'vertical': ''},
        curr_week, prior_week,
    )

    result = pd.concat([summary, detail], ignore_index=True)
    return result.rename(columns={
        'keyword': 'Keyword', 'match_type': 'Match Type',
        'brand': 'Brand', 'brand_category': 'Brand Category',
        'vertical': 'Vertical',
    })


# ---------------------------------------------------------------------------
# Excel formatting
# ---------------------------------------------------------------------------

def format_wow_workbook(path: Path, curr_week: str, prior_week: str):
    """
    Applies WoW-specific formatting:
    - Dark blue header, white bold font
    - TOTAL row in medium blue
    - Alternating row shading from row 3
    - Δ% columns: green text for improvement, red for deterioration
      (direction-aware: lower-is-better metrics invert the colour logic)
    - Number formats per column header content
    - Auto-fit column widths (capped at 28)
    - Freeze panes at D2 (locks brand/category/vertical + header)
    """
    wb = load_workbook(path)

    GREEN      = 'FF006400'
    RED        = 'FFCC0000'
    HDR_FILL   = PatternFill('solid', fgColor='1F4E79')
    TOTAL_FILL = PatternFill('solid', fgColor='2E75B6')
    ALT_FILL   = PatternFill('solid', fgColor='EBF3FB')
    WHITE_FONT = Font(color='FFFFFF', bold=True, size=10)

    for sheet_name in wb.sheetnames:
        ws      = wb[sheet_name]
        max_row = ws.max_row
        max_col = ws.max_column

        # Header row
        for cell in ws[1]:
            cell.fill      = HDR_FILL
            cell.font      = WHITE_FONT
            cell.alignment = Alignment(
                horizontal='center', vertical='center', wrap_text=True)
        ws.row_dimensions[1].height = 48

        # TOTAL row (row 2)
        for cell in ws[2]:
            cell.fill = TOTAL_FILL
            cell.font = Font(color='FFFFFF', bold=True, size=10)
            cell.alignment = Alignment(horizontal='center', vertical='center')

        # Alternating shading from row 3
        for row_idx in range(3, max_row + 1):
            if row_idx % 2 == 0:
                for cell in ws[row_idx]:
                    rgb = getattr(cell.fill.fgColor, 'rgb', None)
                    if not rgb or rgb in ('00000000', 'FFFFFFFF'):
                        cell.fill = ALT_FILL

        # Colour Δ% columns by improvement direction
        for col_idx in range(1, max_col + 1):
            hdr = str(ws.cell(1, col_idx).value or '')
            if not hdr.endswith('Δ%'):
                continue
            metric_name   = hdr.replace(' Δ%', '').strip()
            lower_is_better = metric_name in LOWER_IS_BETTER

            for row_idx in range(3, max_row + 1):
                cell = ws.cell(row_idx, col_idx)
                try:
                    val = float(cell.value)
                except (TypeError, ValueError):
                    continue
                if lower_is_better:
                    is_good = val < -2
                    is_bad  = val >  2
                else:
                    is_good = val >  2
                    is_bad  = val < -2

                if is_good:
                    cell.font = Font(color=GREEN, bold=True, size=10)
                elif is_bad:
                    cell.font = Font(color=RED,   bold=True, size=10)
                else:
                    cell.font = Font(color='FF666666', size=10)

        # Number formatting
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if not isinstance(cell.value, (int, float)):
                    continue
                hdr = str(ws.cell(1, cell.column).value or '').lower()
                if 'δ%' in hdr or 'gp %' in hdr or 'ctr' in hdr or 'cvr' in hdr:
                    cell.number_format = '0.0'
                elif ('$' in hdr or 'revenue' in hdr or 'cost' in hdr
                      or 'gp' in hdr or 'cpl' in hdr or 'rpl' in hdr
                      or 'cpc' in hdr or 'rev/click' in hdr):
                    cell.number_format = '$#,##0.00'
                elif 'roas' in hdr:
                    cell.number_format = '0.00'
                elif 'δ' in hdr:
                    cell.number_format = '+#,##0.00;-#,##0.00;0.00'
                elif ('impression' in hdr or 'click' in hdr or 'lead' in hdr):
                    cell.number_format = '#,##0'

        # Auto-fit columns
        for col in ws.columns:
            max_len    = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                try:
                    val_len = len(str(cell.value)) if cell.value else 0
                    max_len = max(max_len, val_len)
                except Exception:
                    pass
            ws.column_dimensions[col_letter].width = min(max_len + 3, 28)

        ws.freeze_panes = 'D2'

    wb.save(path)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    args       = parse_args()
    cfg        = ACCOUNT_MAP[args.account]
    google_id  = cfg['google_id']
    output_file = OUTPUT_DIR / f"weekly_brand_wow_{args.account}.xlsx"

    conn = get_conn()

    # Resolve weeks
    if args.week and args.prior:
        curr_week  = args.week
        prior_week = args.prior
    elif args.week:
        # Only current week supplied -- auto-detect prior
        recent = get_most_recent_complete_weeks(conn, google_id, n=3)
        if args.week in recent:
            idx = recent.index(args.week)
            prior_week = recent[idx + 1] if idx + 1 < len(recent) else None
        else:
            prior_week = recent[0] if recent else None
        curr_week = args.week
        if not prior_week:
            print("ERROR: Could not auto-detect prior week. Supply --prior explicitly.")
            conn.close()
            return
    else:
        recent = get_most_recent_complete_weeks(conn, google_id, n=2)
        if len(recent) < 2:
            print("ERROR: Not enough complete weeks in database for this account.")
            conn.close()
            return
        curr_week  = recent[0]
        prior_week = recent[1]

    print('=' * 65)
    print(f'Weekly WoW Brand Analysis — {cfg["display_name"]} (Google)')
    print(f'  Account      : {google_id}')
    print(f'  Current week : {curr_week}')
    print(f'  Prior week   : {prior_week}')
    print(f'  Output       : {output_file.name}')
    print('=' * 65)

    # Fetch data
    print('\nFetching current week (search query level)...')
    curr_sq  = fetch_brand_week_performance(conn, google_id, curr_week)
    print(f'  {len(curr_sq):,} brand rows  |  ${curr_sq["revenue"].sum():,.2f} revenue')

    print('Fetching prior week (search query level)...')
    prior_sq = fetch_brand_week_performance(conn, google_id, prior_week)
    print(f'  {len(prior_sq):,} brand rows  |  ${prior_sq["revenue"].sum():,.2f} revenue')

    print('Fetching current week (keyword level)...')
    curr_kw  = fetch_keyword_week_performance(conn, google_id, curr_week)
    print(f'  {len(curr_kw):,} keyword rows  |  ${curr_kw["revenue"].sum():,.2f} revenue')

    print('Fetching prior week (keyword level)...')
    prior_kw = fetch_keyword_week_performance(conn, google_id, prior_week)
    print(f'  {len(prior_kw):,} keyword rows  |  ${prior_kw["revenue"].sum():,.2f} revenue')

    conn.close()

    # Build tables
    print('\nBuilding brand WoW table...')
    brand_wow_df = build_brand_wow(curr_sq, prior_sq, curr_week, prior_week)

    print('Building keyword WoW table...')
    kw_wow_df = build_keyword_wow(curr_kw, prior_kw, curr_week, prior_week)

    # Terminal preview -- TOTAL row
    print('\n--- TOTAL WoW summary ---')
    tot = brand_wow_df.iloc[0]
    preview_metrics = [
        ('Revenue ($)',  'Revenue'),
        ('Cost ($)',     'Cost'),
        ('GP ($)',       'GP'),
        ('GP %',         'GP%'),
        ('ROAS',         'ROAS'),
        ('CPC ($)',      'CPC'),
        ('CPL ($)',      'CPL'),
        ('Leads',        'Leads'),
    ]
    for metric, label in preview_metrics:
        cw  = tot.get(f'{metric} ({curr_week})',  0)
        pw  = tot.get(f'{metric} ({prior_week})', 0)
        pct = tot.get(f'{metric} Δ%', None)
        pct_str = f'{pct:+.1f}%' if pct is not None else 'N/A'
        print(f'  {label:<12} CW={cw:>10.2f}  PW={pw:>10.2f}  WoW={pct_str}')

    # Top 5 GP decliners
    detail = brand_wow_df.iloc[1:]   # exclude TOTAL row
    gp_col = 'GP ($) Δ'
    if gp_col in detail.columns and not detail.empty:
        movers = detail.nsmallest(5, gp_col)
        print(f'\n--- Top 5 GP decliners (WoW) ---')
        for _, r in movers.iterrows():
            brand   = r.get('Brand', '')
            gp_chg  = r.get(gp_col, 0)
            rev_chg = r.get('Revenue ($) Δ', 0)
            gp_pct  = r.get('GP % Δ%', None)
            pct_str = f'{gp_pct:+.1f}%' if gp_pct is not None else ''
            print(f'  {str(brand):<30}  GP Δ={gp_chg:>+10.2f}  '
                  f'Rev Δ={rev_chg:>+10.2f}  GP% WoW={pct_str}')

    # Write Excel
    print(f'\nWriting Excel output...')
    with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
        brand_wow_df.to_excel(writer, sheet_name='WoW by Brand (SQ)', index=False)
        kw_wow_df.to_excel(writer,    sheet_name='WoW by Keyword',     index=False)

    print('Applying formatting...')
    format_wow_workbook(output_file, curr_week, prior_week)

    size_kb = os.path.getsize(output_file) / 1024
    print(f'\nOutput : {output_file}')
    print(f'Size   : {size_kb:.1f} KB')
    print('\nDone.')


if __name__ == '__main__':
    main()
